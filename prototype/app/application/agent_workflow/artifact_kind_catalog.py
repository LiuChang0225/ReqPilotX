"""Small artifact dictionary for task navigation, not a transport or stage allowlist.

Compilers retain their complete bundles. This module validates a saved operation
using the existing domain contracts; it does not judge source semantics.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from app.application.agent_workflow import source_requirement_baseline as core
from app.domain.extraction.source_ir import SourceIRSnapshot

ARTIFACT_KINDS = {
    'VISUAL_SOURCE': {'filename':'source_ir.json','dependencies':['visual-source-capture.json'], 'validator':'SourceIR visual projection + exact predecessor', 'display_role':'reading_evidence'},
    'READING_WORKSPACE': {'filename': 'reading-workspace.json', 'dependencies': ['source_ir.json', 'source.pdf', 'pages/index.json', 'tables/index.json', 'assets/index.json', 'assets/placements.json', 'contact-sheet.png'], 'validator': 'reading workspace source binding', 'display_role': 'page_markdown_and_visual_evidence'},
    'REVIEW_HEAD': {'filename': 'workspace-manifest.json', 'dependencies': [], 'validator': 'existing review head + immutable current bundle', 'display_role': 'working_review'},
    'SOURCE_IR': {'filename': 'source_ir.json', 'dependencies': ['manifest.json'], 'validator': 'SourceIRSnapshot', 'display_role': 'reading_evidence'},
    'DOCUMENT_MAP': {'filename': 'document-map.json', 'dependencies': ['table-previews.json'], 'validator': 'DocumentMapArtifactV1', 'display_role': 'structure_and_original'},
    'SOURCE_ITEMS': {'filename': 'source-requirement-baseline.json', 'dependencies': ['document-map.json', 'table-previews.json', 'validation-receipt.json'], 'validator': 'SourceRequirementBaselineArtifactV1 + receipt replay', 'display_role': 'source_ledger'},
    'SOURCE_FINDINGS': {'filename': 'source-findings.json', 'dependencies': ['document-map.json'], 'validator': 'source-findings/v1 lineage', 'display_role': 'source_review_notices'},
    'ANNOTATIONS': {'filename': 'image-annotations.json', 'dependencies': ['source-requirement-baseline.json'], 'validator': 'build_delivery_snapshot', 'display_role': 'original_number_topics_and_visual_reading'},
    'DECOMPOSITION': {'filename': 'source-item-decomposition.json', 'dependencies': ['source-requirement-baseline.json'], 'validator': 'build_delivery_snapshot', 'display_role': 'derived_requirements'},
    'DELIVERY': {'filename': 'matrix-snapshot.json', 'dependencies': ['matrix.xlsx'], 'validator': 'snapshot + workbook metadata', 'display_role': 'same_version_delivery'},
}
OPERATION_KINDS = {
    'read-pdf-visual': ('VISUAL_SOURCE',), 'revise-visual-text': ('VISUAL_SOURCE',),
    'reading-workspace': ('READING_WORKSPACE',), 'review-open': ('REVIEW_HEAD',), 'capture': ('SOURCE_IR',), 'compile-map': ('DOCUMENT_MAP', 'SOURCE_FINDINGS'),
    'compile-items': ('SOURCE_ITEMS',), 'compile-images': ('SOURCE_ITEMS', 'ANNOTATIONS'),
    'compile-annotations': ('ANNOTATIONS',), 'compile-decomposition': ('DECOMPOSITION',),
    'deliver': ('DELIVERY',), 'review-revise': ('SOURCE_ITEMS',),
}


OPERATION_REQUIRED_FILES = {
    'compile-map': ('native-structure-reference.json', 'scope-requests.json', 'map-compilation-receipt.json', 'map-audit.json', 'map-review.json'),
    'compile-items': ('item-compilation-receipt.json', 'scope-requests.json'),
    'compile-annotations': ('annotation-compilation-receipt.json',),
    'compile-decomposition': ('decomposition-compilation-receipt.json',),
    'deliver': ('delivery-receipt.json',),
}


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def file_sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def inside(root, relative):
    root = Path(root).resolve()
    value = Path(relative)
    if value.is_absolute() or '..' in value.parts:
        raise ValueError('Artifact reference must be relative and remain inside its bundle')
    path = (root / value).resolve()
    if path == root or root not in path.parents:
        raise ValueError('Artifact reference escapes its bundle')
    return path


def validate_bundle(directory, source_sha256, *, operation, artifact_kinds=None, dependency_dirs=()):
    """Validate owned output and exact explicit predecessors without moving files."""
    directory = Path(directory).resolve(strict=True)
    kinds = tuple(artifact_kinds or OPERATION_KINDS.get(operation, ()))
    if not kinds or set(kinds) - ARTIFACT_KINDS.keys():
        raise ValueError('Unsupported workspace operation/artifact kinds: ' + operation)
    if set(OPERATION_KINDS.get(operation, ())) - set(kinds):
        raise ValueError('Operation artifact selection cannot omit its required outputs')
    roots = [directory, *[Path(p).resolve(strict=True) for p in dependency_dirs]]

    def locate(name, *, own=False):
        for root in roots[:1] if own else roots:
            path = inside(root, name)
            if path.is_file():
                return path
        raise ValueError('Missing artifact dependency: ' + name)

    for name in OPERATION_REQUIRED_FILES.get(operation, ()):
        locate(name, own=True)
    for kind in kinds:
        locate(ARTIFACT_KINDS[kind]['filename'], own=True)
        for dependency in ARTIFACT_KINDS[kind]['dependencies']:
            locate(dependency)
    stages, artifacts = {}, []
    if 'REVIEW_HEAD' in kinds:
        head = read_json(directory / 'workspace-manifest.json')
        digest = core.canonical_json_sha256({k: v for k, v in head.items() if k != 'review_sha256'})
        if head.get('schema_version') != 'reqpilot-review-workspace/v1' or head.get('review_sha256') != digest or head.get('source', {}).get('sha256') != source_sha256:
            raise ValueError('Review head source or integrity differs')
        if read_json(inside(directory, 'reviews/' + digest + '.json')) != head or file_sha256(inside(directory, head['source']['path'])) != source_sha256:
            raise ValueError('Review immutable head or original missing')
        current = inside(directory, head['current']['version_path'])
        for name, digest in head['current']['files'].items():
            if file_sha256(inside(current, name)) != digest:
                raise ValueError('Review current files changed')
        checked = validate_bundle(current, source_sha256, operation='review-revise')
        if checked['stages']['MAP']['artifact_sha256'] != head['current']['document_map_sha256'] or checked['stages']['ITEMS']['artifact_sha256'] != head['current']['baseline_sha256']:
            raise ValueError('Review current pointer differs')
        stages.update(checked['stages'])
    if 'READING_WORKSPACE' in kinds:
        manifest = read_json(directory / 'reading-workspace.json')
        snapshot = SourceIRSnapshot.model_validate(read_json(directory / 'source_ir.json'))
        if manifest.get('schema_version') != 'reqpilot-reading-workspace/v1' or manifest.get('source_sha256') != source_sha256 or snapshot.source_sha256 != source_sha256 or file_sha256(directory / 'source.pdf') != source_sha256 or file_sha256(directory / 'source_ir.json') != manifest.get('source_ir_sha256'):
            raise ValueError('Reading workspace source or capture changed')
        address_base = Path(manifest['source_path']).parent
        def local_address(value):
            try: relative = Path(value).relative_to(address_base)
            except ValueError as exc: raise ValueError('Reading reference leaves its bundle') from exc
            path = inside(directory, relative)
            if not path.is_file(): raise ValueError('Missing reading dependency: ' + str(relative))
            return path
        for key in ('source_path', 'source_ir_path', 'contact_sheet_path', 'pages', 'tables', 'assets', 'placements'):
            local_address(manifest[key])
        pages = read_json(directory / 'pages/index.json')
        if len(pages) != manifest.get('counts', {}).get('pages'):
            raise ValueError('Reading page index count changed')
        for page in pages:
            local_address(page['markdown_path'])
            if file_sha256(local_address(page['high_render']['path'])) != page['high_render']['sha256']:
                raise ValueError('Reading page render changed')
        for table in read_json(directory / 'tables/index.json'):
            local_address(table['json_path']); local_address(table['html_path'])
        for asset in read_json(directory / 'assets/index.json'):
            if asset.get('path') and file_sha256(local_address(asset['path'])) != asset['content_sha256']:
                raise ValueError('Reading original asset changed')
        for sidecar in manifest.get('capture_sidecars', {}).values(): local_address(sidecar)
        stages['READING'] = {'path': str(directory / 'source_ir.json'), 'manifest_path': str(directory / 'reading-workspace.json'), 'count': len(snapshot.nodes), 'capture_state': snapshot.status.value}

    map_path = next((root / 'document-map.json' for root in roots if (root / 'document-map.json').is_file()), None)
    mapping = core.DocumentMapArtifactV1.model_validate(read_json(map_path)) if map_path else None
    if mapping and mapping.source_sha256 != source_sha256:
        raise ValueError('Map source differs from task original')
    baseline = None
    if any(kind in kinds for kind in ('SOURCE_ITEMS', 'ANNOTATIONS', 'DECOMPOSITION', 'DELIVERY')):
        baseline = core.SourceRequirementBaselineArtifactV1.model_validate(read_json(locate('source-requirement-baseline.json')))
        receipt = core.SourceRequirementBaselineValidationReceiptV1.model_validate(read_json(locate('validation-receipt.json')))
        if mapping is None:
            raise ValueError('Source items require their exact Map')
        core._assert_validation_receipt_matches_artifact(document_map=mapping, artifact=baseline, receipt=receipt, error_code='WORKSPACE_BASELINE_RECEIPT_MISMATCH')
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot
        optional = {}
        for key, name in [('annotations', 'image-annotations.json'), ('decomposition', 'source-item-decomposition.json'), ('assessments', 'source-item-assessments.json'), ('manual_decisions', 'manual-decisions.json'), ('analysis_policy', 'analysis-policy.json')]:
            # State is an explicit bundle selection, never "latest by kind".
            for root in roots:
                if (root / name).is_file():
                    optional[key] = read_json(root / name)
                    break
        build_delivery_snapshot(mapping, baseline, read_json(locate('table-previews.json')), **optional)
        stages['ITEMS'] = {'path': str(locate('source-requirement-baseline.json')), 'artifact_sha256': baseline.artifact_sha256, 'map_sha256': mapping.artifact_sha256, 'count': len(baseline.records), 'covered_scope_ids': list(baseline.covered_scope_ids)}
        compilation = directory / 'item-compilation-receipt.json'
        if compilation.is_file():
            value = read_json(compilation)
            if value.get('baseline_artifact_sha256') == baseline.artifact_sha256:
                stages['ITEMS']['unprocessed_scope_ids'] = value.get('unprocessed_scope_ids')
    if 'DOCUMENT_MAP' in kinds:
        receipt = read_json(locate('map-compilation-receipt.json', own=True))
        if receipt.get('source_sha256') != source_sha256 or receipt.get('document_map_artifact_sha256') != mapping.artifact_sha256:
            raise ValueError('Map compilation receipt differs from Map')
        inputs = receipt.get('source_inputs') or {}
        for name, hash_name in [('source_path', None), ('source_ir_path', 'source_ir_sha256'), ('decisions_path', 'decisions_sha256')]:
            if not inputs.get(name):
                raise ValueError('Map replay dependency missing: ' + name)
            path = inside(directory, inputs[name])
            expected = source_sha256 if hash_name is None else inputs.get(hash_name)
            if not path.is_file() or file_sha256(path) != expected:
                raise ValueError('Map replay dependency changed: ' + name)
        source_ir = SourceIRSnapshot.model_validate(read_json(inside(directory, inputs['source_ir_path'])))
        if source_ir.source_sha256 != source_sha256:
            raise ValueError('Map SourceIR source differs')
    if mapping:
        stages['MAP'] = {'path': str(map_path), 'artifact_sha256': mapping.artifact_sha256, 'count': len(mapping.nodes)}
    if 'VISUAL_SOURCE' in kinds:
        from app.application.project_visual_source import verify_visual_capture_extension
        snapshot=SourceIRSnapshot.model_validate(read_json(locate('source_ir.json',own=True)))
        receipt=read_json(locate('visual-source-capture.json',own=True))
        if (snapshot.source_sha256!=source_sha256 or receipt.get('source_sha256')!=source_sha256
            or receipt.get('source_ir_sha256')!=file_sha256(directory/'source_ir.json')):
            raise ValueError('Visual reading source or capture changed')
        predecessors=[p for root in roots[1:] for p in root.glob('*.json')
                      if p.is_file() and file_sha256(p)==receipt.get('base_source_ir_sha256')]
        if not predecessors:raise ValueError('Visual reading requires its exact predecessor capture')
        previous=SourceIRSnapshot.model_validate(read_json(predecessors[0]))
        verify_visual_capture_extension(previous,snapshot)
        stages['READING']={'path':str(directory/'source_ir.json'),'source_ir_id':snapshot.source_ir_id,
                          'count':len(snapshot.nodes),'capture_state':snapshot.status.value,
                          'visual_reading_state':receipt.get('status'),
                          'visual_reading_receipt':str(directory/'visual-source-capture.json')}
    if 'SOURCE_IR' in kinds:
        source_ir = SourceIRSnapshot.model_validate(read_json(locate('source_ir.json', own=True)))
        manifest = read_json(locate('manifest.json', own=True))
        if source_ir.source_sha256 != source_sha256 or manifest.get('source', {}).get('sha256') != source_sha256:
            raise ValueError('Capture source differs from task original')
        entries = manifest.get('artifacts') or {}
        original = inside(directory, entries.get('source_document', ''))
        if not original.is_file() or file_sha256(original) != source_sha256:
            raise ValueError('Capture original is missing or changed')
        for value in entries.get('assets', []):
            path = inside(directory, value['path'])
            if not path.is_file() or file_sha256(path) != value['sha256']:
                raise ValueError('Capture asset is missing or changed')
        stages['READING'] = {'path': str(locate('source_ir.json')), 'source_ir_id': source_ir.source_ir_id, 'count': len(source_ir.nodes), 'capture_state': source_ir.status.value}
    if 'SOURCE_FINDINGS' in kinds:
        findings = read_json(locate('source-findings.json', own=True))
        if findings.get('schema_version') != 'source-findings/v1' or findings.get('source_sha256') != source_sha256 or findings.get('document_map_sha256') != mapping.artifact_sha256:
            raise ValueError('Findings lineage differs from Map')
        stages['SOURCE_FINDINGS'] = {'path': str(locate('source-findings.json', own=True)), 'map_sha256': mapping.artifact_sha256, 'finding_count': len(findings.get('findings', [])), 'source_review_state': 'NOT_CLAIMED'}
    if 'DELIVERY' in kinds:
        snapshot = read_json(locate('matrix-snapshot.json', own=True))
        if snapshot.get('delivery_sha256') != core.canonical_json_sha256({k: v for k, v in snapshot.items() if k != 'delivery_sha256'}):
            raise ValueError('Delivery snapshot hash changed')
        if snapshot.get('source_item_baseline', {}).get('source_sha256') != source_sha256:
            raise ValueError('Delivery source differs from task original')
        identity = snapshot['source_item_baseline']
        if identity['artifact_sha256'] != baseline.artifact_sha256 or identity['document_map_artifact_sha256'] != mapping.artifact_sha256:
            raise ValueError('Delivery differs from its selected Map/baseline dependencies')
        delivery_receipt = read_json(locate('delivery-receipt.json', own=True))
        if delivery_receipt.get('source_sha256') != source_sha256 or delivery_receipt.get('baseline_artifact_sha256') != baseline.artifact_sha256 or delivery_receipt.get('xlsx_sha256') != file_sha256(locate('matrix.xlsx', own=True)):
            raise ValueError('Delivery receipt no longer matches the workbook and baseline')
        from openpyxl import load_workbook
        workbook = load_workbook(locate('matrix.xlsx', own=True), read_only=True, data_only=True)
        try:
            if '_ReqPilotXMeta' not in workbook.sheetnames:
                raise ValueError('Workbook has no exact snapshot metadata')
            values = [row[1] for row in workbook['_ReqPilotXMeta'].iter_rows(values_only=True) if len(row) > 1 and row[0] == 'SnapshotSHA256']
            if values != [core.canonical_json_sha256(snapshot)]:
                raise ValueError('Workbook snapshot differs from saved snapshot')
        finally:
            workbook.close()
    for kind in ('ANNOTATIONS', 'DECOMPOSITION'):
        if kind in kinds:
            value = read_json(locate(ARTIFACT_KINDS[kind]['filename'], own=True))
            stages[kind] = {'path': str(locate(ARTIFACT_KINDS[kind]['filename'], own=True)), 'baseline_sha256': baseline.artifact_sha256, 'map_sha256': mapping.artifact_sha256}
            if kind == 'DECOMPOSITION': stages[kind]['review_scope'] = value.get('review_scope')
            receipt_name = 'annotation-compilation-receipt.json' if kind == 'ANNOTATIONS' else 'decomposition-compilation-receipt.json'
            if (directory / receipt_name).is_file():
                receipt = read_json(directory / receipt_name)
                hash_field = 'annotations_sha256' if kind == 'ANNOTATIONS' else 'artifact_sha256'
                if receipt.get('baseline_sha256') != baseline.artifact_sha256 or receipt.get(hash_field) != value.get(hash_field):
                    raise ValueError('Overlay compilation receipt differs from exact baseline or artifact')

    if 'DELIVERY' in kinds:
        stages['DELIVERY'] = {'path': str(locate('matrix-snapshot.json', own=True)), 'baseline_sha256': snapshot['source_item_baseline']['artifact_sha256'], 'map_sha256': snapshot['source_item_baseline']['document_map_artifact_sha256']}
    for kind in kinds:
        path = locate(ARTIFACT_KINDS[kind]['filename'], own=True)
        value = read_json(path)
        if kind == 'REVIEW_HEAD': path = inside(directory, 'reviews/' + value['review_sha256'] + '.json')
        artifacts.append({'kind': kind, 'path': str(path), 'file_sha256': file_sha256(path), 'artifact_sha256': value.get('artifact_sha256'), 'business_digests': {key: value[key] for key in ('artifact_sha256', 'annotations_sha256', 'delivery_sha256', 'review_sha256') if key in value}, 'display_role': ARTIFACT_KINDS[kind]['display_role']})
    return {'artifacts': artifacts, 'stages': stages, 'deterministic_validation': 'PASSED', 'semantic_verification': 'NOT_CLAIMED', 'human_acceptance': 'NOT_CLAIMED'}
