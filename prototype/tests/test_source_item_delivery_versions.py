from .runtime_isolation import activate_test_runtime
activate_test_runtime()
from contextlib import contextmanager
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from app.application.agent_workflow import source_item_delivery_versions as api
from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot
from app.application.errors import ApplicationError
from .test_source_item_delivery_numbers import fixture


class DeliveryVersionsTests(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.path=Path(self.temp.name)/'test.sqlite';self.objects={}
        m,b,d,a=fixture();self.snapshot=build_delivery_snapshot(m,b,{},decomposition=d,annotations=a)
        self.sha=self.snapshot['source_item_baseline']['source_sha256']
        with self.connection() as conn:
            conn.execute('CREATE TABLE sor_documents(id TEXT,project_id TEXT,source_sha256 TEXT)')
            conn.execute('INSERT INTO sor_documents VALUES(?,?,?)',('d','p',self.sha))
        self.patches=[patch.object(api.db,'get_conn',self.connection),patch.object(api.db,'get_project',return_value={'name':'Synthetic'}),
            patch.object(api.bridge,'_require_project_document',return_value={}),
            patch.object(api.bridge,'_build_current_delivery',side_effect=lambda *args:(deepcopy(self.snapshot),'',{'file_name':'synthetic.pdf'},{'business_basis_sha256':'b'*64,'active_jobs':[]})),
            patch.object(api.bridge,'_container',return_value=SimpleNamespace(object_store=self))]
        for p in self.patches:p.start()
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.temp.cleanup()
    @contextmanager
    def connection(self):
        conn=sqlite3.connect(self.path);conn.row_factory=sqlite3.Row
        try:yield conn;conn.commit()
        except Exception:conn.rollback();raise
        finally:conn.close()
    def put(self,*,key,source,expected_sha256,media_type):
        data=source.read();self.assertEqual(sha256(data).hexdigest(),expected_sha256);self.objects[key]=data
    def open(self,key):return BytesIO(self.objects[key])
    def payload(self,**overrides):
        return api.SaveDeliveryVersion(command_id='one',expected_source_sha256=self.sha,
            expected_delivery_sha256=self.snapshot['delivery_sha256'],expected_business_basis_sha256='b'*64,**overrides)
    def sink(self):return SimpleNamespace(headers={})
    def reseal(self):self.snapshot['delivery_sha256']=api.canonical_json_sha256({k:v for k,v in self.snapshot.items() if k!='delivery_sha256'})
    def test_draft_is_immutable_and_exact_retry_returns_same_version_after_new_work(self):
        payload=self.payload();saved=api.save_delivery_version('p','d',payload,self.sink())
        artifact=api.download_delivery_version('p','d',1);workbook=load_workbook(BytesIO(artifact.content))
        self.assertIn('ComplianceMatrix',workbook.sheetnames)
        self.snapshot['nodes'][-1]['display_text']='new working text';self.reseal()
        replay=api.save_delivery_version('p','d',payload,self.sink());self.assertEqual(replay,saved)
        self.assertEqual(api.download_delivery_version('p','d',1).content,artifact.content)
        frozen=api.get_delivery_version('p','d',1,self.sink())
        self.assertNotEqual(frozen['snapshot']['nodes'][-1]['display_text'],'new working text')
        self.assertEqual(len(api.list_delivery_versions('p','d',self.sink())['versions']),1)
        with self.assertRaises(ApplicationError):api.save_delivery_version('p','d',self.payload(),self.sink())
    def test_stale_preview_rejected_and_file_tampering_detected(self):
        payload=self.payload();self.snapshot['nodes'][-1]['display_text']='changed';self.reseal()
        with self.assertRaises(ApplicationError):api.save_delivery_version('p','d',payload,self.sink())
        api.save_delivery_version('p','d',self.payload(),self.sink())
        self.objects[next(iter(self.objects))]=b'tampered'
        with self.assertRaises(ApplicationError):api.download_delivery_version('p','d',1)
        with self.assertRaises(ApplicationError):api.get_delivery_version('p','another',1,self.sink())
    def test_formal_requires_effective_conclusions_and_explicit_snapshot_acceptance(self):
        with self.assertRaises(ApplicationError):api.save_delivery_version('p','d',self.payload(mode='FORMAL',accept_snapshot=True),self.sink())
        for node in self.snapshot['nodes']:
            if node['disposition']=='ANALYZE' and node['analysis_eligibility']=='ELIGIBLE':
                node.update(analysis_status='ASSESSED',analysis_result_state='CURRENT',ai_assessment={'decision':'NON_COMPLIANT'},compliance_status='NON_COMPLIANT')
        self.snapshot['counts']['assessed_nodes'] = 2
        self.reseal()
        self.assertTrue(api.readiness(self.snapshot,{})['can_lock'])
        with self.assertRaises(ApplicationError):api.save_delivery_version('p','d',self.payload(mode='FORMAL'),self.sink())
        saved=api.save_delivery_version('p','d',self.payload(mode='FORMAL',accept_snapshot=True),self.sink())
        self.assertEqual(saved['acceptance']['accepted_delivery_sha256'],self.snapshot['delivery_sha256'])
        self.assertEqual(saved['mode'],'FORMAL')
