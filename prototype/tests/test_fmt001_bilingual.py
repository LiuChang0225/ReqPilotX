from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import copy
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pydantic import ValidationError

from app.adapters.documents.docx_source_ir_adapter import DocxSourceIRAdapter
from app.adapters.documents.preflight import preflight_document
from app.application.extraction.classify_content import classify_content
from app.application.extraction.build_llm_shadow_preview import (
    _serialize_structure_agent,
)
from app.application.extraction.orchestrate_structure_agent import (
    initialize_structure_repair_without_provider,
)
from app.application.extraction.project_requirements import project_requirements
from app.application.extraction.resolve_structure import build_structure_tree
from app.application.extraction.serialize_structure_workbench import (
    serialize_structure_workbench,
)
from app.application.extraction.structure_presentation import (
    STRUCTURE_PRESENTATION_INDEX_SCHEMA_VERSION,
    build_structure_presentation_index,
    validated_structure_presentation_index,
)
from app.domain.extraction.source_ir import (
    DocxInlineBoundaryBasis,
    DocxInlineLanguageEvidence,
    NativeDocumentFormat,
    NativeLocator,
    SourceIRSnapshot,
    SourceIRStatus,
    SourceNode,
    SourceNodeKind,
    stable_document_id,
    stable_node_id,
    stable_source_ir_id,
)
from app.domain.extraction.structure import (
    BilingualCounterpartGroup,
    BilingualLanguageEvidence,
    BilingualMismatchSignal,
    BilingualSourceSpan,
    BilingualSpanBoundaryBasis,
    DocumentLanguagePolicy,
    StructureFindingSeverity,
    StructureRelationKind,
    StructureSourceRef,
    StructureTreeSnapshot,
    stable_bilingual_counterpart_group_id,
    stable_bilingual_mismatch_signal_id,
    stable_bilingual_source_span_id,
    stable_document_language_policy_id,
)


CREATED_AT = datetime(2026, 7, 26, 8, 0, tzinfo=timezone.utc)
SOURCE_HASH = "7" * 64
ADAPTER_ID = "synthetic-docx-fmt001"
ADAPTER_VERSION = "1.0.0"
EN_POLICY = (
    "English and Chinese versions are equally valid. In case of conflict, "
    "the Chinese version shall prevail."
)
ZH_POLICY = "中英文版本具有同等效力。如有冲突，以中文版本为准。"
EN_REQUIREMENT = (
    "The controller shall retain diagnostic records for at least thirty days."
)
ZH_REQUIREMENT = "控制器应至少保留三十天的诊断记录。"


class Fmt001DocxInlineSpanTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)

    def test_run_boundary_preserves_raw_text_and_exact_en_then_zh_spans(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        english = "The controller shall retain all diagnostic records. "
        chinese = "控制器应保留全部诊断记录。"
        paragraph.add_run(english)
        paragraph.add_run(chinese)
        path = self.root / "inline.docx"
        document.save(path)

        snapshot = _extract_docx(path)
        node = next(node for node in snapshot.nodes if node.raw_text)

        self.assertEqual(node.raw_text, english + chinese)
        self.assertEqual(len(node.docx_inline_spans), 2)
        first, second = node.docx_inline_spans
        self.assertEqual(
            (
                first.source_node_id,
                first.start_offset,
                first.end_offset,
                first.language_evidence,
                first.boundary_basis,
            ),
            (
                node.source_node_id,
                0,
                len(english),
                DocxInlineLanguageEvidence.EN,
                DocxInlineBoundaryBasis.DOCX_RUN_BOUNDARY,
            ),
        )
        self.assertEqual(second.start_offset, len(english))
        self.assertEqual(second.end_offset, len(node.raw_text))
        self.assertIs(
            second.language_evidence,
            DocxInlineLanguageEvidence.ZH_CN,
        )

    def test_chinese_first_acronym_and_style_only_runs_do_not_split(self) -> None:
        document = Document()
        chinese_first = document.add_paragraph()
        chinese_first.add_run("控制器应支持")
        chinese_first.add_run(" ABS ")
        chinese_first.add_run("诊断功能。")
        styled_english = document.add_paragraph()
        styled_english.add_run("The controller shall ").bold = True
        styled_english.add_run("retain diagnostic records.").italic = True
        path = self.root / "negative-inline.docx"
        document.save(path)

        snapshot = _extract_docx(path)
        textual_nodes = tuple(node for node in snapshot.nodes if node.raw_text)

        self.assertEqual(len(textual_nodes), 2)
        self.assertTrue(all(not node.docx_inline_spans for node in textual_nodes))

    def test_list_presentation_is_outside_raw_text_offsets(self) -> None:
        document = Document()
        paragraph = document.add_paragraph(style="List Number")
        english = "The unit shall report an alarm. "
        paragraph.add_run(english)
        paragraph.add_run("装置应报告报警。")
        path = self.root / "list-inline.docx"
        document.save(path)

        snapshot = _extract_docx(path)
        node = next(node for node in snapshot.nodes if node.raw_text)

        self.assertIs(node.kind, SourceNodeKind.LIST_ITEM)
        self.assertEqual(node.docx_inline_spans[0].start_offset, 0)
        self.assertEqual(node.docx_inline_spans[0].end_offset, len(english))


class Fmt001StructureAndProjectionTests(unittest.TestCase):
    def test_adjacent_pair_has_policy_relation_and_exactly_one_group_view(self) -> None:
        source = _source_snapshot(
            (EN_POLICY, ZH_POLICY, EN_REQUIREMENT, ZH_REQUIREMENT)
        )

        tree = build_structure_tree(source)
        semantics = classify_content(source, tree)
        view = project_requirements(source, tree, semantics)

        self.assertIsNotNone(tree.document_language_policy)
        assert tree.document_language_policy is not None
        self.assertTrue(tree.document_language_policy.equal_validity)
        self.assertEqual(
            tree.document_language_policy.conflict_precedence,
            "zh-CN",
        )
        self.assertEqual(len(tree.bilingual_counterpart_groups), 1)
        group = tree.bilingual_counterpart_groups[0]
        self.assertEqual(len(group.member_structure_node_ids), 2)
        relation = next(
            relation
            for relation in tree.relations
            if relation.relation_id == group.translation_relation_id
        )
        self.assertIs(relation.kind, StructureRelationKind.TRANSLATION_OF)
        group_views = tuple(
            node
            for node in view.nodes
            if node.owner_structure_node_id in group.member_structure_node_ids
        )
        self.assertEqual(len(group_views), 1)
        self.assertEqual(len(group_views[0].source_semantic_unit_ids), 2)
        self.assertEqual(len(group_views[0].source_span_ids), 2)
        serialized = tree.model_dump_json()
        self.assertNotIn(EN_REQUIREMENT, serialized)
        self.assertNotIn(ZH_REQUIREMENT, serialized)
        self.assertNotIn("replacement_text", serialized)

    def test_inline_pair_is_one_owner_without_self_relation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "inline-pipeline.docx"
            document = Document()
            document.add_paragraph(EN_POLICY)
            document.add_paragraph(ZH_POLICY)
            paragraph = document.add_paragraph()
            paragraph.add_run(EN_REQUIREMENT + " ")
            paragraph.add_run(ZH_REQUIREMENT)
            document.save(path)

            source = _extract_docx(path)
            tree = build_structure_tree(source)
            semantics = classify_content(source, tree)
            view = project_requirements(source, tree, semantics)

        inline_group = next(
            group
            for group in tree.bilingual_counterpart_groups
            if any(
                span.boundary_basis is BilingualSpanBoundaryBasis.DOCX_RUN_BOUNDARY
                for span in (*group.english_spans, *group.chinese_spans)
            )
        )
        self.assertEqual(len(inline_group.member_structure_node_ids), 1)
        self.assertIsNone(inline_group.translation_relation_id)
        self.assertFalse(
            any(
                relation.kind is StructureRelationKind.TRANSLATION_OF
                and relation.source_structure_node_id
                == relation.target_structure_node_id
                for relation in tree.relations
            )
        )
        self.assertEqual(
            len(
                [
                    node
                    for node in view.nodes
                    if node.owner_structure_node_id
                    in inline_group.member_structure_node_ids
                ]
            ),
            1,
        )

    def test_reviewed_table_e_z_e_z_is_one_owner_and_one_view(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "table-ezez.docx"
            document = Document()
            document.add_paragraph(EN_POLICY)
            document.add_paragraph(ZH_POLICY)
            cell = document.add_table(rows=1, cols=1).cell(0, 0)
            cell.paragraphs[0].text = EN_REQUIREMENT
            cell.add_paragraph(ZH_REQUIREMENT)
            cell.add_paragraph(
                "The diagnostic interface shall expose the retained records."
            )
            cell.add_paragraph("诊断接口应提供所保留的记录。")
            document.save(path)

            source = _extract_docx(path)
            tree = build_structure_tree(source)
            semantics = classify_content(source, tree)
            view = project_requirements(source, tree, semantics)

        group = next(
            group
            for group in tree.bilingual_counterpart_groups
            if len(group.english_spans) == 2 and len(group.chinese_spans) == 2
        )
        self.assertEqual(len(group.member_structure_node_ids), 1)
        group_views = tuple(
            node
            for node in view.nodes
            if node.owner_structure_node_id in group.member_structure_node_ids
        )
        self.assertEqual(len(group_views), 1)

    def test_missing_conflicting_and_multiple_policy_fail_closed(self) -> None:
        cases = (
            (
                (EN_REQUIREMENT, ZH_REQUIREMENT),
                "document_language_policy_missing",
            ),
            (
                (
                    "English and Chinese versions are equally valid. In case "
                    "of conflict, the English version shall prevail.",
                    "中英文版本具有同等效力。如有冲突，以英文版本为准。",
                    EN_REQUIREMENT,
                    ZH_REQUIREMENT,
                ),
                "document_language_policy_conflict",
            ),
            (
                (
                    EN_POLICY,
                    EN_POLICY,
                    ZH_POLICY,
                    EN_REQUIREMENT,
                    ZH_REQUIREMENT,
                ),
                "document_language_policy_multiple",
            ),
        )
        for texts, code in cases:
            with self.subTest(code=code):
                tree = build_structure_tree(_source_snapshot(texts))
                self.assertIsNone(tree.document_language_policy)
                self.assertEqual(tree.bilingual_counterpart_groups, ())
                finding = next(item for item in tree.findings if item.code == code)
                self.assertTrue(finding.blocking)
                self.assertIs(finding.severity, StructureFindingSeverity.ERROR)


class Fmt001SafetyContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.source = _source_snapshot(
            (EN_POLICY, ZH_POLICY, EN_REQUIREMENT, ZH_REQUIREMENT)
        )
        self.tree = build_structure_tree(self.source)
        self.group = self.tree.bilingual_counterpart_groups[0]

    def test_mismatch_signal_is_warning_nonblocking_and_double_grounded(self) -> None:
        english_span_ids = tuple(
            span.bilingual_span_id for span in self.group.english_spans
        )
        chinese_span_ids = tuple(
            span.bilingual_span_id for span in self.group.chinese_spans
        )
        signal_id = stable_bilingual_mismatch_signal_id(
            counterpart_group_id=self.group.counterpart_group_id,
            english_span_ids=english_span_ids,
            chinese_span_ids=chinese_span_ids,
        )
        signal = BilingualMismatchSignal(
            mismatch_signal_id=signal_id,
            counterpart_group_id=self.group.counterpart_group_id,
            english_span_ids=english_span_ids,
            chinese_span_ids=chinese_span_ids,
        )
        payload = self.tree.model_dump(mode="python")
        payload["bilingual_mismatch_signals"] = (signal,)
        validated = StructureTreeSnapshot.model_validate(payload)

        self.assertIs(
            validated.bilingual_mismatch_signals[0].severity,
            StructureFindingSeverity.WARNING,
        )
        self.assertFalse(validated.bilingual_mismatch_signals[0].blocking)
        self.assertNotIn(
            "replacement",
            validated.bilingual_mismatch_signals[0].model_dump_json(),
        )
        self.assertEqual(self.tree.bilingual_mismatch_signals, ())

    def test_mismatch_contract_rejects_blocking_missing_side_and_replacement(
        self,
    ) -> None:
        english_span_ids = tuple(
            span.bilingual_span_id for span in self.group.english_spans
        )
        chinese_span_ids = tuple(
            span.bilingual_span_id for span in self.group.chinese_spans
        )
        valid = {
            "mismatch_signal_id": stable_bilingual_mismatch_signal_id(
                counterpart_group_id=self.group.counterpart_group_id,
                english_span_ids=english_span_ids,
                chinese_span_ids=chinese_span_ids,
            ),
            "counterpart_group_id": self.group.counterpart_group_id,
            "english_span_ids": english_span_ids,
            "chinese_span_ids": chinese_span_ids,
        }
        for changes in (
            {"blocking": True},
            {"severity": StructureFindingSeverity.ERROR},
            {"chinese_span_ids": ()},
            {"replacement_text": "forbidden"},
        ):
            with self.subTest(changes=tuple(changes)):
                with self.assertRaises(ValidationError):
                    BilingualMismatchSignal(**(valid | changes))

    def test_group_contract_rejects_missing_side_gap_and_cross_owner_span(self) -> None:
        with self.assertRaises(ValidationError):
            BilingualCounterpartGroup(
                counterpart_group_id="bilingual-group-" + "0" * 32,
                document_language_policy_id=self.group.document_language_policy_id,
                canonical_structure_node_id=self.group.canonical_structure_node_id,
                member_structure_node_ids=self.group.member_structure_node_ids,
                english_spans=self.group.english_spans,
                chinese_spans=(),
                translation_relation_id=self.group.translation_relation_id,
            )

        source_ref = self.group.english_spans[0].source_ref
        english = _span(
            source_ref=source_ref,
            start=0,
            end=5,
            language=BilingualLanguageEvidence.EN,
        )
        chinese = _span(
            source_ref=source_ref,
            start=6,
            end=10,
            language=BilingualLanguageEvidence.ZH_CN,
        )
        with self.assertRaises(ValidationError):
            BilingualCounterpartGroup(
                counterpart_group_id="bilingual-group-" + "0" * 32,
                document_language_policy_id=self.group.document_language_policy_id,
                canonical_structure_node_id=self.group.canonical_structure_node_id,
                member_structure_node_ids=(self.group.canonical_structure_node_id,),
                english_spans=(english,),
                chinese_spans=(chinese,),
            )

        assert self.tree.document_language_policy is not None
        foreign_span = self.tree.document_language_policy.english_declaration_spans[0]
        group_id = stable_bilingual_counterpart_group_id(
            document_language_policy_id=self.group.document_language_policy_id,
            canonical_structure_node_id=self.group.canonical_structure_node_id,
            member_structure_node_ids=self.group.member_structure_node_ids,
            english_spans=(foreign_span,),
            chinese_spans=self.group.chinese_spans,
        )
        cross_owner = BilingualCounterpartGroup(
            counterpart_group_id=group_id,
            document_language_policy_id=self.group.document_language_policy_id,
            canonical_structure_node_id=self.group.canonical_structure_node_id,
            member_structure_node_ids=self.group.member_structure_node_ids,
            english_spans=(foreign_span,),
            chinese_spans=self.group.chinese_spans,
            translation_relation_id=self.group.translation_relation_id,
        )
        payload = self.tree.model_dump(mode="python")
        payload["bilingual_counterpart_groups"] = (cross_owner,)
        with self.assertRaises(ValidationError):
            StructureTreeSnapshot.model_validate(payload)


class PdfBilingualPairedRunTests(unittest.TestCase):
    def test_repeated_numbered_pdf_pairs_create_unknown_typed_groups(self) -> None:
        source = _pdf_paired_run_snapshot(
            (
                (
                    "控制器应保留 3 条诊断记录。",
                    "The controller shall retain 3 diagnostic records.",
                ),
                ("装置应报告故障状态。", "The unit shall report the fault status."),
                (
                    "供应商应提交验证报告。",
                    "The supplier shall submit a validation report.",
                ),
            )
        )

        tree = build_structure_tree(source)

        self.assertEqual(len(tree.bilingual_counterpart_groups), 3)
        self.assertIsNotNone(tree.document_language_policy)
        assert tree.document_language_policy is not None
        self.assertEqual(tree.document_language_policy.authority_basis, "UNKNOWN")
        self.assertEqual(tree.document_language_policy.conflict_precedence, "UNKNOWN")
        self.assertIsNone(tree.document_language_policy.equal_validity)
        self.assertEqual(
            tree.document_language_policy.policy_scope_source_ir_id,
            source.source_ir_id,
        )
        self.assertTrue(
            all(
                len(group.member_structure_node_ids) == 1
                and group.translation_relation_id is None
                and len(group.english_spans) == 1
                and len(group.chinese_spans) == 1
                for group in tree.bilingual_counterpart_groups
            )
        )

    def test_isolated_adjacency_and_non_direct_content_fail_closed(self) -> None:
        isolated = build_structure_tree(
            _pdf_paired_run_snapshot(
                (
                    (
                        "控制器应保留诊断记录。",
                        "The controller shall retain diagnostic records.",
                    ),
                )
            )
        )
        self.assertIsNone(isolated.document_language_policy)
        self.assertEqual(isolated.bilingual_counterpart_groups, ())

        unrelated = build_structure_tree(
            _pdf_paired_run_snapshot(
                (
                    (
                        "控制器应保留诊断记录。",
                        "The controller shall retain diagnostic records.",
                    ),
                    ("装置应报告故障状态。", "The unit shall report the fault status."),
                    (
                        "供应商应提交验证报告。",
                        "The supplier shall submit a validation report.",
                    ),
                ),
                break_direct_owner_at=1,
            )
        )
        self.assertIsNone(unrelated.document_language_policy)
        self.assertEqual(unrelated.bilingual_counterpart_groups, ())

    def test_numeric_mismatch_is_preserved_as_double_grounded_review_signal(
        self,
    ) -> None:
        tree = build_structure_tree(
            _pdf_paired_run_snapshot(
                (
                    (
                        "控制器应保留 3 条诊断记录。",
                        "The controller shall retain 3 diagnostic records.",
                    ),
                    ("装置应报告故障状态。", "The unit shall report the fault status."),
                    ("供应商应提交 6 份报告。", "The supplier shall submit 8 reports."),
                )
            )
        )

        self.assertEqual(len(tree.bilingual_mismatch_signals), 1)
        signal = tree.bilingual_mismatch_signals[0]
        group = next(
            group
            for group in tree.bilingual_counterpart_groups
            if group.counterpart_group_id == signal.counterpart_group_id
        )
        self.assertEqual(
            signal.english_span_ids,
            tuple(span.bilingual_span_id for span in group.english_spans),
        )
        self.assertEqual(
            signal.chinese_span_ids,
            tuple(span.bilingual_span_id for span in group.chinese_spans),
        )

    def test_presentation_round_trip_preserves_typed_bilingual_sidecar(self) -> None:
        source, tree = _pdf_bilingual_mismatch_tree()

        presentation = build_structure_presentation_index(source, structure=tree)
        validated = validated_structure_presentation_index(
            presentation,
            expected_source_sha256=source.source_sha256,
            expected_structure_snapshot_id=tree.structure_snapshot_id,
        )

        self.assertIsNotNone(validated)
        assert validated is not None
        self.assertEqual(
            validated["schema_version"],
            STRUCTURE_PRESENTATION_INDEX_SCHEMA_VERSION,
        )
        sidecar = validated["bilingual_projection"]
        self.assertEqual(len(sidecar["counterpart_groups"]), 3)
        self.assertEqual(len(sidecar["mismatch_signals"]), 1)
        self.assertEqual(
            sidecar["document_language_policy"]["policy_scope_source_ir_id"],
            source.source_ir_id,
        )

    def test_missing_or_tampered_presentation_sidecar_fails_closed(self) -> None:
        source, tree = _pdf_bilingual_mismatch_tree()
        presentation = build_structure_presentation_index(source, structure=tree)

        missing = copy.deepcopy(presentation)
        del missing["bilingual_projection"]
        self.assertIsNone(
            validated_structure_presentation_index(
                missing,
                expected_source_sha256=source.source_sha256,
                expected_structure_snapshot_id=tree.structure_snapshot_id,
            )
        )

        tampered = copy.deepcopy(presentation)
        tampered["bilingual_projection"]["counterpart_groups"][0]["english_spans"][0][
            "end_offset"
        ] += 1
        self.assertIsNone(
            validated_structure_presentation_index(
                tampered,
                expected_source_sha256=source.source_sha256,
                expected_structure_snapshot_id=tree.structure_snapshot_id,
            )
        )

    def test_completed_result_without_source_replay_restores_bilingual_context(
        self,
    ) -> None:
        source, tree = _pdf_bilingual_mismatch_tree()
        repair = initialize_structure_repair_without_provider(
            source,
            tree,
            error_code="structure_audit_not_requested",
        )
        structure_agent = _serialize_structure_agent(
            repair,
            editable=True,
            final=True,
        )
        payload = {
            "schema_version": "extraction-llm-shadow-preview/v5",
            "mode": "LOCAL_STRUCTURE_PROJECTION",
            "source_sha256": source.source_sha256,
            "status": "READY",
            "items": [],
            "requirement_groups": [],
            "review_groups": [],
            "pipeline": {
                "source_ir": {
                    "id": source.source_ir_id,
                    "version": source.adapter_version,
                    "status": "COMPLETE",
                },
                "semantic": {"status": "COMPLETE"},
                "requirement_view": {"status": "SHADOW_READY"},
            },
            "structure_agent": structure_agent,
            "structure_presentation_index": build_structure_presentation_index(
                source,
                structure=repair.effective_structure,
            ),
        }

        workbench = serialize_structure_workbench(
            project_id="project-bilingual-read-path",
            document_id="document-bilingual-read-path",
            document={
                "file_name": "source.pdf",
                "source_sha256": source.source_sha256,
            },
            job={"status": "COMPLETED"},
            payload=payload,
            structure_agent=structure_agent,
            source_ir=None,
            base_tree=None,
        )

        self.assertEqual(
            len(workbench["structure"]["bilingual_projection"]["counterpart_groups"]),
            3,
        )
        bilingual_nodes = tuple(
            node
            for node in workbench["requirement_tree"]["nodes"]
            if node.get("bilingual_context") is not None
        )
        self.assertEqual(len(bilingual_nodes), 3)
        mismatch_node = next(
            node
            for node in bilingual_nodes
            if node["bilingual_context"]["mismatch_review_required"]
        )
        self.assertEqual(mismatch_node["status"], "NEEDS_ATTENTION")
        self.assertEqual(mismatch_node["attention_count"], 1)

    def test_unknown_policy_identity_and_scope_are_fail_closed(self) -> None:
        source_ir_id = "sir-" + "8" * 32
        policy_id = stable_document_language_policy_id(
            english_declaration_spans=(),
            chinese_declaration_spans=(),
            equal_validity=None,
            conflict_precedence="UNKNOWN",
            authority_basis="UNKNOWN",
            policy_scope_source_ir_id=source_ir_id,
        )
        policy = DocumentLanguagePolicy(
            document_language_policy_id=policy_id,
            authority_basis="UNKNOWN",
            policy_scope_source_ir_id=source_ir_id,
            equal_validity=None,
            conflict_precedence="UNKNOWN",
        )

        self.assertEqual(
            DocumentLanguagePolicy.model_validate_json(policy.model_dump_json()),
            policy,
        )
        with self.assertRaises(ValidationError):
            DocumentLanguagePolicy(
                document_language_policy_id=policy_id,
                authority_basis="UNKNOWN",
                policy_scope_source_ir_id=source_ir_id,
                equal_validity=True,
                conflict_precedence="zh-CN",
            )


def _extract_docx(path: Path) -> SourceIRSnapshot:
    return DocxSourceIRAdapter().extract(
        path,
        preflight=preflight_document(path),
        created_at=CREATED_AT,
    )


def _pdf_bilingual_mismatch_tree() -> tuple[
    SourceIRSnapshot,
    StructureTreeSnapshot,
]:
    source = _pdf_paired_run_snapshot(
        (
            (
                "控制器应保留 3 条诊断记录。",
                "The controller shall retain 3 diagnostic records.",
            ),
            (
                "装置应报告故障状态。",
                "The unit shall report the fault status.",
            ),
            (
                "供应商应提交 6 份报告。",
                "The supplier shall submit 8 reports.",
            ),
        )
    )
    return source, build_structure_tree(source)


def _source_snapshot(texts: tuple[str, ...]) -> SourceIRSnapshot:
    document_id = stable_document_id(SOURCE_HASH)
    nodes: list[SourceNode] = []
    for source_order, raw_text in enumerate(texts):
        locator = NativeLocator(
            format=NativeDocumentFormat.DOCX,
            body_element_index=source_order,
            paragraph_index=source_order,
        )
        nodes.append(
            SourceNode(
                source_node_id=stable_node_id(
                    source_sha256=SOURCE_HASH,
                    adapter_id=ADAPTER_ID,
                    adapter_version=ADAPTER_VERSION,
                    kind=SourceNodeKind.PARAGRAPH,
                    native_locator=locator,
                ),
                kind=SourceNodeKind.PARAGRAPH,
                source_order=source_order,
                raw_text=raw_text,
                native_locator=locator,
            )
        )
    return SourceIRSnapshot(
        source_ir_id=stable_source_ir_id(
            document_id=document_id,
            source_sha256=SOURCE_HASH,
            adapter_id=ADAPTER_ID,
            adapter_version=ADAPTER_VERSION,
        ),
        document_id=document_id,
        source_sha256=SOURCE_HASH,
        adapter_id=ADAPTER_ID,
        adapter_version=ADAPTER_VERSION,
        status=SourceIRStatus.COMPLETE,
        nodes=tuple(nodes),
        created_at=CREATED_AT,
    )


def _pdf_paired_run_snapshot(
    pairs: tuple[tuple[str, str], ...],
    *,
    break_direct_owner_at: int | None = None,
) -> SourceIRSnapshot:
    document_id = stable_document_id(SOURCE_HASH)
    nodes: list[SourceNode] = []
    heading_locator = NativeLocator(
        format=NativeDocumentFormat.PDF,
        page_number=1,
        block_index=0,
        bbox=(72.0, 72.0, 300.0, 86.0),
    )
    heading_id = stable_node_id(
        source_sha256=SOURCE_HASH,
        adapter_id=ADAPTER_ID,
        adapter_version=ADAPTER_VERSION,
        kind=SourceNodeKind.HEADING,
        native_locator=heading_locator,
    )
    nodes.append(
        SourceNode(
            source_node_id=heading_id,
            kind=SourceNodeKind.HEADING,
            source_order=0,
            raw_text="2.4 Test notes",
            native_locator=heading_locator,
        )
    )
    for index, (chinese, english) in enumerate(pairs, start=1):
        chinese_locator = NativeLocator(
            format=NativeDocumentFormat.PDF,
            page_number=1,
            block_index=index * 2 - 1,
            bbox=(88.0, 88.0 + index * 40.0, 480.0, 102.0 + index * 40.0),
        )
        chinese_id = stable_node_id(
            source_sha256=SOURCE_HASH,
            adapter_id=ADAPTER_ID,
            adapter_version=ADAPTER_VERSION,
            kind=SourceNodeKind.PARAGRAPH,
            native_locator=chinese_locator,
        )
        nodes.append(
            SourceNode(
                source_node_id=chinese_id,
                kind=SourceNodeKind.PARAGRAPH,
                source_order=index * 2 - 1,
                raw_text=f"{index}） {chinese}",
                native_locator=chinese_locator,
                parent_source_node_id=heading_id,
            )
        )
        english_locator = NativeLocator(
            format=NativeDocumentFormat.PDF,
            page_number=1,
            block_index=index * 2,
            bbox=(98.0, 106.0 + index * 40.0, 500.0, 120.0 + index * 40.0),
        )
        english_id = stable_node_id(
            source_sha256=SOURCE_HASH,
            adapter_id=ADAPTER_ID,
            adapter_version=ADAPTER_VERSION,
            kind=SourceNodeKind.PARAGRAPH,
            native_locator=english_locator,
        )
        nodes.append(
            SourceNode(
                source_node_id=english_id,
                kind=SourceNodeKind.PARAGRAPH,
                source_order=index * 2,
                raw_text=english,
                native_locator=english_locator,
                parent_source_node_id=(
                    heading_id if break_direct_owner_at == index - 1 else chinese_id
                ),
            )
        )
    return SourceIRSnapshot(
        source_ir_id=stable_source_ir_id(
            document_id=document_id,
            source_sha256=SOURCE_HASH,
            adapter_id=ADAPTER_ID,
            adapter_version=ADAPTER_VERSION,
        ),
        document_id=document_id,
        source_sha256=SOURCE_HASH,
        adapter_id=ADAPTER_ID,
        adapter_version=ADAPTER_VERSION,
        status=SourceIRStatus.COMPLETE,
        nodes=tuple(nodes),
        created_at=CREATED_AT,
    )


def _span(
    *,
    source_ref: StructureSourceRef,
    start: int,
    end: int,
    language: BilingualLanguageEvidence,
) -> BilingualSourceSpan:
    span_id = stable_bilingual_source_span_id(
        source_ref=source_ref,
        start_offset=start,
        end_offset=end,
        language_evidence=language,
        boundary_basis=BilingualSpanBoundaryBasis.SOURCE_BLOCK_BOUNDARY,
    )
    return BilingualSourceSpan(
        bilingual_span_id=span_id,
        source_ref=source_ref,
        start_offset=start,
        end_offset=end,
        language_evidence=language,
        boundary_basis=BilingualSpanBoundaryBasis.SOURCE_BLOCK_BOUNDARY,
    )


if __name__ == "__main__":
    unittest.main()
