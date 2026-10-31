"""Actual main HTTP/registration/native-review/export contracts, synthetic only.

The existing source-state loader is bound to a public-compiler synthetic bundle;
this does not exercise a live source/model worker or claim customer semantics.
Database, original capture, main router, core compiler, manual response storage,
canonical XLSX and immutable delivery storage are real implementations.
"""
from .runtime_isolation import activate_test_runtime
activate_test_runtime()

from argparse import Namespace
from copy import deepcopy
from contextlib import contextmanager
from dataclasses import replace
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import sys
import sqlite3
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(REPO / 'plugins/reqpilot/scripts'), str(REPO / 'plugins/reqpilot/tests')]
import artifact_tools
from portable_fixtures import make_fixture, item_decisions, write
from reqpilotx_server import db
from reqpilotx_server.api.app_factory import create_app
from reqpilotx_server.application import runtime, knowledge, openkb_evidence as evidence
from reqpilotx_server.application.agent_workflow import source_item_delivery_bridge as bridge
from reqpilotx_server.application.agent_workflow import source_requirement_reference_bridge as source_bridge
from reqpilotx_server.application.agent_workflow import source_item_manual_decision_api as engineer
from reqpilotx_server.adapters.storage.local_object_store import LocalObjectStore
from reqpilotx_server.infrastructure.runtime_settings import DeploymentMode
from reqpilotx.application.agent_workflow import source_requirement_baseline as core
from reqpilotx.application.agent_workflow.source_item_visual import compile_source_item_annotation_updates

TEXT = b'SYNTHETIC ORIGINAL ONLY. Valve sample A opened in 28 ms and closed in 35 ms at 12 V and 23 C. Not a real validation report.\n'


class EmptyJobs:
    def get(self, *_args): return None
    def get_latest(self, **_kwargs): return None
    def get_latest_succeeded(self, **_kwargs): return None


class NativeEvidenceHostTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(); self.root = Path(self.temp.name)
        self.patches = [patch.object(db, 'DB_PATH', self.root / 'native.sqlite'), patch.object(db, 'get_conn', self.connection)]
        for item in self.patches: item.start()
        db.init_db()
        self.project = db.create_project('Synthetic local project', '')['id']
        self.other = db.create_project('Other isolated project', '')['id']
        self.source, capture, map_input = make_fixture(self.root / 'fixture')
        path = self.root / 'map-input.json'; write(path, map_input)
        map_dir = self.root / 'map'
        artifact_tools.compile_map(Namespace(source=self.source, source_ir=capture / 'source_ir.json', input=path, output_dir=map_dir))
        path = self.root / 'items.json'; write(path, item_decisions(map_dir))
        baseline_dir = self.root / 'baseline'
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=None, output_dir=baseline_dir))
        self.bundle = {'document_map': core.DocumentMapArtifactV1.model_validate_json((baseline_dir / 'document-map.json').read_text()),
            'baseline': core.SourceRequirementBaselineArtifactV1.model_validate_json((baseline_dir / 'source-requirement-baseline.json').read_text()),
            'image_annotations': None, 'analysis_policy': None}
        self.previews = json.loads((baseline_dir / 'table-previews.json').read_text())
        self.document = db.add_sor_document(self.project, self.source.name, 'PDF', str(self.source), '',
            source_sha256=sha256(self.source.read_bytes()).hexdigest())['id']
        self.container = SimpleNamespace(object_store=LocalObjectStore(self.root / 'objects'), jobs=EmptyJobs(), agent_structure_loop_repository=object())
        self.patches += [
            patch.object(runtime.RUNTIME_STATE, 'container', self.container),
            patch.object(bridge, '_current', side_effect=lambda *_args: (db.get_sor_document(self.document), self.bundle)),
            patch.object(bridge, '_source_context', side_effect=lambda *_args: ({}, self.bundle, self.source, self.previews)),
            patch.object(source_bridge, '_load_effective_reference', side_effect=lambda **_kwargs: self.bundle),
            patch.object(knowledge, '_project_knowledge_sync_profile_ids', side_effect=AssertionError('Source-only must not select a model/index profile')),
            patch.object(knowledge, '_sync_knowledge_document_targets', side_effect=AssertionError('No indexing/model call')),
            patch.object(knowledge, '_enqueue_durable_knowledge_sync', side_effect=AssertionError('No background indexing')),
        ]
        for item in self.patches[2:]: item.start()
        self.settings = runtime.RUNTIME_SETTINGS
        self.app = create_app(settings=self.settings)
        self.client = TestClient(self.app)
        self.project_root = f'/api/projects/{self.project}/openkb-evidence'
        self.root_api = f'/api/projects/{self.project}/sor-documents/{self.document}/agent-structure/delivery'

    def tearDown(self):
        self.client.close()
        for item in reversed(self.patches): item.stop()
        self.temp.cleanup()

    @contextmanager
    def connection(self):
        # Real isolated SQLite transactions. Avoid the global test adapter's
        # TRUNCATE checkpoint on every nested read while an outer write is open;
        # that maintenance waits30s per read and is not a product-stage fixture.
        conn = sqlite3.connect(self.root / 'native.sqlite', timeout=3, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    def upload(self, name='reports/test.txt', text=TEXT):
        return self.client.post(self.project_root + '/upload', files=[('files', (name, text, 'text/plain'))],
            data={'doc_type': 'CASE', 'authority': 'INTERNAL_CONTROLLED', 'claim_role': 'TEST_RESULT', 'applicability': 'sample A, 12 V, 23 C only'})

    def session(self, **override):
        status = self.client.get(self.project_root + '/status').json()
        body = {'expected_source_sha256': self.bundle['baseline'].source_sha256,
            'expected_baseline_sha256': self.bundle['baseline'].artifact_sha256,
            'expected_knowledge_base_version': status['knowledge_base_version'],
            'selected_unit_ids': [self.bundle['baseline'].records[0].source_record_id]}
        body.update(override)
        return self.client.post(self.root_api + '/openkb-evidence/sessions', json=body)

    def original(self, session, knowledge_id):
        return self.client.get(self.root_api + f'/openkb-evidence/sessions/{session["session_id"]}/originals/{knowledge_id}')

    def prepare(self, session, ids):
        return self.client.post(self.root_api + f'/openkb-evidence/sessions/{session["session_id"]}/prepare-review',
            json={'expected_packet_sha256': session['packet_sha256'], 'evidence_ids': ids})

    def test_actual_registration_original_capture_prepare_native_review_and_immutable_export(self):
        uploaded = self.upload(); self.assertEqual(uploaded.status_code, 200, uploaded.text)
        value = uploaded.json(); self.assertEqual((value['registered_count'], value['failed_count']), (1, 0))
        self.assertEqual((value['index_status'], value['live_openkb']), ('NOT_REQUESTED', 'NOT_RUN'))
        knowledge_id = value['documents'][0]['document_id']
        raw = self.client.get(f'/api/projects/{self.project}/knowledge-documents/{knowledge_id}/file', params={'expected_source_sha256': sha256(TEXT).hexdigest()})
        self.assertEqual(raw.content, TEXT)
        session = self.session().json()
        original = self.original(session, knowledge_id); self.assertEqual(original.status_code, 200, original.text)
        fact = original.json()['evidence'][0]
        self.assertEqual(fact['text'], TEXT.decode())
        self.assertEqual(fact['reference']['source_sha256'], sha256(TEXT).hexdigest())
        prepared = self.prepare(session, [fact['evidence_id']]); self.assertEqual(prepared.status_code, 200, prepared.text)
        result = prepared.json()
        self.assertEqual((result['status'], result['job_id'], result['live_openkb']), ('HOST_PENDING_DRAFT_SAVED', None, 'NOT_RUN'))
        self.assertEqual(self.prepare(session, [fact['evidence_id']]).json(), result)
        state = self.client.get(self.root_api + '/state'); self.assertEqual(state.status_code, 200, state.text)
        state = state.json()
        self.assertEqual(state['openkb_preparation']['semantic_validation'], 'NOT_PERFORMED')
        self.assertEqual((state['ai_assessed_unit_count'], state['assessed_unit_count'], state['original_prepared_unit_count']), (0, 0, 1))
        self.assertEqual(state['assessments']['artifact_sha256'], result['assessment_sha256'])
        self.assertEqual(next(iter(state['assessments']['records'].values()))['decision'], 'PENDING_EVIDENCE')
        opinions = self.client.get(self.root_api + '/engineer-responses'); self.assertEqual(opinions.status_code, 200, opinions.text)
        view = opinions.json(); uid = session['unit_ids'][0]
        response = self.client.post(self.root_api + '/engineer-responses', json={
            'command_id': 'synthetic-user-opinion', 'expected_source_sha256': session['source_sha256'],
            'expected_baseline_sha256': session['baseline_sha256'], 'expected_response_revision': view['response_revision'],
            'decisions': [{'unit_id': uid, 'expected_unit_context_sha256': view['units'][uid]['unit_context_sha256'],
                'decision': 'PENDING_EVIDENCE', 'reason': 'Synthetic user reviewed original timing; whole scope still needs evidence.',
                'conditions': 'sample A only', 'evidence_requests': 'Controlled whole-unit validation'}]})
        self.assertEqual(response.status_code, 200, response.text)
        preview = self.client.get(self.root_api + '/versions/preview'); self.assertEqual(preview.status_code, 200, preview.text)
        body = {'command_id': 'native-draft', 'expected_source_sha256': preview.json()['source_sha256'],
            'expected_delivery_sha256': preview.json()['delivery_sha256'], 'expected_business_basis_sha256': preview.json()['business_basis_sha256'], 'mode': 'DRAFT'}
        saved = self.client.post(self.root_api + '/versions', json=body); self.assertEqual(saved.status_code, 200, saved.text)
        workbook = self.client.get(self.root_api + '/versions/1/export'); self.assertEqual(workbook.status_code, 200, workbook.text)
        book = load_workbook(BytesIO(workbook.content), data_only=False)
        self.assertIn('内部原件引用', book.sheetnames)
        cells = list(book['内部原件引用'].values)
        self.assertEqual(cells[1][4], sha256(TEXT).hexdigest())
        self.assertEqual(cells[1][7], TEXT.decode())
        self.assertEqual(json.loads(cells[1][6]), fact['source_ref'])
        self.assertEqual(cells[1][14], result['assessment_sha256'])
        self.assertEqual(cells[1][17], fact['evidence_id'])
        self.assertEqual(json.loads(cells[1][19]), fact['reference'])
        book.close()
        old = workbook.content
        db.bump_project_knowledge_base_version(self.project)
        self.assertEqual(self.client.get(self.root_api + '/versions/1/export').content, old)
        self.assertEqual(self.client.post(self.root_api + '/versions', json=body).json(), saved.json())
        self.assertEqual(self.prepare(session, [fact['evidence_id']]).status_code, 409)

    def test_status_and_subset_sessions_share_directory_generation_but_keep_selected_catalog(self):
        kid = self.upload().json()['documents'][0]['document_id']
        status = self.client.get(self.project_root + '/status').json()
        for selected in ([kid], []):
            response = self.session(knowledge_document_ids=selected)
            self.assertEqual(response.status_code, 200, response.text)
            session = response.json()
            self.assertEqual(session['generation'], status['generation'])
            self.assertEqual(session['directory_catalog_sha256'], status['catalog_sha256'])
            self.assertNotEqual(session['catalog_sha256'], status['catalog_sha256'])
            if not selected:
                self.assertEqual(self.original(session, kid).status_code, 409)

    def test_host_origin_peer_and_project_boundaries_precede_upload_or_read(self):
        self.assertEqual(self.client.post(self.project_root + '/upload', files={'files': ('test.txt', TEXT)}, headers={'Host': 'evil.example'}).status_code, 400)
        self.assertEqual(self.client.get(self.project_root + '/status', headers={'Origin': 'https://evil.example'}).status_code, 403)
        self.assertEqual(self.client.get('/api/projects/missing/openkb-evidence/status', headers={'X-User': 'owner'}).status_code, 404)
        foreign = self.root_api.replace(self.project, self.other)
        response = self.client.post(foreign + '/openkb-evidence/sessions', json={
            'expected_source_sha256': self.bundle['baseline'].source_sha256,
            'expected_baseline_sha256': self.bundle['baseline'].artifact_sha256, 'expected_knowledge_base_version': 0}, headers={'X-User': 'owner', 'X-Project': self.project})
        self.assertEqual(response.status_code, 404)
        self.app.state.runtime_settings = replace(self.settings, deployment_mode=DeploymentMode.CLOUD)
        self.assertEqual(self.client.get(self.project_root + '/status').status_code, 403)
        self.app.state.runtime_settings = replace(self.settings, deployment_mode=DeploymentMode.LOCAL)
        self.assertEqual(self.client.get(self.project_root + '/status').status_code, 403)  # testclient peer is not loopback.

    def test_upload_paths_duplicates_budget_and_partial_registration_are_explicit(self):
        self.assertEqual(self.upload('../escape.txt').status_code, 422)
        self.assertEqual(self.upload('/absolute.txt').status_code, 422)
        self.assertEqual(self.upload('empty.txt', b'').status_code, 413)
        self.assertEqual(self.upload('large.txt', b'x' * (10 * 1024 * 1024 + 1)).status_code, 413)
        self.assertEqual(len(db.list_knowledge_documents(self.project)), 0)
        self.assertEqual(self.upload().json()['registered_count'], 1)
        mixed = self.client.post(self.project_root + '/upload', files=[('files', ('reports/test.txt', TEXT)), ('files', ('new.txt', b'Synthetic second source.'))])
        self.assertEqual(mixed.status_code, 200, mixed.text)
        self.assertEqual((mixed.json()['registered_count'], mixed.json()['failed_count']), (1, 1))

    def test_no_caller_identity_path_or_decision_can_be_submitted_as_preparation(self):
        response = self.session(subject='admin', source_path='/etc/passwd')
        self.assertEqual(response.status_code, 422)
        self.upload(); session = self.session().json()
        path = self.root_api + f'/openkb-evidence/sessions/{session["session_id"]}/prepare-review'
        self.assertEqual(self.client.post(path, json={'expected_packet_sha256': session['packet_sha256'], 'evidence_ids': [], 'human_accepted': True}).status_code, 422)
        self.assertEqual(self.prepare(session, ['original-evidence-' + 'f' * 64]).status_code, 409)

    def test_source_byte_annotation_and_knowledge_change_invalidate_issued_session(self):
        kid = self.upload().json()['documents'][0]['document_id']
        session = self.session().json(); evidence_id = self.original(session, kid).json()['evidence'][0]['evidence_id']
        a, b = self.bundle['baseline'].records[:2]
        self.bundle['image_annotations'] = compile_source_item_annotation_updates(self.bundle['document_map'], self.bundle['baseline'],
            [{'source_item_id': a.source_record_id, 'context_source_ids': [b.source_record_id], 'reason': 'Explicit synthetic cross-section condition'}], table_previews=self.previews)
        self.assertEqual(self.prepare(session, [evidence_id]).status_code, 409)
        self.bundle['image_annotations'] = None
        path = Path(db.get_knowledge_document(kid)['file_path']); path.write_bytes(b'Changed unregistered bytes')
        response = self.original(session, kid)
        self.assertEqual(response.status_code, 409, response.text)

    def test_restored_service_reads_persisted_issued_references_without_model_execution(self):
        kid = self.upload().json()['documents'][0]['document_id']
        session = self.session().json(); fact = self.original(session, kid).json()['evidence'][0]
        restarted = TestClient(create_app(settings=self.settings))
        try:
            response = restarted.post(self.root_api + f'/openkb-evidence/sessions/{session["session_id"]}/prepare-review',
                json={'expected_packet_sha256': session['packet_sha256'], 'evidence_ids': [fact['evidence_id']]})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['live_openkb'], 'NOT_RUN')
        finally: restarted.close()

    def test_original_and_status_gets_are_no_store_and_stale_file_overlay_does_not_break_state(self):
        kid = self.upload().json()['documents'][0]['document_id']
        self.assertEqual(self.client.get(self.project_root + '/status').headers['cache-control'], 'no-store')
        session = self.session().json(); original = self.original(session, kid)
        self.assertEqual(original.headers['cache-control'], 'no-store')
        fact = original.json()['evidence'][0]; self.prepare(session, [fact['evidence_id']])
        Path(db.get_knowledge_document(kid)['file_path']).write_bytes(b'Unregistered replaced bytes')
        state = self.client.get(self.root_api + '/state')
        self.assertEqual(state.status_code, 200, state.text)
        self.assertIsNone(state.json()['assessments'])

    def test_wrapper_only_annotation_version_and_resealed_token_mapping_are_rejected(self):
        kid1 = self.upload().json()['documents'][0]['document_id']
        kid2 = self.upload('second.txt', TEXT + b'Second source identity.').json()['documents'][0]['document_id']
        session = self.session().json()
        first = self.original(session, kid1).json()['evidence'][0]
        second = self.original(session, kid2).json()['evidence'][0]
        name = 'sessions/' + session['session_id'] + '/issued.json'
        issued = evidence._get(self.project, self.document, name)
        issued['references'][first['evidence_id']] = second['reference']
        evidence._put(self.project, self.document, name, issued)
        self.assertEqual(self.prepare(session, [first['evidence_id']]).status_code, 409)
        a = self.bundle['baseline'].records[0]
        self.bundle['image_annotations'] = compile_source_item_annotation_updates(self.bundle['document_map'], self.bundle['baseline'],
            [{'source_item_id': a.source_record_id, 'topics': [], 'reason': 'Synthetic annotation wrapper'}], table_previews=self.previews)
        session = self.session().json()
        from reqpilotx.application.agent_workflow.source_item_visual import _seal_annotations
        self.bundle['image_annotations'] = _seal_annotations({**self.bundle['image_annotations'], 'wrapper_review_version': 2})
        self.assertEqual(self.prepare(session, []).status_code, 409)

    def test_preparation_order_uses_instants_not_offset_strings_and_private_budget_is_checked(self):
        self.upload(); session = self.session().json(); self.prepare(session, [])
        head = evidence._get(self.project, self.document, 'head.json')
        head['prepared_at'] = '2026-10-02T09:30:00+00:00'
        evidence._put(self.project, self.document, 'head.json', head)
        self.assertIsNone(evidence.current_preparation(self.project, self.document, self.bundle['baseline'], '',
            newer_than=datetime.fromisoformat('2026-10-02T08:45:00-01:00')))
        with patch.object(evidence, '_write_private') as writer:
            with self.assertRaises(evidence.ApplicationError) as error:
                evidence._write_bounded_private(self.container, project_id=self.project, document_id=self.document,
                    namespace='working', stage='large-test', value={'text': 'x' * (32 * 1024 * 1024)})
            self.assertEqual(error.exception.status_code, 413)
            writer.assert_not_called()
        with patch.object(self.container.object_store, 'put') as writer:
            with self.assertRaises(evidence.ApplicationError) as error:
                evidence._put(self.project, self.document, 'oversize-issued-test.json', {'references': {'text': 'x' * (32 * 1024 * 1024)}})
            self.assertEqual(error.exception.status_code, 413)
            writer.assert_not_called()
