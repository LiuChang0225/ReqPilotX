"""A same-version snapshot is insufficient if its saved links disappear in Excel."""
from __future__ import annotations

# ruff: noqa: E402
from .runtime_isolation import activate_test_runtime
activate_test_runtime()

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook
from app.services.exporter import export_compliance_matrix
from .test_response_matrix import response_snapshot


class ContextLinkExportTests(unittest.TestCase):
    def test_saved_context_is_visible_and_navigable_without_new_source_rows(self):
        data = response_snapshot()
        parent = next(n for n in data['nodes'] if n['node_id'] == 'record')
        context = next(n for n in data['nodes'] if n['node_id'] == 'context')
        parent['source_item'].update(source_item_id='stable-parent', context_source_ids=['stable-context'])
        context['source_item'].update(source_item_id='stable-context', context_source_ids=[])
        before = deepcopy(data)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            workbook = load_workbook(path)
            sheet = workbook['ComplianceMatrix']
            source_row = next(r for r in sheet.iter_rows(min_row=4) if r[0].value == 'SRC-001')
            context_row = next(r for r in sheet.iter_rows(min_row=4) if r[0].value == 'SRC-002')
            self.assertIn('关联原文与条件：SRC-002', source_row[10].value)
            self.assertIn(f'主表第{context_row[0].row}行', source_row[10].value)
            self.assertEqual(source_row[10].hyperlink.location, f"'ComplianceMatrix'!A{context_row[0].row}")
            self.assertNotIn('关联原文与条件', str(source_row[3].value))
            self.assertNotIn('关联原文与条件', str(context_row[10].value))
            self.assertEqual(sheet.max_row - 3, sum(n['matrix_visibility'] != 'SYSTEM_NON_BUSINESS' for n in data['nodes']))
            workbook.close()
        self.assertEqual(data, before)

    def test_export_rejects_an_unresolvable_context_instead_of_silently_hiding_it(self):
        data = response_snapshot()
        parent = next(n for n in data['nodes'] if n['node_id'] == 'record')
        parent['source_item'].update(source_item_id='stable-parent', context_source_ids=['missing-source'])
        with TemporaryDirectory() as folder, self.assertRaisesRegex(ValueError, 'Linked source context'):
            export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
