"""Synthetic adapter tests; UI acceptance stays in its separate case log."""
from copy import deepcopy
import hashlib
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import zipfile

from server import ReviewStore, canonical
from react_contract import api_payload, bootstrap, export_xlsx, matrix, prefix, source_crop


class ReactContractTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        source = b"synthetic evidence identity"
        (self.root/"source.pdf").write_bytes(source)
        def item(item_id, section, order, classes, tags, derived=None):
            return {"id":item_id,"sectionId":section,"label":item_id,"class":"C","text":"=literal text",
                    "sourceRefs":[{"page":1,"rects":[[10,20,100,30]],"precision":"line"}],
                    "derived":derived or [], "reviewMetadata":{"number":item_id,"title":item_id,"sourceOrder":order,
                      "classifications":classes,"topics":tags,"requirementness":"REQUIREMENT","contentRole":"","reason":"Synthetic"}}
        child=item("source/d1","4.1",20,["PRODUCT_SYSTEM.CONSTRAINT"],["REGULATORY"])
        child["parentId"]="source"
        self.bundle={"schema":"reqpilot-target-review/v1","source":{"sha256":hashlib.sha256(source).hexdigest(),"name":"Synthetic.pdf","pageCount":1},
                     "outline":[{"id":"4","parentId":None,"page":1,"title":"Chapter","sourceOrder":0},
                                {"id":"4.1","parentId":"4","page":1,"title":"First","sourceOrder":10},
                                {"id":"4.5","parentId":"4","page":1,"title":"Later","sourceOrder":50}],
                     "items":[item("source","4.1",20,["PRODUCT_SYSTEM.CONSTRAINT"],["REGULATORY"],[child]),
                              item("later-1","4.5",60,["PROJECT"],["LOGISTICS_PACKAGING"]),
                              item("later-2","4.5",70,["PROCESS"],["LOGISTICS_PACKAGING"])],
                     "pages":[{"number":1,"width":595,"height":842}],
                     "ui":{"caseId":"synthetic","initialNodeId":"source","topics":{"REGULATORY":"法规","LOGISTICS_PACKAGING":"物流包装"},"scope":"Synthetic scope"}}
        self.bundle["revision"]=hashlib.sha256(canonical(self.bundle)).hexdigest()
        (self.root/"review-bundle.json").write_bytes(canonical(self.bundle))
        self.store=ReviewStore(self.root,reviewer_kind="qa")
    def tearDown(self):
        self.temp.cleanup()
    def test_source_order_and_one_to_one_identity(self):
        snapshot=matrix(self.store)
        index={node["node_id"]:node for node in snapshot["nodes"]}
        self.assertNotIn("4.1",index)
        self.assertEqual(index["source"]["structure_node_id"],"4.1")
        self.assertEqual(index["4"]["child_ids"],["source","4.5"])
        self.assertEqual(index["source"]["child_ids"],["source/d1"])
        self.assertEqual(snapshot["counts"]["decided_nodes"],0)

    def test_full_candidate_missing_section_order_does_not_precede_earlier_clause(self):
        # A legacy section used to receive array index 2, sorting before the
        # source clause's geometric position 1_002_000.
        bundle = self.store.bundle
        bundle['outline'][2].pop('sourceOrder')
        bundle['lines'] = {'later-heading': {'page': 1, 'bbox': [10, 50, 100, 60], 'text': 'Later'}}
        bundle['outline'][2]['headingLineIds'] = ['later-heading']
        bundle['items'][0]['reviewMetadata']['sourceOrder'] = 1_002_000
        bundle['items'][1]['reviewMetadata']['sourceOrder'] = 1_006_000
        bundle['items'][2]['reviewMetadata']['sourceOrder'] = 1_007_000
        nodes = {n['node_id']: n for n in matrix(self.store)['nodes']}
        self.assertEqual(nodes['4']['child_ids'], ['source', '4.5'])
        self.assertEqual(nodes['4.5']['source_order'], 1_005_000)

    def test_later_appendix_with_no_heading_stays_after_main_chapter(self):
        bundle = self.store.bundle
        bundle['outline'][0]['sourceOrder'] = 1_001_000
        bundle['outline'].append({'id': 'appendix', 'page': 2, 'parentId': None, 'title': 'Appendix'})
        snapshot = matrix(self.store)
        self.assertEqual(snapshot['root_node_ids'], ['4', 'appendix'])

    def test_missing_order_uses_source_heading_not_numeric_title(self):
        bundle = self.store.bundle
        bundle['lines'] = {'a': {'page': 1, 'bbox': [10, 80, 100, 90], 'text': '9'},
                           'b': {'page': 1, 'bbox': [10, 100, 100, 110], 'text': '2'}}
        bundle['outline'].extend([
            {'id': 'two', 'parentId': '4', 'page': 1, 'title': '2', 'headingLineIds': ['b']},
            {'id': 'nine', 'parentId': '4', 'page': 1, 'title': '9', 'headingLineIds': ['a']}])
        nodes = {n['node_id']: n for n in matrix(self.store)['nodes']}
        self.assertLess(nodes['nine']['source_order'], nodes['two']['source_order'])

    def test_trailing_note_cannot_become_parent_of_earlier_subclause(self):
        bundle = self.store.bundle
        bundle['outline'].append({'id': '4.1.1', 'parentId': '4.1', 'page': 1, 'title': 'Earlier subclause', 'sourceOrder': 15})
        bundle['items'][1]['sectionId'] = '4.1.1'
        bundle['items'][1]['reviewMetadata']['sourceOrder'] = 16
        # The single item in 4.1 is at 20, after its subclause at 15.
        nodes = {n['node_id']: n for n in matrix(self.store)['nodes']}
        self.assertIn('4.1', nodes)
        self.assertEqual(nodes['4.1']['child_ids'], ['later-1', 'source'])
        self.assertEqual(nodes['source']['parent_id'], '4.1')

    def test_explicit_whole_table_source_precedes_its_rows(self):
        bundle = self.store.bundle
        wrapper = bundle['items'][2]
        wrapper.update(sectionId='4', matrixTables=['logical'])
        bundle['reviewTables'] = [{'id': 'logical', 'sectionId': '4.1',
            'segments': [{'id': 'physical', 'page': 1, 'cells': []}],
            'sourceRefs': [{'page': 1}], 'records': [{'record_id': 'source'}]}]
        nodes = {n['node_id']: n for n in matrix(self.store)['nodes']}
        self.assertEqual(nodes['4.1']['parent_id'], wrapper['id'])
        self.assertEqual(nodes[wrapper['id']]['child_ids'], ['4.1'])
        self.assertEqual(nodes['source']['parent_id'], '4.1')

    def test_quality_classification_survives_deferred_applicability(self):
        source = self.store.bundle["items"][1]
        source["reviewMetadata"].update(classifications=["PRODUCT_SYSTEM.QUALITY"],
            requirementness="REQUIREMENT", applicabilityStatus="DEFERRED", analysisEligibility="NOT_ELIGIBLE")
        projected = next(n for n in matrix(self.store)["nodes"] if n["node_id"] == source["id"])
        self.assertEqual(projected["classification_paths"], ["PRODUCT_SYSTEM.QUALITY"])
        self.assertEqual(projected["candidate_review"]["requirementness"], "REQUIREMENT")
        self.assertEqual(projected["analysis_eligibility"], "NOT_ELIGIBLE")
        self.assertIsNone(projected["ai_assessment"])

    def test_original_subclauses_stay_under_the_original_parent_not_as_derivatives(self):
        self.store.bundle["outline"].append({"id":"4.1.1","parentId":"4.1","page":1,"title":"Child clause","sourceOrder":30})
        child = self.store.bundle["items"][1]
        child["sectionId"] = "4.1.1"
        projected = matrix(self.store)
        nodes = {n["node_id"]: n for n in projected["nodes"]}
        self.assertNotIn("4.1", nodes)
        self.assertEqual(nodes[child["id"]]["parent_id"], "source")
        self.assertEqual(nodes[child["id"]]["source_kind"], "SOURCE_ITEM")
        self.assertEqual(nodes["source/d1"]["source_kind"], "DERIVED_REQUIREMENT")
        _, _, requirement_set = __import__("react_contract").identity(self.store)
        path = prefix(self.store) + f"/requirement-sets/{requirement_set}/matrix/nodes/{child['id']}/detail"
        detail = api_payload(self.store, path)
        self.assertIn("source", [n["node_id"] for n in detail["path"]])
    def test_feedback_does_not_accept_children_or_change_source(self):
        before=deepcopy(self.store.bundle)
        self.store.append({"mutationId":"synthetic-confirm-0001","expectedRevision":0,"sourceRevision":self.store.source_revision,
                           "targetId":"source","action":"confirm","note":"Synthetic QA"})
        nodes={node["node_id"]:node for node in matrix(self.store)["nodes"]}
        self.assertEqual(nodes["source"]["candidate_review"]["status"],"reviewed")
        self.assertEqual(nodes["source/d1"]["candidate_review"]["status"],"unreviewed")
        self.assertEqual(before,self.store.bundle)
    def test_endpoints_preserve_source_and_deny_unknown_scope(self):
        data=bootstrap(self.store)
        self.assertEqual(data["document"]["source_sha256"], self.store.bundle["source"]["sha256"])
        page = api_payload(self.store, prefix(self.store) + "/extraction/source-pages/1")
        self.assertEqual(page["source_identity"]["source_sha256"], data["document"]["source_sha256"])
        reqset=data["requirementSets"]["active_requirement_set_id"]
        base=prefix(self.store)+f"/requirement-sets/{reqset}/matrix"
        detail=api_payload(self.store,base+"/nodes/source%2Fd1/detail")
        self.assertEqual(detail["parent_clause"]["node_id"],"source")
        self.assertIsNone(api_payload(self.store,base+"/nodes/unknown/detail"))
        self.assertIsNone(api_payload(self.store,"/api/projects/another/sor-documents/other"))

    def test_review_version_label_comes_from_the_frozen_artifact(self):
        self.store.bundle["ui"].update(reviewVersion=6, versionLabel="审阅修订06 · 当前范围")
        version = bootstrap(self.store)["requirementSets"]["requirement_sets"][0]
        self.assertEqual(version["version"], 6)
        self.assertEqual(version["version_label"], "审阅修订06 · 当前范围")

    def test_table_records_use_the_shared_table_and_literal_cell_heading(self):
        source = self.store.bundle["items"][0]
        source["sourcePresentation"] = {"number":"1", "titleRole":"SOURCE_HEADING",
            "titleRanges":[{"tableId":"physical", "cellId":"caption", "start":0, "end":4}]}
        segment = {"id":"physical", "type":"TABLE", "page":1, "page_number":1, "row_count":2, "column_count":2,
            "header_rows":[0], "cells":[{"id":"caption", "row":1, "column":0, "text":"原表项目", "page":1}]}
        self.store.bundle["reviewTables"] = [{"id":"logical", "sectionId":"4.1", "segments":[segment],
            "sourceRefs":[{"page":1,"rects":[[10,10,200,100]]}], "records":[{"record_id":"source", "record_kind":"ROW"}]}]
        snap = matrix(self.store)
        nodes = {n["node_id"]:n for n in snap["nodes"]}
        self.assertEqual(nodes["4.1"]["display_role"], "TABLE")
        self.assertEqual(nodes["source"]["display_title"], "原表项目")
        self.assertEqual(nodes["source"]["table_record_id"], "source")
        self.assertEqual(nodes["source/d1"]["table_record_id"], "source")
        self.assertEqual(nodes["source"]["content_blocks"], [])
        self.assertEqual(snap["logical_tables"][0]["record_count"], 1)
        self.assertEqual(snap["logical_tables"][0]["projection"]["records"][0]["record_id"], "source")

    def test_crop_is_bounded_to_the_source_page(self):
        import fitz
        document = fitz.open()
        page = document.new_page(width=595,height=842)
        page.insert_text((20,30), "original figure")
        document.save(self.root/"source.pdf")
        data = source_crop(self.store,1,[10,10,140,40])
        self.assertTrue(data.startswith(b"\x89PNG"))
        for page,box in [(0,[10,10,20,20]),(1,[-1,0,20,20]),(1,[0,0,600,20]),(1,[0,0,float('nan'),20])]:
            with self.assertRaises(ValueError):
                source_crop(self.store,page,box)
    def test_export_uses_exact_matched_ids_and_safe_text(self):
        body=export_xlsx(self.store,{"sourceRevision":self.store.source_revision,"expectedRevision":0,"nodeIds":["source/d1"]})
        with zipfile.ZipFile(BytesIO(body)) as archive:
            sheet=archive.read("xl/worksheets/sheet1.xml").decode()
            strings=archive.read("xl/sharedStrings.xml").decode()
            self.assertIn('autoFilter ref="A1:M2"',sheet)
            self.assertNotIn("<f>",sheet)
            self.assertIn("=literal text",strings)
            self.assertIn("需求大类",strings)
            self.assertIn("系统需求类型",strings)
            self.assertIn("主题标签",strings)
        for ids in [["unknown"],["source/d1","source/d1"],[{}]]:
            with self.assertRaises(ValueError):
                export_xlsx(self.store,{"sourceRevision":self.store.source_revision,"expectedRevision":0,"nodeIds":ids})
        with self.assertRaises(ValueError):
            export_xlsx(self.store,{"sourceRevision":self.store.source_revision,"expectedRevision":True,"nodeIds":[]})

    def test_source_heading_cannot_be_replaced_by_semantic_caption(self):
        self.store.bundle["lines"] = {"heading": {"text": "7.3 原文标题", "page": 1}}
        source = self.store.bundle["items"][0]
        source["reviewMetadata"]["title"] = "无需分析的状态"
        source["sourcePresentation"] = {"number": "7.3", "titleRole": "SOURCE_HEADING",
            "titleRanges": [{"lineId":"heading", "start":4, "end":8}]}
        projected = next(n for n in matrix(self.store)["nodes"] if n["node_id"] == "source")
        self.assertEqual(projected["display_title"], "原文标题")
        self.assertEqual(projected["source_number"], "7.3")
        unheaded = self.store.bundle["items"][1]
        unheaded["reviewMetadata"]["title"] = "Agent invented heading"
        projected = next(n for n in matrix(self.store)["nodes"] if n["node_id"] == unheaded["id"])
        self.assertEqual(projected["display_title"], unheaded["text"])
        unheaded["sourcePresentation"] = {"titleRole":"AGENT_ASSIGNED", "numberRole":"AGENT_ASSIGNED",
            "number":"7.3(1)", "title":"Agent拆分项", "titleRanges":[]}
        projected = next(n for n in matrix(self.store)["nodes"] if n["node_id"] == unheaded["id"])
        self.assertEqual(projected["display_title"], "Agent拆分项")
        self.assertEqual(projected["source_title_role"], "AGENT_ASSIGNED")
        self.assertEqual(projected["source_number_role"], "AGENT_ASSIGNED")

    def test_excerpts_are_source_slices_including_original_typo(self):
        self.store.bundle["lines"] = {"line": {"text": "背景。试验后不得得出现卡滞。另一要求", "page": 3}}
        child = self.store.bundle["items"][0]["derived"][0]
        child["sourceExcerpts"] = [{"id":"quote", "role":"DIRECT", "text":"Agent paraphrase is ignored",
            "ranges":[{"lineId":"line", "start":3, "end":14}]}]
        projected = next(n for n in matrix(self.store)["nodes"] if n["node_id"] == child["id"])
        self.assertEqual(projected["source_excerpts"], [{"id":"quote", "role":"DIRECT", "text":"试验后不得得出现卡滞。", "pages":[3]}])
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(export_xlsx(self.store, {"sourceRevision":self.store.source_revision,
            "expectedRevision":0, "nodeIds":[child["id"]]})))
        self.assertEqual(book["原文核对清单"]["M2"].value, "试验后不得得出现卡滞。")
        child["sourceExcerpts"][0]["ranges"][0]["end"] = 1000
        with self.assertRaises(ValueError):
            matrix(self.store)

    def test_source_table_export_preserves_columns_merges_and_single_item(self):
        from openpyxl import load_workbook
        source = self.store.bundle["items"][1]
        source["displayBlocks"] = [{"id":"synthetic-table", "type":"TABLE", "projection_mode":"SOURCE_TABLE",
            "page_number":1, "row_count":3, "column_count":4, "header_rows":[0], "cells":[
                {"row":0,"column":i,"text":title} for i,title in enumerate(["No", "Group", "Kind", "Test"])
            ] + [{"row":1,"column":0,"text":"1"}, {"row":2,"column":0,"text":"2"},
                 {"row":1,"column":1,"row_span":2,"text":"Shared"},
                 {"row":1,"column":2,"column_span":2,"text":"=literal"},
                 {"row":2,"column":2,"text":"Type"}, {"row":2,"column":3,"text":"Last column"}]}]
        snapshot = matrix(self.store)
        self.assertEqual(snapshot["logical_tables"][0]["projection_mode"], "SOURCE_TABLE")
        self.assertEqual(snapshot["logical_tables"][0]["record_count"], 2)
        body = export_xlsx(self.store, {"sourceRevision":self.store.source_revision,"expectedRevision":0,"nodeIds":["later-1"]})
        workbook = load_workbook(BytesIO(body))
        self.assertEqual(workbook["原文核对清单"].max_row, 2)
        self.assertNotIn("参考依据", workbook.sheetnames)
        detail = workbook["原表-1"]
        self.assertEqual(detail["D6"].value, "Last column")
        self.assertEqual({str(merge) for merge in detail.merged_cells.ranges}, {"B5:B6", "C5:D5"})
        self.assertEqual(detail["C5"].data_type, "s")
        self.assertEqual(detail.freeze_panes, "A5")
        filtered = load_workbook(BytesIO(export_xlsx(self.store, {
            "sourceRevision":self.store.source_revision,"expectedRevision":0,"nodeIds":["source/d1"]})))
        self.assertNotIn("原表-1", filtered.sheetnames)


if __name__ == "__main__":
    unittest.main()
