from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import hashlib
import shutil
import tempfile
import unittest
import zipfile
from dataclasses import FrozenInstanceError
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from PIL import Image as PILImage
from pypdf import PdfReader, PdfWriter
from pypdf.constants import UserAccessPermissions

from app.adapters.documents import preflight as preflight_module
from app.adapters.documents.preflight import LocalDocumentPreflight, preflight_document
from app.domain.extraction.preflight import (
    AccessStatus,
    AuthenticityStatus,
    DiagnosticCode,
    DocumentFormat,
    IntegrityStatus,
    SupportStatus,
)
from app.ports.document_preflight import DocumentPreflightPort


class DocumentPreflightTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)

    def test_original_byte_identity_is_stable_and_exact(self) -> None:
        path = self.root / "requirements.txt"
        payload = "控制器应在 100 ms 内启动。\n".encode("utf-8")
        path.write_bytes(payload)

        first = preflight_document(path)
        second = preflight_document(path)

        self.assertEqual(first.identity, second.identity)
        self.assertEqual(first.source_size, len(payload))
        self.assertEqual(first.source_sha256, hashlib.sha256(payload).hexdigest())
        self.assertTrue(first.processable)
        with self.assertRaises(FrozenInstanceError):
            first.identity.source_size = 0  # type: ignore[misc]

    def test_domain_contract_rejects_untyped_enum_values(self) -> None:
        report = preflight_document(self._write_pdf("requirements.pdf"))

        with self.assertRaisesRegex(TypeError, "detected_format"):
            type(report.identity)(
                source_name="requirements.pdf",
                source_sha256=report.source_sha256,
                source_size=report.source_size,
                declared_extension=".pdf",
                detected_format="pdf",  # type: ignore[arg-type]
                mime_type="application/pdf",
            )

    def test_first_release_supported_formats_are_identified_by_signature(self) -> None:
        fixtures = {
            DocumentFormat.PDF: self._write_pdf("requirements.pdf"),
            DocumentFormat.DOCX: self._write_docx("requirements.docx"),
            DocumentFormat.XLSX: self._write_xlsx("requirements.xlsx"),
        }
        text_path = self.root / "requirements.txt"
        text_path.write_text("The controller shall start.\n", encoding="utf-8")
        fixtures[DocumentFormat.TXT] = text_path

        for expected_format, path in fixtures.items():
            with self.subTest(expected_format=expected_format):
                report = preflight_document(path)
                self.assertEqual(report.identity.detected_format, expected_format)
                self.assertEqual(report.support_status, SupportStatus.SUPPORTED)
                self.assertEqual(report.integrity_status, IntegrityStatus.VALID)
                self.assertTrue(report.processable)

    def test_utf16_text_is_strictly_validated_without_treating_encoding_bytes_as_binary(
        self,
    ) -> None:
        path = self.root / "utf16-requirements.txt"
        payload = "控制器应在 100 ms 内启动。\n".encode("utf-16")
        path.write_bytes(payload)

        report = preflight_document(path)

        self.assertEqual(report.identity.detected_format, DocumentFormat.TXT)
        self.assertEqual(report.integrity_status, IntegrityStatus.VALID)
        self.assertTrue(report.processable)
        self.assertEqual(report.source_sha256, hashlib.sha256(payload).hexdigest())

    def test_xlsb_and_legacy_doc_are_explicitly_unsupported(self) -> None:
        xlsb = self.root / "requirements.xlsb"
        with zipfile.ZipFile(xlsb, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types />")
            archive.writestr("xl/workbook.bin", b"binary workbook")
        legacy_doc = self.root / "requirements.doc"
        legacy_doc.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512)

        xlsb_report = preflight_document(xlsb)
        doc_report = preflight_document(legacy_doc)

        self.assertEqual(xlsb_report.identity.detected_format, DocumentFormat.XLSB)
        self.assertEqual(doc_report.identity.detected_format, DocumentFormat.LEGACY_DOC)
        self.assertTrue(xlsb_report.controlled_conversion_eligible)
        self.assertFalse(doc_report.controlled_conversion_eligible)
        for report in (xlsb_report, doc_report):
            self.assertEqual(report.authenticity_status, AuthenticityStatus.MATCHED)
            self.assertEqual(report.support_status, SupportStatus.UNSUPPORTED)
            self.assertFalse(report.processable)
            self.assertIn(
                DiagnosticCode.UNSUPPORTED_FORMAT,
                {item.code for item in report.diagnostics},
            )

    def test_tif_is_identified_as_known_attachment_only_format(self) -> None:
        path = self.root / "drawing.tif"
        PILImage.new("1", (2, 2), color=1).save(path, format="TIFF")

        report = preflight_document(path)

        self.assertEqual(report.identity.detected_format, DocumentFormat.TIF)
        self.assertEqual(report.authenticity_status, AuthenticityStatus.MATCHED)
        self.assertEqual(report.support_status, SupportStatus.UNSUPPORTED)
        self.assertEqual(report.integrity_status, IntegrityStatus.VALID)
        self.assertFalse(report.processable)
        self.assertTrue(report.attachment_acceptable)
        self.assertEqual(report.page_count, 1)
        self.assertIn(
            DiagnosticCode.UNSUPPORTED_FORMAT,
            {item.code for item in report.diagnostics},
        )

    def test_corrupt_tif_header_is_not_attachment_acceptable(self) -> None:
        path = self.root / "broken.tif"
        path.write_bytes(b"II*\x00" + b"\x00" * 128)

        report = preflight_document(path)

        self.assertEqual(report.integrity_status, IntegrityStatus.CORRUPT)
        self.assertFalse(report.processable)
        self.assertFalse(report.attachment_acceptable)
        self.assertIn(
            DiagnosticCode.CORRUPT_DOCUMENT,
            {item.code for item in report.diagnostics},
        )

    def test_extension_signature_mismatch_is_diagnostic_and_fail_closed(self) -> None:
        source = self._write_docx("source.docx")
        disguised = self.root / "disguised.pdf"
        shutil.copyfile(source, disguised)

        report = preflight_document(disguised)

        self.assertEqual(report.identity.detected_format, DocumentFormat.DOCX)
        self.assertEqual(report.authenticity_status, AuthenticityStatus.MISMATCHED)
        self.assertFalse(report.processable)
        self.assertIn(
            DiagnosticCode.EXTENSION_SIGNATURE_MISMATCH,
            {item.code for item in report.diagnostics},
        )

    def test_unknown_binary_never_falls_back_to_text(self) -> None:
        path = self.root / "payload.txt"
        path.write_bytes(b"\x00\x01\xffrequirements\x00\x80")

        report = preflight_document(path)

        self.assertEqual(report.identity.detected_format, DocumentFormat.UNKNOWN)
        self.assertEqual(report.integrity_status, IntegrityStatus.UNVERIFIED)
        self.assertFalse(report.processable)
        self.assertIn(
            DiagnosticCode.UNKNOWN_BINARY_FORMAT,
            {item.code for item in report.diagnostics},
        )

    def test_corrupt_supported_documents_have_explicit_integrity_status(self) -> None:
        corrupt_pdf = self.root / "broken.pdf"
        corrupt_pdf.write_bytes(b"%PDF-1.7\nnot a real PDF")
        corrupt_docx = self.root / "broken.docx"
        with zipfile.ZipFile(corrupt_docx, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types />")
            archive.writestr("word/document.xml", "<not-valid")
        corrupt_xlsx = self.root / "broken.xlsx"
        with zipfile.ZipFile(corrupt_xlsx, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types />")
            archive.writestr("xl/workbook.xml", "<not-valid")

        for path in (corrupt_pdf, corrupt_docx, corrupt_xlsx):
            with self.subTest(path=path.name):
                report = preflight_document(path)
                self.assertEqual(report.integrity_status, IntegrityStatus.CORRUPT)
                self.assertFalse(report.processable)
                self.assertIn(
                    DiagnosticCode.CORRUPT_DOCUMENT,
                    {item.code for item in report.diagnostics},
                )

    def test_password_protected_pdf_is_blocked_explicitly(self) -> None:
        source = self._write_pdf("source.pdf")
        encrypted = self.root / "password.pdf"
        self._encrypt_pdf(
            source, encrypted, user_password="secret", owner_password="owner"
        )

        report = preflight_document(encrypted)

        self.assertEqual(report.integrity_status, IntegrityStatus.VALID)
        self.assertEqual(report.access_status, AccessStatus.PASSWORD_REQUIRED)
        self.assertTrue(report.password_required)
        self.assertFalse(report.processable)
        self.assertIn(
            DiagnosticCode.PASSWORD_REQUIRED, {item.code for item in report.diagnostics}
        )

    def test_empty_password_and_permission_restrictions_do_not_block_pdf(self) -> None:
        source = self._write_pdf("source.pdf")
        restricted = self.root / "restricted.pdf"
        self._encrypt_pdf(
            source,
            restricted,
            user_password="",
            owner_password="owner",
            permissions=UserAccessPermissions.EXTRACT,
        )

        report = preflight_document(restricted)

        self.assertEqual(
            report.access_status, AccessStatus.READABLE_WITH_EMPTY_PASSWORD
        )
        self.assertFalse(report.password_required)
        self.assertTrue(report.processable)
        self.assertEqual(report.page_count, 1)
        self.assertIn(
            DiagnosticCode.PDF_SECURITY_RESTRICTIONS,
            {item.code for item in report.diagnostics},
        )

    def test_empty_file_and_damaged_zip_container_are_not_success(self) -> None:
        empty = self.root / "empty.txt"
        empty.write_bytes(b"")
        broken_zip = self.root / "broken.docx"
        broken_zip.write_bytes(b"PK\x03\x04not a zip archive")

        empty_report = preflight_document(empty)
        zip_report = preflight_document(broken_zip)

        self.assertEqual(empty_report.integrity_status, IntegrityStatus.EMPTY)
        self.assertEqual(zip_report.integrity_status, IntegrityStatus.CORRUPT)
        self.assertFalse(empty_report.processable)
        self.assertFalse(zip_report.processable)

    def test_office_container_and_pdf_page_budgets_fail_closed(self) -> None:
        docx = self._write_docx("budget.docx")
        with patch.object(preflight_module, "_MAX_ZIP_MEMBERS", 1):
            docx_report = preflight_document(docx)

        self.assertEqual(docx_report.identity.detected_format, DocumentFormat.DOCX)
        self.assertEqual(docx_report.integrity_status, IntegrityStatus.UNVERIFIED)
        self.assertFalse(docx_report.processable)
        self.assertIn(
            DiagnosticCode.RESOURCE_LIMIT_EXCEEDED,
            {item.code for item in docx_report.diagnostics},
        )

        pdf = self._write_pdf("budget.pdf")
        with patch.object(preflight_module, "_MAX_PDF_PAGES", 0):
            pdf_report = preflight_document(pdf)

        self.assertEqual(pdf_report.identity.detected_format, DocumentFormat.PDF)
        self.assertEqual(pdf_report.integrity_status, IntegrityStatus.UNVERIFIED)
        self.assertFalse(pdf_report.processable)
        self.assertIn(
            DiagnosticCode.RESOURCE_LIMIT_EXCEEDED,
            {item.code for item in pdf_report.diagnostics},
        )

    def test_adapter_satisfies_read_only_port(self) -> None:
        adapter = LocalDocumentPreflight()

        self.assertIsInstance(adapter, DocumentPreflightPort)

    def _write_pdf(self, name: str) -> Path:
        path = self.root / name
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with path.open("wb") as handle:
            writer.write(handle)
        return path

    def _write_docx(self, name: str) -> Path:
        path = self.root / name
        document = Document()
        document.add_paragraph("The controller shall start within 100 ms.")
        document.save(path)
        return path

    def _write_xlsx(self, name: str) -> Path:
        path = self.root / name
        workbook = Workbook()
        workbook.active.append(["ID", "Requirement"])
        workbook.active.append(["SYS-1", "The controller shall start."])
        workbook.save(path)
        workbook.close()
        return path

    @staticmethod
    def _encrypt_pdf(
        source: Path,
        destination: Path,
        *,
        user_password: str,
        owner_password: str,
        permissions: UserAccessPermissions = UserAccessPermissions.PRINT,
    ) -> None:
        reader = PdfReader(str(source))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.encrypt(
            user_password=user_password,
            owner_password=owner_password,
            permissions_flag=permissions,
        )
        with destination.open("wb") as handle:
            writer.write(handle)


if __name__ == "__main__":
    unittest.main()
