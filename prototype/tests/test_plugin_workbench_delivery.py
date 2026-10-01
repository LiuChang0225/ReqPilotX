from .runtime_isolation import activate_test_runtime
activate_test_runtime()
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

from . import test_plugin_workbench as fixtures
from reqpilotx_server.application.agent_workflow import plugin_workbench as work
from reqpilotx_server.application.agent_workflow import plugin_workbench_delivery as api
from reqpilotx_server.application.errors import ApplicationError


class PluginWorkbenchDeliveryTests(unittest.TestCase):
    setUp = fixtures.PluginWorkbenchTests.setUp
    tearDown = fixtures.PluginWorkbenchTests.tearDown
    sink = fixtures.PluginWorkbenchTests.sink
    payload = fixtures.PluginWorkbenchTests.payload
    sync = fixtures.PluginWorkbenchTests.sync
    mapping = fixtures.PluginWorkbenchTests.mapping

    def start(self):
        state = self.sync(self.artifacts, require_map_confirmation=False)
        return state, self.responses()

    def responses(self, revision=None):
        return api.get_plugin_engineer_responses('p', 'd', self.session, self.sink(), revision)

    def response_payload(self, view, **changes):
        fields = dict(command_id='response-one', expected_source_sha256=self.sha,
            expected_revision_id=view['revision_id'], expected_version=view['version'],
            expected_response_revision=view['response_revision'], expected_baseline_sha256=view['baseline_sha256'],
            decisions=[dict(unit_id=uid, expected_unit_context_sha256=unit['unit_context_sha256'],
                decision='COMPLIANT', reason='Explicit synthetic engineer opinion') for uid, unit in view['units'].items()])
        fields.update(changes)
        return api.PluginEngineerResponse(**fields)

    def respond(self, view, **changes):
        return api.post_plugin_engineer_responses('p', 'd', self.session, self.response_payload(view, **changes), self.sink())

    def preview(self):
        return api.preview_plugin_delivery_version('p', 'd', self.session, self.sink())

    def delivery_payload(self, preview, **changes):
        fields = dict(command_id='delivery-one', expected_source_sha256=self.sha, expected_revision_id=preview['revision_id'],
            expected_version=preview['version'], expected_response_revision=preview['response_revision'],
            expected_delivery_sha256=preview['delivery_sha256'])
        fields.update(changes)
        return api.PluginDeliveryVersion(**fields)

    def save(self, preview, **changes):
        return api.save_plugin_delivery_version('p', 'd', self.session, self.delivery_payload(preview, **changes), self.sink())

    def test_response_creates_immutable_result_and_reader_export_use_that_same_snapshot(self):
        state, before = self.start()
        old = work.get_plugin_snapshot('p', 'd', self.session, self.sink(), state['revision_id'])
        updated = self.respond(before)
        self.assertNotEqual(updated['revision_id'], before['revision_id'])
        self.assertEqual(updated['response_revision'], 1)
        self.assertEqual(updated['human_acceptance'], 'NOT_CLAIMED')
        self.assertEqual(self.responses(before['revision_id'])['records'], {})
        self.assertEqual(work.get_plugin_snapshot('p', 'd', self.session, self.sink(), before['revision_id']), old)
        snapshot = work.get_plugin_snapshot('p', 'd', self.session, self.sink(), updated['revision_id'])
        self.assertEqual(snapshot['source_item_baseline']['manual_decisions_revision'], 1)
        self.assertTrue(any(n.get('manual_decision') for n in snapshot['nodes']))
        exported = work.export_plugin_draft('p', 'd', self.session, updated['revision_id'])
        book = load_workbook(BytesIO(exported.content))
        self.assertEqual({row[0]:row[1] for row in book['_ReqPilotXMeta'].values}['SnapshotSHA256'], work.core.canonical_json_sha256(snapshot))
        book.close()
        from reqpilotx_server.application.agent_workflow import plugin_workbench_reader as reader
        view = reader.get_plugin_reader('p', 'd', self.session, updated['revision_id'], self.sink())
        self.assertEqual(view['reasoning']['engineer_responses']['records'], updated['records'])

    def test_response_replay_is_exact_and_conflicts_preserve_history(self):
        _, before = self.start()
        first = self.respond(before)
        second = self.respond(first, command_id='response-two')
        replay = self.respond(before)
        self.assertEqual(replay['revision_id'], first['revision_id'])
        self.assertEqual(replay['records'], first['records'])
        self.assertEqual(replay['current_revision_id'], second['revision_id'])
        for change in ({'command_id':'stale'}, {'expected_version':'a'*64}, {'expected_response_revision':99}):
            with self.assertRaises(ApplicationError):
                self.respond(before, **change)
        self.assertEqual(self.responses()['revision_id'], second['revision_id'])
        self.assertEqual(len(work.get_plugin_workbench('p', 'd', self.session, self.sink())['versions']), 3)

    def test_response_cannot_cross_same_source_session_or_stale_unit_context(self):
        _, before = self.start()
        other = 'plugin-'+'b'*32
        work.sync_plugin_draft('p', 'd', work.PluginDraftUpload(session_id=other,
            expected_source_sha256=self.sha, artifacts=self.artifacts), self.sink())
        with self.assertRaises(ApplicationError) as caught:
            api.post_plugin_engineer_responses('p', 'd', other, self.response_payload(before), self.sink())
        self.assertEqual(caught.exception.detail['code'], 'PLUGIN_VERSION_CHANGED_OR_WRONG_TASK')
        choices = self.response_payload(before).model_dump()['decisions']
        choices[0]['expected_unit_context_sha256'] = '0'*64
        with self.assertRaises(ApplicationError): self.respond(before, decisions=choices)
        self.assertEqual(self.responses()['response_revision'], 0)

    def test_two_response_writers_only_one_can_advance(self):
        _, before = self.start()
        def writer(command):
            try: return self.respond(before, command_id=command)
            except ApplicationError as exc: return exc.status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            values = list(pool.map(writer, ['first', 'second']))
        self.assertEqual(sum(isinstance(value, dict) for value in values), 1)
        self.assertIn(409, values)
        self.assertEqual(self.responses()['response_revision'], 1)

    def test_external_sync_preserves_opinions_and_marks_changed_context_stale_without_revival(self):
        _, before = self.start()
        authored = self.respond(before)
        # An identical producer sync cannot erase the independently saved opinion.
        repeated = self.sync(self.artifacts, before['revision_id'], require_map_confirmation=False)
        self.assertEqual(repeated['revision_id'], authored['revision_id'])
        mapped = self.sync(self.mapping(), authored['revision_id'], require_map_confirmation=False)
        stale = self.responses()
        self.assertFalse(stale['available'])
        self.assertTrue(all(record['state'] == 'STALE' for record in stale['records'].values()))
        self.sync(self.artifacts, mapped['revision_id'], require_map_confirmation=False)
        restored = self.responses()
        self.assertTrue(all(record['state'] == 'STALE' for record in restored['records'].values()))
        self.assertEqual(self.responses(authored['revision_id'])['records'], authored['records'])
        self.assertFalse(self.preview()['readiness']['can_lock'])

    def test_deliveries_freeze_response_snapshot_and_bytes_even_after_source_replacement(self):
        _, before = self.start()
        authored = self.respond(before)
        preview = self.preview()
        saved = self.save(preview)
        frozen = api.get_plugin_delivery_version('p', 'd', self.session, 1, self.sink())
        exported = api.download_plugin_delivery_version('p', 'd', self.session, 1)
        self.assertEqual(frozen['snapshot'], preview['snapshot'])
        book = load_workbook(BytesIO(exported.content))
        self.assertEqual({row[0]:row[1] for row in book['_ReqPilotXMeta'].values}['SnapshotSHA256'], work.core.canonical_json_sha256(frozen['snapshot']))
        book.close()
        self.respond(authored, command_id='new-opinion')
        self.assertEqual(self.save(preview), saved)
        with self.assertRaises(ApplicationError): self.save(preview, command_id='stale-version')
        self.document['source_sha256'] = 'f'*64
        self.assertEqual(api.download_plugin_delivery_version('p', 'd', self.session, 1).content, exported.content)
        self.assertEqual(self.save(preview), saved)
        with self.assertRaises(ApplicationError): self.save(preview, command_id='new-source')
        self.assertEqual(len(api.list_plugin_delivery_versions('p', 'd', self.session, self.sink())['versions']), 1)
        with self.assertRaises(ApplicationError): api.get_plugin_delivery_version('p', 'other', self.session, 1, self.sink())

    def test_formal_requires_human_acceptance_and_all_effective_responses(self):
        _, before = self.start()
        with self.assertRaises(ApplicationError): self.save(self.preview(), mode='FORMAL', accept_snapshot=True)
        self.respond(before)
        self.assertTrue(self.preview()['readiness']['can_lock'])
        with self.assertRaises(ApplicationError): self.save(self.preview(), mode='FORMAL')
        saved = self.save(self.preview(), mode='FORMAL', accept_snapshot=True)
        self.assertEqual(saved['acceptance']['accepted_delivery_sha256'], self.preview()['delivery_sha256'])

    def test_delivery_cannot_cross_same_source_session_or_reuse_command_with_new_input(self):
        self.start()
        preview = self.preview()
        other = 'plugin-'+'b'*32
        work.sync_plugin_draft('p', 'd', work.PluginDraftUpload(session_id=other,
            expected_source_sha256=self.sha, artifacts=self.artifacts), self.sink())
        with self.assertRaises(ApplicationError) as caught:
            api.save_plugin_delivery_version('p', 'd', other, self.delivery_payload(preview), self.sink())
        self.assertEqual(caught.exception.detail['code'], 'PLUGIN_VERSION_CHANGED_OR_WRONG_TASK')
        saved = self.save(preview)
        with self.assertRaises(ApplicationError) as caught:
            self.save(preview, accept_snapshot=True)
        self.assertEqual(caught.exception.detail['code'], 'DELIVERY_COMMAND_REUSED_WITH_DIFFERENT_INPUT')
        self.assertEqual(self.save(preview), saved)
        self.assertEqual(api.list_plugin_delivery_versions('p', 'd', other, self.sink())['versions'], [])

    def test_concurrent_same_command_delivery_is_one_frozen_version(self):
        self.start()
        preview = self.preview()
        with ThreadPoolExecutor(max_workers=2) as pool:
            values = list(pool.map(lambda _: self.save(preview), range(2)))
        self.assertEqual(values[0], values[1])
        self.assertEqual(len(api.list_plugin_delivery_versions('p', 'd', self.session, self.sink())['versions']), 1)
        self.assertEqual(api.download_plugin_delivery_version('p', 'd', self.session, 1).headers['X-ReqPilotX-Export-SHA256'], values[0]['export_sha256'])

    def test_concurrent_response_and_delivery_either_freeze_pinned_snapshot_or_conflict(self):
        _, before = self.start()
        preview = self.preview()
        def deliver():
            try: return self.save(preview)
            except ApplicationError as exc: return exc.status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            delivery = pool.submit(deliver)
            response = pool.submit(self.respond, before)
            saved, authored = delivery.result(), response.result()
        self.assertEqual(self.responses()['revision_id'], authored['revision_id'])
        if isinstance(saved, dict):
            frozen = api.get_plugin_delivery_version('p', 'd', self.session, 1, self.sink())
            self.assertEqual(frozen['revision_id'], preview['revision_id'])
            self.assertEqual(frozen['snapshot'], preview['snapshot'])
        else:
            self.assertEqual(saved, 409)
            self.assertEqual(api.list_plugin_delivery_versions('p', 'd', self.session, self.sink())['versions'], [])

    def test_withdrawal_does_not_retain_effective_conclusion_and_tampering_fails_closed(self):
        _, before = self.start()
        authored = self.respond(before)
        choices = [dict(unit_id=uid, expected_unit_context_sha256=unit['unit_context_sha256'], action='WITHDRAW', reason='Withdraw synthetic opinion')
                   for uid, unit in authored['units'].items()]
        withdrawn = self.respond(authored, command_id='withdraw', decisions=choices)
        self.assertTrue(all(record['state'] == 'WITHDRAWN' for record in withdrawn['records'].values()))
        self.assertFalse(self.preview()['readiness']['can_lock'])
        saved = self.save(self.preview())
        path = work._root('p', 'd', self.session)/'deliveries'/(saved['export_sha256']+'.xlsx')
        path.write_bytes(b'tampered')
        with self.assertRaises(ApplicationError): api.download_plugin_delivery_version('p', 'd', self.session, 1)

    def test_routes_are_typed_and_require_task_bound_version_fields(self):
        self.start()
        from reqpilotx_server.api.routers.source_item_delivery import build
        app = FastAPI(); app.include_router(build())
        root = f'/api/projects/p/sor-documents/d/agent-structure/delivery/plugin-workbench/{self.session}'
        with TestClient(app) as client:
            view = client.get(root+'/engineer-responses').json()
            self.assertEqual(view['session_id'], self.session)
            invalid = self.response_payload(view).model_dump(); invalid.pop('expected_version')
            self.assertEqual(client.post(root+'/engineer-responses', json=invalid).status_code, 422)
            result = client.post(root+'/engineer-responses', json=self.response_payload(view).model_dump())
            self.assertEqual(result.status_code, 200, result.text)
            preview = client.get(root+'/versions/preview').json()
            saved = client.post(root+'/versions', json=self.delivery_payload(preview).model_dump())
            self.assertEqual(saved.status_code, 200, saved.text)
            self.assertEqual(client.get(root+'/versions/1').json()['snapshot'], preview['snapshot'])
            self.assertEqual(client.get(root+'/versions/1/export').status_code, 200)
