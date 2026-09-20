from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import hashlib
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from openpyxl import Workbook
from PIL import Image

import scripts.document_visual_pipeline_export as export_module
from app.application.extraction.classify_content import (
    DEFAULT_SEMANTIC_CLASSIFIER_ID,
    DEFAULT_SEMANTIC_CLASSIFIER_VERSION,
)
from app.application.extraction.project_requirements import (
    DEFAULT_REQUIREMENT_PROJECTOR_ID,
    DEFAULT_REQUIREMENT_PROJECTOR_VERSION,
)
from app.application.extraction.resolve_structure import (
    DEFAULT_STRUCTURE_RESOLVER_ID,
    DEFAULT_STRUCTURE_RESOLVER_VERSION,
)
from app.domain.extraction.source_ir import SourceIRStatus
from scripts.document_visual_pipeline_export import (
    DocumentVisualPipelineExportRejected,
    export_document_visual_pipeline,
)


class DocumentVisualPipelineExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(
            ignore_cleanup_errors=True
        )
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name).resolve()
        self.corpus_root = self.root / "corpus"
        self.corpus_root.mkdir()
        self.output_root = self.root / "controlled-exports"
        self.output_root.mkdir()

    def assert_contract_error(self, code: str):
        return self.assertRaisesRegex(
            DocumentVisualPipelineExportRejected,
            rf"^{re.escape(code)}$",
        )

    def export(
        self,
        source: Path,
        *,
        document_format: str,
        destination_name: str,
        expected_source_sha256: str | None = None,
    ) -> tuple[Path, dict]:
        destination = self.output_root / destination_name
        result = export_document_visual_pipeline(
            source,
            document_format=document_format,
            expected_source_sha256=(
                expected_source_sha256 or _sha256(source)
            ),
            output_dir=destination,
            corpus_root=self.corpus_root,
        )
        return destination, result

    def test_real_docx_and_xlsx_run_four_current_pipeline_stages(self) -> None:
        cases = (
            ("docx", self._docx_source("real-docx")),
            ("xlsx", self._xlsx_source("real-xlsx")),
        )
        for document_format, source in cases:
            with self.subTest(document_format=document_format):
                destination, result = self.export(
                    source,
                    document_format=document_format,
                    destination_name=document_format,
                )
                self.assertEqual(result["pipeline_status"], "COMPLETE")
                profile = result["profile"]
                self.assertEqual(
                    [item["kind"] for item in profile["artifacts"]],
                    [
                        "SOURCE_IR",
                        "STRUCTURE_TREE",
                        "SEMANTICS",
                        "REQUIREMENT_VIEW",
                    ],
                )
                self.assertEqual(
                    profile["pipeline"]["resolver"]["id"],
                    DEFAULT_STRUCTURE_RESOLVER_ID,
                )
                self.assertEqual(
                    profile["pipeline"]["classifier"]["version"],
                    DEFAULT_SEMANTIC_CLASSIFIER_VERSION,
                )
                self.assertEqual(
                    profile["pipeline"]["projector"]["id"],
                    DEFAULT_REQUIREMENT_PROJECTOR_ID,
                )
                for artifact in profile["artifacts"]:
                    artifact_path = destination / artifact["path"]
                    self.assertEqual(_sha256(artifact_path), artifact["sha256"])
                    self.assertIsInstance(
                        json.loads(artifact_path.read_text(encoding="utf-8")),
                        dict,
                    )
                self.assertEqual(
                    _sha256(destination / "pipeline-profile.json"),
                    result["profile_file_sha256"],
                )
                serialized = "".join(
                    path.read_text(encoding="utf-8")
                    for path in destination.iterdir()
                )
                self.assertNotIn(str(source), serialized)
                self.assertNotIn(source.name, serialized)

    def test_tif_runs_real_preflight_and_remains_blocked_not_run(self) -> None:
        source = self._tif_source("opaque-tif")
        destination, result = self.export(
            source,
            document_format="tif",
            destination_name="tif",
        )

        self.assertEqual(result["pipeline_status"], "BLOCKED")
        profile = result["profile"]
        self.assertEqual(
            [artifact["kind"] for artifact in profile["artifacts"]],
            ["PREFLIGHT_REPORT"],
        )
        self.assertEqual(
            profile["pipeline"]["adapter"]["status"],
            "BLOCKED",
        )
        for component in ("resolver", "classifier", "projector"):
            self.assertEqual(
                profile["pipeline"][component]["status"],
                "NOT_RUN",
            )
        evidence = json.loads(
            (destination / "artifact-0001.json").read_text(encoding="utf-8")
        )
        self.assertTrue(evidence["status"]["attachment_acceptable"])
        self.assertFalse(
            evidence["opaque_attachment"]["preservation_executed"]
        )
        self.assertFalse(
            evidence["opaque_attachment"]["downstream_pipeline_executed"]
        )
        serialized = json.dumps(evidence, sort_keys=True)
        self.assertNotIn(source.name, serialized)
        self.assertNotIn(str(source), serialized)
        self.assertEqual(
            {path.name for path in destination.iterdir()},
            {"artifact-0001.json", "pipeline-profile.json"},
        )

    def test_xlsb_route_keeps_partial_adapter_status_and_real_stage_order(
        self,
    ) -> None:
        source = self.corpus_root / "synthetic.xlsb"
        source.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512)
        source_sha256 = _sha256(source)
        stage_order: list[str] = []
        source_ir = _FakeSnapshot(
            {
                "schema_version": "source-ir/v1",
                "source_sha256": source_sha256,
            },
            source_sha256=source_sha256,
            status=SourceIRStatus.PARTIAL,
            adapter_id="controlled-xlsb-to-xlsx",
            adapter_version="test-bound-version",
            source_ir_id="source-ir-" + "1" * 32,
        )
        structure = _FakeSnapshot(
            {"schema_version": "structure-tree/v2"},
            resolver_id=DEFAULT_STRUCTURE_RESOLVER_ID,
            resolver_version=DEFAULT_STRUCTURE_RESOLVER_VERSION,
            structure_snapshot_id="structure-snapshot-" + "2" * 32,
        )
        semantics = _FakeSnapshot(
            {"schema_version": "semantic-classification/v3"},
            classifier_id=DEFAULT_SEMANTIC_CLASSIFIER_ID,
            classifier_version=DEFAULT_SEMANTIC_CLASSIFIER_VERSION,
            semantic_snapshot_id="semantic-snapshot-" + "3" * 32,
        )
        view = _FakeSnapshot(
            {"schema_version": "requirement-view/v3"},
            projector_id=DEFAULT_REQUIREMENT_PROJECTOR_ID,
            projector_version=DEFAULT_REQUIREMENT_PROJECTOR_VERSION,
            requirement_view_snapshot_id="requirement-view-" + "4" * 32,
        )

        def source_stage(*_args, **_kwargs):
            stage_order.append("build_source_ir")
            return source_ir

        def structure_stage(*_args, **_kwargs):
            stage_order.append("build_structure_tree")
            return structure

        def semantic_stage(*_args, **_kwargs):
            stage_order.append("classify_content")
            return semantics

        def projection_stage(*_args, **_kwargs):
            stage_order.append("project_requirements")
            return view

        report = _xlsb_preflight_report(source)
        with (
            mock.patch.object(
                export_module.LocalDocumentPreflight,
                "inspect",
                return_value=report,
            ),
            mock.patch.object(
                export_module,
                "build_source_ir",
                side_effect=source_stage,
            ),
            mock.patch.object(
                export_module,
                "build_structure_tree",
                side_effect=structure_stage,
            ),
            mock.patch.object(
                export_module,
                "classify_content",
                side_effect=semantic_stage,
            ),
            mock.patch.object(
                export_module,
                "project_requirements",
                side_effect=projection_stage,
            ),
        ):
            _, result = self.export(
                source,
                document_format="xlsb",
                destination_name="xlsb",
            )

        self.assertEqual(
            stage_order,
            [
                "build_source_ir",
                "build_structure_tree",
                "classify_content",
                "project_requirements",
            ],
        )
        self.assertEqual(result["pipeline_status"], "PARTIAL")
        self.assertEqual(
            result["profile"]["pipeline"]["adapter"]["status"],
            "PARTIAL",
        )
        self.assertEqual(
            result["profile"]["pipeline"]["adapter"]["version"],
            "test-bound-version",
        )

    def test_stage_failure_is_atomic_and_customer_text_free(self) -> None:
        source = self._docx_source("stage-failure")
        destination = self.output_root / "stage-failure"
        with (
            mock.patch.object(
                export_module,
                "classify_content",
                side_effect=RuntimeError("CUSTOMER SECRET"),
            ),
            self.assert_contract_error(
                "pipeline_export_stage_failed:classify_content"
            ),
        ):
            self.export(
                source,
                document_format="docx",
                destination_name="stage-failure",
            )
        self.assertFalse(destination.exists())
        self.assertEqual(
            [
                path
                for path in self.output_root.iterdir()
                if path.name.startswith(".stage-failure.")
            ],
            [],
        )

    def test_source_pin_output_scope_and_preexisting_are_strict(self) -> None:
        source = self._docx_source("boundaries")
        with self.assert_contract_error("pipeline_export_source_hash_mismatch"):
            self.export(
                source,
                document_format="docx",
                destination_name="wrong-hash",
                expected_source_sha256="f" * 64,
            )

        with self.assert_contract_error("pipeline_export_output_inside_corpus"):
            export_document_visual_pipeline(
                source,
                document_format="docx",
                expected_source_sha256=_sha256(source),
                output_dir=self.corpus_root / "forbidden",
                corpus_root=self.corpus_root,
            )

        destination = self.output_root / "existing"
        destination.mkdir()
        marker = destination / "marker"
        marker.write_text("preserve", encoding="utf-8")
        with self.assert_contract_error("pipeline_export_output_exists"):
            self.export(
                source,
                document_format="docx",
                destination_name="existing",
            )
        self.assertEqual(marker.read_text(encoding="utf-8"), "preserve")

    def test_source_inode_change_before_publish_is_toctou_and_atomic(
        self,
    ) -> None:
        source = self._docx_source("toctou")
        destination = self.output_root / "toctou"
        original_verified_result = export_module._verified_export_result

        def replace_source(**kwargs):
            result = original_verified_result(**kwargs)
            replacement = source.with_name("replacement.docx")
            replacement.write_bytes(source.read_bytes())
            replacement.replace(source)
            return result

        with (
            mock.patch.object(
                export_module,
                "_verified_export_result",
                side_effect=replace_source,
            ),
            self.assert_contract_error(
                "pipeline_export_source_toctou_detected"
            ),
        ):
            self.export(
                source,
                document_format="docx",
                destination_name="toctou",
            )
        self.assertFalse(destination.exists())

    def _docx_source(self, name: str) -> Path:
        path = self.corpus_root / f"{name}.docx"
        document = Document()
        document.add_heading("1 Scope", level=1)
        document.add_paragraph(
            "1.1 The system shall retain the traceability identifier."
        )
        document.save(path)
        return path

    def _xlsx_source(self, name: str) -> Path:
        path = self.corpus_root / f"{name}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Requirements"
        sheet.append(["ID", "Requirement"])
        sheet.append(["REQ-1", "The system shall preserve source ordering."])
        workbook.save(path)
        workbook.close()
        return path

    def _tif_source(self, name: str) -> Path:
        path = self.corpus_root / f"{name}.tif"
        first = Image.new("RGB", (20, 30), color=(100, 20, 20))
        second = Image.new("RGB", (20, 30), color=(20, 100, 20))
        first.save(path, format="TIFF", save_all=True, append_images=[second])
        return path


class _FakeSnapshot:
    def __init__(self, payload: dict, **attributes) -> None:
        self._payload = payload
        for name, value in attributes.items():
            setattr(self, name, value)

    def model_dump(self, *, mode: str) -> dict:
        if mode != "json":
            raise ValueError("test snapshot requires JSON mode")
        return dict(self._payload)


def _xlsb_preflight_report(source: Path):
    from app.domain.extraction.preflight import (
        AccessStatus,
        AuthenticityStatus,
        DiagnosticCode,
        DiagnosticSeverity,
        DocumentFormat,
        DocumentPreflightReport,
        IntegrityStatus,
        PreflightDiagnostic,
        RawDocumentIdentity,
        SupportStatus,
    )

    return DocumentPreflightReport(
        identity=RawDocumentIdentity(
            source_name=source.name,
            source_sha256=_sha256(source),
            source_size=source.stat().st_size,
            declared_extension=".xlsb",
            detected_format=DocumentFormat.XLSB,
            mime_type="application/vnd.ms-excel.sheet.binary.macroenabled.12",
        ),
        authenticity_status=AuthenticityStatus.MATCHED,
        support_status=SupportStatus.UNSUPPORTED,
        integrity_status=IntegrityStatus.VALID,
        access_status=AccessStatus.NOT_INSPECTED,
        diagnostics=(
            PreflightDiagnostic(
                code=DiagnosticCode.UNSUPPORTED_FORMAT,
                severity=DiagnosticSeverity.ERROR,
                message="Synthetic controlled-conversion boundary.",
            ),
        ),
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
