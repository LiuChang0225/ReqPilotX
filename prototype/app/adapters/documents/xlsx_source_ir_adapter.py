from __future__ import annotations

import hashlib
import json
import mimetypes
import posixpath
import warnings
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from os import PathLike
from pathlib import Path
from typing import Any, Callable

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from app.domain.extraction.preflight import DocumentFormat, DocumentPreflightReport
from app.domain.extraction.source_ir import (
    MergedRegion,
    NativeDocumentFormat,
    NativeLocator,
    SourceAssetKind,
    SourceAssetRef,
    SourceDiagnostic,
    SourceDiagnosticSeverity,
    SourceIRSnapshot,
    SourceIRStatus,
    SourceNode,
    SourceNodeKind,
    TableCellIR,
    TableIR,
    stable_cell_id,
    stable_document_id,
    stable_node_id,
    stable_source_ir_id,
    stable_table_id,
)
from app.ports.source_ir_adapter import (
    SourceIRAdapterRejected,
    SourceIRAdapterRejectionCode,
)
from app.adapters.documents.xlsx_cell_images import (
    XlsxCellImageAudit,
    XlsxCellImageReference,
    audit_xlsx_cell_images,
)


_WORKBOOK_PART = "xl/workbook.xml"
_CONTENT_TYPES_PART = "[Content_Types].xml"
_SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_DRAWING_MAIN_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_HASH_CHUNK_BYTES = 1024 * 1024
_MAX_XLSX_MATERIAL_CELLS = 500_000
_MAX_XLSX_MERGED_AREA = 1_000_000
_MAX_XLSX_TABLE_GRID_CELLS = 500_000
_MAX_XLSX_SOURCE_REGIONS = 2_000
_MAX_XLSX_IMAGE_PLACEMENTS = 50_000
_MAX_XLSX_UNIQUE_MEDIA_BYTES = 256 * 1024 * 1024

_NS = {
    "main": _SPREADSHEET_NS,
    "r": _DOCUMENT_REL_NS,
    "xdr": _DRAWING_NS,
    "a": _DRAWING_MAIN_NS,
}
_NON_IMAGE_DRAWING_TAGS = frozenset(
    {"sp", "grpSp", "graphicFrame", "cxnSp", "contentPart"}
)


@dataclass(frozen=True, slots=True)
class _Region:
    min_row: int
    min_column: int
    max_row: int
    max_column: int

    @property
    def cell_range(self) -> str:
        start = f"{get_column_letter(self.min_column)}{self.min_row}"
        end = f"{get_column_letter(self.max_column)}{self.max_row}"
        return start if start == end else f"{start}:{end}"

    @property
    def row_count(self) -> int:
        return self.max_row - self.min_row + 1

    @property
    def column_count(self) -> int:
        return self.max_column - self.min_column + 1

    def contains(self, row: int, column: int) -> bool:
        return (
            self.min_row <= row <= self.max_row
            and self.min_column <= column <= self.max_column
        )


@dataclass(frozen=True, slots=True)
class _Relationship:
    relationship_id: str
    relationship_type: str
    target: str
    external: bool


@dataclass(frozen=True, slots=True)
class _VisualRecord:
    row: int | None
    column: int | None
    locator: NativeLocator | None
    asset: SourceAssetRef
    content_bytes: bytes


@dataclass(frozen=True, slots=True)
class _ImageAssetMetadata:
    content_sha256: str
    media_type: str
    storage_key: str
    content_bytes: bytes


@dataclass(slots=True)
class _VisualResourceBudget:
    image_placements: int = 0
    unique_media_bytes: int = 0


class _XlsxResourceBudgetExceeded(RuntimeError):
    pass


class XlsxSourceIRAdapter:
    """Read-only XLSX-to-SourceIR adapter for the shadow pipeline.

    The exact bytes validated against preflight are parsed from memory.  No
    media is extracted beside the customer workbook and no semantic table or
    requirement classification is attempted here.
    """

    adapter_id = "native-xlsx-ooxml"
    adapter_version = "1.0.8"
    document_format = DocumentFormat.XLSX

    def __init__(
        self,
        *,
        source_asset_publisher: Callable[[SourceAssetRef, bytes], None]
        | None = None,
    ) -> None:
        if source_asset_publisher is not None and not callable(
            source_asset_publisher
        ):
            raise TypeError("source_asset_publisher must be callable")
        self._source_asset_publisher = source_asset_publisher

    def extract(
        self,
        source_path: str | PathLike[str],
        *,
        preflight: DocumentPreflightReport,
        created_at: datetime,
    ) -> SourceIRSnapshot:
        if (
            not isinstance(created_at, datetime)
            or created_at.tzinfo is None
            or created_at.utcoffset() is None
        ):
            raise SourceIRAdapterRejected(
                SourceIRAdapterRejectionCode.INVALID_TIMESTAMP,
                "created_at must be timezone-aware before XLSX extraction.",
            )

        path = Path(source_path)
        source_bytes, source_size, source_sha256 = _read_source_bytes(path)
        self._validate_preflight(
            path=path,
            source_size=source_size,
            source_sha256=source_sha256,
            preflight=preflight,
        )

        try:
            nodes, assets, diagnostics = self._extract_bytes(
                source_bytes,
                source_sha256=source_sha256,
            )
        except _XlsxResourceBudgetExceeded:
            return self._snapshot(
                source_sha256=source_sha256,
                created_at=created_at,
                status=SourceIRStatus.FAILED,
                diagnostics=(
                    SourceDiagnostic(
                        code="xlsx_resource_budget_exceeded",
                        severity=SourceDiagnosticSeverity.ERROR,
                        message=(
                            "The validated XLSX package exceeds a safe worksheet, drawing "
                            "placement, or decompressed-media budget."
                        ),
                    ),
                ),
            )
        except Exception as exc:
            return self._snapshot(
                source_sha256=source_sha256,
                created_at=created_at,
                status=SourceIRStatus.FAILED,
                diagnostics=(
                    SourceDiagnostic(
                        code="xlsx_parse_failed",
                        severity=SourceDiagnosticSeverity.ERROR,
                        message=(
                            "The validated XLSX package could not be mapped to SourceIR "
                            f"({type(exc).__name__})."
                        ),
                    ),
                ),
            )

        status = (
            SourceIRStatus.PARTIAL
            if any(
                item.severity
                in {SourceDiagnosticSeverity.WARNING, SourceDiagnosticSeverity.ERROR}
                for item in diagnostics
            )
            else SourceIRStatus.COMPLETE
        )
        return self._snapshot(
            source_sha256=source_sha256,
            created_at=created_at,
            status=status,
            nodes=nodes,
            assets=assets,
            diagnostics=diagnostics,
        )

    def _validate_preflight(
        self,
        *,
        path: Path,
        source_size: int,
        source_sha256: str,
        preflight: DocumentPreflightReport,
    ) -> None:
        if not isinstance(preflight, DocumentPreflightReport):
            raise SourceIRAdapterRejected(
                SourceIRAdapterRejectionCode.INVALID_PREFLIGHT,
                "A DocumentPreflightReport is required before native XLSX parsing.",
            )
        identity = preflight.identity
        if (
            not preflight.processable
            or identity.detected_format is not self.document_format
            or identity.source_name != path.name
            or identity.declared_extension != path.suffix.lower()
        ):
            raise SourceIRAdapterRejected(
                SourceIRAdapterRejectionCode.INVALID_PREFLIGHT,
                "The preflight report does not authorize this XLSX source path.",
            )
        if identity.source_sha256 != source_sha256 or identity.source_size != source_size:
            raise SourceIRAdapterRejected(
                SourceIRAdapterRejectionCode.SOURCE_CHANGED,
                "XLSX source bytes changed after preflight.",
            )

    def _extract_bytes(
        self,
        source_bytes: bytes,
        *,
        source_sha256: str,
    ) -> tuple[tuple[SourceNode, ...], tuple[SourceAssetRef, ...], tuple[SourceDiagnostic, ...]]:
        formula_workbook = None
        cached_workbook = None
        try:
            _validate_xlsx_resource_budget(source_bytes)
            cell_image_audit = audit_xlsx_cell_images(source_bytes)
            with warnings.catch_warnings(record=True) as reader_warnings:
                warnings.simplefilter("always", UserWarning)
                formula_workbook = load_workbook(
                    BytesIO(source_bytes),
                    read_only=False,
                    data_only=False,
                    keep_links=False,
                    rich_text=True,
                )
                cached_workbook = load_workbook(
                    BytesIO(source_bytes),
                    read_only=False,
                    data_only=True,
                    keep_links=False,
                    rich_text=True,
                )
            with zipfile.ZipFile(BytesIO(source_bytes)) as archive:
                visual_records, visual_diagnostics = _scan_visual_assets(
                    archive,
                    source_sha256=source_sha256,
                    package_sha256=hashlib.sha256(source_bytes).hexdigest(),
                    cell_image_audit=cell_image_audit,
                )

            published_asset_ids: set[str] = set()
            for records in visual_records.values():
                for record in records:
                    if record.asset.asset_id in published_asset_ids:
                        continue
                    self._publish_source_asset(record.asset, record.content_bytes)
                    published_asset_ids.add(record.asset.asset_id)

            nodes: list[SourceNode] = []
            assets_by_id: dict[str, SourceAssetRef] = {}
            diagnostics: list[SourceDiagnostic] = list(visual_diagnostics)
            source_order = 0
            materialized_grid_cells = 0
            user_warning_count = sum(
                1
                for warning in reader_warnings
                if issubclass(warning.category, UserWarning)
            )
            if user_warning_count:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_native_reader_warning",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=(
                            "The native XLSX reader reported unsupported or partially handled "
                            f"package features ({user_warning_count} warning(s)); original OOXML "
                            "parts remain in the source package."
                        ),
                    )
                )

            for worksheet in formula_workbook.worksheets:
                cached_worksheet = cached_workbook[worksheet.title]
                regions = _used_regions(worksheet)
                materialized_grid_cells += sum(_region_area(region) for region in regions)
                if materialized_grid_cells > _MAX_XLSX_TABLE_GRID_CELLS:
                    raise _XlsxResourceBudgetExceeded
                table_nodes: list[tuple[_Region, SourceNode]] = []

                if worksheet.sheet_state != "visible":
                    diagnostics.append(
                        SourceDiagnostic(
                            code="xlsx_sheet_visibility_preserved",
                            severity=SourceDiagnosticSeverity.INFO,
                            message=(
                                f"Worksheet {worksheet.title!r} retains sheet_state="
                                f"{worksheet.sheet_state!r}."
                            ),
                        )
                    )
                if _worksheet_has_comments(worksheet):
                    diagnostics.append(
                        SourceDiagnostic(
                            code="xlsx_comments_ignored_by_policy",
                            severity=SourceDiagnosticSeverity.INFO,
                            message=(
                                f"Cell comments on worksheet {worksheet.title!r} were "
                                "intentionally ignored by the default extraction policy."
                            ),
                        )
                    )

                for region in regions:
                    locator = _xlsx_locator(worksheet.title, region.cell_range)
                    table_payload, table_diagnostics = _build_table(
                        worksheet=worksheet,
                        cached_worksheet=cached_worksheet,
                        region=region,
                        source_sha256=source_sha256,
                        adapter_id=self.adapter_id,
                        adapter_version=self.adapter_version,
                    )
                    table_node = _source_node(
                        source_sha256=source_sha256,
                        adapter_id=self.adapter_id,
                        adapter_version=self.adapter_version,
                        kind=SourceNodeKind.TABLE,
                        source_order=0,
                        locator=locator,
                        table_payload=table_payload,
                    )
                    table_nodes.append((region, table_node))
                    diagnostics.extend(
                        _diagnostic_for_node(item, table_node.source_node_id)
                        for item in table_diagnostics
                    )

                sheet_visuals = visual_records.get(worksheet.title, ())
                objects: list[tuple[int, int, int, object]] = []
                objects.extend(
                    (region.min_row, region.min_column, 0, (region, table_node))
                    for region, table_node in table_nodes
                )
                objects.extend(
                    (
                        record.row if record.row is not None else 2**31 - 1,
                        record.column if record.column is not None else 2**31 - 1,
                        1,
                        record,
                    )
                    for record in sheet_visuals
                )

                for _, _, kind_order, payload in sorted(
                    objects,
                    key=lambda item: (item[0], item[1], item[2]),
                ):
                    if kind_order == 0:
                        _, table_node = payload  # type: ignore[misc]
                        nodes.append(table_node.model_copy(update={"source_order": source_order}))
                        source_order += 1
                        continue

                    record = payload
                    assert isinstance(record, _VisualRecord)
                    assets_by_id.setdefault(record.asset.asset_id, record.asset)
                    if record.locator is None:
                        continue
                    parent_source_node_id = _parent_table_node_id(
                        table_nodes,
                        row=record.row,
                        column=record.column,
                    )
                    image_node = _source_node(
                        source_sha256=source_sha256,
                        adapter_id=self.adapter_id,
                        adapter_version=self.adapter_version,
                        kind=SourceNodeKind.IMAGE,
                        source_order=source_order,
                        locator=record.locator,
                        parent_source_node_id=parent_source_node_id,
                        asset_ref=record.asset.asset_id,
                    )
                    nodes.append(image_node)
                    source_order += 1

                if not regions and not sheet_visuals:
                    diagnostics.append(
                        SourceDiagnostic(
                            code="xlsx_sheet_without_used_region",
                            severity=SourceDiagnosticSeverity.INFO,
                            message=(
                                f"Worksheet {worksheet.title!r} has no non-empty, visibly "
                                "formatted, merged, table, or image region."
                            ),
                        )
                    )

            for records in visual_records.values():
                for record in records:
                    assets_by_id.setdefault(record.asset.asset_id, record.asset)

            return tuple(nodes), tuple(assets_by_id.values()), tuple(diagnostics)
        finally:
            if formula_workbook is not None:
                formula_workbook.close()
            if cached_workbook is not None:
                cached_workbook.close()

    def _publish_source_asset(
        self,
        asset: SourceAssetRef,
        content_bytes: bytes,
    ) -> None:
        if self._source_asset_publisher is None:
            return
        if (
            not content_bytes
            or hashlib.sha256(content_bytes).hexdigest() != asset.content_sha256
        ):
            raise RuntimeError("source asset bytes do not match SourceIR identity")
        self._source_asset_publisher(asset, content_bytes)

    def _snapshot(
        self,
        *,
        source_sha256: str,
        created_at: datetime,
        status: SourceIRStatus,
        nodes: tuple[SourceNode, ...] = (),
        assets: tuple[SourceAssetRef, ...] = (),
        diagnostics: tuple[SourceDiagnostic, ...] = (),
    ) -> SourceIRSnapshot:
        document_id = stable_document_id(source_sha256)
        return SourceIRSnapshot(
            source_ir_id=stable_source_ir_id(
                document_id=document_id,
                source_sha256=source_sha256,
                adapter_id=self.adapter_id,
                adapter_version=self.adapter_version,
            ),
            document_id=document_id,
            source_sha256=source_sha256,
            adapter_id=self.adapter_id,
            adapter_version=self.adapter_version,
            status=status,
            nodes=nodes,
            assets=assets,
            diagnostics=diagnostics,
            created_at=created_at,
        )


def _read_source_bytes(path: Path) -> tuple[bytes, int, str]:
    digest = hashlib.sha256()
    payload = bytearray()
    with path.open("rb") as handle:
        while chunk := handle.read(_HASH_CHUNK_BYTES):
            digest.update(chunk)
            payload.extend(chunk)
    return bytes(payload), len(payload), digest.hexdigest()


def _validate_xlsx_resource_budget(source_bytes: bytes) -> None:
    material_cell_count = 0
    merged_area = 0
    cell_tag = f"{{{_SPREADSHEET_NS}}}c"
    merge_tag = f"{{{_SPREADSHEET_NS}}}mergeCell"
    with zipfile.ZipFile(BytesIO(source_bytes)) as archive:
        worksheet_parts = sorted(
            name
            for name in archive.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
        for worksheet_part in worksheet_parts:
            with archive.open(worksheet_part) as stream:
                elements = etree.iterparse(
                    stream,
                    events=("end",),
                    tag=(cell_tag, merge_tag),
                    resolve_entities=False,
                    no_network=True,
                    huge_tree=False,
                )
                for _, element in elements:
                    if element.tag == cell_tag:
                        material_cell_count += 1
                        if material_cell_count > _MAX_XLSX_MATERIAL_CELLS:
                            raise _XlsxResourceBudgetExceeded
                    else:
                        cell_range = str(element.get("ref") or "")
                        min_column, min_row, max_column, max_row = range_boundaries(
                            cell_range
                        )
                        merged_area += (max_row - min_row + 1) * (
                            max_column - min_column + 1
                        )
                        if merged_area > _MAX_XLSX_MERGED_AREA:
                            raise _XlsxResourceBudgetExceeded
                    element.clear()


def _used_regions(worksheet: Any) -> tuple[_Region, ...]:
    material_cells: set[tuple[int, int]] = set()
    for cell in worksheet._cells.values():
        if cell.value is not None or _has_visible_format(cell):
            material_cells.add((cell.row, cell.column))

    regions = list(_connected_regions(material_cells))
    for merged_range in worksheet.merged_cells.ranges:
        min_column, min_row, max_column, max_row = range_boundaries(str(merged_range))
        regions.append(_Region(min_row, min_column, max_row, max_column))
    for table in worksheet.tables.values():
        min_column, min_row, max_column, max_row = range_boundaries(table.ref)
        regions.append(_Region(min_row, min_column, max_row, max_column))
    if len(regions) > _MAX_XLSX_SOURCE_REGIONS:
        raise _XlsxResourceBudgetExceeded
    coalesced = _coalesce_regions(regions)
    if any(_region_area(region) > _MAX_XLSX_TABLE_GRID_CELLS for region in coalesced):
        raise _XlsxResourceBudgetExceeded
    return coalesced


def _connected_regions(cells: set[tuple[int, int]]) -> tuple[_Region, ...]:
    remaining = set(cells)
    regions: list[_Region] = []
    while remaining:
        seed = min(remaining)
        remaining.remove(seed)
        stack = [seed]
        component = [seed]
        while stack:
            row, column = stack.pop()
            for row_delta in (-1, 0, 1):
                for column_delta in (-1, 0, 1):
                    if row_delta == 0 and column_delta == 0:
                        continue
                    neighbour = (row + row_delta, column + column_delta)
                    if neighbour in remaining:
                        remaining.remove(neighbour)
                        stack.append(neighbour)
                        component.append(neighbour)
        regions.append(
            _Region(
                min(row for row, _ in component),
                min(column for _, column in component),
                max(row for row, _ in component),
                max(column for _, column in component),
            )
        )
        if len(regions) > _MAX_XLSX_SOURCE_REGIONS:
            raise _XlsxResourceBudgetExceeded
    return tuple(regions)


def _coalesce_regions(regions: list[_Region]) -> tuple[_Region, ...]:
    pending = sorted(
        regions,
        key=lambda item: (item.min_row, item.min_column, item.max_row, item.max_column),
    )
    changed = True
    while changed:
        changed = False
        result: list[_Region] = []
        while pending:
            current = pending.pop(0)
            for index, candidate in enumerate(pending):
                if _regions_touch(current, candidate):
                    current = _Region(
                        min(current.min_row, candidate.min_row),
                        min(current.min_column, candidate.min_column),
                        max(current.max_row, candidate.max_row),
                        max(current.max_column, candidate.max_column),
                    )
                    pending.pop(index)
                    pending.insert(0, current)
                    changed = True
                    break
            else:
                result.append(current)
        pending = sorted(
            result,
            key=lambda item: (item.min_row, item.min_column, item.max_row, item.max_column),
        )
    return tuple(pending)


def _regions_touch(first: _Region, second: _Region) -> bool:
    return not (
        first.max_row + 1 < second.min_row
        or second.max_row + 1 < first.min_row
        or first.max_column + 1 < second.min_column
        or second.max_column + 1 < first.min_column
    )


def _region_area(region: _Region) -> int:
    return (region.max_row - region.min_row + 1) * (
        region.max_column - region.min_column + 1
    )


def _has_visible_format(cell: Any) -> bool:
    # ``has_style`` covers number formats, font, alignment, protection, fill,
    # and borders. Style meaning is not yet serialized, but styled blank cells
    # still affect the visible worksheet boundary and must not disappear.
    if bool(getattr(cell, "has_style", False)):
        return True
    fill = getattr(cell, "fill", None)
    if fill is not None and getattr(fill, "fill_type", None) is not None:
        return True
    border = getattr(cell, "border", None)
    if border is None:
        return False
    return any(
        getattr(side, "style", None) is not None
        for side in (
            border.left,
            border.right,
            border.top,
            border.bottom,
            border.diagonal,
            border.vertical,
            border.horizontal,
        )
    )


def _build_table(
    *,
    worksheet: Any,
    cached_worksheet: Any,
    region: _Region,
    source_sha256: str,
    adapter_id: str,
    adapter_version: str,
) -> tuple[TableIR, tuple[SourceDiagnostic, ...]]:
    if _region_area(region) > _MAX_XLSX_TABLE_GRID_CELLS:
        raise _XlsxResourceBudgetExceeded
    locator = _xlsx_locator(worksheet.title, region.cell_range)
    table_id = stable_table_id(
        source_sha256=source_sha256,
        adapter_id=adapter_id,
        adapter_version=adapter_version,
        native_locator=locator,
    )
    merged_regions: list[MergedRegion] = []
    merged_anchors: dict[tuple[int, int], tuple[int, int]] = {}
    covered_merged_cells: set[tuple[int, int]] = set()
    diagnostics: list[SourceDiagnostic] = []

    for merged_range in worksheet.merged_cells.ranges:
        min_column, min_row, max_column, max_row = range_boundaries(str(merged_range))
        if not (
            region.min_row <= min_row <= max_row <= region.max_row
            and region.min_column <= min_column <= max_column <= region.max_column
        ):
            continue
        row_span = max_row - min_row + 1
        column_span = max_column - min_column + 1
        relative_row = min_row - region.min_row
        relative_column = min_column - region.min_column
        merged_regions.append(
            MergedRegion(
                row=relative_row,
                column=relative_column,
                row_span=row_span,
                column_span=column_span,
            )
        )
        merged_anchors[(min_row, min_column)] = (row_span, column_span)
        covered_merged_cells.update(
            (row, column)
            for row in range(min_row, max_row + 1)
            for column in range(min_column, max_column + 1)
            if (row, column) != (min_row, min_column)
        )

    cells: list[TableCellIR] = []
    visibly_formatted = False
    for row in range(region.min_row, region.max_row + 1):
        for column in range(region.min_column, region.max_column + 1):
            if (row, column) in covered_merged_cells:
                continue
            cell = worksheet.cell(row=row, column=column)
            cached_cell = cached_worksheet.cell(row=row, column=column)
            formula = _formula_text(cell)
            raw_text = formula if formula is not None else _scalar_text(cell.value)
            display_value = _optional_scalar_text(cached_cell.value)
            if formula is not None and display_value is None:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_formula_cached_value_unavailable",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=(
                            "The formula was preserved, but this XLSX package contains no "
                            "cached display value for the cell."
                        ),
                        native_locator=_xlsx_locator(worksheet.title, cell.coordinate),
                    )
                )
            row_span, column_span = merged_anchors.get((row, column), (1, 1))
            relative_row = row - region.min_row
            relative_column = column - region.min_column
            cells.append(
                TableCellIR(
                    cell_id=stable_cell_id(table_id, relative_row, relative_column),
                    row=relative_row,
                    column=relative_column,
                    row_span=row_span,
                    column_span=column_span,
                    raw_text=raw_text,
                    formula=formula,
                    display_value=display_value,
                    native_locator=_xlsx_locator(worksheet.title, cell.coordinate),
                )
            )
            visibly_formatted = visibly_formatted or _has_visible_format(cell)

    notes = [
        f"sheet_state={worksheet.sheet_state}",
        f"used_region={region.cell_range}",
    ]
    hidden_rows = tuple(
        row
        for row in range(region.min_row, region.max_row + 1)
        if worksheet.row_dimensions[row].hidden
    )
    hidden_columns = tuple(
        get_column_letter(column)
        for column in range(region.min_column, region.max_column + 1)
        if worksheet.column_dimensions[get_column_letter(column)].hidden
    )
    if hidden_rows:
        notes.append("hidden_rows=" + ",".join(str(row) for row in hidden_rows))
    if hidden_columns:
        notes.append("hidden_columns=" + ",".join(hidden_columns))
    for table in worksheet.tables.values():
        if _range_is_within(table.ref, region):
            notes.append(f"excel_table={table.name}@{table.ref}")

    diagnostics.append(
        SourceDiagnostic(
            code="xlsx_row_locator_derived_from_table",
            severity=SourceDiagnosticSeverity.INFO,
            message=(
                "SourceIR v1 keeps the lossless grid in TableIR cells; row ranges are "
                "deterministically derived from the table range and row offsets instead of "
                "inventing TABLE_ROW raw_text."
            ),
            native_locator=locator,
        )
    )
    if visibly_formatted:
        diagnostics.append(
            SourceDiagnostic(
                code="xlsx_visual_format_not_serialized",
                severity=SourceDiagnosticSeverity.WARNING,
                message=(
                    "Visibly formatted cells participate in this used region, but exact "
                    "workbook styling is not serialized into TableIR v1; visual meaning must "
                    "remain unresolved until a visual rendition is mapped."
                ),
                native_locator=locator,
            )
        )

    return (
        TableIR(
            table_id=table_id,
            row_count=region.row_count,
            column_count=region.column_count,
            cells=tuple(cells),
            notes=tuple(notes),
            merged_regions=tuple(merged_regions),
        ),
        tuple(diagnostics),
    )


def _formula_text(cell: Any) -> str | None:
    if getattr(cell, "data_type", None) != "f":
        return None
    value = cell.value
    if isinstance(value, str):
        return value if value.startswith("=") else f"={value}"
    text = getattr(value, "text", None)
    if isinstance(text, str) and text.strip():
        return text if text.startswith("=") else f"={text}"
    rendered = str(value)
    return rendered if rendered.startswith("=") else f"={rendered}"


def _optional_scalar_text(value: object) -> str | None:
    return None if value is None else _scalar_text(value)


def _scalar_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _range_is_within(cell_range: str, region: _Region) -> bool:
    min_column, min_row, max_column, max_row = range_boundaries(cell_range)
    return (
        region.min_row <= min_row <= max_row <= region.max_row
        and region.min_column <= min_column <= max_column <= region.max_column
    )


def _worksheet_has_comments(worksheet: Any) -> bool:
    return any(getattr(cell, "comment", None) is not None for cell in worksheet._cells.values())


def _scan_visual_assets(
    archive: zipfile.ZipFile,
    *,
    source_sha256: str,
    package_sha256: str | None = None,
    cell_image_audit: XlsxCellImageAudit | None = None,
) -> tuple[dict[str, tuple[_VisualRecord, ...]], tuple[SourceDiagnostic, ...]]:
    member_names = frozenset(archive.namelist())
    content_types = _read_content_types(archive, member_names)
    sheet_parts = _workbook_sheet_parts(archive, member_names)
    records_by_sheet: dict[str, list[_VisualRecord]] = {}
    diagnostics: list[SourceDiagnostic] = []
    resource_budget = _VisualResourceBudget()
    asset_metadata_by_part: dict[str, _ImageAssetMetadata] = {}

    references_by_sheet: dict[
        int,
        tuple[XlsxCellImageReference, ...],
    ] = {}
    if cell_image_audit is not None:
        if cell_image_audit.source_sha256 != (
            package_sha256 or source_sha256
        ):
            raise ValueError("cell-image audit does not bind XLSX source")
        references_by_sheet = {
            sheet_index: tuple(
                item
                for item in cell_image_audit.references
                if item.sheet_index == sheet_index
            )
            for sheet_index in range(len(sheet_parts))
        }

    for sheet_index, (sheet_name, sheet_part) in enumerate(sheet_parts):
        records: list[_VisualRecord] = []
        placement_counts: dict[str, int] = {}
        records_by_sheet[sheet_name] = records
        if sheet_part not in member_names:
            diagnostics.append(
                SourceDiagnostic(
                    code="xlsx_worksheet_part_missing",
                    severity=SourceDiagnosticSeverity.WARNING,
                    message=f"Worksheet {sheet_name!r} references a missing OOXML part.",
                )
            )
            continue
        sheet_root = _parse_xml(archive.read(sheet_part))
        relationships = _read_relationships(archive, member_names, sheet_part)
        comment_relationship_present = any(
            relationship.relationship_type.endswith("/comments")
            for relationship in relationships.values()
        )

        for drawing in sheet_root.findall("main:drawing", namespaces=_NS):
            relationship_id = drawing.get(f"{{{_DOCUMENT_REL_NS}}}id", "")
            relationship = relationships.get(relationship_id)
            if relationship is None or relationship.external:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_drawing_relationship_unavailable",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=(
                            f"Worksheet {sheet_name!r} has a drawing relationship that cannot "
                            "be resolved to an internal OOXML part."
                        ),
                    )
                )
                continue
            drawing_part = _resolve_relationship_target(sheet_part, relationship.target)
            if drawing_part is None or drawing_part not in member_names:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_drawing_part_missing",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=f"Worksheet {sheet_name!r} references a missing drawing part.",
                    )
                )
                continue
            drawing_records, drawing_diagnostics = _scan_drawing_part(
                archive,
                member_names=member_names,
                content_types=content_types,
                sheet_name=sheet_name,
                drawing_part=drawing_part,
                source_sha256=source_sha256,
                placement_counts=placement_counts,
                resource_budget=resource_budget,
                asset_metadata_by_part=asset_metadata_by_part,
            )
            records.extend(drawing_records)
            diagnostics.extend(drawing_diagnostics)

        expected_sheet_hash = hashlib.sha256(sheet_name.encode("utf-8")).hexdigest()
        for reference in references_by_sheet.get(sheet_index, ()):
            if reference.sheet_name_sha256 != expected_sheet_hash:
                raise ValueError("cell-image reference sheet identity mismatch")
            locator = _xlsx_locator(sheet_name, reference.cell_range)
            if reference.disposition == "unresolved_source_reference":
                if reference.binding_kind == "dispimg_formula":
                    diagnostic_code = "xlsx_dispimg_reference_unresolved"
                    diagnostic_message = (
                        "A DISPIMG formula has no source-contained cell-image "
                        "definition and media relationship; no image was guessed or "
                        "fetched."
                    )
                else:
                    diagnostic_code = (
                        "xlsx_rich_value_cell_image_reference_unresolved"
                    )
                    diagnostic_message = (
                        "An Excel rich-value local-image cell has no resolvable "
                        "source-contained media relationship; no image was guessed "
                        "or fetched."
                    )
                diagnostics.append(
                    SourceDiagnostic(
                        code=diagnostic_code,
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=diagnostic_message,
                        native_locator=locator,
                    )
                )
                continue

            resource_budget.image_placements += 1
            if resource_budget.image_placements > _MAX_XLSX_IMAGE_PLACEMENTS:
                raise _XlsxResourceBudgetExceeded
            assert reference.embedded_content_sha256 is not None
            assert reference.embedded_media_type is not None
            assert reference.embedded_storage_key is not None
            assert reference.embedded_bytes is not None
            image_part = reference.embedded_storage_key.removeprefix("ooxml/")
            asset_metadata = asset_metadata_by_part.get(image_part)
            if asset_metadata is None:
                if (
                    resource_budget.unique_media_bytes
                    + len(reference.embedded_bytes)
                    > _MAX_XLSX_UNIQUE_MEDIA_BYTES
                ):
                    raise _XlsxResourceBudgetExceeded
                asset_metadata = _ImageAssetMetadata(
                    content_sha256=reference.embedded_content_sha256,
                    media_type=reference.embedded_media_type,
                    storage_key=_logical_media_storage_key(
                        content_sha256=reference.embedded_content_sha256,
                        part_name=image_part,
                        media_type=reference.embedded_media_type,
                    ),
                    content_bytes=reference.embedded_bytes,
                )
                asset_metadata_by_part[image_part] = asset_metadata
                resource_budget.unique_media_bytes += len(reference.embedded_bytes)
            elif (
                asset_metadata.content_sha256
                != reference.embedded_content_sha256
            ):
                raise ValueError("cell-image media identity mismatch")
            locator_key = locator.model_dump_json(exclude_none=True)
            object_index = placement_counts.get(locator_key, 0)
            placement_counts[locator_key] = object_index + 1
            image_locator = locator.model_copy(
                update={"object_index": object_index}
            )
            asset_id = _stable_reference_id(
                "asset",
                source_sha256,
                asset_metadata.content_sha256,
            )
            asset = SourceAssetRef(
                asset_id=asset_id,
                kind=SourceAssetKind.IMAGE,
                content_sha256=asset_metadata.content_sha256,
                media_type=asset_metadata.media_type,
                storage_key=asset_metadata.storage_key,
                native_locator=image_locator,
            )
            records.append(
                _VisualRecord(
                    row=_single_cell_row(reference.cell_range),
                    column=_single_cell_column(reference.cell_range),
                    locator=image_locator,
                    asset=asset,
                    content_bytes=asset_metadata.content_bytes,
                )
            )
            diagnostics.append(
                SourceDiagnostic(
                    code=(
                        "xlsx_dispimg_embedded_cell_image_preserved"
                        if reference.binding_kind == "dispimg_formula"
                        else "xlsx_rich_value_embedded_cell_image_preserved"
                    ),
                    severity=SourceDiagnosticSeverity.INFO,
                    message=(
                        "A spreadsheet cell-image binding was resolved to its "
                        "source-contained embedded media."
                    ),
                    native_locator=image_locator,
                )
            )

        legacy_drawings = sheet_root.findall("main:legacyDrawing", namespaces=_NS)
        if legacy_drawings and not comment_relationship_present:
            diagnostics.append(
                SourceDiagnostic(
                    code="xlsx_legacy_visual_object_not_mapped",
                    severity=SourceDiagnosticSeverity.WARNING,
                    message=(
                        f"Worksheet {sheet_name!r} contains a VML drawing that is not a "
                        "comment layer; its control state was not inferred."
                    ),
                )
            )
        if (
            sheet_root.find("main:controls", namespaces=_NS) is not None
            or sheet_root.find("main:oleObjects", namespaces=_NS) is not None
        ):
            diagnostics.append(
                SourceDiagnostic(
                    code="xlsx_control_object_not_mapped",
                    severity=SourceDiagnosticSeverity.WARNING,
                    message=(
                        f"Worksheet {sheet_name!r} contains a control or embedded object; "
                        "the object is retained in the source package and no checked state "
                        "was guessed."
                    ),
                )
            )

    if any("featurePropertyBag" in member for member in member_names):
        diagnostics.append(
            SourceDiagnostic(
                code="xlsx_cell_control_state_not_interpreted",
                severity=SourceDiagnosticSeverity.WARNING,
                message=(
                    "The XLSX package contains cell-control feature metadata; SourceIR retained "
                    "the original package and did not infer checkbox state from scalar values."
                ),
            )
        )
    if any(member.startswith("xl/activeX/") for member in member_names):
        diagnostics.append(
            SourceDiagnostic(
                code="xlsx_activex_object_not_mapped",
                severity=SourceDiagnosticSeverity.WARNING,
                message=(
                    "The XLSX package contains ActiveX parts; no control state was inferred."
                ),
            )
        )
    if (
        cell_image_audit is not None
        and cell_image_audit.unused_definition_count
    ):
        diagnostics.append(
            SourceDiagnostic(
                code="xlsx_cell_image_definition_unused",
                severity=SourceDiagnosticSeverity.INFO,
                message=(
                    "The XLSX package contains source-contained cell-image "
                    "definition(s) that are not referenced by a DISPIMG formula "
                    f"({cell_image_audit.unused_definition_count})."
                ),
            )
        )
    return (
        {name: tuple(records) for name, records in records_by_sheet.items()},
        tuple(diagnostics),
    )


def _single_cell_row(cell_range: str) -> int:
    min_column, min_row, max_column, max_row = range_boundaries(cell_range)
    if min_column != max_column or min_row != max_row:
        raise ValueError("cell-image locator must be a single cell")
    return min_row


def _single_cell_column(cell_range: str) -> int:
    min_column, min_row, max_column, max_row = range_boundaries(cell_range)
    if min_column != max_column or min_row != max_row:
        raise ValueError("cell-image locator must be a single cell")
    return min_column


def _scan_drawing_part(
    archive: zipfile.ZipFile,
    *,
    member_names: frozenset[str],
    content_types: dict[str, str],
    sheet_name: str,
    drawing_part: str,
    source_sha256: str,
    placement_counts: dict[str, int],
    resource_budget: _VisualResourceBudget,
    asset_metadata_by_part: dict[str, _ImageAssetMetadata],
) -> tuple[list[_VisualRecord], list[SourceDiagnostic]]:
    root = _parse_xml(archive.read(drawing_part))
    relationships = _read_relationships(archive, member_names, drawing_part)
    records: list[_VisualRecord] = []
    diagnostics: list[SourceDiagnostic] = []

    for anchor in root:
        anchor_name = etree.QName(anchor).localname
        if anchor_name not in {"oneCellAnchor", "twoCellAnchor", "absoluteAnchor"}:
            diagnostics.append(
                SourceDiagnostic(
                    code="xlsx_drawing_anchor_not_mapped",
                    severity=SourceDiagnosticSeverity.WARNING,
                    message=(
                        f"Worksheet {sheet_name!r} contains an unsupported drawing anchor "
                        f"type {anchor_name!r}."
                    ),
                )
            )
            continue
        locator, row, column = _drawing_locator(anchor, sheet_name)
        pics = anchor.findall("xdr:pic", namespaces=_NS)
        for pic in pics:
            resource_budget.image_placements += 1
            if resource_budget.image_placements > _MAX_XLSX_IMAGE_PLACEMENTS:
                raise _XlsxResourceBudgetExceeded
            image_locator = locator
            if locator is not None:
                locator_key = locator.model_dump_json(exclude_none=True)
                object_index = placement_counts.get(locator_key, 0)
                placement_counts[locator_key] = object_index + 1
                image_locator = locator.model_copy(update={"object_index": object_index})
            blip = pic.find(".//a:blip", namespaces=_NS)
            relationship_id = (
                blip.get(f"{{{_DOCUMENT_REL_NS}}}embed", "") if blip is not None else ""
            )
            relationship = relationships.get(relationship_id)
            if relationship is None or relationship.external:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_image_relationship_unavailable",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=(
                            f"Worksheet {sheet_name!r} contains an image whose internal "
                            "relationship cannot be resolved."
                        ),
                        native_locator=image_locator,
                    )
                )
                continue
            image_part = _resolve_relationship_target(drawing_part, relationship.target)
            if image_part is None or image_part not in member_names:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_image_part_missing",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=(
                            f"Worksheet {sheet_name!r} references a missing or unsafe image part."
                        ),
                        native_locator=image_locator,
                    )
                )
                continue
            asset_metadata = asset_metadata_by_part.get(image_part)
            if asset_metadata is None:
                member_size = archive.getinfo(image_part).file_size
                if (
                    resource_budget.unique_media_bytes + member_size
                    > _MAX_XLSX_UNIQUE_MEDIA_BYTES
                ):
                    raise _XlsxResourceBudgetExceeded
                image_bytes = archive.read(image_part)
                if (
                    resource_budget.unique_media_bytes + len(image_bytes)
                    > _MAX_XLSX_UNIQUE_MEDIA_BYTES
                ):
                    raise _XlsxResourceBudgetExceeded
                asset_metadata = _ImageAssetMetadata(
                    content_sha256=hashlib.sha256(image_bytes).hexdigest(),
                    media_type=_media_type(image_part, content_types),
                    storage_key=_logical_media_storage_key(
                        content_sha256=hashlib.sha256(image_bytes).hexdigest(),
                        part_name=image_part,
                        media_type=_media_type(image_part, content_types),
                    ),
                    content_bytes=image_bytes,
                )
                asset_metadata_by_part[image_part] = asset_metadata
                resource_budget.unique_media_bytes += len(image_bytes)
            asset_id = _stable_reference_id(
                "asset",
                source_sha256,
                asset_metadata.content_sha256,
            )
            asset = SourceAssetRef(
                asset_id=asset_id,
                kind=SourceAssetKind.IMAGE,
                content_sha256=asset_metadata.content_sha256,
                media_type=asset_metadata.media_type,
                storage_key=asset_metadata.storage_key,
                native_locator=image_locator,
            )
            records.append(
                _VisualRecord(
                    row=row,
                    column=column,
                    locator=image_locator,
                    asset=asset,
                    content_bytes=asset_metadata.content_bytes,
                )
            )
            if locator is None:
                diagnostics.append(
                    SourceDiagnostic(
                        code="xlsx_image_without_cell_locator",
                        severity=SourceDiagnosticSeverity.WARNING,
                        message=(
                            f"An image on worksheet {sheet_name!r} was retained as an asset, "
                            "but its absolute drawing anchor has no lossless SourceIR v1 cell "
                            "locator."
                        ),
                    )
                )

        non_image_objects = [
            etree.QName(child).localname
            for child in anchor
            if etree.QName(child).localname in _NON_IMAGE_DRAWING_TAGS
        ]
        if non_image_objects:
            diagnostics.append(
                SourceDiagnostic(
                    code="xlsx_visual_object_not_mapped",
                    severity=SourceDiagnosticSeverity.WARNING,
                    message=(
                        f"Worksheet {sheet_name!r} contains drawing object(s) "
                        f"{', '.join(non_image_objects)}; their state was not inferred."
                    ),
                    native_locator=locator,
                )
            )
    return records, diagnostics


def _drawing_locator(
    anchor: etree._Element,
    sheet_name: str,
) -> tuple[NativeLocator | None, int | None, int | None]:
    marker = anchor.find("xdr:from", namespaces=_NS)
    if marker is None:
        return None, None, None
    row = _xml_integer(marker.find("xdr:row", namespaces=_NS)) + 1
    column = _xml_integer(marker.find("xdr:col", namespaces=_NS)) + 1
    end_marker = anchor.find("xdr:to", namespaces=_NS)
    end_row = row
    end_column = column
    if end_marker is not None:
        end_row = max(row, _xml_integer(end_marker.find("xdr:row", namespaces=_NS)) + 1)
        end_column = max(
            column,
            _xml_integer(end_marker.find("xdr:col", namespaces=_NS)) + 1,
        )
    region = _Region(row, column, end_row, end_column)
    return _xlsx_locator(sheet_name, region.cell_range), row, column


def _xml_integer(element: etree._Element | None) -> int:
    if element is None or element.text is None:
        return 0
    try:
        return max(0, int(element.text))
    except ValueError:
        return 0


def _workbook_sheet_parts(
    archive: zipfile.ZipFile,
    member_names: frozenset[str],
) -> tuple[tuple[str, str], ...]:
    if _WORKBOOK_PART not in member_names:
        return ()
    root = _parse_xml(archive.read(_WORKBOOK_PART))
    relationships = _read_relationships(archive, member_names, _WORKBOOK_PART)
    result: list[tuple[str, str]] = []
    for sheet in root.findall("main:sheets/main:sheet", namespaces=_NS):
        name = sheet.get("name", "")
        relationship_id = sheet.get(f"{{{_DOCUMENT_REL_NS}}}id", "")
        relationship = relationships.get(relationship_id)
        if not name or relationship is None or relationship.external:
            continue
        target = _resolve_relationship_target(_WORKBOOK_PART, relationship.target)
        if target is not None:
            result.append((name, target))
    return tuple(result)


def _read_relationships(
    archive: zipfile.ZipFile,
    member_names: frozenset[str],
    source_part: str,
) -> dict[str, _Relationship]:
    relationship_part = _relationship_part_name(source_part)
    if relationship_part not in member_names:
        return {}
    root = _parse_xml(archive.read(relationship_part))
    relationships: dict[str, _Relationship] = {}
    for element in root.findall(f"{{{_PACKAGE_REL_NS}}}Relationship"):
        relationship_id = element.get("Id", "")
        if not relationship_id:
            continue
        relationships[relationship_id] = _Relationship(
            relationship_id=relationship_id,
            relationship_type=element.get("Type", ""),
            target=element.get("Target", ""),
            external=element.get("TargetMode", "").lower() == "external",
        )
    return relationships


def _relationship_part_name(source_part: str) -> str:
    directory, name = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", f"{name}.rels")


def _resolve_relationship_target(source_part: str, target: str) -> str | None:
    if not target or "\\" in target:
        return None
    if target.startswith("/"):
        normalized = posixpath.normpath(target.lstrip("/"))
    else:
        normalized = posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))
    if normalized in {"", ".", ".."} or normalized.startswith("../"):
        return None
    return normalized


def _read_content_types(
    archive: zipfile.ZipFile,
    member_names: frozenset[str],
) -> dict[str, str]:
    if _CONTENT_TYPES_PART not in member_names:
        return {}
    root = _parse_xml(archive.read(_CONTENT_TYPES_PART))
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for element in root:
        local_name = etree.QName(element).localname
        if local_name == "Default":
            defaults[element.get("Extension", "").lower()] = element.get("ContentType", "")
        elif local_name == "Override":
            overrides[element.get("PartName", "").lstrip("/")] = element.get(
                "ContentType", ""
            )
    result = {f"*.{extension}": value for extension, value in defaults.items() if value}
    result.update({name: value for name, value in overrides.items() if value})
    return result


def _media_type(part_name: str, content_types: dict[str, str]) -> str:
    exact = content_types.get(part_name)
    if exact:
        return exact
    extension = posixpath.splitext(part_name)[1].lstrip(".").lower()
    default = content_types.get(f"*.{extension}")
    if default:
        return default
    return mimetypes.guess_type(part_name)[0] or "application/octet-stream"


def _logical_media_storage_key(
    *,
    content_sha256: str,
    part_name: str,
    media_type: str,
) -> str:
    extension = Path(part_name).suffix.lower().lstrip(".")
    if not extension or not extension.isalnum() or len(extension) > 10:
        extension = {
            "image/jpeg": "jpg",
            "image/png": "png",
            "image/gif": "gif",
            "image/tiff": "tif",
            "image/svg+xml": "svg",
        }.get(media_type.lower(), "bin")
    return f"source-assets/{content_sha256[:2]}/{content_sha256}.{extension}"


def _parse_xml(payload: bytes) -> etree._Element:
    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        recover=False,
        huge_tree=False,
    )
    return etree.fromstring(payload, parser=parser)


def _xlsx_locator(sheet_name: str, cell_range: str) -> NativeLocator:
    return NativeLocator(
        format=NativeDocumentFormat.XLSX,
        workbook=_WORKBOOK_PART,
        sheet=sheet_name,
        cell_range=cell_range,
    )


def _source_node(
    *,
    source_sha256: str,
    adapter_id: str,
    adapter_version: str,
    kind: SourceNodeKind,
    source_order: int,
    locator: NativeLocator,
    parent_source_node_id: str | None = None,
    table_payload: TableIR | None = None,
    asset_ref: str | None = None,
) -> SourceNode:
    return SourceNode(
        source_node_id=stable_node_id(
            source_sha256=source_sha256,
            adapter_id=adapter_id,
            adapter_version=adapter_version,
            kind=kind,
            native_locator=locator,
        ),
        kind=kind,
        source_order=source_order,
        raw_text="",
        native_locator=locator,
        parent_source_node_id=parent_source_node_id,
        table_payload=table_payload,
        asset_ref=asset_ref,
    )


def _parent_table_node_id(
    table_nodes: list[tuple[_Region, SourceNode]],
    *,
    row: int | None,
    column: int | None,
) -> str | None:
    if row is None or column is None:
        return None
    matches = [
        node.source_node_id
        for region, node in table_nodes
        if region.contains(row, column)
    ]
    return matches[0] if len(matches) == 1 else None


def _diagnostic_for_node(
    diagnostic: SourceDiagnostic,
    source_node_id: str,
) -> SourceDiagnostic:
    return SourceDiagnostic(
        code=diagnostic.code,
        severity=diagnostic.severity,
        message=diagnostic.message,
        retryable=diagnostic.retryable,
        source_node_id=source_node_id,
        native_locator=diagnostic.native_locator,
    )


def _stable_reference_id(prefix: str, *parts: object) -> str:
    payload = json.dumps(parts, ensure_ascii=False, separators=(",", ":"))
    digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()[:32]
    return f"{prefix}-{digest}"


__all__ = ["XlsxSourceIRAdapter"]
