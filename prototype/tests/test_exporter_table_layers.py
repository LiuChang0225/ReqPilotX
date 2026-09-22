from __future__ import annotations

# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import hashlib
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
import zipfile

import pymupdf
from openpyxl import load_workbook
from PIL import Image as PILImage

from app.application.errors import ApplicationError
from app.application.exports import (
    _dump_matrix_export_snapshot,
    _matrix_decision_nodes,
    _matrix_snapshot_for_export,
    _parse_matrix_export_snapshot,
)
from app.domain.requirement_matrix import (
    RequirementMatrixSnapshotV1,
    matrix_conclusion_contract,
)
from app.services import exporter as exporter_module
from app.services.exporter import (
    _build_matrix_table_catalog,
    _visible_matrix_nodes,
    export_compliance_matrix,
)


OVERVIEW_SHEET = "00_清单概览"
MATRIX_SHEET = "ComplianceMatrix"
MATRIX_REQ_ID_COLUMN = 1
MATRIX_SOURCE_NUMBER_COLUMN = 2
MATRIX_TITLE_COLUMN = 3
MATRIX_BODY_COLUMN = 4
MATRIX_SOURCE_ENTRY_COLUMN = 5
MATRIX_ANALYSIS_SCOPE_COLUMN = 6
MATRIX_ANALYSIS_PROGRESS_COLUMN = 7
MATRIX_CONCLUSION_COLUMN = 8
MATRIX_BASIS_COLUMN = 9
MATRIX_CUSTOMER_CONFIRMATION_COLUMN = 10
DEFAULT_RECORD_BODY = "项目：第一页记录\n要求：12 V"


def _node(
    node_id: str,
    *,
    parent_id: str | None,
    child_ids: list[str],
    depth: int,
    source_order: int,
    display_role: str,
    label: str,
    summary: str = "",
    source_number: str = "",
    display_title: str = "",
    display_text: str = "",
    matrix_visibility: str = "VISIBLE",
    disposition: str = "RETAIN",
    disposition_reason: str = "",
    analysis_eligibility: str = "NOT_ELIGIBLE",
    analysis_status: str = "NOT_APPLICABLE",
    requirement_code: str = "",
    requirement_text: str = "",
    table_record_id: str | None = None,
    logical_table_ids: list[str] | None = None,
    table_refs: list[dict] | None = None,
    figure_refs: list[dict] | None = None,
    ai_assessment: dict | None = None,
    manual_decision: dict | None = None,
) -> dict:
    return {
        "node_id": node_id,
        "parent_id": parent_id,
        "child_ids": child_ids,
        "depth": depth,
        "source_order": source_order,
        "display_role": display_role,
        "source_number": source_number,
        "display_title": display_title,
        "display_text": display_text,
        "label": label,
        "summary": summary,
        "content_blocks": [],
        "table_refs": table_refs or [],
        "figure_refs": figure_refs or [],
        "source_refs": [{"page_number": 1}],
        "page_start": 1,
        "page_end": 1,
        "structure_node_id": f"structure-{node_id}",
        "requirement_entry_id": node_id if requirement_code else None,
        "requirement_id": node_id if requirement_code else None,
        "requirement_code": requirement_code,
        "requirement_text": requirement_text,
        "table_record_id": table_record_id,
        "logical_table_ids": logical_table_ids or [],
        "matrix_visibility": matrix_visibility,
        "disposition": disposition,
        "disposition_reason": disposition_reason,
        "analysis_eligibility": analysis_eligibility,
        "analysis_status": analysis_status,
        "ai_assessment": ai_assessment,
        "manual_decision": manual_decision,
        "compliance_status": (
            str((ai_assessment or {}).get("compliance_status") or "")
        ),
        "save_revision": 3,
        "updated_at": "2026-07-31T10:00:00+08:00",
    }


def _table_part(
    table_id: str,
    *,
    source_order: int,
    page: int,
    image_sha256: str = "",
) -> dict:
    suffix = "第一页记录" if source_order == 0 else "第二页记录"
    cells = [
        {
            "id": f"{table_id}-title",
            "row": 0,
            "column": 0,
            "column_span": 3,
            "raw_text": "性能要求",
        },
        {"id": f"{table_id}-h1", "row": 1, "column": 0, "raw_text": "项目"},
        {"id": f"{table_id}-h2", "row": 1, "column": 1, "raw_text": "要求"},
        {"id": f"{table_id}-h3", "row": 1, "column": 2, "raw_text": "示意"},
        {"id": f"{table_id}-r1c1", "row": 2, "column": 0, "raw_text": suffix},
        {"id": f"{table_id}-r1c2", "row": 2, "column": 1, "raw_text": "12 V"},
    ]
    if image_sha256:
        cells.append(
            {
                "id": f"{table_id}-r1c3",
                "row": 2,
                "column": 2,
                "raw_text": "接线示意",
                "images": [
                    {
                        "asset_id": "asset-cell-image",
                        "content_sha256": image_sha256,
                        "media_type": "image/png",
                        "storage_key": "portable/not-a-local-path.png",
                    }
                ],
            }
        )
    return {
        "logical_table_id": "logical-table-1",
        "table_id": table_id,
        "physical_table_id": table_id,
        "source_order": source_order,
        "source_page": page,
        "source_page_end": page,
        "cells": cells,
        "merged_regions": [{"row": 0, "column": 0, "row_span": 1, "column_span": 3}],
        "header_regions": [{"row": 0, "column": 0, "row_span": 2, "column_span": 3}],
        "notes": [f"第 {page} 页表注"],
    }


def _snapshot(*, image_sha256: str = "") -> dict:
    first_part = _table_part(
        "physical-table-1",
        source_order=0,
        page=1,
        image_sha256=image_sha256,
    )
    second_part = _table_part(
        "physical-table-2",
        source_order=1,
        page=2,
    )
    nodes = [
        _node(
            "root",
            parent_id=None,
            child_ids=["clause", "system", "excluded"],
            depth=0,
            source_order=0,
            display_role="STRUCTURE",
            label="9 技术要求",
            source_number="9",
            display_title="技术要求",
        ),
        _node(
            "clause",
            parent_id="root",
            child_ids=["table", "context"],
            depth=1,
            source_order=1,
            display_role="CLAUSE",
            label="9.16.2 电气性能",
            summary="电气性能应满足下表。",
            source_number="9.16.2",
            display_title="电气性能",
            display_text="电气性能应满足下表。",
        ),
        _node(
            "table",
            parent_id="clause",
            child_ids=["record"],
            depth=2,
            source_order=2,
            display_role="TABLE",
            label="9.16.2 性能参数表",
            summary="跨页性能参数原表",
            source_number="9.16.2",
            logical_table_ids=["logical-table-1"],
            table_refs=[first_part, second_part],
        ),
        _node(
            "record",
            parent_id="table",
            child_ids=["atom-1", "atom-2"],
            depth=3,
            source_order=3,
            display_role="TABLE_RECORD",
            label="工作电压记录",
            summary="工作电压 | 12 V",
            display_title="工作电压记录",
            display_text="工作电压\n12 V",
            table_record_id="record-1",
            logical_table_ids=["logical-table-1"],
            table_refs=[
                {
                    "logical_table_id": "logical-table-1",
                    "table_id": "physical-table-1",
                    "record_id": "record-1",
                    "inherited_headers": ["性能要求", "项目", "要求"],
                    "slice_rows": [["工作电压", "12 V"]],
                }
            ],
        ),
        _node(
            "atom-1",
            parent_id="record",
            child_ids=[],
            depth=4,
            source_order=4,
            display_role="ATOMIC_REQUIREMENT",
            label="工作电压",
            display_title="工作电压",
            requirement_code="REQ-001",
            requirement_text="设备工作电压应为 12 V。",
            table_record_id="record-1",
            logical_table_ids=["logical-table-1"],
            disposition="ANALYZE",
            analysis_eligibility="ELIGIBLE",
            analysis_status="DECIDED",
            ai_assessment={
                "decision": "ACCEPT",
                "compliance_status": "COMPLIANT",
                "reason": "证据完整。",
                "confidence": 0.93,
                "assessment_mode": "LOCAL_RULE",
            },
            manual_decision={
                "decision": "ACCEPT",
                "reason": "工程师确认",
                "decided_at": "2026-07-31T10:00:00+08:00",
            },
        ),
        _node(
            "atom-2",
            parent_id="record",
            child_ids=[],
            depth=4,
            source_order=5,
            display_role="ATOMIC_REQUIREMENT",
            label="REQ-002 电压容差",
            display_title="电压容差",
            requirement_code="REQ-002",
            requirement_text="电压容差应在允许范围内。",
            table_record_id="record-1",
            logical_table_ids=["logical-table-1"],
            disposition="ANALYZE",
            analysis_eligibility="ELIGIBLE",
            analysis_status="ASSESSED",
            ai_assessment={
                "decision": "NEED_CLARIFICATION",
                "compliance_status": "UNKNOWN",
                "reason": "缺少容差值。",
                "analysis_basis": "知识库未检索到明确容差证据。",
                "clarification_question": "请客户确认允许的电压容差。",
                "confidence": 0.61,
                "assessment_mode": "LOCAL_RULE",
            },
        ),
        _node(
            "context",
            parent_id="clause",
            child_ids=[],
            depth=2,
            source_order=6,
            display_role="CONTEXT",
            label="说明",
            summary="本段仅作为分析上下文。",
            display_title="说明",
            display_text="本段仅作为分析上下文。",
        ),
        _node(
            "system",
            parent_id="root",
            child_ids=[],
            depth=1,
            source_order=7,
            display_role="CONTEXT",
            label="封面",
            matrix_visibility="SYSTEM_NON_BUSINESS",
        ),
        _node(
            "excluded",
            parent_id="root",
            child_ids=[],
            depth=1,
            source_order=8,
            display_role="CLAUSE",
            label="人工排除内容",
            matrix_visibility="HUMAN_EXCLUDED",
            disposition="EXCLUDE",
            disposition_reason="客户确认该条款不适用于本项目。",
        ),
    ]
    projection = {
        "schema_version": "table-record-projection/v1",
        "projection_id": "projection-1",
        "mode": "RECORDS",
        "table_fact": {
            "logical_table_id": "logical-table-1",
            "owner_structure_node_id": "structure-table",
            "table_parts": [
                {
                    "table_id": "physical-table-1",
                    "source_order": 0,
                    "source_refs": [{"page_number": 1}],
                },
                {
                    "table_id": "physical-table-2",
                    "source_order": 1,
                    "source_refs": [{"page_number": 2}],
                },
            ],
        },
        "records": [
            {
                "record_id": "record-1",
                "source_order": 0,
                "header_bindings": [
                    {
                        "axis": "COLUMN",
                        "level": 1,
                        "text": "要求",
                        "cell_ids": ["physical-table-1-h2"],
                    },
                    {
                        "axis": "ROW",
                        "level": 0,
                        "text": "性能要求",
                        "cell_ids": ["physical-table-1-title"],
                    },
                ],
                "row_bindings": [
                    {
                        "table_id": "physical-table-1",
                        "row_index": 2,
                        "cell_ids": [
                            "physical-table-1-r1c1",
                            "physical-table-1-r1c2",
                        ],
                    }
                ],
                "shared_merged_values": [],
                "notes": [],
            }
        ],
    }
    return {
        "schema_version": "requirement-matrix-snapshot/v1",
        "state": "READY",
        "project_id": "project-1",
        "document_id": "document-1",
        "requirement_set_id": "set-1",
        "requirement_set_status": "LOCKED",
        "tree_sha256": "a" * 64,
        "root_node_ids": ["root"],
        "nodes": nodes,
        "logical_tables": [
            {
                "logical_table_id": "logical-table-1",
                "projection_id": "projection-1",
                "projection_mode": "RECORDS",
                "record_count": 1,
                "physical_table_ids": ["physical-table-1", "physical-table-2"],
                "node_ids": ["table", "record", "atom-1", "atom-2"],
                "evidence_binding_ids": ["binding-1"],
                "projection": projection,
                "evidence_bindings": [],
            }
        ],
        "counts": {
            "total_nodes": len(nodes),
            "visible_nodes": 7,
            "system_non_business_nodes": 1,
            "human_excluded_nodes": 1,
            "analyze_nodes": 2,
            "retained_nodes": 6,
            "excluded_nodes": 1,
            "analysis_eligible_nodes": 2,
            "assessed_nodes": 2,
            "decided_nodes": 1,
            "logical_tables": 1,
        },
    }


def _find_row(sheet, value: str, *, column: int = 1) -> int:
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=column).value == value:
            return row
    raise AssertionError(f"{value!r} was not found")


def _table_sheet(workbook, *, index: int = 0):
    names = [name for name in workbook.sheetnames if name.startswith("T")]
    if index >= len(names):
        raise AssertionError(f"table sheet {index} was not found")
    return workbook[names[index]]


def _sheet_text(sheet) -> str:
    return "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)


def _first_full_text_field(workbook, field: str) -> str:
    sheet = workbook["全文详情"]
    for row in range(4, sheet.max_row + 1):
        if sheet.cell(row, 2).value != field:
            continue
        chunks = [str(sheet.cell(row, 3).value or "")]
        cursor = row + 1
        while cursor <= sheet.max_row and sheet.cell(cursor, 2).value == "续":
            chunks.append(str(sheet.cell(cursor, 3).value or ""))
            cursor += 1
        return "".join(chunks)
    raise AssertionError(f"full-text field {field!r} was not found")


class ExporterTableLayersTests(unittest.TestCase):
    def _execute_application_export(
        self,
        snapshot: RequirementMatrixSnapshotV1,
        *,
        mode: str = "FORMAL",
        requirement_set_status: str = "LOCKED",
        preview_generation_id: str = "preview-1",
    ) -> tuple[dict, dict]:
        root = Path(self.enterContext(TemporaryDirectory()))
        source_path = root / "source.txt"
        source_path.write_text("source", encoding="utf-8")
        project = {
            "id": "project-1",
            "name": "ReqPilotX",
            "customer_name": "Customer",
        }
        document = {
            "id": "document-1",
            "project_id": "project-1",
            "file_name": "source.txt",
            "file_path": str(source_path),
            "active_requirement_set_id": "set-1",
        }
        requirement_set = {
            "id": "set-1",
            "project_id": "project-1",
            "sor_document_id": "document-1",
            "status": requirement_set_status,
            "version": 3,
            "source_hash": "a" * 64,
            "preview_generation_id": preview_generation_id,
        }
        captured: dict = {}

        def create_export_job(**kwargs):
            captured.update(kwargs)
            return {"id": "export-1"}

        with (
            patch(
                "app.application.projects._require_project",
                return_value=project,
            ),
            patch.object(exports_app.db, "get_sor_document", return_value=document),
            patch.object(
                exports_app.db,
                "get_requirement_set",
                return_value=requirement_set,
            ),
            patch.object(
                exports_app.db,
                "get_latest_completed_ai_analysis_job_for_set",
                return_value=None,
            ),
            patch.object(
                exports_app.db,
                "create_export_job",
                side_effect=create_export_job,
            ),
            patch(
                "app.application.requirement_matrix.build_requirement_matrix_snapshot",
                return_value=snapshot,
            ),
            patch.object(
                exports_app,
                "export_compliance_matrix",
                return_value=("matrix.xlsx", root / "matrix.xlsx"),
            ),
            patch.object(exports_app, "_is_desktop_embedded", return_value=False),
        ):
            result = exports_app._execute_project_export(
                "project-1",
                ExportRequest(
                    document_id="document-1",
                    requirement_set_id="set-1",
                    mode=mode,
                ),
            )
        return result, captured

    def test_formal_export_accepts_current_ai_conclusions_without_manual_clicks(
        self,
    ) -> None:
        snapshot = RequirementMatrixSnapshotV1.model_validate(_snapshot())

        result, _ = self._execute_application_export(snapshot)

        self.assertEqual(result["decision_eligible_count"], 2)
        self.assertEqual(result["human_decision_count"], 1)
        self.assertEqual(result["human_coverage"], 0.5)
        self.assertEqual(result["effective_conclusion_count"], 2)
        self.assertEqual(result["effective_conclusion_coverage"], 1.0)
        self.assertEqual(result["missing_effective_conclusion_count"], 0)
        self.assertEqual(result["ai_default_count"], 1)
        self.assertEqual(result["human_override_count"], 1)
        self.assertEqual(result["human_native_count"], 0)

    def test_formal_export_accepts_human_native_conclusion(self) -> None:
        payload = _snapshot()
        human_native = payload["nodes"][4]
        human_native["analysis_result_state"] = "CURRENT"
        human_native["ai_assessment"] = None
        human_native["compliance_status"] = ""
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        result, _ = self._execute_application_export(snapshot)

        self.assertEqual(result["effective_conclusion_count"], 2)
        self.assertEqual(result["ai_default_count"], 1)
        self.assertEqual(result["human_override_count"], 0)
        self.assertEqual(result["human_native_count"], 1)

    def test_formal_export_keeps_lock_and_stable_preview_hard_gates(self) -> None:
        snapshot = RequirementMatrixSnapshotV1.model_validate(_snapshot())
        cases = (
            (
                "unlocked",
                {"requirement_set_status": "REVIEWING"},
                "REQUIREMENT_SET_NOT_LOCKED",
            ),
            (
                "unstable-preview",
                {"preview_generation_id": ""},
                "PREVIEW_NOT_STABLE",
            ),
        )

        for label, kwargs, expected_code in cases:
            with self.subTest(label=label):
                with self.assertRaises(ApplicationError) as raised:
                    self._execute_application_export(snapshot, **kwargs)
                self.assertIn(
                    expected_code,
                    {
                        item["code"]
                        for item in raised.exception.detail["readiness"][
                            "hard_blockers"
                        ]
                    },
                )

    def test_formal_export_blocks_each_missing_effective_conclusion_state(
        self,
    ) -> None:
        base = _snapshot()
        cases: list[tuple[str, dict]] = []

        pending = deepcopy(base)
        pending_atom = pending["nodes"][5]
        pending_atom["analysis_result_state"] = "NONE"
        pending_atom["analysis_status"] = "NOT_STARTED"
        pending_atom["ai_assessment"] = None
        pending_atom["compliance_status"] = ""
        pending["counts"]["assessed_nodes"] = 1
        cases.append(("pending", pending))

        reason_only = deepcopy(base)
        reason_atom = reason_only["nodes"][5]
        reason_atom["analysis_result_state"] = "CURRENT"
        reason_atom["ai_assessment"]["decision"] = ""
        reason_atom["ai_assessment"]["compliance_status"] = ""
        reason_atom["compliance_status"] = ""
        cases.append(("reason-only", reason_only))

        stale = deepcopy(base)
        stale_atom = stale["nodes"][5]
        stale_atom["analysis_result_state"] = "STALE_REFERENCE"
        stale_atom["analysis_status"] = "NOT_STARTED"
        stale_atom["stale_ai_assessment"] = stale_atom.pop("ai_assessment")
        stale_atom["compliance_status"] = ""
        stale["counts"]["assessed_nodes"] = 1
        stale["counts"]["stale_result_nodes"] = 1
        cases.append(("stale", stale))

        for label, payload in cases:
            with self.subTest(label=label):
                snapshot = RequirementMatrixSnapshotV1.model_validate(payload)
                with self.assertRaises(ApplicationError) as raised:
                    self._execute_application_export(snapshot)
                blockers = raised.exception.detail["readiness"]["hard_blockers"]
                self.assertEqual(
                    {item["code"] for item in blockers},
                    {"MISSING_EFFECTIVE_CONCLUSIONS"},
                )
                self.assertEqual(blockers[0]["count"], 1)

    def test_formal_export_allows_a_trusted_matrix_with_zero_analysis_targets(
        self,
    ) -> None:
        payload = _snapshot()
        for node in payload["nodes"]:
            if node["disposition"] != "ANALYZE":
                continue
            node["disposition"] = "RETAIN"
            node["analysis_eligibility"] = "NOT_ELIGIBLE"
            node["analysis_status"] = "NOT_APPLICABLE"
            node["analysis_result_state"] = "NONE"
            node["ai_assessment"] = None
            node["manual_decision"] = None
            node["compliance_status"] = ""
        payload["counts"].update(
            {
                "analyze_nodes": 0,
                "retained_nodes": 8,
                "analysis_eligible_nodes": 0,
                "assessed_nodes": 0,
                "decided_nodes": 0,
            }
        )
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        result, _ = self._execute_application_export(snapshot)

        self.assertEqual(result["decision_eligible_count"], 0)
        self.assertEqual(result["effective_conclusion_count"], 0)
        self.assertEqual(result["effective_conclusion_coverage"], 1.0)

    def test_rich_table_refs_join_multi_table_nodes_by_stable_physical_id(
        self,
    ) -> None:
        snapshot = _snapshot()
        table_node = next(
            node for node in snapshot["nodes"] if node["node_id"] == "table"
        )
        table_node["logical_table_ids"] = [
            "logical-table-1",
            "logical-table-2",
        ]
        for ref in table_node["table_refs"]:
            ref.pop("logical_table_id", None)
        third_part = _table_part(
            "physical-table-3",
            source_order=2,
            page=3,
        )
        third_part.pop("logical_table_id", None)
        table_node["table_refs"].append(third_part)
        snapshot["logical_tables"].append(
            {
                "logical_table_id": "logical-table-2",
                "projection_id": "projection-2",
                "projection_mode": "WHOLE_TABLE",
                "record_count": 0,
                "physical_table_ids": ["physical-table-3"],
                "node_ids": ["table"],
                "evidence_binding_ids": [],
                "projection": {
                    "projection_id": "projection-2",
                    "mode": "WHOLE_TABLE",
                    "table_fact": {
                        "logical_table_id": "logical-table-2",
                        "table_parts": [
                            {
                                "table_id": "physical-table-3",
                                "source_order": 2,
                            }
                        ],
                    },
                    "records": [],
                },
                "evidence_bindings": [],
            }
        )

        catalog = _build_matrix_table_catalog(
            snapshot,
            _visible_matrix_nodes(snapshot),
        )

        self.assertEqual(
            {
                part.physical_table_id
                for part in catalog.parts_by_logical_id["logical-table-1"]
            },
            {"physical-table-1", "physical-table-2"},
        )
        self.assertEqual(
            {
                part.physical_table_id
                for part in catalog.parts_by_logical_id["logical-table-2"]
            },
            {"physical-table-3"},
        )

    def test_application_exports_and_persists_the_same_complete_snapshot(self) -> None:
        snapshot = RequirementMatrixSnapshotV1.model_validate(_snapshot())
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_path = root / "source.txt"
            source_path.write_text("source", encoding="utf-8")
            project = {
                "id": "project-1",
                "name": "ReqPilotX",
                "customer_name": "Customer",
            }
            document = {
                "id": "document-1",
                "project_id": "project-1",
                "file_name": "source.txt",
                "file_path": str(source_path),
                "active_requirement_set_id": "set-1",
            }
            requirement_set = {
                "id": "set-1",
                "project_id": "project-1",
                "sor_document_id": "document-1",
                "status": "REVIEWING",
                "version": 3,
                "source_hash": "a" * 64,
            }
            captured: dict = {}

            def create_export_job(**kwargs):
                captured.update(kwargs)
                return {"id": "export-1"}

            with (
                patch(
                    "app.application.projects._require_project",
                    return_value=project,
                ),
                patch.object(exports_app.db, "get_sor_document", return_value=document),
                patch.object(
                    exports_app.db, "get_requirement_set", return_value=requirement_set
                ),
                patch.object(
                    exports_app.db,
                    "get_latest_completed_ai_analysis_job_for_set",
                    return_value=None,
                ),
                patch.object(
                    exports_app.db,
                    "create_export_job",
                    side_effect=create_export_job,
                ),
                patch(
                    "app.application.requirement_matrix.build_requirement_matrix_snapshot",
                    return_value=snapshot,
                ),
                patch("app.application.runtime.EXPORTS_DIR", root / "exports"),
                patch.object(exports_app, "_is_desktop_embedded", return_value=False),
            ):
                result = exports_app._execute_project_export(
                    "project-1",
                    ExportRequest(
                        document_id="document-1",
                        requirement_set_id="set-1",
                        mode="DRAFT",
                    ),
                )

            self.assertEqual(result["decision_eligible_count"], 2)
            self.assertEqual(result["human_decision_count"], 1)
            self.assertEqual(result["human_coverage"], 0.5)
            stored_snapshot = _parse_matrix_export_snapshot(captured["snapshot_json"])
            self.assertIsNotNone(stored_snapshot)
            self.assertEqual(len(stored_snapshot["nodes"]), 9)
            self.assertEqual(
                stored_snapshot["nodes"][4]["manual_decision"]["decision"],
                "ACCEPT",
            )
            RequirementMatrixSnapshotV1.model_validate(stored_snapshot)
            workbook = load_workbook(captured["file_path"])
            self.assertEqual(workbook.sheetnames[0], OVERVIEW_SHEET)
            self.assertEqual(workbook.sheetnames[1], MATRIX_SHEET)
            self.assertRegex(workbook.sheetnames[2], r"^T001_9\.16\.2")
            self.assertEqual(workbook.sheetnames[-1], "_ReqPilotXMeta")

    def test_matrix_is_one_hierarchy_preserving_workbook(self) -> None:
        snapshot = RequirementMatrixSnapshotV1.model_validate(_snapshot())
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
                export_mode="FORMAL",
            )

            workbook = load_workbook(file_path)
            self.assertEqual(workbook.sheetnames[0], OVERVIEW_SHEET)
            self.assertEqual(workbook.sheetnames[1], MATRIX_SHEET)
            self.assertRegex(workbook.sheetnames[2], r"^T001_9\.16\.2")
            self.assertEqual(workbook.sheetnames[-1], "_ReqPilotXMeta")
            self.assertEqual(workbook["_ReqPilotXMeta"].sheet_state, "hidden")
            for removed in (
                "ReviewRecords",
                "TableIndex",
                "AnalysisDetails",
                "VersionMeta",
                "DiffVsBaseline",
            ):
                self.assertNotIn(removed, workbook.sheetnames)

            overview = workbook[OVERVIEW_SHEET]
            matrix = workbook[MATRIX_SHEET]
            self.assertEqual(
                [overview.cell(4, column).value for column in range(1, 12)],
                [
                    "对象 / Req ID",
                    "原文编号",
                    "所属章节",
                    "标题",
                    "需求数",
                    "是否分析",
                    "分析进度",
                    "当前结论",
                    "客户确认",
                    "原文 / 附件",
                    "打开 Matrix",
                ],
            )
            self.assertEqual(overview.freeze_panes, "E5")
            self.assertEqual(overview.page_setup.orientation, "landscape")
            self.assertEqual(overview.page_setup.scale, 85)
            self.assertFalse(bool(overview.sheet_properties.pageSetUpPr.fitToPage))
            self.assertEqual(overview.row_dimensions[3].height, 44.0)
            self.assertEqual(
                [matrix.cell(3, column).value for column in range(1, 11)],
                [
                    "Req ID",
                    "原文编号",
                    "章节 / 标题",
                    "需求正文",
                    "原文 / 附件",
                    "是否分析",
                    "分析进度",
                    "当前结论",
                    "分析依据 / 结论来源",
                    "客户确认",
                ],
            )
            identity_rows = [
                row
                for row in range(4, matrix.max_row + 1)
                if any(
                    matrix.cell(row=row, column=column).value not in (None, "")
                    for column in range(1, 11)
                )
            ]
            req_ids = [
                matrix.cell(row=row, column=MATRIX_REQ_ID_COLUMN).value
                for row in identity_rows
            ]
            numbers = [
                matrix.cell(row=row, column=MATRIX_SOURCE_NUMBER_COLUMN).value
                for row in identity_rows
            ]
            titles = [
                matrix.cell(row=row, column=MATRIX_TITLE_COLUMN).value
                for row in identity_rows
            ]
            contents = [
                matrix.cell(row=row, column=MATRIX_BODY_COLUMN).value
                for row in identity_rows
            ]
            self.assertEqual(
                req_ids,
                [None, None, None, None, "REQ-001", "REQ-002", None, None],
            )
            self.assertEqual(
                numbers,
                [
                    "9",
                    "9.16.2",
                    "9.16.2",
                    None,
                    None,
                    None,
                    None,
                    None,
                ],
            )
            self.assertEqual(
                contents,
                [
                    None,
                    "电气性能应满足下表。",
                    "跨页性能参数原表",
                    DEFAULT_RECORD_BODY,
                    "设备工作电压应为 12 V。",
                    "电压容差应在允许范围内。",
                    "本段仅作为分析上下文。",
                    None,
                ],
            )
            self.assertIn("§9", titles[0])
            self.assertIn("技术要求", titles[0])
            self.assertIn("§9 › 9.16.2", titles[1])
            self.assertIn("电气性能", titles[1])
            self.assertIn("性能参数表", titles[2])
            self.assertIn("工作电压", titles[4])
            self.assertIn("电压容差", titles[5])
            self.assertIn("人工排除内容", titles[7])
            customer_facing_text = "\n".join(
                str(value or "")
                for row in range(4, matrix.max_row + 1)
                for value in (
                    matrix.cell(row=row, column=MATRIX_REQ_ID_COLUMN).value,
                    matrix.cell(row=row, column=MATRIX_SOURCE_NUMBER_COLUMN).value,
                    matrix.cell(row=row, column=MATRIX_TITLE_COLUMN).value,
                    matrix.cell(row=row, column=MATRIX_BODY_COLUMN).value,
                )
            )
            self.assertIn("REQ-001", customer_facing_text)
            self.assertIn("REQ-002", customer_facing_text)
            self.assertNotIn(" | ", customer_facing_text)
            self.assertNotIn("封面", customer_facing_text)
            self.assertIn("人工排除内容", customer_facing_text)
            self.assertEqual(
                [matrix.row_dimensions[row].outlineLevel for row in identity_rows],
                [0, 1, 2, 3, 4, 4, 2, 1],
            )
            atom_1_row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            atom_2_row = _find_row(
                matrix, "电压容差应在允许范围内。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(atom_1_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "人工已修改",
            )
            self.assertEqual(
                matrix.cell(atom_2_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "AI 结论已采用",
            )
            self.assertEqual(
                matrix.cell(4, MATRIX_ANALYSIS_SCOPE_COLUMN).value,
                "不分析\n仅保留",
            )
            self.assertEqual(
                matrix.cell(atom_1_row, MATRIX_ANALYSIS_SCOPE_COLUMN).value,
                "需要分析",
            )
            self.assertEqual(
                matrix.cell(atom_1_row, MATRIX_CONCLUSION_COLUMN).value,
                "接受",
            )
            self.assertIn(
                "工程师依据：工程师确认",
                matrix.cell(atom_1_row, MATRIX_BASIS_COLUMN).value,
            )
            detail_text = _sheet_text(workbook["全文详情"])
            for expected_detail in (
                "置信度：93%",
                "知识库未检索到",
                "请客户确认允许的电压容差",
                "人工覆盖 AI 结论",
                "AI 原结论：接受",
                "AI 结论（默认采用）",
                "未人工修改",
            ):
                self.assertIn(expected_detail, detail_text)
            self.assertEqual(
                matrix.cell(atom_1_row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )
            self.assertEqual(
                matrix.cell(atom_2_row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )
            self.assertIsNotNone(
                matrix.cell(atom_1_row, MATRIX_SOURCE_ENTRY_COLUMN).hyperlink
            )
            self.assertIsNotNone(
                matrix.cell(atom_2_row, MATRIX_SOURCE_ENTRY_COLUMN).hyperlink
            )
            excluded_row = next(
                row
                for row in identity_rows
                if "人工排除内容"
                in str(matrix.cell(row, MATRIX_TITLE_COLUMN).value or "")
            )
            self.assertEqual(
                matrix.cell(excluded_row, MATRIX_ANALYSIS_SCOPE_COLUMN).value,
                "不分析\n已排除",
            )
            self.assertEqual(
                matrix.cell(excluded_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "已排除",
            )
            self.assertIsNone(matrix.cell(excluded_row, MATRIX_CONCLUSION_COLUMN).value)
            self.assertIn(
                "处置依据：客户确认该条款不适用于本项目。",
                matrix.cell(excluded_row, MATRIX_BASIS_COLUMN).value,
            )
            self.assertIn(
                "排除原因：客户确认该条款不适用于本项目。",
                detail_text,
            )
            self.assertEqual(matrix.auto_filter.ref, f"A3:J{matrix.max_row}")
            self.assertEqual(matrix.freeze_panes, "D4")
            self.assertFalse(matrix.column_dimensions["H"].hidden)
            self.assertFalse(matrix.column_dimensions["I"].hidden)
            self.assertFalse(
                any(
                    matrix.row_dimensions[row].hidden
                    for row in range(4, matrix.max_row + 1)
                )
            )

            meta = workbook["_ReqPilotXMeta"]
            self.assertEqual(
                meta.cell(_find_row(meta, "ExportMode"), 2).value,
                "FORMAL",
            )
            self.assertEqual(
                meta.cell(
                    _find_row(meta, "CustomerConfirmationCapability"),
                    2,
                ).value,
                "UNAVAILABLE_IN_SNAPSHOT_V1",
            )
            self.assertEqual(
                meta.cell(_find_row(meta, "AnalysisOutputCapability"), 2).value,
                "AVAILABLE",
            )
            meta_header_row = _find_row(meta, "NodeId")
            meta_headers = {
                str(meta.cell(meta_header_row, column).value): column
                for column in range(1, meta.max_column + 1)
                if meta.cell(meta_header_row, column).value
            }
            self.assertTrue(
                {
                    "NodeId",
                    "VisibleParentId",
                    "VisibleDepth",
                    "SourceOrder",
                    "MatrixVisibility",
                    "Disposition",
                    "AnalysisEligibility",
                    "AnalysisStatus",
                    "AnalysisResultState",
                    "EffectiveConclusion",
                    "ConclusionSource",
                    "OverridesAi",
                    "FrozenParentId",
                    "MatrixRowStart",
                    "MatrixRowEnd",
                    "RequirementEntryId",
                    "TableRecordId",
                    "SaveRevision",
                    "UpdatedAt",
                }.issubset(meta_headers)
            )
            meta_rows = []
            for row in range(meta_header_row + 1, meta.max_row + 1):
                node_id = meta.cell(row, meta_headers["NodeId"]).value
                if not node_id:
                    break
                meta_rows.append(
                    {
                        header: meta.cell(row, column).value
                        for header, column in meta_headers.items()
                    }
                )
            self.assertEqual(
                [row["NodeId"] for row in meta_rows],
                [
                    "root",
                    "clause",
                    "table",
                    "record",
                    "atom-1",
                    "atom-2",
                    "context",
                    "excluded",
                ],
            )
            atom_1_meta = next(row for row in meta_rows if row["NodeId"] == "atom-1")
            atom_2_meta = next(row for row in meta_rows if row["NodeId"] == "atom-2")
            excluded_meta = next(
                row for row in meta_rows if row["NodeId"] == "excluded"
            )
            self.assertEqual(atom_1_meta["EffectiveConclusion"], "ACCEPT")
            self.assertEqual(atom_1_meta["ConclusionSource"], "HUMAN_OVERRIDE")
            self.assertTrue(atom_1_meta["OverridesAi"])
            self.assertEqual(atom_2_meta["EffectiveConclusion"], "NEED_CLARIFICATION")
            self.assertEqual(atom_2_meta["ConclusionSource"], "AI_DEFAULT")
            self.assertFalse(atom_2_meta["OverridesAi"])
            self.assertEqual(excluded_meta["MatrixVisibility"], "HUMAN_EXCLUDED")
            self.assertEqual(excluded_meta["Disposition"], "EXCLUDE")
            self.assertEqual(excluded_meta["VisibleParentId"], "root")
            self.assertEqual(excluded_meta["VisibleDepth"], 1)

    def test_stale_analysis_is_exported_as_reanalysis_not_current_result(
        self,
    ) -> None:
        snapshot = _snapshot()
        stale = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        stale["analysis_result_state"] = "STALE_REFERENCE"
        stale["analysis_status"] = "NOT_STARTED"
        stale["stale_ai_assessment"] = stale.pop("ai_assessment")
        stale["stale_manual_decision"] = stale.pop("manual_decision")
        stale["compliance_status"] = ""
        snapshot["counts"].update(
            {
                "assessed_nodes": 1,
                "decided_nodes": 0,
                "stale_result_nodes": 1,
            }
        )

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            stale_row = _find_row(
                matrix,
                "设备工作电压应为 12 V。",
                column=MATRIX_BODY_COLUMN,
            )
            current_row = _find_row(
                matrix,
                "电压容差应在允许范围内。",
                column=MATRIX_BODY_COLUMN,
            )
            self.assertEqual(
                matrix.cell(stale_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "需重新分析",
            )
            self.assertIsNone(matrix.cell(stale_row, MATRIX_CONCLUSION_COLUMN).value)
            self.assertIn(
                "历史结果已失效",
                matrix.cell(stale_row, MATRIX_BASIS_COLUMN).value,
            )
            self.assertEqual(
                matrix.cell(current_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "AI 结论已采用",
            )
            self.assertEqual(
                matrix.cell(current_row, MATRIX_CONCLUSION_COLUMN).value,
                "需要确认",
            )
            self.assertIn(
                "待确认问题：",
                _sheet_text(workbook["全文详情"]),
            )
            self.assertIn(
                "全文详情",
                str(
                    matrix.cell(
                        current_row, MATRIX_SOURCE_ENTRY_COLUMN
                    ).hyperlink.location
                    or ""
                ),
            )
            table_sheet = _table_sheet(workbook)
            rendered_table_text = "\n".join(
                str(cell.value or "")
                for source_row in table_sheet.iter_rows()
                for cell in source_row
            )
            self.assertNotIn("历史分析结果已失效", rendered_table_text)
            self.assertNotIn("表内记录与分析", rendered_table_text)

            meta = workbook["_ReqPilotXMeta"]
            header_row = _find_row(meta, "NodeId")
            headers = {
                str(meta.cell(header_row, column).value): column
                for column in range(1, meta.max_column + 1)
                if meta.cell(header_row, column).value
            }
            stale_meta_row = _find_row(meta, "atom-1")
            self.assertEqual(
                meta.cell(stale_meta_row, headers["AnalysisResultState"]).value,
                "STALE_REFERENCE",
            )
            self.assertIsNone(
                meta.cell(stale_meta_row, headers["EffectiveConclusion"]).value
            )
            self.assertEqual(
                meta.cell(stale_meta_row, headers["ConclusionSource"]).value,
                "STALE_REFERENCE",
            )

    def test_human_native_conclusion_is_not_presented_as_an_ai_override(self) -> None:
        payload = _snapshot()
        native = next(node for node in payload["nodes"] if node["node_id"] == "atom-1")
        native["ai_assessment"] = None
        native["compliance_status"] = ""
        native["analysis_result_state"] = "CURRENT"
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            native_row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(native_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "人工已录入",
            )
            self.assertEqual(
                matrix.cell(native_row, MATRIX_CONCLUSION_COLUMN).value,
                "接受",
            )
            self.assertIn(
                "工程师依据：工程师确认",
                matrix.cell(native_row, MATRIX_BASIS_COLUMN).value,
            )
            self.assertIn(
                "人工录入\n原因：工程师确认",
                _sheet_text(workbook["全文详情"]),
            )
            self.assertNotIn(
                "工程师确认",
                "\n".join(
                    str(cell.value or "")
                    for source_row in _table_sheet(workbook).iter_rows()
                    for cell in source_row
                ),
            )

    def test_reason_only_ai_does_not_turn_human_native_into_override(self) -> None:
        payload = _snapshot()
        native = next(node for node in payload["nodes"] if node["node_id"] == "atom-1")
        native["ai_assessment"]["decision"] = ""
        native["ai_assessment"]["compliance_status"] = ""
        native["compliance_status"] = "COMPLIANT"
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            native_row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(native_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "人工已录入",
            )
            self.assertEqual(
                matrix.cell(native_row, MATRIX_CONCLUSION_COLUMN).value,
                "接受",
            )
            self.assertIn(
                "工程师依据：工程师确认",
                matrix.cell(native_row, MATRIX_BASIS_COLUMN).value,
            )
            self.assertIn(
                "人工录入\n原因：工程师确认",
                _sheet_text(workbook["全文详情"]),
            )
            self.assertNotIn(
                "有效结论：ACCEPT",
                "\n".join(
                    str(cell.value or "")
                    for source_row in _table_sheet(workbook).iter_rows()
                    for cell in source_row
                ),
            )

            meta = workbook["_ReqPilotXMeta"]
            header_row = _find_row(meta, "NodeId")
            headers = {
                str(meta.cell(header_row, column).value): column
                for column in range(1, meta.max_column + 1)
                if meta.cell(header_row, column).value
            }
            native_meta_row = _find_row(meta, "atom-1")
            self.assertEqual(
                meta.cell(native_meta_row, headers["ConclusionSource"]).value,
                "HUMAN_NATIVE",
            )
            self.assertFalse(
                meta.cell(native_meta_row, headers["OverridesAi"]).value,
            )

    def test_human_override_keeps_distinct_ai_and_effective_conclusions(self) -> None:
        payload = _snapshot()
        overridden = next(
            node for node in payload["nodes"] if node["node_id"] == "atom-1"
        )
        overridden["manual_decision"]["decision"] = "DEVIATION_ACCEPT"
        overridden["manual_decision"]["reason"] = "客户书面接受偏差"
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            overridden_row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(overridden_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "人工已修改",
            )
            self.assertEqual(
                matrix.cell(overridden_row, MATRIX_CONCLUSION_COLUMN).value,
                "偏差",
            )
            self.assertIn(
                "工程师依据：客户书面接受偏差",
                matrix.cell(overridden_row, MATRIX_BASIS_COLUMN).value,
            )
            self.assertEqual(
                matrix.cell(overridden_row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )
            detail_text = _sheet_text(workbook["全文详情"])
            self.assertIn(
                "人工覆盖 AI 结论\n原因：客户书面接受偏差",
                detail_text,
            )
            self.assertIn("AI 原结论：接受", detail_text)
            self.assertNotIn(
                "有效结论：DEVIATION_ACCEPT",
                "\n".join(
                    str(cell.value or "")
                    for source_row in _table_sheet(workbook).iter_rows()
                    for cell in source_row
                ),
            )

    def test_pending_analysis_has_no_effective_or_ai_conclusion(self) -> None:
        payload = _snapshot()
        pending = next(node for node in payload["nodes"] if node["node_id"] == "atom-2")
        pending["analysis_status"] = "NOT_STARTED"
        pending["analysis_result_state"] = "NONE"
        pending["ai_assessment"] = None
        pending["compliance_status"] = ""
        payload["counts"]["assessed_nodes"] = 1
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            pending_row = _find_row(
                matrix, "电压容差应在允许范围内。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(pending_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "待分析",
            )
            for column in (MATRIX_CONCLUSION_COLUMN, MATRIX_BASIS_COLUMN):
                self.assertIsNone(matrix.cell(pending_row, column).value)
            self.assertEqual(
                matrix.cell(pending_row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )

    def test_reason_only_analysis_is_exported_as_incomplete_not_pending(
        self,
    ) -> None:
        payload = _snapshot()
        reason_only = next(
            node for node in payload["nodes"] if node["node_id"] == "atom-2"
        )
        reason_only["analysis_result_state"] = "CURRENT"
        reason_only["analysis_status"] = "ASSESSED"
        reason_only["ai_assessment"]["decision"] = ""
        reason_only["ai_assessment"]["compliance_status"] = ""
        reason_only["compliance_status"] = "COMPLIANT"
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            reason_only_row = _find_row(
                matrix,
                "电压容差应在允许范围内。",
                column=MATRIX_BODY_COLUMN,
            )
            self.assertEqual(
                matrix.cell(reason_only_row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "结论不完整",
            )
            for column in (MATRIX_CONCLUSION_COLUMN, MATRIX_BASIS_COLUMN):
                self.assertIsNone(matrix.cell(reason_only_row, column).value)
            self.assertEqual(
                matrix.cell(reason_only_row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )

    def test_system_wrapper_is_hidden_and_promotes_business_children(self) -> None:
        snapshot = _snapshot()
        system = next(node for node in snapshot["nodes"] if node["node_id"] == "system")
        system["child_ids"] = ["promoted"]
        snapshot["nodes"].append(
            _node(
                "promoted",
                parent_id="system",
                child_ids=[],
                depth=2,
                source_order=8,
                display_role="CLAUSE",
                label="被系统包装层包住的业务正文",
            )
        )

        visible = _visible_matrix_nodes(snapshot)

        self.assertNotIn("system", [node["node_id"] for node in visible])
        promoted = next(node for node in visible if node["node_id"] == "promoted")
        self.assertEqual(promoted["_visible_parent_id"], "root")
        self.assertEqual(promoted["depth"], 1)
        self.assertLess(
            [node["node_id"] for node in visible].index("promoted"),
            [node["node_id"] for node in visible].index("excluded"),
        )

    def test_export_handles_a_deep_ready_matrix_without_recursion(self) -> None:
        node_count = 1_105
        nodes = [
            _node(
                f"deep-{index}",
                parent_id=(f"deep-{index - 1}" if index else None),
                child_ids=([f"deep-{index + 1}"] if index + 1 < node_count else []),
                depth=index,
                source_order=index,
                display_role="STRUCTURE",
                label=f"深层结构 {index}",
            )
            for index in range(node_count)
        ]
        snapshot = RequirementMatrixSnapshotV1.model_validate(
            {
                "state": "READY",
                "project_id": "project-deep",
                "document_id": "document-deep",
                "requirement_set_id": "set-deep",
                "requirement_set_status": "LOCKED",
                "tree_sha256": "f" * 64,
                "root_node_ids": ["deep-0"],
                "nodes": nodes,
                "logical_tables": [],
                "counts": {
                    "total_nodes": node_count,
                    "visible_nodes": node_count,
                    "system_non_business_nodes": 0,
                    "human_excluded_nodes": 0,
                    "analyze_nodes": 0,
                    "retained_nodes": node_count,
                    "excluded_nodes": 0,
                    "analysis_eligible_nodes": 0,
                    "assessed_nodes": 0,
                    "decided_nodes": 0,
                    "stale_result_nodes": 0,
                    "logical_tables": 0,
                },
            }
        )

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path, read_only=True)
            meta = workbook["_ReqPilotXMeta"]
            header_row = _find_row(meta, "NodeId")
            exported_node_ids = [
                row[0]
                for row in meta.iter_rows(
                    min_row=header_row + 1,
                    max_row=header_row + node_count,
                    min_col=1,
                    max_col=1,
                    values_only=True,
                )
            ]
            self.assertEqual(
                exported_node_ids,
                [f"deep-{index}" for index in range(node_count)],
            )

    def test_workbook_projection_matches_every_business_node_and_promotes_system_children(
        self,
    ) -> None:
        payload = _snapshot()
        system = next(node for node in payload["nodes"] if node["node_id"] == "system")
        system["child_ids"] = ["promoted"]
        payload["nodes"].append(
            _node(
                "promoted",
                parent_id="system",
                child_ids=[],
                depth=2,
                source_order=8,
                display_role="CLAUSE",
                label="被系统包装层包住的业务正文",
                disposition="RETAIN",
            )
        )
        payload["counts"]["total_nodes"] += 1
        payload["counts"]["visible_nodes"] += 1
        payload["counts"]["retained_nodes"] += 1
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)
        normalized = snapshot.model_dump(mode="json")
        expected_nodes = _visible_matrix_nodes(normalized)

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )

            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            meta = workbook["_ReqPilotXMeta"]
            header_row = _find_row(meta, "NodeId")
            headers = {
                str(meta.cell(header_row, column).value): column
                for column in range(1, meta.max_column + 1)
                if meta.cell(header_row, column).value
            }
            actual_rows: list[dict[str, object]] = []
            for row in range(header_row + 1, meta.max_row + 1):
                if not meta.cell(row, headers["NodeId"]).value:
                    break
                actual_rows.append(
                    {
                        name: meta.cell(row, column).value
                        for name, column in headers.items()
                    }
                )

            self.assertEqual(
                [row["NodeId"] for row in actual_rows],
                [node["node_id"] for node in expected_nodes],
            )
            self.assertNotIn("system", [row["NodeId"] for row in actual_rows])
            matrix_rows: list[int] = []
            for expected, actual in zip(expected_nodes, actual_rows, strict=True):
                conclusion = matrix_conclusion_contract(expected)
                expected_contract = {
                    "VisibleParentId": expected.get("_visible_parent_id") or None,
                    "VisibleDepth": int(expected.get("depth") or 0),
                    "SourceOrder": int(expected.get("source_order") or 0),
                    "MatrixVisibility": expected.get("matrix_visibility") or None,
                    "Disposition": expected.get("disposition") or None,
                    "DispositionReason": expected.get("disposition_reason") or None,
                    "AnalysisEligibility": expected.get("analysis_eligibility") or None,
                    "AnalysisStatus": expected.get("analysis_status") or None,
                    "AnalysisResultState": expected.get("analysis_result_state")
                    or None,
                    "EffectiveConclusion": conclusion.effective_conclusion or None,
                    "ConclusionSource": conclusion.conclusion_source.value,
                    "OverridesAi": conclusion.overrides_ai,
                    "FrozenParentId": expected.get("parent_id") or None,
                }
                self.assertEqual(
                    {key: actual[key] for key in expected_contract},
                    expected_contract,
                    msg=f"Excel parity mismatch for {expected['node_id']}",
                )
                matrix_row = int(actual["MatrixRowStart"])
                matrix_rows.append(matrix_row)
                self.assertTrue(
                    str(
                        matrix.cell(matrix_row, MATRIX_TITLE_COLUMN).value
                        or matrix.cell(matrix_row, MATRIX_BODY_COLUMN).value
                        or ""
                    ).strip()
                )

            self.assertEqual(matrix_rows, sorted(matrix_rows))
            self.assertEqual(len(matrix_rows), len(set(matrix_rows)))
            promoted_row = next(
                row for row in actual_rows if row["NodeId"] == "promoted"
            )
            self.assertEqual(promoted_row["VisibleParentId"], "root")
            self.assertEqual(promoted_row["VisibleDepth"], 1)
            self.assertEqual(
                str(
                    matrix.cell(
                        int(promoted_row["MatrixRowStart"]), MATRIX_TITLE_COLUMN
                    ).value
                    or ""
                ).split("\n")[-1],
                "被系统包装层包住的业务正文",
            )
            rendered_contents = {
                value
                for row in matrix_rows
                for value in (
                    matrix.cell(row, MATRIX_TITLE_COLUMN).value,
                    matrix.cell(row, MATRIX_BODY_COLUMN).value,
                )
            }
            rendered_text = "\n".join(str(value or "") for value in rendered_contents)
            self.assertNotIn("封面", rendered_text)
            self.assertIn("人工排除内容", rendered_text)

    def test_missing_source_warning_counts_business_nodes_and_excludes_system(
        self,
    ) -> None:
        payload = _snapshot()
        for node_id in ("root", "context", "excluded"):
            business_node = next(
                node for node in payload["nodes"] if node["node_id"] == node_id
            )
            business_node["page_start"] = None
            business_node["page_end"] = None
            business_node["source_refs"] = []
        system = next(node for node in payload["nodes"] if node["node_id"] == "system")
        system["page_start"] = None
        system["page_end"] = None
        system["source_refs"] = []
        snapshot = RequirementMatrixSnapshotV1.model_validate(payload)

        with self.assertRaises(ApplicationError) as raised:
            self._execute_application_export(
                snapshot,
                requirement_set_status="REVIEWING",
            )

        warnings = raised.exception.detail["readiness"]["warnings"]
        self.assertIn(
            {"code": "MISSING_SOURCE", "count": 3, "hard": False},
            warnings,
        )

    def test_logical_table_sheet_restores_only_authoritative_native_cells(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / "cell.png"
            PILImage.new("RGB", (96, 64), "#00893D").save(image_path)
            image_bytes = image_path.read_bytes()
            image_sha256 = hashlib.sha256(image_bytes).hexdigest()
            source_docx = root / "source.docx"
            with zipfile.ZipFile(source_docx, "w") as package:
                package.writestr("word/media/image1.png", image_bytes)
            source_sha256 = hashlib.sha256(source_docx.read_bytes()).hexdigest()

            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "exports",
                matrix_snapshot=_snapshot(image_sha256=image_sha256),
                source_file_path=source_docx,
                source_sha256=source_sha256,
            )

            workbook = load_workbook(file_path)
            sheet = _table_sheet(workbook)
            self.assertEqual(len(sheet._images), 1)
            merged_values = {
                sheet.cell(region.min_row, region.min_col).value
                for region in sheet.merged_cells.ranges
            }
            self.assertIn("性能要求", merged_values)
            values = [
                sheet.cell(row=row, column=1).value
                for row in range(1, sheet.max_row + 1)
            ]
            self.assertLess(values.index("第一页记录"), values.index("第二页记录"))
            self.assertIn("表注：第 1 页表注", merged_values)
            self.assertIn("表注：第 2 页表注", merged_values)
            self.assertEqual(sheet["A2"].value, "9.16.2 性能参数表")
            self.assertIn("原表保真视图", sheet["A3"].value)
            self.assertEqual(sheet.freeze_panes, "A4")
            self.assertIsNone(sheet.auto_filter.ref)
            self.assertEqual(sheet.page_setup.orientation, "portrait")
            self.assertIn(sheet.page_setup.scale, {None, 100})
            image_anchor = sheet._images[0].anchor._from
            self.assertGreater(image_anchor.rowOff, 30 * 9525)
            self.assertGreater(sheet.row_dimensions[image_anchor.row + 1].height, 70.0)
            rendered_text = "\n".join(
                str(cell.value or "") for row in sheet.iter_rows() for cell in row
            )
            self.assertNotIn("| ---", rendered_text)
            self.assertNotIn("表内记录与分析", rendered_text)
            self.assertNotIn("设备工作电压应为 12 V。", rendered_text)
            self.assertNotIn("AI 原结论", rendered_text)
            matrix = workbook[MATRIX_SHEET]
            record_row = _find_row(
                matrix,
                DEFAULT_RECORD_BODY,
                column=MATRIX_BODY_COLUMN,
            )
            self.assertIn(
                f"'{sheet.title}'!A7",
                str(
                    matrix.cell(
                        record_row, MATRIX_SOURCE_ENTRY_COLUMN
                    ).hyperlink.location
                    or ""
                ),
            )

    def test_native_table_width_drives_row_height_without_fixed_width_whitespace(
        self,
    ) -> None:
        snapshot = _snapshot()
        table = next(node for node in snapshot["nodes"] if node["node_id"] == "table")
        long_value = "供应商应完成接口定义安全分析验证证据交付物和最终确认"
        table["table_refs"][0]["cells"][4]["raw_text"] = long_value
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            sheet = _table_sheet(load_workbook(file_path))
            row = _find_row(sheet, long_value, column=1)
            self.assertGreaterEqual(sheet.column_dimensions["A"].width, 34.0)
            self.assertGreaterEqual(sheet.row_dimensions[row].height, 40.0)
            self.assertLessEqual(sheet.row_dimensions[row].height, 50.0)

    def test_image_only_table_fallback_uses_full_printable_canvas(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pdf = root / "source.pdf"
            document = pymupdf.open()
            page = document.new_page(width=600, height=420)
            page.draw_rect(
                pymupdf.Rect(40, 50, 560, 360),
                color=(0.0, 0.45, 0.15),
                fill=(0.92, 0.98, 0.94),
            )
            document.save(source_pdf)
            document.close()
            snapshot = _snapshot()
            for node in snapshot["nodes"]:
                node["table_refs"] = []
            table = next(
                node for node in snapshot["nodes"] if node["node_id"] == "table"
            )
            table["table_refs"] = [
                {
                    "logical_table_id": "logical-table-1",
                    "table_id": "image-only-table",
                    "source_order": 0,
                    "location": {
                        "format": "pdf",
                        "page_number": 1,
                        "bbox": [40.0, 50.0, 560.0, 360.0],
                    },
                }
            ]
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "exports",
                matrix_snapshot=snapshot,
                source_file_path=source_pdf,
                source_sha256=hashlib.sha256(source_pdf.read_bytes()).hexdigest(),
            )
            sheet = _table_sheet(load_workbook(file_path))
            self.assertTrue(sheet._images)
            self.assertIn("$H$", str(sheet.print_area))
            self.assertEqual(sheet.max_column, 8)
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_setup.scale, 85)

    def test_body_copy_with_direct_number_does_not_become_requirement_title(
        self,
    ) -> None:
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom.update(
            {
                "source_number": "1.1",
                "display_title": "1.1 设备工作电压应为 12 V。",
                "label": "1.1) 设备工作电压应为 12 V。",
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook[MATRIX_SHEET]
            row = _find_row(matrix, "REQ-001", column=MATRIX_REQ_ID_COLUMN)
            self.assertIn("未提供标题", matrix.cell(row, MATRIX_TITLE_COLUMN).value)
            self.assertNotIn(
                "设备工作电压应为 12 V。",
                matrix.cell(row, MATRIX_TITLE_COLUMN).value,
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_BODY_COLUMN).value,
                "设备工作电压应为 12 V。",
            )
            overview = workbook[OVERVIEW_SHEET]
            overview_row = _find_row(overview, "REQ-001", column=1)
            self.assertEqual(overview.cell(overview_row, 4).value, "未提供标题")

    def test_title_column_strips_common_direct_number_separators(self) -> None:
        node = {"source_number": "1.1"}
        for value in ("1.1) 标题", "1.1）标题", "1.1. 标题", "1.1。标题"):
            with self.subTest(value=value):
                self.assertEqual(
                    exporter_module._matrix_title_without_direct_number(node, value),
                    "标题",
                )

    def test_body_copy_detection_strips_direct_number_on_either_side(self) -> None:
        base = {
            "requirement_id": "requirement-1",
            "source_number": "1.1",
        }
        for title, body in (
            ("1.1。正文", "正文"),
            ("正文", "1.1。正文"),
            ("1.1) 正文", "正文"),
        ):
            with self.subTest(title=title, body=body):
                self.assertTrue(
                    exporter_module._matrix_title_matches_owned_body(
                        {**base, "requirement_text": body},
                        title,
                    )
                )

    def test_pdf_cell_geometry_is_bound_to_the_exact_source_hash(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pdf = root / "source.pdf"
            document = pymupdf.open()
            page = document.new_page(width=420, height=300)
            page.draw_rect(
                pymupdf.Rect(235, 118, 340, 195),
                color=(0.0, 0.45, 0.15),
                fill=(0.0, 0.85, 0.35),
            )
            document.save(source_pdf)
            document.close()
            source_sha256 = hashlib.sha256(source_pdf.read_bytes()).hexdigest()
            snapshot = _snapshot()
            table_node = next(
                node for node in snapshot["nodes"] if node["node_id"] == "table"
            )
            first_part = table_node["table_refs"][0]
            first_part["cells"].append(
                {
                    "id": "pdf-visual",
                    "row": 2,
                    "column": 2,
                    "raw_text": "接线示意",
                    "images": [
                        {
                            "id": "pdf-image",
                            "asset_id": "pdf-image",
                            "location": {
                                "format": "pdf",
                                "page_number": 1,
                                "bbox": [235.0, 118.0, 340.0, 195.0],
                            },
                            "rendering": {
                                "mode": "SOURCE_PDF_CROP",
                                "available": True,
                            },
                        }
                    ],
                }
            )
            _, matched_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "matched",
                matrix_snapshot=snapshot,
                source_file_path=source_pdf,
                source_sha256=source_sha256,
            )
            _, rejected_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "rejected",
                matrix_snapshot=snapshot,
                source_file_path=source_pdf,
                source_sha256="0" * 64,
            )
            self.assertEqual(
                len(_table_sheet(load_workbook(matched_path))._images),
                1,
            )
            self.assertEqual(
                len(_table_sheet(load_workbook(rejected_path))._images),
                0,
            )

    def test_source_object_uses_exact_jpg_storage_key_before_pdf_crop(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pdf = root / "source.pdf"
            document = pymupdf.open()
            document.new_page(width=200, height=200)
            document.save(source_pdf)
            document.close()
            source_sha256 = hashlib.sha256(source_pdf.read_bytes()).hexdigest()

            output = BytesIO()
            PILImage.new("RGB", (12, 8), "green").save(output, format="JPEG")
            image_bytes = output.getvalue()
            content_sha256 = hashlib.sha256(image_bytes).hexdigest()
            storage_key = f"source-assets/{content_sha256[:2]}/{content_sha256}.jpg"
            loader_calls: list[tuple[str, str, int]] = []

            def loader(key: str, expected_sha256: str, max_bytes: int) -> bytes:
                loader_calls.append((key, expected_sha256, max_bytes))
                return image_bytes

            context = exporter_module._build_export_source_image_context(
                source_file_path=source_pdf,
                source_sha256=source_sha256,
                source_asset_loader=loader,
            )
            evidence = {
                "content_sha256": content_sha256,
                "media_type": "image/jpeg",
                "storage_key": storage_key,
                "location": {
                    "format": "pdf",
                    "page_number": 1,
                    "bbox": [20.0, 20.0, 80.0, 80.0],
                },
                "rendering": {
                    "mode": "SOURCE_PDF_CROP",
                    "available": True,
                },
            }
            with patch.object(exporter_module, "render_pdf_source_crop") as crop:
                payload = exporter_module._load_image_payload(
                    evidence,
                    source_image_context=context,
                )

            self.assertIsNotNone(payload)
            assert payload is not None
            self.assertEqual(payload[1:], (12, 8))
            self.assertEqual(loader_calls[0][0], storage_key)
            self.assertTrue(loader_calls[0][0].endswith(".jpg"))
            crop.assert_not_called()

    def test_source_object_wrong_bytes_fail_closed_without_source_crop(self) -> None:
        expected_sha256 = hashlib.sha256(b"expected").hexdigest()
        storage_key = f"source-assets/{expected_sha256[:2]}/{expected_sha256}.png"
        context = exporter_module._build_export_source_image_context(
            source_file_path="",
            source_sha256="",
            source_asset_loader=lambda *_args: b"different",
        )

        self.assertIsNone(
            exporter_module._load_image_payload(
                {
                    "content_sha256": expected_sha256,
                    "media_type": "image/png",
                    "storage_key": storage_key,
                },
                source_image_context=context,
            )
        )

    def test_missing_source_object_falls_back_to_exact_nested_pdf_crop(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pdf = root / "source.pdf"
            document = pymupdf.open()
            page = document.new_page(width=240, height=180)
            page.draw_rect(
                pymupdf.Rect(30, 40, 130, 120),
                color=(0.0, 0.45, 0.15),
                fill=(0.0, 0.85, 0.35),
            )
            document.save(source_pdf)
            document.close()
            source_sha256 = hashlib.sha256(source_pdf.read_bytes()).hexdigest()
            missing_sha256 = hashlib.sha256(b"missing-object").hexdigest()
            storage_key = f"source-assets/{missing_sha256[:2]}/{missing_sha256}.png"
            loader_calls: list[str] = []
            context = exporter_module._build_export_source_image_context(
                source_file_path=source_pdf,
                source_sha256=source_sha256,
                source_asset_loader=lambda key, *_args: (
                    loader_calls.append(key) or None
                ),
            )

            payload = exporter_module._load_image_payload(
                {
                    "kind": "image",
                    "source_refs": [
                        {
                            "content_sha256": missing_sha256,
                            "media_type": "image/png",
                            "storage_key": storage_key,
                            "location": {
                                "format": "pdf",
                                "page_number": 1,
                                "bbox": [30.0, 40.0, 130.0, 120.0],
                            },
                        }
                    ],
                    "rendering": {
                        "mode": "SOURCE_PDF_CROP",
                        "available": True,
                    },
                },
                source_image_context=context,
            )

            self.assertIsNotNone(payload)
            assert payload is not None
            self.assertEqual(loader_calls, [storage_key])
            self.assertTrue(payload[0].getvalue().startswith(b"\x89PNG"))
            self.assertGreater(payload[1], 0)
            self.assertGreater(payload[2], 0)

    def test_pdf_figure_is_discoverable_in_overview_matrix_and_detail(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pdf = root / "source.pdf"
            document = pymupdf.open()
            page = document.new_page(width=420, height=300)
            page.draw_rect(
                pymupdf.Rect(72, 126, 290, 238),
                color=(0.0, 0.45, 0.15),
                fill=(0.0, 0.85, 0.35),
            )
            document.save(source_pdf)
            document.close()
            source_sha256 = hashlib.sha256(source_pdf.read_bytes()).hexdigest()
            figure = {
                "kind": "image",
                "source_id": "source-figure-1",
                "source_order": 12,
                "text": "原 PDF 图像区域",
                "location": {
                    "format": "pdf",
                    "page_number": 1,
                    "bbox": [72.0, 126.0, 290.0, 238.0],
                },
            }
            snapshot = _snapshot()
            clause = next(
                node for node in snapshot["nodes"] if node["node_id"] == "clause"
            )
            # RequirementTree may expose the same immutable placement through
            # more than one content path.  Export it once on the owning row.
            clause["figure_refs"] = [figure, dict(figure)]

            _, matched_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "matched",
                matrix_snapshot=snapshot,
                source_file_path=source_pdf,
                source_sha256=source_sha256,
            )
            _, rejected_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "rejected",
                matrix_snapshot=snapshot,
                source_file_path=source_pdf,
                source_sha256="0" * 64,
            )
            _, without_figure_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "without-figure",
                matrix_snapshot=_snapshot(),
                source_file_path=source_pdf,
                source_sha256=source_sha256,
            )

            matched = load_workbook(matched_path)
            rejected = load_workbook(rejected_path)
            without_figure = load_workbook(without_figure_path)
            matched_matrix = matched["ComplianceMatrix"]
            rejected_matrix = rejected["ComplianceMatrix"]
            without_figure_matrix = without_figure["ComplianceMatrix"]
            clause_row = _find_row(
                matched_matrix,
                "电气性能应满足下表。",
                column=MATRIX_BODY_COLUMN,
            )
            owning_row_images = [
                image
                for image in matched_matrix._images
                if image.anchor._from.col == MATRIX_SOURCE_ENTRY_COLUMN - 1
                and image.anchor._from.row == clause_row - 1
            ]
            self.assertEqual(len(owning_row_images), 1)
            self.assertIn(
                "全文详情",
                str(
                    matched_matrix.cell(
                        clause_row, MATRIX_SOURCE_ENTRY_COLUMN
                    ).hyperlink.location
                    or ""
                ),
            )
            self.assertEqual(
                matched_matrix.cell(clause_row, MATRIX_SOURCE_ENTRY_COLUMN).value,
                "查看图片",
            )
            detail = matched["全文详情"]
            detail_figures = [
                image for image in detail._images if image.anchor._from.col == 2
            ]
            self.assertEqual(len(detail_figures), 1)
            self.assertGreater(detail.row_dimensions[4].height, 80.0)
            self.assertGreater(
                matched_matrix.row_dimensions[clause_row].height,
                without_figure_matrix.row_dimensions[clause_row].height,
            )
            overview = matched[OVERVIEW_SHEET]
            self.assertIn("图片证据 1", str(overview["A3"].value))
            self.assertIn("可渲染 1", str(overview["A3"].value))
            self.assertIn("资源不可用/校验失败 0", str(overview["A3"].value))
            self.assertIn("源信息不完整 0", str(overview["A3"].value))
            self.assertIn("可渲染 0", str(rejected[OVERVIEW_SHEET]["A3"].value))
            self.assertIn(
                "资源不可用/校验失败 1",
                str(rejected[OVERVIEW_SHEET]["A3"].value),
            )
            self.assertIn(
                "源信息不完整 0",
                str(rejected[OVERVIEW_SHEET]["A3"].value),
            )
            figure_row = _find_row(overview, "原 PDF 图像区域", column=4)
            self.assertEqual(overview.cell(figure_row, 1).value, "图片")
            self.assertNotIn("REQ", str(overview.cell(figure_row, 1).value))
            self.assertEqual(overview.cell(figure_row, 5).value, 0)
            self.assertEqual(overview.cell(figure_row, 9).value, "不适用")
            self.assertIn("查看图片", overview.cell(figure_row, 10).value)
            # Both workbooks retain only the shared brand image.  A source-hash
            # mismatch must not leak an unverified crop into the detail sheet.
            self.assertEqual(
                len(rejected_matrix._images),
                len(without_figure_matrix._images),
            )
            rejected_detail = rejected["全文详情"]
            self.assertEqual(
                [
                    image
                    for image in rejected_detail._images
                    if image.anchor._from.col == 2
                ],
                [],
            )
            self.assertIn(
                "图片未能通过来源校验",
                "\n".join(
                    str(cell.value or "")
                    for row in rejected_detail.iter_rows()
                    for cell in row
                ),
            )
            self.assertEqual(len(_table_sheet(matched)._images), 0)

    def test_matrix_thumbnail_starts_below_wrapped_source_entry_label(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pdf = root / "source.pdf"
            document = pymupdf.open()
            page = document.new_page(width=420, height=300)
            page.draw_rect(
                pymupdf.Rect(72, 126, 290, 238),
                color=(0.0, 0.45, 0.15),
                fill=(0.0, 0.85, 0.35),
            )
            document.save(source_pdf)
            document.close()
            figure = {
                "kind": "image",
                "source_id": "source-figure-entry-label",
                "text": "需求原图",
                "location": {
                    "format": "pdf",
                    "page_number": 1,
                    "bbox": [72.0, 126.0, 290.0, 238.0],
                },
            }
            snapshot = _snapshot()
            atom = next(
                node for node in snapshot["nodes"] if node["node_id"] == "atom-1"
            )
            atom["figure_refs"] = [figure]
            atom["requirement_text"] = "\n\n".join(
                "供应商应完成系统设计验证并保留可追溯证据。" for _ in range(18)
            )
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=root / "exports",
                matrix_snapshot=snapshot,
                source_file_path=source_pdf,
                source_sha256=hashlib.sha256(source_pdf.read_bytes()).hexdigest(),
            )
            workbook = load_workbook(file_path)
            matrix = workbook[MATRIX_SHEET]
            row = _find_row(matrix, "REQ-001", column=MATRIX_REQ_ID_COLUMN)
            source_cell = matrix.cell(row, MATRIX_SOURCE_ENTRY_COLUMN)
            self.assertEqual(source_cell.value, "全文 + 图片 + 原表")
            self.assertEqual(source_cell.alignment.vertical, "top")
            thumbnail = next(
                image
                for image in matrix._images
                if image.anchor._from.row == row - 1
                and image.anchor._from.col == MATRIX_SOURCE_ENTRY_COLUMN - 1
            )
            self.assertGreaterEqual(thumbnail.anchor._from.rowOff, 50 * 9525)
            self.assertLessEqual(matrix.row_dimensions[row].height, 96.0)

    def test_tree_order_matches_ui_when_child_ids_are_shuffled(self) -> None:
        snapshot = _snapshot()
        by_id = {node["node_id"]: node for node in snapshot["nodes"]}
        by_id["root"]["child_ids"] = ["excluded", "system", "clause"]
        by_id["clause"]["child_ids"] = ["context", "table"]
        by_id["record"]["child_ids"] = ["atom-2", "atom-1"]
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            # The workbook must preserve the frozen child_ids traversal, not
            # silently re-sort siblings by source_order or opaque identity.
            identity_rows = [
                row
                for row in range(4, matrix.max_row + 1)
                if any(
                    matrix.cell(row=row, column=column).value not in (None, "")
                    for column in range(1, 11)
                )
            ]
            self.assertEqual(
                [
                    (
                        matrix.cell(row=row, column=MATRIX_REQ_ID_COLUMN).value,
                        matrix.cell(row=row, column=MATRIX_SOURCE_NUMBER_COLUMN).value,
                        matrix.cell(row=row, column=MATRIX_BODY_COLUMN).value,
                    )
                    for row in identity_rows
                ],
                [
                    (None, "9", None),
                    (None, None, None),
                    (
                        None,
                        "9.16.2",
                        "电气性能应满足下表。",
                    ),
                    (
                        None,
                        None,
                        "本段仅作为分析上下文。",
                    ),
                    (
                        None,
                        "9.16.2",
                        "跨页性能参数原表",
                    ),
                    (
                        None,
                        None,
                        DEFAULT_RECORD_BODY,
                    ),
                    (
                        "REQ-002",
                        None,
                        "电压容差应在允许范围内。",
                    ),
                    (
                        "REQ-001",
                        None,
                        "设备工作电压应为 12 V。",
                    ),
                ],
            )
            ordered_titles = [
                str(matrix.cell(row, MATRIX_TITLE_COLUMN).value or "")
                for row in identity_rows
            ]
            self.assertIn("技术要求", ordered_titles[0])
            self.assertIn("人工排除内容", ordered_titles[1])
            self.assertIn("电气性能", ordered_titles[2])
            self.assertIn("说明", ordered_titles[3])
            self.assertIn("性能参数表", ordered_titles[4])
            self.assertIn("工作电压记录", ordered_titles[5])
            self.assertIn("电压容差", ordered_titles[6])
            self.assertIn("工作电压", ordered_titles[7])

    def test_whole_table_requirement_keeps_req_id_and_original_table_entry(
        self,
    ) -> None:
        snapshot = _snapshot()
        nodes = {
            node["node_id"]: node
            for node in snapshot["nodes"]
            if node["node_id"] not in {"record", "atom-2"}
        }
        nodes["table"]["child_ids"] = ["atom-1"]
        nodes["atom-1"]["parent_id"] = "table"
        nodes["atom-1"]["depth"] = 3
        nodes["atom-1"]["table_record_id"] = None
        snapshot["nodes"] = list(nodes.values())
        logical_table = snapshot["logical_tables"][0]
        logical_table["projection_mode"] = "WHOLE_TABLE"
        logical_table["record_count"] = 0
        logical_table["node_ids"] = ["table", "atom-1"]
        logical_table["projection"]["mode"] = "WHOLE_TABLE"
        logical_table["projection"]["records"] = []
        snapshot["counts"].update(
            {
                "total_nodes": 7,
                "visible_nodes": 5,
                "analyze_nodes": 1,
                "retained_nodes": 5,
                "analysis_eligible_nodes": 1,
                "assessed_nodes": 1,
                "decided_nodes": 1,
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            atom_row = _find_row(
                matrix,
                "设备工作电压应为 12 V。",
                column=MATRIX_BODY_COLUMN,
            )
            self.assertEqual(
                matrix.cell(atom_row, MATRIX_REQ_ID_COLUMN).value,
                "REQ-001",
            )
            self.assertIn(
                "工作电压",
                matrix.cell(atom_row, MATRIX_TITLE_COLUMN).value,
            )
            table_sheet = _table_sheet(workbook)
            rendered = "\n".join(
                str(cell.value or "")
                for source_row in table_sheet.iter_rows()
                for cell in source_row
            )
            self.assertNotIn("整表分析", rendered)
            table_row = _find_row(matrix, "跨页性能参数原表", column=MATRIX_BODY_COLUMN)
            self.assertIn(
                table_sheet.title,
                str(
                    matrix.cell(
                        table_row, MATRIX_SOURCE_ENTRY_COLUMN
                    ).hyperlink.location
                    or ""
                ),
            )

    def test_draft_preserves_existing_manual_decisions(self) -> None:
        snapshot = _snapshot()
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
                export_mode="DRAFT",
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            self.assertEqual(matrix["A8"].value, "REQ-001")
            self.assertIsNone(matrix["B8"].value)
            self.assertIn("工作电压", matrix["C8"].value)
            self.assertEqual(matrix["D8"].value, "设备工作电压应为 12 V。")
            self.assertEqual(matrix["F8"].value, "需要分析")
            self.assertEqual(matrix["G8"].value, "人工已修改")
            self.assertEqual(matrix["H8"].value, "接受")
            self.assertIn("工程师依据：工程师确认", matrix["I8"].value)
            self.assertEqual(matrix["J8"].value, "上游未提供")
            self.assertIsNone(matrix["A4"].value)
            self.assertEqual(matrix["B4"].value, "9")
            self.assertIn("§9", matrix["C4"].value)
            self.assertIn("技术要求", matrix["C4"].value)
            visible_contract = "\n".join(
                str(matrix.cell(row=row, column=column).value or "")
                for row in range(4, 11)
                for column in (1, 2, 3, 4)
            )
            self.assertIn("REQ-001", visible_contract)
            self.assertIn("REQ-002", visible_contract)
            self.assertNotIn(" | ", visible_contract)
            self.assertIn("AI 原结论：接受", _sheet_text(workbook["全文详情"]))

        persisted = _matrix_snapshot_for_export(snapshot)
        self.assertEqual(
            persisted["nodes"][4]["manual_decision"]["decision"],
            "ACCEPT",
        )
        self.assertIsNotNone(snapshot["nodes"][4]["manual_decision"])
        serialized = _dump_matrix_export_snapshot(persisted)
        self.assertEqual(_parse_matrix_export_snapshot(serialized), persisted)

    def test_current_decision_prefers_engineer_without_faking_customer_confirmation(
        self,
    ) -> None:
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom["manual_decision"] = {
            "decision": "REJECT",
            "reason": "客户规范中的目标值尚未满足。",
            "decided_at": "2026-07-31T10:00:00+08:00",
        }
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(matrix.cell(row, MATRIX_CONCLUSION_COLUMN).value, "拒绝")
            self.assertIn(
                "工程师依据：客户规范",
                matrix.cell(row, MATRIX_BASIS_COLUMN).value,
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )
            detail_text = _sheet_text(workbook["全文详情"])
            self.assertIn("AI 原结论：接受", detail_text)
            self.assertIn("人工覆盖 AI 结论", detail_text)
            self.assertIn(
                "原因：客户规范中的目标值尚未满足。",
                detail_text,
            )

    def test_all_four_business_decisions_use_stable_chinese_labels(self) -> None:
        expected = {
            "ACCEPT": "接受",
            "REJECT": "拒绝",
            "DEVIATION_ACCEPT": "偏差",
            "NEED_CLARIFICATION": "需要确认",
        }
        for token, label in expected.items():
            with self.subTest(token=token), TemporaryDirectory() as tmp:
                snapshot = _snapshot()
                atom = next(
                    node for node in snapshot["nodes"] if node["node_id"] == "atom-1"
                )
                atom["analysis_status"] = "ASSESSED"
                atom["manual_decision"] = None
                atom["ai_assessment"]["decision"] = token
                snapshot["counts"]["decided_nodes"] = 0
                _, file_path = export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    matrix_snapshot=snapshot,
                )
                workbook = load_workbook(file_path)
                matrix = workbook["ComplianceMatrix"]
                row = _find_row(
                    matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
                )
                self.assertEqual(
                    matrix.cell(row, MATRIX_CONCLUSION_COLUMN).value,
                    label,
                )
                self.assertIn(
                    f"AI 原结论：{label}",
                    _sheet_text(workbook["全文详情"]),
                )
                self.assertIn(
                    "AI 结论（默认采用）",
                    _sheet_text(workbook["全文详情"]),
                )
                self.assertEqual(
                    matrix.cell(row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                    "上游未提供",
                )

    def test_pending_analysis_keeps_current_decision_and_basis_blank(self) -> None:
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom.update(
            {
                "analysis_status": "NOT_STARTED",
                "ai_assessment": None,
                "manual_decision": None,
                "compliance_status": "",
            }
        )
        snapshot["counts"].update({"assessed_nodes": 1, "decided_nodes": 0})
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            matrix = load_workbook(file_path)["ComplianceMatrix"]
            row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_ANALYSIS_SCOPE_COLUMN).value,
                "需要分析",
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "待分析",
            )
            for column in (MATRIX_CONCLUSION_COLUMN, MATRIX_BASIS_COLUMN):
                self.assertIsNone(matrix.cell(row, column).value)
            self.assertEqual(
                matrix.cell(row, MATRIX_CUSTOMER_CONFIRMATION_COLUMN).value,
                "上游未提供",
            )

    def test_snapshot_without_current_analysis_collapses_empty_output_columns(
        self,
    ) -> None:
        snapshot = _snapshot()
        for node in snapshot["nodes"]:
            node.update(
                {
                    "ai_assessment": None,
                    "manual_decision": None,
                    "compliance_status": "",
                    "disposition_reason": "",
                }
            )
            if node.get("analysis_eligibility") == "ELIGIBLE":
                node["analysis_status"] = "NOT_STARTED"
        snapshot["counts"].update({"assessed_nodes": 0, "decided_nodes": 0})

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
                export_mode="DRAFT",
            )
            workbook = load_workbook(file_path)
            matrix = workbook[MATRIX_SHEET]
            self.assertTrue(matrix.column_dimensions["H"].hidden)
            self.assertTrue(matrix.column_dimensions["I"].hidden)
            self.assertEqual(matrix.column_dimensions["H"].outlineLevel, 1)
            self.assertEqual(matrix.column_dimensions["I"].outlineLevel, 1)
            self.assertTrue(matrix.column_dimensions["J"].collapsed)
            self.assertIn("H:I 已默认折叠", str(matrix["A2"].value))
            meta = workbook["_ReqPilotXMeta"]
            self.assertEqual(
                meta.cell(_find_row(meta, "AnalysisOutputCapability"), 2).value,
                "NO_CURRENT_RESULTS",
            )

    def test_retained_node_does_not_leak_historical_analysis(self) -> None:
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom.update(
            {
                "disposition": "RETAIN",
                "analysis_eligibility": "ELIGIBLE",
                "analysis_status": "NOT_APPLICABLE",
            }
        )
        snapshot["counts"].update(
            {
                "analyze_nodes": 1,
                "retained_nodes": 7,
                "assessed_nodes": 1,
                "decided_nodes": 0,
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_ANALYSIS_SCOPE_COLUMN).value,
                "不分析\n仅保留",
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "无需分析",
            )
            for column in (MATRIX_CONCLUSION_COLUMN, MATRIX_BASIS_COLUMN):
                self.assertIsNone(matrix.cell(row, column).value)
            table_sheet = _table_sheet(workbook)
            rendered = "\n".join(
                str(cell.value or "")
                for source_row in table_sheet.iter_rows()
                for cell in source_row
            )
            self.assertNotIn("证据完整", rendered)
            self.assertNotIn("工程师确认", rendered)

    def test_excluded_node_does_not_leak_historical_analysis(self) -> None:
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom.update(
            {
                "matrix_visibility": "HUMAN_EXCLUDED",
                "disposition": "EXCLUDE",
                "analysis_eligibility": "ELIGIBLE",
                "analysis_status": "NOT_APPLICABLE",
                "disposition_reason": "工程师确认不属于交付范围",
            }
        )
        snapshot["counts"].update(
            {
                "analyze_nodes": 1,
                "excluded_nodes": 2,
                "human_excluded_nodes": 2,
                "visible_nodes": 6,
                "assessed_nodes": 1,
                "decided_nodes": 0,
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            matrix = load_workbook(file_path)["ComplianceMatrix"]
            row = _find_row(
                matrix, "设备工作电压应为 12 V。", column=MATRIX_BODY_COLUMN
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_ANALYSIS_SCOPE_COLUMN).value,
                "不分析\n已排除",
            )
            self.assertEqual(
                matrix.cell(row, MATRIX_ANALYSIS_PROGRESS_COLUMN).value,
                "已排除",
            )
            self.assertIsNone(matrix.cell(row, MATRIX_CONCLUSION_COLUMN).value)
            self.assertEqual(
                matrix.cell(row, MATRIX_BASIS_COLUMN).value,
                "处置依据：工程师确认不属于交付范围",
            )
            self.assertIn(
                "排除原因：工程师确认不属于交付范围",
                _sheet_text(load_workbook(file_path)["全文详情"]),
            )

    def test_generic_whole_table_uses_trusted_owner_title(self) -> None:
        snapshot = _snapshot()
        clause = next(node for node in snapshot["nodes"] if node["node_id"] == "clause")
        table = next(node for node in snapshot["nodes"] if node["node_id"] == "table")
        clause["logical_table_ids"] = ["logical-table-1"]
        snapshot["logical_tables"][0]["node_ids"].append("clause")
        table.update(
            {
                "label": "整表候选",
                "summary": "整表候选",
                "display_title": "整表候选",
                "display_text": "",
                "requirement_text": "整表候选",
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook[MATRIX_SHEET]
            table_row = [
                row
                for row in range(4, matrix.max_row + 1)
                if matrix.cell(row, MATRIX_SOURCE_NUMBER_COLUMN).value == "9.16.2"
                and matrix.cell(row, MATRIX_SOURCE_ENTRY_COLUMN).value == "打开原表"
            ][-1]
            self.assertIn("电气性能", matrix.cell(table_row, MATRIX_TITLE_COLUMN).value)
            self.assertIsNone(matrix.cell(table_row, MATRIX_BODY_COLUMN).value)
            visible_text = (
                _sheet_text(matrix) + "\n" + _sheet_text(workbook[OVERVIEW_SHEET])
            )
            self.assertNotIn("整表候选", visible_text)

    def test_generic_whole_tables_with_same_owner_get_stable_ordinals(self) -> None:
        generic = {
            "display_role": "TABLE",
            "display_title": "整表候选",
            "label": "整表候选",
            "logical_table_ids": ["logical-table-1"],
        }
        table_details = {
            "logical-table-1": {"title": "项目主要信息"},
            "logical-table-2": {"title": "项目主要信息"},
        }
        self.assertEqual(
            exporter_module._matrix_node_content(
                generic,
                table_details=table_details,
            ),
            "整表：项目主要信息（表 1）",
        )
        self.assertEqual(
            exporter_module._matrix_node_content(
                {**generic, "logical_table_ids": ["logical-table-2"]},
                table_details=table_details,
            ),
            "整表：项目主要信息（表 2）",
        )

    def test_hidden_only_logical_table_stays_in_meta_without_visible_sheet(
        self,
    ) -> None:
        snapshot = _snapshot()
        system_node = next(
            node for node in snapshot["nodes"] if node["node_id"] == "system"
        )
        system_node["logical_table_ids"] = ["logical-hidden"]
        system_node["table_refs"] = [
            {
                "logical_table_id": "logical-hidden",
                "table_id": "physical-hidden",
                "cells": [
                    {"id": "hidden-cell", "row": 0, "column": 0, "raw_text": "隐藏表"}
                ],
            }
        ]
        snapshot["logical_tables"].append(
            {
                "logical_table_id": "logical-hidden",
                "projection_id": "projection-hidden",
                "projection_mode": "WHOLE_TABLE",
                "record_count": 0,
                "physical_table_ids": ["physical-hidden"],
                "node_ids": ["system"],
                "evidence_binding_ids": [],
                "projection": {
                    "projection_id": "projection-hidden",
                    "mode": "WHOLE_TABLE",
                    "table_fact": {
                        "logical_table_id": "logical-hidden",
                        "table_parts": [
                            {
                                "table_id": "physical-hidden",
                                "source_order": 0,
                                "source_refs": [{"page_number": 1}],
                            }
                        ],
                    },
                    "records": [],
                },
                "evidence_bindings": [],
            }
        )
        snapshot["counts"]["logical_tables"] = 2
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            visible_table_sheets = [
                name for name in workbook.sheetnames if name.startswith("T")
            ]
            self.assertEqual(len(visible_table_sheets), 1)
            self.assertRegex(visible_table_sheets[0], r"^T001_9\.16\.2")
            meta = workbook["_ReqPilotXMeta"]
            meta_values = [
                meta.cell(row=row, column=1).value for row in range(1, meta.max_row + 1)
            ]
            self.assertIn("logical-hidden", meta_values)

    def test_formal_denominator_requires_all_three_analysis_dimensions(self) -> None:
        base = _snapshot()["nodes"][4]
        candidates = [
            dict(base),
            {**base, "node_id": "retained", "disposition": "RETAIN"},
            {
                **base,
                "node_id": "hidden",
                "matrix_visibility": "SYSTEM_NON_BUSINESS",
            },
            {
                **base,
                "node_id": "ineligible",
                "analysis_eligibility": "NOT_ELIGIBLE",
            },
        ]
        self.assertEqual(
            [node["node_id"] for node in _matrix_decision_nodes(candidates)],
            ["atom-1"],
        )

    def test_invalid_hierarchy_fails_closed(self) -> None:
        broken = _snapshot()
        broken["nodes"][1]["node_id"] = "root"
        with TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "node IDs must be unique"):
                export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    matrix_snapshot=broken,
                )

    def test_physical_table_refs_do_not_become_logical_table_identities(self) -> None:
        snapshot = _snapshot()
        for node in snapshot["nodes"]:
            for table_ref in node.get("table_refs") or []:
                table_ref.pop("logical_table_id", None)

        frozen_snapshot = RequirementMatrixSnapshotV1.model_validate(snapshot)
        table_node = next(
            node for node in frozen_snapshot.nodes if node.node_id == "table"
        )
        self.assertEqual(table_node.logical_table_ids, ("logical-table-1",))
        self.assertEqual(
            table_node.table_refs[0]["table_id"],
            "physical-table-1",
        )

        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=frozen_snapshot,
            )
            workbook = load_workbook(file_path)
            self.assertTrue(
                any(name.startswith("T001_9.16.2") for name in workbook.sheetnames)
            )
            table_sheet = _table_sheet(workbook)
            rendered_text = {
                str(cell.value or "") for row in table_sheet.iter_rows() for cell in row
            }
            self.assertIn("性能要求", rendered_text)
            self.assertIn("第一页记录", rendered_text)
            self.assertIn("第二页记录", rendered_text)

    def test_unknown_explicit_logical_table_identity_still_fails_closed(self) -> None:
        snapshot = _snapshot()
        snapshot["nodes"][2]["logical_table_ids"] = ["missing-logical-table"]
        with TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "outside logical_tables"):
                export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    matrix_snapshot=snapshot,
                )

    def test_analyze_node_cannot_be_ineligible(self) -> None:
        broken = _snapshot()
        broken["nodes"][4]["analysis_eligibility"] = "NOT_ELIGIBLE"
        with TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "ANALYZE disposition"):
                export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    matrix_snapshot=broken,
                )

    def test_long_matrix_text_keeps_one_matrix_row_and_uses_full_text_detail(
        self,
    ) -> None:
        snapshot = _snapshot()
        long_text = (
            "超长需求内容第一段保持完整。\n\n补充条件第二段保持独立。\n\n" * 1800
        ).strip()
        self.assertGreater(len(long_text), 32_767)
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom["requirement_text"] = long_text
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            start_row = next(
                row
                for row in range(4, matrix.max_row + 1)
                if str(matrix.cell(row, MATRIX_BODY_COLUMN).value or "").startswith(
                    "超长需求内容"
                )
            )
            self.assertEqual(matrix.max_row, 11)
            self.assertIn(
                "完整需求正文见 E 列详情",
                matrix.cell(start_row, MATRIX_BODY_COLUMN).value,
            )
            self.assertIsNotNone(
                matrix.cell(start_row, MATRIX_SOURCE_ENTRY_COLUMN).hyperlink
            )
            self.assertIn(
                "全文详情",
                str(
                    matrix.cell(
                        start_row, MATRIX_SOURCE_ENTRY_COLUMN
                    ).hyperlink.location
                    or ""
                ),
            )
            detail = workbook["全文详情"]
            self.assertEqual(detail.sheet_state, "visible")
            content_rows = [
                row
                for row in range(4, detail.max_row + 1)
                if detail.cell(row, 2).value in {"需求正文", "续"}
            ]
            self.assertEqual(
                "".join(str(detail.cell(row, 3).value or "") for row in content_rows),
                long_text,
            )
            self.assertTrue(
                all(
                    (detail.row_dimensions[row].height or 0) <= 409.5
                    for row in content_rows
                )
            )
            self.assertFalse(
                any(
                    matrix.cell(row, 4).value == "续"
                    for row in range(4, matrix.max_row + 1)
                )
            )

    def test_f06_shaped_clause_stays_in_one_cell_and_keeps_real_paragraphs(
        self,
    ) -> None:
        paragraphs = [
            "本技术资料仅用于受邀供应商进行项目报价，未经许可不得向无关组织或人员转发，相关责任由泄露方承担。",
            "本资料描述供应商在项目中的责任分工、任务安排、工程要求、节点与交付物，供应商承诺范围包括但不限于资料明确列出的内容。",
            "资料仅发布中文版本，其他语言版本由供应商自行翻译并保证准确；全部内容仅供报价阶段参考。",
            "供应商应依据项目最新输入完成报价，如后续输入发生变化，应同步评估影响并按正式流程反馈。",
        ]

        def hard_wrap(value: str) -> str:
            return "\n".join(
                value[index : index + 24] for index in range(0, len(value), 24)
            )

        source_text = "文档范围\n\n" + "\n\n".join(
            hard_wrap(value) for value in paragraphs
        )
        expected = "\n\n".join(paragraphs)
        self.assertGreaterEqual(len(expected), 180)
        self.assertLessEqual(len(expected), 320)
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom.update(
            {
                "label": "文档范围",
                "display_title": "文档范围",
                "requirement_text": source_text,
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            row = _find_row(matrix, "REQ-001", column=MATRIX_REQ_ID_COLUMN)
            self.assertIn("文档范围", matrix.cell(row, MATRIX_TITLE_COLUMN).value)
            self.assertTrue(
                str(matrix.cell(row, MATRIX_BODY_COLUMN).value or "").startswith(
                    paragraphs[0][:24]
                )
            )
            self.assertIn(
                "完整需求正文见 E 列详情",
                matrix.cell(row, MATRIX_BODY_COLUMN).value,
            )
            self.assertLessEqual(matrix.row_dimensions[row].height, 110.0)
            self.assertEqual(matrix.max_row, 11)
            self.assertEqual(
                _first_full_text_field(workbook, "需求正文"),
                expected,
            )
            self.assertFalse(
                any(
                    matrix.cell(candidate, 4).value == "续"
                    for candidate in range(4, matrix.max_row + 1)
                )
            )

    def test_pdf_hard_wrap_cleanup_preserves_unpunctuated_paragraphs_and_lists(
        self,
    ) -> None:
        first_paragraph = (
            "供应商需要基于客户输入完成系统方案评审并形成可追溯交付记录"
            "同时保证全部关键假设均进入后续确认流程"
        )
        second_paragraph = (
            "如输入条件发生变化供应商需要重新评估计划成本以及技术风险"
            "并把影响范围反馈给项目团队"
        )
        first_item = "（1）接口边界需要由双方共同确认并记录责任人"
        second_item = "（2）未确认事项需要保留为开放问题直至客户回复"

        source_text = "\n\n".join(
            (
                first_paragraph,
                second_paragraph,
                f"{first_item}\n{second_item}",
            )
        )
        expected = (
            f"{first_paragraph}\n\n{second_paragraph}\n\n{first_item}\n{second_item}"
        )
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom["requirement_text"] = source_text
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            matrix = workbook["ComplianceMatrix"]
            row = _find_row(matrix, "REQ-001", column=MATRIX_REQ_ID_COLUMN)
            self.assertTrue(
                str(matrix.cell(row, MATRIX_BODY_COLUMN).value or "").startswith(
                    first_paragraph[:24]
                )
            )
            self.assertLessEqual(matrix.row_dimensions[row].height, 110.0)
            self.assertEqual(
                _first_full_text_field(workbook, "需求正文"),
                expected,
            )
            self.assertFalse(
                any(
                    matrix.cell(candidate, 4).value == "续"
                    for candidate in range(4, matrix.max_row + 1)
                )
            )

    def test_same_title_table_records_keep_their_own_values(self) -> None:
        snapshot = _snapshot()
        table = next(node for node in snapshot["nodes"] if node["node_id"] == "table")
        first = next(node for node in snapshot["nodes"] if node["node_id"] == "record")
        first.update(
            {
                "label": "运行状态",
                "display_title": "运行状态",
                "display_text": "I Start",
            }
        )
        first_part = table["table_refs"][0]
        first_part["cells"][4]["raw_text"] = "I Start"
        first_part["cells"].extend(
            [
                {
                    "id": "physical-table-1-r2c1",
                    "row": 3,
                    "column": 0,
                    "raw_text": "I Stop",
                },
                {
                    "id": "physical-table-1-r2c2",
                    "row": 3,
                    "column": 1,
                    "raw_text": "12 V",
                },
            ]
        )
        second = _node(
            "record-2-node",
            parent_id="table",
            child_ids=[],
            depth=3,
            source_order=4,
            display_role="TABLE_RECORD",
            label="运行状态",
            display_title="运行状态",
            display_text="I Stop",
            table_record_id="record-2",
            logical_table_ids=["logical-table-1"],
        )
        table["child_ids"] = ["record", "record-2-node"]
        snapshot["nodes"].append(second)
        logical_table = snapshot["logical_tables"][0]
        logical_table["node_ids"].append("record-2-node")
        second_record = deepcopy(logical_table["projection"]["records"][0])
        second_record["record_id"] = "record-2"
        second_record["source_order"] = 1
        second_record["row_bindings"] = [
            {
                "table_id": "physical-table-1",
                "row_index": 3,
                "cell_ids": [
                    "physical-table-1-r2c1",
                    "physical-table-1-r2c2",
                ],
            }
        ]
        logical_table["projection"]["records"].append(second_record)
        logical_table["record_count"] = 2
        snapshot["counts"]["total_nodes"] += 1
        snapshot["counts"]["visible_nodes"] += 1
        snapshot["counts"]["retained_nodes"] += 1
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            matrix = load_workbook(file_path)["ComplianceMatrix"]
            contents = [
                matrix.cell(row, MATRIX_BODY_COLUMN).value
                for row in range(4, matrix.max_row + 1)
            ]
            start_value = "项目：I Start\n要求：12 V"
            stop_value = "项目：I Stop\n要求：12 V"
            self.assertIn(start_value, contents)
            self.assertIn(stop_value, contents)
            self.assertNotEqual(
                _find_row(matrix, start_value, column=MATRIX_BODY_COLUMN),
                _find_row(matrix, stop_value, column=MATRIX_BODY_COLUMN),
            )
            for value in (start_value, stop_value):
                row = _find_row(matrix, value, column=MATRIX_BODY_COLUMN)
                self.assertIn("运行状态", matrix.cell(row, MATRIX_TITLE_COLUMN).value)

    def test_same_source_node_different_spans_are_not_deduplicated(self) -> None:
        snapshot = _snapshot()
        context = next(
            node for node in snapshot["nodes"] if node["node_id"] == "context"
        )
        context.update(
            {
                "display_text": "",
                "summary": "",
                "content_blocks": [
                    {
                        "kind": "paragraph",
                        "source_id": "shared-source-node",
                        "start_offset": 0,
                        "end_offset": 5,
                        "text": "第一片段",
                    },
                    {
                        "kind": "paragraph",
                        "source_id": "shared-source-node",
                        "start_offset": 5,
                        "end_offset": 10,
                        "text": "第二片段",
                    },
                ],
            }
        )
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            matrix = load_workbook(file_path)["ComplianceMatrix"]
            self.assertIsNotNone(
                _find_row(
                    matrix,
                    "第一片段\n\n第二片段",
                    column=MATRIX_BODY_COLUMN,
                )
            )

    def test_formula_and_url_like_source_text_remain_plain_text(self) -> None:
        snapshot = _snapshot()
        atom = next(node for node in snapshot["nodes"] if node["node_id"] == "atom-1")
        atom["requirement_text"] = "=1+1"
        atom["ai_assessment"]["reason"] = "https://example.com/source"
        table = next(node for node in snapshot["nodes"] if node["node_id"] == "table")
        table["table_refs"][0]["cells"][5]["raw_text"] = "=SUM(1,2)"
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path, data_only=False)
            matrix = workbook["ComplianceMatrix"]
            atom_row = _find_row(matrix, "=1+1", column=MATRIX_BODY_COLUMN)
            self.assertEqual(
                matrix.cell(atom_row, MATRIX_BODY_COLUMN).data_type,
                "s",
            )
            self.assertIsNone(matrix.cell(atom_row, MATRIX_BODY_COLUMN).hyperlink)
            self.assertIn(
                "https://example.com/source",
                _sheet_text(workbook["全文详情"]),
            )
            self.assertIsNone(matrix.cell(atom_row, MATRIX_BASIS_COLUMN).hyperlink)
            source_formula_cell = next(
                cell
                for row in _table_sheet(workbook).iter_rows()
                for cell in row
                if cell.value == "=SUM(1,2)"
            )
            self.assertEqual(source_formula_cell.data_type, "s")

    def test_visible_sheets_have_compact_layout_and_print_contract(self) -> None:
        with TemporaryDirectory() as tmp:
            snapshot = _snapshot()
            long_table_title = (
                "对于有功能安全开发要求的零部件，供应商应完成接口协议、"
                "安全分析、验证证据与交付物确认"
            )
            table_node = next(
                node for node in snapshot["nodes"] if node["node_id"] == "table"
            )
            table_node["label"] = long_table_title
            table_node["display_title"] = long_table_title
            repository_logo = exporter_module._BRAND_LOGO_PATH_CANDIDATES[1]
            with patch.object(
                exporter_module,
                "_BRAND_LOGO_PATH_CANDIDATES",
                (repository_logo,),
            ):
                exporter_module._load_brand_logo_bytes.cache_clear()
                _, file_path = export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    matrix_snapshot=snapshot,
                )
                exporter_module._load_brand_logo_bytes.cache_clear()
            workbook = load_workbook(file_path)
            overview = workbook[OVERVIEW_SHEET]
            matrix = workbook[MATRIX_SHEET]
            self.assertEqual(matrix.freeze_panes, "D4")
            self.assertEqual(
                [matrix.cell(3, column).value for column in range(1, 11)],
                [
                    "Req ID",
                    "原文编号",
                    "章节 / 标题",
                    "需求正文",
                    "原文 / 附件",
                    "是否分析",
                    "分析进度",
                    "当前结论",
                    "分析依据 / 结论来源",
                    "客户确认",
                ],
            )
            total_matrix_width = sum(
                matrix.column_dimensions[column].width for column in "ABCDEFGHIJ"
            )
            self.assertGreaterEqual(matrix.column_dimensions["D"].width, 90.0)
            self.assertGreater(
                matrix.column_dimensions["D"].width / total_matrix_width,
                0.4,
            )
            for branded_sheet in (overview, matrix):
                self.assertTrue(branded_sheet._images)
                self.assertEqual(branded_sheet._images[0].anchor._from.row, 0)
                self.assertEqual(branded_sheet._images[0].anchor._from.col, 0)
            self.assertEqual(matrix["J3"].comment.author, "ReqPilotX")
            self.assertIn("没有独立", matrix["J3"].comment.text)
            meta = workbook["_ReqPilotXMeta"]
            capability = {
                meta.cell(row, 1).value: meta.cell(row, 2).value
                for row in range(1, meta.max_row + 1)
            }
            self.assertEqual(
                capability["CustomerConfirmationCapability"],
                "UNAVAILABLE_IN_SNAPSHOT_V1",
            )
            table_sheet = _table_sheet(workbook)
            self.assertEqual(table_sheet["A2"].value, long_table_title)
            self.assertTrue(table_sheet["A2"].alignment.wrap_text)
            self.assertGreater(table_sheet.row_dimensions[2].height, 30.0)
            overview_table_row = _find_row(
                overview,
                long_table_title,
                column=4,
            )
            self.assertGreater(overview.row_dimensions[overview_table_row].height, 24.0)
            self.assertLessEqual(
                overview.row_dimensions[overview_table_row].height,
                72.0,
            )
            self.assertLessEqual(
                sum(table_sheet.column_dimensions[column].width for column in "ABC"),
                90.0,
            )
            self.assertTrue(str(matrix.print_area))
            self.assertEqual(matrix.page_setup.orientation, "landscape")
            self.assertEqual(matrix.page_setup.scale, 85)
            self.assertEqual(matrix.freeze_panes, "D4")
            self.assertTrue(matrix.print_title_rows)
            self.assertTrue(str(table_sheet.print_area))
            self.assertEqual(table_sheet.page_setup.orientation, "portrait")
            self.assertIn(table_sheet.page_setup.scale, {None, 100})
            self.assertEqual(table_sheet.freeze_panes, "A4")
            self.assertTrue(table_sheet.print_title_rows)
            for sheet in (matrix, table_sheet):
                self.assertFalse(bool(sheet.sheet_properties.pageSetUpPr.fitToPage))
            self.assertIn(
                matrix["D5"].font.name,
                {"PingFang SC", "Microsoft YaHei", "Noto Sans CJK SC"},
            )

    def test_richest_source_segment_wins_and_content_blocks_are_exported(self) -> None:
        snapshot = _snapshot()
        table_node = next(
            node for node in snapshot["nodes"] if node["node_id"] == "table"
        )
        full_ref = deepcopy(table_node["table_refs"][0])
        partial_ref = deepcopy(full_ref)
        partial_ref["cells"] = partial_ref["cells"][:1]
        table_node["content_blocks"] = [full_ref, deepcopy(table_node["table_refs"][1])]
        table_node["table_refs"] = [partial_ref]

        catalog = _build_matrix_table_catalog(
            snapshot,
            _visible_matrix_nodes(snapshot),
        )
        by_physical_id = {
            part.physical_table_id: part
            for part in catalog.parts_by_logical_id["logical-table-1"]
        }
        self.assertEqual(
            {cell["id"] for cell in by_physical_id["physical-table-1"].cells},
            {cell["id"] for cell in full_ref["cells"]},
        )
        self.assertIn("physical-table-2", by_physical_id)

    def test_multi_table_container_can_reach_every_table_sheet(self) -> None:
        snapshot = _snapshot()
        table_node = next(
            node for node in snapshot["nodes"] if node["node_id"] == "table"
        )
        table_node["logical_table_ids"] = ["logical-table-1", "logical-table-2"]
        third_part = _table_part("physical-table-3", source_order=2, page=3)
        third_part["logical_table_id"] = "logical-table-2"
        table_node["table_refs"].append(third_part)
        snapshot["logical_tables"].append(
            {
                "logical_table_id": "logical-table-2",
                "projection_id": "projection-2",
                "projection_mode": "WHOLE_TABLE",
                "record_count": 0,
                "physical_table_ids": ["physical-table-3"],
                "node_ids": ["table"],
                "evidence_binding_ids": [],
                "projection": {
                    "projection_id": "projection-2",
                    "mode": "WHOLE_TABLE",
                    "table_fact": {
                        "logical_table_id": "logical-table-2",
                        "table_parts": [
                            {
                                "table_id": "physical-table-3",
                                "source_order": 2,
                                "source_refs": [{"page_number": 3}],
                            }
                        ],
                    },
                    "records": [],
                },
                "evidence_bindings": [],
            }
        )
        snapshot["counts"]["logical_tables"] = 2
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            workbook = load_workbook(file_path)
            table_sheets = [
                workbook[name] for name in workbook.sheetnames if name.startswith("T")
            ]
            self.assertEqual(len(table_sheets), 2)
            matrix = workbook["ComplianceMatrix"]
            table_row = _find_row(matrix, "跨页性能参数原表", column=MATRIX_BODY_COLUMN)
            self.assertEqual(
                matrix.cell(table_row, MATRIX_SOURCE_ENTRY_COLUMN).value,
                "打开 2 张表",
            )
            self.assertIsNotNone(
                matrix.cell(table_row, MATRIX_SOURCE_ENTRY_COLUMN).hyperlink
            )
            self.assertEqual(
                table_sheets[0].cell(1, 3).hyperlink.location,
                f"'{table_sheets[1].title}'!A1",
            )
            self.assertEqual(
                table_sheets[1].cell(1, 2).hyperlink.location,
                f"'{table_sheets[0].title}'!A1",
            )

    def test_equal_source_rows_are_preserved_without_appending_derived_values(
        self,
    ) -> None:
        snapshot = _snapshot()
        table_node = next(
            node for node in snapshot["nodes"] if node["node_id"] == "table"
        )
        first_ref = table_node["table_refs"][0]
        first_ref["cells"] = []
        first_ref["merged_regions"] = []
        first_ref["headers"] = ["值"]
        first_ref["rows"] = [["相同"], ["相同"]]
        record = snapshot["logical_tables"][0]["projection"]["records"][0]
        record["shared_merged_values"] = [
            {"cell_id": "shared-1", "text": "适用于全部型号"},
            {"cell_id": "shared-2", "text": "适用于全部型号"},
        ]
        catalog = _build_matrix_table_catalog(
            snapshot,
            _visible_matrix_nodes(snapshot),
        )
        part = next(
            item
            for item in catalog.parts_by_logical_id["logical-table-1"]
            if item.physical_table_id == "physical-table-1"
        )
        self.assertEqual(part.rows, [["相同"], ["相同"]])
        with TemporaryDirectory() as tmp:
            _, file_path = export_compliance_matrix(
                project_name="ReqPilotX",
                customer_name="Customer",
                export_dir=Path(tmp),
                matrix_snapshot=snapshot,
            )
            sheet = _table_sheet(load_workbook(file_path))
            source_values = [
                cell.value for source_row in sheet.iter_rows() for cell in source_row
            ]
            self.assertEqual(source_values.count("相同"), 2)
            self.assertNotIn("适用于全部型号", source_values)
            self.assertNotIn("表内记录与分析", source_values)

    def test_oversized_native_source_cell_fails_before_silent_truncation(self) -> None:
        snapshot = _snapshot()
        table_node = next(
            node for node in snapshot["nodes"] if node["node_id"] == "table"
        )
        table_node["table_refs"][0]["cells"][0]["raw_text"] = "x" * 32_768
        with TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "character cell limit"):
                export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    matrix_snapshot=snapshot,
                )

    def test_legacy_requirement_rows_are_not_an_export_input(self) -> None:
        with TemporaryDirectory() as tmp:
            with self.assertRaises(TypeError):
                export_compliance_matrix(
                    project_name="ReqPilotX",
                    customer_name="Customer",
                    export_dir=Path(tmp),
                    requirements=[{"req_code": "REQ-OLD"}],
                )


if __name__ == "__main__":
    unittest.main()
from app.application import exports as exports_app
from app.application.export_contracts import ExportRequest
