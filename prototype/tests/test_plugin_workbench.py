from .runtime_isolation import activate_test_runtime
activate_test_runtime()
from argparse import Namespace
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import hashlib
import json
import sys
import unittest
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'plugins/reqpilot/scripts'),str(ROOT/'plugins/reqpilot/tests')]
from test_portable import make_fixture,item_decisions,write
import artifact_tools
from app.application.agent_workflow import plugin_workbench as api
from app.application.errors import ApplicationError

class PluginWorkbenchTests(unittest.TestCase):
    def test_export_uses_registered_original_name_not_materialized_digest_filename(self):
        saved=self.sync(self.artifacts,require_map_confirmation=False)
        source=self.root/(self.sha+'.pdf');source.write_bytes(self.source.read_bytes())
        registered='客户接口需求原件.pdf';self.document['file_name']=registered
        snapshot=api.get_plugin_snapshot('p','d',self.session,self.sink())
        with patch.object(api,'resolve_source_record_path',return_value=source):
            result=api.export_plugin_draft('p','d',self.session,saved['revision_id'])
        workbook=load_workbook(BytesIO(result.content),read_only=True,data_only=True)
        try:
            subtitle=workbook['ComplianceMatrix']['A2'].value
            self.assertIn(registered,subtitle)
            self.assertNotIn(self.sha+'.pdf',subtitle)
            metadata={row[0]:row[1] for row in workbook['_ReqPilotXMeta'].iter_rows(values_only=True)}
            self.assertEqual(metadata['SnapshotSHA256'],api.core.canonical_json_sha256(snapshot))
        finally:workbook.close()
        self.assertEqual(result.headers['X-ReqPilot-Revision'],saved['revision_id'])
        self.assertEqual(result.headers['X-ReqPilot-Snapshot'],snapshot['delivery_sha256'])
        self.assertEqual(api.get_plugin_snapshot('p','d',self.session,self.sink()),snapshot)

    def test_lists_same_source_tasks_without_latest_winning_or_losing_empty_tasks(self):
        first = self.sync(self.mapping(), goal='First task')
        first_id = self.session
        self.session = 'plugin-' + 'b'*32
        empty = self.sync(goal='Second independent task')
        result = api.list_plugin_workbench_sessions('p', 'd', self.sink())
        self.assertEqual({s['session_id'] for s in result['sessions']}, {first_id, self.session})
        self.assertEqual({s['availability'] for s in result['sessions']}, {'AVAILABLE', 'NO_ARTIFACTS'})
        self.assertIsNone(result['current_session_id'])
        self.assertEqual(result['selection_policy'], 'EXPLICIT_WORK_SELECTION')
        self.assertEqual(result['unavailable'], [])
        self.assertEqual(api.get_plugin_workbench('p', 'd', first_id, self.sink())['revision_id'], first['revision_id'])
        self.assertIsNone(empty['revision_id'])

    def test_bad_session_is_reported_and_source_change_does_not_choose_another_task(self):
        self.sync(self.mapping())
        directory = api._root('p', 'd', 'plugin-' + 'b'*32)
        directory.mkdir(); (directory/'head.json').write_text('{broken')
        result = api.list_plugin_workbench_sessions('p', 'd', self.sink())
        self.assertEqual(len(result['sessions']), 1)
        self.assertEqual(result['unavailable'][0]['reason_code'], 'SAVED_TASK_UNREADABLE')
        with patch.object(api, '_require_project_document', return_value={**self.document, 'source_sha256': 'c'*64}):
            changed = api.list_plugin_workbench_sessions('p', 'd', self.sink())
        self.assertEqual(changed['sessions'], [])
        self.assertIn('SOURCE_VERSION_CHANGED', {s['reason_code'] for s in changed['unavailable']})

    def test_session_listing_does_not_follow_linked_head_or_version_metadata(self):
        import shutil
        state = self.sync(self.mapping())
        directory = api._root('p', 'd', self.session)
        outside = self.root/'outside-head.json'
        head = directory/'head.json'; head.rename(outside); head.symlink_to(outside)
        original_read = api._read
        reads = []
        def read(path):
            reads.append(Path(path))
            return original_read(path)
        with patch.object(api, '_read', side_effect=read):
            result = api.list_plugin_workbench_sessions('p', 'd', self.sink())
        self.assertEqual(reads, [])
        self.assertEqual(result['unavailable'][0]['reason_code'], 'UNSAFE_TASK_PATH')
        head.unlink(); outside.rename(head)
        version = directory/'versions'/state['revision_id']
        escaped = self.root/'outside-version'; shutil.move(version, escaped); version.symlink_to(escaped, target_is_directory=True)
        reads.clear()
        with patch.object(api, '_read', side_effect=read):
            result = api.list_plugin_workbench_sessions('p', 'd', self.sink())
        self.assertEqual(reads, [head])
        self.assertEqual(result['sessions'], [])
        self.assertEqual(result['unavailable'][0]['reason_code'], 'UNSAFE_TASK_PATH')

    def test_incumbent_reader_preserves_tree_ledger_tables_and_same_version_detail(self):
        from app.application.agent_workflow import plugin_workbench_reader as reader
        first = self.sync(self.mapping())
        map_view = reader.get_plugin_reader('p', 'd', self.session, first['revision_id'], self.sink())
        self.assertIsNone(map_view['outline'])
        self.assertEqual(map_view['reference']['structure_projection']['state'], 'DRAFT')
        self.assertFalse(map_view['reference']['continuation']['can_continue'])
        self.confirm(first)
        state = self.sync(self.artifacts, first['revision_id'])
        view = reader.get_plugin_reader('p', 'd', self.session, state['revision_id'], self.sink())
        mapping = self.artifacts['document-map.json']
        baseline = self.artifacts['source-requirement-baseline.json']
        self.assertEqual(view['reference']['document_tree']['root_ids'], mapping['root_node_ids'])
        self.assertEqual([n['node_id'] for n in view['reference']['document_tree']['nodes']], [n['node_id'] for n in mapping['nodes']])
        self.assertEqual([r['source_item_id'] for r in view['outline']['records']], [r['source_record_id'] for r in baseline['records']])
        self.assertIn('source_number', view['outline']['records'][0])
        self.assertEqual(view['scope_summaries'], view['outline']['scope_summaries'])
        self.assertEqual(view['source_findings']['complete_or_unknown'], 'UNKNOWN')
        public = api.get_plugin_workbench('p', 'd', self.session, self.sink())
        self.assertEqual(public['execution']['status'], 'UNREPORTED')
        self.assertFalse(public['capabilities']['start']['supported'])
        self.assertTrue(public['capabilities']['save_feedback']['enabled'])
        self.assertTrue(public['capabilities']['export_draft']['enabled'])
        for scope, table in self.artifacts['table-previews.json'].items():
            self.assertEqual(view['tables'][scope]['fragments'], table['fragments'])
        for item in baseline['records']:
            detail = reader.get_plugin_record('p', 'd', self.session, item['source_record_id'], state['revision_id'], self.sink())
            self.assertEqual(detail['baseline_artifact_sha256'], baseline['artifact_sha256'])
            self.assertEqual(detail['record']['table_fragments'], item['table_fragments'])
            self.assertEqual([p['text'] for p in detail['record']['text_bindings']], [p['text'] for p in item['text_bindings']])
        self.assertIsNone(reader.get_plugin_reader('p', 'd', self.session, first['revision_id'], self.sink())['outline'])
        self.assertEqual(api.get_plugin_workbench('p', 'd', self.session, self.sink())['revision_id'], state['revision_id'])

    def test_plugin_feedback_reappears_in_incumbent_source_detail(self):
        from app.application.agent_workflow import plugin_workbench_reader as reader
        state = self.sync(self.artifacts, require_map_confirmation=False)
        item = self.artifacts['source-requirement-baseline.json']['records'][0]
        self.feedback(state, node_id=item['node_id'], source_item_ids=[item['source_record_id']])
        detail = reader.get_plugin_record('p', 'd', self.session, item['source_record_id'], state['revision_id'], self.sink())
        self.assertEqual(detail['record']['feedbacks'][0]['message'], '请核对这段内容的上级条件。')
        self.assertEqual(detail['record']['feedbacks'][0]['origin_revision_id'], state['revision_id'])
        self.assertEqual(detail['record']['feedbacks'][0]['human_acceptance'], 'NOT_CLAIMED')

    def setUp(self):
        self.temp=TemporaryDirectory();self.root=Path(self.temp.name)
        self.source,capture,decisions=make_fixture(self.root/'fixture')
        write(self.root/'map-input.json',decisions)
        mapping=self.root/'map';artifact_tools.compile_map(Namespace(source=self.source,source_ir=capture/'source_ir.json',input=self.root/'map-input.json',output_dir=mapping))
        write(self.root/'items-input.json',item_decisions(mapping))
        baseline=self.root/'baseline';artifact_tools.compile_items(Namespace(map_dir=mapping,input=self.root/'items-input.json',baseline_dir=None,output_dir=baseline))
        self.artifacts={n:json.loads((baseline/n).read_text()) for n in ('document-map.json','table-previews.json','source-requirement-baseline.json','validation-receipt.json')}
        self.sha=hashlib.sha256(self.source.read_bytes()).hexdigest();self.session='plugin-'+'a'*32
        self.document={'source_sha256':self.sha,'file_name':self.source.name}
        self.patches=[patch.object(api,'_container',return_value=SimpleNamespace(settings=SimpleNamespace(object_root=self.root/'objects'))),
            patch.object(api,'_require_project_document',return_value=self.document),patch.object(api,'resolve_source_record_path',return_value=self.source)]
        for p in self.patches:p.start()
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.temp.cleanup()
    def sink(self):return SimpleNamespace(headers={})
    def payload(self,artifacts=None,expected=None,**extra):
        extra.setdefault('require_map_confirmation',True) # These legacy checkpoint cases explicitly request it.
        return api.PluginDraftUpload(session_id=self.session,expected_source_sha256=self.sha,expected_revision_id=expected,artifacts=artifacts or {},**extra)
    def mapping(self):return {k:v for k,v in self.artifacts.items() if k in {'document-map.json','table-previews.json'}}
    def sync(self,artifacts=None,expected=None,**extra):return api.sync_plugin_draft('p','d',self.payload(artifacts,expected,**extra),self.sink())
    def confirm(self,state):return api.confirm_plugin_map('p','d',self.session,api.PluginMapConfirmation(command_id='confirm',expected_source_sha256=self.sha,expected_map_sha256=state['stages'][0]['artifact_sha256'],expected_revision_id=state['revision_id']),self.sink())
    def changed_map(self,summary):
        data=deepcopy(self.mapping());m=data['document-map.json'];m['document_summary']=summary
        m['artifact_sha256']=api.core.canonical_json_sha256({k:v for k,v in m.items() if k not in {'artifact_id','artifact_sha256'}});m['artifact_id']='document-map-'+m['artifact_sha256'][:32];return data
    def feedback(self,state,**overrides):
        fields={'command_id':'feedback','expected_source_sha256':self.sha,'expected_revision_id':state['revision_id'],'source_item_ids':[],'text':'请核对这段内容的上级条件。'};fields.update(overrides)
        return api.post_plugin_feedback('p','d',self.session,api.PluginFeedback(**fields),self.sink())
    def test_map_confirmation_then_items_uses_real_shared_snapshot_and_excel(self):
        initial=self.sync();self.assertEqual(initial['stages'],[]);self.assertEqual(initial['model_dispatch'],'NOT_PERFORMED')
        mapping=self.sync(self.mapping());self.assertTrue(mapping['map_confirmation_required'])
        with self.assertRaises(ApplicationError):self.sync(self.artifacts,mapping['revision_id'])
        self.confirm(mapping);items=self.sync(self.artifacts,mapping['revision_id'])
        self.assertEqual([s['stage'] for s in items['stages']],['MAP','ITEMS'])
        latest=api.get_latest_plugin_workbench('p','d',self.sink())
        self.assertEqual(latest['revision_id'],items['revision_id']);self.assertEqual(latest['session_id'],self.session)
        snapshot=api.get_plugin_snapshot('p','d',self.session,self.sink());exported=api.export_plugin_draft('p','d',self.session,items['revision_id'])
        workbook=load_workbook(BytesIO(exported.content));metadata={r[0]:r[1] for r in workbook['_ReqPilotXMeta'].values}
        self.assertEqual(metadata['SnapshotSHA256'],api.core.canonical_json_sha256(snapshot))
        self.assertEqual(snapshot['source_item_baseline']['artifact_sha256'],self.artifacts['source-requirement-baseline.json']['artifact_sha256'])
        self.assertEqual(api.get_plugin_stage('p','d',self.session,'MAP',self.sink(),revision=mapping['revision_id'])['artifact_sha256'],mapping['stages'][0]['artifact_sha256'])

    def test_new_default_is_autonomous_and_existing_checkpoint_policy_is_inherited_when_omitted(self):
        automatic=api.PluginDraftUpload(session_id=self.session,expected_source_sha256=self.sha,artifacts=self.artifacts)
        self.assertFalse(automatic.require_map_confirmation)
        self.assertNotIn('require_map_confirmation',automatic.model_fields_set)
        saved=api.sync_plugin_draft('p','d',automatic,self.sink())
        self.assertFalse(saved['map_confirmation_required'])
        self.assertEqual([stage['stage'] for stage in saved['stages']],['MAP','ITEMS'])
        self.assertTrue(api.export_plugin_draft('p','d',self.session,saved['revision_id']).content)
        with self.assertRaises(ApplicationError):
            api.sync_plugin_draft('p','d',api.PluginDraftUpload(session_id=self.session,expected_source_sha256=self.sha,require_map_confirmation=True),self.sink())

        legacy='plugin-'+'b'*32
        first=api.sync_plugin_draft('p','d',api.PluginDraftUpload(session_id=legacy,expected_source_sha256=self.sha,
            require_map_confirmation=True,artifacts=self.mapping()),self.sink())
        omitted=api.PluginDraftUpload(session_id=legacy,expected_source_sha256=self.sha)
        self.assertTrue(api.sync_plugin_draft('p','d',omitted,self.sink())['map_confirmation_required'])
        with self.assertRaises(ApplicationError):
            api.sync_plugin_draft('p','d',api.PluginDraftUpload(session_id=legacy,expected_source_sha256=self.sha,
                expected_revision_id=first['revision_id'],artifacts=self.artifacts),self.sink())
        with self.assertRaises(ApplicationError):
            api.sync_plugin_draft('p','d',api.PluginDraftUpload(session_id=legacy,expected_source_sha256=self.sha,require_map_confirmation=False),self.sink())
        api.confirm_plugin_map('p','d',legacy,api.PluginMapConfirmation(command_id='explicit-confirm',expected_source_sha256=self.sha,
            expected_revision_id=first['revision_id'],expected_map_sha256=first['stages'][0]['artifact_sha256']),self.sink())
        after=api.sync_plugin_draft('p','d',api.PluginDraftUpload(session_id=legacy,expected_source_sha256=self.sha,
            expected_revision_id=first['revision_id'],artifacts=self.artifacts),self.sink())
        self.assertEqual([stage['stage'] for stage in after['stages']],['MAP','ITEMS'])
        self.assertTrue(json.loads((api._root('p','d',legacy)/'head.json').read_text())['require_map_confirmation'])
    def test_wrong_source_tampered_artifacts_and_unsupported_state_are_rejected(self):
        with self.assertRaises(ApplicationError):api.sync_plugin_draft('p','d',self.payload().model_copy(update={'expected_source_sha256':'0'*64}),self.sink())
        bad=self.mapping();bad['document-map.json']=deepcopy(bad['document-map.json']);bad['document-map.json']['document_summary']='tampered'
        with self.assertRaises(ApplicationError):self.sync(bad)
        with self.assertRaises(ApplicationError):self.sync({**self.artifacts,'manual-decisions.json':{'records':{}}})
        self.assertFalse((api._root('p','d',self.session)/'head.json').exists())
    def test_feedback_is_version_bound_idempotent_and_not_lost_on_new_revision(self):
        first=self.sync(self.mapping());saved=self.feedback(first);self.assertEqual(len(self.feedback(first)['feedback']),1)
        with self.assertRaises(ApplicationError):self.feedback(first,text='different command retry')
        second=self.sync(self.changed_map('Synthetic second draft'),first['revision_id']);self.assertEqual(second['feedback'],saved['feedback'])
        with self.assertRaises(ApplicationError):self.feedback(first,command_id='stale-feedback')
        with self.assertRaises(ApplicationError):self.confirm(first)
        with self.assertRaises(ApplicationError):self.feedback(second,command_id='bad-scope',node_id='does-not-exist')
    def test_exact_retry_does_not_grow_history_and_return_to_old_content_is_new_revision(self):
        original=self.mapping();first=self.sync(original);self.assertEqual(self.sync(original,first['revision_id'])['revision_id'],first['revision_id'])
        old_path=api._root('p','d',self.session)/'versions'/first['revision_id']/'version.json';before=old_path.read_bytes()
        second=self.sync(self.changed_map('Changed note'),first['revision_id']);third=self.sync(original,second['revision_id'])
        self.assertNotEqual(first['revision_id'],third['revision_id']);self.assertEqual(third['revision_number'],3);self.assertEqual(before,old_path.read_bytes())
    def test_two_writers_cannot_overwrite_each_other(self):
        first=self.sync(self.mapping())
        def writer(text):
            try:return self.sync(self.changed_map(text),first['revision_id'])
            except ApplicationError as exc:return exc.status_code
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(writer,['First new note','Second new note']))
        self.assertEqual(sum(isinstance(r,dict) for r in results),1);self.assertIn(409,results)
    def test_file_change_and_registered_source_change_are_detected(self):
        saved=self.sync(self.mapping());api.get_plugin_snapshot('p','d',self.session,self.sink())
        path=api._root('p','d',self.session)/'versions'/saved['revision_id']/'matrix-snapshot.json';path.write_text('{}')
        with self.assertRaises(ApplicationError):api.get_plugin_snapshot('p','d',self.session,self.sink())
        self.document['source_sha256']='0'*64
        with self.assertRaises(ApplicationError):api.get_plugin_workbench('p','d',self.session,self.sink())
        self.assertIsNone(api.get_latest_plugin_workbench('p','d',self.sink()))
    def test_reported_feedback_resolution_is_bound_to_a_new_draft_not_human_acceptance(self):
        first=self.sync(self.mapping());self.confirm(first);items=self.sync(self.artifacts,first['revision_id'])
        saved=self.feedback(items);feedback_id=saved['feedback'][0]['feedback_id']
        result=self.sync(self.artifacts,items['revision_id'],feedback_resolutions={feedback_id:'已按原文核对并修订工作草稿。'})
        resolution=result['feedback_resolutions'][feedback_id]
        self.assertEqual(resolution['revision_id'],result['revision_id'])
        self.assertEqual(resolution['human_acceptance'],'NOT_CLAIMED')
        self.assertEqual(result['feedback'],saved['feedback'])
        with self.assertRaises(ApplicationError):self.sync(self.artifacts,result['revision_id'],feedback_resolutions={'unknown':'假反馈'})
        # An exact confirmation retry remains acknowledged even after item publication.
        self.assertEqual(self.confirm(first)['revision_id'],result['revision_id'])
