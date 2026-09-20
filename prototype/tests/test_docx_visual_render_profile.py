from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import hashlib
import json
import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from PIL import Image

from scripts.document_visual_review_bundle import _validate_render_profile
from scripts.docx_visual_render_profile import (
    DocxPdfRenderResult,
    DocxVisualRenderConfig,
    DocxVisualRenderProfileRejected,
    produce_docx_visual_render_profile,
)


def _pdf_bytes(page_count: int = 2) -> bytes:
    document = pymupdf.open()
    try:
        for page_number in range(1, page_count + 1):
            page = document.new_page(width=300, height=420)
            page.insert_text((40, 50), f"review page {page_number}")
        return document.tobytes()
    finally:
        document.close()


class _FakeRenderer:
    def __init__(
        self,
        *,
        pdf_bytes: bytes | None = None,
        mutate_source: Path | None = None,
    ) -> None:
        self.pdf_bytes = pdf_bytes if pdf_bytes is not None else _pdf_bytes()
        self.mutate_source = mutate_source
        self.call_count = 0

    def render(self, source_bytes: bytes) -> DocxPdfRenderResult:
        self.call_count += 1
        if self.mutate_source is not None:
            changed = Document()
            changed.add_paragraph("source changed while rendering")
            changed.save(self.mutate_source)
        return DocxPdfRenderResult(
            pdf_bytes=self.pdf_bytes,
            pdf_sha256=hashlib.sha256(self.pdf_bytes).hexdigest(),
            renderer_id="test-docx-pdf",
            renderer_version="9.8.7.6",
            renderer_config_sha256=hashlib.sha256(
                b"test renderer config"
            ).hexdigest(),
        )


class DocxVisualRenderProfileTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)
        self.source = self.root / "customer-source.docx"
        document = Document()
        document.add_heading("Controlled visual review", level=1)
        document.add_paragraph("First source paragraph.")
        document.add_page_break()
        document.add_paragraph("Second source paragraph.")
        document.save(self.source)
        self.source_sha256 = hashlib.sha256(
            self.source.read_bytes()
        ).hexdigest()

    def test_produces_generic_full_page_profile_without_source_path_leak(
        self,
    ) -> None:
        renderer = _FakeRenderer()
        output = self.root / "render"

        result = produce_docx_visual_render_profile(
            self.source,
            expected_source_sha256=self.source_sha256,
            output_dir=output,
            renderer=renderer,
        )

        self.assertEqual(renderer.call_count, 1)
        self.assertEqual(result.root, output.resolve())
        self.assertEqual(result.profile["page_count"], 2)
        self.assertEqual(
            result.profile["source"],
            {
                "source_sha256": self.source_sha256,
                "format": "docx",
            },
        )
        self.assertEqual(
            [page["page_number"] for page in result.profile["pages"]],
            [1, 2],
        )
        for page in result.profile["pages"]:
            self.assertEqual(page["locator"], {"kind": "OFFICE_PAGE"})
            artifact_path = output / page["artifact"]["path"]
            self.assertTrue(artifact_path.is_file())
            self.assertEqual(
                hashlib.sha256(artifact_path.read_bytes()).hexdigest(),
                page["artifact"]["sha256"],
            )
            with Image.open(artifact_path) as image:
                self.assertEqual(
                    image.size,
                    (
                        page["artifact"]["width_px"],
                        page["artifact"]["height_px"],
                    ),
                )
        _validate_render_profile(
            result.profile,
            source_sha256=self.source_sha256,
            document_format="docx",
        )
        profile_bytes = result.profile_path.read_bytes()
        self.assertEqual(
            hashlib.sha256(profile_bytes).hexdigest(),
            result.profile_file_sha256,
        )
        serialized = profile_bytes.decode("utf-8")
        self.assertNotIn(self.source.name, serialized)
        self.assertNotIn(str(self.source), serialized)
        self.assertNotIn("First source paragraph", serialized)
        self.assertEqual(
            json.loads(serialized)["profile_sha256"],
            result.profile["profile_sha256"],
        )
        self.assertFalse(
            any(
                item.name.startswith(".render.")
                for item in self.root.iterdir()
            )
        )

        with self.assertRaises(DocxVisualRenderProfileRejected):
            produce_docx_visual_render_profile(
                self.source,
                expected_source_sha256=self.source_sha256,
                output_dir=output,
                renderer=renderer,
            )

    def test_source_hash_symlink_and_mid_render_drift_fail_closed(self) -> None:
        never_called = _FakeRenderer()
        with self.assertRaises(DocxVisualRenderProfileRejected):
            produce_docx_visual_render_profile(
                self.source,
                expected_source_sha256="0" * 64,
                output_dir=self.root / "wrong-hash",
                renderer=never_called,
            )
        self.assertEqual(never_called.call_count, 0)
        self.assertFalse((self.root / "wrong-hash").exists())

        source_link = self.root / "linked.docx"
        source_link.symlink_to(self.source)
        with self.assertRaises(DocxVisualRenderProfileRejected):
            produce_docx_visual_render_profile(
                source_link,
                expected_source_sha256=self.source_sha256,
                output_dir=self.root / "symlink",
                renderer=never_called,
            )
        self.assertEqual(never_called.call_count, 0)

        mutating_renderer = _FakeRenderer(mutate_source=self.source)
        with self.assertRaisesRegex(
            DocxVisualRenderProfileRejected,
            "changed during rendering",
        ):
            produce_docx_visual_render_profile(
                self.source,
                expected_source_sha256=self.source_sha256,
                output_dir=self.root / "drift",
                renderer=mutating_renderer,
            )
        self.assertEqual(mutating_renderer.call_count, 1)
        self.assertFalse((self.root / "drift").exists())
        self.assertFalse(
            any(
                item.name.startswith(".drift.")
                for item in self.root.iterdir()
            )
        )

    def test_invalid_renderer_pdf_and_invalid_docx_fail_closed(self) -> None:
        invalid_pdf = _FakeRenderer(pdf_bytes=b"not a PDF")
        with self.assertRaisesRegex(
            DocxVisualRenderProfileRejected,
            "could not be opened",
        ):
            produce_docx_visual_render_profile(
                self.source,
                expected_source_sha256=self.source_sha256,
                output_dir=self.root / "invalid-pdf",
                renderer=invalid_pdf,
            )
        self.assertFalse((self.root / "invalid-pdf").exists())

        invalid_docx = self.root / "invalid.docx"
        invalid_docx.write_bytes(b"not a DOCX")
        invalid_docx_sha256 = hashlib.sha256(
            invalid_docx.read_bytes()
        ).hexdigest()
        with self.assertRaisesRegex(
            DocxVisualRenderProfileRejected,
            "package is invalid",
        ):
            produce_docx_visual_render_profile(
                invalid_docx,
                expected_source_sha256=invalid_docx_sha256,
                output_dir=self.root / "invalid-docx-output",
                renderer=_FakeRenderer(),
            )
        self.assertFalse((self.root / "invalid-docx-output").exists())

    def test_render_config_bounds_are_strict(self) -> None:
        with self.assertRaises(ValueError):
            DocxVisualRenderConfig(png_dpi=71)
        with self.assertRaises(ValueError):
            DocxVisualRenderConfig(png_dpi=301)
        with self.assertRaises(ValueError):
            DocxVisualRenderConfig(timeout_seconds=0)
        with self.assertRaises(ValueError):
            DocxVisualRenderConfig(timeout_seconds=601)


if __name__ == "__main__":
    unittest.main()
