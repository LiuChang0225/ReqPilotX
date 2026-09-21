"""Portable feedback loop around the shared ReqPilotX artifact compilers.

The manifest is a local review pointer, not a second source ledger. Its immutable
snapshots preserve feedback authorship and reject stale working results. Semantic
judgments always arrive as explicit Agent decisions for the shared compilers.
"""
from __future__ import annotations

from contextlib import contextmanager
from copy import deepcopy
import hashlib
import json
import os
import shutil
import tempfile
from pathlib import Path

from reqpilotx.application.agent_workflow import source_requirement_baseline as core
from artifact_tools import read, save, required, new_output, analysis_policy_input


BASE_FILES = ("document-map.json", "source-requirement-baseline.json", "table-previews.json", "validation-receipt.json")
OPTIONAL_FILES = ("scope-requests.json", "map-compilation-receipt.json", "item-compilation-receipt.json", "image-annotations.json")
STATE_FILES = {"analysis_policy": "analysis-policy.json", "annotations": "image-annotations.json", "decomposition": "source-item-decomposition.json", "assessments": "source-item-assessments.json", "analysis_packet": "analysis-packet.json", "knowledge_catalog": "knowledge-catalog.json", "manual_decisions": "manual-decisions.json"}
JUDGMENT_FIELDS = {"requirementness", "requirement_kinds", "non_requirement_role", "requirement_display_id", "confidence", "decision_summary", "reason_codes"}


def digest_file(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def published_workbook_snapshot_sha256(path):
    """Read the snapshot pin from a workbook whose bytes were already verified."""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if '_ReqPilotXMeta' not in workbook.sheetnames:
            raise ValueError('Published workbook has no immutable snapshot metadata')
        values = [row[1] for row in workbook['_ReqPilotXMeta'].iter_rows(values_only=True)
                  if len(row) > 1 and row[0] == 'SnapshotSHA256']
        if len(values) != 1 or not isinstance(values[0], str) or len(values[0]) != 64:
            raise ValueError('Published workbook snapshot hash is missing or ambiguous')
        return values[0]
    finally:
        workbook.close()


def inside(root, value):
    root = root.resolve()
    path = (root / value).resolve()
    if path == root or root not in path.parents:
        raise ValueError("Artifact path must stay inside the review workspace")
    return path


def relative(root, path):
    return str(path.resolve().relative_to(root.resolve()))


def seal(value):
    value = {k: v for k, v in value.items() if k != "review_sha256"}
    return {**value, "review_sha256": core.canonical_json_sha256(value)}


@contextmanager
def locked(root):
    """OS locks release on process exit; no orphan lock recovery guesses."""
    with (root / ".review.lock").open("a+b") as stream:
        if os.name == "nt":
            import msvcrt
            stream.seek(0); stream.write(b"0"); stream.flush(); stream.seek(0)
            msvcrt.locking(stream.fileno(), msvcrt.LK_LOCK, 1)
        else:
            import fcntl
            fcntl.flock(stream.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            if os.name == "nt":
                stream.seek(0); msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(stream.fileno(), fcntl.LOCK_UN)


def write_head(root, value):
    value = seal(value)
    snapshots = root / "reviews"; snapshots.mkdir(exist_ok=True, mode=0o700)
    name = value["review_sha256"] + ".json"
    try:
        save(snapshots, name, value)
    except FileExistsError:
        # A prior close may have saved immutable history just before process
        # exit. Reuse only the exact content; the caller still holds head CAS.
        existing = read(snapshots / name)
        if existing != value or seal(existing) != existing:
            raise ValueError("Existing review history differs from the intended head")
    fd, name = tempfile.mkstemp(prefix=".review-head-", dir=root)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
            stream.write("\n"); stream.flush(); os.fsync(stream.fileno())
        os.replace(name, root / "workspace-manifest.json")
    finally:
        if os.path.exists(name):
            os.unlink(name)
    return value


def load_head(root, expected=None):
    value = read(root / "workspace-manifest.json")
    if value.get("schema_version") != "reqpilot-review-workspace/v1" or seal(value)["review_sha256"] != value.get("review_sha256"):
        raise ValueError("Review manifest integrity changed")
    if expected is not None and expected != value["review_sha256"]:
        raise ValueError("Review changed since the packet was read; reload current feedback before revision")
    if digest_file(inside(root, value["source"]["path"])) != value["source"]["sha256"]:
        raise ValueError("Original source bytes changed")
    return value


def bundle(directory):
    document_map = core.DocumentMapArtifactV1.model_validate(read(directory / "document-map.json"))
    baseline = core.SourceRequirementBaselineArtifactV1.model_validate(read(directory / "source-requirement-baseline.json"))
    receipt = core.SourceRequirementBaselineValidationReceiptV1.model_validate(read(directory / "validation-receipt.json"))
    if (baseline.source_sha256 != document_map.source_sha256 or baseline.document_map_artifact_sha256 != document_map.artifact_sha256 or receipt.artifact_sha256 != baseline.artifact_sha256):
        raise ValueError("Imported Map, baseline and receipt lineage differ")
    return document_map, baseline, receipt, read(directory / "table-previews.json")


def state(directory):
    return {key: read(directory / name) for key, name in STATE_FILES.items() if (directory / name).is_file()}


def validate_state(document_map, baseline, tables, values, evidence_path=None):
    from reqpilotx.application.agent_workflow.source_item_delivery import build_delivery_snapshot
    from reqpilotx.application.agent_workflow.source_item_analysis import validate_source_item_assessments
    if values.get("assessments"):
        if not values.get("analysis_packet"):
            raise ValueError("Assessment import requires its exact analysis packet")
        validate_source_item_assessments(values["analysis_packet"], values["assessments"])
        if values.get("knowledge_catalog") is None:
            values["knowledge_catalog"] = empty_catalog_from_packet(values["analysis_packet"])
        if values["knowledge_catalog"]["catalog_sha256"] != values["analysis_packet"]["catalog_sha256"]:
            raise ValueError("Imported analysis packet and knowledge catalog differ")
    if values.get("knowledge_catalog"):
        from reqpilotx.application.agent_workflow.source_item_analysis import validate_knowledge_catalog
        validate_knowledge_catalog(values["knowledge_catalog"])
        if values["knowledge_catalog"].get("evidence_catalog"):
            from evidence_tools import verify_catalog
            if evidence_path is None or verify_catalog(evidence_path) != values["knowledge_catalog"]:
                raise ValueError("Nonempty review knowledge requires actual source-bound evidence registration")
    if values.get("analysis_packet"):
        from evidence_tools import verify_source_packet
        catalog = values.get("knowledge_catalog") or empty_catalog_from_packet(values["analysis_packet"])
        verify_source_packet(document_map, baseline, tables, values["analysis_packet"], catalog,
            annotations=values.get("annotations"), decomposition=values.get("decomposition"))
    build_delivery_snapshot(document_map, baseline, tables, annotations=values.get("annotations"), decomposition=values.get("decomposition"), assessments=values.get("assessments"), manual_decisions=values.get('manual_decisions'), analysis_policy=values.get('analysis_policy'))


def empty_catalog_from_packet(packet):
    catalog = {"schema_version": "project-knowledge-facts/v1", "project_id": "portable", "knowledge_base_version": "NO_CATALOG_SUPPLIED", "documents": [], "evidence_catalog": {},
        "gaps": [{"code": "NO_REGISTERED_KNOWLEDGE_CATALOG", "message": "No registered internal evidence catalog was supplied to this command; a compliance conclusion cannot be formed."}], "selection": "NO_REGISTERED_CATALOG_SUPPLIED"}
    catalog["catalog_sha256"] = core.canonical_json_sha256(catalog)
    if packet["catalog_sha256"] != catalog["catalog_sha256"]:
        raise ValueError("Existing analysis requires its exact registered --knowledge-catalog; only the canonical empty portable catalog can be recovered without that file")
    return catalog


def fingerprint(root, directory, baseline, document_map):
    values = state(directory)
    assessment = values.get("assessments", {})
    decomposition = values.get("decomposition", {})
    return {"version_path": relative(root, directory), "baseline_sha256": baseline.artifact_sha256, "document_map_sha256": document_map.artifact_sha256,
            "review_state": {"decomposition_status": decomposition.get("status", "ABSENT"), "derived_parent_count": len(decomposition.get("records", {})),
                "assessment_status": assessment.get("status", "ABSENT"), "assessed_unit_count": len(assessment.get("records", {})),
                "unassessed_unit_count": assessment.get("unassessed_unit_count"), "pending_unit_count": assessment.get("pending_unit_count")},
            "files": {path.name: digest_file(path) for path in sorted(directory.iterdir()) if path.is_file() and path.suffix == ".json"}}


def current_bundle(root, head):
    directory = inside(root, head["current"]["version_path"])
    for name, digest in head["current"].get("files", {}).items():
        if digest_file(inside(directory, name)) != digest:
            raise ValueError("Current version artifact changed after registration: " + name)
    result = bundle(directory)
    if result[1].artifact_sha256 != head["current"]["baseline_sha256"]:
        raise ValueError("Current baseline differs from the review manifest")
    return directory, result


def _identity_watermark(head, baseline, directory):
    floor = core._baseline_display_high_watermark(baseline)
    if "scope-revision-receipt.json" in head["current"].get("files", {}):
        proof = core.SourceItemScopeRevisionReceiptV1.model_validate(read(directory / "scope-revision-receipt.json"))
        floor = max(floor, core._resolved_source_display_high_watermark(
            artifact=baseline, previous_scope_revision_receipt=proof))
    return floor


def _identity_context(head, baseline, floor, prefix):
    context = {"policy": core.SOURCE_IDENTITY_ALLOCATION_POLICY,
        "review_sha256": head["review_sha256"], "history_prefix_sha256": core.canonical_json_sha256(prefix),
        "source_sha256": baseline.source_sha256, "previous_baseline_artifact_id": baseline.artifact_id,
        "previous_baseline_sha256": baseline.artifact_sha256, "high_watermark": floor}
    return {**context, "allocation_sha256": core.canonical_json_sha256(context)}


def _verified_source_identity_allocation(root, review_sha256, previous):
    """Replay one fixed immutable prefix once; never recursively verify prefixes."""
    chain, seen, digest = [], set(), review_sha256
    while digest:
        if digest in seen or len(digest) != 64 or any(c not in "0123456789abcdef" for c in digest):
            raise ValueError("Identity history integrity or ancestry changed")
        seen.add(digest)
        head = read(root / "reviews" / (digest + ".json"))
        if head.get("review_sha256") != digest or seal(head)["review_sha256"] != digest:
            raise ValueError("Identity history snapshot changed")
        directory, (document_map, baseline, _, tables) = current_bundle(root, head)
        if baseline.source_sha256 != previous.source_sha256 or head["source"]["sha256"] != previous.source_sha256:
            raise ValueError("Identity history replaced the original source")
        chain.append((head, baseline, directory, document_map, tables))
        digest = head.get("predecessor_review_sha256")
    chain.reverse()
    if not chain or chain[-1][1].artifact_sha256 != previous.artifact_sha256:
        raise ValueError("Identity history does not end at the actual predecessor")
    first = chain[0]
    floor = _identity_watermark(first[0], first[1], first[2]); prefix = [first[0]["review_sha256"]]
    for before, after in zip(chain, chain[1:]):
        old, old_baseline, old_dir, _, _ = before
        new, baseline, directory, document_map, tables = after
        if new["review_revision"] != old["review_revision"] + 1:
            raise ValueError("Identity history revision sequence is not contiguous")
        if baseline.artifact_sha256 != old_baseline.artifact_sha256 and baseline.predecessor_artifact_sha256 != old_baseline.artifact_sha256:
            raise ValueError("Identity history baseline predecessor changed")
        verify_transition(root, old_dir, directory, document_map, old_baseline, baseline, tables,
            expected_allocation=_identity_context(old, old_baseline, floor, prefix), expected_review_sha256=old["review_sha256"])
        floor = max(floor, _identity_watermark(new, baseline, directory))
        prefix.append(new["review_sha256"])
    return _identity_context(chain[-1][0], previous, floor, prefix)


def _new_source_identity_allocation(root, head, previous):
    return _verified_source_identity_allocation(root, head["review_sha256"], previous)


def summary(status, root, head, **extra):
    return {"status": status, "review_dir": str(root), "review_sha256": head["review_sha256"], "baseline_sha256": head["current"]["baseline_sha256"], "current_version_path": head["current"]["version_path"], **extra}


def open_review(args):
    if getattr(args, 'manual_decisions', None) and os.environ.get('REQPILOT_NATIVE_AGENT_TOOLS') == '1':
        raise ValueError('Native Agent tools cannot import engineer opinions')
    source_dir = required(args.baseline_dir, "--baseline-dir")
    document_map, baseline, receipt, tables = bundle(source_dir)
    source = required(args.source, "--source")
    if digest_file(source) != baseline.source_sha256:
        raise ValueError("Review source differs from the imported baseline")
    if document_map.source_format:
        from reqpilotx.application.agent_workflow.office_source_evidence import verify_office_map_against_source
        verify_office_map_against_source(source, document_map, tables, baseline=baseline)
    values = state(source_dir)
    for key in STATE_FILES:
        if getattr(args, key, None):
            values[key] = read(required(getattr(args, key), "--" + key.replace("_", "-")))
    if getattr(args,'analysis_policy',None):
        values['analysis_policy']=analysis_policy_input(values['analysis_policy'],baseline.source_sha256,baseline.artifact_sha256)
    evidence_path = getattr(args, "knowledge_catalog", None) or source_dir / "knowledge-catalog.json"
    validate_state(document_map, baseline, tables, values, evidence_path)
    root = new_output(args.review_dir)
    directory = root / "versions/initial"; directory.mkdir(parents=True, mode=0o700)
    for name in (*BASE_FILES, *OPTIONAL_FILES):
        if (source_dir / name).is_file() and name not in STATE_FILES.values():
            shutil.copyfile(source_dir / name, directory / name)
    for key, value in values.items():
        save(directory, STATE_FILES[key], value)
    if evidence_path.is_file():
        from evidence_tools import copy_registration
        copy_registration(evidence_path, directory)
    originals = root / "source"; originals.mkdir(mode=0o700)
    target = originals / source.name; shutil.copyfile(source, target); target.chmod(0o600)
    head = write_head(root, {"schema_version": "reqpilot-review-workspace/v1", "source": {"path": relative(root, target), "sha256": baseline.source_sha256},
        "current": fingerprint(root, directory, baseline, document_map), "feedback": [], "review_revision": 1, "predecessor_review_sha256": "", "human_acceptance": "NOT_CLAIMED"})
    return summary("REVIEW_OPENED", root, head, analysis_policy_status=(values.get('analysis_policy') or {}).get('status','UNSPECIFIED'))


def feedback_review(args):
    root = required(args.review_dir, "--review-dir"); value = read(required(args.input, "--input"))
    with locked(root):
        head = load_head(root, value["expected_review_sha256"])
        _, (_, baseline, _, _) = current_bundle(root, head)
        known = {record.source_record_id for record in baseline.records}
        entries = value.get("feedback", [])
        if not entries:
            raise ValueError("Supply actual feedback or a concrete self-review finding")
        updated = deepcopy(head)
        for entry in entries:
            if entry.get("origin") not in {"USER", "AGENT_SELF_REVIEW"} or not str(entry.get("text", "")).strip():
                raise ValueError("Feedback needs an explicit USER or AGENT_SELF_REVIEW origin and text")
            ids = entry.get("source_item_ids", [])
            if not isinstance(ids, list) or len(ids) != len(set(ids)) or set(ids) - known:
                raise ValueError("Feedback targets must resolve to current SourceItem IDs")
            item = {"origin": entry["origin"], "text": entry["text"], "source_item_ids": ids, "rationale": str(entry.get("rationale", "")),
                "bound_baseline_sha256": baseline.artifact_sha256, "status": "OPEN", "created_review_revision": head["review_revision"] + 1}
            item["feedback_id"] = "feedback-" + core.canonical_json_sha256(item)[:32]
            if any(old["feedback_id"] == item["feedback_id"] for old in updated["feedback"]):
                raise ValueError("Duplicate feedback entry")
            updated["feedback"].append(item)
        updated.update(review_revision=head["review_revision"] + 1, predecessor_review_sha256=head["review_sha256"])
        updated = write_head(root, updated)
    return summary("FEEDBACK_RECORDED", root, updated, added_feedback_ids=[entry["feedback_id"] for entry in updated["feedback"][len(head["feedback"]):]])


def packet_review(args):
    root = required(args.review_dir, "--review-dir"); head = load_head(root)
    directory, (document_map, baseline, _, tables) = current_bundle(root, head)
    selected = set(args.source_item_id or [])
    known = {r.source_record_id for r in baseline.records}
    if selected - known:
        raise ValueError("Unknown current SourceItem selection")
    records = [r for r in baseline.records if not selected or r.source_record_id in selected]
    node_ids = {r.node_id for r in records}
    by_id = {n.node_id: n for n in document_map.nodes}
    for node_id in list(node_ids):
        parent = by_id[node_id].parent_id
        while parent:
            node_ids.add(parent); parent = by_id[parent].parent_id
    values = state(directory)
    from scope_tools import scope_template
    selected_scopes = sorted({core._record_scope_id(record) for record in records if not record.image_bindings})
    payload = {"schema_version": "reqpilot-review-packet/v1", "review_sha256": head["review_sha256"], "source": head["source"], "current": head["current"],
        "feedback": [item for item in head["feedback"] if not selected or not item["source_item_ids"] or selected.intersection(item["source_item_ids"])],
        "records": [r.model_dump(mode="json") for r in records], "map_nodes": [n.model_dump(mode="json") for n in document_map.nodes if n.node_id in node_ids],
        "tables": {key: value for key, value in tables.items() if any(r.logical_table_scope_id == key for r in records)},
        "annotations": {sid: value for sid, value in values.get("annotations", {}).get("records", {}).items() if sid in {r.source_record_id for r in records}},
        "related_artifacts": {key: relative(root, directory / name) for key, name in STATE_FILES.items() if (directory / name).is_file()},
        "scope_revision_templates": [scope_template(document_map, baseline, tables, scope_id) for scope_id in selected_scopes],
        "source_scope": "Selection is a reading convenience; all original source, Map, tables and images remain available inside this workspace.", "human_acceptance": "NOT_CLAIMED"}
    table_presentations = {t.scope_id: core.document_map_table_presentation(t) for t in document_map.logical_tables
        if t.scope_id in payload["tables"] and t.agent_display_label is not None}
    if table_presentations:
        payload["table_presentations"] = table_presentations
    payload["packet_sha256"] = core.canonical_json_sha256(payload)
    if args.output_dir:
        output = new_output(args.output_dir); save(output, "review-packet.json", payload)
        return summary("REVIEW_PACKET_WRITTEN", root, head, packet_path=str(output / "review-packet.json"), packet_sha256=payload["packet_sha256"])
    return payload


def source_diff(previous, current, old_annotations=None, new_annotations=None):
    before = {r.source_record_id: r.model_dump(mode="json") for r in previous.records}
    after = {r.source_record_id: r.model_dump(mode="json") for r in current.records}
    if set(before) != set(after):
        raise ValueError("This classification/visual revision cannot add or remove original SourceItems")
    changes = []
    for sid, old in before.items():
        new = after[sid]
        if {k: v for k, v in old.items() if k not in JUDGMENT_FIELDS} != {k: v for k, v in new.items() if k not in JUDGMENT_FIELDS}:
            raise ValueError("Original source evidence, order or identity changed: " + sid)
        fields = {key: {"before": old.get(key), "after": new.get(key)} for key in sorted(JUDGMENT_FIELDS) if old.get(key) != new.get(key)}
        def reading(value):
            item = deepcopy((value or {}).get("records", {}).get(sid, {}))
            if isinstance(item.get("visual_interpretation"), dict):
                item["visual_interpretation"].pop("baseline_sha256", None)
            return item
        if reading(old_annotations) != reading(new_annotations):
            fields["annotations"] = {"before": reading(old_annotations), "after": reading(new_annotations)}
        if fields:
            changes.append({"source_item_id": sid, "source_display_id": old["source_display_id"], "fields": fields})
    return {"schema_version": "reqpilot-revision-diff/v1", "source_sha256": current.source_sha256, "predecessor_baseline_sha256": previous.artifact_sha256,
        "baseline_sha256": current.artifact_sha256, "preserved_source_item_count": len(before), "changed_source_item_count": len(changes), "changes": changes,
        "source_identity_order_and_bindings": "UNCHANGED", "semantic_verification": "REQUIRES_SOURCE_REVIEW", "human_acceptance": "NOT_CLAIMED"}


def annotation_diff(before, after):
    def content(value):
        if isinstance(value, dict):
            ignored = {"baseline_sha256", "annotations_sha256", "source_evidence_sha256", "annotation_source_evidence_sha256",
                "image_evidence_sha256", "document_map_artifact_sha256", "document_map_artifact_id", "node_id"}
            return {key: content(item) for key, item in value.items() if key not in ignored}
        if isinstance(value, list):
            return [content(item) for item in value]
        return value
    old, new = (before or {}).get("records", {}), (after or {}).get("records", {})
    return [{"source_item_id": sid, "before": content(old.get(sid)), "after": content(new.get(sid))}
        for sid in sorted(set(old) | set(new)) if content(old.get(sid)) != content(new.get(sid))]


def revise_review(args):
    from reqpilotx.application.agent_workflow.source_item_visual import rebind_source_item_annotations, compile_source_item_visual_readings, compile_source_item_annotation_updates
    root = required(args.review_dir, "--review-dir"); value = read(required(args.input, "--input"))
    if set(value) & {'manual_decisions', 'engineer_responses', 'manual_decision_import'}:
        raise ValueError('Inline review JSON cannot create or withdraw engineer opinions; use an explicit user import')
    if getattr(args, 'manual_decisions', None) and os.environ.get('REQPILOT_NATIVE_AGENT_TOOLS') == '1':
        raise ValueError('Native Agent revisions cannot replace engineer opinions')
    if "source_identity_allocation" in value or "source_display_high_watermark_floor" in value:
        raise ValueError("Source identity allocation is computed by the program from review history; callers must not supply it")
    head = load_head(root, value["expected_review_sha256"])
    directory, (document_map, previous, receipt, tables) = current_bundle(root, head)
    if value.get("schema_version") != "reqpilot-review-revision/v1" or value.get("expected_baseline_sha256") != previous.artifact_sha256:
        raise ValueError("Revision decisions must bind the current review and baseline")
    replacements = {key: read(required(getattr(args, key), "--" + key.replace("_", "-"))) for key in STATE_FILES if getattr(args, key, None)}
    if 'analysis_policy' in replacements:
        replacements['analysis_policy']=analysis_policy_input(replacements['analysis_policy'],previous.source_sha256,previous.artifact_sha256)
    if 'manual_decisions' in replacements and (len(replacements) != 1 or any(value.get(k) for k in ('decisions','scope_revision','map_revision','visual_readings','annotation_updates'))):
        raise ValueError('Import user opinions against the current version in a separate review revision')
    if not value.get("decisions") and not value.get("scope_revision") and not value.get("map_revision") and not value.get("visual_readings") and not value.get("annotation_updates") and not replacements:
        raise ValueError("Supply classification, image-reading changes or shared-compiler review artifacts")
    addressed = value.get("addressed_feedback_ids", [])
    if set(addressed) - {f["feedback_id"] for f in head["feedback"] if f["status"] in {"OPEN", "UNRESOLVED"}}:
        raise ValueError("Revision references absent or already addressed feedback")
    allocation = _new_source_identity_allocation(root, head, previous)
    if allocation is not None:
        value["source_identity_allocation"] = allocation
    allocation_floor = core.validate_source_identity_allocation(previous, allocation)
    old_values = state(directory); values = deepcopy(old_values)
    previous_map, previous_tables = document_map, tables
    map_result, map_rebind = None, None
    scope_revision, scope_proof = None, None
    if value.get("map_revision"):
        if value.get("decisions"):
            raise ValueError("Map migration preserves existing judgments; revise classifications separately")
        from map_tools import compile_map_transition
        candidate_map_dir = inside(root, value["map_revision"]["candidate_map_path"])
        map_result, map_rebind = compile_map_transition(document_map, previous, receipt, tables, candidate_map_dir,
            text_scope_recovery=value.get("scope_revision"), feedback_ids=addressed, source_identity_allocation=allocation)
        document_map = core.DocumentMapArtifactV1.model_validate(map_result["document_map"])
        baseline = core.SourceRequirementBaselineArtifactV1.model_validate(map_rebind["baseline"])
        validation = core.SourceRequirementBaselineValidationReceiptV1.model_validate(map_rebind["validation_receipt"])
        tables = map_rebind["table_previews"]
        coordination, proof = None, None
    elif value.get("scope_revision"):
        if value.get("decisions"):
            raise ValueError("A scope candidate already contains its explicit classifications; do not mix a second coordination operation")
        baseline, validation, scope_revision, scope_proof = compile_scope_transition(root, directory, document_map, previous, receipt, tables, value)
        coordination, proof = None, None
    elif value.get("decisions"):
        if not str(value.get("model_id", "")).strip():
            raise ValueError("Classification needs the actual Agent model identity")
        decisions = tuple(core.SourceRequirementClassificationCoordinationDecisionV1.model_validate(item) for item in value["decisions"])
        baseline, validation, coordination, proof = core.compile_source_requirement_classification_coordination(previous_artifact=previous, previous_validation_receipt=receipt, decisions=decisions, model_id=value["model_id"], coordinator_role="PORTABLE_AGENT", source_display_high_watermark_floor=allocation_floor)
    else:
        baseline, validation, coordination, proof = previous, receipt, None, None
    annotation_carry = None
    if values.get("annotations"):
        annotation_carry = rebind_source_item_annotations(previous_map, document_map, previous, baseline, values["annotations"], table_previews=tables,
            previous_table_previews=previous_tables, map_revision_receipt=map_rebind["receipt"] if map_rebind else None)
        values["annotations"] = annotation_carry["artifact"]
    if value.get("visual_readings"):
        values["annotations"] = compile_source_item_visual_readings(document_map, baseline, value["visual_readings"], base_annotations=values.get("annotations"))
    if value.get("annotation_updates"):
        values["annotations"] = compile_source_item_annotation_updates(document_map, baseline, value["annotation_updates"], base_annotations=values.get("annotations"), table_previews=tables)
    carry = carry_state(document_map, previous, baseline, tables, old_values, values, previous_tables=previous_tables, map_revision_receipt=map_rebind["receipt"] if map_rebind else None)
    if annotation_carry:
        carry["annotations"] = {key: val for key, val in annotation_carry.items() if key != "artifact"}
    values.update(replacements)
    if old_values.get('manual_decisions') is not None or values.get('manual_decisions') is not None:
        from reqpilotx.application.agent_workflow.source_item_manual_decisions import build_manual_decision_context, rebind_manual_decisions
        manual_context = build_manual_decision_context(document_map, baseline, tables,
            decomposition=values.get('decomposition'), annotations=values.get('annotations'))
        if 'manual_decisions' in replacements:
            from manual_tools import verify_import
            values['manual_decisions'] = verify_import(args.manual_decisions, manual_context, old_values.get('manual_decisions'))
        else:
            before_manual_context = build_manual_decision_context(previous_map, previous, previous_tables,
                decomposition=old_values.get('decomposition'), annotations=old_values.get('annotations'))
            values['manual_decisions'] = rebind_manual_decisions(old_values['manual_decisions'], manual_context,
                previous_context=before_manual_context, require_previous_context=True)
    carry["explicit_replacement_artifact_kinds"] = sorted(replacements)
    evidence_path = getattr(args, "knowledge_catalog", None) or directory / "knowledge-catalog.json"
    validate_state(document_map, baseline, tables, values, evidence_path)
    if map_rebind:
        recovery = map_rebind["receipt"].get("text_scope_recovery") or {}
        diff = {"schema_version": "reqpilot-revision-diff/v1", "transition_kind": "DOCUMENT_MAP_REVISION", "source_sha256": baseline.source_sha256,
            "predecessor_baseline_sha256": previous.artifact_sha256, "baseline_sha256": baseline.artifact_sha256,
            "preserved_source_item_count": len(previous.records), "changed_source_item_count": len(map_rebind["receipt"]["rebound_source_item_ids"]) + len(recovery.get("added_source_item_ids", [])),
            "created_source_item_ids": recovery.get("added_source_item_ids", []),
            "changes": map_result["receipt"], "source_rebinding": map_rebind["receipt"], "semantic_verification": "REQUIRES_SOURCE_REVIEW", "human_acceptance": "NOT_CLAIMED"}
    elif scope_proof:
        from scope_tools import scope_diff
        diff = scope_diff(previous, baseline, scope_proof)
    else:
        diff = source_diff(previous, baseline, old_values.get("annotations"), values.get("annotations"))
    diff["artifact_changes"] = {key: {"before_sha256": core.canonical_json_sha256(old_values.get(key)), "after_sha256": core.canonical_json_sha256(values.get(key))}
                                for key in set(old_values) | set(values) if old_values.get(key) != values.get(key)}
    diff["annotation_changes"] = annotation_diff(old_values.get("annotations"), values.get("annotations"))
    diff["annotation_changed_source_item_count"] = len(diff["annotation_changes"])
    baseline_affected = {change["source_item_id"] for change in diff.get("changes", [])} if isinstance(diff.get("changes"), list) else set((diff.get("source_rebinding") or {}).get("rebound_source_item_ids", []))
    baseline_affected.update(diff.get("created_source_item_ids", [])); baseline_affected.update(diff.get("retired_source_item_ids", [])); baseline_affected.update(diff.get("restored_source_item_ids", []))
    diff["total_affected_source_item_count"] = len(baseline_affected | {change["source_item_id"] for change in diff["annotation_changes"]})
    output = args.output_dir.expanduser().resolve() if args.output_dir else None
    if output is None or root not in output.parents:
        raise ValueError("Candidate output must be a new directory inside the review workspace")
    output = new_output(output)
    for name in ("document-map.json", "table-previews.json", "scope-requests.json", "map-compilation-receipt.json", "item-compilation-receipt.json"):
        if (directory / name).is_file() and not map_rebind:
            shutil.copyfile(directory / name, output / name)
    if map_rebind:
        save(output, "document-map.json", document_map.model_dump(mode="json")); save(output, "table-previews.json", tables)
        save(output, "native-structure-reference.json", map_result["native_reference"])
        save(output, "map-revision-receipt.json", map_result["receipt"]); save(output, "map-rebind-receipt.json", map_rebind["receipt"])
        text_requests = [core.build_source_item_text_scope_request(document_map, node_id=node.node_id).model_dump(mode="json") for node in document_map.nodes if node.source_content]
        table_requests = [core.build_source_requirement_scope_request(document_map, preview).model_dump(mode="json") for preview in tables.values()]
        save(output, "scope-requests.json", {"text_requests": text_requests, "table_requests": table_requests})
    save(output, "source-requirement-baseline.json", baseline.model_dump(mode="json")); save(output, "validation-receipt.json", validation.model_dump(mode="json"))
    for key, item in values.items():
        save(output, STATE_FILES[key], item)
    if 'manual_decisions' in replacements:
        for name in ('manual-decision-import-receipt.json', 'user-decisions.json', 'previous-manual-decisions.json'):
            original = args.manual_decisions.parent / name
            if original.is_file(): shutil.copyfile(original, output / name)
    if evidence_path.is_file():
        from evidence_tools import copy_registration
        copy_registration(evidence_path, output)
    save(output, "revision-diff.json", diff); save(output, "carry-forward.json", carry)
    if coordination:
        save(output, "classification-coordination.json", coordination.model_dump(mode="json")); save(output, "classification-proof.json", proof.model_dump(mode="json"))
    if scope_revision:
        save(output, "scope-revision.json", scope_revision.model_dump(mode="json")); save(output, "scope-revision-receipt.json", scope_proof.model_dump(mode="json"))
    save(output, "revision-input.json", value)
    save(output, "revision-receipt.json", {"schema_version": "reqpilot-review-revision-receipt/v1", "review_sha256": head["review_sha256"], "predecessor_baseline_sha256": previous.artifact_sha256,
        "baseline_sha256": baseline.artifact_sha256, "candidate_path": relative(root, output), "addressed_feedback_ids": addressed,
        "files": {path.name: digest_file(path) for path in sorted(output.iterdir()) if path.is_file()}, "model_execution": "CALLING_AGENT", "human_acceptance": "NOT_CLAIMED"})
    return summary("REVISION_COMPILED", root, head, candidate_path=relative(root, output), candidate_baseline_sha256=baseline.artifact_sha256,
        changed_source_item_count=diff["changed_source_item_count"], annotation_changed_source_item_count=diff["annotation_changed_source_item_count"],
        total_affected_source_item_count=diff["total_affected_source_item_count"], carry_forward=carry)


def compile_scope_transition(root, directory, document_map, previous, receipt, tables, value):
    from scope_tools import compile_scope
    scope = value["scope_revision"]
    old_receipt = None; undo_target = None
    receipt_path = directory / "scope-revision-receipt.json"
    if receipt_path.is_file():
        old_receipt = core.SourceItemScopeRevisionReceiptV1.model_validate(read(receipt_path))
        if old_receipt.artifact_sha256 != previous.artifact_sha256:
            raise ValueError("Scope identity receipt belongs to a different predecessor; do not re-sign an old proof")
    if scope.get("operation", "REPLACE") == "UNDO":
        if not scope.get("undo_target_path"):
            raise ValueError("Scope undo requires the immutable direct predecessor's workspace path")
        _, undo_target, _, _ = bundle(inside(root, scope["undo_target_path"]))
    return compile_scope(document_map, previous, receipt, tables, scope, feedback_ids=value.get("addressed_feedback_ids", []),
        previous_scope_receipt=old_receipt, undo_target=undo_target,
        source_display_high_watermark_floor=core.validate_source_identity_allocation(previous, value.get("source_identity_allocation")))


def verify_transition(root, previous_dir, current_dir, document_map, previous, current, tables,
        *, expected_allocation=None, expected_review_sha256=None):
    """Only shared-compiler replay can authorize changed source partitions."""
    if previous_dir != current_dir:
        from manual_tools import verify_manual_transition
        verify_manual_transition(previous_dir, current_dir)
    input_path = current_dir / "revision-input.json"
    if not input_path.is_file() or (previous.artifact_sha256 == current.artifact_sha256 and previous_dir == current_dir):
        return source_diff(previous, current, state(previous_dir).get("annotations"), state(current_dir).get("annotations"))
    value = read(input_path)
    allocation = value.get("source_identity_allocation")
    allocation_floor = core.validate_source_identity_allocation(previous, allocation)
    if allocation is not None:
        if expected_review_sha256 is not None and value.get("expected_review_sha256") != expected_review_sha256:
            raise ValueError("Source identity allocation is not bound to the actual previous review")
        if expected_allocation is None:
            expected_allocation = _verified_source_identity_allocation(root, value["expected_review_sha256"], previous)
        if allocation != expected_allocation:
            raise ValueError("Source identity allocation differs from replaying the immutable review history")
    if previous.artifact_sha256 == current.artifact_sha256:
        return source_diff(previous, current, state(previous_dir).get("annotations"), state(current_dir).get("annotations"))
    if value.get("map_revision"):
        from map_tools import compile_map_transition
        previous_map, _, previous_receipt, previous_tables = bundle(previous_dir)
        mapped, rebound = compile_map_transition(previous_map, previous, previous_receipt, previous_tables, inside(root, value["map_revision"]["candidate_map_path"]),
            text_scope_recovery=value.get("scope_revision"), feedback_ids=value.get("addressed_feedback_ids", []), source_identity_allocation=allocation)
        if (mapped["document_map"] != document_map.model_dump(mode="json") or rebound["baseline"] != current.model_dump(mode="json")
                or rebound["table_previews"] != tables or read(current_dir / "map-revision-receipt.json") != mapped["receipt"]
                or read(current_dir / "map-rebind-receipt.json") != rebound["receipt"]):
            raise ValueError("Map successor differs from replaying original source decisions and the shared lineage compiler")
        return {"transition_kind": "DOCUMENT_MAP_REVISION", "map_receipt": mapped["receipt"], "rebind_receipt": rebound["receipt"]}
    if value.get("scope_revision"):
        _, _, previous_receipt, _ = bundle(previous_dir)
        expected, validation, revision, proof = compile_scope_transition(root, previous_dir, document_map, previous, previous_receipt, tables, value)
        if (expected != current or read(current_dir / "scope-revision.json") != revision.model_dump(mode="json")
                or read(current_dir / "scope-revision-receipt.json") != proof.model_dump(mode="json")
                or read(current_dir / "validation-receipt.json") != validation.model_dump(mode="json")):
            raise ValueError("Source-scope successor differs from replaying the actual bound decisions through the shared compiler")
        from scope_tools import scope_diff
        return scope_diff(previous, current, proof)
    if value.get("decisions") and allocation is not None:
        _, _, previous_receipt, _ = bundle(previous_dir)
        expected, validation, coordination, proof = core.compile_source_requirement_classification_coordination(
            previous_artifact=previous, previous_validation_receipt=previous_receipt,
            decisions=tuple(core.SourceRequirementClassificationCoordinationDecisionV1.model_validate(item) for item in value["decisions"]),
            model_id=value["model_id"], coordinator_role="PORTABLE_AGENT", source_display_high_watermark_floor=allocation_floor)
        if (expected != current or read(current_dir / "classification-coordination.json") != coordination.model_dump(mode="json")
                or read(current_dir / "classification-proof.json") != proof.model_dump(mode="json")
                or read(current_dir / "validation-receipt.json") != validation.model_dump(mode="json")):
            raise ValueError("Classification successor differs from the history-bound compiler replay")
    return source_diff(previous, current, state(previous_dir).get("annotations"), state(current_dir).get("annotations"))


def carry_state(document_map, previous, baseline, tables, old_values, values, *, previous_tables=None, map_revision_receipt=None):
    """Invoke shared lineage-aware carry; keep obsolete state in prior version."""
    report = {"schema_version": "reqpilot-carry-forward/v1", "decomposition": "ABSENT", "assessments": "ABSENT", "stale_parent_ids": [], "stale_unit_ids": []}
    if old_values.get("decomposition") and previous.artifact_sha256 != baseline.artifact_sha256:
        from reqpilotx.application.agent_workflow.source_item_decomposition import rebind_source_item_decomposition
        details = rebind_source_item_decomposition(previous, baseline, old_values["decomposition"], table_previews=tables,
            previous_table_previews=previous_tables, map_revision_receipt=map_revision_receipt)
        values["decomposition"] = details["artifact"]
        report.update(decomposition="REBOUND", stale_parent_ids=details["stale_parent_ids"])
    elif values.get("decomposition"):
        report["decomposition"] = "UNCHANGED"
    if old_values.get("assessments"):
        from reqpilotx.application.agent_workflow.source_item_analysis import build_source_item_analysis_packet, rebind_source_item_assessments
        old_packet = old_values["analysis_packet"]
        catalog = values.get("knowledge_catalog")
        if catalog is None:
            catalog = empty_catalog_from_packet(old_packet)
        packet = build_source_item_analysis_packet(document_map, baseline, tables, catalog, decomposition=values.get("decomposition"), image_annotations=values.get("annotations"))
        details = rebind_source_item_assessments(old_packet, old_values["assessments"], packet, map_revision_receipt=map_revision_receipt, next_baseline=baseline)
        values["assessments"] = details["artifact"]
        values["analysis_packet"] = packet
        report.update(assessments="REBOUND", stale_unit_ids=details["stale_unit_ids"])
    return report


def close_review(args):
    root = required(args.review_dir, "--review-dir"); value = read(required(args.input, "--input"))
    if getattr(args, 'manual_decisions', None):
        raise ValueError('Engineer opinions must enter through an explicit user-import review revision before close')
    with locked(root):
        head = load_head(root, value["expected_review_sha256"])
        previous_dir, (old_map, previous, _, _) = current_bundle(root, head)
        candidate = inside(root, value["candidate_path"]); document_map, baseline, _, tables = bundle(candidate)
        receipt = read(candidate / "revision-receipt.json")
        if receipt.get("review_sha256") != head["review_sha256"] or receipt.get("predecessor_baseline_sha256") != previous.artifact_sha256:
            raise ValueError("Candidate was produced against a different review or predecessor")
        for name, digest in receipt["files"].items():
            if digest_file(inside(candidate, name)) != digest:
                raise ValueError("Candidate changed after compilation: " + name)
        if baseline.source_sha256 != previous.source_sha256:
            raise ValueError("A review successor cannot replace the original source")
        if old_map.artifact_sha256 != document_map.artifact_sha256 and not read(candidate / "revision-input.json").get("map_revision"):
            raise ValueError("A changed Map requires explicit source decisions and shared Map lineage replay")
        verify_transition(root, previous_dir, candidate, document_map, previous, baseline, tables, expected_review_sha256=head["review_sha256"])
        values = state(candidate)
        replacements = {}
        for key in STATE_FILES:
            if getattr(args, key, None):
                replacements[key] = read(required(getattr(args, key), "--" + key.replace("_", "-")))
        values.update(replacements)
        evidence_path = getattr(args, "knowledge_catalog", None) or candidate / "knowledge-catalog.json"
        validate_state(document_map, baseline, tables, values, evidence_path)
        delivery = inside(root, value["delivery_path"]); exported = read(delivery / "delivery-receipt.json")
        if exported.get("baseline_artifact_sha256") != baseline.artifact_sha256 or exported.get("source_sha256") != baseline.source_sha256 or not exported.get("source_document_bound_for_visual_export") or digest_file(delivery / "matrix.xlsx") != exported.get("xlsx_sha256"):
            raise ValueError("Delivery must bind the exact candidate and unchanged original source")
        from reqpilotx.application.agent_workflow.source_item_delivery import build_delivery_snapshot
        expected_snapshot = build_delivery_snapshot(document_map, baseline, tables, annotations=values.get("annotations"), decomposition=values.get("decomposition"), assessments=values.get("assessments"), manual_decisions=values.get('manual_decisions'), analysis_policy=values.get('analysis_policy'))
        if read(delivery / "matrix-snapshot.json") != expected_snapshot:
            raise ValueError("Delivery omits or changes the candidate's current analysis, derivation or annotation state")
        updated = deepcopy(head); by_id = {item["feedback_id"]: item for item in updated["feedback"]}; seen = set()
        for resolution in value.get("resolutions", []):
            fid = resolution["feedback_id"]
            if fid not in by_id or fid in seen or resolution.get("disposition") not in {"ADDRESSED", "UNRESOLVED"} or not str(resolution.get("reason", "")).strip():
                raise ValueError("Resolution needs unique existing feedback, a disposition and reason")
            if resolution["disposition"] == "ADDRESSED" and fid not in receipt["addressed_feedback_ids"]:
                raise ValueError("Addressed feedback must be bound to the compiled revision")
            seen.add(fid); by_id[fid].update(status=resolution["disposition"], resolution_reason=resolution["reason"], resolved_in_baseline_sha256=baseline.artifact_sha256, resolved_in_review_revision=head["review_revision"] + 1)
        if replacements:
            # New compiler outputs are composed as a new immutable version; the
            # earlier candidate and its hash-bound receipt remain unchanged.
            final_dir = root / "versions" / ("final-" + core.canonical_json_sha256({"review": head["review_sha256"], "candidate": value["candidate_path"], "state": values})[:24])
            final_dir.mkdir(mode=0o700)
            for path in candidate.iterdir():
                if path.is_file() and path.name not in {*STATE_FILES.values(), "evidence-registration.json"}:
                    shutil.copyfile(path, final_dir / path.name)
            for key, item in values.items():
                save(final_dir, STATE_FILES[key], item)
            if evidence_path.is_file():
                from evidence_tools import copy_registration
                copy_registration(evidence_path, final_dir)
            save(final_dir, "composition-receipt.json", {"source_candidate_path": value["candidate_path"], "replaced_artifact_kinds": sorted(replacements), "source_candidate_receipt_sha256": digest_file(candidate / "revision-receipt.json")})
            candidate = final_dir
        updated.update(current=fingerprint(root, candidate, baseline, document_map), review_revision=head["review_revision"] + 1, predecessor_review_sha256=head["review_sha256"],
            delivery={"path": relative(root, delivery), "xlsx_sha256": exported["xlsx_sha256"], "baseline_sha256": baseline.artifact_sha256})
        updated.pop('delivery_attachment',None) # This delivery belongs to the actual new candidate, not a former unchanged draft.
        updated = write_head(root, updated)
    return summary("REVIEW_ADVANCED", root, updated, delivery_path=relative(root, delivery), open_feedback_count=sum(f["status"] != "ADDRESSED" for f in updated["feedback"]), human_acceptance="NOT_CLAIMED")


def _delivery_files(directory):
    files={}
    for path in sorted(directory.rglob('*')):
        if path.is_symlink():raise ValueError('Delivery attachment cannot follow symlinks')
        if path.is_file():files[str(path.relative_to(directory))]=digest_file(path)
    return files


def _check_current_delivery(root,head,directory):
    """Rebuild every current state layer and check the exact exported bytes."""
    from reqpilotx.application.agent_workflow.source_item_delivery import build_delivery_snapshot
    current,(mapping,baseline,_,tables)=current_bundle(root,head)
    values=state(current)
    validate_state(mapping,baseline,tables,values,current/'knowledge-catalog.json')
    exported=read(directory/'delivery-receipt.json')
    if (exported.get('source_sha256')!=head['source']['sha256']
            or exported.get('baseline_artifact_sha256')!=baseline.artifact_sha256
            or not exported.get('source_document_bound_for_visual_export')
            or exported.get('xlsx_sha256')!=digest_file(directory/'matrix.xlsx')):
        raise ValueError('Delivery must bind the current baseline, original source and exact workbook bytes')
    if mapping.source_format:
        from reqpilotx.application.agent_workflow.office_source_evidence import verify_office_map_against_source
        verify_office_map_against_source(inside(root,head['source']['path']),mapping,tables,baseline=baseline)
    expected=build_delivery_snapshot(mapping,baseline,tables,annotations=values.get('annotations'),decomposition=values.get('decomposition'),
        assessments=values.get('assessments'),manual_decisions=values.get('manual_decisions'),analysis_policy=values.get('analysis_policy'))
    if read(directory/'matrix-snapshot.json')!=expected:
        raise ValueError('Delivery differs from the current Map/baseline or annotation, decomposition, assessment, opinion or policy state')
    snapshot_sha=core.canonical_json_sha256(expected)
    if published_workbook_snapshot_sha256(directory/'matrix.xlsx')!=snapshot_sha:
        raise ValueError('Workbook snapshot metadata differs from the current saved snapshot')
    return exported,snapshot_sha


def attach_current_delivery(args):
    """Register an unchanged current review's delivery, without a fake candidate."""
    root=required(args.review_dir,'--review-dir').resolve(strict=True)
    delivery=required(getattr(args,'delivery_dir',None),'--delivery-dir').resolve(strict=True)
    expected=getattr(args,'expected_review_sha256',None)
    if not isinstance(expected,str) or len(expected)!=64:
        raise ValueError('Delivery attachment requires the exact --expected-review-sha256')
    if any(getattr(args,key,None) for key in (*STATE_FILES,'input')):
        raise ValueError('Delivery attachment cannot introduce replacement state; use a real review revision for changes')
    if not delivery.is_dir() or delivery==root or delivery in root.parents:
        raise ValueError('Select the exact saved delivery directory')
    # When the export belongs to a task, only its committed operation is eligible.
    # Do this outside the review lock: core inspection may read this same owner.
    from reqpilotx.application.agent_workflow.task_workspace import find_task_workspace,workspace_status
    explicit_task=getattr(args,'workspace',None)
    task=find_task_workspace(start=delivery,explicit=explicit_task)
    if explicit_task is not None and task is None:
        raise ValueError('Explicit --workspace has no task-context.json; it cannot fall back to a standalone review')
    operation=None
    if task:
        task_root,_=task;task_state=workspace_status(task_root,refresh=False)
        owner=task_state['current_working']
        if owner.get('owner')!='portable_review' or Path(owner.get('review_root','')).resolve()!=root:
            raise ValueError('Task current review owner differs from the requested attachment')
        operation_path=delivery/'operation-receipt.json'
        if operation_path.is_file():
            operation=read(operation_path)
            entry=next((v for v in task_state['index']['committed'] if v['operation_id']==operation.get('operation_id')),None)
            if (not entry or entry['state']!='COMMITTED' or entry['operation']!='deliver'
                    or ((task_root/entry['receipt_path']).resolve()!=operation_path and not delivery.is_relative_to(root))
                    or digest_file(task_root/entry['receipt_path'])!=digest_file(operation_path)
                    or entry['receipt_sha256']!=operation.get('receipt_sha256')):
                raise ValueError('Only the exact COMMITTED task delivery can be attached; a candidate or unregistered output cannot advance the owner')
        elif delivery.is_relative_to(task_root) and not delivery.is_relative_to(root):
            raise ValueError('Task output has no COMMITTED delivery operation receipt')
    with locked(root):
        head=load_head(root)
        prior=head.get('delivery_attachment') or {}
        retry=(head.get('predecessor_review_sha256')==expected and prior.get('expected_review_sha256')==expected)
        if head['review_sha256']!=expected and not retry:
            raise ValueError('Review changed since delivery was prepared; reload the current owner before attaching')
        exported,snapshot_sha=_check_current_delivery(root,head,delivery)
        files=_delivery_files(delivery)
        identity={'source_sha256':head['source']['sha256'],'baseline_sha256':head['current']['baseline_sha256'],
            'document_map_sha256':head['current']['document_map_sha256'],'snapshot_sha256':snapshot_sha,'files':files}
        attachment_sha=core.canonical_json_sha256(identity)
        same=prior.get('attachment_sha256')==attachment_sha and head.get('delivery',{}).get('xlsx_sha256')==exported['xlsx_sha256']
        if retry and not same:
            raise ValueError('A retry with the old review SHA must reference the exact previously attached delivery')
        if operation and operation.get('expected_review_sha256') not in {expected,prior.get('expected_review_sha256') if same else expected}:
            raise ValueError('Task delivery was compiled against a different review head')
        if same:
            owned=inside(root,head['delivery']['path'])
            if _delivery_files(owned)!=files:raise ValueError('Previously attached delivery bytes changed')
            _check_current_delivery(root,head,owned)
            return summary('CURRENT_DELIVERY_ATTACHED',root,head,delivery_path=head['delivery']['path'],idempotent=True,
                deterministic_validation='PASSED',human_acceptance='NOT_CLAIMED',business_revision_created=False)
        destination=root/'deliveries'/('current-'+attachment_sha)
        destination.parent.mkdir(exist_ok=True,mode=0o700)
        if destination.exists():
            if _delivery_files(destination)!=files:raise ValueError('Content-addressed review delivery changed')
        else:
            staging=Path(tempfile.mkdtemp(prefix='.attach-',dir=destination.parent))
            try:
                for name in files:
                    target=staging/name;target.parent.mkdir(parents=True,exist_ok=True)
                    shutil.copyfile(delivery/name,target);target.chmod(0o600)
                if _delivery_files(staging)!=files:raise ValueError('Delivery changed while copying its immutable attachment')
                _check_current_delivery(root,head,staging)
                os.rename(staging,destination)
            finally:
                if staging.exists():shutil.rmtree(staging)
        updated=deepcopy(head)
        updated.update(review_revision=head['review_revision']+1,predecessor_review_sha256=head['review_sha256'],
            delivery={'path':relative(root,destination),'xlsx_sha256':exported['xlsx_sha256'],'baseline_sha256':head['current']['baseline_sha256']},
            delivery_attachment={'schema_version':'reqpilot-current-delivery-attachment/v1','expected_review_sha256':head['review_sha256'],
                'attachment_sha256':attachment_sha,'snapshot_sha256':snapshot_sha,'current_version_path':head['current']['version_path'],
                'source_operation_id':operation.get('operation_id') if operation else None,'human_acceptance':'NOT_CLAIMED'})
        updated=write_head(root,updated)
    return summary('CURRENT_DELIVERY_ATTACHED',root,updated,delivery_path=relative(root,destination),idempotent=False,
        deterministic_validation='PASSED',human_acceptance='NOT_CLAIMED',business_revision_created=False,
        next='The unchanged current draft and exact Excel are registered. Use review-verify for the handoff check; create a revision only for an actual supported change.')


def verify_review(args):
    """Read-only publication evidence; shared source/analysis/export validation."""
    from reqpilotx.application.agent_workflow.source_item_delivery import build_delivery_snapshot
    root = required(args.review_dir, "--review-dir")
    with locked(root):
        head = load_head(root)
        chain, seen, current = [], set(), head
        while current:
            digest = current["review_sha256"]
            if digest in seen or seal(current)["review_sha256"] != digest:
                raise ValueError("Review history integrity or ancestry changed")
            seen.add(digest)
            snapshot = read(root / "reviews" / (digest + ".json"))
            if snapshot != current:
                raise ValueError("Review head does not match its immutable snapshot")
            directory, (document_map, baseline, _, tables) = current_bundle(root, current)
            if baseline.source_sha256 != head["source"]["sha256"]:
                raise ValueError("Review history replaced the original source")
            if document_map.source_format:
                from reqpilotx.application.agent_workflow.office_source_evidence import verify_office_map_against_source
                verify_office_map_against_source(inside(root, head["source"]["path"]), document_map, tables, baseline=baseline)
            values = state(directory)
            validate_state(document_map, baseline, tables, values, directory / "knowledge-catalog.json")
            chain.append((current, baseline, values))
            predecessor = current.get("predecessor_review_sha256")
            current = read(root / "reviews" / (predecessor + ".json")) if predecessor else None
        chain.reverse()
        first_head, first_baseline, _ = chain[0]
        identity_floor = _identity_watermark(first_head, first_baseline, inside(root, first_head["current"]["version_path"]))
        identity_prefix = [first_head["review_sha256"]]
        for before, after in zip(chain, chain[1:]):
            old, previous, old_values = before; new, baseline, values = after
            if new["review_revision"] != old["review_revision"] + 1:
                raise ValueError("Review revision sequence is not contiguous")
            if baseline.artifact_sha256 != previous.artifact_sha256 and baseline.predecessor_artifact_sha256 != previous.artifact_sha256:
                raise ValueError("Review baseline predecessor is not the actual former working version")
            previous_dir = inside(root, old["current"]["version_path"]); current_dir = inside(root, new["current"]["version_path"])
            verify_transition(root, previous_dir, current_dir, core.DocumentMapArtifactV1.model_validate(read(current_dir / "document-map.json")), previous, baseline, read(current_dir / "table-previews.json"),
                expected_allocation=_identity_context(old, previous, identity_floor, identity_prefix), expected_review_sha256=old["review_sha256"])
            identity_floor = max(identity_floor, _identity_watermark(new, baseline, current_dir))
            identity_prefix.append(new["review_sha256"])
            old_feedback = {f["feedback_id"]: f for f in old["feedback"]}
            new_feedback = {f["feedback_id"]: f for f in new["feedback"]}
            if set(old_feedback) - set(new_feedback):
                raise ValueError("Review history lost feedback")
            for fid, item in old_feedback.items():
                for key in ("origin", "text", "source_item_ids", "bound_baseline_sha256", "created_review_revision"):
                    if item.get(key) != new_feedback[fid].get(key):
                        raise ValueError("Feedback provenance changed across reviews")
        if not head.get("delivery"):
            raise ValueError("Review has no registered current delivery. Use workspace run deliver or review-attach-delivery for an unchanged current draft; use review-close only for a real revision candidate.")
        directory, (document_map, baseline, _, tables) = current_bundle(root, head); values = state(directory)
        delivery = inside(root, head["delivery"]["path"]); receipt = read(delivery / "delivery-receipt.json")
        if (receipt.get("source_sha256") != baseline.source_sha256 or receipt.get("baseline_artifact_sha256") != baseline.artifact_sha256
                or not receipt.get("source_document_bound_for_visual_export") or digest_file(delivery / "matrix.xlsx") != receipt.get("xlsx_sha256")
                or head["delivery"].get("xlsx_sha256") != receipt.get("xlsx_sha256") or head["delivery"].get("baseline_sha256") != baseline.artifact_sha256):
            raise ValueError("Current exported workbook lineage or bytes changed")
        snapshot = build_delivery_snapshot(document_map, baseline, tables, annotations=values.get("annotations"), decomposition=values.get("decomposition"), assessments=values.get("assessments"), manual_decisions=values.get('manual_decisions'), analysis_policy=values.get('analysis_policy'))
        published = read(delivery / "matrix-snapshot.json")
        snapshot_pin = published_workbook_snapshot_sha256(delivery / "matrix.xlsx")
        compatibility = None
        if published != snapshot:
            from reqpilotx.application.agent_workflow.delivery_snapshot_compatibility import verify_published_delivery_snapshot
            compatibility = verify_published_delivery_snapshot(document_map=document_map, baseline=baseline,
                published_snapshot=published, current_snapshot=snapshot, expected_snapshot_sha256=snapshot_pin,
                annotations=values.get('annotations'))
        elif core.canonical_json_sha256(published) != snapshot_pin:
            raise ValueError('Current export snapshot differs from its immutable workbook metadata')
        nodes = {(node.get("source_item") or {}).get("source_item_id", node["node_id"]): node for node in snapshot["nodes"]}
        findings = [{"code": "NUMBERED_CLAUSE_NUMBER_MISSING", "source_item_id": record.source_record_id, "source_display_id": record.source_display_id,
            "state": "PENDING_SOURCE_REVIEW", "message": "This item is labelled NUMBERED_CLAUSE but its independent original-number projection is empty; read its own original bounds or correct the source form."}
            for record in baseline.records if record.source_form.value == "NUMBERED_CLAUSE" and not str(nodes.get(record.source_record_id, {}).get("source_number", "")).strip()]
        from reqpilotx.application.agent_workflow.source_item_workbench_metadata import numbering_findings
        numbering = numbering_findings(document_map)
        findings.extend(numbering['findings'])
        return {"status": "REVIEW_VERIFIED", "review_sha256": head["review_sha256"], "source_sha256": head["source"]["sha256"], "current": head["current"],
            "delivery": head["delivery"], "feedback": head["feedback"], "version_chain": [{"review_sha256": item[0]["review_sha256"], "review_revision": item[0]["review_revision"], **item[0]["current"]} for item in chain],
            "source_consistency": "HAS_FINDINGS" if findings else "NO_FINDINGS_IN_CHECKED_SCOPE", "consistency_findings": findings,
            "source_findings": numbering,
            "semantic_verification": "REQUIRES_SOURCE_REVIEW", "human_acceptance": "NOT_CLAIMED",
            **({'published_snapshot_compatibility': compatibility} if compatibility else {})}


def run_review(args):
    return {"review-open": open_review, "review-feedback": feedback_review, "review-packet": packet_review, "review-revise": revise_review, "review-close": close_review, "review-attach-delivery": attach_current_delivery, "review-verify": verify_review}[args.command](args)
