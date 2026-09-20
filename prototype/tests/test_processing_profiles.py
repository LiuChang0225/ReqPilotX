from __future__ import annotations

# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import fitz
from docx import Document

from app import db
from app.application import processing_profiles
from app.application.extraction import build_processing_plan_preview
from app.application.extraction import processing_pipeline
from app.application.errors import ApplicationError
from app.application.extraction import document_jobs, structure_workspace
from app.application.extraction.build_llm_shadow_preview import (
    LlmShadowPreviewRejected,
)
from app.application.extraction.processing_pipeline import (
    LOCAL_STRUCTURE_PROJECTION_MODE,
    build_local_structure_document_result,
    processing_plan_visual_strategy,
)
from app.application.extraction.source_page_facts import (
    SOURCE_PAGE_PREVIEW_KEY,
)
from app.application.extraction.structure_presentation import (
    STRUCTURE_PRESENTATION_INDEX_SCHEMA_VERSION,
)
from app.domain.extraction.processing_profile import (
    CapabilityState,
    DocumentProcessingProfileV1,
    ProjectProcessingProfileUpdateV1,
    ResolvedProviderBindingsV1,
    RunProfileOverrideV1,
    RuntimeCapabilityV1,
    build_local_rules_profile,
    build_resolved_processing_plan,
    processing_plan_requires_connection,
)


def _capability(
    capability_id: str,
    state: CapabilityState = CapabilityState.READY,
) -> RuntimeCapabilityV1:
    return RuntimeCapabilityV1(
        capability_id=capability_id,
        category="document_parser",
        provider="test",
        label=capability_id,
        state=state,
        reason_code=("" if state is CapabilityState.READY else "not_configured"),
        message=("" if state is CapabilityState.READY else "Not configured."),
        blocked_by=(() if state is CapabilityState.READY else ("setting",)),
        connection_revision=1,
    )


class ProcessingProfileIdentityTests(unittest.TestCase):
    def test_fast_bailian_ocr_is_a_distinct_frozen_enhancement(self) -> None:
        document = DocumentProcessingProfileV1(
            primary_parser="pymupdf",
            enhancement_chain=("bailian_ocr",),
            complex_vision=None,
            auto_escalation=False,
        )

        self.assertEqual(document.enhancement_chain, ("bailian_ocr",))
        self.assertEqual(
            processing_profiles.required_capability_ids(
                document=document,
                llm_provider="none",
                knowledge_backend="none",
            ),
            (
                processing_profiles.CAPABILITY_LOCAL_PYMUPDF,
                processing_profiles.CAPABILITY_BAILIAN_OCR,
            ),
        )
        with (
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                return_value=SimpleNamespace(chat_model="qwen3.8-max"),
            ),
            patch.object(
                processing_profiles,
                "get_runtime_ai_config_value",
                side_effect=lambda _name, default=None: default or "",
            ),
        ):
            bindings = processing_profiles._resolve_provider_bindings(
                document=document,
                llm_provider="none",
            )
        self.assertEqual(
            bindings.enhancement_providers,
            ("aliyun:bailian-ocr:qwen3.5-ocr",),
        )

        with self.assertRaisesRegex(ValueError, "mutually exclusive"):
            DocumentProcessingProfileV1(
                primary_parser="pymupdf",
                enhancement_chain=("bailian_ocr", "bailian_layout"),
            )

        with self.assertRaisesRegex(ValueError, "requires pymupdf"):
            DocumentProcessingProfileV1(
                primary_parser="azure_document_intelligence",
                enhancement_chain=("bailian_ocr",),
            )

    def test_visual_provider_telemetry_reaches_the_actual_composable_builder(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.pdf"
            document = fitz.open()
            document.new_page().insert_text((72, 72), "Visual page")
            document.save(source)
            document.close()
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            profile = build_local_rules_profile(project_id="project-telemetry")
            plan = build_resolved_processing_plan(
                project_id="project-telemetry",
                document_id="document-telemetry",
                source_sha256=source_sha256,
                profile=profile,
                document=DocumentProcessingProfileV1(
                    primary_parser="pymupdf",
                    enhancement_chain=("bailian_ocr",),
                    auto_escalation=False,
                ),
                llm_provider="none",
                provider_bindings=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf",
                    enhancement_providers=("aliyun:bailian-ocr:qwen3.5-ocr",),
                ),
                connection_revision=1,
            )
            telemetry_callback = Mock()
            adapter = object()
            source_ir = object()
            with (
                patch(
                    "app.application.extraction.source_asset_objects."
                    "runtime_source_asset_publisher",
                    return_value=None,
                ),
                patch(
                    "app.application.extraction.processing_visual_chain."
                    "build_composable_pdf_source_ir_adapter",
                    return_value=adapter,
                ) as build_adapter,
                patch(
                    "app.application.extraction.processing_visual_chain."
                    "page_processing_receipts_from_adapter",
                    return_value=(),
                ),
                patch.object(
                    processing_pipeline,
                    "build_source_ir",
                    return_value=source_ir,
                ),
            ):
                actual_source_ir, receipts = (
                    processing_pipeline.build_source_ir_and_receipts_for_plan(
                        source,
                        plan=plan,
                        provider_telemetry_callback=telemetry_callback,
                    )
                )

        self.assertIs(actual_source_ir, source_ir)
        self.assertEqual(receipts, ())
        self.assertIs(
            build_adapter.call_args.kwargs["provider_telemetry_callback"],
            telemetry_callback,
        )

    def test_processing_plan_preview_forwards_visual_provider_telemetry(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.pdf"
            document = fitz.open()
            document.new_page().insert_text((72, 72), "Visual page")
            document.save(source)
            document.close()
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            profile = build_local_rules_profile(project_id="project-telemetry")
            plan = build_resolved_processing_plan(
                project_id="project-telemetry",
                document_id="document-telemetry",
                source_sha256=source_sha256,
                profile=profile,
                document=DocumentProcessingProfileV1(
                    primary_parser="pymupdf",
                    enhancement_chain=("bailian_ocr",),
                    auto_escalation=False,
                ),
                llm_provider="aliyun",
                provider_bindings=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf",
                    enhancement_providers=("aliyun:bailian-ocr:qwen3.5-ocr",),
                    llm_transport_provider="aliyun",
                ),
                connection_revision=1,
            )
            telemetry_callback = Mock()
            with (
                patch.object(
                    build_processing_plan_preview,
                    "_validate_document_request",
                    return_value=(source_sha256, ()),
                ),
                patch.object(
                    build_processing_plan_preview,
                    "build_source_ir_and_receipts_for_plan",
                    side_effect=RuntimeError("stop after callback seam"),
                ) as build_source,
                self.assertRaises(LlmShadowPreviewRejected),
            ):
                build_processing_plan_preview.build_processing_plan_document_preview(
                    source,
                    plan=plan,
                    provider=object(),  # type: ignore[arg-type]
                    provider_telemetry_callback=telemetry_callback,
                )

        self.assertIs(
            build_source.call_args.kwargs["provider_telemetry_callback"],
            telemetry_callback,
        )

    def test_processing_plan_preview_blocks_visual_gaps_before_semantic_provider(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.pdf"
            document = fitz.open()
            document.new_page().insert_text((72, 72), "Visual page")
            document.save(source)
            document.close()
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            profile = build_local_rules_profile(project_id="project-visual-gate")
            plan = build_resolved_processing_plan(
                project_id="project-visual-gate",
                document_id="document-visual-gate",
                source_sha256=source_sha256,
                profile=profile,
                document=DocumentProcessingProfileV1(
                    primary_parser="pymupdf",
                    enhancement_chain=("bailian_ocr",),
                    auto_escalation=False,
                ),
                llm_provider="aliyun",
                provider_bindings=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf",
                    enhancement_providers=("aliyun:bailian-ocr:qwen3.5-ocr",),
                    llm_transport_provider="aliyun",
                ),
                connection_revision=1,
            )

            for cause_code in (
                "pdf_visual_text_authentication_failed",
                "pdf_ocr_authority_text_unresolved",
                "pdf_ocr_authority_table_text_unresolved",
            ):
                with self.subTest(cause_code=cause_code):
                    source_ir = object()
                    tree = object()
                    semantic_provider = object()
                    with (
                        patch.object(
                            build_processing_plan_preview,
                            "_validate_document_request",
                            return_value=(source_sha256, ()),
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "build_source_ir_and_receipts_for_plan",
                            return_value=(source_ir, ()),
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "build_structure_tree",
                            return_value=tree,
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "initialize_structure_repair_without_provider",
                            return_value=SimpleNamespace(effective_structure=tree),
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "locate_llm_document_scope",
                            return_value=SimpleNamespace(
                                scope_structure_node_ids=("scope",)
                            ),
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "_scope_source_page_numbers",
                            return_value=(1,),
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "_scoped_visual_incompleteness_cause_code",
                            return_value=cause_code,
                        ),
                        patch.object(
                            build_processing_plan_preview,
                            "extract_requirement_document_with_provider",
                        ) as extract,
                        self.assertRaises(LlmShadowPreviewRejected) as caught,
                    ):
                        build_processing_plan_preview.build_processing_plan_document_preview(
                            source,
                            plan=plan,
                            provider=semantic_provider,  # type: ignore[arg-type]
                        )

                    self.assertEqual(caught.exception.stage, "scoped_visual_gate")
                    self.assertEqual(caught.exception.cause_code, cause_code)
                    extract.assert_not_called()

    def test_profile_and_resolved_plan_identity_are_canonical(self) -> None:
        profile = build_local_rules_profile(project_id="project-a")
        bindings = ResolvedProviderBindingsV1(primary_parser_provider="local:pymupdf")
        first = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=profile.document,
            llm_provider="none",
            provider_bindings=bindings,
            connection_revision=7,
        )
        replay = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=profile.document,
            llm_provider="none",
            provider_bindings=bindings,
            connection_revision=7,
        )
        forced = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=DocumentProcessingProfileV1(
                primary_parser="pymupdf",
                complex_vision="azure_gpt56_sol",
            ),
            llm_provider="none",
            provider_bindings=ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf",
                complex_vision_provider="azure-openai:gpt56-sol:vision",
                complex_vision_deployment="sol-deployment",
            ),
            connection_revision=7,
            force_visual_pages=(4,),
        )

        self.assertEqual(first, replay)
        self.assertEqual(first.plan_id, replay.plan_id)
        self.assertEqual(first.canonical_sha256, replay.canonical_sha256)
        self.assertNotEqual(first.plan_id, forced.plan_id)
        serialized = json.dumps(forced.model_dump(mode="json")).lower()
        self.assertNotIn("api_key", serialized)
        self.assertNotIn("credential", serialized)
        self.assertNotIn("bearer", serialized)

    def test_connection_dependency_is_derived_from_frozen_provider_bindings(
        self,
    ) -> None:
        profile = build_local_rules_profile(project_id="project-a")
        local = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=profile.document,
            llm_provider="none",
            provider_bindings=ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf"
            ),
            connection_revision=17,
        )
        remote_llm = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=profile.document,
            llm_provider="none",
            provider_bindings=ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf",
                llm_transport_provider="aliyun",
            ),
            connection_revision=17,
        )

        self.assertFalse(processing_plan_requires_connection(local))
        self.assertTrue(processing_plan_requires_connection(remote_llm))

        remote_bindings = (
            ResolvedProviderBindingsV1(
                primary_parser_provider="azure:document-intelligence:prebuilt-layout"
            ),
            ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf",
                enhancement_providers=("aliyun:bailian-ocr:qwen3.5-ocr",),
            ),
            ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf",
                complex_vision_provider="aliyun:bailian-vision",
            ),
        )
        for bindings in remote_bindings:
            with self.subTest(bindings=bindings):
                plan = build_resolved_processing_plan(
                    project_id="project-a",
                    document_id="document-a",
                    source_sha256="a" * 64,
                    profile=profile,
                    document=profile.document,
                    llm_provider="none",
                    provider_bindings=bindings,
                    connection_revision=17,
                )
                self.assertTrue(processing_plan_requires_connection(plan))

    def test_resolved_local_plan_does_not_bind_global_connection_revision(self) -> None:
        profile = build_local_rules_profile(project_id="project-a")
        document = {
            "id": "document-a",
            "project_id": "project-a",
            "source_sha256": "a" * 64,
        }
        with (
            patch.object(db, "get_sor_document", return_value=document),
            patch.object(
                processing_profiles,
                "get_project_processing_profile",
                return_value=profile,
            ),
            patch.object(processing_profiles, "require_capabilities"),
            patch.object(
                processing_profiles,
                "_resolve_provider_bindings",
                return_value=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf"
                ),
            ),
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=27,
            ),
        ):
            plan = processing_profiles.resolve_processing_plan(
                project_id="project-a",
                document_id="document-a",
            )

        self.assertEqual(plan.connection_revision, 0)
        self.assertFalse(processing_plan_requires_connection(plan))

    def test_durable_execution_requirements_are_recomputed_from_the_plan(self) -> None:
        profile = build_local_rules_profile(project_id="project-a")
        local = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=profile.document,
            llm_provider="none",
            provider_bindings=ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf"
            ),
            connection_revision=19,
        )
        remote = build_resolved_processing_plan(
            project_id="project-a",
            document_id="document-a",
            source_sha256="a" * 64,
            profile=profile,
            document=profile.document,
            llm_provider="none",
            provider_bindings=ResolvedProviderBindingsV1(
                primary_parser_provider="local:pymupdf",
                enhancement_providers=("aliyun:bailian-ocr:qwen3.5-ocr",),
            ),
            connection_revision=19,
        )

        self.assertEqual(
            document_jobs._document_extraction_execution_requirements(local),
            {
                "schema_version": ("document-extraction-execution-requirements/v1"),
                "connection_revision_required": False,
                "required_connection_revision": None,
            },
        )
        self.assertEqual(
            document_jobs._document_extraction_execution_requirements(remote),
            {
                "schema_version": ("document-extraction-execution-requirements/v1"),
                "connection_revision_required": True,
                "required_connection_revision": 19,
            },
        )

    def test_run_override_rejects_unsorted_duplicate_or_zero_pages(self) -> None:
        with self.assertRaises(ValueError):
            RunProfileOverrideV1(force_visual_pages=(2, 1))
        with self.assertRaises(ValueError):
            RunProfileOverrideV1(force_visual_pages=(1, 1))
        with self.assertRaises(ValueError):
            RunProfileOverrideV1(force_visual_pages=(0,))


class ProjectProcessingProfilePersistenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.db_patch = patch.object(db, "DB_PATH", self.root / "app.db")
        self.storage_patch = patch.object(db, "STORAGE_DIR", self.root)
        self.db_patch.start()
        self.storage_patch.start()
        db.init_db()
        self.project = db.create_project("Processing profile", "Customer")

    def tearDown(self) -> None:
        self.storage_patch.stop()
        self.db_patch.stop()
        self.temporary.cleanup()

    def test_isolated_default_is_local_rules_and_update_is_optimistic(self) -> None:
        default = processing_profiles.get_project_processing_profile(self.project["id"])
        self.assertEqual(default.revision, 0)
        self.assertEqual(default.llm_provider, "none")
        self.assertEqual(default.knowledge_backend, "none")
        self.assertEqual(default.document.primary_parser, "pymupdf")

        with patch.object(
            processing_profiles,
            "build_runtime_capabilities",
            return_value=(_capability(processing_profiles.CAPABILITY_LOCAL_PYMUPDF),),
        ):
            updated = processing_profiles.update_project_processing_profile_route(
                self.project["id"],
                ProjectProcessingProfileUpdateV1(
                    expected_revision=0,
                    document=default.document,
                ),
            )["profile"]
            self.assertEqual(updated["revision"], 1)
            self.assertRegex(updated["canonical_sha256"], r"^[0-9a-f]{64}$")

            with self.assertRaises(ApplicationError) as caught:
                processing_profiles.update_project_processing_profile_route(
                    self.project["id"],
                    ProjectProcessingProfileUpdateV1(
                        expected_revision=0,
                        document=default.document,
                    ),
                )
        self.assertEqual(caught.exception.status_code, 409)
        self.assertEqual(
            caught.exception.detail["reason_code"],
            "stale_processing_profile_revision",
        )

    def test_knowledge_backend_change_requires_explicit_reindex(self) -> None:
        current = processing_profiles.get_project_processing_profile(self.project["id"])
        with self.assertRaises(ApplicationError) as caught:
            processing_profiles.update_project_processing_profile_route(
                self.project["id"],
                ProjectProcessingProfileUpdateV1(
                    expected_revision=0,
                    document=current.document,
                    knowledge_backend="aliyun_qdrant",
                    reindex_knowledge=False,
                ),
            )
        self.assertEqual(caught.exception.status_code, 409)
        self.assertEqual(
            caught.exception.detail["reason_code"],
            "knowledge_reindex_required",
        )


class RuntimeCapabilityTests(unittest.TestCase):
    @staticmethod
    def _azure_config() -> SimpleNamespace:
        return SimpleNamespace(
            azure_openai_embedding_deployment="",
            azure_search_endpoint="",
            azure_search_index_name="",
        )

    @staticmethod
    def _aliyun_config() -> SimpleNamespace:
        return SimpleNamespace(
            dashscope_api_key="configured-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            chat_model="qwen3.8-max",
            screening_model="qwen3.7-flash",
            embedding_model="text-embedding-v4",
            retrieval_backend="qdrant",
            qdrant_url="http://127.0.0.1:6333",
            qdrant_collection_name="reqpilot-knowledge-aliyun-local",
        )

    def test_missing_components_are_reported_independently(self) -> None:
        azure = SimpleNamespace(
            azure_openai_embedding_deployment="",
            azure_search_endpoint="",
            azure_search_index_name="",
        )
        aliyun = SimpleNamespace(
            dashscope_api_key="",
            dashscope_base_url="",
            chat_model="",
            screening_model="",
            embedding_model="",
            retrieval_backend="disabled",
            qdrant_url="",
            qdrant_collection_name="",
        )
        processing_profiles.invalidate_runtime_capabilities()
        with (
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=5,
            ),
            patch.object(
                processing_profiles,
                "get_runtime_ai_config_value",
                side_effect=lambda _name, default=None: default or "",
            ),
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                side_effect=lambda profile_id: (
                    azure if profile_id == "azure" else aliyun
                ),
            ),
        ):
            capabilities = {
                item.capability_id: item
                for item in processing_profiles.build_runtime_capabilities(probe=False)
            }

        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_AZURE_DI].state,
            CapabilityState.NOT_CONFIGURED,
        )
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_AZURE_SOL].state,
            CapabilityState.NOT_CONFIGURED,
        )
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_AZURE_TERRA].state,
            CapabilityState.NOT_CONFIGURED,
        )
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_ALIYUN_EMBEDDING].state,
            CapabilityState.NOT_CONFIGURED,
        )
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_QDRANT].state,
            CapabilityState.NOT_CONFIGURED,
        )

    def test_capability_gate_maps_configuration_to_409_and_auth_to_503(
        self,
    ) -> None:
        with patch.object(
            processing_profiles,
            "build_runtime_capabilities",
            return_value=(
                _capability(
                    processing_profiles.CAPABILITY_AZURE_SOL,
                    CapabilityState.MODEL_NOT_DEPLOYED,
                ),
            ),
        ):
            with self.assertRaises(ApplicationError) as missing:
                processing_profiles.require_capabilities(
                    (processing_profiles.CAPABILITY_AZURE_SOL,),
                    operation="test",
                )
        self.assertEqual(missing.exception.status_code, 409)

        with patch.object(
            processing_profiles,
            "build_runtime_capabilities",
            return_value=(
                _capability(
                    processing_profiles.CAPABILITY_AZURE_SOL,
                    CapabilityState.AUTH_FAILED,
                ),
            ),
        ):
            with self.assertRaises(ApplicationError) as auth:
                processing_profiles.require_capabilities(
                    (processing_profiles.CAPABILITY_AZURE_SOL,),
                    operation="test",
                )
        self.assertEqual(auth.exception.status_code, 503)

    def test_observed_provider_failure_survives_read_only_capability_probe(
        self,
    ) -> None:
        azure = SimpleNamespace(
            azure_openai_embedding_deployment="",
            azure_search_endpoint="",
            azure_search_index_name="",
        )
        aliyun = SimpleNamespace(
            dashscope_api_key="configured-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            chat_model="qwen3.8-max",
            screening_model="qwen3.7-flash",
            embedding_model="text-embedding-v4",
            retrieval_backend="disabled",
            qdrant_url="",
            qdrant_collection_name="",
        )
        processing_profiles.invalidate_runtime_capabilities()
        with (
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=9,
            ),
            patch.object(
                processing_profiles,
                "get_runtime_ai_config_value",
                side_effect=lambda _name, default=None: default or "",
            ),
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                side_effect=lambda profile_id: (
                    azure if profile_id == "azure" else aliyun
                ),
            ),
            patch.object(
                processing_profiles,
                "_probe_runtime_components",
                return_value={},
            ),
        ):
            processing_profiles.record_runtime_capability_failure(
                capability_id=processing_profiles.CAPABILITY_BAILIAN_LAYOUT,
                state=CapabilityState.AUTH_FAILED,
                reason_code="bailian_layout_auth_failed",
                message="Invalid key.",
            )
            capabilities = {
                item.capability_id: item
                for item in processing_profiles.build_runtime_capabilities(probe=True)
            }

        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_BAILIAN_LAYOUT].state,
            CapabilityState.AUTH_FAILED,
        )

    def test_configured_aliyun_capabilities_are_checking_until_live_probe(self) -> None:
        processing_profiles.invalidate_runtime_capabilities()
        azure = self._azure_config()
        aliyun = self._aliyun_config()
        with (
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=21,
            ),
            patch.object(
                processing_profiles,
                "get_runtime_ai_config_value",
                side_effect=lambda _name, default=None: default or "",
            ),
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                side_effect=lambda profile_id: azure
                if profile_id == "azure"
                else aliyun,
            ),
        ):
            capabilities = {
                item.capability_id: item
                for item in processing_profiles.build_runtime_capabilities(probe=False)
            }

        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_ALIYUN_LLM].state,
            CapabilityState.CHECKING,
        )
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_ALIYUN_EMBEDDING].state,
            CapabilityState.CHECKING,
        )
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_QDRANT].state,
            CapabilityState.CHECKING,
        )

    def test_explicit_aliyun_probe_binds_live_results_to_current_revision(self) -> None:
        processing_profiles.invalidate_runtime_capabilities()
        azure = self._azure_config()
        aliyun = self._aliyun_config()
        live_checks = {
            "screening": {
                "state": "READY",
                "ready": True,
                "reason_code": "",
                "message": "",
                "model": "qwen3.7-flash",
                "role": "screening",
            },
            "assessment": {
                "state": "READY",
                "ready": True,
                "reason_code": "",
                "message": "",
                "model": "qwen3.8-max",
                "role": "assessment",
            },
            "ocr": {
                "state": "READY",
                "ready": True,
                "reason_code": "",
                "message": "",
                "model": "qwen3.5-ocr",
                "role": "ocr",
                "request_count": 1,
            },
            "embedding": {
                "state": "READY",
                "ready": True,
                "reason_code": "",
                "message": "",
                "model": "text-embedding-v4",
                "role": "embedding",
                "dimensions": 1024,
            },
        }
        with (
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=22,
            ),
            patch.object(
                processing_profiles,
                "get_runtime_ai_config_value",
                side_effect=lambda _name, default=None: default or "",
            ),
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                side_effect=lambda profile_id: azure
                if profile_id == "azure"
                else aliyun,
            ),
            patch.object(
                processing_profiles,
                "probe_runtime_aliyun_models",
                return_value=live_checks,
            ) as model_probe,
            patch(
                "app.infrastructure.local_qdrant.ensure_local_qdrant_running",
                return_value=SimpleNamespace(ready=True, state="ready_owned"),
            ) as qdrant_start,
            patch.object(
                processing_profiles,
                "build_ai_profile_check",
                return_value={"qdrant_probe": {"status": "ok"}},
            ),
        ):
            payload = processing_profiles.probe_runtime_aliyun_connection()

        model_probe.assert_called_once_with(aliyun)
        qdrant_start.assert_called_once()
        self.assertEqual(
            payload["schema_version"], "runtime-aliyun-connection-probe/v2"
        )
        self.assertEqual(payload["connection_revision"], 22)
        self.assertTrue(payload["ready"])
        self.assertEqual(
            set(payload["checks"]),
            {"screening", "assessment", "ocr", "embedding", "qdrant"},
        )
        by_id = {item["capability_id"]: item for item in payload["capabilities"]}
        self.assertTrue(by_id[processing_profiles.CAPABILITY_ALIYUN_LLM]["ready"])
        self.assertTrue(by_id[processing_profiles.CAPABILITY_BAILIAN_OCR]["ready"])
        self.assertTrue(by_id[processing_profiles.CAPABILITY_ALIYUN_EMBEDDING]["ready"])
        self.assertTrue(by_id[processing_profiles.CAPABILITY_QDRANT]["ready"])
        serialized = json.dumps(payload)
        self.assertNotIn("configured-key", serialized)
        self.assertNotIn("dashscope.aliyuncs.com", serialized)

    def test_latest_aliyun_probe_receipt_is_readable_only_for_current_revision(
        self,
    ) -> None:
        processing_profiles.invalidate_runtime_capabilities()
        with patch.object(
            processing_profiles,
            "get_runtime_connection_revision",
            return_value=22,
        ):
            self.assertIsNone(
                processing_profiles.get_runtime_aliyun_connection_probe_receipt()
            )

        receipt = {
            "schema_version": "runtime-aliyun-connection-probe/v2",
            "provider": "aliyun",
            "connection_revision": 22,
            "checked_at": "2026-08-18T00:00:00+00:00",
            "ready": True,
            "checks": {"assessment": {"state": "READY", "ready": True}},
        }
        with processing_profiles._CAPABILITY_CACHE_LOCK:
            processing_profiles._CAPABILITY_CACHE["aliyun_probe_receipt"] = receipt

        with patch.object(
            processing_profiles,
            "get_runtime_connection_revision",
            return_value=22,
        ):
            restored = processing_profiles.get_runtime_aliyun_connection_probe_receipt()
        self.assertEqual(restored, receipt)
        self.assertIsNot(restored, receipt)

        with patch.object(
            processing_profiles,
            "get_runtime_connection_revision",
            return_value=23,
        ):
            self.assertIsNone(
                processing_profiles.get_runtime_aliyun_connection_probe_receipt()
            )

    def test_explicit_aliyun_probe_sanitizes_lifecycle_and_provider_failures(
        self,
    ) -> None:
        processing_profiles.invalidate_runtime_capabilities()
        azure = self._azure_config()
        aliyun = self._aliyun_config()
        provider_failure = {
            "state": "AUTH_FAILED",
            "ready": False,
            "reason_code": "aliyun_auth_failed",
            "message": "DashScope rejected the configured API key.",
            "model": "qwen3.8-max",
            "role": "assessment",
        }
        live_checks = {
            "screening": {
                **provider_failure,
                "model": "qwen3.7-flash",
                "role": "screening",
            },
            "assessment": provider_failure,
            "ocr": {
                **provider_failure,
                "model": "qwen3.5-ocr",
                "role": "ocr",
            },
            "embedding": {
                **provider_failure,
                "model": "text-embedding-v4",
                "role": "embedding",
            },
        }
        with (
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=23,
            ),
            patch.object(
                processing_profiles,
                "get_runtime_ai_config_value",
                side_effect=lambda _name, default=None: default or "",
            ),
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                side_effect=lambda profile_id: azure
                if profile_id == "azure"
                else aliyun,
            ),
            patch.object(
                processing_profiles,
                "probe_runtime_aliyun_models",
                return_value=live_checks,
            ),
            patch(
                "app.infrastructure.local_qdrant.ensure_local_qdrant_running",
                side_effect=OSError(
                    "secret-lifecycle-detail configured-key dashscope.aliyuncs.com"
                ),
            ),
        ):
            payload = processing_profiles.probe_runtime_aliyun_connection()

        self.assertFalse(payload["ready"])
        self.assertEqual(payload["checks"]["qdrant"]["state"], "DEPENDENCY_UNAVAILABLE")
        self.assertEqual(
            payload["checks"]["qdrant"]["reason_code"], "qdrant_start_failed"
        )
        serialized = json.dumps(payload)
        self.assertNotIn("secret-lifecycle-detail", serialized)
        self.assertNotIn("configured-key", serialized)
        self.assertNotIn("dashscope.aliyuncs.com", serialized)

    def test_bailian_ocr_capability_uses_only_the_real_ocr_probe(self) -> None:
        processing_profiles.invalidate_runtime_capabilities()
        aliyun = self._aliyun_config()
        azure = self._azure_config()
        with (
            patch.object(
                processing_profiles,
                "get_runtime_connection_revision",
                return_value=24,
            ),
            patch.object(
                processing_profiles,
                "load_ai_copilot_config",
                side_effect=lambda profile_id: azure
                if profile_id == "azure"
                else aliyun,
            ),
        ):
            with processing_profiles._CAPABILITY_CACHE_LOCK:
                processing_profiles._CAPABILITY_CACHE.update(
                    {
                        "connection_revision": 24,
                        "checked_at_monotonic": processing_profiles.time.monotonic(),
                        "probe": {
                            "aliyun_screening": processing_profiles._ProbeResult(
                                processing_profiles.CapabilityState.READY
                            ),
                            "aliyun_assessment": processing_profiles._ProbeResult(
                                processing_profiles.CapabilityState.READY
                            ),
                            "aliyun_ocr": processing_profiles._ProbeResult(
                                processing_profiles.CapabilityState.MODEL_NOT_DEPLOYED,
                                "aliyun_model_not_deployed",
                                "The configured OCR model is unavailable.",
                            ),
                        },
                    }
                )
            capabilities = {
                item.capability_id: item
                for item in processing_profiles.build_runtime_capabilities(probe=False)
            }

        self.assertTrue(capabilities[processing_profiles.CAPABILITY_ALIYUN_LLM].ready)
        self.assertFalse(capabilities[processing_profiles.CAPABILITY_BAILIAN_OCR].ready)
        self.assertEqual(
            capabilities[processing_profiles.CAPABILITY_BAILIAN_OCR].state,
            processing_profiles.CapabilityState.MODEL_NOT_DEPLOYED,
        )

    def test_capability_gate_fails_closed_for_unknown_identifier(self) -> None:
        with patch.object(
            processing_profiles,
            "build_runtime_capabilities",
            return_value=(),
        ):
            with self.assertRaises(ApplicationError) as caught:
                processing_profiles.require_capabilities(
                    ("runtime.capability.does_not_exist",),
                    operation="test",
                )

        self.assertEqual(caught.exception.status_code, 409)
        self.assertEqual(
            caught.exception.detail["reason_code"],
            "unknown_runtime_capability",
        )


class LocalRulesPipelineTests(unittest.TestCase):
    def test_office_results_require_current_adapter_and_runtime_binding(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "requirements.docx"
            document = Document()
            document.add_heading("1 Scope", level=1)
            document.add_paragraph("The supplier shall provide a pump.")
            document.save(source)
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            profile = build_local_rules_profile(project_id="project-office")
            plan = build_resolved_processing_plan(
                project_id="project-office",
                document_id="document-office",
                source_sha256=source_sha256,
                profile=profile,
                document=profile.document,
                llm_provider="none",
                provider_bindings=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf"
                ),
                connection_revision=0,
            )
            result_runtime_sha256 = (
                document_jobs._document_extraction_result_runtime_sha256(plan)
            )
            result = build_local_structure_document_result(
                source,
                plan=plan,
                result_runtime_sha256=result_runtime_sha256,
            )

        self.assertEqual(
            result["pipeline"]["source_ir"]["adapter_id"],
            "native-docx-ooxml",
        )
        self.assertTrue(
            document_jobs._local_structure_extraction_result_matches_runtime(result)
        )
        missing_runtime = json.loads(json.dumps(result))
        missing_runtime.pop("result_runtime_sha256")
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                missing_runtime
            )
        )
        stale_adapter = json.loads(json.dumps(result))
        stale_adapter["pipeline"]["source_ir"]["version"] = "1.1.0"
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                stale_adapter
            )
        )

    def test_local_structure_does_not_invoke_requirement_semantics(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "1 Scope")
            page.insert_text((72, 104), "The supplier shall provide a pump.")
            document.save(source)
            document.close()
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            profile = build_local_rules_profile(project_id="project-local")
            plan = build_resolved_processing_plan(
                project_id="project-local",
                document_id="document-local-rejected-projection",
                source_sha256=source_sha256,
                profile=profile,
                document=profile.document,
                llm_provider="none",
                provider_bindings=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf"
                ),
                connection_revision=0,
            )
            with (
                patch(
                    "app.application.extraction.classify_content.classify_content",
                    side_effect=AssertionError("local path classified semantics"),
                ) as semantic_classifier,
                patch(
                    "app.application.extraction.project_requirements."
                    "project_requirements",
                    side_effect=AssertionError("local path projected requirements"),
                ) as requirement_projector,
                patch(
                    "app.application.extraction.processing_pipeline."
                    "build_table_record_projections",
                    wraps=processing_pipeline.build_table_record_projections,
                ) as table_projection_builder,
            ):
                result = build_local_structure_document_result(
                    source,
                    plan=plan,
                    result_runtime_sha256=(
                        document_jobs._document_extraction_result_runtime_sha256(plan)
                    ),
                )

            table_projection_builder.assert_called_once()
            semantic_classifier.assert_not_called()
            requirement_projector.assert_not_called()
            self.assertEqual(result["status"], "STRUCTURE_READY")
            self.assertGreater(result["counts"]["source_nodes"], 0)
            self.assertGreater(result["counts"]["structure_nodes"], 0)
            self.assertEqual(
                result["counts"]["coverage_status"],
                "COMPLETE_LOCAL_STRUCTURE",
            )
            self.assertEqual(result["items"], [])
            self.assertEqual(result["requirement_groups"], [])
            self.assertEqual(result["review_groups"], [])
            self.assertEqual(result["pipeline"]["semantic"]["status"], "NOT_RUN")
            self.assertEqual(result["pipeline"]["semantic"]["authority"], "NONE")
            self.assertEqual(
                result["pipeline"]["requirement_view"]["status"],
                "NOT_RUN",
            )
            self.assertIsNone(result["pipeline"]["requirement_view"]["id"])
            self.assertEqual(
                result["pipeline"]["requirement_view"]["counts"],
                {
                    "nodes": 0,
                    "requirement_groups": 0,
                    "review_groups": 0,
                },
            )
            stages = {
                item["stage"]: item["status"]
                for item in result["structure_agent"]["stages"]
            }
            self.assertEqual(
                stages["REQUIREMENT_CLASSIFICATION"],
                "SKIPPED",
            )
            self.assertEqual(stages["VALIDATION"], "COMPLETE")
            self.assertTrue(
                document_jobs._local_structure_extraction_result_matches_runtime(result)
            )
            revised_result = json.loads(json.dumps(result))
            semantic_stage_before = json.loads(
                json.dumps(revised_result["pipeline"]["semantic"])
            )
            requirement_view_stage_before = json.loads(
                json.dumps(revised_result["pipeline"]["requirement_view"])
            )
            structure_workspace._mark_structure_derivations(
                revised_result,
                previous_effective_snapshot_id="effective-structure-before",
                current_effective_snapshot_id="effective-structure-after",
                stale_structure_node_ids=("structure-node-reparented",),
                reusable_structure_node_ids=(),
            )
            self.assertEqual(
                revised_result["pipeline"]["semantic"],
                semantic_stage_before,
            )
            self.assertEqual(
                revised_result["pipeline"]["requirement_view"],
                requirement_view_stage_before,
            )
            self.assertTrue(
                document_jobs._local_structure_extraction_result_matches_runtime(
                    revised_result
                )
            )
            tampered_rejection = json.loads(json.dumps(result))
            tampered_rejection["pipeline"]["semantic"]["status"] = "COMPLETE"
            self.assertFalse(
                document_jobs._local_structure_extraction_result_matches_runtime(
                    tampered_rejection
                )
            )

    def test_local_structure_exposes_semantic_neutral_content_roles(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "1 Scope")
            page.insert_text((72, 104), "The supplier shall provide a pump.")
            document.save(source)
            document.close()
            source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
            profile = build_local_rules_profile(project_id="project-local")
            plan = build_resolved_processing_plan(
                project_id="project-local",
                document_id="document-local",
                source_sha256=source_sha256,
                profile=profile,
                document=profile.document,
                llm_provider="none",
                provider_bindings=ResolvedProviderBindingsV1(
                    primary_parser_provider="local:pymupdf"
                ),
                connection_revision=0,
            )

            with patch(
                "app.application.extraction.provider_support."
                "_build_llm_shadow_semantic_provider",
                side_effect=AssertionError("local rules constructed an LLM"),
            ) as remote_builder:
                result_runtime_sha256 = (
                    document_jobs._document_extraction_result_runtime_sha256(plan)
                )
                first = build_local_structure_document_result(
                    source,
                    plan=plan,
                    result_runtime_sha256=result_runtime_sha256,
                )
                replay = build_local_structure_document_result(
                    source,
                    plan=plan,
                    result_runtime_sha256=result_runtime_sha256,
                )
                read_model = build_local_structure_document_result(
                    source,
                    plan=plan,
                    result_runtime_sha256=result_runtime_sha256,
                    include_source_read_model=True,
                )

        remote_builder.assert_not_called()
        self.assertEqual(first["mode"], LOCAL_STRUCTURE_PROJECTION_MODE)
        self.assertEqual(first["pipeline"]["semantic"]["status"], "NOT_RUN")
        self.assertEqual(first["pipeline"]["semantic"]["authority"], "NONE")
        self.assertEqual(first["pipeline"]["requirement_view"]["status"], "NOT_RUN")
        self.assertEqual(first["items"], [])
        self.assertEqual(first["requirement_groups"], [])
        self.assertEqual(first["review_groups"], [])
        self.assertEqual(
            first["counts"]["items"],
            len(first["items"]),
        )
        self.assertEqual(
            first["counts"]["requirement_groups"],
            len(first["requirement_groups"]),
        )
        self.assertEqual(
            first["counts"]["review_groups"],
            len(first["review_groups"]),
        )
        self.assertEqual(
            first["counts"]["table_record_projections"],
            len(first["table_record_projections"]),
        )
        self.assertEqual(
            first["counts"]["table_evidence_bindings"],
            len(first["table_evidence_bindings"]),
        )
        self.assertEqual(
            first["counts"]["table_records"],
            sum(len(item["records"]) for item in first["table_record_projections"]),
        )
        self.assertEqual(
            first["pipeline"]["semantic"]["structure_snapshot_id"],
            first["pipeline"]["structure"]["id"],
        )
        inspection = first["structure_inspection"]
        self.assertEqual(inspection["basis"], "BASE_STRUCTURE")
        self.assertEqual(inspection["basis_revision"], 0)
        self.assertEqual(inspection["effective_revision"], 0)
        self.assertFalse(inspection["stale"])
        self.assertEqual(inspection["status"], "COMPLETE")
        self.assertEqual(
            inspection["resolver_version"],
            first["pipeline"]["structure"]["version"],
        )
        self.assertEqual(
            inspection["inspection_id"],
            replay["structure_inspection"]["inspection_id"],
        )
        self.assertEqual(
            inspection["canonical_sha256"],
            replay["structure_inspection"]["canonical_sha256"],
        )
        self.assertEqual(
            first["structure_agent"]["deterministic_inspection"],
            inspection,
        )
        self.assertTrue(
            first["structure_agent"]["local_capabilities"][
                "deterministic_structure_inspection"
            ]
        )
        presentation = read_model["structure_presentation_index"]
        self.assertNotIn("structure_source_index", read_model)
        self.assertEqual(
            presentation["schema_version"],
            STRUCTURE_PRESENTATION_INDEX_SCHEMA_VERSION,
        )
        self.assertTrue(presentation["nodes"])
        serialized_presentation = json.dumps(presentation)
        self.assertNotIn('"text_lines"', serialized_presentation)
        self.assertNotIn('"bbox"', serialized_presentation)
        page_preview = read_model[SOURCE_PAGE_PREVIEW_KEY]
        self.assertEqual(page_preview["document"]["page_count"], 1)
        page_facts = page_preview["page_comparison"]["pages"][0]["blocks"]
        self.assertTrue(
            any(
                "supplier shall provide" in block["text"].lower()
                for block in page_facts
            )
        )
        self.assertTrue(
            any(block["locatable"] and block["bbox"] for block in page_facts)
        )
        self.assertEqual(
            first["pipeline"]["requirement_view"]["structure_snapshot_id"],
            first["pipeline"]["structure"]["id"],
        )
        self.assertEqual(
            first["generation"],
            {
                "mode": "DETERMINISTIC_STRUCTURE",
                "remote_provider_calls": 0,
                "ai_review_generated": False,
                "semantic_authority": "NONE",
                "atomization_status": "NOT_RUN",
                "refinement_status": "NOT_RUN",
            },
        )
        self.assertIsNone(first["llm_review"])
        self.assertIsNone(first["requirement_set_id"])
        self.assertTrue(
            document_jobs._local_structure_extraction_result_matches_runtime(first)
        )
        self.assertEqual(first["result_runtime_sha256"], result_runtime_sha256)
        missing_result_runtime = json.loads(json.dumps(first))
        missing_result_runtime.pop("result_runtime_sha256")
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                missing_result_runtime
            )
        )
        stale_result_runtime = json.loads(json.dumps(first))
        stale_result_runtime["result_runtime_sha256"] = "0" * 64
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                stale_result_runtime
            )
        )
        self.assertEqual(
            first["requirement_tree_projection_runtime_version"],
            document_jobs.REQUIREMENT_TREE_PROJECTION_RUNTIME_VERSION,
        )
        missing_tree_runtime = json.loads(json.dumps(first))
        missing_tree_runtime.pop("requirement_tree_projection_runtime_version")
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                missing_tree_runtime
            )
        )
        stale_tree_runtime = json.loads(json.dumps(first))
        stale_tree_runtime["requirement_tree_projection_runtime_version"] = (
            "requirement-tree-projection/v10"
        )
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                stale_tree_runtime
            )
        )
        missing_inspection = json.loads(json.dumps(first))
        missing_inspection.pop("structure_inspection")
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                missing_inspection
            )
        )
        mismatched_inspection = json.loads(json.dumps(first))
        mismatched_inspection["structure_agent"]["deterministic_inspection"][
            "canonical_sha256"
        ] = "f" * 64
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                mismatched_inspection
            )
        )
        self.assertEqual(
            first["pipeline"]["source_ir"]["canonical_sha256"],
            replay["pipeline"]["source_ir"]["canonical_sha256"],
        )
        self.assertEqual(
            first["pipeline"]["structure"]["canonical_sha256"],
            replay["pipeline"]["structure"]["canonical_sha256"],
        )
        self.assertNotIn("canonical_sha256", first["pipeline"]["semantic"])
        self.assertNotIn(
            "canonical_sha256",
            first["pipeline"]["requirement_view"],
        )
        self.assertEqual(first["items"], replay["items"])
        self.assertEqual(
            first["requirement_groups"],
            replay["requirement_groups"],
        )
        self.assertEqual(first["review_groups"], replay["review_groups"])
        self.assertEqual(
            first["table_record_projections"],
            replay["table_record_projections"],
        )
        self.assertEqual(
            first["table_evidence_bindings"],
            replay["table_evidence_bindings"],
        )

        stale_structure_only = json.loads(json.dumps(first))
        stale_structure_only["mode"] = "LOCAL_RULES_STRUCTURE"
        stale_structure_only["pipeline"]["semantic"] = {
            "status": "NOT_RUN",
            "reason_code": "llm_provider_none",
        }
        stale_structure_only["pipeline"].pop("requirement_view", None)
        stale_structure_only["items"] = []
        stale_structure_only["requirement_groups"] = []
        stale_structure_only["review_groups"] = []
        self.assertFalse(
            document_jobs._local_structure_extraction_result_matches_runtime(
                stale_structure_only
            )
        )
        visual_strategy = processing_plan_visual_strategy(plan)
        strategy_sha256 = hashlib.sha256(
            json.dumps(
                visual_strategy,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        result_runtime_sha256 = (
            document_jobs._document_extraction_result_runtime_sha256(plan)
        )
        with patch.object(
            document_jobs,
            "PDF_SOURCE_IR_ADAPTER_VERSION",
            "999.0.0",
        ):
            source_ir_adapter_changed_sha256 = (
                document_jobs._document_extraction_result_runtime_sha256(plan)
            )
        with patch.object(
            document_jobs,
            "STRUCTURE_INSPECTION_RUNTIME_VERSION",
            "999.0.0",
        ):
            inspection_changed_sha256 = (
                document_jobs._document_extraction_result_runtime_sha256(plan)
            )
        with patch.object(
            document_jobs,
            "STRUCTURE_BOUNDARY_RULESET_VERSION",
            "999.0.0",
        ):
            boundary_rules_changed_sha256 = (
                document_jobs._document_extraction_result_runtime_sha256(plan)
            )
        with patch.object(
            document_jobs,
            "REQUIREMENT_TREE_PROJECTION_RUNTIME_VERSION",
            "requirement-tree-projection/999",
        ):
            requirement_tree_changed_sha256 = (
                document_jobs._document_extraction_result_runtime_sha256(plan)
            )
        self.assertNotEqual(
            result_runtime_sha256,
            inspection_changed_sha256,
        )
        self.assertNotEqual(
            result_runtime_sha256,
            source_ir_adapter_changed_sha256,
        )
        self.assertNotEqual(
            result_runtime_sha256,
            boundary_rules_changed_sha256,
        )
        self.assertNotEqual(
            result_runtime_sha256,
            requirement_tree_changed_sha256,
        )
        current_idempotency_key = (
            document_jobs._document_extraction_durable_idempotency_key(
                source_sha256=source_sha256,
                processing_plan_sha256=plan.canonical_sha256,
                visual_strategy_sha256=strategy_sha256,
                structure_revision_chain_sha256="",
                result_runtime_sha256=result_runtime_sha256,
            )
        )
        legacy_structure_only_key = (
            f"{source_sha256[:20]}:{plan.canonical_sha256[:20]}:"
            f"{strategy_sha256[:20]}:base"
        )
        self.assertRegex(current_idempotency_key, r"^[0-9a-f]{64}$")
        self.assertNotEqual(
            current_idempotency_key,
            legacy_structure_only_key,
        )


if __name__ == "__main__":
    unittest.main()
