from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import copy
import hashlib
import json
import re
import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from PIL import Image

from scripts.document_visual_review_orchestrator import (
    DocumentVisualReviewOrchestrationRejected,
    execute_docx_review_bundle_batch,
    execute_office_pipeline_profile_batch,
    execute_tif_not_evaluated_review,
    export_document_visual_review_plan,
)
from scripts.docx_visual_render_profile import (
    DocxPdfRenderResult,
    produce_docx_visual_render_profile,
)


class DocumentVisualReviewOrchestratorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(
            ignore_cleanup_errors=True
        )
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name).resolve()
        self.corpus_root = self.root / "corpus"
        self.corpus_root.mkdir()
        self.controlled_root = self.root / "controlled"
        self.controlled_root.mkdir()
        sources = self._sources()
        self.manifest = {
            "schema_version": "extraction-case-manifest/v1",
            "corpus": {
                "root_uri": "external://training-corpus/",
                "mapping_environment_variable": "TEST_CORPUS_ROOT",
                "contains_customer_content": True,
                "content_committed_to_repository": False,
            },
            "documents": {},
            "cases": {},
        }
        for index, (document_id, document_format, path) in enumerate(
            sources,
            start=1,
        ):
            self.manifest["documents"][document_id] = {
                "display_name": f"redacted-{index}",
                "document_ref": (
                    "external://training-corpus/" + path.name
                ),
                "format": document_format,
                "resolution_status": "resolved",
                "source_sha256": _sha256(path),
                "required": True,
                "duplicate_of": None,
            }
        duplicate_path = sources[0][2]
        self.manifest["documents"]["DOC-PDF-DUP"] = {
            "display_name": "redacted-duplicate",
            "document_ref": (
                "external://training-corpus/" + duplicate_path.name
            ),
            "format": "pdf",
            "resolution_status": "resolved",
            "source_sha256": _sha256(duplicate_path),
            "required": True,
            "duplicate_of": "DOC-PDF",
        }
        self.manifest_path = self.root / "manifest.json"
        self._write_manifest()

    def assert_contract_error(self, code: str):
        return self.assertRaisesRegex(
            DocumentVisualReviewOrchestrationRejected,
            rf"^{re.escape(code)}$",
        )

    def test_plan_verifies_refs_hashes_and_deduplicates_exact_source(self) -> None:
        destination = self.controlled_root / "plan"
        plan = export_document_visual_review_plan(
            self.manifest_path,
            expected_manifest_sha256=_sha256(self.manifest_path),
            corpus_root=self.corpus_root,
            output_dir=destination,
        )

        self.assertEqual(plan["source_slot_count"], 6)
        self.assertEqual(plan["unique_source_count"], 5)
        self.assertEqual(plan["duplicate_slot_count"], 1)
        self.assertEqual(
            plan["format_counts"],
            {"docx": 1, "pdf": 1, "tif": 1, "xlsb": 1, "xlsx": 1},
        )
        pdf = next(
            item for item in plan["documents"] if item["format"] == "pdf"
        )
        self.assertEqual(pdf["slot_count"], 2)
        self.assertEqual(pdf["slot_document_ids"], ["DOC-PDF", "DOC-PDF-DUP"])
        serialized = (destination / "orchestration-plan.json").read_text(
            encoding="utf-8"
        )
        self.assertNotIn("external://", serialized)
        for path in self.corpus_root.iterdir():
            self.assertNotIn(path.name, serialized)

    def test_real_tif_closure_is_blocked_and_all_pages_not_evaluated(
        self,
    ) -> None:
        result = execute_tif_not_evaluated_review(
            self.manifest_path,
            expected_manifest_sha256=_sha256(self.manifest_path),
            corpus_root=self.corpus_root,
            controlled_root=self.controlled_root,
            reviewed_at="2026-07-26T12:00:00+08:00",
        )

        self.assertEqual(result["pipeline_status"], "BLOCKED")
        self.assertEqual(result["ai_pre_review_verdict"], "NOT_EVALUATED")
        destination = self.controlled_root / (
            "tif-review-" + result["source_sha256"][:16]
        )
        ledger = json.loads(
            (destination / "ai-pre-review.json").read_text(encoding="utf-8")
        )
        self.assertEqual(
            {page["verdict"] for page in ledger["pages"]},
            {"NOT_EVALUATED"},
        )
        self.assertEqual(
            {tuple(page["issue_codes"]) for page in ledger["pages"]},
            {("PIPELINE_NOT_RUN",)},
        )
        self.assertTrue((destination / "bundle" / "manifest.json").is_file())

    def test_docx_only_batch_does_not_run_other_office_formats(self) -> None:
        index = execute_office_pipeline_profile_batch(
            self.manifest_path,
            expected_manifest_sha256=_sha256(self.manifest_path),
            corpus_root=self.corpus_root,
            controlled_root=self.controlled_root,
            document_formats=("docx",),
        )

        self.assertEqual(index["document_count"], 1)
        self.assertEqual(index["formats"], ["docx"])
        self.assertEqual(index["documents"][0]["format"], "docx")
        self.assertEqual(
            index["documents"][0]["pipeline_status"],
            "COMPLETE",
        )

    def test_docx_bundle_batch_joins_profiles_and_reverifies_publish(
        self,
    ) -> None:
        render_root = self.controlled_root / "docx-render-profiles"
        render_root.mkdir()
        source = self.corpus_root / "source.docx"
        rendered = produce_docx_visual_render_profile(
            source,
            expected_source_sha256=_sha256(source),
            output_dir=render_root / "DOC-DOCX",
            renderer=_FakeDocxRenderer(),
        )
        pipeline_index = execute_office_pipeline_profile_batch(
            self.manifest_path,
            expected_manifest_sha256=_sha256(self.manifest_path),
            corpus_root=self.corpus_root,
            controlled_root=self.controlled_root,
            document_formats=("docx",),
            output_name="docx-pipeline-profiles",
        )

        index = execute_docx_review_bundle_batch(
            self.manifest_path,
            expected_manifest_sha256=_sha256(self.manifest_path),
            corpus_root=self.corpus_root,
            controlled_root=self.controlled_root,
            render_profiles_root=render_root,
            pipeline_profiles_root=(
                self.controlled_root / "docx-pipeline-profiles"
            ),
            expected_pipeline_batch_sha256=pipeline_index["batch_sha256"],
        )

        self.assertEqual(index["document_count"], 1)
        self.assertEqual(index["page_count"], 2)
        self.assertEqual(index["pipeline_status_counts"], {"COMPLETE": 1})
        row = index["documents"][0]
        self.assertEqual(row["document_id"], "DOC-DOCX")
        self.assertEqual(row["page_count"], 2)
        self.assertEqual(row["pipeline_status"], "COMPLETE")
        self.assertEqual(row["source_ir_diagnostic_code_counts"], {})
        self.assertEqual(
            row["render_profile_file_sha256"],
            rendered.profile_file_sha256,
        )
        destination = self.controlled_root / "docx-review-bundles-v1"
        manifest = json.loads(
            (
                destination
                / row["review_id"]
                / "manifest.json"
            ).read_text(encoding="utf-8")
        )
        self.assertEqual(manifest["bundle_sha256"], row["bundle_sha256"])
        self.assertFalse(any(destination.rglob("*ai-pre-review*")))
        serialized = (destination / "batch-index.json").read_text(
            encoding="utf-8"
        )
        self.assertNotIn("external://", serialized)
        self.assertNotIn(str(source), serialized)
        self.assertNotIn(source.name, serialized)

    def test_hash_traversal_and_duplicate_binding_fail_before_publish(
        self,
    ) -> None:
        wrong_hash = copy.deepcopy(self.manifest)
        wrong_hash["documents"]["DOC-TIF"]["source_sha256"] = "f" * 64
        self.manifest = wrong_hash
        self._write_manifest()
        with self.assert_contract_error(
            "orchestrator_document_source_hash_mismatch"
        ):
            export_document_visual_review_plan(
                self.manifest_path,
                expected_manifest_sha256=_sha256(self.manifest_path),
                corpus_root=self.corpus_root,
                output_dir=self.controlled_root / "wrong-hash",
            )

        traversal = copy.deepcopy(self.manifest)
        traversal["documents"]["DOC-TIF"]["source_sha256"] = _sha256(
            self.corpus_root / "drawing.tif"
        )
        traversal["documents"]["DOC-TIF"]["document_ref"] = (
            "external://training-corpus/../drawing.tif"
        )
        self.manifest = traversal
        self._write_manifest()
        with self.assert_contract_error("orchestrator_document_ref_invalid"):
            export_document_visual_review_plan(
                self.manifest_path,
                expected_manifest_sha256=_sha256(self.manifest_path),
                corpus_root=self.corpus_root,
                output_dir=self.controlled_root / "traversal",
            )

    def _sources(self):
        pdf = self.corpus_root / "source.pdf"
        pdf.write_bytes(b"%PDF-1.4\n%%EOF")
        docx = self.corpus_root / "source.docx"
        document = Document()
        document.add_heading("1 Scope", level=1)
        document.add_paragraph("The system shall preserve traceability.")
        document.save(docx)
        xlsx = self.corpus_root / "source.xlsx"
        workbook = Workbook()
        workbook.active.append(["ID", "Requirement"])
        workbook.active.append(["R1", "The system shall retain order."])
        workbook.save(xlsx)
        workbook.close()
        xlsb = self.corpus_root / "source.xlsb"
        xlsb.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 128)
        tif = self.corpus_root / "drawing.tif"
        first = Image.new("1", (21, 31), color=1)
        second = Image.new("1", (22, 32), color=0)
        first.save(tif, format="TIFF", save_all=True, append_images=[second])
        return (
            ("DOC-PDF", "pdf", pdf),
            ("DOC-DOCX", "docx", docx),
            ("DOC-XLSX", "xlsx", xlsx),
            ("DOC-XLSB", "xlsb", xlsb),
            ("DOC-TIF", "tif", tif),
        )

    def _write_manifest(self) -> None:
        self.manifest_path.write_text(
            json.dumps(
                self.manifest,
                ensure_ascii=False,
                sort_keys=True,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class _FakeDocxRenderer:
    def render(self, source_bytes: bytes) -> DocxPdfRenderResult:
        del source_bytes
        document = pymupdf.open()
        try:
            for page_number in (1, 2):
                page = document.new_page(width=300, height=420)
                page.insert_text((40, 50), f"review page {page_number}")
            payload = document.tobytes()
        finally:
            document.close()
        return DocxPdfRenderResult(
            pdf_bytes=payload,
            pdf_sha256=hashlib.sha256(payload).hexdigest(),
            renderer_id="test-docx-pdf",
            renderer_version="9.8.7.6",
            renderer_config_sha256=hashlib.sha256(
                b"test renderer config"
            ).hexdigest(),
        )
