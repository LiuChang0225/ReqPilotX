from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import base64
import shutil
import tempfile
import unittest
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.styles import PatternFill

from app.adapters.documents.preflight import preflight_document
from app.adapters.documents import xlsx_source_ir_adapter as xlsx_adapter_module
from app.adapters.documents.xlsx_source_ir_adapter import XlsxSourceIRAdapter
from app.domain.extraction.preflight import DocumentFormat
from app.domain.extraction.source_ir import SourceIRStatus, SourceNodeKind
from app.ports.source_ir_adapter import (
    SourceIRAdapterPort,
    SourceIRAdapterRejected,
    SourceIRAdapterRejectionCode,
)


CREATED_AT = datetime(2026, 7, 20, 14, 0, tzinfo=timezone.utc)
_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class XlsxSourceIRAdapterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)
        self.adapter = XlsxSourceIRAdapter()

    def test_regions_merges_formula_cache_sheet_order_image_and_comments_are_faithful(
        self,
    ) -> None:
        path = self._write_rich_workbook("requirements.xlsx")
        before = _relative_inventory(self.root)
        preflight = preflight_document(path)

        first = self.adapter.extract(path, preflight=preflight, created_at=CREATED_AT)
        second = self.adapter.extract(path, preflight=preflight, created_at=CREATED_AT)

        self.assertEqual(first.status, SourceIRStatus.PARTIAL)
        self.assertEqual(first, second)
        self.assertEqual(
            [node.kind for node in first.nodes],
            [
                SourceNodeKind.TABLE,
                SourceNodeKind.IMAGE,
                SourceNodeKind.TABLE,
                SourceNodeKind.TABLE,
            ],
        )
        self.assertEqual([node.source_order for node in first.nodes], list(range(4)))
        self.assertFalse(
            any(
                node.kind in {SourceNodeKind.TABLE_ROW, SourceNodeKind.CELL}
                for node in first.nodes
            )
        )

        first_table_node = first.nodes[0]
        self.assertEqual(first_table_node.native_locator.workbook, "xl/workbook.xml")
        self.assertEqual(first_table_node.native_locator.sheet, "Visible")
        self.assertEqual(first_table_node.native_locator.cell_range, "A1:D2")
        table = first_table_node.table_payload
        assert table is not None
        self.assertEqual((table.row_count, table.column_count), (2, 4))
        self.assertEqual(len(table.cells), 7)
        cells = {(cell.row, cell.column): cell for cell in table.cells}
        self.assertEqual(cells[(0, 0)].raw_text, "共享标准")
        self.assertEqual(cells[(0, 0)].row_span, 2)
        self.assertNotIn((1, 0), cells)
        self.assertEqual(
            {
                (region.row, region.column, region.row_span, region.column_span)
                for region in table.merged_regions
            },
            {(0, 0, 2, 1)},
        )
        formula_cell = cells[(1, 2)]
        self.assertEqual(formula_cell.raw_text, "=C1*2")
        self.assertEqual(formula_cell.formula, "=C1*2")
        self.assertEqual(formula_cell.display_value, "20")
        self.assertEqual(formula_cell.native_locator.cell_range, "C2")
        self.assertIn((0, 3), cells)
        self.assertEqual(cells[(0, 3)].raw_text, "")
        self.assertIn((1, 3), cells)
        self.assertIn("sheet_state=visible", table.notes)
        self.assertIn("used_region=A1:D2", table.notes)

        image_node = first.nodes[1]
        self.assertEqual(image_node.native_locator.sheet, "Visible")
        self.assertEqual(image_node.native_locator.cell_range, "B4")
        self.assertEqual(len(first.assets), 1)
        self.assertEqual(image_node.asset_ref, first.assets[0].asset_id)
        self.assertEqual(first.assets[0].media_type, "image/png")
        self.assertTrue(first.assets[0].storage_key.startswith("source-assets/"))
        self.assertIn(first.assets[0].content_sha256, first.assets[0].storage_key)

        self.assertEqual(first.nodes[2].native_locator.sheet, "Visible")
        self.assertEqual(first.nodes[2].native_locator.cell_range, "F5:G5")
        self.assertEqual(first.nodes[3].native_locator.sheet, "Hidden Knowledge")
        self.assertEqual(first.nodes[3].native_locator.cell_range, "A1")
        hidden_table = first.nodes[3].table_payload
        assert hidden_table is not None
        self.assertIn("sheet_state=hidden", hidden_table.notes)

        diagnostic_codes = {item.code for item in first.diagnostics}
        self.assertIn("xlsx_comments_ignored_by_policy", diagnostic_codes)
        self.assertIn("xlsx_sheet_visibility_preserved", diagnostic_codes)
        self.assertIn("xlsx_row_locator_derived_from_table", diagnostic_codes)
        self.assertIn("xlsx_visual_format_not_serialized", diagnostic_codes)
        self.assertNotIn("xlsx_formula_cached_value_unavailable", diagnostic_codes)
        self.assertNotIn("工程师批注", first.model_dump_json())
        self.assertEqual(_relative_inventory(self.root), before)

    def test_formula_without_cached_value_is_preserved_and_explicitly_partial(self) -> None:
        path = self.root / "uncached-formula.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = 4
        worksheet["A2"] = "=A1+1"
        workbook.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.PARTIAL)
        table = snapshot.nodes[0].table_payload
        assert table is not None
        formula_cell = next(cell for cell in table.cells if cell.formula is not None)
        self.assertEqual(formula_cell.formula, "=A1+1")
        self.assertEqual(formula_cell.raw_text, "=A1+1")
        self.assertIsNone(formula_cell.display_value)
        self.assertIn(
            "xlsx_formula_cached_value_unavailable",
            {item.code for item in snapshot.diagnostics},
        )

    def test_chart_is_not_silently_treated_as_text_or_checkbox_state(self) -> None:
        path = self.root / "chart.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Chart"
        worksheet.append(("项目", "值"))
        worksheet.append(("A", 1))
        worksheet.append(("B", 2))
        chart = BarChart()
        chart.add_data(Reference(worksheet, min_col=2, min_row=1, max_row=3))
        worksheet.add_chart(chart, "D2")
        workbook.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.PARTIAL)
        self.assertIn(
            "xlsx_visual_object_not_mapped",
            {item.code for item in snapshot.diagnostics},
        )
        self.assertFalse(any(node.kind is SourceNodeKind.IMAGE for node in snapshot.nodes))

    def test_repeated_image_placements_at_one_anchor_remain_distinct(self) -> None:
        path, shared_image_part = self._write_repeated_image_workbook(
            "repeated-images.xlsx"
        )
        preflight = preflight_document(path)
        media_reads: list[str] = []
        with zipfile.ZipFile(path) as archive:
            original_read = archive.read

            def recording_read(name: object, pwd: bytes | None = None) -> bytes:
                member_name = getattr(name, "filename", str(name))
                if member_name.startswith("xl/media/"):
                    media_reads.append(member_name)
                return original_read(name, pwd)  # type: ignore[arg-type]

            with patch.object(archive, "read", side_effect=recording_read):
                records_by_sheet, _ = xlsx_adapter_module._scan_visual_assets(
                    archive,
                    source_sha256=preflight.identity.source_sha256,
                )

        self.assertEqual(len(records_by_sheet["Sheet"]), 2)
        self.assertEqual(media_reads, [shared_image_part])

        snapshot = self.adapter.extract(
            path,
            preflight=preflight,
            created_at=CREATED_AT,
        )

        image_nodes = [node for node in snapshot.nodes if node.kind is SourceNodeKind.IMAGE]
        self.assertEqual(len(image_nodes), 2)
        self.assertEqual(len(snapshot.assets), 1)
        self.assertEqual(
            [node.native_locator.object_index for node in image_nodes],
            [0, 1],
        )
        self.assertEqual(image_nodes[0].asset_ref, image_nodes[1].asset_ref)
        self.assertNotEqual(image_nodes[0].source_node_id, image_nodes[1].source_node_id)
        self.assertNotIn(
            "xlsx_multiple_images_share_locator",
            {item.code for item in snapshot.diagnostics},
        )

    def test_image_placement_and_unique_media_budgets_fail_closed(self) -> None:
        path, shared_image_part = self._write_repeated_image_workbook(
            "image-budgets.xlsx"
        )
        preflight = preflight_document(path)

        with patch.object(xlsx_adapter_module, "_MAX_XLSX_IMAGE_PLACEMENTS", 1):
            placement_snapshot = self.adapter.extract(
                path,
                preflight=preflight,
                created_at=CREATED_AT,
            )

        self.assertEqual(placement_snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(
            {item.code for item in placement_snapshot.diagnostics},
            {"xlsx_resource_budget_exceeded"},
        )

        with zipfile.ZipFile(path) as archive:
            shared_image_size = archive.getinfo(shared_image_part).file_size
        with patch.object(
            xlsx_adapter_module,
            "_MAX_XLSX_UNIQUE_MEDIA_BYTES",
            shared_image_size - 1,
        ):
            media_snapshot = self.adapter.extract(
                path,
                preflight=preflight,
                created_at=CREATED_AT,
            )

        self.assertEqual(media_snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(
            {item.code for item in media_snapshot.diagnostics},
            {"xlsx_resource_budget_exceeded"},
        )

    def test_cell_control_feature_is_flagged_without_guessing_boolean_state(self) -> None:
        path = self.root / "cell-control.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = True
        workbook.save(path)
        _inject_package_part(
            path,
            "xl/featurePropertyBag/featurePropertyBag.xml",
            b"<?xml version='1.0' encoding='UTF-8'?><featurePropertyBags/>",
        )

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.PARTIAL)
        table = snapshot.nodes[0].table_payload
        assert table is not None
        self.assertEqual(table.cells[0].raw_text, "TRUE")
        self.assertIn(
            "xlsx_cell_control_state_not_interpreted",
            {item.code for item in snapshot.diagnostics},
        )

    def test_invalid_preflight_source_change_and_naive_timestamp_fail_closed(self) -> None:
        source = self._write_simple_workbook("source.xlsx")
        disguised = self.root / "disguised.pdf"
        shutil.copyfile(source, disguised)
        with self.assertRaises(SourceIRAdapterRejected) as invalid_context:
            self.adapter.extract(
                disguised,
                preflight=preflight_document(disguised),
                created_at=CREATED_AT,
            )
        self.assertEqual(
            invalid_context.exception.code,
            SourceIRAdapterRejectionCode.INVALID_PREFLIGHT,
        )

        preflight = preflight_document(source)
        with source.open("ab") as handle:
            handle.write(b"changed-after-preflight")
        with self.assertRaises(SourceIRAdapterRejected) as changed_context:
            self.adapter.extract(source, preflight=preflight, created_at=CREATED_AT)
        self.assertEqual(
            changed_context.exception.code,
            SourceIRAdapterRejectionCode.SOURCE_CHANGED,
        )

        timestamp_source = self._write_simple_workbook("timestamp.xlsx")
        with self.assertRaises(SourceIRAdapterRejected) as timestamp_context:
            self.adapter.extract(
                timestamp_source,
                preflight=preflight_document(timestamp_source),
                created_at=datetime(2026, 7, 20, 14, 0),
            )
        self.assertEqual(
            timestamp_context.exception.code,
            SourceIRAdapterRejectionCode.INVALID_TIMESTAMP,
        )

    def test_native_mapping_failure_returns_failed_snapshot_without_detail_leak(self) -> None:
        path = self._write_simple_workbook("failure.xlsx")
        with patch.object(
            self.adapter,
            "_extract_bytes",
            side_effect=RuntimeError("sensitive workbook detail"),
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight_document(path),
                created_at=CREATED_AT,
            )

        self.assertEqual(snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(snapshot.nodes, ())
        self.assertEqual(snapshot.assets, ())
        self.assertEqual(snapshot.diagnostics[0].code, "xlsx_parse_failed")
        self.assertNotIn(str(path), snapshot.diagnostics[0].message)
        self.assertNotIn("sensitive workbook detail", snapshot.diagnostics[0].message)

    def test_cell_and_merged_area_budgets_fail_before_workbook_materialization(self) -> None:
        cell_path = self.root / "cell-budget.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "one"
        worksheet["A2"] = "two"
        workbook.save(cell_path)
        workbook.close()

        with patch.object(xlsx_adapter_module, "_MAX_XLSX_MATERIAL_CELLS", 1):
            cell_snapshot = self.adapter.extract(
                cell_path,
                preflight=preflight_document(cell_path),
                created_at=CREATED_AT,
            )

        self.assertEqual(cell_snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(
            {item.code for item in cell_snapshot.diagnostics},
            {"xlsx_resource_budget_exceeded"},
        )

        merge_path = self.root / "merge-budget.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.merge_cells("A1:B1")
        worksheet["A1"] = "shared"
        workbook.save(merge_path)
        workbook.close()

        with patch.object(xlsx_adapter_module, "_MAX_XLSX_MERGED_AREA", 1):
            merge_snapshot = self.adapter.extract(
                merge_path,
                preflight=preflight_document(merge_path),
                created_at=CREATED_AT,
            )

        self.assertEqual(merge_snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(
            {item.code for item in merge_snapshot.diagnostics},
            {"xlsx_resource_budget_exceeded"},
        )

    def test_sparse_grid_and_region_count_cannot_bypass_materialization_budget(
        self,
    ) -> None:
        sparse_path = self.root / "sparse-grid-budget.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "one"
        worksheet["B2"] = "two"
        worksheet["C3"] = "three"
        workbook.save(sparse_path)
        workbook.close()

        with patch.object(xlsx_adapter_module, "_MAX_XLSX_TABLE_GRID_CELLS", 4):
            sparse_snapshot = self.adapter.extract(
                sparse_path,
                preflight=preflight_document(sparse_path),
                created_at=CREATED_AT,
            )

        self.assertEqual(sparse_snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(
            {item.code for item in sparse_snapshot.diagnostics},
            {"xlsx_resource_budget_exceeded"},
        )

        regions_path = self.root / "region-count-budget.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "one"
        worksheet["C3"] = "two"
        workbook.save(regions_path)
        workbook.close()

        with patch.object(xlsx_adapter_module, "_MAX_XLSX_SOURCE_REGIONS", 1):
            regions_snapshot = self.adapter.extract(
                regions_path,
                preflight=preflight_document(regions_path),
                created_at=CREATED_AT,
            )

        self.assertEqual(regions_snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(
            {item.code for item in regions_snapshot.diagnostics},
            {"xlsx_resource_budget_exceeded"},
        )

    def test_adapter_satisfies_source_ir_port(self) -> None:
        self.assertIsInstance(self.adapter, SourceIRAdapterPort)
        self.assertEqual(self.adapter.document_format, DocumentFormat.XLSX)

    def _write_simple_workbook(self, name: str) -> Path:
        path = self.root / name
        workbook = Workbook()
        workbook.active["A1"] = "原文"
        workbook.save(path)
        return path

    def _write_rich_workbook(self, name: str) -> Path:
        path = self.root / name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Visible"
        worksheet.merge_cells("A1:A2")
        worksheet["A1"] = "共享标准"
        worksheet["B1"] = "参数一"
        worksheet["C1"] = 10
        worksheet["D1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        worksheet["B2"] = "参数二"
        worksheet["C2"] = "=C1*2"
        worksheet["B1"].comment = Comment("工程师批注", "Engineer")
        worksheet["F5"] = "排期"
        worksheet["G5"] = "2天"
        image_stream = BytesIO(_ONE_PIXEL_PNG)
        worksheet.add_image(Image(image_stream), "B4")

        hidden = workbook.create_sheet("Hidden Knowledge")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "内部知识"
        try:
            workbook.save(path)
        finally:
            image_stream.close()
        _set_cached_formula_value(
            path,
            worksheet_part="xl/worksheets/sheet1.xml",
            cell="C2",
            value="20",
        )
        return path

    def _write_repeated_image_workbook(self, name: str) -> tuple[Path, str]:
        image_path = self.root / f"{name}.png"
        image_path.write_bytes(_ONE_PIXEL_PNG)
        path = self.root / name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "图片要求"
        worksheet.add_image(Image(str(image_path)), "B2")
        worksheet.add_image(Image(str(image_path)), "B2")
        try:
            workbook.save(path)
        finally:
            workbook.close()
            image_path.unlink()
        return path, _share_drawing_image_target(path)


def _set_cached_formula_value(
    path: Path,
    *,
    worksheet_part: str,
    cell: str,
    value: str,
) -> None:
    with zipfile.ZipFile(path) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for info, payload in members:
            if info.filename == worksheet_part:
                root = etree.fromstring(payload)
                cell_element = root.find(
                    f".//{{{_MAIN_NS}}}c[@r='{cell}']",
                )
                assert cell_element is not None
                cached = cell_element.find(f"{{{_MAIN_NS}}}v")
                if cached is None:
                    cached = etree.SubElement(cell_element, f"{{{_MAIN_NS}}}v")
                cached.text = value
                payload = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            archive.writestr(info, payload)
    path.write_bytes(output.getvalue())


def _inject_package_part(path: Path, member_name: str, payload: bytes) -> None:
    with zipfile.ZipFile(path) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for info, member_payload in members:
            archive.writestr(info, member_payload)
        archive.writestr(member_name, payload)
    path.write_bytes(output.getvalue())


def _share_drawing_image_target(path: Path) -> str:
    relationship_part = "xl/drawings/_rels/drawing1.xml.rels"
    with zipfile.ZipFile(path) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    output = BytesIO()
    shared_image_part = ""
    with zipfile.ZipFile(output, "w") as archive:
        for info, payload in members:
            if info.filename == relationship_part:
                root = etree.fromstring(payload)
                image_relationships = [
                    relationship
                    for relationship in root.findall(
                        f"{{{_PACKAGE_REL_NS}}}Relationship"
                    )
                    if relationship.get("Type", "").endswith("/image")
                ]
                assert len(image_relationships) == 2
                target = image_relationships[0].get("Target")
                assert target is not None
                image_relationships[1].set("Target", target)
                shared_image_part = target.lstrip("/")
                payload = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            archive.writestr(info, payload)
    assert shared_image_part
    path.write_bytes(output.getvalue())
    return shared_image_part


def _relative_inventory(root: Path) -> tuple[str, ...]:
    return tuple(sorted(str(item.relative_to(root)) for item in root.rglob("*")))


if __name__ == "__main__":
    unittest.main()
