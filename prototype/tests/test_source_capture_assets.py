from __future__ import annotations

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import base64
import hashlib
import shutil
import tempfile
import unittest
from contextlib import ExitStack
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image

from app import db
from app.adapters.persistence.sql_source_asset_publication import (
    SqlSourceAssetPublicationRepository,
)
from app.adapters.persistence.sql_source_capture import SqlSourceCaptureRepository
from app.adapters.persistence.sql_source_capture import SqlSourceFactsReader
from app.adapters.storage.local_object_store import LocalObjectStore
from app.application.source_asset_publication import release_source_asset_candidates
from app.application.source_asset_publication import SourceAssetPublicationAttempt
from app.application.source_capture import capture_current_document_source
from app.domain.extraction.source_ir import SourceAssetKind, SourceAssetRef


_PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEElEQVR4nGP8zwACTGCS"
    "AQANHQEDgslx/wAAAABJRU5ErkJggg=="
)
_NOW = datetime(2026, 8, 13, 12, 0, tzinfo=timezone.utc)


class SourceCaptureAssetLifecycleTests(unittest.TestCase):
    def setUp(self) -> None:
        self.stack = ExitStack()
        self.root = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        self.stack.enter_context(patch.object(db, "DB_PATH", self.root / "app.db"))
        db.init_db()
        self.project = db.create_project("Source capture assets", "Customer")
        self.object_store = LocalObjectStore(self.root / "objects")

    def tearDown(self) -> None:
        self.stack.close()

    def test_pdf_docx_and_xlsx_capture_publish_every_asset_before_binding(self) -> None:
        sources = (
            self._write_pdf("source.pdf"),
            self._write_docx("source.docx"),
            self._write_xlsx("source.xlsx"),
        )
        captures = []
        for source in sources:
            document = self._add_document(source)
            captures.append(
                capture_current_document_source(
                    project_id=self.project["id"],
                    sor_document_id=document["id"],
                    expected_source_sha256=document["source_sha256"],
                    repository=SqlSourceCaptureRepository(),
                    object_store=self.object_store,
                    clock=lambda: _NOW,
                )
            )

        for capture in captures:
            self.assertTrue(capture.source_ir.assets)
            for asset in capture.source_ir.assets:
                stored = self.object_store.stat(asset.storage_key)
                self.assertIsNotNone(stored)
                assert stored is not None
                self.assertEqual(stored.sha256, asset.content_sha256)
                self.assertTrue(asset.storage_key.startswith("source-assets/"))
        with db.get_conn() as connection:
            self.assertEqual(
                int(
                    connection.execute(
                        "SELECT COUNT(*) AS count FROM source_capture_asset_publications"
                    ).fetchone()["count"]
                ),
                0,
            )
            self.assertEqual(
                int(
                    connection.execute(
                        "SELECT COUNT(*) AS count FROM source_capture_asset_staging_references"
                    ).fetchone()["count"]
                ),
                0,
            )
            self.assertEqual(
                int(
                    connection.execute(
                        "SELECT COUNT(*) AS count FROM source_capture_asset_references"
                    ).fetchone()["count"]
                ),
                sum(len(capture.source_ir.assets) for capture in captures),
            )

    def test_capture_with_assets_cannot_commit_without_publication(self) -> None:
        document = self._add_document(self._write_pdf("unfenced.pdf"))
        capture = capture_current_document_source(
            project_id=self.project["id"],
            sor_document_id=document["id"],
            repository=SqlSourceCaptureRepository(),
            object_store=self.object_store,
            clock=lambda: _NOW,
        )
        other_document = self._add_document(self._write_pdf("other.pdf"))
        invalid = capture.__class__(
            capture_id="unfenced-capture",
            project_id=self.project["id"],
            sor_document_id=other_document["id"],
            source_sha256=capture.source_sha256,
            source_ir=capture.source_ir,
            created_at_utc=_NOW,
        )
        with self.assertRaisesRegex(RuntimeError, "require an open"):
            SqlSourceCaptureRepository().append_capture(invalid)

    def test_reader_rejects_capture_when_asset_reference_metadata_is_missing(self) -> None:
        document = self._add_document(self._write_pdf("corrupt.pdf"))
        capture = capture_current_document_source(
            project_id=self.project["id"],
            sor_document_id=document["id"],
            repository=SqlSourceCaptureRepository(),
            object_store=self.object_store,
            clock=lambda: _NOW,
        )
        with db.get_conn() as connection:
            connection.execute(
                "DELETE FROM source_capture_asset_references WHERE capture_id = ?",
                (capture.capture_id,),
            )
        with self.assertRaisesRegex(RuntimeError, "do not exact-match"):
            SqlSourceFactsReader().get_capture(
                project_id=self.project["id"],
                sor_document_id=document["id"],
                source_sha256=document["source_sha256"],
            )

    def test_shared_capture_asset_is_deleted_only_after_last_binding(self) -> None:
        first_path = self._write_pdf("shared-a.pdf")
        second_path = self.root / "shared-b.pdf"
        shutil.copyfile(first_path, second_path)
        first = self._add_document(first_path)
        second = self._add_document(second_path)
        repository = SqlSourceCaptureRepository()
        first_capture = capture_current_document_source(
            project_id=self.project["id"],
            sor_document_id=first["id"],
            repository=repository,
            object_store=self.object_store,
            clock=lambda: _NOW,
        )
        second_capture = capture_current_document_source(
            project_id=self.project["id"],
            sor_document_id=second["id"],
            repository=repository,
            object_store=self.object_store,
            clock=lambda: _NOW,
        )
        self.assertEqual(first_capture.capture_id, second_capture.capture_id)
        key = first_capture.source_ir.assets[0].storage_key

        first_deleted = db.delete_sor_document(first["id"])
        assert first_deleted is not None
        release_source_asset_candidates(
            tuple(first_deleted["_released_source_asset_storage_keys"]),
            object_store=self.object_store,
        )
        self.assertIsNotNone(self.object_store.stat(key))

        second_deleted = db.delete_sor_document(second["id"])
        assert second_deleted is not None
        release_source_asset_candidates(
            tuple(second_deleted["_released_source_asset_storage_keys"]),
            object_store=self.object_store,
        )
        self.assertIsNone(self.object_store.stat(key))
        self.assertNotIn(
            key,
            SqlSourceAssetPublicationRepository().list_unreferenced(),
        )

    def test_aborted_publication_removes_unreferenced_bytes_and_fence(self) -> None:
        document = self._add_document(self._write_pdf("aborted.pdf"))
        payload = b"private-source-asset"
        digest = hashlib.sha256(payload).hexdigest()
        asset = SourceAssetRef(
            asset_id="asset-aborted",
            kind=SourceAssetKind.IMAGE,
            content_sha256=digest,
            media_type="image/png",
            storage_key=f"source-assets/{digest[:2]}/{digest}.png",
        )
        attempt = SourceAssetPublicationAttempt(
            project_id=self.project["id"],
            sor_document_id=document["id"],
            source_sha256=document["source_sha256"],
            object_store=self.object_store,
            created_at_utc=_NOW,
        )
        attempt.publish(asset, payload)
        self.assertIsNotNone(self.object_store.stat(asset.storage_key))

        attempt.abort()

        self.assertIsNone(self.object_store.stat(asset.storage_key))
        with db.get_conn() as connection:
            self.assertEqual(
                int(
                    connection.execute(
                        "SELECT COUNT(*) AS count FROM source_capture_asset_publications"
                    ).fetchone()["count"]
                ),
                0,
            )
            self.assertEqual(
                int(
                    connection.execute(
                        "SELECT COUNT(*) AS count FROM stored_objects"
                    ).fetchone()["count"]
                ),
                0,
            )

    def test_postgresql_migrations_and_global_gc_include_capture_references(self) -> None:
        reading_migration = (
            Path(__file__).resolve().parents[1]
            / "migrations"
            / "postgresql"
            / "0024_reading_workspace.sql"
        ).read_text(encoding="utf-8")
        asset_migration = (
            Path(__file__).resolve().parents[1]
            / "migrations"
            / "postgresql"
            / "0025_source_capture_asset_lifecycle.sql"
        ).read_text(encoding="utf-8")
        self.assertIn("source_capture_snapshots", reading_migration)
        self.assertIn("reading_sessions", reading_migration)
        self.assertNotIn("source_capture_asset_publications", reading_migration)
        self.assertIn("source_capture_asset_publications", asset_migration)
        self.assertIn("source_capture_asset_staging_references", asset_migration)
        self.assertIn("source_capture_asset_references", asset_migration)
        self.assertIn("REFERENCES stored_objects", asset_migration)
        self.assertIn("REFERENCES sor_documents", asset_migration)
        self.assertIn(
            "CREATE OR REPLACE VIEW source_asset_reference_counts",
            asset_migration,
        )

        db_source = (Path(db.__file__)).read_text(encoding="utf-8")
        gc_start = db_source.index(
            "def claim_unreferenced_source_asset_object_deletion"
        )
        gc_function = db_source[gc_start : db_source.index("def create_project", gc_start)]
        self.assertIn("SELECT storage_key FROM source_capture_asset_references", gc_function)
        self.assertIn(
            "SELECT storage_key FROM source_capture_asset_staging_references",
            gc_function,
        )

    def _add_document(self, path: Path) -> dict:
        payload = path.read_bytes()
        suffix = path.suffix.lower()
        detected = suffix.lstrip(".")
        return db.add_sor_document(
            self.project["id"],
            path.name,
            detected.upper(),
            str(path),
            "",
            source_sha256=hashlib.sha256(payload).hexdigest(),
            source_size=len(payload),
            declared_extension=suffix,
            detected_format=detected,
            preflight_schema_version="document-preflight/v1",
        )

    def _write_pdf(self, name: str) -> Path:
        path = self.root / name
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), "Source facts")
        page.insert_image(fitz.Rect(72, 100, 140, 168), stream=_PNG_BYTES)
        document.save(path)
        document.close()
        return path

    def _write_docx(self, name: str) -> Path:
        path = self.root / name
        image_path = self.root / f"{name}.png"
        image_path.write_bytes(_PNG_BYTES)
        document = Document()
        document.add_paragraph("Source facts")
        document.add_picture(str(image_path))
        document.save(path)
        image_path.unlink()
        return path

    def _write_xlsx(self, name: str) -> Path:
        path = self.root / name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "Source facts"
        image_stream = BytesIO(_PNG_BYTES)
        worksheet.add_image(Image(image_stream), "B2")
        try:
            workbook.save(path)
        finally:
            image_stream.close()
            workbook.close()
        return path


if __name__ == "__main__":
    unittest.main()
