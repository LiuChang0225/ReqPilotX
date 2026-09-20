
if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

from argparse import Namespace
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import hashlib
import json
import sys
import unittest

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'plugins/reqpilot/scripts'))
import artifact_tools
from app.adapters.documents.docx_source_ir_adapter import DocxSourceIRAdapter
from app.adapters.documents.xlsx_source_ir_adapter import XlsxSourceIRAdapter
from app.adapters.documents.preflight import LocalDocumentPreflight
from app.application.agent_workflow import source_requirement_baseline as core


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')


def fixture(root, suffix):
    stream=BytesIO();Image.new('RGB',(120,60),(24,92,66)).save(stream,format='PNG');payload=stream.getvalue()
    path=root/('requirements.'+suffix)
    if suffix=='docx':
        doc=Document();doc.add_heading('1 Interface limits',1)
        doc.add_paragraph('The controller shall report readiness within 20 ms and retain the last fault code.')
        table=doc.add_table(rows=3,cols=3)
        table.cell(0,0).merge(table.cell(0,1)).text='Operating condition';table.cell(0,2).text='Limit'
        table.cell(1,0).merge(table.cell(2,0)).text='12 V';table.cell(1,1).text='Startup';table.cell(1,2).text='20 ms'
        table.cell(2,1).text='Recovery';table.cell(2,2).text='30 ms'
        doc.add_picture(BytesIO(payload),width=Inches(1));doc.save(path)
        adapter=DocxSourceIRAdapter()
    else:
        book=Workbook();sheet=book.active;sheet.title='Z-first';sheet.merge_cells('A1:B1');sheet['A1']='Operating limits';sheet['C1']='Criterion'
        sheet.append(['Startup','12 V','20 ms']);sheet.append(['Recovery','12 V','30 ms'])
        # Separate same-sheet region, with a merged header and real blank grid between.
        sheet.merge_cells('E1:F1');sheet['E1']='Delivery';sheet['E2']='Sample units';sheet['F2']=3
        sheet.add_image(SheetImage(BytesIO(payload)),'H2')
        sheet=book.create_sheet('A-second');sheet.append(['Object','Limit']);sheet.append(['Reset','40 ms']);book.save(path)
        adapter=XlsxSourceIRAdapter()
    snapshot=adapter.extract(path,preflight=LocalDocumentPreflight().inspect(path),created_at=datetime(2026,9,8,tzinfo=timezone.utc))
    write(root/'source-ir.json',snapshot.model_dump(mode='json'))
    return path,snapshot,payload


def map_decisions(source,snapshot):
    headings=[node for node in snapshot.nodes if node.kind.value=='heading' and not node.parent_source_node_id]
    visible=[node for node in snapshot.nodes if not node.parent_source_node_id and node.table_payload is None]
    tables=[node for node in snapshot.nodes if node.table_payload is not None]
    return {'schema_version':'reqpilot-map-decisions/v1','source_sha256':snapshot.source_sha256,
        'nodes':[{'node_id':'document','kind':'DOCUMENT','title':headings[0].raw_text if headings else source.name,
            'title_source_ids':[node.source_node_id for node in headings],
            'content_source_ids':[node.source_node_id for node in visible if node not in headings]}],
        'logical_tables':[{'scope_id':f'table-{index}','owner_node_id':'document','title':f'Table {index}',
            'role':'REQUIREMENTS','fragments':[{'source_node_id':node.source_node_id}]} for index,node in enumerate(tables)]}


def item_decisions(map_dir):
    mapping=json.loads((map_dir/'document-map.json').read_text());scopes=json.loads((map_dir/'scope-requests.json').read_text())
    judgment={'requirementness':'REQUIREMENT','requirement_kinds':['FUNCTIONAL'],'confidence':0.9,
        'decision_summary':'Explicit synthetic fixture judgment.','reason_codes':['fixture_requirement']}
    texts=[]
    for request in scopes['text_requests']:
        items=[];structure=[]
        for index,content in enumerate(request['source_content']):
            span={'content_id':content['content_id'],'start_offset':0,'end_offset':len(content['text'])}
            if content['is_title']:structure.append({'text_span':span,'reason_codes':['source_title']})
            else:items.append({'item_key':f'paragraph-{index}','label':content['text'],'source_form':'PARAGRAPH','text_spans':[span],**judgment})
        texts.append({'scope_id':request['scope_id'],'items':items,'structure_only_spans':structure})
    tables=[]
    for request in scopes['table_requests']:
        records=[];contexts=[]
        for fragment in request['fragments']:
            def group(lo,hi):return {'fragment_id':fragment['fragment_id'],'row_start':lo,'row_end':hi}
            contexts.append({'role':'HEADER','row_group':group(0,1),'reason_codes':['table_header']})
            for row in range(1,fragment['row_count']):
                records.append({'record_key':f'row-{row}','label':f'Original row {row}','source_form':'TABLE_RECORD','row_groups':[group(row,row+1)],**judgment})
        tables.append({'logical_table_scope_id':request['logical_table_scope_id'],'records':records,'context_row_groups':contexts})
    return {'schema_version':'reqpilot-item-decisions/v1','source_sha256':mapping['source_sha256'],
        'document_map_artifact_sha256':mapping['artifact_sha256'],'text_decisions':texts,'table_decisions':tables}


class OfficeSourcePipelineTests(unittest.TestCase):
    def test_native_map_and_items_preserve_source_locator_and_order(self):
        for suffix in ['docx','xlsx']:
            with self.subTest(format=suffix),TemporaryDirectory() as tmp:
                root=Path(tmp);source,snapshot,payload=fixture(root,suffix)
                write(root/'map-input.json',map_decisions(source,snapshot))
                result=artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
                self.assertEqual(result['evidence_state'],'SUFFICIENT')
                write(root/'items.json',item_decisions(root/'map'))
                result=artifact_tools.compile_items(Namespace(map_dir=root/'map',input=root/'items.json',output_dir=root/'items',baseline_dir=None))
                baseline=core.SourceRequirementBaselineArtifactV1.model_validate_json((root/'items/source-requirement-baseline.json').read_text())
                self.assertEqual(baseline.source_sha256,hashlib.sha256(source.read_bytes()).hexdigest())
                self.assertGreater(len(baseline.records),1)
                self.assertTrue(all(not record.pages for record in baseline.records))
                for record in baseline.records:
                    for binding in (*record.table_fragments,*record.text_bindings):
                        self.assertEqual(binding.native_locator.format.value,suffix)
                        self.assertIsNone(binding.page);self.assertIsNone(binding.bbox);self.assertIsNone(binding.source_bounds)
                from app.application.agent_workflow.source_item_ledger import source_item_position
                positions=[source_item_position(record) for record in baseline.records]
                self.assertEqual(positions,sorted(positions))
                self.assertEqual(len(set(positions)),len(positions))

    def test_native_images_analysis_and_program_excel(self):
        from app.application.agent_workflow.source_item_visual import build_source_item_image_scope_packet, compile_source_item_images
        from app.application.agent_workflow.source_item_decomposition import build_source_item_decomposition_packet, compile_source_item_decomposition
        from app.application.agent_workflow.source_item_analysis import build_source_item_analysis_packet, compile_source_item_assessments
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot, export_delivery_xlsx
        from zipfile import ZipFile
        from openpyxl import load_workbook
        for suffix in ['docx','xlsx']:
            with self.subTest(format=suffix),TemporaryDirectory() as tmp:
                root=Path(tmp);source,snapshot,payload=fixture(root,suffix)
                write(root/'map-input.json',map_decisions(source,snapshot))
                artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
                write(root/'items.json',item_decisions(root/'map'))
                artifact_tools.compile_items(Namespace(map_dir=root/'map',input=root/'items.json',output_dir=root/'items',baseline_dir=None))
                mapping=core.DocumentMapArtifactV1.model_validate_json((root/'map/document-map.json').read_text())
                baseline=core.SourceRequirementBaselineArtifactV1.model_validate_json((root/'items/source-requirement-baseline.json').read_text())
                tables=json.loads((root/'map/table-previews.json').read_text())
                packet=build_source_item_image_scope_packet(mapping,node_id='document',baseline=baseline,table_previews=tables)
                baseline,receipt,annotations=compile_source_item_images(mapping,[{'node_id':'document','packet_sha256':packet['packet_sha256'],'items':[
                    {'item_key':'source-diagram','display_label':'Synthetic diagram','image_anchor_ids':[value['anchor_id'] for value in packet['images']],
                     'requirementness':'NON_REQUIREMENT','non_requirement_role':'CONTEXT','confidence':0.9,'decision_summary':'Fixture image retained as source context.',
                     'reason_codes':['source_illustration'],'visual_interpretation':'Solid green rectangular image.'}]}],previous_baseline=baseline,table_previews=tables)
                decomposition=None
                if suffix=='docx':
                    parent=next(record for record in baseline.records if record.text_bindings)
                    packet=build_source_item_decomposition_packet(baseline,parent.source_record_id,table_previews=tables)
                    text=packet['segments'][0]['text'];cut=text.index(' and ')
                    decomposition=compile_source_item_decomposition(baseline,[{'source_item_id':parent.source_record_id,'action':'SPLIT','reason':'Two explicit, independent required outcomes in the fixture.','children':[
                        {'text':text[:cut], 'source_ranges':[{'segment_index':0,'start_offset':0,'end_offset':cut}],'requirement_kinds':['FUNCTIONAL']},
                        {'text':text[cut:].strip(), 'source_ranges':[{'segment_index':0,'start_offset':cut,'end_offset':len(text)}],'requirement_kinds':['FUNCTIONAL']}]}],table_previews=tables)
                catalog={'schema_version':'project-knowledge-facts/v1','project_id':'office-fixture','knowledge_base_version':0,'evidence_catalog':{},'documents':[],'gaps':[]}
                catalog['catalog_sha256']=core.canonical_json_sha256(catalog)
                packet=build_source_item_analysis_packet(mapping,baseline,tables,catalog,decomposition=decomposition,image_annotations=annotations)
                assessments=compile_source_item_assessments(packet,{'decisions':[{'unit_id':unit['unit_id'],'decision':'PENDING_EVIDENCE',
                    'reason':'No internal product or test document was provided for this fixture.','evidence':[],
                    'clarification_question':'Please provide the current product test report for '+unit['requirement_display_id']+'.'} for unit in packet['units']]})
                delivery=build_delivery_snapshot(mapping,baseline,tables,annotations=annotations,decomposition=decomposition,assessments=assessments)
                report={};data=export_delivery_xlsx(delivery,source_file_path=str(source),source_sha256=baseline.source_sha256,image_export_receipt=report)
                target=root/'matrix.xlsx';target.write_bytes(data)
                with ZipFile(target) as package:
                    media=[package.read(name) for name in package.namelist() if name.startswith('xl/media/')]
                    self.assertIn(payload,media)
                book=load_workbook(target);self.assertEqual(sum(len(sheet._images) for sheet in book),1)
                self.assertTrue(all(node['page_start'] is None and node['page_end'] is None for node in delivery['nodes']))
                self.assertEqual(report['status'],'VERIFIED')

    def test_original_office_locator_and_image_tampering_is_rejected(self):
        from copy import deepcopy
        from app.application.agent_workflow.office_source_evidence import verify_office_source_ir, validate_office_image_reference
        for suffix in ['docx','xlsx']:
            with self.subTest(format=suffix),TemporaryDirectory() as tmp:
                root=Path(tmp);source,snapshot,payload=fixture(root,suffix)
                value=snapshot.model_copy(update={'nodes':(snapshot.nodes[0].model_copy(update={'raw_text':'An invented customer statement'}),*snapshot.nodes[1:])})
                with self.assertRaisesRegex(ValueError,'differs from replaying'):
                    verify_office_source_ir(source,value)
                image=next(node for node in snapshot.nodes if node.kind.value=='image')
                asset=next(asset for asset in snapshot.assets if asset.asset_id==image.asset_ref)
                figure={'source_sha256':snapshot.source_sha256,'source_handle':image.source_node_id,'anchor_id':image.source_node_id,
                    'asset_id':asset.asset_id,'content_sha256':asset.content_sha256,'location':image.native_locator.model_dump(mode='json',exclude_none=True)}
                self.assertTrue(validate_office_image_reference(source,snapshot.source_sha256,figure))
                wrong=deepcopy(figure)
                if suffix=='xlsx':wrong['location']['cell_range']='A99'
                else:wrong['location']['paragraph_index']+=1
                with self.assertRaisesRegex(ValueError,'actual original object'):
                    validate_office_image_reference(source,snapshot.source_sha256,wrong)

    def test_office_map_successor_replays_original_locations(self):
        from app.application.agent_workflow.document_map_revision import compile_document_map_revision,rebind_baseline_for_document_map_revision
        for suffix in ['docx','xlsx']:
            with self.subTest(format=suffix),TemporaryDirectory() as tmp:
                root=Path(tmp);source,snapshot,payload=fixture(root,suffix)
                decisions=map_decisions(source,snapshot);write(root/'map-input.json',decisions)
                artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
                write(root/'items.json',item_decisions(root/'map'))
                artifact_tools.compile_items(Namespace(map_dir=root/'map',input=root/'items.json',output_dir=root/'items',baseline_dir=None))
                old=core.DocumentMapArtifactV1.model_validate_json((root/'map/document-map.json').read_text())
                baseline=core.SourceRequirementBaselineArtifactV1.model_validate_json((root/'items/source-requirement-baseline.json').read_text())
                receipt=core.SourceRequirementBaselineValidationReceiptV1.model_validate_json((root/'items/validation-receipt.json').read_text())
                tables=json.loads((root/'map/table-previews.json').read_text())
                decisions['logical_tables'][0]['agent_display_label']='Explicit fixture display label'
                write(root/'next-map-input.json',decisions)
                artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'next-map-input.json',output_dir=root/'next-map'))
                candidate=json.loads((root/'next-map/document-map.json').read_text())
                revised=compile_document_map_revision(old,candidate)
                rebound=rebind_baseline_for_document_map_revision(old,revised['document_map'],baseline,receipt,tables)
                next_base=core.SourceRequirementBaselineArtifactV1.model_validate(rebound['baseline'])
                self.assertEqual([record.source_record_id for record in baseline.records],[record.source_record_id for record in next_base.records])
                self.assertEqual([record.text_bindings for record in baseline.records],[record.text_bindings for record in next_base.records])
                self.assertEqual([record.table_fragments for record in baseline.records],[record.table_fragments for record in next_base.records])

    def test_number_annotation_uses_original_native_paragraph_without_fake_geometry(self):
        from app.application.agent_workflow.office_source_evidence import capture_native_office_source
        from app.application.agent_workflow.source_item_visual import compile_source_item_annotation_updates,validate_source_item_number_evidence
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source,_,_=fixture(root,'docx');doc=Document(source)
            doc.paragraphs[1].text='1.1 The controller shall retain the last fault code.';doc.save(source)
            snapshot=capture_native_office_source(source,hashlib.sha256(source.read_bytes()).hexdigest())
            write(root/'source-ir.json',snapshot.model_dump(mode='json'));write(root/'map-input.json',map_decisions(source,snapshot))
            artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
            write(root/'items.json',item_decisions(root/'map'))
            artifact_tools.compile_items(Namespace(map_dir=root/'map',input=root/'items.json',output_dir=root/'items',baseline_dir=None))
            mapping=core.DocumentMapArtifactV1.model_validate_json((root/'map/document-map.json').read_text())
            baseline=core.SourceRequirementBaselineArtifactV1.model_validate_json((root/'items/source-requirement-baseline.json').read_text())
            record=next(record for record in baseline.records if record.text_bindings)
            annotations=compile_source_item_annotation_updates(mapping,baseline,[{'source_item_id':record.source_record_id,'source_number':'1.1',
                'source_number_ref':{'kind':'TEXT','content_id':record.text_bindings[0].content_id,'start_offset':0,'end_offset':3},'reason':'The exact leading identifier occurs in the authored original paragraph.'}])
            evidence=annotations['records'][record.source_record_id]['source_number_evidence']
            self.assertEqual(evidence['native_locator']['paragraph_index'],1);self.assertIsNone(evidence['page']);self.assertIsNone(evidence['bbox'])
            self.assertEqual(validate_source_item_number_evidence(mapping,baseline,record.source_record_id,'1.1',evidence),evidence)

    def test_imported_native_map_and_table_preview_must_match_original_file(self):
        from copy import deepcopy
        from app.application.agent_workflow.office_source_evidence import verify_office_map_against_source
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source,snapshot,payload=fixture(root,'xlsx');write(root/'map-input.json',map_decisions(source,snapshot))
            artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
            mapping=core.DocumentMapArtifactV1.model_validate_json((root/'map/document-map.json').read_text());tables=json.loads((root/'map/table-previews.json').read_text())
            self.assertEqual(verify_office_map_against_source(source,mapping,tables)['status'],'ORIGINAL_NATIVE_LOCATIONS_VERIFIED')
            bad=deepcopy(tables);next(iter(bad.values()))['fragments'][0]['cells'][0]['text']='An invented customer condition'
            with self.assertRaisesRegex(ValueError,'actual original value'):
                verify_office_map_against_source(source,mapping,bad)
            bad=deepcopy(tables);next(iter(bad.values()))['fragments'][0]['cells'][0]['native_locator']['cell_range']='A9'
            with self.assertRaisesRegex(ValueError,'native sheet/range/offset'):
                verify_office_map_against_source(source,mapping,bad)

    def test_docx_cell_image_has_one_actual_owner_and_exports_original_pixels(self):
        from app.application.agent_workflow.office_source_evidence import capture_native_office_source
        from app.application.agent_workflow.source_item_visual import build_source_item_image_scope_packet,compile_source_item_images
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot,export_delivery_xlsx
        from zipfile import ZipFile
        with TemporaryDirectory() as tmp:
            root=Path(tmp);source,_,payload=fixture(root,'docx');doc=Document(source)
            doc.tables[0].cell(1,1).paragraphs[0].add_run().add_picture(BytesIO(payload),width=Inches(.5));doc.save(source)
            snapshot=capture_native_office_source(source,hashlib.sha256(source.read_bytes()).hexdigest())
            decisions=map_decisions(source,snapshot)
            decisions['nodes'][0]['content_source_ids'] += [node.source_node_id for node in snapshot.nodes if node.parent_source_node_id and node.kind.value=='image']
            write(root/'source-ir.json',snapshot.model_dump(mode='json'));write(root/'map-input.json',decisions)
            artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
            write(root/'items.json',item_decisions(root/'map'));artifact_tools.compile_items(Namespace(map_dir=root/'map',input=root/'items.json',output_dir=root/'items',baseline_dir=None))
            mapping=core.DocumentMapArtifactV1.model_validate_json((root/'map/document-map.json').read_text());base=core.SourceRequirementBaselineArtifactV1.model_validate_json((root/'items/source-requirement-baseline.json').read_text());tables=json.loads((root/'map/table-previews.json').read_text())
            packet=build_source_item_image_scope_packet(mapping,node_id='document',baseline=base,table_previews=tables)
            self.assertEqual(len(packet['table_image_relations']),1);self.assertEqual(len(packet['images']),1)
            baseline,_,annotations=compile_source_item_images(mapping,[{'node_id':'document','packet_sha256':packet['packet_sha256'],'items':[
                {'item_key':'remaining-image','display_label':'Original independent diagram','image_anchor_ids':[packet['images'][0]['anchor_id']],
                 'requirementness':'NON_REQUIREMENT','non_requirement_role':'CONTEXT','confidence':0.9,'decision_summary':'Synthetic graphical context.','reason_codes':['fixture_image']}]}],previous_baseline=base,table_previews=tables)
            self.assertEqual(len(baseline.table_image_relations),1)
            delivery=build_delivery_snapshot(mapping,baseline,tables,annotations=annotations);report={}
            data=export_delivery_xlsx(delivery,source_file_path=str(source),source_sha256=baseline.source_sha256,image_export_receipt=report)
            with ZipFile(BytesIO(data)) as archive:self.assertIn(payload,[archive.read(name) for name in archive.namelist() if name.startswith('xl/media/')])
            self.assertEqual(report['status'],'VERIFIED')
            self.assertEqual(len({image['anchor_id'] for image in report['references']}),2)
