from __future__ import annotations

import asyncio
import io
import json
import os
import tempfile
import time
import unittest
import sqlite3
from contextlib import ExitStack
from pathlib import Path
from threading import Event
from unittest.mock import patch

import fitz
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image as PILImage
from starlette.datastructures import UploadFile

from app import db, main
from app.services.complex_table_mode import _row_shadow_signature
from app.services import exporter, extraction_preview, preview_cache, table_assets, extractor
from app.services.ai_copilot import AICopilotConfig, clear_runtime_ai_provider_overrides
from app.services.exporter import export_compliance_matrix


class ImmediateThread:
    def __init__(self, target=None, args=(), kwargs=None, name=None, daemon=None):
        self._target = target
        self._args = args
        self._kwargs = kwargs or {}
        self._alive = False

    def start(self) -> None:
        self._alive = True
        try:
            if self._target:
                self._target(*self._args, **self._kwargs)
        finally:
            self._alive = False

    def is_alive(self) -> bool:
        return self._alive


def _fake_preview() -> dict:
    return {
        "summary": {
            "dedupe_version": "heading_requirement_v1",
            "dedup_exact_count": 1,
            "dedup_prefix_count": 0,
            "hybrid_entries_count": 0,
        },
        "sections": [
            {
                "scope": "Main",
                "section_code": "1",
                "title": "概述",
                "page_span": "1-1",
                "requirement_count": 1,
            }
        ],
        "complete_requirement_entries": [
            {
                "entry_id": "REQ0001",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "供应商需要满足总成试验要求。",
                "raw_text": "供应商需要满足总成试验要求。",
                "requirement_text": "供应商需要满足总成试验要求。",
                "source_page": 1,
                "source_page_end": 1,
                "source_anchor": "Main > 1",
                "parent_section_code": "1",
                "parent_section_title": "概述",
                "parent_section_label": "1 概述",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
                "text_relation": "distinct",
                "is_deduped": False,
                "merged_from_req_codes": [],
                "merged_from_entry_ids": [],
            }
        ],
        "complete_requirement_entries_raw": [
            {
                "entry_id": "H-1",
                "req_code": "",
                "entry_kind": "HEADING",
                "display_text": "1 概述",
                "raw_text": "1 概述",
                "requirement_text": "1 概述",
                "source_page": 1,
                "source_page_end": 1,
                "source_anchor": "Main > 1",
                "parent_section_code": "1",
                "parent_section_title": "概述",
                "parent_section_label": "1 概述",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
                "text_relation": "exact_match",
                "is_deduped": True,
                "merged_from_req_codes": [],
                "merged_from_entry_ids": [],
            },
            {
                "entry_id": "REQ0001",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "供应商需要满足总成试验要求。",
                "raw_text": "供应商需要满足总成试验要求。",
                "requirement_text": "供应商需要满足总成试验要求。",
                "source_page": 1,
                "source_page_end": 1,
                "source_anchor": "Main > 1",
                "parent_section_code": "1",
                "parent_section_title": "概述",
                "parent_section_label": "1 概述",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
                "text_relation": "distinct",
                "is_deduped": False,
                "merged_from_req_codes": [],
                "merged_from_entry_ids": [],
            },
        ],
    }


def _fake_preview_with_reversed_req_codes() -> dict:
    return {
        "summary": {"dedupe_version": "heading_requirement_v1"},
        "sections": [],
        "complete_requirement_entries": [
            {
                "entry_id": "REQ0024",
                "req_code": "REQ0024",
                "entry_kind": "REQUIREMENT",
                "display_text": "5.1 首条需求",
                "raw_text": "5.1 首条需求",
                "requirement_text": "5.1 首条需求",
                "source_page": 5,
                "source_page_end": 5,
                "source_anchor": "Main > 5.1",
                "parent_section_code": "5.1",
                "parent_section_title": "首条需求",
                "parent_section_label": "5.1 首条需求",
            },
            {
                "entry_id": "REQ0023",
                "req_code": "REQ0023",
                "entry_kind": "REQUIREMENT",
                "display_text": "5.7 次条需求",
                "raw_text": "5.7 次条需求",
                "requirement_text": "5.7 次条需求",
                "source_page": 5,
                "source_page_end": 5,
                "source_anchor": "Main > 5.7",
                "parent_section_code": "5.7",
                "parent_section_title": "次条需求",
                "parent_section_label": "5.7 次条需求",
            },
        ],
        "complete_requirement_entries_raw": [],
    }


def _fake_preview_with_reused_req_codes() -> dict:
    return {
        "summary": {"dedupe_version": "heading_requirement_v1"},
        "sections": [],
        "complete_requirement_entries": [
            {
                "entry_id": "REQ0001",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "5.2 第二条需求",
                "raw_text": "5.2 第二条需求",
                "requirement_text": "5.2 第二条需求",
                "source_page": 5,
                "source_page_end": 5,
                "source_anchor": "Main > 5.2",
                "parent_section_code": "5.2",
                "parent_section_title": "第二条需求",
                "parent_section_label": "5.2 第二条需求",
            },
            {
                "entry_id": "REQ0999",
                "req_code": "REQ0999",
                "entry_kind": "REQUIREMENT",
                "display_text": "5.1 第一条需求",
                "raw_text": "5.1 第一条需求",
                "requirement_text": "5.1 第一条需求",
                "source_page": 5,
                "source_page_end": 5,
                "source_anchor": "Main > 5.1",
                "parent_section_code": "5.1",
                "parent_section_title": "第一条需求",
                "parent_section_label": "5.1 第一条需求",
            },
        ],
        "complete_requirement_entries_raw": [],
    }


def _fake_preview_with_same_code_local_clause_in_section_six() -> dict:
    section_six_title = "投标说明"
    clause_titles = {
        1: "技术标书为简体中文",
        2: "技术标书须包含完整技术方案",
        3: "技术方案需包含关键器件选型",
        4: "技术方案需包含测试验证计划",
        5: "技术方案需包含质量控制方案",
        6: "技术标书中应包含生产工艺部分的描述，针对技术标书中提出的技术方案，论证工艺可行性",
        7: "技术标书需说明量产能力",
        8: "技术标书需说明项目进度计划",
        9: "技术标书需说明交流文件要求",
        10: "商务报价",
    }

    def _heading_row() -> dict:
        return {
            "entry_id": "H-6",
            "req_code": "",
            "entry_kind": "HEADING",
            "section_code": "6",
            "section_key": "6",
            "section_instance_key": "6",
            "section_title": section_six_title,
            "section_label": "",
            "owner_section_code": "6",
            "owner_section_title": section_six_title,
            "owner_section_label": "6",
            "parent_section_code": "",
            "parent_section_title": "",
            "parent_section_label": "",
            "structural_parent_section_code": "",
            "structural_parent_section_key": "",
            "structural_parent_section_title": "",
            "structural_parent_section_label": "",
            "display_text": "6 投标说明",
            "raw_text": "6 投标说明",
            "requirement_text": "6 投标说明",
            "source_page": 9,
            "source_page_end": 9,
            "source_anchor": "Page 9 | Section 6 投标说明",
            "text_relation": "exact_match",
            "is_deduped": True,
            "merged_from_req_codes": [],
            "merged_from_entry_ids": [],
        }

    def _requirement_row(number: int) -> dict:
        title = clause_titles[number]
        text = f"{number}）{title}"
        return {
            "entry_id": f"REQ06{number:02d}",
            "req_code": f"REQ06{number:02d}",
            "entry_kind": "REQUIREMENT",
            "section_code": str(number),
            "section_key": f"6::{number}",
            "section_instance_key": f"6::{number}",
            "section_title": title,
            "section_label": f"{number}）",
            "owner_section_code": str(number),
            "owner_section_title": title,
            "owner_section_label": f"{number}）",
            "parent_section_code": str(number),
            "parent_section_title": title,
            "parent_section_label": f"{number} {title}",
            "structural_parent_section_code": "6",
            "structural_parent_section_key": "6",
            "structural_parent_section_title": section_six_title,
            "structural_parent_section_label": "6",
            "display_text": text,
            "raw_text": text,
            "requirement_text": text,
            "source_page": 9,
            "source_page_end": 9,
            "source_anchor": f"Page 9 | Section 6 投标说明 | Clause {number} {title}",
            "requirement_type": "FUNCTIONAL",
            "priority": "MEDIUM",
            "text_relation": "distinct",
            "is_deduped": False,
            "merged_from_req_codes": [],
            "merged_from_entry_ids": [],
        }

    section_tree = [
        {
            "appendix_scope": "",
            "section_code": "6",
            "section_key": "6",
            "section_instance_key": "6",
            "section_title": section_six_title,
            "owner_section_code": "6",
            "owner_section_title": section_six_title,
            "owner_section_label": "6",
            "parent_section_code": "",
            "parent_section_key": "",
            "structural_parent_section_code": "",
            "structural_parent_section_key": "",
            "source_page_start": 9,
            "source_page_end": 9,
            "requirement_count": 10,
            "tree_order": 0,
        }
    ]
    section_tree.extend(
        {
            "appendix_scope": "",
            "section_code": str(number),
            "section_key": f"6::{number}",
            "section_instance_key": f"6::{number}",
            "section_title": clause_titles[number],
            "section_label": f"{number}）",
            "owner_section_code": str(number),
            "owner_section_title": clause_titles[number],
            "owner_section_label": f"{number}）",
            "parent_section_code": "6",
            "parent_section_key": "6",
            "structural_parent_section_code": "6",
            "structural_parent_section_key": "6",
            "source_page_start": 9,
            "source_page_end": 9,
            "requirement_count": 1,
            "tree_order": number,
        }
        for number in range(1, 11)
    )
    return {
        "summary": {"dedupe_version": "heading_requirement_v1"},
        "sections": section_tree,
        "section_tree": section_tree,
        "complete_requirement_entries": [_heading_row(), *[_requirement_row(number) for number in range(1, 11)]],
        "complete_requirement_entries_raw": [_heading_row(), *[_requirement_row(number) for number in range(1, 11)]],
    }


def _fake_preview_with_table_markdown() -> dict:
    table_text = (
        "试验条件如下：\n\n"
        "| Item | Condition |\n"
        "| --- | --- |\n"
        "| Voltage | 12V |\n"
        "| Current | 3A |"
    )
    return {
        "summary": {
            "dedupe_version": "heading_requirement_v1",
            "dedup_exact_count": 0,
            "dedup_prefix_count": 0,
            "hybrid_entries_count": 0,
        },
        "sections": [
            {
                "scope": "Main",
                "section_code": "5.1",
                "title": "Test Setup",
                "page_span": "5-5",
                "requirement_count": 1,
            }
        ],
        "complete_requirement_entries": [
            {
                "entry_id": "REQ0008",
                "req_code": "REQ0008",
                "entry_kind": "REQUIREMENT",
                "display_text": table_text,
                "raw_text": table_text,
                "requirement_text": table_text,
                "source_page": 5,
                "source_page_end": 5,
                "source_anchor": "Main > 5.1",
                "parent_section_code": "5.1",
                "parent_section_title": "Test Setup",
                "parent_section_label": "5.1 Test Setup",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
                "text_relation": "distinct",
                "is_deduped": False,
                "merged_from_req_codes": [],
                "merged_from_entry_ids": [],
            }
        ],
        "complete_requirement_entries_raw": [],
    }


def _fake_preview_with_single_row_slice_table() -> dict:
    table_text = (
        "1 外观检查\n"
        "|1|外观检查|-|目视外观检查|表面无划伤|DV/PV|供应商|6件/OTS|1天|"
    )
    return {
        "summary": {"dedupe_version": "heading_requirement_v1"},
        "sections": [],
        "complete_requirement_entries": [
            {
                "entry_id": "REQ0040",
                "req_code": "REQ0040",
                "entry_kind": "REQUIREMENT",
                "display_text": table_text,
                "raw_text": table_text,
                "requirement_text": table_text,
                "source_page": 20,
                "source_page_end": 20,
                "source_anchor": "Appendix Y > 1",
                "parent_section_code": "Y.1",
                "parent_section_title": "用电器试验大纲",
                "parent_section_label": "Y.1 用电器试验大纲",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
            }
        ],
        "complete_requirement_entries_raw": [],
    }


def _fake_preview_with_complex_table_markdown() -> dict:
    table_text = (
        "5.7 整车用电器与线束匹配评估表\n\n"
        "| Group | Item |\n"
        "| --- | --- |\n"
        "| 测试分组 | Test Groups |\n"
        "| 整车用电器与线束匹配评估表 | 评估项 |"
    )
    return {
        "summary": {"dedupe_version": "heading_requirement_v1"},
        "sections": [],
        "complete_requirement_entries": [
            {
                "entry_id": "REQ0057",
                "req_code": "REQ0057",
                "entry_kind": "REQUIREMENT",
                "display_text": table_text,
                "raw_text": table_text,
                "requirement_text": table_text,
                "source_page": 1,
                "source_page_end": 2,
                "source_anchor": "Main > 5.7",
                "parent_section_code": "5.7",
                "parent_section_title": "整车用电器与线束匹配评估表",
                "parent_section_label": "5.7 整车用电器与线束匹配评估表",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
            }
        ],
        "complete_requirement_entries_raw": [],
    }


def _complex_group_review_row(*, source_snapshot_path: str, shadow_rows: list[dict] | None = None) -> dict:
    shadow_signatures = [_row_shadow_signature(row) for row in (shadow_rows or []) if _row_shadow_signature(row)]
    return {
        "entry_id": "REQ-CX-cx-merged-requirement-matrix-p17-p22-network-security",
        "req_code": "REQ-CX-cx-merged-requirement-matrix-p17-p22-network-security",
        "entry_kind": "REQUIREMENT",
        "display_text": "网络安全\n复杂跨页表分组，内部条目未自动细拆，请结合原文截图人工处理。",
        "raw_text": "网络安全\n复杂跨页表分组，内部条目未自动细拆，请结合原文截图人工处理。",
        "requirement_text": "网络安全\n复杂跨页表分组，内部条目未自动细拆，请结合原文截图人工处理。",
        "source_page": 17,
        "source_page_end": 20,
        "source_anchor": "Complex Table Group cx-merged-requirement-matrix-p17-p22-network-security",
        "parent_section_code": "网络安全",
        "parent_section_title": "网络安全",
        "parent_section_label": "网络安全",
        "requirement_type": "OTHER",
        "priority": "MEDIUM",
        "text_relation": "distinct",
        "is_deduped": False,
        "merged_from_req_codes": ["REQ-CX-cx-merged-requirement-matrix-p17-p22-network-security"],
        "merged_from_entry_ids": ["REQ-CX-cx-merged-requirement-matrix-p17-p22-network-security"],
        "dedupe_version": "heading_requirement_v1",
        "complex_table_mode": True,
        "complex_zone_id": "cx-merged-requirement-matrix-p17-p22",
        "complex_family": "merged_requirement_matrix",
        "complex_role": "group",
        "complex_group_label": "网络安全",
        "complex_split_level": "group",
        "complex_warnings": ["跨页复杂表，包含纵向合并单元格"],
        "complex_row_policy": "exportable_group",
        "table_refs": [
            {
                "table_id": "tbl_complex_group",
                "table_title": "网络安全",
                "source_snapshot_path": source_snapshot_path,
                "slice_snapshot_path": source_snapshot_path,
                "slice_render_mode": "snapshot_only",
                "inherited_headers": [],
                "slice_rows": [],
                "complex_table_mode": True,
                "complex_zone_id": "cx-merged-requirement-matrix-p17-p22",
                "complex_family": "merged_requirement_matrix",
                "complex_split_level": "group",
                "complex_group_label": "网络安全",
                "complex_shadow_signatures": shadow_signatures,
                "export_variant": "source",
                "recommended_export_variant": "source",
                "export_variant_locked": True,
            }
        ],
        "content_blocks": [
            {"kind": "paragraph", "text": "网络安全\n复杂跨页表分组，内部条目未自动细拆，请结合原文截图人工处理。"},
            {"kind": "table", "table_id": "tbl_complex_group"},
        ],
        "has_table_asset": True,
        "table_render_mode": "snapshot_only",
        "keep": True,
    }


class UnifiedWorkspaceTests(unittest.TestCase):
    def setUp(self) -> None:
        clear_runtime_ai_provider_overrides()
        self.tmpdir = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self.tmpdir.name)
        self.exit_stack = ExitStack()
        self.exit_stack.enter_context(patch.object(db, "DB_PATH", self.tmp_path / "app.db"))
        self.exit_stack.enter_context(patch.object(db, "STORAGE_DIR", self.tmp_path))
        self.exit_stack.enter_context(patch.object(main, "EXPORTS_DIR", self.tmp_path / "exports"))
        db.init_db()

    def tearDown(self) -> None:
        clear_runtime_ai_provider_overrides()
        self.exit_stack.close()
        db_file = self.tmp_path / "app.db"
        for _ in range(5):
            try:
                for suffix in ("-wal", "-shm"):
                    sidecar = Path(f"{db_file}{suffix}")
                    if sidecar.exists():
                        sidecar.unlink()
                self.tmpdir.cleanup()
                return
            except OSError:
                time.sleep(0.05)
        self.tmpdir.cleanup()

    def _create_project_with_document(self, suffix: str = "") -> tuple[dict, dict]:
        sample_path = self.tmp_path / f"sample{suffix}.txt"
        sample_path.write_text("sample", encoding="utf-8")
        project = db.create_project(name=f"Workspace Test{suffix}", customer_name="Customer")
        document = db.add_sor_document(
            project_id=project["id"],
            file_name=f"sample{suffix}.txt",
            file_type="txt",
            file_path=str(sample_path),
            extracted_text="sample",
        )
        return project, document

    def _create_project_with_pdf_document(self, suffix: str = "") -> tuple[dict, dict]:
        pdf_path = self.tmp_path / f"sample{suffix}.pdf"
        pdf = fitz.open()
        for label in ("5.7 整车用电器与线束匹配评估表", "测试分组 / Test Groups"):
            page = pdf.new_page()
            page.insert_text((72, 72), label, fontsize=18)
            page.insert_text((92, 150), f"Source snapshot page: {label}", fontsize=12)
            x_positions = [72, 180, 300, 420, 520]
            y_positions = [180, 225, 270, 315, 360, 405, 450]
            for x in x_positions:
                page.draw_line((x, y_positions[0]), (x, y_positions[-1]), color=(0.2, 0.4, 0.7), width=1)
            for y in y_positions:
                page.draw_line((x_positions[0], y), (x_positions[-1], y), color=(0.2, 0.4, 0.7), width=1)
            headers = ["Group", "Item", "Method", "Req"]
            for idx, header in enumerate(headers):
                page.insert_text((x_positions[idx] + 8, y_positions[0] + 20), header, fontsize=10)
            for row_index in range(1, 6):
                values = [
                    f"G{row_index}",
                    "测试分组" if row_index == 1 else f"条目 {row_index}",
                    "Visual",
                    "Pass",
                ]
                base_y = y_positions[row_index] + 20
                for col_index, value in enumerate(values):
                    page.insert_text((x_positions[col_index] + 8, base_y), value, fontsize=10)
        pdf.save(pdf_path)
        pdf.close()
        project = db.create_project(name=f"Workspace PDF Test{suffix}", customer_name="Customer")
        document = db.add_sor_document(
            project_id=project["id"],
            file_name=f"sample{suffix}.pdf",
            file_type="pdf",
            file_path=str(pdf_path),
            extracted_text="sample pdf",
        )
        return project, document

    def _add_knowledge_document(self, project_id: str, suffix: str = "") -> dict:
        knowledge_path = self.tmp_path / f"knowledge{suffix}.txt"
        knowledge_path.write_text("knowledge", encoding="utf-8")
        return db.add_knowledge_document(
            project_id=project_id,
            file_name=f"knowledge{suffix}.txt",
            file_type="txt",
            file_path=str(knowledge_path),
            extracted_text="knowledge",
        )

    def _version_meta_values(self, workbook) -> dict[str, object]:
        sheet = workbook["VersionMeta"]
        return {
            str(sheet.cell(row=row, column=1).value): sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value is not None
        }

    def test_extraction_review_redirects_to_workspace_review(self) -> None:
        response = main.extraction_review_page()
        self.assertEqual(response.status_code, 307)
        self.assertEqual(response.headers.get("location"), "/workspace?stage=review")

    def test_review_entry_req_codes_are_unique_per_document(self) -> None:
        project = db.create_project(name="ReqCode Test", customer_name="Customer")
        file_a = self.tmp_path / "a.txt"
        file_b = self.tmp_path / "b.txt"
        file_a.write_text("a", encoding="utf-8")
        file_b.write_text("b", encoding="utf-8")
        doc_a = db.add_sor_document(
            project_id=project["id"],
            file_name="a.txt",
            file_type="txt",
            file_path=str(file_a),
            extracted_text="a",
        )
        doc_b = db.add_sor_document(
            project_id=project["id"],
            file_name="b.txt",
            file_type="txt",
            file_path=str(file_b),
            extracted_text="b",
        )
        payload = {
            "entry_id": "REQ0001",
            "req_code": "REQ0001",
            "entry_kind": "REQUIREMENT",
            "source_page": 1,
            "source_anchor": "Main > 1",
            "requirement_text": "同一项目下的两个文档都可以复用相同 req_code。",
            "display_text": "同一项目下的两个文档都可以复用相同 req_code。",
            "raw_text": "同一项目下的两个文档都可以复用相同 req_code。",
            "requirement_type": "FUNCTIONAL",
            "priority": "MEDIUM",
            "keep": True,
        }

        res_a = main.save_review_entry(
            project["id"],
            doc_a["id"],
            main.ReviewedRequirementEntry(**payload),
        )
        res_b = main.save_review_entry(
            project["id"],
            doc_b["id"],
            main.ReviewedRequirementEntry(**payload),
        )

        self.assertEqual(res_a["review_save"]["action"], "inserted")
        self.assertEqual(res_b["review_save"]["action"], "inserted")
        self.assertEqual(res_a["saved_entry"]["req_code"], "REQ0001")
        self.assertEqual(res_b["saved_entry"]["req_code"], "REQ0001")
        rows_a = db.list_requirements(project["id"], doc_a["id"])
        rows_b = db.list_requirements(project["id"], doc_b["id"])
        self.assertEqual([row["req_code"] for row in rows_a], ["REQ0001"])
        self.assertEqual([row["req_code"] for row in rows_b], ["REQ0001"])

    def test_workspace_include_preview_returns_review_entries(self) -> None:
        project, document = self._create_project_with_document("-preview")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            analyze = main.analyze_document(
                project["id"],
                main.AnalyzeRequest(document_id=document["id"]),
            )
            workspace = main.get_document_workspace(
                project["id"],
                document["id"],
                include_preview=True,
            )

        self.assertEqual(analyze["requirements_generated"], 1)
        self.assertEqual(workspace["document"]["workflow_stage"], "EXTRACTED")
        self.assertEqual(workspace["document"]["requirement_set_version"], 1)
        self.assertEqual(len(workspace["requirements"]), 1)
        self.assertEqual(len(workspace["review_entries"]), 1)
        self.assertEqual(len(workspace["review_entries_raw"]), 2)
        self.assertTrue(workspace["review_entries"][0]["keep"])
        self.assertFalse(workspace["review_entries_raw"][0]["keep"])
        self.assertEqual(workspace["summary"]["order_source"], "review_overlay")
        self.assertEqual(workspace["summary"]["order_mismatch_count"], 0)
        self.assertEqual(workspace["summary"]["preview_status"], "READY")
        self.assertTrue(workspace["summary"]["preview_cache_hit"])

    def test_workspace_without_preview_does_not_rebuild_preview(self) -> None:
        project, document = self._create_project_with_document("-no-preview-build")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        with patch.object(main, "build_extraction_preview", side_effect=AssertionError("workspace rebuilt preview")):
            workspace = main.get_document_workspace(project["id"], document["id"], include_preview=False)

        self.assertEqual(len(workspace["requirements"]), 1)
        self.assertEqual(workspace["summary"]["preview_status"], "READY")
        self.assertTrue(workspace["summary"]["preview_cache_hit"])

    def test_extraction_preview_reads_cached_preview(self) -> None:
        project, document = self._create_project_with_document("-cached-preview")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        with patch.object(main, "build_extraction_preview", side_effect=AssertionError("preview rebuilt")):
            payload = main.get_sor_extraction_preview(project["id"], document["id"], response=main.Response())

        self.assertEqual(payload["preview_status"], "READY")
        self.assertTrue(payload["preview_cache_hit"])
        self.assertEqual(len(payload["review_entries"]), 1)
        self.assertEqual(len(payload["review_entries_raw"]), 2)

    def test_review_entry_roundtrip_updates_workspace_and_version(self) -> None:
        project, document = self._create_project_with_document("-roundtrip")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            analyze = main.analyze_document(
                project["id"],
                main.AnalyzeRequest(document_id=document["id"]),
            )
            self.assertEqual(analyze["requirements_generated"], 1)

        insert_heading = main.save_review_entry(
            project["id"],
            document["id"],
            main.ReviewedRequirementEntry(
                entry_id="H-1",
                req_code="",
                entry_kind="HEADING",
                source_page=1,
                source_anchor="Main > 1",
                requirement_text="1 概述",
                display_text="1 概述",
                raw_text="1 概述",
                requirement_type="FUNCTIONAL",
                priority="MEDIUM",
                text_relation="exact_match",
                is_deduped=True,
                parent_section_code="1",
                parent_section_title="概述",
                parent_section_label="1 概述",
                keep=True,
            ),
        )
        insert_payload = insert_heading
        self.assertEqual(insert_payload["review_save"]["action"], "inserted")
        self.assertEqual(insert_payload["review_save"]["req_code"], "H-1")
        self.assertEqual(insert_payload["saved_entry"]["req_code"], "H-1")
        self.assertEqual(insert_payload["saved_entry"]["keep"], True)
        self.assertEqual(insert_payload["document"]["requirement_set_version"], 2)
        self.assertEqual(len(insert_payload["requirements"]), 2)

        update_heading = main.save_review_entry(
            project["id"],
            document["id"],
            main.ReviewedRequirementEntry(
                entry_id="H-1",
                req_code="",
                entry_kind="HYBRID",
                source_page=1,
                source_anchor="Main > 1",
                requirement_text="1 概述",
                display_text="1 概述",
                raw_text="1 概述",
                requirement_type="FUNCTIONAL",
                priority="MEDIUM",
                text_relation="exact_match",
                is_deduped=True,
                parent_section_code="1",
                parent_section_title="概述",
                parent_section_label="1 概述",
                keep=True,
            ),
        )
        update_payload = update_heading
        self.assertEqual(update_payload["review_save"]["action"], "updated")
        self.assertEqual(update_payload["saved_entry"]["entry_kind"], "HYBRID")
        self.assertEqual(update_payload["saved_entry"]["req_code"], "H-1")
        self.assertEqual(update_payload["document"]["requirement_set_version"], 3)

        remove_heading = main.save_review_entry(
            project["id"],
            document["id"],
            main.ReviewedRequirementEntry(
                entry_id="H-1",
                req_code="",
                entry_kind="HEADING",
                source_page=1,
                source_anchor="Main > 1",
                requirement_text="1 概述",
                display_text="1 概述",
                raw_text="1 概述",
                requirement_type="FUNCTIONAL",
                priority="MEDIUM",
                text_relation="exact_match",
                is_deduped=True,
                parent_section_code="1",
                parent_section_title="概述",
                parent_section_label="1 概述",
                keep=False,
            ),
        )
        remove_payload = remove_heading
        self.assertEqual(remove_payload["review_save"]["action"], "removed")
        self.assertEqual(remove_payload["review_save"]["req_code"], "H-1")
        self.assertEqual(remove_payload["saved_entry"]["keep"], False)
        self.assertEqual(remove_payload["document"]["requirement_set_version"], 4)
        self.assertEqual(len(remove_payload["requirements"]), 1)

    def test_review_save_uses_cached_preview_and_does_not_rebuild(self) -> None:
        project, document = self._create_project_with_document("-review-cache")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        with patch.object(main, "build_extraction_preview", side_effect=AssertionError("review save rebuilt preview")):
            saved = main.save_review_entry(
                project["id"],
                document["id"],
                main.ReviewedRequirementEntry(
                    entry_id="REQ0001",
                    req_code="REQ0001",
                    entry_kind="REQUIREMENT",
                    source_page=1,
                    source_page_end=1,
                    source_anchor="Main > 1",
                    requirement_text="供应商需要满足总成试验要求（缓存更新）。",
                    display_text="供应商需要满足总成试验要求（缓存更新）。",
                    raw_text="供应商需要满足总成试验要求（缓存更新）。",
                    requirement_type="FUNCTIONAL",
                    priority="MEDIUM",
                    parent_section_code="1",
                    parent_section_title="概述",
                    parent_section_label="1 概述",
                    keep=True,
                ),
            )

        self.assertEqual(saved["document"]["workflow_stage"], "REVIEWING")
        self.assertEqual(saved["summary"]["preview_status"], "READY")
        self.assertTrue(saved["summary"]["preview_cache_hit"])
        self.assertEqual(saved["saved_entry"]["raw_text"], "供应商需要满足总成试验要求（缓存更新）。")

    def test_regenerate_requirement_set_clears_lock_and_invalidates_ai(self) -> None:
        project, document = self._create_project_with_document("-regen")
        preview = _fake_preview()
        with patch.object(main, "build_extraction_preview", return_value=preview):
            first = main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))
        self.assertEqual(first["requirement_set_version"], 1)

        locked = main.lock_document_extraction(project["id"], document["id"])
        self.assertEqual(locked["document"]["workflow_stage"], "EXTRACTION_LOCKED")

        db.update_sor_document_ai_state(
            document["id"],
            ai_status="READY",
            ai_assessed_at="2026-03-08T00:00:00+00:00",
            ai_requirement_set_version=1,
            ai_knowledge_base_version=0,
            workflow_stage="AI_READY",
        )

        with patch.object(main, "build_extraction_preview", return_value=preview):
            second = main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        current = db.get_sor_document(document["id"]) or {}
        self.assertEqual(second["requirement_set_version"], 2)
        self.assertEqual(current.get("workflow_stage"), "EXTRACTED")
        self.assertEqual(str(current.get("extraction_locked_at") or ""), "")
        self.assertEqual(str(current.get("ai_status") or ""), "STALE")

    def test_review_save_auto_unlocks_and_disables_export(self) -> None:
        project, document = self._create_project_with_document("-auto-unlock")
        preview = _fake_preview()
        with patch.object(main, "build_extraction_preview", return_value=preview):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))
            main.lock_document_extraction(project["id"], document["id"])
            saved = main.save_review_entry(
                project["id"],
                document["id"],
                main.ReviewedRequirementEntry(
                    entry_id="REQ0001",
                    req_code="REQ0001",
                    entry_kind="REQUIREMENT",
                    source_page=1,
                    source_page_end=1,
                    source_anchor="Main > 1",
                    requirement_text="供应商需要满足总成试验要求（更新）。",
                    display_text="供应商需要满足总成试验要求（更新）。",
                    raw_text="供应商需要满足总成试验要求（更新）。",
                    requirement_type="FUNCTIONAL",
                    priority="MEDIUM",
                    parent_section_code="1",
                    parent_section_title="概述",
                    parent_section_label="1 概述",
                    keep=True,
                ),
            )

        self.assertEqual(saved["document"]["workflow_stage"], "REVIEWING")
        self.assertFalse(saved["summary"]["locked"])
        self.assertFalse(saved["actions"]["can_export"])
        self.assertEqual(saved["saved_entry"]["raw_text"], "供应商需要满足总成试验要求（更新）。")

    def test_lock_and_ai_gate_behavior(self) -> None:
        project, empty_document = self._create_project_with_document("-empty")
        with self.assertRaises(HTTPException) as lock_exc:
            main.lock_document_extraction(project["id"], empty_document["id"])
        self.assertEqual(lock_exc.exception.status_code, 409)

        project, document = self._create_project_with_document("-gate")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        workspace_before_lock = main.get_document_workspace(project["id"], document["id"], include_preview=False)
        self.assertFalse(workspace_before_lock["actions"]["can_export"])

        with self.assertRaises(HTTPException) as ai_before_lock_exc:
            main.run_document_ai_analysis(project["id"], document["id"])
        self.assertEqual(ai_before_lock_exc.exception.status_code, 409)

        workspace_locked = main.lock_document_extraction(project["id"], document["id"])
        self.assertTrue(workspace_locked["summary"]["locked"])
        self.assertTrue(workspace_locked["actions"]["can_export"])
        self.assertTrue(workspace_locked["actions"]["can_run_ai_analysis"])

        ai_without_kb = main.run_document_ai_analysis(project["id"], document["id"])
        self.assertEqual(ai_without_kb["ai_processed_requirements"], 0)
        self.assertGreater(ai_without_kb["ai_screened_requirements"], 0)
        self.assertEqual(ai_without_kb["knowledge_docs_used"], 0)
        self.assertGreater(ai_without_kb["summary"]["ai_eligible_count"], 0)
        self.assertGreater(ai_without_kb["summary"]["assessment_target_count"], 0)

        self._add_knowledge_document(project["id"], "-gate")
        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=False)
        self.assertTrue(workspace["actions"]["can_run_ai_analysis"])

    def test_workspace_summary_counts_runnable_profiles_from_ready_flag(self) -> None:
        project, document = self._create_project_with_document("-ready-flag")
        requirements = [
            {
                "id": "req-ready-1",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "供应商需要满足总成试验要求。",
                "raw_text": "供应商需要满足总成试验要求。",
                "requirement_text": "供应商需要满足总成试验要求。",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
                "source_page": 1,
                "source_anchor": "Main > 1",
                "parent_section_code": "1",
                "parent_section_title": "概述",
                "parent_section_label": "1 概述",
            }
        ]
        locked_document = {
            **document,
            "workflow_stage": "EXTRACTION_LOCKED",
            "requirement_set_version": 1,
            "extraction_locked_at": "2026-04-01T00:00:00+00:00",
        }

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_RETRIEVAL_BACKEND": "local",
            },
            clear=True,
        ), patch.object(
            main,
            "_build_ai_profile_catalog",
            return_value=[
                {
                    "id": "aliyun",
                    "provider": "aliyun",
                    "label": "Aliyun",
                    "assessment_mode": "llm_shadow",
                    "retrieval_backend": "qdrant",
                    "ready": True,
                    "provider_ready": False,
                    "llm_ready": True,
                    "retrieval_ready": True,
                    "missing": [],
                    "knowledge_sync": {"status": "NOT_REQUIRED", "ready_count": 0, "total_count": 0},
                }
            ],
        ):
            workspace = main._build_document_workspace(
                project=project,
                document=locked_document,
                requirements=requirements,
                knowledge_docs_count=0,
                preview=None,
                preview_meta={},
            )

        self.assertEqual(workspace["summary"]["runnable_ai_profile_count"], 1)
        self.assertTrue(workspace["actions"]["can_run_ai_analysis"])

    def test_refresh_preview_endpoint_builds_and_persists_cache(self) -> None:
        project, document = self._create_project_with_document("-refresh-preview")
        cache_path = preview_cache.preview_cache_path(document["id"])
        if cache_path.exists():
            cache_path.unlink()

        with patch.object(preview_cache._PREVIEW_EXECUTOR, "submit", return_value=None):
            response = main.refresh_document_preview(project["id"], document["id"])
        self.assertEqual(response["preview_status"], "BUILDING")

        with patch.object(preview_cache, "build_extraction_preview", return_value=_fake_preview()):
            preview_cache._run_preview_refresh(document["id"], document["file_path"])

        refreshed = main.get_sor_extraction_preview(project["id"], document["id"], response=main.Response())
        self.assertEqual(refreshed["preview_status"], "READY")
        self.assertTrue(refreshed["preview_cache_hit"])
        self.assertTrue(cache_path.exists())

    def test_load_cached_preview_marks_legacy_cache_as_stale(self) -> None:
        project, document = self._create_project_with_document("-legacy-preview-cache")
        cache_path = preview_cache.preview_cache_path(document["id"])
        cache_path.write_text(json.dumps(_fake_preview(), ensure_ascii=False), encoding="utf-8")
        db.update_sor_document_preview_state(
            document["id"],
            preview_status="READY",
            preview_generated_at="2026-03-24T10:00:00+00:00",
            preview_file_mtime=preview_cache._source_mtime_token(document["file_path"]),
            preview_error="",
        )

        preview, meta = preview_cache.load_cached_preview(db.get_sor_document(document["id"]))

        self.assertIsNone(preview)
        self.assertEqual(meta["preview_status"], "STALE")
        self.assertFalse(meta["preview_cache_hit"])

    def test_load_cached_preview_marks_old_version_cache_as_stale(self) -> None:
        project, document = self._create_project_with_document("-old-version-preview-cache")
        cache_path = preview_cache.preview_cache_path(document["id"])
        payload = _fake_preview()
        payload["_preview_cache_version"] = "preview-cache-v3"
        cache_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        db.update_sor_document_preview_state(
            document["id"],
            preview_status="READY",
            preview_generated_at="2026-03-24T10:00:00+00:00",
            preview_file_mtime=preview_cache._source_mtime_token(document["file_path"]),
            preview_error="",
        )

        preview, meta = preview_cache.load_cached_preview(db.get_sor_document(document["id"]))

        self.assertIsNone(preview)
        self.assertEqual(meta["preview_status"], "STALE")
        self.assertFalse(meta["preview_cache_hit"])
        self.assertEqual(meta["preview_error"], "cached preview is outdated")

    def test_load_cached_preview_marks_old_build_token_as_stale(self) -> None:
        project, document = self._create_project_with_document("-old-build-token-preview-cache")
        cache_path = preview_cache.preview_cache_path(document["id"])
        payload = _fake_preview()
        payload["_preview_cache_version"] = preview_cache._PREVIEW_CACHE_VERSION
        payload["_preview_build_token"] = "preview-cache-v28-numbered-heading-anchor-fragments|source-crop-v56"
        cache_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        db.update_sor_document_preview_state(
            document["id"],
            preview_status="READY",
            preview_generated_at="2026-03-24T10:00:00+00:00",
            preview_file_mtime=preview_cache._source_mtime_token(document["file_path"]),
            preview_error="",
        )

        preview, meta = preview_cache.load_cached_preview(db.get_sor_document(document["id"]))

        self.assertIsNone(preview)
        self.assertEqual(meta["preview_status"], "STALE")
        self.assertFalse(meta["preview_cache_hit"])
        self.assertEqual(meta["preview_error"], "cached preview is outdated")

    def test_preview_refresh_purges_stale_snapshot_root_before_rebuild(self) -> None:
        project, document = self._create_project_with_document("-refresh-purge")
        cache_path = preview_cache.preview_cache_path(document["id"])
        cache_path.write_text("{}", encoding="utf-8")
        snapshot_root = preview_cache.document_snapshot_root(document["id"])
        stale_file = snapshot_root / "stale.png"
        stale_file.write_text("old", encoding="utf-8")

        def _fake_build(file_path: str, *, document_id: str = "", snapshot_root: Path | None = None) -> dict:
            self.assertEqual(document_id, document["id"])
            self.assertIsNotNone(snapshot_root)
            self.assertFalse(cache_path.exists())
            self.assertFalse(stale_file.exists())
            self.assertTrue(Path(snapshot_root or "").exists())
            return _fake_preview()

        with patch.object(preview_cache, "build_extraction_preview", side_effect=_fake_build):
            preview_cache._run_preview_refresh(document["id"], document["file_path"])

        self.assertFalse(stale_file.exists())
        self.assertTrue(cache_path.exists())

    def test_sync_document_table_assets_keeps_preview_snapshot_root_intact(self) -> None:
        project, document = self._create_project_with_document("-isolated-snapshots")
        preview_root = preview_cache.document_snapshot_root(document["id"])
        persisted_root = main._table_snapshot_root(document["id"])
        preview_file = preview_root / "preview.png"
        persisted_file = persisted_root / "persisted.png"
        preview_file.write_text("preview", encoding="utf-8")
        persisted_file.write_text("persisted", encoding="utf-8")

        with patch.object(main, "_materialize_document_table_assets", return_value=([], [])):
            main._sync_document_table_assets(project["id"], document["id"])

        self.assertTrue(preview_file.exists())
        self.assertFalse(persisted_file.exists())

    def test_workspace_reorders_matrix_display_ids_by_preview_order(self) -> None:
        project, document = self._create_project_with_document("-ordering")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0023",
            source_page=5,
            source_anchor="Main > 5.7",
            requirement_text="5.7 次条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0024",
            source_page=5,
            source_anchor="Main > 5.1",
            requirement_text="5.1 首条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")

        preview_cache.persist_preview_cache(
            db.get_sor_document(document["id"]) or {},
            _fake_preview_with_reversed_req_codes(),
        )
        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=True)

        self.assertEqual([row["req_code"] for row in workspace["requirements"]], ["REQ0024", "REQ0023"])
        self.assertEqual([row["display_req_code"] for row in workspace["requirements"]], ["REQ0001", "REQ0002"])
        self.assertEqual([row["display_sort_index"] for row in workspace["requirements"]], [1, 2])
        self.assertEqual([row["source_order"] for row in workspace["requirements"]], [1, 2])
        self.assertEqual(workspace["summary"]["order_source"], "review_overlay")
        self.assertEqual(workspace["summary"]["order_mismatch_count"], 0)

    def test_workspace_uses_cached_preview_order_even_without_preview_payload(self) -> None:
        project, document = self._create_project_with_document("-ordering-no-preview-payload")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0023",
            source_page=5,
            source_anchor="Main > 5.7",
            requirement_text="5.7 次条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0024",
            source_page=5,
            source_anchor="Main > 5.1",
            requirement_text="5.1 首条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")

        preview_cache.persist_preview_cache(
            db.get_sor_document(document["id"]) or {},
            _fake_preview_with_reversed_req_codes(),
        )
        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=False)

        self.assertEqual([row["req_code"] for row in workspace["requirements"]], ["REQ0024", "REQ0023"])
        self.assertEqual(workspace["summary"]["order_source"], "review_overlay")
        self.assertEqual(workspace["summary"]["order_mismatch_count"], 0)
        self.assertNotIn("review_entries", workspace)
        self.assertNotIn("preview", workspace)

    def test_workspace_order_prefers_source_signature_over_stale_req_code(self) -> None:
        project, document = self._create_project_with_document("-ordering-stale-req")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=5,
            source_anchor="Main > 5.1",
            requirement_text="5.1 第一条需求",
            raw_text="5.1 第一条需求",
            display_text="5.1 第一条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0002",
            source_page=5,
            source_anchor="Main > 5.2",
            requirement_text="5.2 第二条需求",
            raw_text="5.2 第二条需求",
            display_text="5.2 第二条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")

        preview_cache.persist_preview_cache(
            db.get_sor_document(document["id"]) or {},
            _fake_preview_with_reused_req_codes(),
        )
        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=True)

        self.assertEqual(
            [str(row.get("display_text") or row.get("requirement_text") or "") for row in workspace["requirements"]],
            ["5.2 第二条需求", "5.1 第一条需求"],
        )
        self.assertEqual([row["display_req_code"] for row in workspace["requirements"]], ["REQ0001", "REQ0002"])
        self.assertEqual(workspace["summary"]["order_source"], "review_overlay")
        self.assertEqual(workspace["summary"]["order_mismatch_count"], 0)

    def test_workspace_keeps_same_code_local_clause_in_preview_order(self) -> None:
        project, document = self._create_project_with_document("-same-code-local-clause")
        preview = _fake_preview_with_same_code_local_clause_in_section_six()
        for number in range(10, 0, -1):
            preview_row = preview["complete_requirement_entries"][number]
            db.insert_requirement(
                project_id=project["id"],
                sor_document_id=document["id"],
                req_code=str(preview_row["req_code"]),
                source_page=9,
                source_anchor=str(preview_row["source_anchor"]),
                requirement_text=str(preview_row["display_text"]),
                raw_text=str(preview_row["raw_text"]),
                display_text=str(preview_row["display_text"]),
                requirement_type="FUNCTIONAL",
                priority="MEDIUM",
                parent_section_code=str(preview_row["parent_section_code"]),
                parent_section_title=str(preview_row["parent_section_title"]),
                parent_section_label=str(preview_row["parent_section_label"]),
            )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        preview_cache.persist_preview_cache(db.get_sor_document(document["id"]) or {}, preview)

        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=True)

        self.assertEqual(
            [row["section_key"] for row in workspace["requirements"]],
            [f"6::{number}" for number in range(1, 11)],
        )
        self.assertEqual(
            [row["owner_section_label"] for row in workspace["requirements"]],
            [f"{number}）" for number in range(1, 11)],
        )
        self.assertEqual(workspace["requirements"][5]["section_key"], "6::6")
        self.assertEqual(workspace["requirements"][5]["owner_section_label"], "6）")
        self.assertEqual(workspace["summary"]["order_source"], "review_overlay")
        self.assertEqual(workspace["summary"]["order_mismatch_count"], 0)

    def test_workspace_review_entries_preserve_section_six_subtree_order_and_owner_labels(self) -> None:
        project, document = self._create_project_with_document("-review-same-code-local-clause")
        preview = _fake_preview_with_same_code_local_clause_in_section_six()
        for number in range(10, 0, -1):
            preview_row = preview["complete_requirement_entries"][number]
            db.insert_requirement(
                project_id=project["id"],
                sor_document_id=document["id"],
                req_code=str(preview_row["req_code"]),
                source_page=9,
                source_anchor=str(preview_row["source_anchor"]),
                requirement_text=str(preview_row["display_text"]),
                raw_text=str(preview_row["raw_text"]),
                display_text=str(preview_row["display_text"]),
                requirement_type="FUNCTIONAL",
                priority="MEDIUM",
                parent_section_code=str(preview_row["parent_section_code"]),
                parent_section_title=str(preview_row["parent_section_title"]),
                parent_section_label=str(preview_row["parent_section_label"]),
            )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        preview_cache.persist_preview_cache(db.get_sor_document(document["id"]) or {}, preview)

        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=True)

        self.assertEqual(
            [row["section_key"] for row in workspace["review_entries"]],
            ["6", *[f"6::{number}" for number in range(1, 11)]],
        )
        self.assertEqual(workspace["review_entries"][0]["owner_section_label"], "6")
        self.assertEqual(workspace["review_entries"][6]["section_key"], "6::6")
        self.assertEqual(workspace["review_entries"][6]["owner_section_label"], "6）")
        self.assertEqual(workspace["review_entries"][6]["structural_parent_section_label"], "6")

    def test_workspace_matrix_rows_prefer_current_review_text_over_stale_persisted_text(self) -> None:
        project, document = self._create_project_with_document("-matrix-preview")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 1",
            requirement_text="旧的错误内容",
            raw_text="旧的错误内容",
            display_text="旧的错误内容",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")

        preview = {
            "summary": {"dedupe_version": "heading_requirement_v1"},
            "sections": [],
            "complete_requirement_entries": [
                {
                    "entry_id": "REQ0001",
                    "req_code": "REQ0001",
                    "entry_kind": "REQUIREMENT",
                    "display_text": "新的正确内容",
                    "raw_text": "新的正确内容",
                    "requirement_text": "新的正确内容",
                    "source_page": 1,
                    "source_page_end": 1,
                    "source_anchor": "Page 1 | Section 1",
                    "parent_section_code": "1",
                    "parent_section_title": "概述",
                    "parent_section_label": "1 概述",
                }
            ],
            "complete_requirement_entries_raw": [],
        }

        workspace = main._build_document_workspace(
            project=project,
            document=db.get_sor_document(document["id"]) or {},
            requirements=db.list_requirements(project["id"], document["id"]),
            knowledge_docs_count=0,
            preview=preview,
        )

        self.assertEqual(len(workspace["requirements"]), 1)
        self.assertEqual(workspace["requirements"][0]["raw_text"], "新的正确内容")
        self.assertEqual(workspace["requirements"][0]["req_code"], "REQ0001")

    def test_overlay_review_entries_skips_stale_req_code_match_when_preview_fingerprint_changed(self) -> None:
        preview_rows = [
            {
                "entry_id": "REQ0001",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "新的正确内容",
                "raw_text": "新的正确内容",
                "requirement_text": "新的正确内容",
                "source_page": 11,
                "source_page_end": 11,
                "source_anchor": "Page 11 | Heading 软件开发要求",
                "parent_section_code": "",
                "parent_section_title": "软件开发要求",
                "parent_section_label": "软件开发要求",
            }
        ]
        persisted_rows = [
            {
                "id": "persisted-1",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "旧的错误内容",
                "raw_text": "旧的错误内容",
                "requirement_text": "旧的错误内容",
                "source_page": 12,
                "source_page_end": 12,
                "source_anchor": "Page 12 | Section 1.1 FMEA编制与更新",
                "parent_section_code": "1.1",
                "parent_section_title": "FMEA编制与更新",
                "parent_section_label": "1.1 FMEA编制与更新",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
            }
        ]

        rows = main._overlay_review_entries(preview_rows, persisted_rows, prefer_raw=False)

        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0]["keep"])
        self.assertEqual(rows[0]["parent_section_label"], "软件开发要求")
        self.assertEqual(rows[0]["requirement_text"], "新的正确内容")

    def test_overlay_review_entries_matches_same_source_when_req_code_changed(self) -> None:
        preview_rows = [
            {
                "entry_id": "REQ0999",
                "req_code": "REQ0999",
                "entry_kind": "REQUIREMENT",
                "display_text": "新的正确内容",
                "raw_text": "新的正确内容",
                "requirement_text": "新的正确内容",
                "source_page": 11,
                "source_page_end": 11,
                "source_anchor": "Page 11 | Heading 软件开发要求",
                "parent_section_code": "",
                "parent_section_title": "软件开发要求",
                "parent_section_label": "软件开发要求",
            }
        ]
        persisted_rows = [
            {
                "id": "persisted-1",
                "req_code": "REQ0001",
                "entry_kind": "REQUIREMENT",
                "display_text": "旧的内容",
                "raw_text": "旧的内容",
                "requirement_text": "旧的内容",
                "source_page": 11,
                "source_page_end": 11,
                "source_anchor": "Page 11 | Heading 软件开发要求",
                "parent_section_code": "",
                "parent_section_title": "软件开发要求",
                "parent_section_label": "软件开发要求",
                "status": "DRAFT",
            }
        ]

        rows = main._overlay_review_entries(preview_rows, persisted_rows, prefer_raw=True)

        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["keep"])
        self.assertEqual(rows[0]["id"], "persisted-1")
        self.assertEqual(rows[0]["req_code"], "REQ0999")
        self.assertEqual(rows[0]["raw_text"], "新的正确内容")

    def test_overlay_review_entries_keeps_review_state_for_same_preview_fingerprint(self) -> None:
        preview_rows = [
            {
                "entry_id": "REQ0002",
                "req_code": "REQ0002",
                "entry_kind": "REQUIREMENT",
                "display_text": "供应商需要满足总成试验要求。",
                "raw_text": "供应商需要满足总成试验要求。",
                "requirement_text": "供应商需要满足总成试验要求。",
                "source_page": 1,
                "source_page_end": 1,
                "source_anchor": "Main > 1",
                "parent_section_code": "1",
                "parent_section_title": "概述",
                "parent_section_label": "1 概述",
            }
        ]
        persisted_rows = [
            {
                "id": "persisted-2",
                "req_code": "REQ0002",
                "entry_kind": "REQUIREMENT",
                "display_text": "供应商需要满足总成试验要求。",
                "raw_text": "供应商需要满足总成试验要求。",
                "requirement_text": "供应商需要满足总成试验要求。",
                "source_page": 1,
                "source_page_end": 1,
                "source_anchor": "Main > 1",
                "parent_section_code": "1",
                "parent_section_title": "概述",
                "parent_section_label": "1 概述",
                "requirement_type": "FUNCTIONAL",
                "priority": "HIGH",
                "status": "reviewed",
            }
        ]

        rows = main._overlay_review_entries(preview_rows, persisted_rows, prefer_raw=False)

        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["keep"])
        self.assertEqual(rows[0]["priority"], "HIGH")
        self.assertEqual(rows[0]["parent_section_label"], "1 概述")

    def test_overlay_review_entries_keeps_preview_structure_for_matched_review_row(self) -> None:
        preview_rows = [
            {
                "entry_id": "REQ0103",
                "req_code": "REQ-0103",
                "entry_kind": "REQUIREMENT",
                "display_text": "10）商务报价应至少包含设计开发费用。",
                "raw_text": "10）商务报价应至少包含设计开发费用。",
                "requirement_text": "10）商务报价应至少包含设计开发费用。",
                "source_page": 10,
                "source_page_end": 10,
                "source_anchor": "Main > 6",
                "parent_section_code": "6",
                "parent_section_title": "投标说明",
                "parent_section_label": "6 投标说明",
            }
        ]
        persisted_rows = [
            {
                "id": "persisted-3",
                "req_code": "REQ-0103",
                "entry_kind": "REQUIREMENT",
                "display_text": "10）商务报价应至少包含设计开发费用。",
                "raw_text": "10）商务报价应至少包含设计开发费用。",
                "requirement_text": "10）商务报价应至少包含设计开发费用。",
                "source_page": 10,
                "source_page_end": 10,
                "source_anchor": "Main > 7",
                "parent_section_code": "7",
                "parent_section_title": "附件目录",
                "parent_section_label": "7 附件目录",
                "requirement_type": "FUNCTIONAL",
                "priority": "HIGH",
                "status": "reviewed",
            }
        ]

        rows = main._overlay_review_entries(preview_rows, persisted_rows, prefer_raw=False)

        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["keep"])
        self.assertEqual(rows[0]["priority"], "HIGH")
        self.assertEqual(rows[0]["source_anchor"], "Main > 6")
        self.assertEqual(rows[0]["parent_section_label"], "6 投标说明")

    def test_find_requirement_for_review_entry_prefers_source_match_over_stale_req_code(self) -> None:
        requirements = [
            {
                "id": "persisted-a",
                "req_code": "REQ0001",
                "source_page": 11,
                "source_anchor": "Page 11 | Section 3.1",
                "raw_text": "3.1 旧条目",
                "display_text": "3.1 旧条目",
                "requirement_text": "3.1 旧条目",
            },
            {
                "id": "persisted-b",
                "req_code": "REQ0002",
                "source_page": 12,
                "source_anchor": "Page 12 | Section 3.2",
                "raw_text": "3.2 正确条目",
                "display_text": "3.2 正确条目",
                "requirement_text": "3.2 正确条目",
            },
        ]
        normalized = {
            "entry_id": "REQ0001",
            "req_code": "REQ0001",
            "source_page": 12,
            "source_anchor": "Page 12 | Section 3.2",
            "raw_text": "3.2 正确条目",
            "display_text": "3.2 正确条目",
            "requirement_text": "3.2 正确条目",
        }

        matched = main._find_requirement_for_review_entry(requirements, normalized)

        self.assertIsNotNone(matched)
        self.assertEqual(matched["id"], "persisted-b")

    def test_enrich_requirements_with_preview_context_matches_by_source_signature(self) -> None:
        project, document = self._create_project_with_document("-preview-context-source-match")
        rows = [
            {
                "id": "persisted-b",
                "project_id": project["id"],
                "sor_document_id": document["id"],
                "req_code": "REQ0002",
                "source_page": 5,
                "source_page_end": 5,
                "source_anchor": "Main > 5.2",
                "raw_text": "5.2 第二条需求",
                "display_text": "5.2 第二条需求",
                "requirement_text": "5.2 第二条需求",
                "parent_section_code": "",
                "parent_section_title": "",
                "parent_section_label": "",
            }
        ]

        with patch.object(main, "load_cached_preview", return_value=(_fake_preview_with_reused_req_codes(), {})):
            enriched = main._enrich_requirements_with_preview_context(rows, str(document["file_path"]))

        self.assertEqual(enriched[0]["parent_section_code"], "5.2")
        self.assertEqual(enriched[0]["parent_section_label"], "5.2 第二条需求")

    def test_materialize_document_table_assets_uses_preview_source_match_not_req_code(self) -> None:
        project, document = self._create_project_with_document("-materialize-source-match")
        rows = [
            {
                "id": "persisted-b",
                "project_id": project["id"],
                "sor_document_id": document["id"],
                "req_code": "REQ0002",
                "source_page": 0,
                "source_page_end": 0,
                "source_anchor": "Main > 5.2",
                "raw_text": "5.2 第二条需求",
                "display_text": "5.2 第二条需求",
                "requirement_text": "5.2 第二条需求",
                "table_refs": [],
                "content_blocks": [],
                "has_table_asset": False,
                "table_render_mode": "",
            }
        ]

        def _fake_enrich(items, **kwargs):
            return items, []

        with patch.object(main, "load_cached_preview", return_value=(_fake_preview_with_reused_req_codes(), {})):
            with patch.object(main, "enrich_entries_with_table_assets", side_effect=_fake_enrich):
                materialized, assets = main._materialize_document_table_assets(rows, sor_document_id=document["id"])

        self.assertEqual(assets, [])
        self.assertEqual(materialized[0]["source_page"], 5)
        self.assertEqual(materialized[0]["source_page_end"], 5)

    def test_workspace_exposes_table_refs_and_snapshot_urls(self) -> None:
        project, document = self._create_project_with_document("-table-assets")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview_with_table_markdown()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))
            workspace = main.get_document_workspace(project["id"], document["id"], include_preview=True)

        requirement = workspace["requirements"][0]
        review_entry = workspace["review_entries"][0]
        self.assertTrue(requirement["has_table_asset"])
        self.assertEqual(requirement["table_render_mode"], "structured")
        self.assertEqual(len(requirement["table_refs"]), 1)
        self.assertEqual(requirement["table_refs"][0]["table_id"], review_entry["table_refs"][0]["table_id"])
        self.assertEqual(requirement["table_refs"][0]["row_range"], [1, 2])
        self.assertIn("/snapshot", requirement["table_refs"][0]["slice_snapshot_url"])
        self.assertEqual(requirement["table_refs"][0]["recommended_export_variant"], "slice")
        self.assertEqual(requirement["table_refs"][0]["export_variant"], "slice")
        self.assertEqual(requirement["table_refs"][0]["export_variant_reasons"], [])
        self.assertEqual(
            [block["kind"] for block in requirement["content_blocks"]],
            ["paragraph", "table"],
        )

    def test_table_ref_recommendation_prefers_source_for_snapshot_only_with_images(self) -> None:
        ref = table_assets._normalize_table_ref(
            {
                "table_id": "tbl_complex",
                "slice_render_mode": "snapshot_only",
                "source_snapshot_path": str(self.tmp_path / "source.png"),
                "source_page": 5,
                "source_page_end": 6,
                "source_image_block_count": 1,
                "inherited_headers": ["项目", "要求"],
                "slice_rows": [["测试分组", "含图片和复杂布局"]],
                "inherited_notes": ["跨页续表"],
            }
        )

        self.assertEqual(ref["recommended_export_variant"], "source")
        self.assertEqual(ref["export_variant"], "source")
        self.assertIn("含图片", ref["export_variant_reasons"])
        self.assertIn("结构低置信", ref["export_variant_reasons"])

    def test_table_ref_recommendation_prefers_source_for_placeholder_symbol_table(self) -> None:
        ref = table_assets._normalize_table_ref(
            {
                "table_id": "tbl_placeholder",
                "slice_render_mode": "structured",
                "source_snapshot_path": str(self.tmp_path / "source-placeholder.png"),
                "source_page": 12,
                "source_page_end": 12,
                "inherited_headers": ["项目", "内容", "备注"],
                "slice_rows": [
                    ["安全相关关键特性", "法规相关关键特性", "功能相关重要特性"],
                    ["#REF!", "#REF!", "#REF!"],
                ],
                "inherited_notes": ["2.2 特殊特性分类与标识"],
            }
        )

        self.assertEqual(ref["recommended_export_variant"], "source")
        self.assertEqual(ref["export_variant"], "source")
        self.assertIn("切片可读性差", ref["export_variant_reasons"])

    def test_table_asset_api_returns_asset_and_snapshot(self) -> None:
        project, document = self._create_project_with_document("-table-api")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview_with_table_markdown()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=False)
        table_id = workspace["requirements"][0]["table_refs"][0]["table_id"]
        asset = main.get_document_table_asset(project["id"], document["id"], table_id)
        snapshot = main.get_document_table_asset_snapshot(project["id"], document["id"], table_id)

        self.assertEqual(asset["table_id"], table_id)
        self.assertEqual(asset["headers"], ["Item", "Condition"])
        self.assertEqual(asset["rows"], [["Voltage", "12V"], ["Current", "3A"]])
        self.assertIn("/snapshot", asset["snapshot_url"])
        self.assertTrue(Path(snapshot.path).exists())
        self.assertEqual(snapshot.headers.get("cache-control"), "no-store, max-age=0")
        self.assertEqual(snapshot.headers.get("pragma"), "no-cache")

    def test_table_asset_api_exposes_source_snapshot_for_pdf_page_span(self) -> None:
        project, document = self._create_project_with_pdf_document("-source-snapshot")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview_with_complex_table_markdown()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=False)
        table_ref = workspace["requirements"][0]["table_refs"][0]
        table_id = table_ref["table_id"]

        asset = main.get_document_table_asset(project["id"], document["id"], table_id)
        source_snapshot = main.get_document_table_asset_source_snapshot(project["id"], document["id"], table_id)

        self.assertIn("/source-snapshot", table_ref["source_snapshot_url"])
        self.assertIn("/source-snapshot", asset["source_snapshot_url"])
        self.assertTrue(Path(source_snapshot.path).exists())
        self.assertEqual(source_snapshot.headers.get("cache-control"), "no-store, max-age=0")
        self.assertEqual(source_snapshot.headers.get("pragma"), "no-cache")
        with PILImage.open(source_snapshot.path) as image:
            self.assertGreaterEqual(image.width, 900)
            self.assertGreater(image.height, image.width)
            self.assertGreater(image.height, 1500)

    def test_table_asset_api_prefers_preview_asset_over_stale_db_asset(self) -> None:
        project, document = self._create_project_with_document("-preview-asset-priority")
        stale_snapshot = self.tmp_path / "stale.png"
        stale_snapshot.write_text("old", encoding="utf-8")
        fresh_snapshot = self.tmp_path / "fresh.png"
        fresh_snapshot.write_bytes(b"fresh")
        fresh_source = self.tmp_path / "fresh_source.png"
        fresh_source.write_bytes(b"fresh-source")

        db.replace_document_table_assets(
            project["id"],
            document["id"],
            [
                {
                    "table_id": "tbl_preview_first",
                    "document_id": document["id"],
                    "headers": ["Old", "Header"],
                    "rows": [["old", "row"]],
                    "render_mode": "structured",
                    "snapshot_path": str(stale_snapshot),
                    "source_snapshot_path": str(self.tmp_path / "missing_source.png"),
                }
            ],
        )
        preview_cache.persist_preview_cache(
            db.get_sor_document(document["id"]) or document,
            {
                "summary": {"dedupe_version": "heading_requirement_v1"},
                "table_assets": [
                    {
                        "table_id": "tbl_preview_first",
                        "document_id": document["id"],
                        "headers": ["Fresh", "Header"],
                        "rows": [["fresh", "row"]],
                        "render_mode": "structured",
                        "snapshot_path": str(fresh_snapshot),
                        "source_snapshot_path": str(fresh_source),
                    }
                ],
                "complete_requirement_entries": [],
                "complete_requirement_entries_raw": [],
            },
        )

        asset = main.get_document_table_asset(project["id"], document["id"], "tbl_preview_first")
        source_snapshot = main.get_document_table_asset_source_snapshot(project["id"], document["id"], "tbl_preview_first")

        self.assertEqual(asset["headers"], ["Fresh", "Header"])
        self.assertEqual(Path(source_snapshot.path), fresh_source)

    def test_table_asset_source_snapshot_falls_back_to_persisted_asset_when_preview_source_is_missing(self) -> None:
        project, document = self._create_project_with_document("-preview-source-fallback")
        preview_snapshot = self.tmp_path / "preview.png"
        preview_snapshot.write_bytes(b"preview")
        db_snapshot = self.tmp_path / "db.png"
        db_snapshot.write_bytes(b"db")
        db_source = self.tmp_path / "db_source.png"
        db_source.write_bytes(b"db-source")

        db.replace_document_table_assets(
            project["id"],
            document["id"],
            [
                {
                    "table_id": "tbl_source_fallback",
                    "document_id": document["id"],
                    "headers": ["Persisted"],
                    "rows": [["ok"]],
                    "render_mode": "structured",
                    "snapshot_path": str(db_snapshot),
                    "source_snapshot_path": str(db_source),
                }
            ],
        )
        preview_cache.persist_preview_cache(
            db.get_sor_document(document["id"]) or document,
            {
                "summary": {"dedupe_version": "heading_requirement_v1"},
                "table_assets": [
                    {
                        "table_id": "tbl_source_fallback",
                        "document_id": document["id"],
                        "headers": ["Preview"],
                        "rows": [["broken"]],
                        "render_mode": "structured",
                        "snapshot_path": str(preview_snapshot),
                        "source_snapshot_path": str(self.tmp_path / "missing_preview_source.png"),
                    }
                ],
                "complete_requirement_entries": [],
                "complete_requirement_entries_raw": [],
            },
        )

        source_snapshot = main.get_document_table_asset_source_snapshot(project["id"], document["id"], "tbl_source_fallback")

        self.assertEqual(Path(source_snapshot.path), db_source)

    def test_review_entry_save_persists_table_export_variant_override(self) -> None:
        project, document = self._create_project_with_pdf_document("-table-choice")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview_with_complex_table_markdown()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))

        workspace = main.get_document_workspace(project["id"], document["id"], include_preview=False)
        row = workspace["requirements"][0]
        table_ref = dict(row["table_refs"][0], export_variant="slice")
        payload = main.ReviewedRequirementEntry(
            entry_id=row.get("entry_id") or row["req_code"],
            req_code=row["req_code"],
            entry_kind=row["entry_kind"],
            source_page=row["source_page"],
            source_page_end=row.get("source_page_end"),
            source_anchor=row["source_anchor"],
            requirement_text=row["requirement_text"],
            display_text=row["display_text"],
            raw_text=row["raw_text"],
            requirement_type=row["requirement_type"],
            priority=row["priority"],
            text_relation=row.get("text_relation") or "distinct",
            is_deduped=bool(row.get("is_deduped")),
            heading_text=row.get("heading_text") or "",
            parent_section_code=row.get("parent_section_code") or "",
            parent_section_title=row.get("parent_section_title") or "",
            parent_section_label=row.get("parent_section_label") or "",
            merged_from_req_codes=row.get("merged_from_req_codes") or [],
            merged_from_entry_ids=row.get("merged_from_entry_ids") or [],
            dedupe_version=row.get("dedupe_version") or "",
            table_refs=[table_ref],
            content_blocks=row.get("content_blocks") or [],
            has_table_asset=bool(row.get("has_table_asset")),
            table_render_mode=row.get("table_render_mode") or "",
            keep=True,
        )

        main.save_review_entry(project["id"], document["id"], payload)
        refreshed = main.get_document_workspace(project["id"], document["id"], include_preview=False)
        refreshed_ref = refreshed["requirements"][0]["table_refs"][0]

        self.assertEqual(refreshed_ref["export_variant"], "slice")
        self.assertTrue(refreshed_ref["export_variant_overridden"])
        self.assertEqual(refreshed_ref["recommended_export_variant"], "source")

    def test_preview_recommendation_replaces_non_overridden_persisted_variant(self) -> None:
        merged_refs, _, _, _ = main._merge_preview_table_assets_for_review_row(
            preview_row={
                "table_refs": [
                    {
                        "table_id": "tbl_preview",
                        "recommended_export_variant": "source",
                        "export_variant": "source",
                        "export_variant_overridden": False,
                    }
                ],
                "content_blocks": [{"kind": "table", "table_id": "tbl_preview"}],
                "has_table_asset": True,
                "table_render_mode": "structured",
            },
            persisted_row={
                "table_refs": [
                    {
                        "table_id": "tbl_preview",
                        "recommended_export_variant": "slice",
                        "export_variant": "slice",
                        "export_variant_overridden": False,
                    }
                ],
                "content_blocks": [{"kind": "table", "table_id": "tbl_preview"}],
                "has_table_asset": True,
                "table_render_mode": "structured",
            },
        )

        self.assertEqual(merged_refs[0]["export_variant"], "source")
        self.assertFalse(merged_refs[0]["export_variant_overridden"])

    def test_row_slice_bbox_matches_merged_cell_slice_with_shared_prefix_and_suffix(self) -> None:
        candidate = {
            "bbox": (0.0, 0.0, 240.0, 250.0),
            "rows_text": [
                ["7", "机械振动和冲击", "外观 Visual Inspection", "USCAR-2 Rev 6 5.1.8", "目视检查", "无外观损伤", "PV", "5", "C Sample"],
                ["连接器或端子插拔 Connector and/or Terminal Cycling", "USCAR-2 Rev 6 5.1.7", "完成一对连接器或端子的10次插拔", "连接器无变形等外观损伤", "PV", "5", "C Sample"],
                ["干电路电阻 Dry Circuit Resistance", "U5SC3AR12REV6", "测量和记录之间的电阻", "电阻值满足5.3.2.4的要求", "PV", "5", "C Sample"],
                ["电压降 Voltage Drop", "U5SC3AR22REV6", "测量并记录测试点之间的毫伏电压降", "电压降满足5.3.2.4的要求", "PV", "5", "C Sample"],
                ["外观 Visual Inspection", "USCAR-2 Rev 6 5.1.8", "目视检查", "无外观损伤", "PV", "5", "C Sample"],
            ],
            "row_bboxes": [
                (0.0, 0.0, 240.0, 50.0),
                (0.0, 50.0, 240.0, 100.0),
                (0.0, 100.0, 240.0, 150.0),
                (0.0, 150.0, 240.0, 200.0),
                (0.0, 200.0, 240.0, 250.0),
            ],
        }

        bbox = table_assets._match_row_slice_bbox(
            candidate,
            headers=[],
            slice_rows=[
                ["7", "机械振动和冲击", "连接器或端子插拔 Connector and/or Terminal Cycling", "USCAR-2 Rev 6 5.1.7", "完成一对连接器或端子的10次插拔", "连接器无变形等外观损伤", "PV", "5", "C Sample"],
                ["7", "机械振动和冲击", "干电路电阻 Dry Circuit Resistance", "USCAR-2 Rev 6 5.3.1", "测量和记录之间的电阻", "电阻值满足5.3.2.4的要求", "PV", "5", "C Sample"],
                ["7", "机械振动和冲击", "电压降 Voltage Drop", "USCAR-2 Rev 6 5.3.2", "测量并记录测试点之间的毫伏电压降", "电压降满足5.3.2.4的要求", "PV", "5", "C Sample"],
                ["7", "机械振动和冲击", "外观 Visual Inspection", "USCAR-2 Rev 6 5.1.8", "目视检查", "无外观损伤", "PV", "5", "C Sample"],
            ],
        )

        self.assertEqual(bbox, (0.0, 50.0, 240.0, 250.0))

    def test_row_slice_bbox_includes_first_group_row_when_slice_starts_with_item_one(self) -> None:
        candidate = {
            "bbox": (0.0, 0.0, 240.0, 200.0),
            "rows_text": [
                ["1", "连接器结合力", "外观 Visual Inspection", "USCAR-2 Rev 6 5.1.8", "目视检查", "无外观损伤", "PV", "15/OTS"],
                ["连接器插入力 Connector-Connector Mating Force", "U5SC4AR22REV6", "插拔速度50mm/Min", "插入力 / Engage force：75N Max.", "PV", "15/OTS"],
                ["连接器拔出力 Connector-Connector Unmating Force", "U5SC4AR22REV6", "破坏卡扣插拔速度50mm/Min", "拔出力 / Disengage force：15N~75N", "PV", "5/OTS"],
                ["主锁解锁力 Lock Deflection", "U5SC4AR22REV6", "向主锁施加解锁力至最大51N", "解锁力 / Unlock force：6N~51N", "PV", "5/OTS"],
            ],
            "row_bboxes": [
                (0.0, 0.0, 240.0, 50.0),
                (0.0, 50.0, 240.0, 100.0),
                (0.0, 100.0, 240.0, 150.0),
                (0.0, 150.0, 240.0, 200.0),
            ],
        }

        bbox = table_assets._match_row_slice_bbox(
            candidate,
            headers=[],
            slice_rows=[
                ["1", "连接器结合力", "连接器插入力 Connector-Connector Mating Force", "USCAR-2 Rev 6 5.4.2", "插拔速度50mm/Min", "插入力 / Engage force：75N Max.", "PV", "15/OTS"],
                ["1", "连接器结合力", "连接器拔出力 Connector-Connector Unmating Force", "USCAR-2 Rev 6 5.4.2", "破坏卡扣插拔速度50mm/Min", "拔出力 / Disengage force：15N~75N", "PV", "5/OTS"],
                ["1", "连接器结合力", "主锁解锁力 Lock Deflection", "USCAR-2 Rev 6 5.4.2", "向主锁施加解锁力至最大51N", "解锁力 / Unlock force：6N~51N", "PV", "5/OTS"],
            ],
        )

        self.assertEqual(bbox, (0.0, 0.0, 240.0, 200.0))

    def test_row_slice_bbox_uses_full_section_span_for_blank_leading_continuations(self) -> None:
        candidate = {
            "bbox": (0.0, 0.0, 240.0, 200.0),
            "rows_text": [
                ["62", "耐化学试液", "USCAR-2 Rev 6 5.1.8", "目视检查"],
                [None, None, "USCAR-2 Rev 6 5.5.1", "绝缘电阻"],
                [None, None, "USCAR-2 Rev 6 5.6.4", "耐化学试液"],
                [None, None, "USCAR-2 Rev 6 5.5.1", "绝缘电阻"],
            ],
            "row_bboxes": [
                (0.0, 0.0, 240.0, 50.0),
                (0.0, 50.0, 240.0, 100.0),
                (0.0, 100.0, 240.0, 150.0),
                (0.0, 150.0, 240.0, 200.0),
            ],
        }

        bbox = table_assets._match_row_slice_bbox(
            candidate,
            headers=["62", "耐化学试液", "USCAR-2 Rev 6 5.1.8", "目视检查"],
            slice_rows=[
                ["62", "耐化学试液", "USCAR-2 Rev 6 5.5.1", "绝缘电阻"],
                ["62", "耐化学试液", "USCAR-2 Rev 6 5.6.4", "耐化学试液"],
                ["62", "耐化学试液", "USCAR-2 Rev 6 5.5.1", "绝缘电阻"],
            ],
        )

        self.assertEqual(bbox, (0.0, 0.0, 240.0, 200.0))

    def test_row_slice_bbox_skips_full_standalone_table_sequences(self) -> None:
        candidate = {
            "bbox": (0.0, 0.0, 240.0, 250.0),
            "rows_text": [
                ["1", "基本尺寸", "要求一"],
                ["2", "电磁阀", "要求二"],
                ["3", "空簧接口", "要求三"],
                ["4", "下叉臂接口", "要求四"],
                ["5", "活塞杆接口", "要求五"],
            ],
            "row_bboxes": [
                (0.0, 0.0, 240.0, 50.0),
                (0.0, 50.0, 240.0, 100.0),
                (0.0, 100.0, 240.0, 150.0),
                (0.0, 150.0, 240.0, 200.0),
                (0.0, 200.0, 240.0, 250.0),
            ],
        }

        bbox = table_assets._match_row_slice_bbox(
            candidate,
            headers=["序号", "项目", "要求"],
            slice_rows=[
                ["1", "基本尺寸", "要求一"],
                ["2", "电磁阀", "要求二"],
                ["3", "空簧接口", "要求三"],
                ["4", "下叉臂接口", "要求四"],
                ["5", "活塞杆接口", "要求五"],
            ],
        )

        self.assertIsNone(bbox)

    def test_context_anchor_prefers_later_page_for_before_context(self) -> None:
        page_lines = {
            0: [
                {
                    "text": "[√]前减振器、蓄能器等需供应商发货",
                    "norm": "前减振器蓄能器等需供应商发货",
                    "bbox": (0.0, 500.0, 120.0, 510.0),
                }
            ],
            1: [
                {
                    "text": "前减振器",
                    "norm": "前减振器",
                    "bbox": (0.0, 180.0, 80.0, 190.0),
                }
            ],
        }

        def fake_load_lines(document, *, source_key, page_index, detection_cache):
            return page_lines.get(page_index, [])

        with patch.object(table_assets, "_load_pdf_page_line_records", side_effect=fake_load_lines):
            anchor = table_assets._find_context_anchor(
                object(),
                source_key="dummy",
                page_indices=[0, 1],
                candidates=table_assets._build_context_anchor_fragments("前减振器", from_end=True),
                detection_cache={},
                prefer_last=True,
            )

        self.assertIsNotNone(anchor)
        self.assertEqual(anchor["page_index"], 1)
        self.assertEqual(anchor["text"], "前减振器")

    def test_context_anchor_prefers_later_exact_self_anchor_over_earlier_noise(self) -> None:
        page_lines = {
            0: [
                {
                    "text": "[√]前减振器、蓄能器等需供应商发货",
                    "norm": "前减振器蓄能器等需供应商发货",
                    "bbox": (0.0, 500.0, 120.0, 510.0),
                }
            ],
            1: [
                {
                    "text": "前减振器",
                    "norm": "前减振器",
                    "bbox": (0.0, 180.0, 80.0, 190.0),
                }
            ],
        }

        def fake_load_lines(document, *, source_key, page_index, detection_cache):
            return page_lines.get(page_index, [])

        with patch.object(table_assets, "_load_pdf_page_line_records", side_effect=fake_load_lines):
            anchor = table_assets._find_context_anchor(
                object(),
                source_key="dummy",
                page_indices=[0, 1],
                candidates=table_assets._build_current_entry_anchor_fragments(
                    table_text="|1|2|2905100 X 8D1110|0.0|左前减振器总成|",
                    table_title="前减振器",
                    before_context_text="前减振器",
                ),
                detection_cache={},
                prefer_score_across_pages=True,
            )

        self.assertIsNotNone(anchor)
        self.assertEqual(anchor["page_index"], 1)
        self.assertEqual(anchor["text"], "前减振器")

    def test_context_anchor_fragments_include_short_label_prefix(self) -> None:
        fragments = table_assets._build_context_anchor_fragments(
            "备注：0、该表格零部件号与OEM系统内零部件号一致。",
            from_end=False,
        )
        values = [value for value, _ in fragments]

        self.assertIn("备注", values)

    def test_preview_promotes_local_reset_clause_heading(self) -> None:
        self.assertTrue(
            extraction_preview._should_promote_local_clause_to_section(
                parent_section_code="3.2.2",
                clause_code="6.6",
            )
        )
        self.assertTrue(
            extraction_preview._should_promote_local_clause_to_section(
                parent_section_code="3.2",
                clause_code="3.2.1",
            )
        )

    def test_fidelity_section_count_skips_main_doc_synthetic_local_split(self) -> None:
        self.assertFalse(
            extraction_preview.counts_toward_fidelity_section(
                {
                    "section_code": "6.1",
                    "appendix_scope": "",
                    "structural_kind": "synthetic",
                }
            )
        )
        self.assertTrue(
            extraction_preview.counts_toward_fidelity_section(
                {
                    "section_code": "6",
                    "appendix_scope": "",
                    "structural_kind": "outline",
                }
            )
        )
        self.assertTrue(
            extraction_preview.counts_toward_fidelity_section(
                {
                    "section_code": "1",
                    "appendix_scope": "Y",
                    "structural_kind": "synthetic",
                }
            )
        )

    def test_split_section_key_preserves_local_code_for_scoped_clause_keys(self) -> None:
        self.assertEqual(extraction_preview._split_section_key("3.2.2::11.12"), ("", "11.12"))
        self.assertEqual(extraction_preview._split_section_key("SUB1:6::10"), ("SUB1", "10"))

    def test_preview_keeps_standard_ids_inside_parent_section(self) -> None:
        rows = extraction_preview._split_requirement_preview_rows(
            item={
                "req_code": "REQ-0001",
                "source_anchor": "Page 7 | Section 4.4 依据标准",
                "requirement_type": "FUNCTIONAL",
                "priority": "MEDIUM",
                "source_line_index": 15,
                "source_line_index_end": 18,
            },
            source_page=7,
            source_page_end=7,
            appendix_scope="",
            section_code="4.4",
            section_title_hint="依据标准",
            explicit_section_code="4.4",
            text=(
                "4.4 依据标准\n"
                "GB/T 30512 汽车禁用物质要求\n"
                "GB/T 28046.5 道路车辆电气及电子设备的环境条件和试验第5部分：化学负荷\n"
                "GB/T 28046.4 道路车辆电气及电子设备的环境条件和试验第4部分：气候负荷"
            ),
            req_order_start=0,
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["section_code"], "4.4")
        self.assertIn("28046.5", rows[0]["requirement_text"])
        self.assertIn("28046.4", rows[0]["requirement_text"])

    def test_serial_table_bare_numeric_rows_are_not_promoted_to_sections(self) -> None:
        line_entries = [
            {
                "text": "3.2.2 产品详细描述",
                "source_page": 1,
                "source_anchor": "Page 1",
                "table_count": 0,
                "source_line_index": 0,
                "global_line_index": 0,
            },
            {
                "text": "2）后减振器布置要求",
                "source_page": 1,
                "source_anchor": "Page 1",
                "table_count": 0,
                "source_line_index": 1,
                "global_line_index": 1,
            },
            {
                "text": "|序号|项目|要求|",
                "source_page": 1,
                "source_anchor": "Page 1",
                "table_count": 1,
                "source_line_index": 2,
                "global_line_index": 2,
            },
            {
                "text": "12 减振器强度要求",
                "source_page": 1,
                "source_anchor": "Page 1",
                "table_count": 1,
                "source_line_index": 3,
                "global_line_index": 3,
            },
            {
                "text": "3）减振器阻尼力基础要求",
                "source_page": 1,
                "source_anchor": "Page 1",
                "table_count": 0,
                "source_line_index": 4,
                "global_line_index": 4,
            },
        ]

        blocks = extractor._extract_section_blocks(line_entries)

        self.assertNotIn("12", [str(block.get("section_code") or "") for block in blocks])

    def test_merge_duplicate_section_block_recomputes_line_indices_from_entries(self) -> None:
        base = {
            "appendix_scope": "",
            "section_code": "6.9",
            "section_title": "商务报价",
            "section_level": 2,
            "_discovery_order": 1,
            "source_page": 9,
            "source_page_end": 9,
            "source_line_index": 0,
            "source_line_index_end": 20,
            "entries": [
                {"source_page": 9, "source_line_index": 18, "text": "6.9 商务报价"},
                {"source_page": 9, "source_line_index": 19, "text": "产过程费用"},
                {"source_page": 9, "source_line_index": 20, "text": "包含以上所有费用。"},
            ],
        }
        incoming = {
            "appendix_scope": "",
            "section_code": "6.9",
            "section_title": "商务报价",
            "section_level": 2,
            "_discovery_order": 2,
            "source_page": 9,
            "source_page_end": 9,
            "source_line_index": 18,
            "source_line_index_end": 20,
            "entries": [
                {"source_page": 9, "source_line_index": 18, "text": "6.9 商务报价"},
                {"source_page": 9, "source_line_index": 19, "text": "产过程费用"},
                {"source_page": 9, "source_line_index": 20, "text": "包含以上所有费用。"},
            ],
        }

        merged = extractor._merge_duplicate_section_block(base, incoming)

        self.assertEqual(merged["source_line_index"], 18)
        self.assertEqual(merged["source_line_index_end"], 20)

    def test_numbered_outline_rows_prefer_row_slice_snapshot(self) -> None:
        self.assertFalse(
            table_assets._should_prefer_source_clause_window(
                table_title="12 活塞杆弯曲强度",
                render_mode="snapshot_only",
                headers=["序号", "项目", "参考标准", "试验方法", "要求"],
                rows=[["12", "活塞杆弯曲强度", "-", "台架", "通过"]],
                source_page=11,
                source_page_end=11,
                notes=[],
                trailing_context_text="",
            )
        )

    def test_visual_snapshot_infers_numbered_outline_row_slice(self) -> None:
        self.assertEqual(
            table_assets._derive_visual_snapshot_slice_rows("12 活塞杆弯曲强度"),
            [["12", "活塞杆弯曲强度"]],
        )

    def test_trial_outline_single_row_header_inference_applies_to_appendix_t(self) -> None:
        headers = table_assets._infer_contextual_single_row_header(
            lines=[
                "|12|活塞杆弯曲强度|-|台架试验|通过|DV/PV|供应商|6件/OTS|1天|"
            ],
            row={
                "appendix_scope": "T",
                "parent_section_code": "12",
                "parent_section_title": "主动前减振器试验大纲",
                "source_anchor": "Appendix T | 主动前减振器试验大纲",
            },
            table_title="12 活塞杆弯曲强度",
        )

        self.assertEqual(
            headers,
            ["序号", "项目", "参考标准", "试验方法", "要求", "DV/PV", "责任方", "样件", "周期"],
        )

    def test_canonical_section_tree_keeps_parent_before_child_without_direct_requirements(self) -> None:
        sections, stats = extraction_preview._build_canonical_section_rows(
            [
                {
                    "appendix_scope": "",
                    "section_code": "1",
                    "section_key": "1",
                    "section_instance_key": "1",
                    "section_title": "概述",
                    "section_level": 1,
                    "parent_section_code": "",
                    "source_page_start": 1,
                    "source_page_end": 1,
                    "source_line_index_start": 13,
                    "source_line_index_end": 13,
                    "has_children": True,
                    "is_leaf": False,
                    "requirement_count": 0,
                    "first_req_order": None,
                    "outline_order": 0,
                },
                {
                    "appendix_scope": "",
                    "section_code": "1.1",
                    "section_key": "1.1",
                    "section_instance_key": "1.1",
                    "section_title": "范围",
                    "section_level": 2,
                    "parent_section_code": "1",
                    "source_page_start": 1,
                    "source_page_end": 1,
                    "source_line_index_start": 14,
                    "source_line_index_end": 14,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 0,
                    "outline_order": 1,
                },
                {
                    "appendix_scope": "",
                    "section_code": "1.5",
                    "section_key": "1.5",
                    "section_instance_key": "1.5",
                    "section_title": "其他",
                    "section_level": 2,
                    "parent_section_code": "1",
                    "source_page_start": 1,
                    "source_page_end": 1,
                    "source_line_index_start": 34,
                    "source_line_index_end": 34,
                    "has_children": True,
                    "is_leaf": False,
                    "requirement_count": 0,
                    "first_req_order": None,
                    "outline_order": 5,
                },
                {
                    "appendix_scope": "",
                    "section_code": "1.5.1",
                    "section_key": "1.5.1",
                    "section_instance_key": "1.5.1",
                    "section_title": "简体中文",
                    "section_level": 3,
                    "parent_section_code": "1.5",
                    "source_page_start": 1,
                    "source_page_end": 1,
                    "source_line_index_start": 35,
                    "source_line_index_end": 35,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 4,
                    "outline_order": 6,
                },
            ]
        )

        self.assertEqual([row["section_code"] for row in sections], ["1", "1.1", "1.5", "1.5.1"])
        self.assertEqual([row["tree_order"] for row in sections], [0, 1, 2, 3])
        self.assertEqual(sections[1]["structural_parent_key"], "1")
        self.assertEqual(sections[3]["structural_parent_key"], "1.5")
        self.assertEqual(stats["zero_direct_requirement_parent_count"], 2)

    def test_canonical_section_tree_preserves_outline_order_over_stale_line_indices(self) -> None:
        sections, _ = extraction_preview._build_canonical_section_rows(
            [
                {
                    "appendix_scope": "",
                    "section_code": "5.8",
                    "section_key": "5.8",
                    "section_instance_key": "5.8",
                    "section_title": "禁用物质要求",
                    "section_level": 2,
                    "parent_section_code": "5",
                    "source_page_start": 9,
                    "source_page_end": 9,
                    "source_line_index_start": 2,
                    "source_line_index_end": 2,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 62,
                    "outline_order": 37,
                },
                {
                    "appendix_scope": "",
                    "section_code": "6.9",
                    "section_key": "6.9",
                    "section_instance_key": "6.9",
                    "section_title": "商务报价",
                    "section_level": 2,
                    "parent_section_code": "6",
                    "source_page_start": 9,
                    "source_page_end": 9,
                    "source_line_index_start": 0,
                    "source_line_index_end": 0,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 71,
                    "outline_order": 47,
                },
            ]
        )

        self.assertEqual([row["section_code"] for row in sections], ["5.8", "6.9"])

    def test_canonical_section_tree_attaches_scoped_local_clauses_under_formal_parent(self) -> None:
        sections, _ = extraction_preview._build_canonical_section_rows(
            [
                {
                    "appendix_scope": "",
                    "section_code": "3.2.2",
                    "section_key": "3.2.2",
                    "section_instance_key": "3.2.2",
                    "section_title": "产品详细描述",
                    "section_level": 3,
                    "parent_section_code": "3.2",
                    "parent_section_key": "3.2",
                    "source_page_start": 2,
                    "source_page_end": 8,
                    "source_line_index_start": 0,
                    "source_line_index_end": 72,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 14,
                    "outline_order": 21,
                },
                {
                    "appendix_scope": "",
                    "section_code": "1",
                    "section_key": "3.2.2::1",
                    "section_instance_key": "3.2.2::1",
                    "section_title": "油泵性能要求",
                    "section_label": "1）",
                    "section_level": 1,
                    "parent_section_code": "3.2.2",
                    "parent_section_key": "3.2.2",
                    "source_page_start": 3,
                    "source_page_end": 3,
                    "source_line_index_start": 1,
                    "source_line_index_end": 1,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 15,
                    "outline_order": None,
                },
                {
                    "appendix_scope": "",
                    "section_code": "6",
                    "section_key": "3.2.2::6",
                    "section_instance_key": "3.2.2::6",
                    "section_title": "双电机控制器总成要求",
                    "section_label": "6）",
                    "section_level": 1,
                    "parent_section_code": "3.2.2",
                    "parent_section_key": "3.2.2",
                    "source_page_start": 5,
                    "source_page_end": 5,
                    "source_line_index_start": 20,
                    "source_line_index_end": 20,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 20,
                    "outline_order": None,
                },
                {
                    "appendix_scope": "",
                    "section_code": "6.1",
                    "section_key": "3.2.2::6.1",
                    "section_instance_key": "3.2.2::6.1",
                    "section_title": "简述",
                    "section_label": "6.1）",
                    "section_level": 2,
                    "parent_section_code": "6",
                    "parent_section_key": "3.2.2::6",
                    "source_page_start": 5,
                    "source_page_end": 5,
                    "source_line_index_start": 21,
                    "source_line_index_end": 21,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 21,
                    "outline_order": None,
                },
                {
                    "appendix_scope": "",
                    "section_code": "6.6",
                    "section_key": "3.2.2::6.6",
                    "section_instance_key": "3.2.2::6.6",
                    "section_title": "模式管理",
                    "section_label": "6.6）",
                    "section_level": 2,
                    "parent_section_code": "6",
                    "parent_section_key": "3.2.2::6",
                    "source_page_start": 6,
                    "source_page_end": 6,
                    "source_line_index_start": 10,
                    "source_line_index_end": 10,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 26,
                    "outline_order": None,
                },
                {
                    "appendix_scope": "",
                    "section_code": "6.6.4",
                    "section_key": "3.2.2::6.6.4",
                    "section_instance_key": "3.2.2::6.6.4",
                    "section_title": "OTA要求",
                    "section_label": "6.6.4）",
                    "section_level": 3,
                    "parent_section_code": "6.6",
                    "parent_section_key": "3.2.2::6.6",
                    "source_page_start": 7,
                    "source_page_end": 7,
                    "source_line_index_start": 1,
                    "source_line_index_end": 1,
                    "has_children": False,
                    "is_leaf": True,
                    "requirement_count": 1,
                    "first_req_order": 30,
                    "outline_order": None,
                },
            ]
        )

        by_key = {row["section_key"]: row for row in sections}
        self.assertEqual(by_key["3.2.2::1"]["structural_parent_key"], "3.2.2")
        self.assertEqual(by_key["3.2.2::6"]["structural_parent_key"], "3.2.2")
        self.assertEqual(by_key["3.2.2::6.1"]["structural_parent_key"], "3.2.2::6")
        self.assertEqual(by_key["3.2.2::6.6.4"]["structural_parent_key"], "3.2.2::6.6")
        self.assertTrue(by_key["3.2.2"]["has_children"])
        self.assertFalse(by_key["3.2.2"]["is_leaf"])

    def test_preview_section_sort_key_uses_tree_order_when_available(self) -> None:
        earlier = extraction_preview._preview_section_row_sort_key(
            {
                "section_code": "1",
                "section_key": "1",
                "tree_order": 0,
                "source_page_start": 1,
                "source_line_index_start": 13,
                "first_req_order": None,
                "outline_order": 0,
                "section_level": 1,
                "structural_kind": "outline",
            }
        )
        later = extraction_preview._preview_section_row_sort_key(
            {
                "section_code": "1.1",
                "section_key": "1.1",
                "tree_order": 1,
                "source_page_start": 1,
                "source_line_index_start": 14,
                "first_req_order": 0,
                "outline_order": 1,
                "section_level": 2,
                "structural_kind": "outline",
            }
        )

        self.assertLess(earlier, later)

    def test_exporter_uses_table_refs_without_markdown_roundtrip(self) -> None:
        _, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0100",
                    "req_code": "REQ0100",
                    "parent_section_label": "9.1 Asset Render",
                    "raw_text": "试验条件如下：",
                    "content_blocks": [
                        {"kind": "paragraph", "text": "试验条件如下："},
                        {"kind": "table", "table_id": "tbl_asset_render"},
                    ],
                    "table_refs": [
                        {
                            "table_id": "tbl_asset_render",
                            "row_range": [1, 2],
                            "inherited_headers": ["Item", "Condition"],
                            "slice_rows": [["Voltage", "12V"], ["Current", "3A"]],
                            "slice_render_mode": "structured",
                            "slice_ai_projection": "table_id: tbl_asset_render",
                        }
                    ],
                    "has_table_asset": True,
                    "table_render_mode": "structured",
                    "final_decision": "ACCEPT",
                    "engineer_reason": "Asset render",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        workbook = load_workbook(file_path)
        sheet = workbook["ComplianceMatrix"]
        meta = self._version_meta_values(workbook)

        self.assertEqual(sheet["C4"].value, "试验条件如下：")
        self.assertIsNone(sheet["C5"].value)
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(meta["TableImageCount"], 1)
        self.assertGreater(sheet.column_dimensions["C"].width or 0, 170)

    def test_table_ref_snapshot_prefers_source_variant(self) -> None:
        slice_path = self.tmp_path / "slice.png"
        source_path = self.tmp_path / "source.png"
        PILImage.new("RGB", (120, 120), "red").save(slice_path)
        PILImage.new("RGB", (280, 90), "blue").save(source_path)

        buffer, image_width, image_height, warning = exporter._table_ref_snapshot(
            {
                "export_variant": "source",
                "slice_render_mode": "snapshot_only",
                "slice_snapshot_path": str(slice_path),
                "source_snapshot_path": str(source_path),
                "inherited_headers": [],
                "slice_rows": [],
            },
            640,
        )

        self.assertIsNotNone(buffer)
        self.assertEqual((image_width, image_height), (280, 90))
        self.assertEqual(warning, "")

    def test_exporter_records_warning_when_source_snapshot_missing(self) -> None:
        _, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0200",
                    "req_code": "REQ0200",
                    "parent_section_label": "9.2 Source Fallback",
                    "raw_text": "导出应回退。",
                    "content_blocks": [
                        {"kind": "paragraph", "text": "导出应回退。"},
                        {"kind": "table", "table_id": "tbl_source_fallback"},
                    ],
                    "table_refs": [
                        {
                            "table_id": "tbl_source_fallback",
                            "row_range": [1, 1],
                            "inherited_headers": ["Item", "Condition"],
                            "slice_rows": [["Voltage", "12V"]],
                            "slice_render_mode": "structured",
                            "slice_ai_projection": "table_id: tbl_source_fallback",
                            "export_variant": "source",
                            "recommended_export_variant": "source",
                            "source_snapshot_path": "",
                        }
                    ],
                    "has_table_asset": True,
                    "table_render_mode": "structured",
                    "final_decision": "ACCEPT",
                    "engineer_reason": "Fallback",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        workbook = load_workbook(file_path)
        meta = self._version_meta_values(workbook)

        self.assertEqual(meta["TableExportWarningCount"], 1)
        self.assertIn("source snapshot missing", str(meta["TableExportWarnings"]))

    def test_single_row_pipe_table_slice_is_materialized_as_table_asset(self) -> None:
        project, document = self._create_project_with_document("-single-row-slice")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview_with_single_row_slice_table()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))
            workspace = main.get_document_workspace(project["id"], document["id"], include_preview=True)

        requirement = workspace["requirements"][0]
        self.assertTrue(requirement["has_table_asset"])
        self.assertEqual(requirement["table_render_mode"], "structured")
        self.assertEqual(len(requirement["table_refs"]), 1)
        self.assertEqual(
            requirement["table_refs"][0]["inherited_headers"],
            ["序号", "项目", "参考标准", "试验方法", "要求", "DV/PV", "责任方", "样件", "周期"],
        )
        self.assertEqual(
            requirement["table_refs"][0]["slice_rows"],
            [["1", "外观检查", "-", "目视外观检查", "表面无划伤", "DV/PV", "供应商", "6件/OTS", "1天"]],
        )

    def test_build_analysis_rows_from_preview_skips_complex_review_only_summary(self) -> None:
        preview = {
            "complete_requirement_entries": [
                {
                    "entry_id": "REQ-CX-summary",
                    "req_code": "REQ-CX-summary",
                    "entry_kind": "REQUIREMENT",
                    "display_text": "复杂跨页表整表总览",
                    "raw_text": "复杂跨页表整表总览",
                    "requirement_text": "复杂跨页表整表总览",
                    "source_page": 17,
                    "source_page_end": 22,
                    "complex_table_mode": True,
                    "complex_role": "zone_summary",
                    "complex_row_policy": "review_only",
                },
                _complex_group_review_row(source_snapshot_path=str(self.tmp_path / "source.png")),
            ]
        }

        rows = main._build_analysis_rows_from_preview(preview)

        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["complex_table_mode"])
        self.assertEqual(rows[0]["complex_role"], "group")

    def test_workspace_suppresses_fallback_rows_shadowed_by_complex_group(self) -> None:
        swallowed = {
            "req_code": "REQ-LEGACY",
            "entry_id": "REQ-LEGACY",
            "entry_kind": "REQUIREMENT",
            "source_page": 18,
            "source_page_end": 18,
            "source_anchor": "Page 18 | legacy",
            "display_text": "legacy split row",
            "raw_text": "legacy split row",
            "requirement_text": "legacy split row",
        }
        unrelated_same_page = {
            "req_code": "REQ-LEGIT",
            "entry_id": "REQ-LEGIT",
            "entry_kind": "REQUIREMENT",
            "source_page": 18,
            "source_page_end": 18,
            "source_anchor": "Page 18 | legit",
            "section_code": "8.1",
            "display_text": "legit boundary row",
            "raw_text": "legit boundary row",
            "requirement_text": "legit boundary row",
        }
        complex_row = _complex_group_review_row(
            source_snapshot_path=str(self.tmp_path / "source.png"),
            shadow_rows=[swallowed],
        )

        rows, _ = main._build_workspace_requirements_from_review_entries(
            [complex_row],
            fallback_requirements=[swallowed, unrelated_same_page],
            preview=None,
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["req_code"], complex_row["req_code"])
        self.assertTrue(any(row["req_code"] == "REQ-LEGIT" for row in rows))

    def test_workspace_review_entries_skip_rows_marked_keep_false(self) -> None:
        excluded_heading = {
            "entry_id": "H-1",
            "req_code": "H-1",
            "entry_kind": "HEADING",
            "keep": False,
            "source_page": 1,
            "source_page_end": 1,
            "source_anchor": "Main > 1",
            "display_text": "1 概述",
            "raw_text": "1 概述",
            "requirement_text": "1 概述",
            "parent_section_code": "1",
            "parent_section_title": "概述",
            "parent_section_label": "1 概述",
        }
        kept_requirement = {
            "entry_id": "REQ0024",
            "req_code": "REQ0024",
            "entry_kind": "REQUIREMENT",
            "keep": True,
            "source_page": 1,
            "source_page_end": 1,
            "source_anchor": "Main > 1.1",
            "display_text": "5.1 首条需求",
            "raw_text": "5.1 首条需求",
            "requirement_text": "5.1 首条需求",
            "parent_section_code": "1.1",
            "parent_section_title": "首条需求",
            "parent_section_label": "1.1 首条需求",
            "requirement_type": "FUNCTIONAL",
            "priority": "MEDIUM",
        }

        rows, _ = main._build_workspace_requirements_from_review_entries(
            [excluded_heading, kept_requirement],
            fallback_requirements=[],
            preview=None,
        )

        self.assertEqual([row["req_code"] for row in rows], ["REQ0024"])
        self.assertTrue(all(bool(row.get("keep", True)) for row in rows))
        self.assertTrue(all(str(row.get("entry_kind") or "") != "HEADING" for row in rows))

    def test_save_review_entry_preserves_complex_metadata_via_table_refs(self) -> None:
        project, document = self._create_project_with_document("-complex-save")
        source_path = self.tmp_path / "complex-source.png"
        PILImage.new("RGB", (320, 180), "navy").save(source_path)

        payload = main.ReviewedRequirementEntry(**_complex_group_review_row(source_snapshot_path=str(source_path)))
        with patch.object(main, "request_preview_refresh", return_value={"accepted": False}):
            result = main.save_review_entry(project["id"], document["id"], payload)

        self.assertEqual(result["review_save"]["action"], "inserted")
        self.assertTrue(result["saved_entry"]["complex_table_mode"])
        self.assertEqual(result["saved_entry"]["complex_zone_id"], "cx-merged-requirement-matrix-p17-p22")

        persisted = db.list_requirements(project["id"], document["id"])
        self.assertEqual(len(persisted), 1)
        self.assertTrue(persisted[0]["complex_table_mode"])
        self.assertTrue(persisted[0]["table_refs"][0]["complex_table_mode"])
        self.assertTrue(persisted[0]["table_refs"][0]["export_variant_locked"])
        self.assertEqual(len(db.list_document_table_assets(project["id"], document["id"])), 1)
        asset = main.get_document_table_asset(project["id"], document["id"], "tbl_complex_group")
        self.assertEqual(asset["table_id"], "tbl_complex_group")
        source_snapshot = main.get_document_table_asset_source_snapshot(project["id"], document["id"], "tbl_complex_group")
        self.assertTrue(Path(source_snapshot.path).exists())

    def test_export_requires_locked_extraction(self) -> None:
        project, document = self._create_project_with_document("-export-lock")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 1",
            requirement_text="需要先锁定条目集后才能导出。",
            raw_text="需要先锁定条目集后才能导出。",
            display_text="需要先锁定条目集后才能导出。",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")

        with self.assertRaises(HTTPException) as export_exc:
            main.export_project(project["id"], main.ExportRequest(document_id=document["id"]))

        self.assertEqual(export_exc.exception.status_code, 409)
        self.assertIn("locked extraction", str(export_exc.exception.detail))

        main.lock_document_extraction(project["id"], document["id"])
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            result = main.export_project(project["id"], main.ExportRequest(document_id=document["id"]))
        self.assertTrue((self.tmp_path / "exports" / result["file_name"]).exists())

    def test_run_ai_analysis_persists_extended_assessment_fields_and_runs(self) -> None:
        project, document = self._create_project_with_document("-ai-copilot")
        table_refs = json.dumps(
            [
                {
                    "table_id": "tbl_ai",
                    "slice_ai_projection": "|Temperature|85C|Voltage|12V|",
                }
            ],
            ensure_ascii=False,
        )
        content_blocks = json.dumps(
            [{"kind": "paragraph", "text": "适用于整车冷启动工况。"}],
            ensure_ascii=False,
        )
        requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 5.1",
            requirement_text="ECU shall support wake-up time less than 100ms.",
            raw_text="ECU shall support wake-up time less than 100ms.",
            display_text="ECU shall support wake-up time less than 100ms.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="5.1",
            parent_section_title="Startup",
            parent_section_label="5.1 Startup",
            table_refs=table_refs,
            content_blocks=content_blocks,
            has_table_asset=True,
            table_render_mode="structured",
        )
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(self.tmp_path / "platform-spec.txt"),
            extracted_text=(
                "5.1 Startup\n"
                "The ECU shall support wake-up time less than 100ms.\n"
                "Measured wake-up time is 80ms at -20C.\n"
            ),
            doc_version="v1.0",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_AI_TRACE_DIR": str(self.tmp_path / "ai_traces"),
            },
            clear=False,
        ):
            result = main.run_document_ai_analysis(project["id"], document["id"])

        row = next(item for item in result["requirements"] if item["id"] == requirement["id"])
        self.assertTrue(row["ai_decision"])
        self.assertEqual(row["assessment_mode"], "heuristic")
        self.assertEqual(row["analysis_label"], "REQUIREMENT")
        self.assertTrue(row["normalized_requirement"])
        self.assertIn(row["risk_level"], {"LOW", "MEDIUM", "HIGH"})
        self.assertIn("source_doc_version", row["evidences"][0])

        runs = db.list_assessment_runs(requirement["id"])
        self.assertEqual(len(runs), 2)
        self.assertEqual({run["mode"] for run in runs}, {"screening_heuristic", "heuristic"})
        backend_by_mode = {run["mode"]: run["retrieval_backend"] for run in runs}
        self.assertEqual(backend_by_mode["screening_heuristic"], "screening_only")
        self.assertEqual(backend_by_mode["heuristic"], "local")

    def test_ai_analysis_job_api_exposes_status_events_and_requirement_detail(self) -> None:
        project, document = self._create_project_with_document("-ai-job-api")
        requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 5.1",
            requirement_text="ECU shall support wake-up time less than 100ms.",
            raw_text="ECU shall support wake-up time less than 100ms.",
            display_text="ECU shall support wake-up time less than 100ms.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="5.1",
            parent_section_title="Startup",
            parent_section_label="5.1 Startup",
        )
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(self.tmp_path / "platform-spec-job-api.txt"),
            extracted_text=(
                "5.1 Startup\n"
                "The ECU shall support wake-up time less than 100ms.\n"
                "Measured wake-up time is 80ms at -20C.\n"
            ),
            doc_version="v1.0",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_AI_TRACE_DIR": str(self.tmp_path / "ai_traces_job_api"),
            },
            clear=False,
        ):
            with patch.object(main, "Thread", ImmediateThread):
                with TestClient(main.app) as client:
                    start_response = client.post(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/start",
                        json={"profile_id": "azure"},
                    )
                    self.assertEqual(start_response.status_code, 200)
                    start_payload = start_response.json()
                    self.assertTrue(start_payload["job"]["id"])
                    self.assertEqual(start_payload["job"]["profile_id"], "azure")
                    self.assertTrue(any(item["id"] == "azure" for item in start_payload["ai_profiles"]))

                    status_response = client.get(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/ai-analysis-status"
                    )
                    self.assertEqual(status_response.status_code, 200)
                    status_payload = status_response.json()
                    self.assertEqual(status_payload["job"]["status"], "COMPLETED")
                    self.assertEqual(status_payload["job"]["profile_id"], "azure")
                    self.assertTrue(
                        any(item["event_type"] == "requirement_completed" for item in status_payload["events"])
                    )
                    self.assertTrue(any(item["event_type"] == "job_completed" for item in status_payload["events"]))

                    detail_response = client.get(f"/api/requirements/{requirement['id']}/analysis-detail")
                    self.assertEqual(detail_response.status_code, 200)
                    detail_payload = detail_response.json()
                    self.assertEqual(detail_payload["requirement"]["id"], requirement["id"])
                    self.assertEqual(detail_payload["assessment"]["analysis_label"], "REQUIREMENT")
                    self.assertGreaterEqual(len(detail_payload["assessment_runs"]), 1)
                    self.assertIsInstance(detail_payload["evidences"], list)

    def test_ai_analysis_detail_and_events_use_display_req_codes_from_workspace_order(self) -> None:
        project, document = self._create_project_with_document("-ai-display-order")
        later_requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0023",
            source_page=5,
            source_anchor="Main > 5.7",
            requirement_text="5.7 次条需求",
            raw_text="5.7 次条需求",
            display_text="5.7 次条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="5.7",
            parent_section_title="次条需求",
            parent_section_label="5.7 次条需求",
        )
        first_requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0024",
            source_page=5,
            source_anchor="Main > 5.1",
            requirement_text="5.1 首条需求",
            raw_text="5.1 首条需求",
            display_text="5.1 首条需求",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="5.1",
            parent_section_title="首条需求",
            parent_section_label="5.1 首条需求",
        )
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(self.tmp_path / "platform-spec-display-order.txt"),
            extracted_text=(
                "5.1 首条需求\n"
                "系统应满足首条需求。\n"
                "5.7 次条需求\n"
                "系统应满足次条需求。\n"
            ),
            doc_version="v1.0",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        preview_cache.persist_preview_cache(
            db.get_sor_document(document["id"]) or {},
            _fake_preview_with_reversed_req_codes(),
        )
        main.lock_document_extraction(project["id"], document["id"])

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_AI_TRACE_DIR": str(self.tmp_path / "ai_traces_display_order"),
            },
            clear=False,
        ):
            with patch.object(main, "Thread", ImmediateThread):
                with TestClient(main.app) as client:
                    start_response = client.post(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/start",
                        json={"profile_id": "azure"},
                    )
                    self.assertEqual(start_response.status_code, 200)

                    status_response = client.get(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/ai-analysis-status"
                    )
                    self.assertEqual(status_response.status_code, 200)
                    status_payload = status_response.json()
                    completed_codes = {
                        str(item.get("requirement_code") or "")
                        for item in status_payload["events"]
                        if item.get("event_type") == "requirement_completed"
                    }

                    detail_response = client.get(f"/api/requirements/{first_requirement['id']}/analysis-detail")
                    self.assertEqual(detail_response.status_code, 200)
                    detail_payload = detail_response.json()

        self.assertEqual(status_payload["job"]["status"], "COMPLETED")
        self.assertEqual(completed_codes, {"REQ0001", "REQ0002"})
        self.assertNotIn("REQ0023", completed_codes)
        self.assertNotIn("REQ0024", completed_codes)
        self.assertEqual(detail_payload["requirement"]["req_code"], first_requirement["req_code"])
        self.assertEqual(detail_payload["requirement"]["display_req_code"], "REQ0001")
        self.assertEqual(detail_payload["requirement"]["display_sort_index"], 1)
        self.assertEqual(later_requirement["req_code"], "REQ0023")

    def test_ai_analysis_job_api_reports_failed_runs(self) -> None:
        project, document = self._create_project_with_document("-ai-job-failed")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 5.1",
            requirement_text="ECU shall support wake-up time less than 100ms.",
            raw_text="ECU shall support wake-up time less than 100ms.",
            display_text="ECU shall support wake-up time less than 100ms.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="5.1",
            parent_section_title="Startup",
            parent_section_label="5.1 Startup",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        with patch.object(main, "run_ai_copilot_assessment", side_effect=RuntimeError("simulated ai failure")):
            with patch.object(main, "Thread", ImmediateThread):
                with TestClient(main.app) as client:
                    start_response = client.post(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/start"
                    )
                    self.assertEqual(start_response.status_code, 200)

                    status_response = client.get(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/ai-analysis-status"
                    )
                    self.assertEqual(status_response.status_code, 200)
                    status_payload = status_response.json()
                    self.assertEqual(status_payload["job"]["status"], "FAILED")
                    self.assertIn("simulated ai failure", status_payload["job"]["error_message"])
                    self.assertEqual(status_payload["document"]["ai_status"], "STALE")
                    self.assertTrue(any(item["event_type"] == "job_failed" for item in status_payload["events"]))

    def test_ai_analysis_job_api_surfaces_screening_and_analysis_failed_statuses(self) -> None:
        project, document = self._create_project_with_document("-ai-job-fail-closed")
        screening_requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 5.1",
            requirement_text="Supplier shall provide EMC validation evidence.",
            raw_text="Supplier shall provide EMC validation evidence.",
            display_text="Supplier shall provide EMC validation evidence.",
            requirement_type="REGULATION",
            priority="MEDIUM",
            parent_section_code="5.1",
            parent_section_title="Compliance",
            parent_section_label="5.1 Compliance",
        )
        analysis_requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0002",
            source_page=2,
            source_anchor="Page 2 | Section 5.2",
            requirement_text="System shall support wake-up time less than 100ms.",
            raw_text="System shall support wake-up time less than 100ms.",
            display_text="System shall support wake-up time less than 100ms.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="5.2",
            parent_section_title="Startup",
            parent_section_label="5.2 Startup",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])
        ai_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="qdrant",
            top_k=5,
            trace_dir=self.tmp_path / "ai_traces_fail_closed",
            azure_auth_mode="",
            azure_allow_interactive_login=False,
            azure_search_endpoint="",
            azure_search_api_key="",
            azure_search_index_name="",
            azure_openai_endpoint="",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="",
            azure_openai_embedding_deployment="",
            azure_openai_embedding_dimensions=0,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="aliyun",
            provider="aliyun",
            profile_label="Aliyun",
            chat_model="qwen3.5-plus",
            screening_model="qwen3.5-flash",
            embedding_model="text-embedding-v4",
            dashscope_api_key="aliyun-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            qdrant_url="http://localhost:6333",
            qdrant_collection_name="reqpilot-knowledge-aliyun-local",
        )
        eligible_rows = db.list_requirements(project["id"], document["id"])
        prepared = (project, document, eligible_rows, eligible_rows, [], ai_config)
        screening_failed_result = {
            "assessment": {
                "profile_id": "aliyun",
                "provider": "aliyun",
                "analysis_label": "REQUIREMENT",
                "analysis_basis": "LLM screening failed; analysis stopped. Request timed out.",
                "ai_decision": "",
                "ai_reason": "",
                "ai_confidence": 0.0,
                "risk_level": "",
                "clarification_question": "",
                "normalized_requirement": "Supplier shall provide EMC validation evidence.",
                "assessment_mode": "screening_failed",
                "retrieval_backend": "screening_only",
                "model_name": "qwen3.5-flash",
                "prompt_version": "v2.1-requirement-screening-aliyun",
            },
            "evidences": [],
            "runs": [
                {
                    "profile_id": "aliyun",
                    "provider": "aliyun",
                    "mode": "screening_llm_shadow",
                    "retrieval_backend": "screening_only",
                    "model_name": "qwen3.5-flash",
                    "prompt_version": "v2.1-requirement-screening-aliyun",
                    "input_payload_json": "{}",
                    "output_payload_json": "",
                    "validation_status": "failed",
                    "fallback_reason": "Request timed out.",
                }
            ],
            "trace_path": str(self.tmp_path / "trace-screening-failed.json"),
        }
        analysis_failed_result = {
            "assessment": {
                "profile_id": "aliyun",
                "provider": "aliyun",
                "analysis_label": "REQUIREMENT",
                "analysis_basis": "LLM requirement analysis failed; analysis stopped. Connection error.",
                "ai_decision": "",
                "ai_reason": "",
                "ai_confidence": 0.0,
                "risk_level": "",
                "clarification_question": "",
                "normalized_requirement": "System shall support wake-up time less than 100ms.",
                "assessment_mode": "analysis_failed",
                "retrieval_backend": "qdrant",
                "model_name": "qwen3.5-plus",
                "prompt_version": "v2.0-aliyun-rag-copilot",
            },
            "evidences": [],
            "runs": [
                {
                    "profile_id": "aliyun",
                    "provider": "aliyun",
                    "mode": "screening_llm_shadow",
                    "retrieval_backend": "screening_only",
                    "model_name": "qwen3.5-flash",
                    "prompt_version": "v2.1-requirement-screening-aliyun",
                    "input_payload_json": "{}",
                    "output_payload_json": "{\"analysis_label\":\"REQUIREMENT\"}",
                    "validation_status": "accepted",
                    "fallback_reason": "",
                },
                {
                    "profile_id": "aliyun",
                    "provider": "aliyun",
                    "mode": "llm_shadow",
                    "retrieval_backend": "qdrant",
                    "model_name": "qwen3.5-plus",
                    "prompt_version": "v2.0-aliyun-rag-copilot",
                    "input_payload_json": "{}",
                    "output_payload_json": "",
                    "validation_status": "failed",
                    "fallback_reason": "Connection error.",
                },
            ],
            "trace_path": str(self.tmp_path / "trace-analysis-failed.json"),
        }

        with patch.object(main, "_prepare_document_ai_analysis_request", return_value=prepared), patch.object(
            main,
            "run_ai_copilot_assessment",
            side_effect=[screening_failed_result, analysis_failed_result],
        ), patch.object(main, "Thread", ImmediateThread):
            with TestClient(main.app) as client:
                start_response = client.post(
                    f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/start",
                    json={"profile_id": "aliyun"},
                )
                self.assertEqual(start_response.status_code, 200)
                status_response = client.get(
                    f"/api/projects/{project['id']}/sor-documents/{document['id']}/ai-analysis-status"
                )
                self.assertEqual(status_response.status_code, 200)
                status_payload = status_response.json()

        completed_messages = {
            str(item.get("message") or "")
            for item in status_payload["events"]
            if item.get("event_type") == "requirement_completed"
        }
        self.assertIn("完成 REQ0001：筛选失败。", completed_messages)
        self.assertIn("完成 REQ0002：分析失败。", completed_messages)
        screening_row = db.get_requirement(screening_requirement["id"]) or {}
        analysis_row = db.get_requirement(analysis_requirement["id"]) or {}
        self.assertEqual(screening_row.get("assessment_mode"), "screening_failed")
        self.assertEqual(analysis_row.get("assessment_mode"), "analysis_failed")

    def test_run_ai_analysis_marks_info_rows_and_skips_manual_decision(self) -> None:
        project, document = self._create_project_with_document("-ai-info")
        info_requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 4.5",
            requirement_text="Functional status A",
            raw_text="Functional status A",
            display_text="Functional status A",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="4.5",
            parent_section_title="Status",
            parent_section_label="4.5 Status",
        )
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(self.tmp_path / "platform-spec-info.txt"),
            extracted_text="Status naming reference only.",
            doc_version="v1.0",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_AI_TRACE_DIR": str(self.tmp_path / "ai_traces_info"),
            },
            clear=False,
        ):
            result = main.run_document_ai_analysis(project["id"], document["id"])

        row = next(item for item in result["requirements"] if item["id"] == info_requirement["id"])
        self.assertEqual(row["analysis_label"], "INFO")
        self.assertTrue(row["analysis_basis"])
        self.assertEqual(row["ai_decision"], "")
        self.assertEqual(row["evidences"], [])
        self.assertEqual(result["ai_processed_requirements"], 0)
        self.assertEqual(result["ai_info_requirements"], 1)

    def test_ai_analysis_job_api_exposes_reasoning_stream_events(self) -> None:
        project, document = self._create_project_with_document("-ai-reasoning-events")
        requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 4.5",
            requirement_text="Supplier shall submit CAMDS data before SOP.",
            raw_text="Supplier shall submit CAMDS data before SOP.",
            display_text="Supplier shall submit CAMDS data before SOP.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="4.5",
            parent_section_title="Compliance",
            parent_section_label="4.5 Compliance",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])
        ai_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="qdrant",
            top_k=5,
            trace_dir=self.tmp_path / "ai_traces_reasoning_events",
            azure_auth_mode="",
            azure_allow_interactive_login=False,
            azure_search_endpoint="",
            azure_search_api_key="",
            azure_search_index_name="",
            azure_openai_endpoint="",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="",
            azure_openai_embedding_deployment="",
            azure_openai_embedding_dimensions=0,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="aliyun",
            provider="aliyun",
            profile_label="Aliyun",
            chat_model="qwen3.5-plus",
            screening_model="qwen3.5-flash",
            embedding_model="text-embedding-v4",
            dashscope_api_key="aliyun-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            qdrant_url="http://localhost:6333",
            qdrant_collection_name="reqpilot-knowledge-aliyun-local",
        )
        eligible_rows = db.list_requirements(project["id"], document["id"])
        prepared = (project, document, eligible_rows, eligible_rows, [], ai_config)

        def fake_run_ai_copilot_assessment(
            requirement_row: dict[str, object],
            knowledge_docs: list[dict[str, object]],
            *,
            config: AICopilotConfig | None = None,
            forced_analysis_label: str | None = None,
            forced_analysis_basis: str = "",
            progress_callback=None,
        ) -> dict[str, object]:
            if progress_callback:
                progress_callback(
                    {
                        "type": "stage",
                        "stage": "screening",
                        "mode": "screening_llm_shadow",
                        "model_name": "qwen3.5-flash",
                        "message": "正在进行 requirement / INFO 筛选。",
                    }
                )
                progress_callback(
                    {
                        "type": "reasoning",
                        "stage": "screening",
                        "mode": "screening_llm_shadow",
                        "model_name": "qwen3.5-flash",
                        "delta": "先判断这条是否属于真正需要合规结论的 requirement。",
                        "content": "先判断这条是否属于真正需要合规结论的 requirement。",
                        "characters_emitted": 24,
                    }
                )
                progress_callback(
                    {
                        "type": "answer",
                        "stage": "screening",
                        "mode": "screening_llm_shadow",
                        "model_name": "qwen3.5-flash",
                        "delta": "{\"analysis_label\":\"REQUIREMENT\",\"reason\":\"该条涉及 CAMDS 填报义务。\"}",
                        "content": "{\"analysis_label\":\"REQUIREMENT\",\"reason\":\"该条涉及 CAMDS 填报义务。\"}",
                        "characters_emitted": 58,
                    }
                )
                progress_callback(
                    {
                        "type": "stage",
                        "stage": "retrieval",
                        "mode": "screening_llm_shadow",
                        "model_name": "qwen3.5-flash",
                        "message": "筛选完成，正在检索候选证据。",
                    }
                )
                progress_callback(
                    {
                        "type": "stage",
                        "stage": "assessment",
                        "mode": "llm_shadow",
                        "model_name": "qwen3.5-plus",
                        "message": "候选证据已就绪，正在生成结构化结论。",
                    }
                )
                progress_callback(
                    {
                        "type": "reasoning",
                        "stage": "assessment",
                        "mode": "llm_shadow",
                        "model_name": "qwen3.5-plus",
                        "delta": "先检查 requirement 的约束条件。",
                        "content": "先检查 requirement 的约束条件。",
                        "characters_emitted": 16,
                    }
                )
                progress_callback(
                    {
                        "type": "reasoning",
                        "stage": "assessment",
                        "mode": "llm_shadow",
                        "model_name": "qwen3.5-plus",
                        "delta": " 再核对候选证据是否真正覆盖 CAMDS 义务。",
                        "content": "先检查 requirement 的约束条件。 再核对候选证据是否真正覆盖 CAMDS 义务。",
                        "characters_emitted": 40,
                    }
                )
                progress_callback(
                    {
                        "type": "answer",
                        "stage": "assessment",
                        "mode": "llm_shadow",
                        "model_name": "qwen3.5-plus",
                        "delta": "{\"normalized_requirement\":\"Supplier shall submit CAMDS data before SOP.\",",
                        "content": "{\"normalized_requirement\":\"Supplier shall submit CAMDS data before SOP.\",",
                        "characters_emitted": 72,
                    }
                )
            return {
                "assessment": {
                    "profile_id": "aliyun",
                    "provider": "aliyun",
                    "analysis_label": "REQUIREMENT",
                    "analysis_basis": "已完成LLM分析。",
                    "ai_decision": "NEED_CLARIFICATION",
                    "ai_reason": "候选证据无法直接覆盖 CAMDS 数据填报义务。",
                    "ai_confidence": 0.91,
                    "risk_level": "HIGH",
                    "clarification_question": "请补充 CAMDS 数据提交记录。",
                    "normalized_requirement": "Supplier shall submit CAMDS data before SOP.",
                    "assessment_mode": "llm_shadow",
                    "retrieval_backend": "qdrant",
                    "model_name": "qwen3.5-plus",
                    "prompt_version": "v2.0-aliyun-rag-copilot",
                },
                "evidences": [],
                "runs": [
                    {
                        "profile_id": "aliyun",
                        "provider": "aliyun",
                        "mode": "screening_llm_shadow",
                        "retrieval_backend": "screening_only",
                        "model_name": "qwen3.5-flash",
                        "prompt_version": "v2.1-requirement-screening-aliyun",
                        "input_payload_json": "{}",
                        "output_payload_json": "{\"analysis_label\":\"REQUIREMENT\"}",
                        "validation_status": "accepted",
                        "fallback_reason": "",
                    },
                    {
                        "profile_id": "aliyun",
                        "provider": "aliyun",
                        "mode": "llm_shadow",
                        "retrieval_backend": "qdrant",
                        "model_name": "qwen3.5-plus",
                        "prompt_version": "v2.0-aliyun-rag-copilot",
                        "input_payload_json": "{}",
                        "output_payload_json": "{\"proposed_decision\":\"NEED_CLARIFICATION\"}",
                        "validation_status": "accepted",
                        "fallback_reason": "",
                    },
                ],
                "trace_path": str(self.tmp_path / "trace-reasoning-events.json"),
            }

        with patch.object(main, "_prepare_document_ai_analysis_request", return_value=prepared), patch.object(
            main,
            "run_ai_copilot_assessment",
            side_effect=fake_run_ai_copilot_assessment,
        ), patch.object(main, "Thread", ImmediateThread):
            with TestClient(main.app) as client:
                start_response = client.post(
                    f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/start",
                    json={"profile_id": "aliyun"},
                )
                self.assertEqual(start_response.status_code, 200)
                status_response = client.get(
                    f"/api/projects/{project['id']}/sor-documents/{document['id']}/ai-analysis-status"
                )
                self.assertEqual(status_response.status_code, 200)
                status_payload = status_response.json()

        reasoning_events = [
            item for item in status_payload["events"] if item.get("event_type") == "requirement_reasoning"
        ]
        self.assertTrue(reasoning_events)
        self.assertTrue(
            any("CAMDS" in str(item.get("payload", {}).get("content_preview") or "") for item in reasoning_events)
        )
        screening_reasoning_events = [
            item for item in status_payload["events"] if item.get("event_type") == "requirement_screening_reasoning"
        ]
        self.assertTrue(screening_reasoning_events)
        self.assertTrue(
            any("REQUIREMENT" in str(item.get("payload", {}).get("content_preview") or "") for item in status_payload["events"] if item.get("event_type") == "requirement_screening_output")
        )
        self.assertTrue(
            any(item.get("event_type") == "requirement_stage_update" for item in status_payload["events"])
        )
        self.assertTrue(
            any(item.get("event_type") == "requirement_output" for item in status_payload["events"])
        )
        self.assertTrue(
            any(item.get("event_type") == "requirement_answer_started" for item in status_payload["events"])
        )
        self.assertTrue(
            any(item.get("event_type") == "requirement_completed" for item in status_payload["events"])
        )
        updated_row = db.get_requirement(requirement["id"]) or {}
        self.assertEqual(updated_row.get("assessment_mode"), "llm_shadow")

    def test_ai_analysis_job_can_be_stopped_from_api(self) -> None:
        project, document = self._create_project_with_document("-ai-stop")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 4.5",
            requirement_text="Supplier shall submit CAMDS data before SOP.",
            raw_text="Supplier shall submit CAMDS data before SOP.",
            display_text="Supplier shall submit CAMDS data before SOP.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="4.5",
            parent_section_title="Compliance",
            parent_section_label="4.5 Compliance",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])
        ai_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="qdrant",
            top_k=5,
            trace_dir=self.tmp_path / "ai_traces_stop",
            azure_auth_mode="",
            azure_allow_interactive_login=False,
            azure_search_endpoint="",
            azure_search_api_key="",
            azure_search_index_name="",
            azure_openai_endpoint="",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="",
            azure_openai_embedding_deployment="",
            azure_openai_embedding_dimensions=0,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="aliyun",
            provider="aliyun",
            profile_label="Aliyun",
            chat_model="qwen3.5-plus",
            screening_model="qwen3.5-flash",
            embedding_model="text-embedding-v4",
            dashscope_api_key="aliyun-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            qdrant_url="http://localhost:6333",
            qdrant_collection_name="reqpilot-knowledge-aliyun-local",
        )
        eligible_rows = db.list_requirements(project["id"], document["id"])
        prepared = (project, document, eligible_rows, eligible_rows, [], ai_config)
        started = Event()

        def fake_run_ai_copilot_assessment(
            requirement_row: dict[str, object],
            knowledge_docs: list[dict[str, object]],
            *,
            config: AICopilotConfig | None = None,
            forced_analysis_label: str | None = None,
            forced_analysis_basis: str = "",
            progress_callback=None,
        ) -> dict[str, object]:
            if progress_callback:
                progress_callback(
                    {
                        "type": "stage",
                        "stage": "assessment",
                        "mode": "llm_shadow",
                        "model_name": "qwen3.5-plus",
                        "message": "候选证据已就绪，正在生成结构化结论。",
                    }
                )
                started.set()
                for index in range(200):
                    progress_callback(
                        {
                            "type": "reasoning",
                            "stage": "assessment",
                            "mode": "llm_shadow",
                            "model_name": "qwen3.5-plus",
                            "delta": f"thinking-{index}-" + ("x" * 200),
                            "content": f"thinking-{index}-" + ("x" * 200),
                            "characters_emitted": (index + 1) * 210,
                        }
                    )
                    time.sleep(0.02)
            return {
                "assessment": {
                    "profile_id": "aliyun",
                    "provider": "aliyun",
                    "analysis_label": "REQUIREMENT",
                    "analysis_basis": "已完成LLM分析。",
                    "ai_decision": "NEED_CLARIFICATION",
                    "ai_reason": "候选证据无法直接覆盖 CAMDS 数据填报义务。",
                    "ai_confidence": 0.91,
                    "risk_level": "HIGH",
                    "clarification_question": "请补充 CAMDS 数据提交记录。",
                    "normalized_requirement": "Supplier shall submit CAMDS data before SOP.",
                    "assessment_mode": "llm_shadow",
                    "retrieval_backend": "qdrant",
                    "model_name": "qwen3.5-plus",
                    "prompt_version": "v2.0-aliyun-rag-copilot",
                },
                "evidences": [],
                "runs": [],
                "trace_path": str(self.tmp_path / "trace-stop.json"),
            }

        with patch.object(main, "_prepare_document_ai_analysis_request", return_value=prepared), patch.object(
            main,
            "run_ai_copilot_assessment",
            side_effect=fake_run_ai_copilot_assessment,
        ):
            with TestClient(main.app) as client:
                start_response = client.post(
                    f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/start",
                    json={"profile_id": "aliyun"},
                )
                self.assertEqual(start_response.status_code, 200)
                self.assertTrue(started.wait(timeout=1.5))

                stop_response = client.post(
                    f"/api/projects/{project['id']}/sor-documents/{document['id']}/run-ai-analysis/stop"
                )
                self.assertEqual(stop_response.status_code, 200)
                self.assertEqual(stop_response.json()["job"]["status"], "CANCEL_REQUESTED")

                deadline = time.time() + 3.0
                status_payload = stop_response.json()
                while time.time() < deadline:
                    status_response = client.get(
                        f"/api/projects/{project['id']}/sor-documents/{document['id']}/ai-analysis-status"
                    )
                    self.assertEqual(status_response.status_code, 200)
                    status_payload = status_response.json()
                    if str(status_payload.get("job", {}).get("status") or "").upper() == "CANCELLED":
                        break
                    time.sleep(0.05)

        self.assertEqual(str(status_payload.get("job", {}).get("status") or "").upper(), "CANCELLED")
        self.assertEqual(str(status_payload.get("job", {}).get("current_stage") or "").lower(), "cancelled")
        event_types = [str(item.get("event_type") or "") for item in status_payload.get("events") or []]
        self.assertIn("job_cancel_requested", event_types)
        self.assertIn("job_cancelled", event_types)
        self.assertTrue(
            any("已停止" in str(item.get("message") or "") for item in status_payload.get("events") or [])
        )
        refreshed_document = db.get_sor_document(document["id"]) or {}
        self.assertEqual(str(refreshed_document.get("ai_status") or "").upper(), "STALE")

    def test_override_analysis_label_promotes_info_row_back_to_requirement(self) -> None:
        project, document = self._create_project_with_document("-override-info")
        info_requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 4.5",
            requirement_text="Functional status A",
            raw_text="Functional status A",
            display_text="Functional status A",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="4.5",
            parent_section_title="Status",
            parent_section_label="4.5 Status",
        )
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="status-spec.txt",
            file_type="txt",
            file_path=str(self.tmp_path / "status-spec.txt"),
            extracted_text=(
                "4.5 Status\n"
                "Functional status A shall be supported during and after the test.\n"
                "Status A validation passed in historical project baseline.\n"
            ),
            doc_version="v1.0",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_AI_TRACE_DIR": str(self.tmp_path / "ai_traces_override"),
            },
            clear=False,
        ):
            main.run_document_ai_analysis(project["id"], document["id"])
            row_before = db.get_requirement(info_requirement["id"])
            self.assertEqual(row_before["analysis_label"], "INFO")

            result = main.override_analysis_label(
                info_requirement["id"],
                main.AnalysisLabelOverrideRequest(analysis_label="REQUIREMENT", rerun_assessment=True),
            )

        updated = db.get_requirement(info_requirement["id"])
        self.assertEqual(updated["analysis_label"], "REQUIREMENT")
        self.assertTrue(updated["analysis_basis"])
        self.assertTrue(updated["ai_decision"])
        self.assertEqual(result["assessment"]["analysis_label"], "REQUIREMENT")
        self.assertGreaterEqual(result["assessment_runs_recorded"], 1)
        runs = db.list_assessment_runs(info_requirement["id"])
        self.assertIn("screening_manual_override", {run["mode"] for run in runs})

    def test_run_ai_analysis_without_knowledge_docs_only_screens_requirement_rows(self) -> None:
        project, document = self._create_project_with_document("-screen-only")
        requirement = db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 4.5",
            requirement_text="ECU shall support wake-up time less than 100ms.",
            raw_text="ECU shall support wake-up time less than 100ms.",
            display_text="ECU shall support wake-up time less than 100ms.",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
            parent_section_code="4.5",
            parent_section_title="Status",
            parent_section_label="4.5 Status",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        with patch.dict(
            os.environ,
            {
                "REQPILOT_AI_ASSESSMENT_MODE": "heuristic",
                "REQPILOT_AI_TRACE_DIR": str(self.tmp_path / "ai_traces_screen_only"),
            },
            clear=False,
        ):
            result = main.run_document_ai_analysis(project["id"], document["id"])

        row = next(item for item in result["requirements"] if item["id"] == requirement["id"])
        self.assertEqual(row["analysis_label"], "REQUIREMENT")
        self.assertEqual(row["assessment_mode"], "screening_only")
        self.assertTrue(row["analysis_basis"])
        self.assertEqual(row["ai_decision"], "")
        self.assertEqual(row["evidences"], [])
        self.assertEqual(result["ai_processed_requirements"], 0)
        self.assertEqual(result["ai_info_requirements"], 0)
        self.assertEqual(result["ai_screened_requirements"], 1)

    def test_upload_knowledge_document_accepts_metadata_fields(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-metadata")
        knowledge_dir = self.tmp_path / "knowledge"
        knowledge_dir.mkdir(parents=True, exist_ok=True)
        upload = UploadFile(filename="spec.txt", file=io.BytesIO("Spec content".encode("utf-8")))

        with patch.object(main, "KNOWLEDGE_DIR", knowledge_dir):
            created = asyncio.run(
                main.upload_knowledge_document(
                    project["id"],
                    file=upload,
                    doc_type="STANDARD",
                    doc_version="V1.2",
                    status="ACTIVE",
                    effective_from="2026-01-01",
                    effective_to="2026-12-31",
                )
            )

        listed = main.list_knowledge_documents(project["id"])
        self.assertEqual(created["doc_type"], "STANDARD")
        self.assertEqual(created["doc_version"], "V1.2")
        self.assertEqual(created["status"], "ACTIVE")
        self.assertEqual(created["effective_from"], "2026-01-01")
        self.assertEqual(created["effective_to"], "2026-12-31")
        self.assertEqual(listed[0]["doc_type"], "STANDARD")
        self.assertEqual(listed[0]["doc_version"], "V1.2")

    def test_upload_knowledge_document_syncs_only_default_profile_on_upload(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-default-sync")
        knowledge_dir = self.tmp_path / "knowledge_default_sync"
        knowledge_dir.mkdir(parents=True, exist_ok=True)
        upload = UploadFile(filename="spec.txt", file=io.BytesIO("Spec content".encode("utf-8")))
        captured: dict[str, object] = {}

        def fake_sync(project_id: str, document: dict, *, profile_ids: list[str] | None = None) -> dict:
            captured["project_id"] = project_id
            captured["profile_ids"] = list(profile_ids or [])
            return {
                **document,
                "retrieval_sync_status": "READY",
                "retrieval_sync_error": "",
                "retrieval_synced_at": "2026-04-02T00:00:00+00:00",
                "retrieval_syncs": [
                    {
                        "profile_id": "azure",
                        "provider": "azure",
                        "retrieval_backend": "azure_sdk",
                        "status": "READY",
                        "synced_at": "2026-04-02T00:00:00+00:00",
                        "chunk_count": 1,
                        "payload_path": str(self.tmp_path / "payload.json"),
                    }
                ],
            }

        with patch.object(main, "KNOWLEDGE_DIR", knowledge_dir), patch.object(
            main,
            "_default_knowledge_sync_profile_ids",
            return_value=["azure"],
        ), patch.object(
            main,
            "_sync_knowledge_document_targets",
            side_effect=fake_sync,
        ):
            created = asyncio.run(
                main.upload_knowledge_document(
                    project["id"],
                    file=upload,
                    doc_type="REFERENCE",
                    doc_version="",
                    status="ACTIVE",
                    effective_from="",
                    effective_to="",
                )
            )

        self.assertEqual(captured["project_id"], project["id"])
        self.assertEqual(captured["profile_ids"], ["azure"])
        self.assertEqual(created["retrieval_syncs"][0]["profile_id"], "azure")

    def test_runtime_ai_provider_setup_endpoint_can_switch_default_profile(self) -> None:
        with TestClient(main.app) as client:
            response = client.post(
                "/api/runtime/ai-provider-setup",
                json={
                    "default_profile_id": "aliyun",
                    "aliyun": {
                        "api_key": "aliyun-key",
                        "chat_model": "qwen3.5-plus",
                        "embedding_model": "text-embedding-v4",
                        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
                        "qdrant_url": "http://localhost:6333",
                        "qdrant_collection": "reqpilot-knowledge-aliyun-local",
                    },
                },
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["default_profile_id"], "aliyun")
        self.assertTrue(payload["aliyun"]["api_key_configured"])
        self.assertEqual(payload["aliyun"]["embedding_model"], "text-embedding-v4")

    def test_runtime_ai_provider_setup_endpoint_rejects_non_ascii_aliyun_api_key(self) -> None:
        with TestClient(main.app) as client:
            response = client.post(
                "/api/runtime/ai-provider-setup",
                json={
                    "default_profile_id": "aliyun",
                    "aliyun": {
                        "api_key": "sk-测试key",
                        "chat_model": "qwen3.5-plus",
                        "embedding_model": "text-embedding-v4",
                        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
                        "qdrant_url": "http://localhost:6333",
                        "qdrant_collection": "reqpilot-knowledge-aliyun-local",
                    },
                },
            )

        self.assertEqual(response.status_code, 422)
        self.assertIn("DashScope API Key contains non-ASCII characters", response.json()["detail"])

    def test_upload_knowledge_document_requires_ready_default_profile(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-profile-gate")
        upload = UploadFile(filename="spec.txt", file=io.BytesIO("Spec content".encode("utf-8")))

        with patch.dict(
            os.environ,
            {
                "REQPILOT_ALIYUN_API_KEY": "",
                "REQPILOT_DASHSCOPE_API_KEY": "",
                "DASHSCOPE_API_KEY": "",
                "REQPILOT_ALIYUN_QDRANT_API_KEY": "",
                "REQPILOT_QDRANT_API_KEY": "",
                "QDRANT_API_KEY": "",
            },
            clear=False,
        ):
            with TestClient(main.app) as client:
                setup_resp = client.post(
                    "/api/runtime/ai-provider-setup",
                    json={
                        "default_profile_id": "aliyun",
                        "aliyun": {
                            "chat_model": "qwen3.5-plus",
                            "embedding_model": "text-embedding-v4",
                            "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
                            "qdrant_url": "http://localhost:6333",
                            "qdrant_collection": "reqpilot-knowledge-aliyun-local",
                        },
                    },
                )
                self.assertEqual(setup_resp.status_code, 200)

            with self.assertRaises(HTTPException) as ctx:
                asyncio.run(
                    main.upload_knowledge_document(
                        project["id"],
                        file=upload,
                        doc_type="REFERENCE",
                        doc_version="",
                        status="ACTIVE",
                        effective_from="",
                        effective_to="",
                    )
                )

        self.assertEqual(ctx.exception.status_code, 409)
        self.assertIn("Aliyun knowledge sync is not ready yet", str(ctx.exception.detail))

    def test_build_ai_profile_catalog_marks_unsynced_profile_as_pending(self) -> None:
        azure_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="azure_sdk",
            top_k=5,
            trace_dir=self.tmp_path / "trace-azure",
            azure_auth_mode="rbac",
            azure_allow_interactive_login=False,
            azure_search_endpoint="https://search.example.search.windows.net",
            azure_search_api_key="",
            azure_search_index_name="reqpilot-knowledge",
            azure_openai_endpoint="https://openai.example.openai.azure.com",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="gpt-5.4",
            azure_openai_embedding_deployment="text-embedding-3-large",
            azure_openai_embedding_dimensions=1536,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="azure",
            provider="azure",
            profile_label="Azure",
            chat_model="gpt-5.4",
            embedding_model="text-embedding-3-large",
        )
        aliyun_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="qdrant",
            top_k=5,
            trace_dir=self.tmp_path / "trace-aliyun",
            azure_auth_mode="",
            azure_allow_interactive_login=False,
            azure_search_endpoint="",
            azure_search_api_key="",
            azure_search_index_name="",
            azure_openai_endpoint="",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="",
            azure_openai_embedding_deployment="",
            azure_openai_embedding_dimensions=0,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="aliyun",
            provider="aliyun",
            profile_label="Aliyun",
            chat_model="qwen3.5-plus",
            embedding_model="text-embedding-v4",
            dashscope_api_key="aliyun-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            qdrant_url="http://localhost:6333",
            qdrant_collection_name="reqpilot-knowledge-aliyun-local",
        )

        with patch.object(main, "list_ai_copilot_profiles", return_value=[azure_config, aliyun_config]), patch.object(
            main,
            "load_ai_copilot_config",
            return_value=azure_config,
        ), patch.object(
            main,
            "build_ai_profile_check",
            side_effect=[
                {"status": "ok", "detail": "Azure ready.", "runnable": True, "missing": []},
                {"status": "ok", "detail": "Aliyun ready.", "runnable": True, "missing": []},
            ],
        ):
            payload = main._build_ai_profile_catalog(
                [
                    {
                        "id": "kb-1",
                        "retrieval_syncs": [
                            {"profile_id": "azure", "status": "READY"},
                        ],
                    }
                ]
            )

        by_id = {item["id"]: item for item in payload}
        self.assertEqual(by_id["azure"]["knowledge_sync"]["status"], "READY")
        self.assertEqual(by_id["aliyun"]["knowledge_sync"]["status"], "PENDING")
        self.assertTrue(by_id["aliyun"]["ready"])

    def test_build_ai_profile_catalog_uses_runtime_readiness_for_ready_flag(self) -> None:
        azure_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="azure_sdk",
            top_k=5,
            trace_dir=self.tmp_path / "trace-azure-runtime",
            azure_auth_mode="rbac",
            azure_allow_interactive_login=False,
            azure_search_endpoint="https://search.example.search.windows.net",
            azure_search_api_key="",
            azure_search_index_name="reqpilot-knowledge",
            azure_openai_endpoint="https://openai.example.openai.azure.com",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="gpt-5.4",
            azure_openai_embedding_deployment="text-embedding-3-large",
            azure_openai_embedding_dimensions=1536,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="azure",
            provider="azure",
            profile_label="Azure",
            chat_model="gpt-5.4",
            embedding_model="text-embedding-3-large",
        )
        aliyun_config = AICopilotConfig(
            assessment_mode="llm_shadow",
            retrieval_backend="qdrant",
            top_k=5,
            trace_dir=self.tmp_path / "trace-aliyun-runtime",
            azure_auth_mode="",
            azure_allow_interactive_login=False,
            azure_search_endpoint="",
            azure_search_api_key="",
            azure_search_index_name="",
            azure_openai_endpoint="",
            azure_openai_api_key="",
            azure_openai_api_version="2024-10-21",
            azure_openai_chat_deployment="",
            azure_openai_embedding_deployment="",
            azure_openai_embedding_dimensions=0,
            azure_search_semantic_configuration="",
            llm_timeout_seconds=45.0,
            profile_id="aliyun",
            provider="aliyun",
            profile_label="Aliyun",
            chat_model="qwen3.5-plus",
            embedding_model="text-embedding-v4",
            dashscope_api_key="aliyun-key",
            dashscope_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
            qdrant_url="http://localhost:6333",
            qdrant_collection_name="reqpilot-knowledge-aliyun-local",
        )

        with patch.object(main, "list_ai_copilot_profiles", return_value=[azure_config, aliyun_config]), patch.object(
            main,
            "load_ai_copilot_config",
            return_value=aliyun_config,
        ), patch.object(
            main,
            "build_ai_profile_check",
            side_effect=[
                {"status": "ok", "detail": "Azure ready.", "runnable": True, "missing": []},
                {
                    "status": "error",
                    "detail": "Qdrant local service is unreachable.",
                    "runnable": False,
                    "missing": ["ALIYUN_QDRANT_UNREACHABLE"],
                },
            ],
        ):
            payload = main._build_ai_profile_catalog(
                [
                    {
                        "id": "kb-1",
                        "retrieval_syncs": [],
                    }
                ]
            )

        by_id = {item["id"]: item for item in payload}
        self.assertFalse(by_id["aliyun"]["ready"])
        self.assertEqual(by_id["aliyun"]["runtime_status"], "error")
        self.assertEqual(by_id["aliyun"]["runtime_missing"], ["ALIYUN_QDRANT_UNREACHABLE"])
        self.assertEqual(by_id["aliyun"]["missing"], ["ALIYUN_QDRANT_UNREACHABLE"])

    def test_resync_knowledge_document_route_uses_default_profile(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-resync")
        knowledge_path = self.tmp_path / "knowledge-resync.txt"
        knowledge_path.write_text("knowledge", encoding="utf-8")
        knowledge_document = db.add_knowledge_document(
            project["id"],
            "knowledge-resync.txt",
            "txt",
            str(knowledge_path),
            "knowledge",
        )

        captured: dict[str, object] = {}

        def fake_sync(project_id: str, document: dict, *, profile_ids: list[str] | None = None) -> dict:
            captured["project_id"] = project_id
            captured["document_id"] = document["id"]
            captured["profile_ids"] = list(profile_ids or [])
            return {
                **document,
                "retrieval_syncs": [
                    {
                        "profile_id": "aliyun",
                        "status": "READY",
                    }
                ],
            }

        with patch.object(main, "_default_knowledge_sync_profile_ids", return_value=["aliyun"]), patch.object(
            main,
            "_sync_knowledge_document_targets",
            side_effect=fake_sync,
        ):
            with TestClient(main.app) as client:
                response = client.post(
                    f"/api/projects/{project['id']}/knowledge-documents/{knowledge_document['id']}/resync"
                )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(captured["project_id"], project["id"])
        self.assertEqual(captured["document_id"], knowledge_document["id"])
        self.assertEqual(captured["profile_ids"], ["aliyun"])
        self.assertEqual(response.json()["retrieval_syncs"][0]["status"], "READY")

    def test_upload_knowledge_document_route_accepts_multipart_payload(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-route")
        knowledge_dir = self.tmp_path / "knowledge_route"
        knowledge_dir.mkdir(parents=True, exist_ok=True)

        with patch.object(main, "KNOWLEDGE_DIR", knowledge_dir):
            with TestClient(main.app) as client:
                response = client.post(
                    f"/api/projects/{project['id']}/knowledge-documents",
                    files={"file": ("spec.txt", io.BytesIO(b"Spec content"), "text/plain")},
                    data={
                        "doc_type": "STANDARD",
                        "doc_version": "V1.2",
                        "status": "ACTIVE",
                        "effective_from": "2026-01-01",
                        "effective_to": "2026-12-31",
                    },
                )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["doc_type"], "STANDARD")
        self.assertEqual(payload["doc_version"], "V1.2")
        self.assertEqual(payload["status"], "ACTIVE")

    def test_upload_knowledge_document_route_rejects_duplicate_file_name(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-duplicate")
        knowledge_dir = self.tmp_path / "knowledge_duplicate"
        knowledge_dir.mkdir(parents=True, exist_ok=True)

        with patch.object(main, "KNOWLEDGE_DIR", knowledge_dir):
            with TestClient(main.app) as client:
                first = client.post(
                    f"/api/projects/{project['id']}/knowledge-documents",
                    files={"file": ("spec.txt", io.BytesIO(b"Spec content"), "text/plain")},
                )
                duplicate = client.post(
                    f"/api/projects/{project['id']}/knowledge-documents",
                    files={"file": ("spec.txt", io.BytesIO(b"Updated spec content"), "text/plain")},
                )

        self.assertEqual(first.status_code, 200)
        self.assertEqual(duplicate.status_code, 409)
        self.assertIn("already exists", duplicate.json()["detail"])
        self.assertEqual(len(db.list_knowledge_documents(project["id"])), 1)

    def test_delete_knowledge_document_route_removes_files_and_bumps_version(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-delete")
        knowledge_path = self.tmp_path / "knowledge-delete.txt"
        knowledge_path.write_text("knowledge", encoding="utf-8")
        payload_dir = self.tmp_path / "knowledge_payloads"
        payload_dir.mkdir(parents=True, exist_ok=True)
        payload_path = payload_dir / "payload-delete.json"
        payload_path.write_text('{"chunks": []}', encoding="utf-8")
        knowledge_document = db.add_knowledge_document(
            project_id=project["id"],
            file_name="knowledge-delete.txt",
            file_type="txt",
            file_path=str(knowledge_path),
            extracted_text="knowledge",
        )
        db.upsert_knowledge_document_sync(
            knowledge_document["id"],
            profile_id="azure",
            provider="azure",
            retrieval_backend="azure_sdk",
            status="READY",
            synced_at="2026-04-03T06:00:00+00:00",
            chunk_count=4,
            payload_path=str(payload_path),
        )

        with patch.object(main, "INDEX_PAYLOAD_DIR", payload_dir):
            with TestClient(main.app) as client:
                response = client.delete(
                    f"/api/projects/{project['id']}/knowledge-documents/{knowledge_document['id']}"
                )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["removed_count"], 1)
        self.assertEqual(payload["knowledge_base_version"], 1)
        self.assertFalse(knowledge_path.exists())
        self.assertFalse(payload_path.exists())
        self.assertEqual(db.list_knowledge_documents(project["id"]), [])
        self.assertEqual(db.list_knowledge_document_syncs(knowledge_document["id"]), [])

    def test_delete_knowledge_document_route_removes_legacy_same_name_duplicates(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-delete-duplicates")
        first_path = self.tmp_path / "legacy-dup-a.txt"
        second_path = self.tmp_path / "legacy-dup-b.txt"
        first_path.write_text("alpha", encoding="utf-8")
        second_path.write_text("beta", encoding="utf-8")
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(first_path),
            extracted_text="alpha",
        )
        latest_document = db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(second_path),
            extracted_text="beta",
        )

        with TestClient(main.app) as client:
            response = client.delete(
                f"/api/projects/{project['id']}/knowledge-documents/{latest_document['id']}"
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["removed_count"], 2)
        self.assertFalse(first_path.exists())
        self.assertFalse(second_path.exists())
        self.assertEqual(db.list_knowledge_documents(project["id"]), [])

    def test_list_knowledge_documents_collapses_existing_duplicate_file_names(self) -> None:
        project, _document = self._create_project_with_document("-knowledge-list-dedupe")
        first_path = self.tmp_path / "dup-a.txt"
        second_path = self.tmp_path / "dup-b.txt"
        first_path.write_text("alpha", encoding="utf-8")
        second_path.write_text("beta", encoding="utf-8")
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(first_path),
            extracted_text="alpha",
        )
        db.add_knowledge_document(
            project_id=project["id"],
            file_name="platform-spec.txt",
            file_type="txt",
            file_path=str(second_path),
            extracted_text="beta",
        )

        listed = main.list_knowledge_documents(project["id"])

        self.assertEqual(len(db.list_knowledge_documents(project["id"])), 2)
        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0]["file_name"], "platform-spec.txt")

    def test_export_works_without_cached_preview(self) -> None:
        project, document = self._create_project_with_document("-export-no-preview")
        db.insert_requirement(
            project_id=project["id"],
            sor_document_id=document["id"],
            req_code="REQ0001",
            source_page=1,
            source_anchor="Page 1 | Section 1",
            requirement_text="无 preview 缓存时仍可导出。",
            raw_text="无 preview 缓存时仍可导出。",
            display_text="无 preview 缓存时仍可导出。",
            requirement_type="FUNCTIONAL",
            priority="MEDIUM",
        )
        db.increment_sor_document_requirement_set_version(document["id"], "REVIEWING")
        main.lock_document_extraction(project["id"], document["id"])

        result = main.export_project(project["id"], main.ExportRequest(document_id=document["id"]))
        self.assertTrue((self.tmp_path / "exports" / result["file_name"]).exists())

    def test_export_translates_database_lock_to_retryable_http_error(self) -> None:
        project, document = self._create_project_with_document("-export-busy")
        with patch.object(main, "build_extraction_preview", return_value=_fake_preview()):
            main.analyze_document(project["id"], main.AnalyzeRequest(document_id=document["id"]))
            main.lock_document_extraction(project["id"], document["id"])
            with patch.object(db, "create_export_job", side_effect=sqlite3.OperationalError("database is locked")):
                with self.assertRaises(HTTPException) as export_exc:
                    main.export_project(project["id"], main.ExportRequest(document_id=document["id"]))

        self.assertEqual(export_exc.exception.status_code, 503)
        self.assertIn("busy", str(export_exc.exception.detail).lower())

    def test_exporter_renders_markdown_tables_as_images(self) -> None:
        file_name, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0001",
                    "req_code": "REQ0001",
                    "parent_section_label": "5.1 Test Setup",
                    "raw_text": (
                        "试验条件如下：\n\n"
                        "| Item | Condition |\n"
                        "| --- | --- |\n"
                        "| Voltage | 12V |\n"
                        "| Current | 3A |"
                    ),
                    "final_decision": "ACCEPT",
                    "engineer_reason": "Reviewed",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        self.assertTrue(file_name.endswith(".xlsx"))
        workbook = load_workbook(file_path)
        sheet = workbook["ComplianceMatrix"]
        meta = self._version_meta_values(workbook)

        self.assertEqual(sheet["A4"].value, "REQ0001")
        self.assertEqual(sheet["C4"].value, "试验条件如下：")
        self.assertIsNone(sheet["C5"].value)
        self.assertEqual(sheet["D4"].value, "ACCEPT")
        self.assertEqual(sheet["E4"].value, "Reviewed")
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(sheet._images[0].anchor._from.col, 2)
        self.assertEqual(sheet._images[0].anchor._from.row, 4)
        merged_ranges = {str(item) for item in sheet.merged_cells.ranges}
        self.assertNotIn("C3:D3", merged_ranges)
        self.assertEqual(meta["TableImageCount"], 1)
        self.assertEqual(meta["TableFallbackTextCount"], 0)

    def test_exporter_preserves_paragraph_table_paragraph_order(self) -> None:
        _, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0002",
                    "req_code": "REQ0002",
                    "parent_section_label": "5.2 Sequence",
                    "raw_text": (
                        "第一段说明。\n\n"
                        "| Step | Value |\n"
                        "| --- | --- |\n"
                        "| Torque | 25Nm |\n"
                        "| Time | 30s |\n\n"
                        "第二段说明。"
                    ),
                    "final_decision": "CHECK",
                    "engineer_reason": "Ordered",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        workbook = load_workbook(file_path)
        sheet = workbook["ComplianceMatrix"]

        self.assertEqual(sheet["A4"].value, "REQ0002")
        self.assertEqual(sheet["C4"].value, "第一段说明。")
        self.assertIsNone(sheet["C5"].value)
        self.assertEqual(sheet["C6"].value, "第二段说明。")
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(sheet._images[0].anchor._from.col, 2)
        self.assertEqual(sheet._images[0].anchor._from.row, 4)
        merged_ranges = {str(item) for item in sheet.merged_cells.ranges}
        self.assertIn("A4:A6", merged_ranges)
        self.assertIn("B4:B6", merged_ranges)
        self.assertIn("D4:D6", merged_ranges)
        self.assertIn("E4:E6", merged_ranges)

    def test_exporter_invalid_markdown_table_falls_back_to_text(self) -> None:
        _, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0003",
                    "req_code": "REQ0003",
                    "parent_section_label": "5.3 Fallback",
                    "raw_text": "| Item | Condition |\n| --- | --- |\n| Too | Many | Cells |",
                    "final_decision": "REVIEW",
                    "engineer_reason": "Fallback",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        workbook = load_workbook(file_path)
        sheet = workbook["ComplianceMatrix"]
        meta = self._version_meta_values(workbook)

        self.assertEqual(len(sheet._images), 0)
        self.assertEqual(
            sheet["C4"].value,
            "| Item | Condition |\n| --- | --- |\n| Too | Many | Cells |",
        )
        self.assertEqual(meta["TableImageCount"], 0)
        self.assertEqual(meta["TableFallbackTextCount"], 1)

    def test_exporter_renders_pipe_tables_without_separator_as_images(self) -> None:
        _, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0004",
                    "req_code": "REQ0004",
                    "parent_section_label": "3.2.2 产品详细描述",
                    "raw_text": (
                        "1）前减振器布置要求\n"
                        "|序号|项目|要求|\n"
                        "|1|基本尺寸|活塞桶直径Φ32-48mm|\n"
                        "|2|电磁阀|前后电磁阀处于同一高度|\n"
                        "表格后说明。"
                    ),
                    "final_decision": "ACCEPT",
                    "engineer_reason": "Pipe table",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        workbook = load_workbook(file_path)
        sheet = workbook["ComplianceMatrix"]
        meta = self._version_meta_values(workbook)

        self.assertEqual(sheet["C4"].value, "1）前减振器布置要求")
        self.assertIsNone(sheet["C5"].value)
        self.assertEqual(sheet["C6"].value, "表格后说明。")
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(sheet._images[0].anchor._from.col, 2)
        self.assertEqual(sheet._images[0].anchor._from.row, 4)
        self.assertEqual(meta["TableImageCount"], 1)
        self.assertEqual(meta["TableFallbackTextCount"], 0)

    def test_exporter_keeps_multiple_table_images_in_requirement_order(self) -> None:
        _, file_path = export_compliance_matrix(
            project_name="Workspace Test",
            customer_name="Customer",
            requirements=[
                {
                    "display_req_code": "REQ0005",
                    "req_code": "REQ0005",
                    "parent_section_label": "6.1 多表顺序",
                    "raw_text": (
                        "第一段说明。\n"
                        "|项|值|\n"
                        "|A|1|\n"
                        "|B|2|\n"
                        "中间说明。\n"
                        "|项|值|\n"
                        "|C|3|\n"
                        "|D|4|\n"
                        "最后说明。"
                    ),
                    "final_decision": "ACCEPT",
                    "engineer_reason": "Multi table",
                }
            ],
            export_dir=self.tmp_path / "exports",
        )

        workbook = load_workbook(file_path)
        sheet = workbook["ComplianceMatrix"]

        self.assertEqual(sheet.max_column, 5)
        self.assertEqual(sheet["C4"].value, "第一段说明。")
        self.assertEqual(sheet["C6"].value, "中间说明。")
        self.assertEqual(sheet["C8"].value, "最后说明。")
        self.assertEqual(len(sheet._images), 2)
        self.assertEqual(sheet._images[0].anchor._from.col, 2)
        self.assertEqual(sheet._images[0].anchor._from.row, 4)
        self.assertEqual(sheet._images[1].anchor._from.col, 2)
        self.assertEqual(sheet._images[1].anchor._from.row, 6)


if __name__ == "__main__":
    unittest.main()
