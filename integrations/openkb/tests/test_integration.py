from argparse import Namespace
from copy import deepcopy
from datetime import date
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from fastapi import FastAPI, Header, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

import artifact_tools
from portable_fixtures import make_fixture, item_decisions, write
from reqpilotx_openkb.api import create_router
from reqpilotx_openkb.errors import AdapterError
from reqpilotx_openkb.knowledge_store import FolderKnowledgeStore
from reqpilotx_openkb.service import OpenKBAnalysisService

TEXT = "SYNTHETIC TEST ONLY. Valve opened in 28 ms and closed in 35 ms at 12 V and 23 C. This invented result applies to sample A only.\n"
META = {"test.txt": {"doc_type": "CASE", "authority": "INTERNAL_CONTROLLED", "claim_role": "TEST_RESULT", "applicability": "SYNTHETIC sample A, 12 V, 23 C only"}}


@pytest.fixture
def service(tmp_path, request):
    source, capture, map_input = make_fixture(tmp_path / "source", **getattr(request, "param", {}))
    inp = tmp_path / "map-input.json"; write(inp, map_input)
    map_dir = tmp_path / "map"
    artifact_tools.compile_map(Namespace(source=source, source_ir=capture / "source_ir.json", input=inp, output_dir=map_dir))
    items = tmp_path / "items.json"; write(items, item_decisions(map_dir))
    baseline = tmp_path / "baseline"
    artifact_tools.compile_items(Namespace(map_dir=map_dir, input=items, baseline_dir=None, output_dir=baseline))
    store = FolderKnowledgeStore(tmp_path / "kb", "synthetic-project", "synthetic-kb", as_of=date(2026, 10, 2))
    store.upload_files({"test.txt": TEXT.encode()}, metadata=META, declaration="HOST synthetic integration fixture; no customer evidence")
    return OpenKBAnalysisService(store, baseline, source, tmp_path / "analysis")


def begin(service):
    return service.begin(expected_generation=service.store.current_generation,
        expected_baseline_sha256=json.loads((service.baseline_dir / "source-requirement-baseline.json").read_text())["artifact_sha256"])


def envelope(session, evidence_id=None):
    decisions = []
    for index, unit_id in enumerate(session["unit_ids"]):
        matched = index == 0 and evidence_id is not None
        decisions.append({"unit_id": unit_id, "decision": "COMPLIANT" if matched else "PENDING_EVIDENCE",
            "reason": "SYNTHETIC: original measured timing supports both fixture thresholds." if matched else "缺少本单元的内部产品或项目依据。",
            "evidence_ids": [evidence_id] if matched else [], "conditions": ["SYNTHETIC sample A, 12 V, 23 C only"] if matched else [],
            "conflicts": [], "missing": [] if matched else ["适用的受控产品或项目证据"],
            "clarification_question": "" if matched else "请补充本单元对应的受控产品或项目证据。"})
    return {"schema_version": "reqpilotx-openkb-assessment/v1", "project_id": session["project_id"], "kb": session["kb"],
        "generation": session["generation"], "packet_sha256": session["packet_sha256"], "decisions": decisions}


def read_original(service, session):
    index = service.source_index(session["session_id"])
    return service.read_original(session["session_id"], index["documents"][0]["document_id"])


def test_actual_native_analysis_and_same_version_excel(service):
    before = {p.name: p.read_bytes() for p in service.baseline_dir.iterdir() if p.is_file()}
    session = begin(service)
    original = read_original(service, session)
    evidence = original["evidence"][0]
    result = service.compile(session["session_id"], envelope(session, evidence["evidence_id"]))
    assert result["status"] == "COMPLETE" and result["pending_unit_count"] == len(session["unit_ids"]) - 1
    assert result["human_acceptance"] == "NOT_CLAIMED"
    directory, receipt = service.export(session["session_id"], result["assessment_sha256"])
    assert receipt["generation"] == session["generation"]
    assert receipt["baseline_sha256"] == session["source_binding"]["baseline_sha256"]
    assert hashlib.sha256((directory / "matrix.xlsx").read_bytes()).hexdigest() == receipt["xlsx_sha256"]
    native_receipt = json.loads((directory / "delivery-receipt.json").read_text())
    assert native_receipt["xlsx_sha256"] == receipt["native_xlsx_sha256"]
    snapshot = json.loads((directory / "matrix-snapshot.json").read_text())
    assert any((n.get("ai_assessment") or {}).get("decision") == "COMPLIANT" for n in snapshot["nodes"])
    book = load_workbook(directory / "matrix.xlsx")
    sheet = book["OpenKB原件依据"]
    values = list(sheet.values)
    actual = next(row for row in values[1:] if row[2] == evidence["evidence_id"])
    assert actual[7] == evidence["text"]
    assert json.loads(actual[6]) == evidence["source_ref"]
    assert actual[5] == evidence["reference"]["source_sha256"]
    assert actual[14] == session["generation"]
    assert actual[15] == result["assessment_sha256"] and actual[16] == "NOT_CLAIMED"
    # Native sheet values remain unchanged; original figure pixels survive annexing.
    native = load_workbook(directory / "matrix-native.xlsx")
    for name in native.sheetnames:
        assert list(book[name].values) == list(native[name].values)
    with ZipFile(directory / "matrix-native.xlsx") as z:
        native_images = {hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith("xl/media/")}
    with ZipFile(directory / "matrix.xlsx") as z:
        final_images = {hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith("xl/media/")}
    assert native_images and native_images == final_images
    assert before == {p.name: p.read_bytes() for p in service.baseline_dir.iterdir() if p.is_file()}
    prior = (directory / "matrix.xlsx").read_bytes()
    assert service.export(session["session_id"], result["assessment_sha256"])[0] == directory
    assert (directory / "matrix.xlsx").read_bytes() == prior
    book.close(); native.close()


@pytest.mark.parametrize("mutation", ["unread", "foreign_kb", "foreign_packet", "invented_fact", "human_acceptance", "duplicate_unit", "wiki_as_fact", "malformed_decision", "missing_full_compliance"])
def test_model_cannot_invent_original_or_acceptance(service, mutation):
    session = begin(service)
    original = read_original(service, session)
    proposal = envelope(session, original["evidence"][0]["evidence_id"])
    first = proposal["decisions"][0]
    if mutation == "unread": first["evidence_ids"] = ["original-evidence-unread"]
    elif mutation == "foreign_kb": proposal["kb"] = "another-kb"
    elif mutation == "foreign_packet": proposal["packet_sha256"] = "f" * 64
    elif mutation == "invented_fact": first["source_ref"] = {"page": 99}
    elif mutation == "human_acceptance": first["human_accepted"] = True
    elif mutation == "duplicate_unit": proposal["decisions"][1] = deepcopy(first)
    elif mutation == "wiki_as_fact": first["evidence_ids"] = ["summaries/test"]
    elif mutation == "malformed_decision": first["decision"] = ["COMPLIANT"]
    else: first["missing"] = ["Required cold-temperature validation is absent"]
    with pytest.raises(AdapterError): service.compile(session["session_id"], proposal)
    assert not (service.root / "sessions" / session["session_id"] / "results").exists()


def test_issued_ids_survive_service_restart(service):
    session = begin(service); original = read_original(service, session)
    restarted = OpenKBAnalysisService(service.store, service.baseline_dir, service.source_file, service.root)
    result = restarted.compile(session["session_id"], envelope(session, original["evidence"][0]["evidence_id"]))
    assert result["status"] == "COMPLETE"


@pytest.mark.parametrize("change", ["updated", "deleted", "baseline_file", "review_file", "original_bytes"])
def test_updates_invalidate_active_sessions_and_export(service, change):
    session = begin(service); original = read_original(service, session)
    result = service.compile(session["session_id"], envelope(session, original["evidence"][0]["evidence_id"]))
    output, _ = service.export(session["session_id"], result["assessment_sha256"])
    old_bytes = (output / "matrix.xlsx").read_bytes()
    if change == "updated": service.upload({"test.txt": (TEXT + "Different fixture version.").encode()}, metadata=META, declaration="Explicit fixture revision")
    elif change == "deleted": service.store.delete_document(original["document_id"], declaration="Explicit fixture removal")
    elif change == "baseline_file":
        path = service.baseline_dir / "table-previews.json"; path.write_text(path.read_text() + "\n")
    elif change == "review_file": (service.baseline_dir / "analysis-policy.json").write_text('{}')
    else: service.source_file.write_bytes(service.source_file.read_bytes() + b"\n")
    with pytest.raises(AdapterError): service.export(session["session_id"], result["assessment_sha256"])
    assert (output / "matrix.xlsx").read_bytes() == old_bytes  # History remains immutable.
    with pytest.raises(AdapterError): service.compile(session["session_id"], envelope(session))


def test_unknown_authority_can_only_form_pending_proposal(service):
    service.upload({"test.txt": TEXT.encode()}, metadata={"test.txt": {"doc_type": "CASE"}}, declaration="Unknown-authority fixture")
    session = begin(service); original = read_original(service, session)
    with pytest.raises(ValueError, match="authority-verified"):
        service.compile(session["session_id"], envelope(session, original["evidence"][0]["evidence_id"]))
    result = service.compile(session["session_id"], envelope(session))
    assert result["pending_unit_count"] == len(session["unit_ids"])


def test_tampered_registered_original_is_not_compilable(service):
    session = begin(service); original = read_original(service, session)
    manifest, _ = service.store.snapshot()
    path = service.store.inside(manifest["documents"][0]["source_path"])
    path.chmod(0o600); path.write_bytes(b"Forged edited result")
    with pytest.raises(AdapterError): service.compile(session["session_id"], envelope(session, original["evidence"][0]["evidence_id"]))


def test_excel_injection_stays_literal_original(service):
    text = '=HYPERLINK("https://invalid.example", "untrusted original")\n'
    service.upload({"test.txt": text.encode()}, metadata=META, declaration="Formula-injection fixture")
    session = begin(service); original = read_original(service, session)
    proposal = envelope(session)
    proposal["decisions"][0]["evidence_ids"] = [original["evidence"][0]["evidence_id"]]
    result = service.compile(session["session_id"], proposal)
    output, _ = service.export(session["session_id"], result["assessment_sha256"])
    book = load_workbook(output / "matrix.xlsx", data_only=False)
    cell = next(row[7] for row in book["OpenKB原件依据"].iter_rows(min_row=2) if row[2].value)
    assert cell.value.startswith("=") and cell.data_type == "s"
    book.close()


def test_api_requires_host_authorization_and_exports_exact_bound_result(service):
    def authorize(x_synthetic_access: str | None = Header(None)):
        if x_synthetic_access != "fixture-approved-project":
            raise HTTPException(status_code=403, detail="Project access required")
    app = FastAPI(); app.include_router(create_router(service, authorize=authorize))
    client = TestClient(app); base = "/api/v1/openkb-evidence"
    generation = service.store.current_generation
    binding = service._source_binding()
    body = {"expected_generation": generation, "expected_baseline_sha256": binding["baseline_sha256"]}
    assert client.post(base + "/sessions", json=body).status_code == 403
    headers = {"x-synthetic-access": "fixture-approved-project"}
    response = client.post(base + "/sessions", json=body, headers=headers)
    assert response.status_code == 200
    session = response.json(); path = base + "/sessions/" + session["session_id"]
    index = client.get(path + "/sources", headers=headers).json()
    document = index["documents"][0]["document_id"]
    original = client.get(path + "/originals/" + document, headers=headers).json()
    result = client.post(path + "/assessments", json=envelope(session, original["evidence"][0]["evidence_id"]), headers=headers)
    assert result.status_code == 200, result.text
    version = result.json()["assessment_sha256"]
    receipt = client.get(path + "/assessments/" + version + "/export-receipt", headers=headers)
    workbook = client.get(path + "/assessments/" + version + "/matrix.xlsx", headers=headers)
    assert workbook.status_code == 200 and hashlib.sha256(workbook.content).hexdigest() == receipt.json()["xlsx_sha256"]
    assert client.get(path + "/assessments/" + version + "/matrix.xlsx").status_code == 403
    assert client.post(base + "/sessions", json={**body, "source_path": "/etc/passwd"}, headers=headers).status_code == 422


def test_api_upload_rejects_traversal_and_identity_override(service):
    app = FastAPI(); app.include_router(create_router(service, authorize=lambda: None))
    client = TestClient(app); path = "/api/v1/openkb-evidence/sources/upload"
    old = service.store.current_generation
    response = client.post(path, files={"files": ("../escape.txt", b"Untrusted text")}, data={"declaration": "Test upload"})
    assert response.status_code == 409 and service.store.current_generation == old
    response = client.post(path, files={"files": ("test.txt", TEXT.encode())}, data={"declaration": "Test upload", "metadata": json.dumps({"test.txt": {"project_id": "other"}})})
    assert response.status_code == 409 and service.store.current_generation == old


def test_optional_router_never_has_implicit_authentication(service):
    with pytest.raises(TypeError): create_router(service, authorize=None)


def test_optional_live_runtime_unavailable_is_an_explicit_gap(service):
    from reqpilotx_openkb.runtime import build_analysis_query_agent
    session = begin(service)
    with pytest.raises(AdapterError, match="runtime is unavailable"):
        build_analysis_query_agent(service, session["session_id"], "host-approved-model")


def test_reused_export_revalidates_native_assessment_bytes(service):
    session = begin(service); original = read_original(service, session)
    result = service.compile(session["session_id"], envelope(session, original["evidence"][0]["evidence_id"]))
    output, _ = service.export(session["session_id"], result["assessment_sha256"])
    path = output.parent / "source-item-assessments.json"
    artifact = json.loads(path.read_text()); artifact["human_accepted"] = True
    path.write_text(json.dumps(artifact))
    with pytest.raises(ValueError): service.export(session["session_id"], result["assessment_sha256"])
