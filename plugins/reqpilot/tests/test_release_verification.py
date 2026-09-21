"""Regression tests for the R1 release verifier's own failure boundaries."""
from __future__ import annotations

import json
import shutil
import sys
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile


ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts"))

from run_required_unittest import evaluate_required_tests
from verify_engine_release import compare_business_outputs, release_status


class ReleaseVerificationTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(prefix="reqpilot-release-verifier-")
        self.root = Path(self.temp.name)
        self.left = self.root / "left"
        self.right = self.root / "right"
        self._write_result(self.left)
        shutil.copytree(self.left, self.right)

    def tearDown(self):
        self.temp.cleanup()

    def _write_result(self, root: Path):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.drawing.image import Image as WorksheetImage
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
        from PIL import Image

        for relative in ("map/document-map.json", "baseline/source-requirement-baseline.json", "delivery/matrix-snapshot.json"):
            path = root / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps({"contract": "same", "revision": 7}) + "\n", encoding="utf-8")
        workbook = Workbook()
        requirements = workbook.active
        requirements.title = "Requirements"
        requirements["A1"] = "2026-09-21T10:00:00"
        requirements["B1"] = "Stable business value"
        requirements["B1"].font = Font(name="Arial", bold=True, color="FF112233")
        requirements["B1"].fill = PatternFill("solid", fgColor="FFCCDD00")
        requirements["B1"].border = Border(left=Side(style="thin", color="FF000000"))
        requirements["B1"].alignment = Alignment(horizontal="center", wrap_text=True)
        requirements["B1"].protection = Protection(locked=True, hidden=False)
        requirements["C1"] = "=1+1"
        requirements.merge_cells("D1:E1")
        requirements["D1"] = "Merged"
        for row, value in enumerate((1, 2, 3), start=2):
            requirements.cell(row, 1, value)
        chart = BarChart()
        chart.title = "Stable chart"
        chart.add_data(Reference(requirements, min_col=1, min_row=2, max_row=4))
        requirements.add_chart(chart, "G2")
        image_buffer = BytesIO()
        with Image.new("RGB", (4, 4), (10, 20, 30)) as image:
            image.save(image_buffer, format="PNG")
        image_buffer.seek(0)
        requirements.add_image(WorksheetImage(image_buffer), "F1")
        second_image_buffer = BytesIO()
        with Image.new("RGB", (4, 4), (70, 80, 90)) as image:
            image.save(second_image_buffer, format="PNG")
        second_image_buffer.seek(0)
        requirements.add_image(WorksheetImage(second_image_buffer), "F10")
        requirements.row_dimensions[1].height = 24
        requirements.column_dimensions["B"].width = 32
        metadata = workbook.create_sheet("_ReqPilotXMeta")
        metadata["A30"] = "ExportedAt"
        metadata["B30"] = "2026-09-21T10:00:00"
        workbook.save(root / "delivery/matrix.xlsx")
        if not image_buffer.closed:
            image_buffer.close()
        if not second_image_buffer.closed:
            second_image_buffer.close()

    def _save_mutation(self, mutate):
        from openpyxl import load_workbook
        path = self.right / "delivery/matrix.xlsx"
        workbook = load_workbook(path)
        mutate(workbook)
        workbook.save(path)

    def assert_rejected(self):
        with self.assertRaises(ValueError):
            compare_business_outputs(self.left, self.right)

    def _rewrite_zip_member(self, path: Path, target: str, transform):
        replacement = path.with_suffix(".changed.xlsx")
        with ZipFile(path) as source, ZipFile(replacement, "w", ZIP_DEFLATED) as output:
            for item in source.infolist():
                data = source.read(item.filename)
                changed = transform(data) if item.filename == target else data
                if item.filename == target and changed == data:
                    raise AssertionError(f"Mutation did not change {target}")
                output.writestr(item, changed)
        replacement.replace(path)

    def _drawing_part(self, path: Path) -> str:
        with ZipFile(path) as archive:
            return next(name for name in archive.namelist() if name.startswith("xl/drawings/drawing") and name.endswith(".xml"))

    def test_only_contracted_export_timestamp_is_normalized(self):
        self._save_mutation(lambda wb: setattr(wb["_ReqPilotXMeta"]["B30"], "value", "2026-09-21T10:00:01"))
        result = compare_business_outputs(self.left, self.right)
        self.assertEqual(result["xlsx_normalized_timestamp_cells"], ["_ReqPilotXMeta!B30"])

    def test_business_iso_date_and_wrong_timestamp_label_are_rejected(self):
        self._save_mutation(lambda wb: setattr(wb["Requirements"]["A1"], "value", "2036-09-21T10:00:00"))
        self.assert_rejected()
        shutil.rmtree(self.right); shutil.copytree(self.left, self.right)
        self._save_mutation(lambda wb: setattr(wb["_ReqPilotXMeta"]["A30"], "value", "NotExportedAt"))
        self.assert_rejected()

    def test_font_fill_border_alignment_and_protection_changes_are_rejected(self):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
        mutations = (
            lambda cell: setattr(cell, "font", Font(name="Arial", italic=True)),
            lambda cell: setattr(cell, "fill", PatternFill("solid", fgColor="FFFF0000")),
            lambda cell: setattr(cell, "border", Border(right=Side(style="double"))),
            lambda cell: setattr(cell, "alignment", Alignment(horizontal="right")),
            lambda cell: setattr(cell, "protection", Protection(locked=False, hidden=True)),
        )
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                shutil.rmtree(self.right); shutil.copytree(self.left, self.right)
                self._save_mutation(lambda wb, operation=mutation: operation(wb["Requirements"]["B1"]))
                self.assert_rejected()

    def test_value_formula_merge_and_layout_changes_are_rejected(self):
        mutations = (
            lambda wb: setattr(wb["Requirements"]["B1"], "value", "Changed"),
            lambda wb: setattr(wb["Requirements"]["C1"], "value", "=2+2"),
            lambda wb: wb["Requirements"].unmerge_cells("D1:E1"),
            lambda wb: setattr(wb["Requirements"].row_dimensions[1], "height", 42),
        )
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                shutil.rmtree(self.right); shutil.copytree(self.left, self.right)
                self._save_mutation(mutation)
                self.assert_rejected()

    def test_media_and_chart_content_changes_are_rejected(self):
        workbook = self.right / "delivery/matrix.xlsx"
        with ZipFile(workbook) as archive:
            media = next(name for name in archive.namelist() if name.startswith("xl/media/"))
            chart = next(name for name in archive.namelist() if name.startswith("xl/charts/"))
        self._rewrite_zip_member(workbook, media, lambda data: data + b"changed")
        self.assert_rejected()
        shutil.rmtree(self.right); shutil.copytree(self.left, self.right)
        self._rewrite_zip_member(self.right / "delivery/matrix.xlsx", chart, lambda data: data.replace(b"Stable chart", b"Changed chart"))
        self.assert_rejected()

    def test_image_and_chart_position_changes_are_rejected(self):
        drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"

        def move(data: bytes, object_tag: str) -> bytes:
            root = ET.fromstring(data)
            anchor = next(item for item in list(root) if item.find(f"{{{drawing_ns}}}{object_tag}") is not None)
            row = anchor.find(f"{{{drawing_ns}}}from/{{{drawing_ns}}}row")
            row.text = str(int(row.text) + 20)
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)

        workbook = self.right / "delivery/matrix.xlsx"
        drawing = self._drawing_part(workbook)
        self._rewrite_zip_member(workbook, drawing, lambda data: move(data, "pic"))
        self.assert_rejected()
        shutil.rmtree(self.right); shutil.copytree(self.left, self.right)
        workbook = self.right / "delivery/matrix.xlsx"; drawing = self._drawing_part(workbook)
        self._rewrite_zip_member(workbook, drawing, lambda data: move(data, "graphicFrame"))
        self.assert_rejected()

    def test_image_resource_relationship_swap_is_rejected_with_same_media_set(self):
        drawing_main = "http://schemas.openxmlformats.org/drawingml/2006/main"
        relationships = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

        def swap(data: bytes) -> bytes:
            root = ET.fromstring(data)
            blips = root.findall(f".//{{{drawing_main}}}blip")
            self.assertEqual(len(blips), 2)
            key = f"{{{relationships}}}embed"
            blips[0].attrib[key], blips[1].attrib[key] = blips[1].attrib[key], blips[0].attrib[key]
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)

        workbook = self.right / "delivery/matrix.xlsx"
        drawing = self._drawing_part(workbook)
        self._rewrite_zip_member(workbook, drawing, swap)
        self.assert_rejected()

    def test_required_test_evaluation_rejects_skip_zero_partial_and_failure(self):
        required = ("Suite.test_a", "Suite.test_b")
        cases = {
            "all_pass": ({"tests_run": 2, "successful": ["pkg.Suite.test_a", "pkg.Suite.test_b"]}, "PASSED"),
            "all_skip": ({"tests_run": 2, "successful": [], "skipped": [{"test": "pkg.Suite.test_a", "reason": "missing"}, {"test": "pkg.Suite.test_b", "reason": "missing"}]}, "NOT_VERIFIED"),
            "partial_skip": ({"tests_run": 2, "successful": ["pkg.Suite.test_a"], "skipped": [{"test": "pkg.Suite.test_b", "reason": "missing"}]}, "NOT_VERIFIED"),
            "zero": ({"tests_run": 0, "successful": []}, "NOT_VERIFIED"),
            "failure": ({"tests_run": 2, "successful": ["pkg.Suite.test_a"], "failures": [{"test": "pkg.Suite.test_b", "detail": "assertion"}]}, "NOT_VERIFIED"),
        }
        for name, (summary, expected) in cases.items():
            with self.subTest(name=name):
                result = evaluate_required_tests(summary, required)
                self.assertEqual(result["status"], expected)
                self.assertEqual(release_status(result), "PASSED" if expected == "PASSED" else "PARTIAL")


if __name__ == "__main__":
    unittest.main()
