from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import asyncio
import hashlib
import io
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from app.application.errors import ApplicationError as HTTPException
from openpyxl import Workbook
from starlette.datastructures import UploadFile

from app import db
if __package__:
    from .application_under_test import main
else:
    from application_under_test import main
from app.main import app as http_app

from app.application.extraction import (
    DocumentIngestionRejected,
    IngestionRejectionCode,
    inspect_document_for_ingestion,
    prepare_document_for_ingestion,
)
from app.ports.document_adapter import ParsedDocumentPayload


main.app = http_app


class DocumentIngestionBoundaryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)

    def test_xlsx_is_authentic_but_blocked_until_safe_serving_adapter_exists(self) -> None:
        path = self.root / "requirements.xlsx"
        workbook = Workbook()
        workbook.active.append(["ID", "Requirement"])
        workbook.active.append(["SYS-1", "The controller shall start."])
        workbook.save(path)
        workbook.close()

        with self.assertRaises(DocumentIngestionRejected) as captured:
            prepare_document_for_ingestion(path)

        rejection = captured.exception
        self.assertEqual(rejection.code, IngestionRejectionCode.ADAPTER_UNAVAILABLE)
        self.assertTrue(rejection.report.processable)
        self.assertEqual(rejection.safe_detail()["detected_format"], "xlsx")

    def test_preflight_only_accepts_xlsx_without_invoking_compatibility_parser(
        self,
    ) -> None:
        path = self.root / "requirements.xlsx"
        workbook = Workbook()
        workbook.active.append(["ID", "Requirement"])
        workbook.active.append(["SYS-1", "The controller shall start."])
        workbook.save(path)
        workbook.close()

        with patch(
            "app.adapters.documents.isolated_test_parser.parse_document_with_diagnostics"
        ) as parser:
            report = inspect_document_for_ingestion(path)

        self.assertTrue(report.processable)
        self.assertEqual(report.identity.detected_format.value, "xlsx")
        parser.assert_not_called()

    def test_unknown_binary_is_rejected_before_legacy_parser_runs(self) -> None:
        path = self.root / "requirements.txt"
        path.write_bytes(b"\x00\x01\xffrequirements\x00")

        with patch(
            "app.adapters.documents.isolated_test_parser.parse_document_with_diagnostics"
        ) as parser:
            with self.assertRaises(DocumentIngestionRejected) as captured:
                prepare_document_for_ingestion(path)

        self.assertEqual(captured.exception.code, IngestionRejectionCode.PREFLIGHT_FAILED)
        parser.assert_not_called()

    def test_legacy_parser_text_fallback_is_rejected(self) -> None:
        path = self.root / "requirements.txt"
        path.write_text("The controller shall start.\n", encoding="utf-8")

        with patch(
            "app.adapters.documents.isolated_test_parser.parse_document_with_diagnostics",
            return_value={
                "text": "unsafe",
                "segments": [{"text": "unsafe"}],
                "file_type": "TEXT",
                "diagnostics": {"mode": "fallback_text", "error": "parser failed"},
            },
        ):
            with self.assertRaises(DocumentIngestionRejected) as captured:
                prepare_document_for_ingestion(path)

        self.assertEqual(captured.exception.code, IngestionRejectionCode.PARSE_FAILED)

    def test_valid_text_returns_raw_identity_and_legacy_payload(self) -> None:
        path = self.root / "requirements.txt"
        payload = "The controller shall start.\n".encode("utf-8")
        path.write_bytes(payload)

        prepared = prepare_document_for_ingestion(path)

        self.assertEqual(prepared.preflight.source_sha256, hashlib.sha256(payload).hexdigest())
        self.assertEqual(prepared.parsed.file_type, "TEXT")
        self.assertIn("controller shall start", prepared.parsed.text)

    def test_utf16_text_uses_the_same_strict_decoder_after_preflight(self) -> None:
        path = self.root / "requirements.txt"
        expected = "控制器应在 100 ms 内启动。"
        payload = f"{expected}\n".encode("utf-16")
        path.write_bytes(payload)

        prepared = prepare_document_for_ingestion(path)

        self.assertEqual(prepared.preflight.source_sha256, hashlib.sha256(payload).hexdigest())
        self.assertEqual(prepared.parsed.text, expected)
        self.assertNotIn("�", prepared.parsed.text)

    def test_source_changed_by_serving_parser_is_rejected_after_parse(self) -> None:
        path = self.root / "requirements.txt"
        path.write_text("The controller shall start.\n", encoding="utf-8")

        def mutate_source(
            _: bytes,
            *,
            source_name: str,
        ) -> ParsedDocumentPayload:
            self.assertEqual(source_name, path.name)
            path.write_text("Changed after preflight.\n", encoding="utf-8")
            return ParsedDocumentPayload(
                text="The controller shall start.",
                segments=(),
                file_type="TEXT",
                diagnostics={},
            )

        with patch(
            "app.adapters.documents.isolated_test_parser."
            "IsolatedTestDocumentAdapter.parse_bytes",
            side_effect=mutate_source,
        ):
            with self.assertRaises(DocumentIngestionRejected) as captured:
                prepare_document_for_ingestion(path)

        self.assertEqual(
            captured.exception.code,
            IngestionRejectionCode.SOURCE_CHANGED,
        )


class SorUploadPreflightIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.temporary_directory.name)
        self.exit_stack = ExitStack()
        self.addCleanup(self._cleanup)
        self.exit_stack.enter_context(patch.object(db, "STORAGE_DIR", self.root))
        self.exit_stack.enter_context(patch.object(db, "DB_PATH", self.root / "app.db"))
        self.exit_stack.enter_context(patch.object(main, "UPLOADS_DIR", self.root / "uploads"))
        db.init_db()
        self.project = db.create_project("Preflight", "Customer")

    def _cleanup(self) -> None:
        self.exit_stack.close()
        self.temporary_directory.cleanup()

    def test_rejected_upload_is_not_persisted_and_temporary_file_is_removed(self) -> None:
        payload = self._xlsx_bytes()
        upload = UploadFile(filename="requirements.xlsx", file=io.BytesIO(payload))

        with self.assertRaises(HTTPException) as captured:
            asyncio.run(main.upload_sor_document(self.project["id"], upload))

        self.assertEqual(captured.exception.status_code, 422)
        self.assertEqual(captured.exception.detail["code"], "adapter_unavailable")
        self.assertEqual(db.list_sor_documents(self.project["id"]), [])
        self.assertEqual(list((self.root / "uploads").glob("*")), [])

    def test_accepted_upload_persists_raw_identity_and_requirement_set_uses_it(self) -> None:
        payload = "The controller shall start.\n".encode("utf-8")
        upload = UploadFile(filename="requirements.txt", file=io.BytesIO(payload))

        document = asyncio.run(main.upload_sor_document(self.project["id"], upload))

        expected_hash = hashlib.sha256(payload).hexdigest()
        self.assertEqual(document["source_sha256"], expected_hash)
        self.assertEqual(document["source_size"], len(payload))
        self.assertEqual(document["declared_extension"], ".txt")
        self.assertEqual(document["detected_format"], "txt")
        self.assertEqual(document["preflight_schema_version"], "document-preflight/v2")
        requirement_set = db.create_requirement_set(self.project["id"], document["id"])
        self.assertEqual(requirement_set["source_hash"], expected_hash)

    def test_oversized_upload_is_rejected_without_target_or_part_file(self) -> None:
        payload = b"0123456789"
        upload = UploadFile(filename="requirements.txt", file=io.BytesIO(payload))

        with patch.object(main, "MAX_DOCUMENT_UPLOAD_BYTES", 8):
            with self.assertRaises(HTTPException) as captured:
                asyncio.run(main.upload_sor_document(self.project["id"], upload))

        self.assertEqual(captured.exception.status_code, 413)
        self.assertEqual(captured.exception.detail["code"], "upload_too_large")
        self.assertEqual(db.list_sor_documents(self.project["id"]), [])
        self.assertEqual(list((self.root / "uploads").glob("*")), [])

    def test_storage_name_normalization_cannot_fabricate_a_declared_extension(
        self,
    ) -> None:
        payload = "The controller shall start.\n".encode("utf-8")
        upload = UploadFile(filename="requirements.t!xt", file=io.BytesIO(payload))

        with self.assertRaises(HTTPException) as captured:
            asyncio.run(main.upload_sor_document(self.project["id"], upload))

        self.assertEqual(captured.exception.status_code, 422)
        self.assertEqual(
            captured.exception.detail["code"],
            "invalid_declared_extension",
        )
        self.assertEqual(db.list_sor_documents(self.project["id"]), [])
        self.assertEqual(list((self.root / "uploads").glob("*")), [])

    def test_database_failure_removes_validated_upload_file(self) -> None:
        payload = "The controller shall start.\n".encode("utf-8")
        upload = UploadFile(filename="requirements.txt", file=io.BytesIO(payload))

        with patch.object(db, "add_sor_document", side_effect=RuntimeError("db failed")):
            with self.assertRaisesRegex(RuntimeError, "db failed"):
                asyncio.run(main.upload_sor_document(self.project["id"], upload))

        self.assertEqual(list((self.root / "uploads").glob("*")), [])

    @staticmethod
    def _xlsx_bytes() -> bytes:
        workbook = Workbook()
        workbook.active.append(["ID", "Requirement"])
        workbook.active.append(["SYS-1", "The controller shall start."])
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()


if __name__ == "__main__":
    unittest.main()
