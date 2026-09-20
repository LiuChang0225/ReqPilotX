from __future__ import annotations
from .runtime_isolation import activate_test_runtime
activate_test_runtime()
from copy import deepcopy
from io import BytesIO
import unittest
from openpyxl import load_workbook
from app.application.agent_workflow import source_item_analysis as analysis
from app.application.agent_workflow import source_item_decomposition as decomposition
from app.application.agent_workflow import source_item_delivery as delivery
from app.application.agent_workflow import source_requirement_baseline as source
from app.services.response_matrix import _result
from .test_source_item_lineage import source_fixture
from .test_source_item_analysis import catalog, decision
from .test_source_requirement_baseline import _table_preview, _fragment


def table_context():
    return {
        'lt-bom': _table_preview(scope_id='lt-bom', title='产品信息', fragments=(
            _fragment('bom-fragment','T0001',page=2,rows=(('序号','名称'),('1','后减振器总成'),('2','后减振器子件'),('','前减振器'),('1','前减振器总成'))),)),
        'lt-dvp': _table_preview(scope_id='lt-dvp', title='主动前减振器试验大纲', fragments=(
            _fragment('dvp-one','T0002',page=10,rows=(('试验编号','试验名称'),('1','外观检查'),('2','跨页试验'),('','跨页试验子行'))),
            _fragment('dvp-two','T0003',page=11,rows=(('','跨页试验续行'),('3','耐久试验')))))
    }


class CarryPresentationTests(unittest.TestCase):
    def setUp(self):
        self.text='Standby current shall be <= 5 mA at 12 V and 25 C.'
        self.doc, _, _, self.baseline, self.receipt=source_fixture(self.text)
        self.sid=self.baseline.records[0].source_record_id
        self.tables=table_context()
        self.old_packet=analysis.build_source_item_analysis_packet(self.doc,self.baseline,self.tables,catalog())
        self.assessment=analysis.compile_source_item_assessments(self.old_packet,{'decisions':[decision(self.sid)]})
    def successor(self, kinds=None):
        change=source.SourceRequirementClassificationCoordinationDecisionV1(source_record_id=self.sid,
            requirementness='REQUIREMENT',requirement_kinds=kinds or self.baseline.records[0].requirement_kinds,
            confidence=.95,decision_summary='Reviewed source rationale.',reason_codes=('source_review',))
        return source.compile_source_requirement_classification_coordination(previous_artifact=self.baseline,
            previous_validation_receipt=self.receipt,decisions=(change,),model_id='synthetic')[0]
    def rebound_snapshot(self,kinds=None):
        new=self.successor(kinds)
        packet=analysis.build_source_item_analysis_packet(self.doc,new,self.tables,catalog())
        carry=analysis.rebind_source_item_assessments(self.old_packet,self.assessment,packet)
        snap=delivery.build_delivery_snapshot(self.doc,new,self.tables,assessments=carry['artifact'])
        return snap,next(n for n in snap['nodes'] if (n.get('source_item')or{}).get('source_item_id')==self.sid),carry['artifact']
    def workbook_row(self,snapshot):
        book=load_workbook(BytesIO(delivery.export_delivery_xlsx(snapshot)))
        row=next(r for r in book['ComplianceMatrix'].iter_rows() if r[0].value=='SRC-0001')
        return '\n'.join(str(c.value or '') for c in row)
    def test_inherited_current_advice_keeps_original_decision_and_exports_origin(self):
        snapshot,node,artifact=self.rebound_snapshot()
        self.assertEqual(node['analysis_result_state'],'CURRENT')
        self.assertEqual(artifact['records'][self.sid],self.assessment['records'][self.sid])
        self.assertEqual(node['ai_assessment']['reason'],self.assessment['records'][self.sid]['reason'])
        note=node['source_item']['revision_carry'][0]
        self.assertEqual(note['operation'],'SOURCE_VALIDATED_REBIND')
        self.assertEqual(note['origin_baseline_sha256'],self.baseline.artifact_sha256)
        row=self.workbook_row(snapshot)
        self.assertIn('沿用前版分析建议',row)
        self.assertIn('本轮未重新分析',row)
        self.assertIn(self.baseline.artifact_sha256,row)
    def test_stale_analysis_is_visible_without_becoming_a_new_pending_assessment(self):
        snapshot,node,_=self.rebound_snapshot(('QUALITY',))
        self.assertEqual(node['analysis_result_state'],'STALE_REFERENCE')
        self.assertIsNone(node['ai_assessment'])
        self.assertEqual(_result(node)[0],'需重新分析')
        row=self.workbook_row(snapshot)
        self.assertIn('输入一致性核验',row)
        self.assertNotIn('AI建议：符合',row)
        self.assertEqual(node['display_text'],self.text)
    def test_removed_derived_analysis_is_explained_at_its_original_parent(self):
        child='singular-requirement-'+'1'*32
        dep={'records':{},'identity_registry':{self.sid:{child:'REQ-0001-S01'}},
             'carry_forward':{'operation':'SOURCE_VALIDATED_REBIND','stale_parent_ids':[self.sid],'from_baseline_sha256':'b'*64}}
        art={'records':{},'carry_forward':{'operation':'SOURCE_VALIDATED_REBIND','stale_unit_ids':[child],'from_baseline_sha256':'b'*64}}
        notes=delivery.source_item_carry_notes(dep,art)
        self.assertEqual(notes[self.sid][1]['stale_unit_ids'],[child])
        self.assertIn('REQ-0001-S01',notes[self.sid][1]['message'])
        self.assertNotIn(child,notes)
    def test_stale_split_keeps_history_in_the_matrix_without_readding_old_children(self):
        text='Open the valve. Close the valve.'
        doc,_,_,base,receipt=source_fixture(text);sid=base.records[0].source_record_id
        split=decomposition.compile_source_item_decomposition(base,[{'source_item_id':sid,'action':'SPLIT','reason':'Distinct behaviors',
            'children':[{'text':text[a:b],'source_ranges':[{'segment_index':0,'start_offset':a,'end_offset':b}]} for a,b in [(0,15),(16,len(text))]]}])
        change=source.SourceRequirementClassificationCoordinationDecisionV1(source_record_id=sid,requirementness='REQUIREMENT',requirement_kinds=('QUALITY',),confidence=.95,decision_summary='Changed class.',reason_codes=('source_review',))
        new=source.compile_source_requirement_classification_coordination(previous_artifact=base,previous_validation_receipt=receipt,decisions=(change,),model_id='synthetic')[0]
        carried=decomposition.rebind_source_item_decomposition(base,new,split)['artifact']
        snap=delivery.build_delivery_snapshot(doc,new,self.tables,decomposition=carried)
        node=next(n for n in snap['nodes'] if (n.get('source_item')or{}).get('source_item_id')==sid)
        self.assertEqual(node['child_ids'],[])
        self.assertEqual(node['analysis_result_state'],'STALE_REFERENCE')
        self.assertIn('需检查拆分',self.workbook_row(snap))
    def test_retained_decomposition_notice_survives_parent_group_export(self):
        text='Open the valve. Close the valve.'
        doc,_,_,base,receipt=source_fixture(text);sid=base.records[0].source_record_id
        split=decomposition.compile_source_item_decomposition(base,[{'source_item_id':sid,'action':'SPLIT','reason':'Distinct behaviors',
            'children':[{'text':text[a:b],'source_ranges':[{'segment_index':0,'start_offset':a,'end_offset':b}]} for a,b in [(0,15),(16,len(text))]]}])
        change=source.SourceRequirementClassificationCoordinationDecisionV1(source_record_id=sid,requirementness='REQUIREMENT',requirement_kinds=base.records[0].requirement_kinds,confidence=.95,decision_summary='Clarified rationale.',reason_codes=('source_review',))
        new=source.compile_source_requirement_classification_coordination(previous_artifact=base,previous_validation_receipt=receipt,decisions=(change,),model_id='synthetic')[0]
        carried=decomposition.rebind_source_item_decomposition(base,new,split)['artifact']
        snap=delivery.build_delivery_snapshot(doc,new,self.tables,decomposition=carried)
        self.assertIn('沿用前版派生',self.workbook_row(snap))
        self.assertIn(base.artifact_sha256,self.workbook_row(snap))

    def test_fresh_result_and_unknown_operation_are_not_labelled_inherited(self):
        snap=delivery.build_delivery_snapshot(self.doc,self.baseline,self.tables,assessments=self.assessment)
        node=next(n for n in snap['nodes'] if (n.get('source_item')or{}).get('source_item_id')==self.sid)
        self.assertNotIn('revision_carry',node['source_item'])
        fake=deepcopy(self.assessment);fake['unit_input_receipts'][self.sid]['operation']='OTHER'
        self.assertEqual(delivery.source_item_carry_notes(assessments=fake),{})
