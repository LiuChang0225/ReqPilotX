from __future__ import annotations

from .runtime_isolation import activate_test_runtime
activate_test_runtime()

from copy import deepcopy
from io import BytesIO
import unittest

from openpyxl import load_workbook
from app.application.agent_workflow import source_requirement_baseline as core
from app.application.agent_workflow.source_item_decomposition import compile_source_item_decomposition
from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot, export_delivery_xlsx, _derived_source_number
from app.application.agent_workflow.source_item_visual import compile_source_item_annotation_updates, _seal_annotations
from .test_source_requirement_baseline import _reference_with_text_node


def fixture(*, body_only=False, extents=False, numbers=('6.1', '6.2')):
    texts = [(f'part-{i}', f'{number} {statement}', False)
             for i, (number, statement) in enumerate(zip(numbers, ('Open the valve.', 'Close the valve.')))]
    reference = _reference_with_text_node(node_id='clause', label='6', title='Commands', contents=tuple(texts))
    reference['document_tree']['nodes'] = [node for node in reference['document_tree']['nodes'] if node['node_id'] in {'root', 'clause'}]
    reference['logical_tables'] = []
    document_map = core.compile_document_map_artifact(reference)
    request = core.build_source_item_text_scope_request(document_map, node_id='clause')
    candidate = core.SourceItemTextSliceCandidateV1(source_sha256=request.source_sha256,
        document_map_artifact_id=document_map.artifact_id, document_map_artifact_sha256=document_map.artifact_sha256,
        request_sha256=request.request_sha256, scope_id=request.scope_id,
        items=(core.SourceItemTextBoundaryCandidateV1(item_key='commands', label='Valve commands', source_form='NUMBERED_CLAUSE',
            text_spans=tuple(core.SourceItemTextSpanV1(content_id=c.content_id, start_offset=0, end_offset=len(c.text)) for c in request.source_content),
            requirementness='REQUIREMENT', requirement_kinds=('FUNCTIONAL',), confidence=0.95,
            decision_summary='Keep complete independent operating commands.', reason_codes=('source_review',)),))
    baseline, _ = core.compile_source_requirement_baseline(document_map=document_map, requests=(), candidates=(),
        text_requests=(request,), text_candidates=(candidate,))
    source_id = baseline.records[0].source_record_id
    children, contexts, number_entries = [], [], []
    for i, content in enumerate(request.source_content):
        start = len(numbers[i]) + 1 if body_only else 0
        children.append({'text': content.text[start:], 'source_ranges': [{'segment_index': i, 'start_offset': start, 'end_offset': len(content.text)}]})
        if start:
            contexts.append({'source_range': {'segment_index': i, 'start_offset': 0, 'end_offset': start}, 'disposition': 'CONTEXT'})
        reference = {'kind': 'TEXT', 'content_id': content.content_id, 'start_offset': 0, 'end_offset': len(numbers[i])}
        entry = {'value': numbers[i], 'source_number_ref': reference}
        if extents:
            entry['source_extent'] = [{**reference, 'end_offset': len(content.text)}]
        number_entries.append(entry)
    decomposition = compile_source_item_decomposition(baseline, [{'source_item_id': source_id, 'action': 'SPLIT',
        'reason': 'Each command has a separate observable outcome.', 'children': children, 'context_ranges': contexts}])
    annotations = compile_source_item_annotation_updates(document_map, baseline, [{'source_item_id': source_id,
        'source_numbers': number_entries, 'reason': 'Preserve the exact original clause numbers.'}])
    return document_map, baseline, decomposition, annotations


class DeliveryOriginalNumbersTests(unittest.TestCase):
    def snapshot(self, values):
        document_map, baseline, decomposition, annotations = values
        return build_delivery_snapshot(document_map, baseline, {}, decomposition=decomposition, annotations=annotations)

    def test_parent_collection_and_direct_number_occurrences_share_typed_ui_and_export_fields(self):
        values = fixture()
        original = deepcopy(values[2:])
        snapshot = self.snapshot(values)
        parent = next(n for n in snapshot['nodes'] if n.get('source_item'))
        children = [n for n in snapshot['nodes'] if n.get('decomposition')]
        self.assertEqual(parent['source_number'], '原条号 6.1、6.2')
        self.assertEqual([n['source_number'] for n in children], ['6.1', '6.2'])
        self.assertEqual([n['requirement_text'] for n in children], [c['text'] for c in next(iter(values[2]['records'].values()))])
        self.assertEqual(values[2:], original)
        workbook = load_workbook(BytesIO(export_delivery_xlsx(snapshot)))
        rows = [row for row in workbook['ComplianceMatrix'].iter_rows(min_row=4) if '-S' in str(row[1].value)]
        self.assertEqual([row[2].value for row in rows], ['6.1', '6.2'])

    def test_explicit_extents_map_numberless_child_bodies_to_their_exact_original_clauses(self):
        snapshot = self.snapshot(fixture(body_only=True, extents=True))
        children = [n for n in snapshot['nodes'] if n.get('decomposition')]
        self.assertEqual([n['source_number'] for n in children], ['6.1', '6.2'])
        self.assertTrue(all(n['decomposition']['source_number_projection']['role'] == 'EXACT_SOURCE_NUMBER' for n in children))

    def test_multiple_numbers_without_a_unique_child_binding_are_explicit_parent_context(self):
        snapshot = self.snapshot(fixture(body_only=True))
        children = [n for n in snapshot['nodes'] if n.get('decomposition')]
        self.assertTrue(all(n['source_number'] == '父项原条号 6.1、6.2；本项见原文依据' for n in children))
        self.assertTrue(all(n['decomposition']['source_number_projection']['role'] == 'PARENT_NUMBER_CONTEXT' for n in children))
        self.assertTrue(all('-S' not in n['source_number'] for n in children))
        workbook = load_workbook(BytesIO(export_delivery_xlsx(snapshot)))
        displayed = [row[2].value for row in workbook['ComplianceMatrix'].iter_rows(min_row=4) if '-S' in str(row[1].value)]
        self.assertEqual(displayed, [n['source_number'] for n in children])

    def test_repeated_original_numbers_keep_distinct_occurrence_evidence(self):
        snapshot = self.snapshot(fixture(numbers=('1', '1')))
        children = [n for n in snapshot['nodes'] if n.get('decomposition')]
        self.assertEqual([n['source_number'] for n in children], ['1', '1'])
        refs = [n['decomposition']['source_number_projection']['numbers'][0]['evidence']['reference']['content_id'] for n in children]
        self.assertEqual(len(set(refs)), 2)

    def test_number_evidence_is_revalidated_before_delivery_even_if_annotation_was_resigned(self):
        values = list(fixture())
        bad = deepcopy(values[-1])
        next(iter(bad['records'].values()))['source_numbers'][0]['evidence']['reference']['start_offset'] = 1
        values[-1] = _seal_annotations(bad)
        with self.assertRaises(ValueError):
            self.snapshot(values)

    def test_explicit_scalar_evidence_takes_precedence_over_a_coarser_structure_label(self):
        document_map, baseline, _, _ = fixture()
        sid = baseline.records[0].source_record_id
        annotation = compile_source_item_annotation_updates(document_map, baseline, [{'source_item_id': sid,
            'source_number': '6.1', 'source_number_ref': {'kind': 'TEXT', 'content_id': 'part-0', 'start_offset': 0, 'end_offset': 3},
            'reason': 'Use the exact first original clause number instead of the broader structure label.'}])
        snapshot = build_delivery_snapshot(document_map, baseline, {}, annotations=annotation)
        parent = next(n for n in snapshot['nodes'] if n.get('source_item'))
        self.assertEqual(parent['source_number'], '6.1')

    def test_table_fragment_id_matches_child_table_id_without_using_snn_as_source_number(self):
        child = {'source_excerpts': [{'source_ref': {'kind': 'TABLE_CELL', 'table_id': 'physical-fragment', 'cell_id': 'body', 'start_offset': 4, 'end_offset': 15}}]}
        entries = [{'value': '7', 'evidence': {'reference': {'kind': 'TABLE_CELL', 'fragment_id': 'physical-fragment', 'cell_id': 'number', 'start_offset': 0, 'end_offset': 1},
            'source_extent': [{'reference': {'kind': 'TABLE_CELL', 'fragment_id': 'physical-fragment', 'cell_id': 'body', 'start_offset': 0, 'end_offset': 20}}]}},
            {'value': '8', 'evidence': {'reference': {'kind': 'TABLE_CELL', 'fragment_id': 'physical-fragment', 'cell_id': 'other-number', 'start_offset': 0, 'end_offset': 1}}}]
        self.assertEqual(_derived_source_number(child, entries, '')[0], '7')
        entries[0]['evidence']['source_extent'][0]['reference']['end_offset'] = 10
        self.assertIn('父项原条号', _derived_source_number(child, entries, '')[0])

    def test_text_extent_uses_absolute_source_offsets_not_child_segment_offsets(self):
        child = {'source_ranges': [{'segment_index': 0, 'start_offset': 10, 'end_offset': 30}],
                 'source_excerpts': [{'source_ref': {'kind': 'TEXT', 'content_id': 'parser-block', 'start_offset': 210, 'end_offset': 230}}]}
        entries = [{'value': value, 'evidence': {'reference': {'kind': 'TEXT', 'content_id': 'parser-block', 'start_offset': start, 'end_offset': start + 3},
            'source_extent': [{'reference': {'kind': 'TEXT', 'content_id': 'parser-block', 'start_offset': start, 'end_offset': start + 40}}]}}
            for value, start in [('6.1', 100), ('6.2', 200)]]
        self.assertEqual(_derived_source_number(child, entries, '')[0], '6.2')

    def test_two_matching_occurrences_are_ambiguous_even_if_their_number_values_repeat(self):
        child = {'source_excerpts': [{'source_ref': {'kind': 'TEXT', 'content_id': 'body', 'start_offset': 0, 'end_offset': 40}}]}
        entries = [{'value': '1', 'evidence': {'reference': {'kind': 'TEXT', 'content_id': 'body', 'start_offset': offset, 'end_offset': offset + 1}}} for offset in (0, 20)]
        number, projection = _derived_source_number(child, entries, '')
        self.assertEqual(number, '父项原条号 1、1；本项见原文依据')
        self.assertEqual(projection['role'], 'PARENT_NUMBER_CONTEXT')

    def test_single_parent_number_is_marked_as_inherited_and_long_collections_remain_complete_in_metadata(self):
        child = {'source_excerpts': [{'source_ref': {'kind': 'TEXT', 'content_id': 'body', 'start_offset': 4, 'end_offset': 20}}]}
        entry = {'value': '6.1', 'evidence': {'reference': {'kind': 'TEXT', 'content_id': 'body', 'start_offset': 0, 'end_offset': 3}}}
        number, projection = _derived_source_number(child, [entry], '')
        self.assertEqual((number, projection['role']), ('6.1', 'SINGLE_PARENT_NUMBER'))
        many = [{'value': f'123.456.{i}', 'evidence': {'reference': {'kind': 'TEXT', 'content_id': f'other-{i}', 'start_offset': 0, 'end_offset': 3}}} for i in range(60)]
        number, projection = _derived_source_number(child, many, '')
        self.assertLessEqual(len(number), 160)
        self.assertIn('共60处', number)
        self.assertEqual(projection['numbers'], many)


if __name__ == '__main__':
    unittest.main()
