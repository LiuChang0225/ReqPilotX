from __future__ import annotations

from .runtime_isolation import activate_test_runtime
activate_test_runtime()

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import re
import unittest

from openpyxl import load_workbook
from openpyxl.utils.escape import unescape
from app.services.exporter import export_compliance_matrix
from .test_response_matrix import response_snapshot


def detail_sections(workbook, main_row):
    navigation = main_row[12]
    match = re.fullmatch(r"'([^']+)'!A(\d+)", navigation.hyperlink.location)
    sheet, start = workbook[match[1]], int(match[2])
    back = f"'ComplianceMatrix'!A{main_row[0].row}"
    sections, current = {}, None
    for row in sheet.iter_rows(min_row=start):
        if not row[2].hyperlink or row[2].hyperlink.location != back:
            break
        text = unescape(row[1].value or '')
        if text in {'原文全文', '说明全文', 'AI分析建议（保留）'}:
            current = text; sections[current] = []
        else:
            sections[current].append(text)
            assert sheet.row_dimensions[row[0].row].height <= 320
    return {name: ''.join(parts) for name, parts in sections.items()}


class ResponseFullContentTests(unittest.TestCase):
    def test_combined_engineer_and_ai_over_excel_cell_limit_are_lossless(self):
        data = response_snapshot()
        human = '工程师明确的工作意见，包含条件。\n' * 850
        ai = '原有AI建议，尚待实际内部证据。\n' * 1000
        child = next(n for n in data['nodes'] if n['node_id'] == 'atom-1')
        child.update(analysis_result_state='CURRENT', analysis_status='DECIDED',
            manual_decision={'decision':'NON_COMPLIANT','reason':human},
            ai_assessment={'decision':'PENDING_EVIDENCE','reason':ai})
        child['decomposition']['engineer_response']={'state':'CURRENT'}
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example','',Path(folder),matrix_snapshot=data,export_mode='DRAFT')
            workbook = load_workbook(path)
            row = next(r for r in workbook['ComplianceMatrix'].iter_rows(min_row=4)
                if str(r[9].value).startswith('工程师工作意见：'))
            self.assertIn(human,row[10].value)
            self.assertIn('完整内容',row[10].value)
            self.assertLess(len(row[10].value),32767)
            sections = detail_sections(workbook,row)
            self.assertEqual(sections['说明全文'],row[10].value)
            self.assertEqual(sections['AI分析建议（保留）'],
                'AI分析建议（保留，不覆盖人工意见）：待内部证据\n'+ai)
            self.assertIsNotNone(row[11].hyperlink)

    def test_long_lines_and_punctuation_stay_in_main_and_lossless_detail_rows(self):
        data = response_snapshot()
        source = '\n'.join(f'{i}）中文原文，保留；条件与例外。 END-{i}' for i in range(38)) + '\r\n尾段。'
        context = next(n for n in data['nodes'] if n['node_id'] == 'context')
        context.update(display_text=source, requirement_text=source)
        before = deepcopy(data)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            workbook = load_workbook(path); main = workbook['ComplianceMatrix']
            row = next(row for row in main.iter_rows(min_row=4) if row[0].value == 'SRC-002')
            self.assertEqual(unescape(row[3].value), source)
            self.assertEqual(row[12].value, '查看完整原文')
            self.assertEqual(detail_sections(workbook, row)['原文全文'], source)
            self.assertEqual(main.cell(3, 12).value, '来源 / 原表')
            self.assertEqual(main.cell(3, 13).value, '完整内容')
        self.assertEqual(data, before)

    def test_source_over_3000_characters_is_never_replaced_by_the_old_240_character_preview(self):
        data = response_snapshot()
        source = '准确原文与句号。\n' * 420
        context = next(n for n in data['nodes'] if n['node_id'] == 'context')
        context.update(display_text=source)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            workbook = load_workbook(path)
            row = next(row for row in workbook['ComplianceMatrix'].iter_rows(min_row=4) if row[0].value == 'SRC-002')
            self.assertEqual(row[3].value, source)
            self.assertIsNotNone(row[11].hyperlink)
            self.assertEqual(detail_sections(workbook, row)['原文全文'], source)

    def test_long_response_has_complete_detail_without_overwriting_the_original_table_link(self):
        data = response_snapshot()
        response = '\n'.join(f'证据{i}：尚无内部试验报告；请核对完整条件，不构成符合判断。' for i in range(42))
        child = next(n for n in data['nodes'] if n['node_id'] == 'atom-1')
        child.update(analysis_result_state='CURRENT', analysis_status='ASSESSED', manual_decision=None,
                     ai_assessment={'decision': 'PENDING_EVIDENCE', 'reason': response})
        data['counts'].update(assessed_nodes=2, decided_nodes=0)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            workbook = load_workbook(path)
            row = next(row for row in workbook['ComplianceMatrix'].iter_rows(min_row=4)
                       if row[4].value == '派生条目' and response in str(row[10].value))
            self.assertEqual(row[9].value, '待内部证据')
            self.assertTrue(row[11].hyperlink.location.startswith("'T001"))
            self.assertEqual(row[12].value, '查看完整说明')
            self.assertEqual(detail_sections(workbook, row)['说明全文'], row[10].value)

    def test_long_visual_reading_remains_whole_in_the_explanation_and_detail(self):
        from .test_response_matrix import with_reading
        data = response_snapshot(); reading = '图中尺寸与原图同源，暂待工程师核对。\n' * 90
        with_reading(data, 'context', reading)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            workbook = load_workbook(path)
            row = next(row for row in workbook['ComplianceMatrix'].iter_rows(min_row=4) if row[0].value == 'SRC-002')
            self.assertIn(reading, row[10].value)
            self.assertEqual(detail_sections(workbook, row)['说明全文'], row[10].value)
            self.assertIn(reading, row[10].comment.text)


if __name__ == '__main__':
    unittest.main()
