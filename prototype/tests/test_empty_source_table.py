
if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from argparse import Namespace
from datetime import datetime, timezone
import json
import unittest

from docx import Document
from .test_office_source_pipeline import write, map_decisions
from app.adapters.documents.docx_source_ir_adapter import DocxSourceIRAdapter
from app.adapters.documents.preflight import LocalDocumentPreflight
from app.application.agent_workflow import source_requirement_baseline as core
import artifact_tools
import scope_tools


class EmptySourceTableTests(unittest.TestCase):
    def prepare(self, root, body=''):
        source=root/'template.docx';doc=Document();doc.add_heading('Material information',1)
        table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Material';table.cell(0,1).text='Supplier'
        table.cell(1,0).text=body;table.cell(1,1).text=body;doc.save(source)
        snapshot=DocxSourceIRAdapter().extract(source,preflight=LocalDocumentPreflight().inspect(source),created_at=datetime(2026,9,8,tzinfo=timezone.utc))
        write(root/'source-ir.json',snapshot.model_dump(mode='json'));write(root/'map-input.json',map_decisions(source,snapshot))
        artifact_tools.compile_map(Namespace(source=source,source_ir=root/'source-ir.json',input=root/'map-input.json',output_dir=root/'map'))
        mapping=core.DocumentMapArtifactV1.model_validate_json((root/'map/document-map.json').read_text())
        request=core.SourceRequirementScopeRequestV1.model_validate(json.loads((root/'map/scope-requests.json').read_text())['table_requests'][0])
        contexts=[{'role':role,'row_group':{'fragment_id':request.fragments[0].fragment_id,'row_start':index,'row_end':index+1},'reason_codes':['table_header' if index==0 else 'empty_template']} for index,role in enumerate(['HEADER','METADATA'])]
        values={'source_sha256':request.source_sha256,'document_map_artifact_id':mapping.artifact_id,'document_map_artifact_sha256':mapping.artifact_sha256,
            'request_sha256':request.request_sha256,'logical_table_scope_id':request.logical_table_scope_id,'records':[],'context_row_groups':contexts}
        return mapping,request,values

    def test_empty_body_keeps_grid_without_fabricated_src(self):
        with TemporaryDirectory() as tmp:
            mapping,request,values=self.prepare(Path(tmp));candidate=core.SourceRequirementSliceCandidateV1.model_validate(values)
            baseline,receipt=core.compile_source_requirement_baseline(document_map=mapping,requests=(request,),candidates=(candidate,))
            self.assertEqual(baseline.records,());self.assertEqual(receipt.context_cell_count,4)
            self.assertEqual(len(baseline.context_dispositions),2)
            self.assertNotIn('exact_placeholder_values',baseline.model_dump(mode='json')['context_dispositions'][0])
            bad=deepcopy(values);bad['context_row_groups'].pop()
            with self.assertRaisesRegex(ValueError,'COVERAGE_INCOMPLETE'):
                core.validate_source_requirement_slice_candidate(request,core.SourceRequirementSliceCandidateV1.model_validate(bad))

    def test_placeholder_requires_explicit_exact_value_and_source_reason(self):
        for body in ['/', '0', 'N/A', 'The product shall retain a fault code.']:
            with self.subTest(body=body),TemporaryDirectory() as tmp:
                mapping,request,values=self.prepare(Path(tmp),body)
                with self.assertRaisesRegex(ValueError,'HAS_BODY_TEXT'):
                    core.validate_source_requirement_slice_candidate(request,core.SourceRequirementSliceCandidateV1.model_validate(values))
                if body=='/':
                    values['context_row_groups'][1].update(exact_placeholder_values=['/'],empty_template_reason='The supplied template repeats the slash in each unfilled field; retain the exact original grid as context.')
                    baseline,receipt=core.compile_source_requirement_baseline(document_map=mapping,requests=(request,),candidates=(core.SourceRequirementSliceCandidateV1.model_validate(values),))
                    self.assertEqual(baseline.records,());self.assertEqual(baseline.context_dispositions[1].exact_placeholder_values,('/',))
                    tables=json.loads((Path(tmp)/'map/table-previews.json').read_text())
                    template=scope_tools.scope_template(mapping,baseline,tables,request.logical_table_scope_id)['candidate_template']
                    self.assertEqual(template['context_row_groups'][1]['exact_placeholder_values'],['/'])

    def test_zero_record_revision_cannot_reclassify_existing_body_as_header(self):
        with TemporaryDirectory() as tmp:
            mapping,request,values=self.prepare(Path(tmp),'Source body requirement')
            initial=deepcopy(values);initial['context_row_groups']=initial['context_row_groups'][:1]
            initial['records']=[{'record_key':'body','label':'Original table body','source_form':'TABLE_RECORD',
                'row_groups':[{'fragment_id':request.fragments[0].fragment_id,'row_start':1,'row_end':2}],
                'requirementness':'REQUIREMENT','requirement_kinds':['FUNCTIONAL'],'confidence':0.9,'decision_summary':'Explicit source obligation.', 'reason_codes':['source_obligation']}]
            base,receipt=core.compile_source_requirement_baseline(document_map=mapping,requests=(request,),candidates=(core.SourceRequirementSliceCandidateV1.model_validate(initial),))
            bad=deepcopy(values);bad['context_row_groups']=[{'role':'HEADER','row_group':{'fragment_id':request.fragments[0].fragment_id,'row_start':0,'row_end':2},'reason_codes':['misleading_header']}]
            candidate=core.SourceRequirementSliceCandidateV1.model_validate(bad)
            revision=core.build_source_item_scope_revision(document_map=mapping,current_artifact=base,current_validation_receipt=receipt,operation='REPLACE',affected_scope_ids=(request.logical_table_scope_id,),comparison_scope_ids=(),feedback_ids=('synthetic-check',),requests=(request,),candidates=(candidate,))
            with self.assertRaisesRegex(ValueError,'HEADER_BOUNDARY_CHANGED'):
                core.compile_source_item_scope_revision(document_map=mapping,current_artifact=base,current_validation_receipt=receipt,scope_revision=revision,requests=(request,),candidates=(candidate,))
