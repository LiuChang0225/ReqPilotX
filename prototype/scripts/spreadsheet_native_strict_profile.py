"""Build a text-free, source-bound strict profile for one native XLSX.

The profile is an evaluator input, not approval.  It binds the immutable
OOXML package, workbook/sheet structure, cached/formula/display projections,
visual pages, and the current SourceIR snapshot without copying customer text
into the report.
"""

from __future__ import annotations

# The repository-local app path must be installed before application imports.
# ruff: noqa: E402

import argparse
import hashlib
import json
import math
import sys
import zipfile
from datetime import date, datetime, time, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from app.adapters.documents.xlsx_source_ir_adapter import XlsxSourceIRAdapter
from app.application.extraction.build_source_ir import build_source_ir
from app.domain.extraction.source_ir import SourceIRSnapshot


SCHEMA_VERSION = "xlsx-native-strict-profile/v1"
REVIEW_GRADE = "AI_SILVER_PENDING"
_FIXED_SOURCE_IR_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
_KNOWN_PACKAGE_PARTS = {
    "xl/workbook.xml": "workbook_xml",
    "xl/styles.xml": "styles_xml",
    "xl/sharedStrings.xml": "shared_strings_xml",
    "xl/calcChain.xml": "calc_chain_xml",
    "xl/theme/theme1.xml": "theme_xml",
    "docProps/core.xml": "core_properties_xml",
    "docProps/app.xml": "app_properties_xml",
}
_PACKAGE_PART_CLASSES = {
    "worksheet": "xl/worksheets/",
    "drawing": "xl/drawings/",
    "image": "xl/media/",
    "chart": "xl/charts/",
    "comment": "xl/comments",
}


class SpreadsheetNativeStrictProfileRejected(RuntimeError):
    pass


def build_xlsx_native_strict_profile(
    source_path: Path,
    *,
    visual_profile: dict[str, Any],
    case_id: str,
    source_ir_snapshot: SourceIRSnapshot | None = None,
) -> dict[str, Any]:
    source_path = source_path.resolve(strict=True)
    if source_path.suffix.casefold() != ".xlsx":
        raise SpreadsheetNativeStrictProfileRejected(
            "native strict profile requires an XLSX source"
        )
    if not case_id.strip():
        raise SpreadsheetNativeStrictProfileRejected("case_id must not be blank")
    source_bytes = source_path.read_bytes()
    source_sha256 = _sha256(source_bytes)
    _validate_visual_profile(
        visual_profile,
        source_sha256=source_sha256,
        source_size=len(source_bytes),
    )
    snapshot = source_ir_snapshot or build_source_ir(
        source_path,
        adapters=(XlsxSourceIRAdapter(),),
        clock=lambda: _FIXED_SOURCE_IR_TIME,
    )
    if snapshot.source_sha256 != source_sha256:
        raise SpreadsheetNativeStrictProfileRejected(
            "SourceIR snapshot is not bound to the XLSX source"
        )

    formula_book = load_workbook(
        BytesIO(source_bytes),
        read_only=False,
        data_only=False,
        keep_links=False,
    )
    cached_book = load_workbook(
        BytesIO(source_bytes),
        read_only=False,
        data_only=True,
        keep_links=False,
    )
    try:
        sheets = tuple(
            _sheet_fact(index, worksheet, cached_book.worksheets[index])
            for index, worksheet in enumerate(formula_book.worksheets)
        )
        original_ooxml = _ooxml_fact(source_bytes, formula_book=formula_book)
    except Exception as exc:
        raise SpreadsheetNativeStrictProfileRejected(
            "XLSX native facts could not be profiled safely"
        ) from exc
    finally:
        formula_book.close()
        cached_book.close()

    profile: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "case_id": case_id,
        "source_format": "xlsx",
        "source_sha256": source_sha256,
        "source_size": len(source_bytes),
        "original_ooxml": original_ooxml,
        "sheets": list(sheets),
        "source_ir": _source_ir_fact(snapshot),
        "visual_profile": _visual_fact(visual_profile),
        "coverage": {
            "sheet_order_state": True,
            "used_ranges": True,
            "nonempty_cells_formulas_cached_display_projection": True,
            "merged_ranges": True,
            "hidden_rows_columns": True,
            "images_drawings_charts_comments": True,
            "styles_number_formats": True,
            "page_print_setup": True,
            "direct_visual_pages": True,
            "source_ir_snapshot": True,
        },
        "review_grade": REVIEW_GRADE,
        "human_reviewed": False,
        "serving_eligible": False,
        "contains_customer_text": False,
        "contains_source_path": False,
        "contains_source_filename": False,
    }
    profile["profile_sha256"] = _canonical_sha256(profile)
    return profile


def _ooxml_fact(
    source_bytes: bytes,
    *,
    formula_book: Any,
) -> dict[str, Any]:
    with zipfile.ZipFile(BytesIO(source_bytes)) as archive:
        members = sorted(archive.infolist(), key=lambda item: item.filename)
        if len({item.filename for item in members}) != len(members):
            raise SpreadsheetNativeStrictProfileRejected(
                "OOXML package contains duplicate members"
            )
        package_facts: list[dict[str, Any]] = []
        known_part_hashes: dict[str, str] = {}
        class_facts: dict[str, list[dict[str, Any]]] = {
            key: [] for key in _PACKAGE_PART_CLASSES
        }
        for member in members:
            payload = archive.read(member)
            fact = {
                "name_sha256": _text_sha256(member.filename),
                "size": len(payload),
                "content_sha256": _sha256(payload),
                "compress_type": member.compress_type,
            }
            package_facts.append(fact)
            safe_name = _KNOWN_PACKAGE_PARTS.get(member.filename)
            if safe_name is not None:
                known_part_hashes[safe_name] = fact["content_sha256"]
            for class_name, prefix in _PACKAGE_PART_CLASSES.items():
                if member.filename.startswith(prefix) and not member.filename.endswith(
                    ".rels"
                ):
                    class_facts[class_name].append(fact)
        part_classes = {
            class_name: {
                "count": len(facts),
                "facts_sha256": _canonical_sha256(facts),
            }
            for class_name, facts in class_facts.items()
        }

    defined_names = [
        {
            "name_sha256": _text_sha256(str(item.name)),
            "type_sha256": _text_sha256(str(item.type)),
            "attr_text_sha256": _text_sha256(str(item.attr_text or "")),
            "local_sheet_id": item.localSheetId,
            "hidden": bool(item.hidden),
        }
        for item in formula_book.defined_names.values()
    ]
    number_formats = [
        _text_sha256(str(value))
        for value in getattr(formula_book, "_number_formats", ())
    ]
    return {
        "package_member_count": len(package_facts),
        "package_manifest_sha256": _canonical_sha256(package_facts),
        "known_part_hashes": known_part_hashes,
        "part_classes": part_classes,
        "external_link_count": len(
            getattr(formula_book, "_external_links", ())
        ),
        "workbook_style_count": len(
            getattr(formula_book, "_cell_styles", ())
        ),
        "workbook_number_format_count": len(number_formats),
        "workbook_number_formats_sha256": _canonical_sha256(number_formats),
        "defined_names_count": len(defined_names),
        "defined_names_sha256": _canonical_sha256(defined_names),
        "active_sheet_index": formula_book.index(formula_book.active),
    }


def _sheet_fact(index: int, worksheet: Any, cached_worksheet: Any) -> dict[str, Any]:
    nonempty_cells: list[dict[str, Any]] = []
    formulas: list[dict[str, Any]] = []
    styled_cells: list[dict[str, Any]] = []
    comments: list[dict[str, Any]] = []
    for (row, column), cell in sorted(worksheet._cells.items()):
        coordinate_sha256 = _text_sha256(cell.coordinate)
        cached_cell = cached_worksheet.cell(row=row, column=column)
        if cell.value is not None:
            is_formula = cell.data_type == "f"
            raw_value = _value_fact(cell.value)
            cached_value = _value_fact(cached_cell.value)
            display_value = cached_value if is_formula else raw_value
            number_format_sha256 = _text_sha256(cell.number_format)
            cell_fact = {
                "coordinate_sha256": coordinate_sha256,
                "data_type": cell.data_type,
                "formula": is_formula,
                "raw_value": raw_value,
                "cached_value": cached_value,
                "number_format_sha256": number_format_sha256,
                "display_projection_sha256": _canonical_sha256(
                    {
                        "cached_or_value": display_value,
                        "number_format_sha256": number_format_sha256,
                    }
                ),
            }
            nonempty_cells.append(cell_fact)
            if is_formula:
                formulas.append(
                    {
                        "coordinate_sha256": coordinate_sha256,
                        "formula_sha256": raw_value["sha256"],
                        "cached_value": cached_value,
                    }
                )
        if cell.has_style:
            styled_cells.append(
                {
                    "coordinate_sha256": coordinate_sha256,
                    "style_id": cell.style_id,
                    "number_format_sha256": _text_sha256(cell.number_format),
                }
            )
        if cell.comment is not None:
            comments.append(
                {
                    "coordinate_sha256": coordinate_sha256,
                    "author_sha256": _text_sha256(cell.comment.author or ""),
                    "text_sha256": _text_sha256(cell.comment.text or ""),
                }
            )

    hidden_rows = [
        {
            "row": row,
            "height": dimension.height,
            "outline_level": dimension.outlineLevel,
            "collapsed": bool(dimension.collapsed),
        }
        for row, dimension in sorted(worksheet.row_dimensions.items())
        if dimension.hidden
    ]
    hidden_columns = [
        {
            "column_sha256": _text_sha256(str(column)),
            "min": dimension.min,
            "max": dimension.max,
            "width": dimension.width,
            "outline_level": dimension.outlineLevel,
            "collapsed": bool(dimension.collapsed),
        }
        for column, dimension in sorted(worksheet.column_dimensions.items())
        if dimension.hidden
    ]
    images = [_image_fact(item) for item in getattr(worksheet, "_images", ())]
    charts = [_chart_fact(item) for item in getattr(worksheet, "_charts", ())]
    merged_ranges = sorted(str(item) for item in worksheet.merged_cells.ranges)
    page_and_print = _page_and_print_fact(worksheet)
    sheet: dict[str, Any] = {
        "sheet_index": index,
        "sheet_name_sha256": _text_sha256(worksheet.title),
        "sheet_state": worksheet.sheet_state,
        "used_range": worksheet.calculate_dimension(),
        "max_row": worksheet.max_row,
        "max_column": worksheet.max_column,
        "nonempty_cell_count": len(nonempty_cells),
        "nonempty_cells_sha256": _canonical_sha256(nonempty_cells),
        "formula_count": len(formulas),
        "formulas_sha256": _canonical_sha256(formulas),
        "display_projection_sha256": _canonical_sha256(
            [item["display_projection_sha256"] for item in nonempty_cells]
        ),
        "styled_cell_count": len(styled_cells),
        "styled_cells_sha256": _canonical_sha256(styled_cells),
        "number_format_count": len(
            {item["number_format_sha256"] for item in styled_cells}
        ),
        "merged_ranges": merged_ranges,
        "merged_ranges_sha256": _canonical_sha256(merged_ranges),
        "hidden_rows": hidden_rows,
        "hidden_columns": hidden_columns,
        "images": images,
        "images_sha256": _canonical_sha256(images),
        "charts": charts,
        "charts_sha256": _canonical_sha256(charts),
        "comments": comments,
        "comments_sha256": _canonical_sha256(comments),
        "page_and_print": page_and_print,
        "page_and_print_sha256": _canonical_sha256(page_and_print),
    }
    sheet["sheet_fact_sha256"] = _canonical_sha256(sheet)
    return sheet


def _image_fact(image: Any) -> dict[str, Any]:
    try:
        payload = image._data()
    except Exception as exc:
        raise SpreadsheetNativeStrictProfileRejected(
            "embedded image bytes are unavailable"
        ) from exc
    return {
        "content_sha256": _sha256(payload),
        "size": len(payload),
        "format": getattr(image, "format", None),
        "width": image.width,
        "height": image.height,
        "anchor": _anchor_fact(image.anchor),
    }


def _chart_fact(chart: Any) -> dict[str, Any]:
    try:
        payload = etree.tostring(chart._write())
    except Exception as exc:
        raise SpreadsheetNativeStrictProfileRejected(
            "chart XML could not be canonicalized"
        ) from exc
    return {
        "type": type(chart).__name__,
        "content_sha256": _sha256(payload),
        "anchor": _anchor_fact(chart.anchor),
    }


def _anchor_fact(anchor: Any) -> dict[str, Any]:
    fact: dict[str, Any] = {"kind": type(anchor).__name__}
    for label, attribute in (("from", "_from"), ("to", "to")):
        marker = getattr(anchor, attribute, None)
        if marker is not None:
            fact[label] = {
                "col": marker.col,
                "row": marker.row,
                "col_offset": marker.colOff,
                "row_offset": marker.rowOff,
            }
    extent = getattr(anchor, "ext", None)
    if extent is not None:
        fact["extent"] = {"cx": extent.cx, "cy": extent.cy}
    return fact


def _page_and_print_fact(worksheet: Any) -> dict[str, Any]:
    return {
        "setup": _primitive_attributes(
            worksheet.page_setup,
            (
                "orientation",
                "paperSize",
                "scale",
                "fitToWidth",
                "fitToHeight",
                "firstPageNumber",
                "useFirstPageNumber",
                "blackAndWhite",
                "draft",
                "cellComments",
                "errors",
                "horizontalDpi",
                "verticalDpi",
                "copies",
                "id",
            ),
        ),
        "margins": _primitive_attributes(
            worksheet.page_margins,
            ("left", "right", "top", "bottom", "header", "footer"),
        ),
        "print_options": _primitive_attributes(
            worksheet.print_options,
            (
                "horizontalCentered",
                "verticalCentered",
                "headings",
                "gridLines",
                "gridLinesSet",
            ),
        ),
        "print_area_sha256": _text_sha256(str(worksheet.print_area or "")),
        "print_title_rows_sha256": _text_sha256(
            str(worksheet.print_title_rows or "")
        ),
        "print_title_cols_sha256": _text_sha256(
            str(worksheet.print_title_cols or "")
        ),
        "freeze_panes_sha256": _text_sha256(str(worksheet.freeze_panes or "")),
        "auto_filter_ref_sha256": _text_sha256(
            str(worksheet.auto_filter.ref or "")
        ),
        "row_breaks": [item.id for item in worksheet.row_breaks.brk],
        "column_breaks": [item.id for item in worksheet.col_breaks.brk],
    }


def _primitive_attributes(item: Any, names: tuple[str, ...]) -> dict[str, Any]:
    fact: dict[str, Any] = {}
    for name in names:
        value = getattr(item, name, None)
        if value is None or isinstance(value, (str, int, float, bool)):
            fact[name] = value
        else:
            fact[name] = _text_sha256(str(value))
    return fact


def _source_ir_fact(snapshot: SourceIRSnapshot) -> dict[str, Any]:
    snapshot_payload = snapshot.model_dump(
        mode="json",
        exclude={"created_at"},
    )
    fact: dict[str, Any] = {
        "source_ir_id": snapshot.source_ir_id,
        "adapter_id": snapshot.adapter_id,
        "adapter_version": snapshot.adapter_version,
        "status": snapshot.status.value,
        "node_count": len(snapshot.nodes),
        "asset_count": len(snapshot.assets),
        "diagnostic_codes": [item.code for item in snapshot.diagnostics],
        "snapshot_sha256": _canonical_sha256(snapshot_payload),
        "node_facts": [
            {
                "source_node_id": node.source_node_id,
                "kind": node.kind.value,
                "fact_sha256": _canonical_sha256(node.model_dump(mode="json")),
            }
            for node in snapshot.nodes
        ],
        "asset_facts": [
            {
                "asset_id": asset.asset_id,
                "kind": asset.kind.value,
                "content_sha256": asset.content_sha256,
                "fact_sha256": _canonical_sha256(asset.model_dump(mode="json")),
            }
            for asset in snapshot.assets
        ],
        "page_render_asset_count": sum(
            asset.kind.value == "page_render" for asset in snapshot.assets
        ),
    }
    fact["projection_sha256"] = _canonical_sha256(fact)
    return fact


def _visual_fact(profile: dict[str, Any]) -> dict[str, Any]:
    pages = profile["pages"]
    return {
        "profile_sha256": profile["profile_sha256"],
        "pdf_sha256": profile["pdf_sha256"],
        "pdf_page_count": profile["pdf_page_count"],
        "sheets": profile["sheets"],
        "pages_sha256": _canonical_sha256(pages),
        "renderer_id": profile["renderer_id"],
        "renderer_version": profile["renderer_version"],
        "libreoffice_version": profile["libreoffice_version"],
        "config_sha256": profile["config_sha256"],
    }


def _validate_visual_profile(
    profile: dict[str, Any],
    *,
    source_sha256: str,
    source_size: int,
) -> None:
    if (
        profile.get("schema_version") != "spreadsheet-visible-profile/v1"
        or profile.get("source_format") != "xlsx"
        or profile.get("source_sha256") != source_sha256
        or profile.get("source_size") != source_size
        or profile.get("profile_sha256")
        != _canonical_sha256(
            {
                key: value
                for key, value in profile.items()
                if key != "profile_sha256"
            }
        )
    ):
        raise SpreadsheetNativeStrictProfileRejected(
            "visual profile is not bound to the XLSX source"
        )
    pages = profile.get("pages")
    if not isinstance(pages, list) or [
        item.get("page_number") for item in pages
    ] != list(range(1, len(pages) + 1)):
        raise SpreadsheetNativeStrictProfileRejected(
            "visual profile page coverage is not exact"
        )


def _value_fact(value: Any) -> dict[str, str]:
    if value is None:
        kind, normalized = "none", ""
    elif isinstance(value, bool):
        kind, normalized = "boolean", "true" if value else "false"
    elif isinstance(value, int):
        kind, normalized = "integer", str(value)
    elif isinstance(value, float):
        kind = "float"
        if math.isnan(value):
            normalized = "nan"
        elif value == math.inf:
            normalized = "inf"
        elif value == -math.inf:
            normalized = "-inf"
        else:
            normalized = repr(value)
    elif isinstance(value, Decimal):
        kind, normalized = "decimal", str(value)
    elif isinstance(value, datetime):
        kind, normalized = "datetime", value.isoformat()
    elif isinstance(value, date):
        kind, normalized = "date", value.isoformat()
    elif isinstance(value, time):
        kind, normalized = "time", value.isoformat()
    elif isinstance(value, bytes):
        kind, normalized = "bytes", _sha256(value)
    else:
        kind, normalized = "string", str(value)
    return {
        "kind": kind,
        "sha256": _text_sha256(f"{kind}:{normalized}"),
    }


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _text_sha256(value: str) -> str:
    return _sha256(value.encode("utf-8"))


def _canonical_sha256(value: Any) -> str:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return _sha256(payload)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--visual-profile", type=Path, required=True)
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args()
    visual_profile = json.loads(
        arguments.visual_profile.read_text(encoding="utf-8")
    )
    profile = build_xlsx_native_strict_profile(
        arguments.source,
        visual_profile=visual_profile,
        case_id=arguments.case_id,
    )
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(
        json.dumps(profile, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )
    print(profile["profile_sha256"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
