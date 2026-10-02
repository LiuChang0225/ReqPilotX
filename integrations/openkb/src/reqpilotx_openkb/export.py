"""Add an explicitly version-bound original-evidence annex to native Excel.

The untouched canonical workbook remains alongside the extended deliverable.
No wiki output is copied into the original quote column.
"""
from __future__ import annotations

import json
from pathlib import Path


def add_original_evidence_annex(directory, analysis_receipt):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    directory = Path(directory)
    native = directory / "matrix-native.xlsx"
    (directory / "matrix.xlsx").rename(native)
    book = load_workbook(native)
    sheet = book.create_sheet("OpenKB原件依据")
    headers = ["分析单元", "Agent判断", "证据ID", "原生fact_id", "原件名称", "原件SHA256", "原生定位JSON", "原文摘录", "摘录分段", "适用条件", "冲突", "缺失依据", "依据权威", "依据类型", "知识版本", "分析版本", "人工接受"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
    for unit_id, decision in sorted(analysis_receipt["decisions"].items()):
        originals = decision["resolved_originals"] or [None]
        for original in originals:
            original = original or {}
            quote = original.get("excerpt", "")
            parts = [quote[i:i + 8000] for i in range(0, len(quote), 8000)] or [""]
            for number, part in enumerate(parts, 1):
                values = [unit_id, decision["decision"], original.get("evidence_id", ""), original.get("fact_id", ""),
                    original.get("source_doc_name", ""), original.get("source_sha256", ""),
                    json.dumps(original.get("source_ref", {}), ensure_ascii=False, sort_keys=True), part,
                    str(number) + "/" + str(len(parts)), "\n".join(decision["conditions"]), "\n".join(decision["conflicts"]),
                    "\n".join(decision["missing"]), original.get("authority", ""), original.get("claim_role", ""),
                    analysis_receipt["generation"], analysis_receipt["assessment_sha256"], "NOT_CLAIMED"]
                if any(len(str(value)) > 32767 for value in values):
                    raise ValueError("Evidence annex exceeds Excel cell limit; no silently truncated export")
                sheet.append(values)
                for cell in sheet[sheet.max_row]:
                    cell.data_type = "s"  # An original starting '=' stays literal text.
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for name in ("A", "B", "C", "D", "E"):
        sheet.column_dimensions[name].width = 28
    for name in ("F", "G", "H", "J", "K", "L", "O", "P"):
        sheet.column_dimensions[name].width = 48
    book.save(directory / "matrix.xlsx")
    book.close()
