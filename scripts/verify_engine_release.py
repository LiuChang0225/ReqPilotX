#!/usr/bin/env python3
"""Verify the R1 engine release outside the source checkout.

Dependency preparation is explicit. Runtime checks then clear source-path
injection, use the non-editable wheel, rebuild from the sdist, exercise the
wheel-derived plugin and compare one exact synthetic business input.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile


ROOT = Path(__file__).resolve().parents[1]
PLUGIN_SCRIPTS = ROOT / "plugins/reqpilot/scripts"
sys.path.insert(0, str(PLUGIN_SCRIPTS))

from build_bundle import (  # noqa: E402
    _assert_same_selection,
    _install_direct_selection,
    _validated_wheel_manifest,
    _wheel_content_inventory,
    build as build_bundle,
)
from build_engine import build as build_engine  # noqa: E402
from portable_core import select_core  # noqa: E402
from run_required_unittest import evaluate_required_tests  # noqa: E402


REQUIRED_MCP_TESTS = (
    "McpTaskAdapterTests.test_region_pixels_reuse_native_geometry_without_rebuilding_all_page_views",
    "McpTaskAdapterTests.test_visual_revision_is_callable_through_real_stdio_and_fixed_cli",
    "McpTaskAdapterTests.test_typed_cli_and_stdio_mcp_produce_equal_business_artifacts",
    "McpTaskAdapterTests.test_actual_stdio_discovery_create_capture_replay_and_restart",
    "McpTaskAdapterTests.test_actual_image_content_matches_verified_page_bytes",
)
CODE_FINGERPRINT_PATHS = (
    "packages/engine/pyproject.toml",
    "plugins/reqpilot/scripts/portable_core.py",
    "plugins/reqpilot/scripts/build_engine.py",
    "plugins/reqpilot/scripts/build_bundle.py",
    "plugins/reqpilot/scripts/runtime_origin_audit.py",
    "plugins/reqpilot/tests/portable_fixtures.py",
    "plugins/reqpilot/tests/test_engine_distribution.py",
    "plugins/reqpilot/tests/test_release_verification.py",
    "scripts/run_required_unittest.py",
    "scripts/verify_engine_release.py",
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(command: list[str], *, cwd: Path, env: dict[str, str] | None = None, capture: bool = False) -> subprocess.CompletedProcess:
    return subprocess.run(command, cwd=cwd, env=env, check=True, text=True, capture_output=capture)


def isolated_env() -> dict[str, str]:
    env = {
        key: value for key, value in os.environ.items()
        if key not in {"PYTHONPATH", "PYTHONHOME"}
        and not key.startswith(("REQPILOT_", "REQPILOTX_", "OPENAI_", "AZURE_", "DASHSCOPE_", "ALIYUN_", "ALIBABA_CLOUD_"))
    }
    env.update({
        "PYTHONNOUSERSITE": "1",
        "REQPILOT_TEST_MODE": "1",
        "REQPILOT_ENV": "test",
        "REQPILOT_AGENT_RUNTIME": "none",
        "HF_HUB_OFFLINE": "1",
        "TRANSFORMERS_OFFLINE": "1",
        "HF_DATASETS_OFFLINE": "1",
        "DO_NOT_TRACK": "1",
    })
    return env


EXPORT_TIMESTAMP_SHEET = "_ReqPilotXMeta"
EXPORT_TIMESTAMP_LABEL_CELL = "A30"
EXPORT_TIMESTAMP_VALUE_CELL = "B30"
EXPORT_TIMESTAMP_LABEL = "ExportedAt"
ISO_TIMESTAMP = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})?$")
PACKAGE_RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_RELATIONSHIPS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SPREADSHEET = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
DRAWING_MAIN = "http://schemas.openxmlformats.org/drawingml/2006/main"
CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _style_xml(value) -> str:
    from xml.etree.ElementTree import tostring
    return tostring(value.to_tree(), encoding="unicode")


def _relationship_part(source_part: str) -> str:
    directory, name = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", f"{name}.rels")


def _resolve_relationship_target(source_part: str, target: str) -> str:
    if not target or "\\" in target:
        raise ValueError(f"Invalid OOXML relationship target: {target!r}")
    resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(posixpath.dirname(source_part), target))
    if resolved == ".." or resolved.startswith("../") or resolved.startswith("/"):
        raise ValueError(f"OOXML relationship escapes the workbook: {target!r}")
    return resolved


def _relationships(archive: ZipFile, source_part: str) -> dict[str, dict]:
    relation_part = _relationship_part(source_part)
    if relation_part not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(relation_part))
    result = {}
    for node in root.findall(f"{{{PACKAGE_RELATIONSHIPS}}}Relationship"):
        identity = node.attrib.get("Id")
        if not identity or identity in result:
            raise ValueError(f"Invalid or duplicate OOXML relationship ID in {relation_part}")
        target_mode = node.attrib.get("TargetMode", "Internal")
        target = node.attrib.get("Target", "")
        result[identity] = {
            "type": node.attrib.get("Type", "").rsplit("/", 1)[-1],
            "target_mode": target_mode,
            "target": target if target_mode == "External" else _resolve_relationship_target(source_part, target),
        }
    return result


def _direct_child(node: ET.Element, namespace: str, local_name: str) -> ET.Element | None:
    return node.find(f"{{{namespace}}}{local_name}")


def _marker(anchor: ET.Element, name: str) -> dict | None:
    node = _direct_child(anchor, DRAWING, name)
    if node is None:
        return None
    fields = {}
    for field in ("col", "colOff", "row", "rowOff"):
        value = _direct_child(node, DRAWING, field)
        if value is None or value.text is None:
            raise ValueError(f"Drawing anchor {name} is missing {field}")
        fields[field] = int(value.text)
    return fields


def drawing_bindings(path: Path) -> list[dict]:
    bindings = []
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        workbook_part = "xl/workbook.xml"
        workbook = ET.fromstring(archive.read(workbook_part))
        workbook_rels = _relationships(archive, workbook_part)
        for sheet in workbook.findall(f".//{{{SPREADSHEET}}}sheet"):
            sheet_name = sheet.attrib.get("name", "")
            sheet_rel_id = sheet.attrib.get(f"{{{OFFICE_RELATIONSHIPS}}}id")
            if not sheet_rel_id or sheet_rel_id not in workbook_rels:
                raise ValueError(f"Worksheet {sheet_name!r} has no resolvable workbook relationship")
            sheet_part = workbook_rels[sheet_rel_id]["target"]
            if sheet_part not in names:
                raise ValueError(f"Worksheet part is missing: {sheet_part}")
            worksheet = ET.fromstring(archive.read(sheet_part))
            sheet_rels = _relationships(archive, sheet_part)
            for drawing_ref in worksheet.findall(f"{{{SPREADSHEET}}}drawing"):
                drawing_rel_id = drawing_ref.attrib.get(f"{{{OFFICE_RELATIONSHIPS}}}id")
                if not drawing_rel_id or drawing_rel_id not in sheet_rels:
                    raise ValueError(f"Worksheet {sheet_name!r} has an unresolved drawing relationship")
                drawing_part = sheet_rels[drawing_rel_id]["target"]
                if drawing_part not in names:
                    raise ValueError(f"Drawing part is missing: {drawing_part}")
                drawing = ET.fromstring(archive.read(drawing_part))
                drawing_rels = _relationships(archive, drawing_part)
                for anchor in list(drawing):
                    anchor_type = anchor.tag.rsplit("}", 1)[-1]
                    if anchor_type not in {"oneCellAnchor", "twoCellAnchor", "absoluteAnchor"}:
                        continue
                    anchor_value = {
                        "type": anchor_type,
                        "attributes": dict(sorted(anchor.attrib.items())),
                        "from": _marker(anchor, "from"),
                        "to": _marker(anchor, "to"),
                    }
                    position = _direct_child(anchor, DRAWING, "pos")
                    extent = _direct_child(anchor, DRAWING, "ext")
                    anchor_value["position"] = dict(sorted(position.attrib.items())) if position is not None else None
                    anchor_value["extent"] = dict(sorted(extent.attrib.items())) if extent is not None else None
                    objects = []
                    for object_type in ("pic", "graphicFrame"):
                        obj = _direct_child(anchor, DRAWING, object_type)
                        if obj is None:
                            continue
                        nonvisual = obj.find(f".//{{{DRAWING}}}cNvPr")
                        references = []
                        if object_type == "pic":
                            blip = obj.find(f".//{{{DRAWING_MAIN}}}blip")
                            relation_ids = [] if blip is None else [
                                (kind, blip.attrib.get(f"{{{OFFICE_RELATIONSHIPS}}}{kind}")) for kind in ("embed", "link")
                                if blip.attrib.get(f"{{{OFFICE_RELATIONSHIPS}}}{kind}")
                            ]
                        else:
                            chart = obj.find(f".//{{{CHART}}}chart")
                            relation_ids = [] if chart is None else [("chart", chart.attrib.get(f"{{{OFFICE_RELATIONSHIPS}}}id"))]
                        for reference_kind, relation_id in relation_ids:
                            if not relation_id or relation_id not in drawing_rels:
                                raise ValueError(f"Drawing object has an unresolved resource relationship: {relation_id}")
                            relation = drawing_rels[relation_id]
                            target = relation["target"]
                            content_sha256 = None
                            if relation["target_mode"] != "External":
                                if target not in names:
                                    raise ValueError(f"Drawing resource part is missing: {target}")
                                content_sha256 = hashlib.sha256(archive.read(target)).hexdigest()
                            references.append({
                                "reference_kind": reference_kind,
                                "relationship_type": relation["type"],
                                "target_mode": relation["target_mode"],
                                "content_sha256": content_sha256,
                                "external_target": target if relation["target_mode"] == "External" else None,
                            })
                        objects.append({
                            "object_type": "IMAGE" if object_type == "pic" else "CHART",
                            "object_name": nonvisual.attrib.get("name") if nonvisual is not None else None,
                            "references": sorted(references, key=lambda item: json.dumps(item, sort_keys=True)),
                        })
                    for obj in objects:
                        bindings.append({"worksheet": sheet_name, "anchor": anchor_value, **obj})
    return sorted(bindings, key=lambda item: json.dumps(item, sort_keys=True, separators=(",", ":")))


def workbook_semantics(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    sheets = []
    timestamp_cells = []
    for sheet in workbook.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style and cell.hyperlink is None:
                    continue
                value = cell.value
                if sheet.title == EXPORT_TIMESTAMP_SHEET and cell.coordinate == EXPORT_TIMESTAMP_VALUE_CELL:
                    if sheet[EXPORT_TIMESTAMP_LABEL_CELL].value != EXPORT_TIMESTAMP_LABEL:
                        raise ValueError("Export timestamp coordinate no longer has the required ExportedAt label")
                    if not isinstance(value, str) or not ISO_TIMESTAMP.fullmatch(value):
                        raise ValueError("Export timestamp value is not the contracted ISO timestamp")
                    timestamp_cells.append(f"{sheet.title}!{cell.coordinate}")
                    value = "<EXPORT_TIMESTAMP>"
                cells.append({
                    "coordinate": cell.coordinate,
                    "value": value,
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                    "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                    "font": _style_xml(cell.font),
                    "fill": _style_xml(cell.fill),
                    "border": _style_xml(cell.border),
                    "alignment": _style_xml(cell.alignment),
                    "protection": _style_xml(cell.protection),
                })
        row_dimensions = [{
            "index": index, "height": value.height, "hidden": value.hidden,
            "outline_level": value.outlineLevel, "collapsed": value.collapsed,
            "thick_top": value.thickTop, "thick_bottom": value.thickBot,
        } for index, value in sorted(sheet.row_dimensions.items())]
        column_dimensions = [{
            "index": index, "width": value.width, "hidden": value.hidden,
            "outline_level": value.outlineLevel, "collapsed": value.collapsed,
            "best_fit": value.bestFit, "min": value.min, "max": value.max,
        } for index, value in sorted(sheet.column_dimensions.items())]
        sheets.append({
            "title": sheet.title,
            "state": sheet.sheet_state,
            "dimension": sheet.calculate_dimension(),
            "freeze_panes": str(sheet.freeze_panes or ""),
            "auto_filter": sheet.auto_filter.ref,
            "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
            "row_dimensions": row_dimensions,
            "column_dimensions": column_dimensions,
            "cells": cells,
        })
    with ZipFile(path) as archive:
        charts = {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist()) if name.startswith("xl/charts/")
        }
        media = sorted(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith("xl/media/")
        )
    return {
        "sheets": sheets,
        "charts": charts,
        "media_sha256": media,
        "drawing_bindings": drawing_bindings(path),
        "normalized_timestamp_cells": timestamp_cells,
    }


def first_difference(left, right, path="workbook"):
    if type(left) is not type(right):
        return f"{path}: type {type(left).__name__} != {type(right).__name__}"
    if isinstance(left, dict):
        if left.keys() != right.keys():
            return f"{path}: keys {sorted(left)} != {sorted(right)}"
        for key in left:
            difference = first_difference(left[key], right[key], f"{path}.{key}")
            if difference:
                return difference
        return None
    if isinstance(left, list):
        if len(left) != len(right):
            return f"{path}: length {len(left)} != {len(right)}"
        for index, (left_item, right_item) in enumerate(zip(left, right)):
            difference = first_difference(left_item, right_item, f"{path}[{index}]")
            if difference:
                return difference
        return None
    return None if left == right else f"{path}: {left!r} != {right!r}"


def compare_business_outputs(development: Path, release: Path) -> dict:
    json_paths = (
        "map/document-map.json",
        "baseline/source-requirement-baseline.json",
        "delivery/matrix-snapshot.json",
    )
    json_equal = {}
    for relative in json_paths:
        left = json.loads((development / relative).read_text(encoding="utf-8"))
        right = json.loads((release / relative).read_text(encoding="utf-8"))
        json_equal[relative] = left == right
    left_workbook = workbook_semantics(development / "delivery/matrix.xlsx")
    right_workbook = workbook_semantics(release / "delivery/matrix.xlsx")
    workbook_equal = left_workbook == right_workbook
    if not all(json_equal.values()) or not workbook_equal:
        raise ValueError(
            "Installed release changed deterministic business output: "
            f"json={json_equal}, workbook_difference={first_difference(left_workbook, right_workbook)}"
        )
    return {
        "json_equal": json_equal,
        "xlsx_cells_formulas_formats_merges_equal": True,
        "xlsx_charts_equal": True,
        "xlsx_media_equal": True,
        "xlsx_drawing_bindings_equal": True,
        "xlsx_sheet_count": len(left_workbook["sheets"]),
        "xlsx_media_count": len(left_workbook["media_sha256"]),
        "xlsx_drawing_binding_count": len(left_workbook["drawing_bindings"]),
        "xlsx_normalized_timestamp_cells": left_workbook["normalized_timestamp_cells"],
    }


def venv_python(venv: Path) -> Path:
    return venv / ("Scripts/python.exe" if os.name == "nt" else "bin/python")


def prepare_venv(uv: str, root: Path, name: str) -> Path:
    venv = root / name
    run([uv, "venv", "--python", sys.executable, str(venv)], cwd=root)
    return venv_python(venv)


def product_origins() -> str:
    return """
origins={}
for name,module in sorted(sys.modules.items()):
    if name == 'app' or name.startswith('app.') or name == 'reqpilotx_engine' or name.startswith('reqpilotx_engine.'):
        source=getattr(module,'__file__',None)
        origins[name]=[str(Path(source).resolve())] if source else [str(Path(item).resolve()) for item in getattr(module,'__path__',[])]
outside={name:paths for name,paths in origins.items() if not paths or any(not Path(source).is_relative_to(allowed) for source in paths)}
"""


def run_installed_example(
    python: Path,
    fixture_dir: Path,
    output: Path,
    *,
    source: Path | None,
    allowed_root: Path,
    cwd: Path,
    env: dict[str, str],
    prepend_product_root: bool = False,
) -> dict:
    source_expression = "None" if source is None else f"Path({json.dumps(str(source))})"
    code = f"""
import contextlib,json,sys
from pathlib import Path
allowed=Path({json.dumps(str(allowed_root))}).resolve()
if {str(prepend_product_root)}: sys.path.insert(0,str(allowed))
sys.path.insert(0,{json.dumps(str(fixture_dir))})
from worked_example import generate
with contextlib.redirect_stdout(sys.stderr):
    result=generate(None,Path({json.dumps(str(output))}),{source_expression})
{product_origins()}
if outside: raise RuntimeError('Installed business modules escaped site-packages: '+json.dumps(outside,sort_keys=True))
print(json.dumps({{'result':result,'module_origins':origins}}))
"""
    process = subprocess.run([str(python), "-I", "-c", code], cwd=cwd, env=env, text=True, capture_output=True)
    if process.returncode:
        raise RuntimeError((process.stderr or process.stdout)[-8000:])
    return json.loads(process.stdout)


def run_plugin_cli(
    python: Path,
    entrypoint: Path,
    arguments: list[str],
    *,
    allowed_root: Path,
    origin_dir: Path,
    cwd: Path,
    env: dict[str, str],
    expect_success: bool = True,
) -> subprocess.CompletedProcess:
    command_env = dict(env)
    command_env.update({
        "REQPILOT_TEST_ALLOWED_PRODUCT_ROOT": str(allowed_root.resolve()),
        "REQPILOT_TEST_ORIGIN_DIR": str(origin_dir.resolve()),
    })
    result = subprocess.run([str(python), str(entrypoint), *arguments], cwd=cwd, env=command_env, text=True, capture_output=True)
    if expect_success and result.returncode:
        raise ValueError(f"Plugin CLI failed: {(result.stderr or result.stdout)[-4000:]}")
    if not expect_success and result.returncode == 0:
        raise ValueError("Plugin CLI unexpectedly succeeded after a required module was removed")
    return result


def read_origin_receipts(directory: Path, *, require_components: set[str]) -> dict:
    receipts = [json.loads(path.read_text(encoding="utf-8")) for path in sorted(directory.glob("*.json"))]
    if not receipts:
        raise ValueError("No module-origin receipts were produced")
    components = {item.get("component") for item in receipts}
    if not require_components <= components:
        raise ValueError(f"Module-origin receipts are missing components: {sorted(require_components - components)}")
    outside = {item["component"]: item.get("outside_allowed_root") for item in receipts if item.get("outside_allowed_root")}
    if outside:
        raise ValueError(f"Product modules escaped the isolated implementation root: {outside}")
    observed = {
        name: source
        for item in receipts
        for name, source in (item.get("module_origins") or {}).items()
    }
    if not observed:
        raise ValueError("Origin receipts did not observe any ReqPilotX business modules")
    summary = summarize_module_origins(observed, Path(receipts[0]["allowed_product_root"]))
    summary.update({"receipt_count": len(receipts), "components": sorted(components)})
    return {"summary": summary, "details": receipts}


def summarize_module_origins(origins: dict[str, list[str]], allowed_root: Path) -> dict:
    representative_names = (
        "app.application.fast_source_capture",
        "app.application.agent_workflow.source_requirement_baseline",
        "app.application.agent_workflow.source_item_delivery",
        "app.services.exporter",
    )
    return {
        "allowed_root": str(allowed_root.resolve()),
        "module_count": len(origins),
        "origins_sha256": hashlib.sha256(json.dumps(origins, sort_keys=True, separators=(",", ":")).encode()).hexdigest(),
        "representative": {name: origins[name] for name in representative_names if name in origins},
    }


def release_status(mcp_result: dict) -> str:
    return "PASSED" if mcp_result.get("status") == "PASSED" else "PARTIAL"


def verify(repo: Path, *, keep_output: Path | None = None, skip_mcp: bool = False) -> dict:
    repo = repo.expanduser().resolve(strict=True)
    uv = shutil.which("uv")
    if uv is None:
        raise ValueError("Release verification requires the prepared uv frontend")
    temporary = tempfile.TemporaryDirectory(prefix="reqpilotx-engine-verify-")
    root = Path(temporary.name)
    try:
        selection = select_core(repo)
        engine_output = root / "engine-release"
        engine_receipt = build_engine(repo, engine_output)
        wheel = next(engine_output.glob("*.whl"))
        sdist = next(engine_output.glob("*.tar.gz"))
        original_manifest = _validated_wheel_manifest(wheel)
        _assert_same_selection(selection, original_manifest)

        rebuilt = root / "rebuilt-from-sdist"
        run([
            uv, "build", "--wheel", "--no-sources", "--no-python-downloads", "--no-build-logs",
            "--python", sys.executable, "--out-dir", str(rebuilt), str(sdist),
        ], cwd=root)
        rebuilt_wheel = next(rebuilt.glob("*.whl"))
        rebuilt_manifest = _validated_wheel_manifest(rebuilt_wheel)
        _assert_same_selection(selection, rebuilt_manifest)
        if original_manifest != rebuilt_manifest:
            raise ValueError("Wheel rebuilt from the sdist changed the engine identity manifest")
        if _wheel_content_inventory(wheel) != _wheel_content_inventory(rebuilt_wheel):
            raise ValueError("Wheel rebuilt from the sdist changed its member set or member content")

        engine_python = prepare_venv(uv, root, "engine-venv")
        run([
            uv, "pip", "install", "--python", str(engine_python), str(rebuilt_wheel),
        ], cwd=root)

        plugin_python = prepare_venv(uv, root, "plugin-venv")
        plugin_requirements = root / "plugin-dependencies.txt"
        mcp_requirements = [
            line.strip() for line in (repo / "plugins/reqpilot/requirements-mcp.txt").read_text(encoding="utf-8").splitlines()
            if line.strip() and not line.lstrip().startswith("#")
        ]
        plugin_requirements.write_text("\n".join([*original_manifest["requirements"], *mcp_requirements]) + "\n", encoding="utf-8")
        run([uv, "pip", "install", "--python", str(plugin_python), "-r", str(plugin_requirements)], cwd=root)

        env = isolated_env()
        probe_code = f"""
import importlib.util, json, sys
from pathlib import Path
import reqpilotx_engine
manifest=reqpilotx_engine.core_manifest()
site_root=Path(reqpilotx_engine.__file__).resolve().parent.parent
repo=Path({json.dumps(str(repo))}).resolve()
assert 'site-packages' in str(site_root), site_root
assert not site_root.is_relative_to(repo), site_root
assert importlib.util.find_spec('app.adapters.persistence') is None
assert importlib.util.find_spec('app.infrastructure') is None
assert not {{'fastapi','psycopg','openai','qdrant_client'}} & set(sys.modules)
print(json.dumps({{'core_identity':manifest['core_identity'],'site_root':str(site_root)}}))
"""
        origin = json.loads(run([str(engine_python), "-I", "-c", probe_code], cwd=root, env=env, capture=True).stdout)
        if origin["core_identity"] != engine_receipt["core_identity"]:
            raise ValueError("Installed engine reports a different core identity")
        engine_site_root = Path(origin["site_root"])
        no_plugin_engine = """
import importlib.util
assert importlib.util.find_spec('app') is None
assert importlib.util.find_spec('reqpilotx_engine') is None
"""
        run([str(plugin_python), "-I", "-c", no_plugin_engine], cwd=root, env=env)

        plugin = root / "plugin"
        build_bundle(repo, plugin, engine_wheel=wheel)
        plugin_receipt = json.loads((plugin / "build-receipt.json").read_text(encoding="utf-8"))
        if plugin_receipt["engine_release"]["core_identity"] != engine_receipt["core_identity"]:
            raise ValueError("Plugin does not consume the built engine identity")
        doctor = run([str(plugin_python), str(plugin / "scripts/reqpilot.py"), "doctor"], cwd=root, env=env, capture=True)
        doctor_result = json.loads(doctor.stdout)
        if doctor_result.get("status") != "READY" or doctor_result.get("integrity_mismatches"):
            raise ValueError(f"Wheel-derived plugin doctor failed: {doctor_result}")

        direct_plugin = root / "direct-selection-plugin"
        shutil.copytree(
            repo / "plugins/reqpilot",
            direct_plugin,
            ignore=shutil.ignore_patterns("runtime", "runtime-data", "tests", "__pycache__", "*.pyc", "*.zip", ".pytest_cache"),
        )
        _install_direct_selection(select_core(repo), direct_plugin / "runtime/prototype")
        fixture_dir = root / "installed-fixture"
        shutil.copytree(
            repo / "plugins/reqpilot/scripts", fixture_dir,
            ignore=shutil.ignore_patterns("__pycache__", "*.pyc"),
        )
        (fixture_dir / "reqpilot.py").write_text("def load_runtime(root=None):\n    return None\n", encoding="utf-8")

        source_origins = root / "source-origins"; source_origins.mkdir()
        plugin_origins = root / "plugin-origins"; plugin_origins.mkdir()
        development = root / "development-example"
        run_plugin_cli(
            plugin_python, direct_plugin / "scripts/reqpilot.py", ["worked-example", "--output-dir", str(development)],
            allowed_root=direct_plugin / "runtime/prototype", origin_dir=source_origins, cwd=root, env=env,
        )
        installed = root / "installed-example"
        installed_result = run_installed_example(
            engine_python, fixture_dir, installed, source=development / "example.pdf",
            allowed_root=engine_site_root, cwd=root, env=env,
        )
        release = root / "release-example"
        run_plugin_cli(
            plugin_python, plugin / "scripts/reqpilot.py",
            ["worked-example", "--source", str(development / "example.pdf"), "--output-dir", str(release)],
            allowed_root=plugin / "runtime/prototype", origin_dir=plugin_origins, cwd=root, env=env,
        )
        engine_comparison = compare_business_outputs(development, installed)
        plugin_comparison = compare_business_outputs(development, release)
        source_origin_result = read_origin_receipts(source_origins, require_components={"reqpilot-cli"})
        plugin_origin_result = read_origin_receipts(plugin_origins, require_components={"reqpilot-cli"})

        broken_plugin = root / "broken-plugin"
        shutil.copytree(plugin, broken_plugin)
        (broken_plugin / "runtime/prototype/app/application/agent_workflow/source_requirement_baseline.py").unlink()
        broken_origins = root / "broken-origins"; broken_origins.mkdir()
        run_plugin_cli(
            plugin_python, broken_plugin / "scripts/reqpilot.py",
            ["worked-example", "--source", str(development / "example.pdf"), "--output-dir", str(root / "broken-plugin-example")],
            allowed_root=broken_plugin / "runtime/prototype", origin_dir=broken_origins, cwd=root, env=env,
            expect_success=False,
        )
        broken_target = root / "broken-engine-target"
        run([uv, "pip", "install", "--python", str(plugin_python), "--target", str(broken_target), "--no-deps", str(rebuilt_wheel)], cwd=root)
        (broken_target / "app/application/agent_workflow/source_requirement_baseline.py").unlink()
        try:
            run_installed_example(
                plugin_python, fixture_dir, root / "broken-engine-example", source=development / "example.pdf",
                allowed_root=broken_target, cwd=root, env=env, prepend_product_root=True,
            )
        except RuntimeError:
            missing_module_negative = "PASSED_NO_FALLBACK"
        else:
            raise ValueError("Installed engine unexpectedly succeeded after a required module was removed")

        mcp_result = {
            "status": "NOT_VERIFIED",
            "reason": "SKIPPED_BY_REQUEST",
            "required": {name: {"status": "NOT_RUN"} for name in REQUIRED_MCP_TESTS},
        }
        mcp_origin_result = None
        if not skip_mcp:
            mcp_origins = root / "mcp-origins"; mcp_origins.mkdir()
            mcp_env = dict(env)
            mcp_env.update({
                "REQPILOT_TEST_PROJECT_ROOT": str(root / "no-checkout"),
                "REQPILOT_TEST_PLUGIN_SCRIPTS": str(plugin / "scripts"),
                "REQPILOT_TEST_RUNTIME": str(plugin / "runtime/prototype"),
                "REQPILOT_TEST_ALLOWED_PRODUCT_ROOT": str(plugin / "runtime/prototype"),
                "REQPILOT_TEST_ORIGIN_DIR": str(mcp_origins),
            })
            (root / "no-checkout").mkdir()
            structured = root / "mcp-results.json"
            process = subprocess.run([
                str(plugin_python), str(repo / "scripts/run_required_unittest.py"),
                "--test-file", str(repo / "plugins/reqpilot/tests/test_mcp_task_adapter.py"),
                "--result", str(structured),
            ], cwd=root, env=mcp_env, text=True)
            if not structured.is_file():
                raise ValueError("Structured MCP test result was not written")
            mcp_summary = json.loads(structured.read_text(encoding="utf-8"))
            mcp_result = evaluate_required_tests(mcp_summary, REQUIRED_MCP_TESTS)
            if process.returncode or mcp_result["status"] != "PASSED" or mcp_summary.get("outside_allowed_root"):
                raise ValueError(f"Required real stdio MCP tests were not all executed from the plugin runtime: {mcp_result}")
            mcp_result["client_module_origins"] = summarize_module_origins(
                mcp_summary.get("module_origins", {}), plugin / "runtime/prototype"
            )
            mcp_origin_result = read_origin_receipts(mcp_origins, require_components={"reqpilot-cli", "mcp-server"})

        overall_status = release_status(mcp_result)
        result = {
            "schema_version": "reqpilotx-engine-release-verification/v1",
            "status": overall_status,
            "source_commit": engine_receipt["source_commit"],
            "source_tree": engine_receipt["source_tree"],
            "core_identity": engine_receipt["core_identity"],
            "runtime_file_count": engine_receipt["runtime_file_count"],
            "environment": {
                "python": sys.version.split()[0],
                "platform": platform.system(),
                "machine": platform.machine(),
                "uv": run([uv, "--version"], cwd=root, capture=True).stdout.strip(),
            },
            "code_fingerprints": {
                relative: sha256(repo / relative) for relative in CODE_FINGERPRINT_PATHS
            },
            "wheel_sha256": sha256(wheel),
            "sdist_sha256": sha256(sdist),
            "sdist_rebuild": "PASSED_WITHOUT_CHECKOUT_AND_CONTENT_EQUAL",
            "installed_module_origins": summarize_module_origins(installed_result["module_origins"], engine_site_root),
            "source_control_module_origins": source_origin_result["summary"],
            "plugin_module_origins": plugin_origin_result["summary"],
            "mcp_module_origins": mcp_origin_result["summary"] if mcp_origin_result else None,
            "missing_module_negative": missing_module_negative,
            "plugin_bundle_sha256": plugin_receipt["bundle_sha256"],
            "plugin_engine_identity_matches": True,
            "doctor": "READY",
            "installed_engine_business_comparison": engine_comparison,
            "plugin_business_comparison": plugin_comparison,
            "mcp": mcp_result,
            "network_during_business_execution": "NO NETWORK STEP REQUESTED; PROVIDER DOWNLOADS DISABLED; DEPENDENCY PREPARATION COMPLETED BEFORE BUSINESS CHECKS",
            "customer_or_production_data_used": False,
        }
        if keep_output is not None:
            keep_output = keep_output.expanduser().resolve()
            if keep_output.exists():
                raise ValueError("Verification output must not exist")
            keep_output.mkdir(parents=True)
            (keep_output / "verification-receipt.json").write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            origin_details = {
                "installed_engine": installed_result["module_origins"],
                "source_control": source_origin_result["details"],
                "plugin": plugin_origin_result["details"],
                "mcp": mcp_origin_result["details"] if mcp_origin_result else None,
            }
            (keep_output / "module-origin-details.json").write_text(json.dumps(origin_details, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        return result
    finally:
        temporary.cleanup()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, help="Optional new directory retaining the non-sensitive verification receipt")
    parser.add_argument("--skip-mcp", action="store_true", help="Explicitly leave real stdio MCP unverified")
    args = parser.parse_args()
    try:
        print(json.dumps(verify(args.repo, keep_output=args.output, skip_mcp=args.skip_mcp), indent=2, ensure_ascii=False))
    except (ValueError, RuntimeError, OSError, subprocess.CalledProcessError) as exc:
        print(json.dumps({"status": "FAILED", "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
