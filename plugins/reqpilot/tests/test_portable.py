"""Behavioral adapter checks with an invented PDF, never a customer fixture."""
from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from copy import deepcopy
from zipfile import ZipFile
from io import BytesIO
from argparse import Namespace
from pathlib import Path

PLUGIN = Path(__file__).resolve().parents[1]
REPO = PLUGIN.parents[1]
sys.path[:0] = [str(PLUGIN / "scripts"), str(REPO / "prototype")]

import pymupdf
import artifact_tools
from app.application.fast_source_capture import capture_fast_source


def write(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def make_fixture(directory: Path, *, table_image=False, numbered_targets=False) -> tuple[Path, Path, dict]:
    """Raw invented obligations cover all five classes and a complete test table."""
    directory.mkdir(parents=True, exist_ok=True)
    source = directory / "synthetic.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    lines = [
        (50, "Valve module specification"),
        (90, "1 Technical targets"),
        (120, "The valve shall open within 30 ms at 12 V and 23 C.\nThe valve shall close within 40 ms at 12 V and 23 C."),
        (155, "2 Project delivery"),
        (185, "The supplier shall deliver drawings before sample approval."),
        (220, "3 Review process"),
        (250, "The reviewer shall record the review decision."),
        (285, "4 Construction"),
        (315, "The housing shall use aluminium."),
        (350, "5 Functional behavior"),
        (380, "The module shall report an overtemperature warning."),
        (415, "6 Technical test"),
        # Deliberately emitted before the table in native paint order.
        (650, "Note: Packaging shown in Appendix A is illustrative."),
    ]
    if numbered_targets:
        lines[2] = (120, "1.1 The valve shall open within 30 ms at 12 V and 23 C.\n1.2 The valve shall close within 40 ms at 12 V and 23 C.")
    for y, text in lines:
        page.insert_text((40, y), text, fontsize=11)
    for x in (40, 190, 420, 555):
        page.draw_line((x, 440), (x, 590))
    for y in (440, 475, 530, 590):
        page.draw_line((40, y), (555, y))
    for x, y, text in [(45, 460, "Test"), (195, 460, "Method and conditions"), (425, 460, "Criterion"), (45, 495, "Response"), (195, 495, "Apply 12 V at 23 C."), (425, 495, "Open <= 30 ms"), (45, 550, "Durability"), (195, 550, "Repeat 1000 cycles."), (425, 550, "No leakage")]:
        page.insert_text((x, y), text, fontsize=10)
    from PIL import Image, ImageDraw
    if table_image:
        figure = Image.new("RGB", (100, 21), (130, 90, 70))
        ImageDraw.Draw(figure).text((5, 5), "Original test diagram", fill="white")
        image_bytes = BytesIO(); figure.save(image_bytes, format="PNG")
        page.insert_image(pymupdf.Rect(310, 503, 410, 524), stream=image_bytes.getvalue())
    for x, caption in ((40, "WIDTH <= 18 mm"), (260, "HEIGHT <= 6 mm")):
        figure = Image.new("RGB", (160, 80), (100, 100, 100))
        ImageDraw.Draw(figure).text((12, 32), caption, fill=(255, 255, 255))
        image_bytes = BytesIO()
        figure.save(image_bytes, format="PNG")
        page.insert_image(pymupdf.Rect(x, 700, x + 160, 780), stream=image_bytes.getvalue())
    pdf.save(source)
    pdf.close()
    capture = directory / "capture"
    capture_fast_source(source, output_dir=capture)
    ir = json.loads((capture / "source_ir.json").read_text())
    by_text = {node["raw_text"]: node for node in ir["nodes"]}
    nodes = [{"node_id": "main", "parent_id": None, "kind": "DOCUMENT", "title": lines[0][1], "label": "", "title_source_ids": [by_text[lines[0][1]]["source_node_id"]], "content_source_ids": []}]
    for number in range(1, 7):
        heading = lines[number * 2 - 1][1]
        body = lines[number * 2][1] if number < 6 else lines[-1][1]
        nodes.append({"node_id": f"s{number}", "parent_id": "main", "kind": "SECTION", "title": heading.split(" ", 1)[1], "label": str(number), "title_source_ids": [by_text[heading]["source_node_id"]], "content_source_ids": [by_text[body]["source_node_id"]]})
    table = next(node for node in ir["nodes"] if node.get("table_payload"))
    nodes[-1]["content_source_ids"].extend(node["source_node_id"] for node in ir["nodes"] if node["kind"] in {"image", "figure"})
    duplicate_ids = [node["source_node_id"] for node in ir["nodes"] if node["kind"] not in {"table", "image", "figure"} and 440 < node["native_locator"]["bbox"][1] < 590]
    decisions = {"schema_version": "reqpilot-map-decisions/v1", "source_sha256": ir["source_sha256"], "nodes": nodes, "logical_tables": [{"scope_id": "technical-test", "owner_node_id": "s6", "title": "Technical test", "role": "TEST_MATRIX", "fragments": [{"source_node_id": table["source_node_id"], "row_range": [0, 3]}]}], "source_dispositions": [{"source_node_id": value, "disposition": "DUPLICATE_CAPTURE", "reason": "This native text block is visibly the same cells in the captured table region."} for value in duplicate_ids]}
    return source, capture, decisions


def item_decisions(map_dir: Path) -> dict:
    artifact = json.loads((map_dir / "document-map.json").read_text())
    requests = json.loads((map_dir / "scope-requests.json").read_text())
    kinds = {"s1": "QUALITY", "s2": "PROJECT", "s3": "PROCESS", "s4": "CONSTRAINT", "s5": "FUNCTIONAL"}
    decisions = {"schema_version": "reqpilot-item-decisions/v1", "source_sha256": artifact["source_sha256"], "document_map_artifact_sha256": artifact["artifact_sha256"], "text_decisions": [], "table_decisions": []}
    for request in requests["text_requests"]:
        items, structure = [], []
        for content in request["source_content"]:
            span = {"content_id": content["content_id"], "start_offset": 0, "end_offset": len(content["text"])}
            if content["is_title"]:
                structure.append({"text_span": span, "reason_codes": ["source_heading"]})
            else:
                kind = kinds.get(request["node_id"])
                items.append({"item_key": "body", "label": content["text"], "source_form": "NUMBERED_CLAUSE", "text_spans": [span], "requirementness": "REQUIREMENT" if kind else "NON_REQUIREMENT", "requirement_kinds": [kind] if kind else [], "non_requirement_role": None if kind else "NOTE", "confidence": 0.9, "decision_summary": "Synthetic original provision preserves complete wording." if kind else "The note states illustrative context.", "reason_codes": ["source_meaning"]})
        decisions["text_decisions"].append({"scope_id": request["scope_id"], "items": items, "structure_only_spans": structure})
    for request in requests["table_requests"]:
        fragment = request["fragments"][0]
        records = []
        for row, label in ((1, "Response"), (2, "Durability")):
            records.append({"record_key": label.lower(), "label": label, "source_form": "TABLE_RECORD", "row_groups": [{"fragment_id": fragment["fragment_id"], "row_start": row, "row_end": row + 1}], "requirementness": "REQUIREMENT", "requirement_kinds": ["QUALITY"], "confidence": 0.9, "decision_summary": "Method, conditions and criterion belong to the same complete test.", "reason_codes": ["complete_test"]})
        decisions["table_decisions"].append({"logical_table_scope_id": request["logical_table_scope_id"], "records": records, "context_row_groups": [{"role": "HEADER", "row_group": {"fragment_id": fragment["fragment_id"], "row_start": 0, "row_end": 1}, "reason_codes": ["shared_header"]}]})
    return decisions


class PortableCompilationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="reqpilot-unit-")
        self.directory = Path(self.temp.name)
        self.source, self.capture, self.map_input = make_fixture(self.directory)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def compile_map(self) -> Path:
        input_path = self.directory / "map-input.json"
        write(input_path, self.map_input)
        output = self.directory / "map"
        artifact_tools.compile_map(Namespace(source=self.source, source_ir=self.capture / "source_ir.json", input=input_path, output_dir=output))
        return output

    def test_source_bound_round_trip_keeps_classes_table_and_order(self) -> None:
        map_dir = self.compile_map()
        input_path = self.directory / "items.json"
        write(input_path, item_decisions(map_dir))
        output = self.directory / "baseline"
        result = artifact_tools.compile_items(Namespace(map_dir=map_dir, input=input_path, baseline_dir=None, output_dir=output))
        baseline = json.loads((output / "source-requirement-baseline.json").read_text())
        self.assertEqual((result["source_item_count"], result["requirement_count"], result["non_requirement_count"]), (8, 7, 1))
        records = baseline["records"]
        self.assertEqual([item["source_display_id"] for item in records], [f"SRC-{number:04d}" for number in range(1, 9)])
        self.assertIsNone(records[-1]["requirement_display_id"])
        self.assertEqual(records[1]["requirement_kinds"], ["PROJECT"])
        self.assertEqual(records[2]["requirement_kinds"], ["PROCESS"])
        self.assertEqual([item["label"] for item in records[5:7]], ["Response", "Durability"])
        self.assertEqual(len(records[5]["shared_header_cell_ids"]), 3)
        self.assertTrue(records[-1]["text_bindings"][0]["text"].startswith("Note:"))
        self.assertEqual(result["unprocessed_scope_count"], 0)
        self.assertEqual(result["semantic_verification"], "NOT_RUN")

    def test_reason_code_order_is_transport_work_and_invalid_codes_still_fail(self) -> None:
        map_dir = self.compile_map(); decisions = item_decisions(map_dir)
        item = next(x['items'][0] for x in decisions['text_decisions'] if x['items'])
        item['reason_codes'] = ['source_meaning', 'complete_unit', 'source_meaning']
        decisions['table_decisions'][0]['records'][0]['reason_codes'] = ['test_method', 'complete_test']
        path = self.directory / 'unordered.json'; write(path, decisions); before = path.read_bytes()
        first = self.directory / 'unordered-baseline'
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=None, output_dir=first))
        ordered = artifact_tools._canonical_item_reason_codes(decisions)
        ordered_path = self.directory / 'ordered.json'; write(ordered_path, ordered)
        second = self.directory / 'ordered-baseline'
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=ordered_path, baseline_dir=None, output_dir=second))
        self.assertEqual((first/'source-requirement-baseline.json').read_bytes(), (second/'source-requirement-baseline.json').read_bytes())
        self.assertEqual(path.read_bytes(), before)
        item['reason_codes'] = ['NOT A MACHINE CODE']; write(path, decisions)
        with self.assertRaisesRegex(ValueError, 'lowercase machine codes'):
            artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=None, output_dir=self.directory/'invalid-codes'))

    def test_wrong_source_revision_is_rejected(self) -> None:
        self.map_input["source_sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "same source SHA"):
            self.compile_map()

    def test_fabricated_heading_is_rejected(self) -> None:
        self.map_input["nodes"][1]["title"] = "Invented technology container"
        with self.assertRaisesRegex(ValueError, "verbatim"):
            self.compile_map()

    def test_unknown_handle_is_rejected(self) -> None:
        self.map_input["nodes"][1]["content_source_ids"] = ["guessed-source"]
        with self.assertRaisesRegex(ValueError, "Unknown source handle"):
            self.compile_map()

    def test_uncovered_original_is_reported_not_discarded(self) -> None:
        self.map_input["nodes"][1]["content_source_ids"] = []
        map_dir = self.compile_map()
        receipt = json.loads((map_dir / "map-compilation-receipt.json").read_text())
        self.assertEqual(receipt["evidence_state"], "INSUFFICIENT")
        self.assertEqual(len(receipt["uncovered_source_node_ids"]), 1)

    def test_text_coverage_gap_cannot_be_called_compiled(self) -> None:
        map_dir = self.compile_map()
        decisions = item_decisions(map_dir)
        body_scope = next(item for item in decisions["text_decisions"] if item["scope_id"] == "text:s1")
        body_scope["items"][0]["text_spans"][0]["end_offset"] -= 3
        input_path = self.directory / "items.json"
        write(input_path, decisions)
        with self.assertRaisesRegex(ValueError, "COVERAGE_INCOMPLETE"):
            artifact_tools.compile_items(Namespace(map_dir=map_dir, input=input_path, baseline_dir=None, output_dir=self.directory / "bad-baseline"))

    def test_shared_delivery_exports_source_identity_and_pending_evidence(self) -> None:
        from openpyxl import load_workbook
        map_dir = self.compile_map()
        decisions_path = self.directory / "items.json"
        write(decisions_path, item_decisions(map_dir))
        baseline_dir = self.directory / "baseline"
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=decisions_path, baseline_dir=None, output_dir=baseline_dir))
        export_dir = self.directory / "delivery"
        result = artifact_tools.deliver(Namespace(baseline_dir=baseline_dir, input=None, source=self.source, output_dir=export_dir))
        snapshot = json.loads((export_dir / "matrix-snapshot.json").read_text())
        self.assertEqual(result["status"], "DRAFT_EXPORTED")
        self.assertEqual(snapshot["source_item_baseline"]["source_item_count"], 8)
        self.assertFalse(any(node.get("ai_assessment") for node in snapshot["nodes"]))
        workbook = load_workbook(export_dir / "matrix.xlsx", data_only=False)
        texts = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
        joined = "\n".join(texts)
        for value in ("SRC-0001", "REQ-0001", "SRC-0008"):
            self.assertIn(value, joined)
        self.assertTrue(any("PROJECT" in node.get("classification_paths", []) for node in snapshot["nodes"]))
        self.assertTrue(any("PROCESS" in node.get("classification_paths", []) for node in snapshot["nodes"]))
        self.assertIn("No leakage", joined)
        with ZipFile(export_dir / "matrix.xlsx") as archive:
            from PIL import Image
            source_crops = []
            for name in archive.namelist():
                if not name.startswith("xl/media/"):
                    continue
                with Image.open(BytesIO(archive.read(name))) as embedded:
                    rgb = embedded.convert("RGB")
                    colors = rgb.getcolors(rgb.width * rgb.height) or []
                    source_gray = sum(count for count, color in colors if all(abs(channel - 100) <= 2 for channel in color))
                    if 1.8 <= rgb.width / rgb.height <= 2.2 and source_gray / (rgb.width * rgb.height) >= 0.95:
                        source_crops.append(name)
            self.assertTrue(source_crops, "Workbook must contain the actual source's 2:1 gray figure pixels, not just a brand logo")

    def test_incremental_scope_growth_preserves_existing_records_and_rejects_overlap(self) -> None:
        map_dir = self.compile_map()
        path = self.directory / "items.json"
        decisions = item_decisions(map_dir)
        initial = deepcopy(decisions)
        initial["text_decisions"] = [item for item in decisions["text_decisions"] if item["scope_id"] == "text:s2"]
        initial["table_decisions"] = []
        write(path, initial)
        first = self.directory / "baseline-1"
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=None, output_dir=first))
        second = self.directory / "baseline-2"
        additional = deepcopy(decisions)
        additional["text_decisions"] = [item for item in decisions["text_decisions"] if item["scope_id"] != "text:s2"]
        write(path, additional)
        result = artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=first, output_dir=second))
        before = json.loads((first / "source-requirement-baseline.json").read_text())
        after = json.loads((second / "source-requirement-baseline.json").read_text())
        self.assertEqual(after["revision"], 2)
        self.assertEqual(after["predecessor_artifact_id"], before["artifact_id"])
        self.assertEqual(result["compilation_mode"], "SOURCE_ORDER_GROWTH")
        self.assertEqual(result["unprocessed_scope_count"], 0)
        old = before["records"][0]; kept = next(r for r in after["records"] if r["source_record_id"] == old["source_record_id"])
        self.assertEqual({k:v for k,v in kept.items() if k != "source_order"}, {k:v for k,v in old.items() if k != "source_order"})
        self.assertEqual(kept["source_display_id"], old["source_display_id"])
        write(path, decisions)
        with self.assertRaisesRegex(ValueError, "explicit review-revise scope_revision"):
            artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=second, output_dir=self.directory / "bad-successor"))

    def decomposition_input(self):
        from app.application.agent_workflow.source_requirement_baseline import SourceRequirementBaselineArtifactV1
        from app.application.agent_workflow.source_item_decomposition import build_source_item_decomposition_packet
        map_dir = self.compile_map()
        path = self.directory / "items.json"
        write(path, item_decisions(map_dir))
        directory = self.directory / "baseline"
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=None, output_dir=directory))
        baseline = SourceRequirementBaselineArtifactV1.model_validate_json((directory / "source-requirement-baseline.json").read_text())
        tables = json.loads((directory / "table-previews.json").read_text())
        packet = build_source_item_decomposition_packet(baseline, baseline.records[0].source_record_id, table_previews=tables)
        text = packet["segments"][0]["text"]
        split = text.index("\n")
        children = [{"text": text[lo:hi], "derivation_type": "VERBATIM_SPAN", "source_ranges": [{"segment_index": 0, "start_offset": lo, "end_offset": hi}], "requirement_kinds": ["QUALITY"]} for lo, hi in ((0, split), (split + 1, len(text)))]
        decision = {"source_item_id": baseline.records[0].source_record_id, "action": "SPLIT", "reason": "Opening and closing response have separate acceptance outcomes.", "children": children, "context_ranges": []}
        return baseline, tables, decision

    def test_explicit_decomposition_keeps_parent_and_supports_undo(self) -> None:
        from app.application.agent_workflow.source_item_decomposition import compile_source_item_decomposition, validate_source_item_decomposition
        baseline, tables, decision = self.decomposition_input()
        before = baseline.model_dump_json()
        artifact = compile_source_item_decomposition(baseline, [decision], table_previews=tables)
        children = artifact["records"][decision["source_item_id"]]
        self.assertEqual([child["requirement_display_id"] for child in children], ["REQ-0001-S01", "REQ-0001-S02"])
        self.assertEqual(baseline.model_dump_json(), before)
        self.assertFalse(artifact["human_accepted"])
        validate_source_item_decomposition(baseline, artifact, table_previews=tables)
        undo = compile_source_item_decomposition(baseline, [{"source_item_id": decision["source_item_id"], "action": "UNDO", "reason": "Restore analysis of the complete original parent."}], table_previews=tables, previous=artifact)
        self.assertEqual(undo["records"], {})
        self.assertEqual(len(undo["retired_singular_ids"]), 2)
        self.assertEqual(undo["predecessor_sha256"], artifact["artifact_sha256"])

    def test_decomposition_rejects_source_omission(self) -> None:
        from app.application.agent_workflow.source_item_decomposition import compile_source_item_decomposition
        baseline, tables, decision = self.decomposition_input()
        decision["children"][1]["source_ranges"][0]["end_offset"] -= 3
        with self.assertRaises(ValueError):
            compile_source_item_decomposition(baseline, [decision], table_previews=tables)

    def test_decomposition_rejects_invented_number_and_pseudo_s01(self) -> None:
        from app.application.agent_workflow.source_item_decomposition import compile_source_item_decomposition
        baseline, tables, decision = self.decomposition_input()
        invented = deepcopy(decision)
        invented["children"][1]["text"] = invented["children"][1]["text"].replace("40 ms", "41 ms")
        invented["children"][1]["derivation_type"] = "NORMALIZED_DERIVATION"
        with self.assertRaisesRegex(ValueError, "unsupported numeric"):
            compile_source_item_decomposition(baseline, [invented], table_previews=tables)
        decision["children"] = decision["children"][:1]
        with self.assertRaisesRegex(ValueError, "pseudo S01"):
            compile_source_item_decomposition(baseline, [decision], table_previews=tables)

    def test_decomposition_rejects_modified_sealed_artifact(self) -> None:
        from app.application.agent_workflow.source_item_decomposition import compile_source_item_decomposition, validate_source_item_decomposition
        baseline, tables, decision = self.decomposition_input()
        artifact = compile_source_item_decomposition(baseline, [decision], table_previews=tables)
        artifact["records"][decision["source_item_id"]][0]["text"] += " plus another obligation"
        with self.assertRaisesRegex(ValueError, "modified after compilation"):
            validate_source_item_decomposition(baseline, artifact, table_previews=tables)

    def test_new_children_after_undo_do_not_reuse_retired_display_ids(self) -> None:
        from app.application.agent_workflow.source_item_decomposition import compile_source_item_decomposition
        baseline, tables, decision = self.decomposition_input()
        first = compile_source_item_decomposition(baseline, [decision], table_previews=tables)
        undo = compile_source_item_decomposition(baseline, [{"source_item_id": decision["source_item_id"], "action": "UNDO", "reason": "Reassess the response boundaries."}], table_previews=tables, previous=first)
        changed = deepcopy(decision)
        changed["children"][0]["text"] = "Opening response must be within 30 ms at 12 V and 23 C."
        changed["children"][0]["derivation_type"] = "NORMALIZED_DERIVATION"
        third = compile_source_item_decomposition(baseline, [changed], table_previews=tables, previous=undo)
        former = {child["requirement_display_id"]: child["singular_requirement_id"] for child in first["records"][decision["source_item_id"]]}
        for child in third["records"][decision["source_item_id"]]:
            if child["requirement_display_id"] in former:
                self.assertEqual(child["singular_requirement_id"], former[child["requirement_display_id"]], "A retired Snn must not identify a different child after undo")

    def test_compiled_children_and_cited_pending_analysis_reach_excel(self) -> None:
        from openpyxl import load_workbook
        baseline, tables, decision = self.decomposition_input()
        path = self.directory / "decomposition-decisions.json"
        envelope = {"schema_version": "reqpilot-decomposition-decisions/v1", "baseline_sha256": baseline.artifact_sha256, "authorization": {"request": "Separate the opening and closing responses for review.", "source_item_ids": [decision["source_item_id"]]}, "decisions": [decision]}
        write(path, envelope)
        output = self.directory / "decomposition"
        artifact_tools.decomposition_operation(Namespace(command="compile-decomposition", baseline_dir=self.directory / "baseline", input=path, previous_decomposition=None, output_dir=output))
        sealed = output / "source-item-decomposition.json"
        artifact = json.loads(sealed.read_text())
        children = artifact["records"][decision["source_item_id"]]
        quote = children[0]["source_excerpts"][0]["text"]
        review = {"assessments": {"baseline_sha256": baseline.artifact_sha256, "records": {children[0]["singular_requirement_id"]: {"decision": "UNKNOWN", "reason": "The customer defines the target; internal product evidence has not been supplied.", "evidence": [{"fact_id": "customer-clause", "quote": quote}]}}, "evidence_catalog": {"customer-clause": {"role": "CUSTOMER_SOURCE", "text": quote}}}}
        review_path = self.directory / "review.json"
        write(review_path, review)
        delivery = self.directory / "with-children"
        with self.assertRaisesRegex(ValueError, "sealed artifact"):
            artifact_tools.deliver(Namespace(baseline_dir=self.directory / "baseline", input=review_path, source=self.source, decomposition=sealed, output_dir=delivery))
        packet_dir = self.directory / "child-analysis"
        artifact_tools.analysis_operation(Namespace(command="analysis-packet", baseline_dir=self.directory / "baseline", knowledge_catalog=None, decomposition=sealed, unit_id=None, output_dir=packet_dir))
        packet = json.loads((packet_dir / "analysis-packet.json").read_text())
        response = {"decisions": [{"unit_id": unit["unit_id"], "decision": "PENDING_EVIDENCE", "reason": "No actual product evidence was supplied for these source conditions.", "evidence": [], "clarification_question": "Provide the corresponding product design or test evidence."} for unit in packet["units"]]}
        response_path = self.directory / "child-assessment-input.json"; write(response_path, response)
        assessment_dir = self.directory / "child-assessments"
        artifact_tools.analysis_operation(Namespace(command="compile-assessments", analysis_packet=packet_dir / "analysis-packet.json", input=response_path, previous_assessments=None, output_dir=assessment_dir))
        artifact_tools.deliver(Namespace(baseline_dir=self.directory / "baseline", input=None, source=self.source, decomposition=sealed, assessments=assessment_dir / "source-item-assessments.json", output_dir=delivery))
        snapshot = json.loads((delivery / "matrix-snapshot.json").read_text())
        child_nodes = [node for node in snapshot["nodes"] if node["source_kind"] == "SINGULAR_REQUIREMENT"]
        self.assertEqual(len(child_nodes), 2)
        self.assertEqual(child_nodes[0]["ai_assessment"]["decision"], "PENDING_EVIDENCE")
        self.assertEqual(snapshot["source_item_baseline"]["source_item_count"], 8)
        workbook = load_workbook(delivery / "matrix.xlsx")
        values = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
        for value in ("SRC-0001", "REQ-0001-S01", "REQ-0001-S02"):
            self.assertIn(value, values)
        self.assertLess(values.index("REQ-0001-S01"), values.index("REQ-0001-S02"))
        total = snapshot['counts']['analysis_eligible_nodes']
        self.assertTrue(any(f"已分析 {total}/{total}" in value and f"有效判断 0/{total}" in value and f"待内部证据 {total}" in value for value in values),
                        "Completed pending analysis must remain distinct from a supported technical conclusion")
        review["assessments"]["records"][children[0]["singular_requirement_id"]]["decision"] = "COMPLIANT"
        write(review_path, review)
        with self.assertRaisesRegex(ValueError, "sealed artifact"):
            artifact_tools.deliver(Namespace(baseline_dir=self.directory / "baseline", input=review_path, source=self.source, decomposition=sealed, output_dir=self.directory / "false-compliance"))

    def image_input(self):
        from app.application.agent_workflow.source_requirement_baseline import DocumentMapArtifactV1, SourceRequirementBaselineArtifactV1
        from app.application.agent_workflow.source_item_visual import build_source_item_image_scope_packet
        map_dir = self.compile_map()
        path = self.directory / "items.json"
        write(path, item_decisions(map_dir))
        directory = self.directory / "baseline"
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=path, baseline_dir=None, output_dir=directory))
        baseline = SourceRequirementBaselineArtifactV1.model_validate_json((directory / "source-requirement-baseline.json").read_text())
        document_map = DocumentMapArtifactV1.model_validate_json((directory / "document-map.json").read_text())
        packet = build_source_item_image_scope_packet(document_map, node_id="s6")
        self.assertEqual(len(packet["images"]), 2)
        candidate = {"node_id": "s6", "packet_sha256": packet["packet_sha256"], "items": []}
        for index, image in enumerate(packet["images"]):
            candidate["items"].append({"item_key": f"drawing-{index+1}", "display_label": f"Dimensional drawing {index+1}", "image_anchor_ids": [image["anchor_id"]], "requirementness": "REQUIREMENT", "requirement_kinds": ["QUALITY"], "non_requirement_role": None, "confidence": 0.9, "decision_summary": "The original drawing specifies a dimensional limit.", "reason_codes": ["normative_drawing"], "visual_interpretation": "Agent reading of the visible dimension; this is not customer-verbatim text."})
        return document_map, baseline, candidate, packet

    def test_table_owned_picture_is_context_not_a_duplicate_image_src(self):
        from app.application.agent_workflow import source_requirement_baseline as core
        from app.application.agent_workflow.source_item_visual import build_source_item_image_scope_packet, compile_source_item_images
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot
        self.source, self.capture, self.map_input = make_fixture(self.directory / "with-table-image", table_image=True)
        map_dir = self.compile_map()
        decisions = self.directory / "items.json"; write(decisions, item_decisions(map_dir))
        baseline_dir = self.directory / "baseline"
        artifact_tools.compile_items(Namespace(map_dir=map_dir, input=decisions, baseline_dir=None, output_dir=baseline_dir))
        document_map = core.DocumentMapArtifactV1.model_validate_json((map_dir / "document-map.json").read_text())
        previous = core.SourceRequirementBaselineArtifactV1.model_validate_json((baseline_dir / "source-requirement-baseline.json").read_text())
        tables = json.loads((map_dir / "table-previews.json").read_text())
        with self.assertRaisesRegex(ValueError, "compile their D3"):
            build_source_item_image_scope_packet(document_map, node_id="s6")
        packet = build_source_item_image_scope_packet(document_map, node_id="s6", baseline=previous, table_previews=tables)
        self.assertEqual((len(packet["images"]), len(packet["table_image_relations"])), (2, 1))
        candidate = {"node_id":"s6", "packet_sha256":packet["packet_sha256"], "items":[{
            "item_key":"standalone", "display_label":"Separate original drawings", "image_anchor_ids":[x["anchor_id"] for x in packet["images"]],
            "requirementness":"REQUIREMENT", "requirement_kinds":["QUALITY"], "non_requirement_role":None,
            "confidence":0.9, "decision_summary":"Complete standalone drawings.", "reason_codes":["original_drawing"], "visual_interpretation":"Agent reading only."}]}
        baseline, _, _ = compile_source_item_images(document_map, [candidate], previous_baseline=previous, table_previews=tables)
        self.assertEqual(len(baseline.records), len(previous.records)+1)
        self.assertEqual(len(baseline.table_image_relations), 1)
        relation = baseline.table_image_relations[0]
        owner = next(r for r in baseline.records if r.source_record_id == relation.table_source_record_id)
        self.assertTrue(owner.table_fragments); self.assertFalse(owner.image_bindings)
        self.assertFalse(any(relation.image_binding.anchor_id == b.anchor_id for r in baseline.records for b in r.image_bindings))
        snapshot = build_delivery_snapshot(document_map, baseline, tables)
        self.assertTrue(next(n for n in snapshot["nodes"] if n["node_id"] == owner.source_record_id)["figure_refs"])
        core._validate_table_image_relations(document_map, baseline.records, baseline.table_image_relations)
        invalid = relation.model_copy(update={"cell_ids": ("unowned-cell",)})
        with self.assertRaisesRegex(ValueError, "CELL_OWNERSHIP"):
            core._validate_table_image_relations(document_map, baseline.records, (invalid,))
        changed = tuple(r for r in baseline.records if r.source_record_id != owner.source_record_id)
        with self.assertRaisesRegex(ValueError, "REQUIRES_RECOMPILATION"):
            core._inherit_table_image_relations(baseline, changed)

    def test_image_src_keeps_pixels_separate_and_is_analyzable(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot, export_delivery_xlsx
        from openpyxl import load_workbook
        document_map, previous, candidate, _ = self.image_input()
        baseline, validation, annotations = compile_source_item_images(document_map, [candidate], previous_baseline=previous)
        images = [record for record in baseline.records if record.image_bindings]
        self.assertEqual((len(baseline.records), validation.requirement_count, validation.owned_image_count), (10, 9, 2))
        self.assertEqual([record.source_display_id for record in images], ["SRC-0009", "SRC-0010"])
        self.assertTrue(all(not record.text_bindings and not record.table_fragments for record in images))
        self.assertEqual({record.source_record_id for record in previous.records}, {record.source_record_id for record in baseline.records if not record.image_bindings})
        tables = json.loads((self.directory / "baseline/table-previews.json").read_text())
        assessment = {"baseline_sha256": baseline.artifact_sha256, "records": {record.source_record_id: {"decision": "UNKNOWN", "reason": "Image specification is known; internal design evidence is missing.", "evidence": []} for record in images}}
        snapshot = build_delivery_snapshot(document_map, baseline, tables, annotations=annotations, assessments=assessment)
        image_nodes = [node for node in snapshot["nodes"] if node.get("source_item") and node["source_item"]["source_form"] == "IMAGE"]
        self.assertEqual(len(image_nodes), 2)
        for node in image_nodes:
            self.assertEqual(node["display_text"], "")
            self.assertEqual(node["requirement_text"], "")
            self.assertEqual(node["source_item"]["original_representation"], "IMAGE")
            self.assertEqual(node["source_item"]["visual_interpretation"]["authority"], "AGENT_VISUAL_INTERPRETATION")
            self.assertEqual(node["analysis_eligibility"], "ELIGIBLE")
            self.assertEqual(node["ai_assessment"]["decision"], "UNKNOWN")
            self.assertEqual(len(node["figure_refs"]), 1)
        bound_ids = {binding.anchor_id for record in images for binding in record.image_bindings}
        self.assertFalse(any(ref.get("anchor_id") in bound_ids for node in snapshot["nodes"] if not node.get("source_item") for ref in node.get("figure_refs", [])))
        workbook = self.directory / "image-src.xlsx"
        workbook.write_bytes(export_delivery_xlsx(snapshot, source_file_path=str(self.source), source_sha256=baseline.source_sha256))
        with ZipFile(workbook) as archive:
            from PIL import Image
            source_payloads = []
            for name in archive.namelist():
                if not name.startswith("xl/media/"): continue
                with Image.open(BytesIO(archive.read(name))) as image:
                    image = image.convert("RGB")
                    gray = sum(count for count, color in (image.getcolors(image.width*image.height) or []) if all(abs(channel-100)<=2 for channel in color))
                    if 1.8 <= image.width/image.height <= 2.2 and gray/(image.width*image.height) > 0.95:
                        source_payloads.append(archive.read(name))
            self.assertEqual(len(set(source_payloads)), 2, "Both distinct original drawing payloads must reach the workbook")
        values = [str(cell.value) for sheet in load_workbook(workbook) for row in sheet for cell in row if cell.value is not None]
        self.assertIn("SRC-0009", values)
        self.assertIn("REQ-0010", values)
        readings = [cell for sheet in load_workbook(workbook) for row in sheet for cell in row
                    if cell.value and "Agent reading of the visible dimension" in str(cell.value)]
        self.assertEqual(len(readings), 2)
        self.assertTrue(all(cell.column == 11 and "Agent图片释读（以原图为准）" in cell.value for cell in readings),
                        "Agent readings belong in the analysis column, with their authority visible")

    def test_image_compilation_rejects_unknown_anchors_and_omission(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        document_map, baseline, candidate, _ = self.image_input()
        omitted = deepcopy(candidate); omitted["items"] = omitted["items"][:1]
        with self.assertRaisesRegex(ValueError, "Every image anchor"):
            compile_source_item_images(document_map, [omitted], previous_baseline=baseline)
        candidate["items"][0]["image_anchor_ids"] = ["invented-image-anchor"]
        with self.assertRaisesRegex(ValueError, "do not resolve"):
            compile_source_item_images(document_map, [candidate], previous_baseline=baseline)

    def test_image_scope_revisions_preserve_historical_src(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        document_map, baseline, candidate, packet = self.image_input()
        first, _, annotations = compile_source_item_images(document_map, [candidate], previous_baseline=baseline)
        merged = deepcopy(candidate)
        merged["items"] = merged["items"][:1]
        merged["items"][0]["image_anchor_ids"] = [image["anchor_id"] for image in packet["images"]]
        second, _, second_annotations = compile_source_item_images(document_map, [merged], previous_baseline=first, previous_annotations=annotations)
        third, _, _ = compile_source_item_images(document_map, [candidate], previous_baseline=second, previous_annotations=second_annotations)
        self.assertEqual({r.source_record_id:r.source_display_id for r in first.records}, {r.source_record_id:r.source_display_id for r in third.records})
        self.assertEqual([r.source_display_id for r in second.records if r.image_bindings], ["SRC-0011"])

    def test_image_binding_cannot_drift_from_map_or_mix_with_text(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        from app.application.agent_workflow.source_requirement_baseline import SourceRequirementRecordV1, validate_source_image_record_against_map
        document_map, baseline, candidate, _ = self.image_input()
        compiled, _, _ = compile_source_item_images(document_map, [candidate], previous_baseline=baseline)
        record = next(record for record in compiled.records if record.image_bindings)
        changed = record.model_copy(update={"image_bindings": (record.image_bindings[0].model_copy(update={"page": 2}),)})
        with self.assertRaisesRegex(ValueError, "BINDING_DRIFT"):
            validate_source_image_record_against_map(document_map, changed)
        mixed = record.model_dump(mode="json"); mixed["text_bindings"] = baseline.records[0].model_dump(mode="json")["text_bindings"]
        with self.assertRaisesRegex(ValueError, "exactly one"):
            SourceRequirementRecordV1.model_validate(mixed)

    def test_image_only_baseline_needs_no_fake_transcribed_text(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        document_map, _, candidate, _ = self.image_input()
        baseline, receipt, _ = compile_source_item_images(document_map, [candidate])
        self.assertEqual((len(baseline.records), receipt.owned_text_character_count, receipt.owned_image_count), (2, 0, 2))

    def test_mixed_visual_revision_and_image_merge_restore_preserve_history(self):
        from app.application.agent_workflow.source_item_visual import compile_source_item_images, compile_source_item_visual_readings, rebind_source_item_image_annotations
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot
        document_map, previous, candidate, packet = self.image_input()
        tables = json.loads((self.directory / "baseline/table-previews.json").read_text())
        first, _, first_annotations = compile_source_item_images(document_map, [candidate], previous_baseline=previous, table_previews=tables)
        original_image_ids = {record.source_record_id: record.source_display_id for record in first.records if record.image_bindings}
        table = next(record for record in first.records if record.table_fragments)
        reading = {"source_item_id": table.source_record_id, "image_anchor_ids": [packet["images"][0]["anchor_id"]], "text": "This same-section original diagram provides declared context for the table review.", "association_reason": "Explicit test association to the existing original image; no new source item."}
        mixed = compile_source_item_visual_readings(document_map, first, [reading], base_annotations=first_annotations)
        rebound = rebind_source_item_image_annotations(document_map, first, mixed)
        self.assertEqual(rebound["image_identity_registry"], first_annotations["image_identity_registry"])
        for sid in original_image_ids:
            self.assertEqual(rebound["records"][sid]["visual_interpretation"]["text"], first_annotations["records"][sid]["visual_interpretation"]["text"])
            self.assertEqual(rebound["records"][sid]["visual_interpretation"]["source_item_id"], sid)
        merged_candidate = deepcopy(candidate); merged_candidate["items"] = merged_candidate["items"][:1]
        merged_candidate["items"][0]["image_anchor_ids"] = [image["anchor_id"] for image in packet["images"]]
        merged, _, merged_annotations = compile_source_item_images(document_map, [merged_candidate], previous_baseline=first, previous_annotations=mixed, table_previews=tables)
        self.assertEqual(merged_annotations["records"][table.source_record_id]["visual_interpretation"]["text"], reading["text"])
        self.assertGreater(merged_annotations["source_display_high_watermark"], first_annotations["source_display_high_watermark"])
        revised_reading = {**reading, "text": "The same declared diagram context is retained after source-driven review."}
        updated = compile_source_item_visual_readings(document_map, merged, [revised_reading], base_annotations=merged_annotations)
        restored, _, final_annotations = compile_source_item_images(document_map, [candidate], previous_baseline=merged, previous_annotations=updated, table_previews=tables)
        self.assertEqual({record.source_record_id: record.source_display_id for record in restored.records if record.image_bindings}, original_image_ids)
        self.assertEqual(final_annotations["source_display_high_watermark"], merged_annotations["source_display_high_watermark"])
        self.assertTrue(set(merged_annotations["image_identity_registry"]).issubset(final_annotations["image_identity_registry"]))
        self.assertEqual(final_annotations["records"][table.source_record_id]["visual_interpretation"]["text"], revised_reading["text"])
        snapshot = build_delivery_snapshot(document_map, restored, tables, annotations=final_annotations)
        self.assertEqual(snapshot["source_item_baseline"]["source_item_count"], len(first.records))
        self.assertTrue(next(node for node in snapshot["nodes"] if node["node_id"] == table.source_record_id)["source_item"]["visual_interpretation"])

    def test_mixed_visual_revision_keeps_an_unread_image_without_inventing_text(self):
        from app.application.agent_workflow.source_item_visual import compile_source_item_images, compile_source_item_visual_readings, rebind_source_item_image_annotations
        document_map, previous, candidate, packet = self.image_input()
        candidate["items"][0]["visual_interpretation"] = ""
        tables = json.loads((self.directory / "baseline/table-previews.json").read_text())
        baseline, _, annotations = compile_source_item_images(document_map, [candidate], previous_baseline=previous, table_previews=tables)
        table = next(record for record in baseline.records if record.table_fragments)
        mixed = compile_source_item_visual_readings(document_map, baseline, [{"source_item_id": table.source_record_id, "image_anchor_ids": [packet["images"][0]["anchor_id"]], "text": "An explicit reading of the table's shared graphical context.", "association_reason": "Declared same-node context."}], base_annotations=annotations)
        rebind_source_item_image_annotations(document_map, baseline, mixed)
        unread = next(sid for sid, value in annotations["records"].items() if not value["visual_interpretation"]["text"])
        self.assertIsNone(mixed["records"][unread].get("visual_interpretation"))

    def test_text_compilation_reserves_retired_image_src_and_req_numbers(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        from app.application.agent_workflow.source_requirement_baseline import SourceRequirementBaselineArtifactV1
        document_map, _, candidate, packet = self.image_input()
        first, _, annotations = compile_source_item_images(document_map, [candidate])
        merged = deepcopy(candidate); merged["items"] = merged["items"][:1]
        merged["items"][0]["image_anchor_ids"] = [image["anchor_id"] for image in packet["images"]]
        second, _, ann2 = compile_source_item_images(document_map, [merged], previous_baseline=first, previous_annotations=annotations)
        restored, receipt, ann3 = compile_source_item_images(document_map, [candidate], previous_baseline=second, previous_annotations=ann2)
        self.assertEqual(restored.source_display_high_watermark, 3)
        previous = self.directory / "images-only"; previous.mkdir()
        write(previous / "source-requirement-baseline.json", restored.model_dump(mode="json"))
        write(previous / "validation-receipt.json", receipt.model_dump(mode="json"))
        write(previous / "image-annotations.json", ann3)
        decisions = self.directory / "all-text.json"; write(decisions, item_decisions(self.directory / "map"))
        output = self.directory / "images-and-text"
        artifact_tools.compile_items(Namespace(map_dir=self.directory / "map", input=decisions, baseline_dir=previous, output_dir=output))
        after = SourceRequirementBaselineArtifactV1.model_validate_json((output / "source-requirement-baseline.json").read_text())
        self.assertFalse(any(record.source_display_id == "SRC-0003" or record.requirement_display_id == "REQ-0003" for record in after.records if not record.image_bindings))
        rebound = json.loads((output / "image-annotations.json").read_text())
        again, _, _ = compile_source_item_images(document_map, [merged], previous_baseline=after, previous_annotations=rebound)
        image = next(record for record in again.records if record.image_bindings)
        self.assertEqual((image.source_display_id, image.requirement_display_id), ("SRC-0003", "REQ-0003"))

    def test_delivery_rejects_modified_sealed_image_reading(self) -> None:
        from app.application.agent_workflow.source_item_visual import compile_source_item_images
        from app.application.agent_workflow.source_item_delivery import build_delivery_snapshot
        document_map, previous, candidate, _ = self.image_input()
        baseline, _, annotations = compile_source_item_images(document_map, [candidate], previous_baseline=previous)
        source_id = next(iter(annotations["records"]))
        annotations["records"][source_id]["visual_interpretation"]["text"] = "Changed after sealing"
        tables = json.loads((self.directory / "baseline/table-previews.json").read_text())
        with self.assertRaisesRegex(ValueError, "changed after compilation"):
            build_delivery_snapshot(document_map, baseline, tables, annotations=annotations)

    def test_image_cli_and_frozen_assessment_pipeline(self) -> None:
        from openpyxl import load_workbook
        document_map, previous, candidate, _ = self.image_input()
        decisions = self.directory / "images.json"
        write(decisions, {"schema_version": "reqpilot-image-decisions/v1", "source_sha256": previous.source_sha256, "document_map_artifact_sha256": document_map.artifact_sha256, "scopes": [candidate]})
        image_dir = self.directory / "image-baseline"
        artifact_tools.image_operation(Namespace(command="compile-images", map_dir=self.directory / "map", baseline_dir=self.directory / "baseline", input=decisions, output_dir=image_dir))
        packet_dir = self.directory / "analysis"
        result = artifact_tools.analysis_operation(Namespace(command="analysis-packet", baseline_dir=image_dir, knowledge_catalog=None, decomposition=None, unit_id=None, output_dir=packet_dir))
        self.assertEqual((result["unit_count"], result["registered_fact_count"]), (9, 0))
        packet = json.loads((packet_dir / "analysis-packet.json").read_text())
        self.assertEqual(sum(bool(unit["image_refs"]) for unit in packet["units"]), 2)
        response = {"decisions": [{"unit_id": unit["unit_id"], "decision": "PENDING_EVIDENCE", "reason": "原文已明确要求，尚缺该产品的设计或验证证据。", "evidence": [], "clarification_question": "请提供与该要求条件一致的产品设计或测试结果。"} for unit in packet["units"]]}
        response_path = self.directory / "assessments.json"; write(response_path, response)
        result_dir = self.directory / "assessment-result"
        compiled = artifact_tools.analysis_operation(Namespace(command="compile-assessments", analysis_packet=packet_dir / "analysis-packet.json", input=response_path, previous_assessments=None, output_dir=result_dir))
        self.assertEqual(compiled["pending_unit_count"], 9)
        delivery_dir = self.directory / "final"
        artifact_tools.deliver(Namespace(baseline_dir=image_dir, input=None, source=self.source, decomposition=None, assessments=result_dir / "source-item-assessments.json", analysis_packet=None, output_dir=delivery_dir))
        values = [str(cell.value) for sheet in load_workbook(delivery_dir / "matrix.xlsx") for row in sheet for cell in row if cell.value is not None]
        self.assertTrue(any("已分析 9/9" in value and "有效判断 0/9" in value and "待内部证据 9" in value for value in values))
        self.assertIn("REQ-0010", values)
        response["evidence_catalog"] = {"fabricated": {"text": "Product passed"}}
        write(response_path, response)
        with self.assertRaisesRegex(ValueError, "decisions only"):
            artifact_tools.analysis_operation(Namespace(command="compile-assessments", analysis_packet=packet_dir / "analysis-packet.json", input=response_path, previous_assessments=None, output_dir=self.directory / "bad-assessment"))


if __name__ == "__main__":
    unittest.main()
