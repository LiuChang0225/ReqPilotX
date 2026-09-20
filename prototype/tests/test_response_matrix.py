from __future__ import annotations

# ruff: noqa: E402
from .runtime_isolation import activate_test_runtime
activate_test_runtime()

from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
import unittest

from openpyxl import load_workbook
from app.services.exporter import export_compliance_matrix
from app.services.response_matrix import response_units, _result, group_assessment, structure_label, visual_reading, VISUAL_READING_LABEL
from .test_exporter_table_layers import _snapshot


def response_snapshot():
    data = _snapshot()
    data['source_item_baseline'] = {'source_item_count': 2, 'requirement_count': 1, 'non_requirement_count': 1,
                                   'revision': 1, 'source_sha256': 'a'*64}
    shared = {'content_id': 'shared-condition', 'text': 'Apply at 12 V and 25 C.', 'page': 1, 'bbox': [0,0,100,20]}
    for node in data['nodes']:
        if node['node_id'] == 'record':
            node.update(requirement_id='record', requirement_entry_id='record', requirement_code='REQ-001')
            node['source_item'] = {'source_display_id': 'SRC-001', 'requirement_display_id': 'REQ-001',
                'requirementness': 'REQUIREMENT', 'requirement_kinds': ['QUALITY'], 'shared_source_content': [shared]}
        elif node['node_id'] == 'context':
            node.update(display_text='Retained customer background.')
            node['source_item'] = {'source_display_id': 'SRC-002', 'requirement_display_id': None,
                'requirementness': 'NON_REQUIREMENT', 'non_requirement_role': 'CONTEXT',
                'requirement_kinds': [], 'shared_source_content': []}
        elif node['node_id'].startswith('atom-'):
            node['decomposition'] = {'source_item_id': 'record', 'shared_source_context': [shared]}
            node['classification_paths'] = ['PRODUCT_SYSTEM.QUALITY']
            node['topic_tags'] = ['VERIFICATION']
    return data


def with_reading(data, node_id, text='The source drawing shows a mounting direction.'):
    node = next(n for n in data['nodes'] if n['node_id'] == node_id)
    data['source_item_baseline']['artifact_sha256'] = 'b'*64
    node['source_item'].update(source_item_id=node_id, image_bindings=[{'anchor_id': 'image-a'}],
        visual_interpretation={'text': text, 'authority': 'AGENT_VISUAL_INTERPRETATION',
            'source_sha256': 'a'*64, 'baseline_sha256': 'b'*64, 'source_item_id': node_id,
            'image_anchor_ids': ['image-a'], 'review_state': 'UNREVIEWED',
            'unresolved': ['Dimension notation requires confirmation.']})
    return node


class ResponseMatrixTests(unittest.TestCase):
    def test_visual_reading_does_not_cross_source_item_version_or_image_bindings(self):
        data = response_snapshot(); node = with_reading(data, 'record')
        valid = node['source_item']['visual_interpretation']; baseline = data['source_item_baseline']
        self.assertIn(valid['text'], visual_reading(node, baseline))
        for field, value in [('source_sha256','c'*64), ('baseline_sha256','c'*64),
                             ('source_item_id','other-item'), ('image_anchor_ids',['other-image']),
                             ('authority','SOURCE_IMAGE_PIXELS')]:
            changed = deepcopy(node); changed['source_item']['visual_interpretation'][field] = value
            with self.subTest(field=field):
                displayed = visual_reading(changed, baseline)
                self.assertNotIn(valid['text'], displayed)
                self.assertIn('关联不匹配', displayed)
        legacy = deepcopy(node)
        for field in ('baseline_sha256','source_item_id'):legacy['source_item']['visual_interpretation'].pop(field)
        self.assertIn(valid['text'], visual_reading(legacy, baseline))

    def test_table_reading_has_its_own_column_and_does_not_rewrite_original_content(self):
        data = response_snapshot(); node = with_reading(data, 'record'); original = deepcopy(data)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            wb = load_workbook(path); main = wb['ComplianceMatrix']
            row = next(row for row in main.iter_rows(min_row=4) if row[0].value == 'SRC-001')
            reading = node['source_item']['visual_interpretation']['text']
            self.assertNotIn(reading, row[3].value)
            self.assertEqual(row[1].value, 'REQ-001')
            table = next(ws for ws in wb if ws.title.startswith('T001'))
            columns = {c.value:c.column for c in table[3] if c.value}
            self.assertIn(VISUAL_READING_LABEL, columns)
            values = '\n'.join(str(row[columns[VISUAL_READING_LABEL]-1].value or '') for row in table.iter_rows(min_row=4))
            self.assertIn(reading, values); self.assertIn('待人工复核', values)
            self.assertIn('待核对：Dimension notation', values)
            self.assertEqual(main.max_row-3, sum(n['matrix_visibility']!='SYSTEM_NON_BUSINESS' for n in data['nodes']))
        self.assertEqual(data, original)

    def test_long_non_table_reading_is_labelled_and_preserved_in_comment(self):
        data = response_snapshot(); text = 'Illustration interpretation. '*100
        with_reading(data, 'context', text)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            main = load_workbook(path)['ComplianceMatrix']; columns = {c.value:c.column-1 for c in main[3]}
            row = next(row for row in main.iter_rows(min_row=4) if row[0].value=='SRC-002')
            self.assertEqual(row[3].value, 'Retained customer background.')
            note = row[columns['说明 / 待补充']]
            self.assertIn(VISUAL_READING_LABEL, note.value)
            self.assertIn('完整释读见本单元格批注', note.value)
            self.assertIn(text, note.comment.text)
            self.assertEqual(row[columns['内容判定']].value, '非需求')

    def test_document_table_and_requirement_identity_are_independent(self):
        self.assertEqual(structure_label({'source_kind':'DOCUMENT_UNIT', 'display_role':'TABLE'}), '文档')
        self.assertEqual(structure_label({'source_kind':'SECTION', 'display_role':'STRUCTURE'}), '章节')
        self.assertEqual(structure_label({'source_item':{'source_form':'TABLE_BRANCH'}, 'display_role':'TABLE_RECORD'}), '表格分支')
        self.assertEqual(structure_label({'source_item':{'requirementness':'NON_REQUIREMENT'}, 'display_role':'TABLE_RECORD'}), '表格条目')

    def test_pending_human_response_does_not_fall_back_to_ai_compliance(self):
        data = response_snapshot()
        node = next(n for n in data['nodes'] if n['node_id'] == 'atom-1')
        node.update(analysis_result_state='CURRENT', analysis_status='DECIDED',
            manual_decision={'decision': 'PENDING_EVIDENCE', 'reason': 'Need the released report.'},
            ai_assessment={'decision': 'COMPLIANT', 'reason': 'Previous advice.'})
        label, reason = _result(node)
        self.assertEqual(label, '人工：待内部证据')
        self.assertEqual(reason, 'Need the released report.')

    def test_ordinary_child_structure_does_not_remove_an_original_requirement(self):
        nodes = [{'node_id': 'original', 'requirement_id': 'req', 'child_ids': ['note'],
                  'matrix_visibility': 'VISIBLE', 'disposition': 'ANALYZE', 'analysis_eligibility': 'ELIGIBLE'},
                 {'node_id': 'note', 'requirement_id': None, 'child_ids': []}]
        self.assertEqual(response_units(nodes), [nodes[0]])

    def test_original_or_children_are_response_rows_without_structure_or_parent_duplicates(self):
        data = response_snapshot()
        self.assertEqual([n['node_id'] for n in response_units(data['nodes'])], ['atom-1', 'atom-2'])

    def test_export_count_uses_analysis_scope_when_source_and_display_ids_differ(self):
        data = response_snapshot()
        parent = next(n for n in data['nodes'] if n['node_id'] == 'record')
        parent['source_item']['source_item_id'] = 'source-record-original'
        for node in data['nodes']:
            if node.get('decomposition'):
                node['decomposition']['source_item_id'] = 'source-record-original'
        self.assertEqual([n['node_id'] for n in response_units(data['nodes'])], ['atom-1', 'atom-2'])
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            main = load_workbook(path)['ComplianceMatrix']
            self.assertIn('2 个分析单元', main['A2'].value)
            parent_row = next(row for row in main.iter_rows(min_row=4) if row[1].value == 'REQ-001')
            self.assertIn('逐项判断见派生子项', parent_row[10].value)

    def test_complete_hierarchy_and_content_share_one_matrix_without_repeated_context(self):
        data = response_snapshot(); original = deepcopy(data)
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames[0], 'ComplianceMatrix')
            self.assertNotIn('00_清单概览', wb.sheetnames)
            self.assertNotIn('全文详情', wb.sheetnames)
            self.assertNotIn('原文索引', wb.sheetnames)
            self.assertNotIn('共用条件', wb.sheetnames)
            main = wb['ComplianceMatrix']
            visible = [n for n in data['nodes'] if n['matrix_visibility'] != 'SYSTEM_NON_BUSINESS']
            self.assertEqual(main.max_row-3, len(visible))
            combined = '\n'.join(str(c.value or '') for row in main for c in row)
            self.assertEqual(combined.count('Apply at 12 V and 25 C.'), 1)
            self.assertIn('Retained customer background.', combined)
            self.assertIn('REQ-001', combined)  # Complete original parent stays visible.
            self.assertGreater(max(d.outlineLevel for d in main.row_dimensions.values()), 0)
            self.assertFalse(any(d.hidden for d in main.row_dimensions.values()))
            source_column = next(c.column - 1 for c in main[3] if c.value == '来源 / 原表')
            table_links = [row[source_column] for row in main.iter_rows(min_row=4) if row[source_column].hyperlink and (row[source_column].hyperlink.location or '').startswith("'T001")]
            self.assertGreaterEqual(len(table_links), 3)
            original_row = next(row[0].row for row in main.iter_rows(min_row=4) if row[1].value == 'REQ-001')
            table_sheet = next(ws for ws in wb if ws.title.startswith('T001'))
            self.assertTrue(any(c.hyperlink and c.hyperlink.location == f"'ComplianceMatrix'!A{original_row}"
                                for row in table_sheet for c in row))
            analysis_columns = {c.value:c.column for c in table_sheet[3] if c.value}
            self.assertIn('判定依据（需求提取）', analysis_columns)
            result_column = analysis_columns['响应判断']
            result_cells = [row[result_column-1] for row in table_sheet.iter_rows(min_row=4)
                            if row[result_column-1].data_type == 'f']
            self.assertTrue(result_cells)
            self.assertTrue(all('ComplianceMatrix' in c.value for c in result_cells))
            cached = load_workbook(path, data_only=True)[table_sheet.title]
            self.assertEqual(cached[result_cells[0].coordinate].value, '已完成逐项判断')
            self.assertIn('REQ-002', '\n'.join(str(c.value or '') for row in table_sheet for c in row
                                             if c.column == analysis_columns['响应分析 / 依据 / 缺口']))
            self.assertEqual(wb['_ReqPilotXMeta'].sheet_state, 'hidden')
        self.assertEqual(data, original, 'Export must not rewrite the frozen source/derivation graph')

    def test_requirementness_content_role_and_source_structure_remain_separate(self):
        data = response_snapshot()
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            main = load_workbook(path)['ComplianceMatrix']
            columns = {cell.value: cell.column-1 for cell in main[3]}
            judgment, role = columns['内容判定'], columns['内容角色']
            rows = list(main.iter_rows(min_row=4))
            source = next(row for row in rows if row[0].value == 'SRC-001')
            context = next(row for row in rows if row[0].value == 'SRC-002')
            self.assertEqual(source[judgment].value, '需求')
            self.assertIsNone(source[role].value)
            self.assertEqual(context[judgment].value, '非需求')
            self.assertEqual(context[role].value, '背景/上下文')
            for row in rows:
                if not row[0].value and not row[1].value:
                    self.assertIsNone(row[judgment].value, 'Structure is not a content judgment')
                    self.assertIsNone(row[role].value)
            self.assertEqual({row[judgment].value for row in rows}, {None, '需求', '非需求'})
            derived = [row for row in rows if not row[0].value and row[1].value]
            self.assertTrue(derived)
            self.assertTrue(all(row[judgment].value == '需求' for row in derived))

    def test_shared_source_reuses_existing_reading_row_and_never_reprints_raw_paint_text(self):
        data = response_snapshot()
        context = next(n for n in data['nodes'] if n['node_id'] == 'context')
        context['source_refs'] = [{'kind': 'TEXT', 'content_id': 'shared-condition', 'start': 0,
                                   'end': len('Apply at 12 V and 25 C.')}]
        context['display_text'] = 'Proved reading view of the existing source.'
        context['structure_node_id'] = context['node_id']
        context['display_title'] = 'Parser fallback must not become source text'
        parent = next(n for n in data['nodes'] if n['node_id'] == 'record')
        heading = next(n for n in data['nodes'] if n['node_id'] == 'table')
        heading['display_title'] = '未命名原文节点'
        parent['source_item']['shared_source_content'].append(
            {'content_id': 'title', 'text': 'Heading\nHeading\nHeading', 'is_title': True})
        with TemporaryDirectory() as folder:
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            main = load_workbook(path)['ComplianceMatrix']
            combined = '\n'.join(str(c.value or '') for row in main for c in row)
            self.assertEqual(combined.count(context['display_text']), 1)
            self.assertNotIn('Apply at 12 V and 25 C.', combined)
            self.assertNotIn('Heading\nHeading', combined)
            self.assertNotIn(context['display_title'], combined)
            self.assertNotIn('未命名原文节点', combined)
            self.assertIn('共用原文见主表行', combined)

    def test_table_summary_exposes_partial_failure_and_pending_units(self):
        nodes = response_snapshot()['nodes']; by_id = {n['node_id']: n for n in nodes}
        by_id['atom-1'].update(analysis_result_state='CURRENT', analysis_status='ASSESSED',
            ai_assessment={'decision': 'NON_COMPLIANT', 'reason': 'Different threshold.'}, manual_decision=None)
        by_id['atom-2'].update(analysis_result_state='NONE', analysis_status='NOT_STARTED',
            ai_assessment=None, manual_decision=None, compliance_status='')
        status, detail = group_assessment(by_id['table'], by_id)
        self.assertEqual(status, '存在不符合')
        self.assertIn('1/2', detail)
        self.assertIn('待分析 1', detail)

    def test_pending_evidence_and_clarification_are_analyzed_not_unstarted(self):
        nodes = response_snapshot()['nodes']; by_id = {n['node_id']: n for n in nodes}
        for key, outcome in (('atom-1', 'PENDING_EVIDENCE'), ('atom-2', 'NEEDS_CLARIFICATION')):
            by_id[key].update(analysis_result_state='CURRENT', analysis_status='ASSESSED',
                ai_assessment={'decision': outcome, 'reason': 'Source has been assessed; required input remains missing.'},
                manual_decision=None)
        status, detail = group_assessment(by_id['table'], by_id)
        self.assertEqual(status, '已分析 · 待证据/澄清')
        self.assertIn('已分析 2/2', detail)
        self.assertIn('有效判断 0/2', detail)
        self.assertIn('待内部证据 1', detail)
        self.assertIn('待澄清 1', detail)
        self.assertNotIn('待分析', detail)
        with TemporaryDirectory() as folder:
            data = response_snapshot(); data['nodes'] = nodes
            data['counts'].update(assessed_nodes=2, decided_nodes=0)
            _, path = export_compliance_matrix('Example', '', Path(folder), matrix_snapshot=data, export_mode='DRAFT')
            workbook = load_workbook(path)
            texts = '\n'.join(str(cell.value or '') for sheet in workbook for row in sheet for cell in row)
            self.assertIn('已分析 2/2', texts)
            self.assertNotIn('待完成逐项判断', texts)

    def test_current_and_stale_results_have_separate_analysis_progress(self):
        nodes = response_snapshot()['nodes']; by_id = {n['node_id']: n for n in nodes}
        by_id['atom-1'].update(analysis_result_state='CURRENT', analysis_status='ASSESSED',
            ai_assessment={'decision': 'PENDING_EVIDENCE', 'reason': 'Missing internal evidence.'}, manual_decision=None)
        by_id['atom-2'].update(analysis_result_state='STALE_REFERENCE', analysis_status='ASSESSED',
            ai_assessment={'decision': 'COMPLIANT', 'reason': 'Old result.'}, manual_decision=None)
        status, detail = group_assessment(by_id['table'], by_id)
        self.assertEqual(status, '部分已分析 · 仍需补全')
        self.assertIn('已分析 1/2', detail)
        self.assertIn('需重新分析 1', detail)
        self.assertIn('有效判断 0/2', detail)


if __name__ == '__main__':
    unittest.main()
