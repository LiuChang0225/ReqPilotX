"""Actual PDF/compiler/import/review/Excel path for declared user work opinions."""
import json
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import test_review_loop as fixture
from test_portable import artifact_tools
import manual_tools
from app.application.agent_workflow.source_item_decomposition import build_source_item_decomposition_packet, compile_source_item_decomposition
args,loop=fixture.args,fixture.loop


class ManualReviewTests(unittest.TestCase):
    setUp=fixture.ReviewLoopTests.setUp
    tearDown=fixture.ReviewLoopTests.tearDown
    put=fixture.ReviewLoopTests.put

    def prepare(self):
        _,base,_,tables=loop.bundle(self.baseline_dir)
        src=base.records[0].source_record_id
        text=build_source_item_decomposition_packet(base,src,table_previews=tables)['segments'][0]['text'];cut=text.index('\n')+1
        decomp=compile_source_item_decomposition(base,[{'source_item_id':src,'action':'SPLIT','reason':'Synthetic distinct opening and closing obligations.',
            'children':[{'text':text[:cut].strip(),'source_ranges':[{'segment_index':0,'start_offset':0,'end_offset':cut}],'requirement_kinds':['QUALITY']},
                        {'text':text[cut:].strip(),'source_ranges':[{'segment_index':0,'start_offset':cut,'end_offset':len(text)}],'requirement_kinds':['QUALITY']}]}],table_previews=tables)
        dp=self.put('decomposition.json',decomp)
        packetdir=self.root/'analysis'
        artifact_tools.analysis_operation(args(command='analysis-packet',baseline_dir=self.baseline_dir,decomposition=dp,unit_id=None,output_dir=packetdir))
        packet=json.loads((packetdir/'analysis-packet.json').read_text())
        assessed=self.root/'assessed'
        artifact_tools.analysis_operation(args(command='compile-assessments',analysis_packet=packetdir/'analysis-packet.json',
            input=self.put('ai.json',{'decisions':[{'unit_id':u['unit_id'],'decision':'PENDING_EVIDENCE','reason':'原AI建议：尚无内部试验记录。','evidence':[],
                'clarification_question':'请提供内部验证。'} for u in packet['units']]}),previous_assessments=None,output_dir=assessed))
        self.review=self.root/'manual-review'
        head=loop.open_review(args(review_dir=self.review,baseline_dir=self.baseline_dir,source=self.source,decomposition=dp,
            assessments=assessed/'source-item-assessments.json',analysis_packet=assessed/'analysis-packet.json'))
        self.current_dir=self.review/'versions/initial';self.context=manual_tools.context_for(self.current_dir)
        self.unit=decomp['records'][src][0]['singular_requirement_id']
        self.direct=next(u for u in self.context['units'] if u in {r.source_record_id for r in base.records})
        return head

    def import_and_advance(self,head,name,decisions):
        imported=self.root/(name+'-import')
        manual_tools.run_manual(args(command='import-manual-decisions',baseline_dir=self.current_dir,input=self.put(name+'-user.json',{
            'schema_version':'source-item-user-decisions/v1','expected_source_sha256':self.context['source_sha256'],
            'expected_baseline_sha256':self.context['baseline_sha256'],'user_input_reference':'Explicit synthetic user work opinions', 'decisions':decisions}),output_dir=imported))
        candidate=self.review/'versions'/name
        loop.revise_review(args(review_dir=self.review,input=self.put(name+'-revision.json',{'schema_version':'reqpilot-review-revision/v1',
            'expected_review_sha256':head['review_sha256'],'expected_baseline_sha256':head['baseline_sha256']}),
            manual_decisions=imported/'manual-decisions.json',output_dir=candidate))
        delivery=self.review/'deliveries'/name
        artifact_tools.deliver(args(baseline_dir=candidate,source=self.source,decomposition=candidate/'source-item-decomposition.json',
            assessments=candidate/'source-item-assessments.json',analysis_packet=candidate/'analysis-packet.json',output_dir=delivery))
        advanced=loop.close_review(args(review_dir=self.review,input=self.put(name+'-close.json',{'expected_review_sha256':head['review_sha256'],
            'candidate_path':str(candidate.relative_to(self.review)),'delivery_path':str(delivery.relative_to(self.review))})))
        self.current_dir=candidate
        return advanced,delivery

    def test_user_src_and_snn_import_priority_ai_retention_withdraw_and_source_stale(self):
        head=self.prepare()
        head,delivery=self.import_and_advance(head,'user-opinion',[{'unit_id':u,'decision':'NON_COMPLIANT','reason':'合成工程师工作意见',
            'conditions':'仅限当前接口方案','evidence_requests':'请补充正式图纸'} for u in (self.unit,self.direct)])
        snapshot=json.loads((delivery/'matrix-snapshot.json').read_text())
        for uid in (self.unit,self.direct):
            node=next(n for n in snapshot['nodes'] if n.get('requirement_id')==uid)
            self.assertEqual(node['manual_decision']['decision'],'NON_COMPLIANT')
            self.assertEqual(node['ai_assessment']['decision'],'PENDING_EVIDENCE')
        wb=load_workbook(delivery/'matrix.xlsx');main=wb['ComplianceMatrix'];target=next(n for n in snapshot['nodes'] if n['requirement_id']==self.unit)
        row=next(r for r in main.iter_rows() if r[1].value==target['requirement_code'])
        self.assertIn('工程师工作意见：不符合',row[9].value)
        self.assertIn('仅限当前接口方案',row[10].value);self.assertIn('原AI建议',row[10].value)
        self.assertEqual(loop.verify_review(args(review_dir=self.review))['status'],'REVIEW_VERIFIED')
        head,withdrawal=self.import_and_advance(head,'withdrawal',[{'unit_id':self.unit,'action':'WITHDRAW','reason':'用户撤回本条工作意见'}])
        final=json.loads((withdrawal/'matrix-snapshot.json').read_text());node=next(n for n in final['nodes'] if n['requirement_id']==self.unit)
        self.assertIsNone(node['manual_decision']);self.assertEqual(node['ai_assessment']['decision'],'PENDING_EVIDENCE')
        self.assertEqual(loop.verify_review(args(review_dir=self.review))['status'],'REVIEW_VERIFIED')
        # A genuine classification change creates a new source context. Human
        # opinions are program-carried as STALE, not silently relabelled current.
        candidate=self.review/'versions/classified'
        loop.revise_review(args(review_dir=self.review,input=self.put('classify.json',{'schema_version':'reqpilot-review-revision/v1',
            'expected_review_sha256':head['review_sha256'],'expected_baseline_sha256':head['baseline_sha256'],'model_id':'synthetic-review',
            'decisions':[{'source_record_id':self.direct,'requirementness':'REQUIREMENT','requirement_kinds':['CONSTRAINT'],'confidence':.9,
                'decision_summary':'Explicit synthetic classification correction.','reason_codes':['source_review']}]}),output_dir=candidate))
        opinions=json.loads((candidate/'manual-decisions.json').read_text())
        self.assertEqual(opinions['records'][self.direct]['state'],'STALE')
        loop.verify_transition(self.review,self.current_dir,candidate,*self._transition_args(candidate))

    def _transition_args(self,candidate):
        dm,current,_,tables=loop.bundle(candidate);previous=loop.bundle(self.current_dir)[1]
        return dm,previous,current,tables

    def test_native_import_and_inline_manual_are_rejected(self):
        head=self.prepare()
        with patch.dict('os.environ',{'REQPILOT_NATIVE_AGENT_TOOLS':'1'}):
            with self.assertRaisesRegex(ValueError,'cannot create'):
                manual_tools.run_manual(args(command='import-manual-decisions',baseline_dir=self.current_dir,input=self.put('none.json',{}),output_dir=self.root/'forbidden'))
        with self.assertRaisesRegex(ValueError,'Inline review JSON'):
            loop.revise_review(args(review_dir=self.review,input=self.put('inline.json',{'schema_version':'reqpilot-review-revision/v1',
                'expected_review_sha256':head['review_sha256'],'expected_baseline_sha256':head['baseline_sha256'],'manual_decisions':{}}),output_dir=self.review/'versions/forged'))


if __name__=='__main__':unittest.main()
