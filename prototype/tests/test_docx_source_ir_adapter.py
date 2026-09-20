from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

import base64
import hashlib
import shutil
import tempfile
import unittest
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app.adapters.documents.docx_source_ir_adapter import DocxSourceIRAdapter
from app.adapters.documents.preflight import preflight_document
from app.domain.extraction.preflight import DocumentFormat
from app.domain.extraction.source_ir import (
    SourceIRStatus,
    SourceNodeKind,
)
from app.ports.source_ir_adapter import (
    SourceIRAdapterPort,
    SourceIRAdapterRejected,
    SourceIRAdapterRejectionCode,
)


CREATED_AT = datetime(2026, 7, 20, 12, 0, tzinfo=timezone.utc)
_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


class DocxSourceIRAdapterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)
        self.adapter = DocxSourceIRAdapter()

    def test_body_order_lists_tables_merges_images_and_comments_are_faithful(self) -> None:
        path = self._write_rich_docx("requirements.docx")
        preflight = preflight_document(path)
        source_directory_before = sorted(item.name for item in self.root.iterdir())

        first = self.adapter.extract(path, preflight=preflight, created_at=CREATED_AT)
        second = self.adapter.extract(path, preflight=preflight, created_at=CREATED_AT)

        self.assertEqual(first.status, SourceIRStatus.COMPLETE)
        self.assertEqual(first, second)
        self.assertEqual(
            [node.kind for node in first.nodes],
            [
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.LIST_ITEM,
                SourceNodeKind.TABLE,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.PARAGRAPH,
                SourceNodeKind.IMAGE,
            ],
        )
        self.assertEqual(
            [node.source_order for node in first.nodes],
            list(range(len(first.nodes))),
        )
        self.assertEqual(first.nodes[0].raw_text, "父需求原文")
        self.assertEqual(first.nodes[0].native_locator.body_element_index, 0)
        self.assertEqual(first.nodes[0].native_locator.paragraph_index, 0)
        self.assertEqual(first.nodes[1].raw_text, "同级编号 1")
        self.assertEqual(first.nodes[1].docx_list_metadata.num_id, 1)
        self.assertEqual(first.nodes[1].docx_list_metadata.ilvl, 0)
        self.assertEqual(first.nodes[1].native_locator.body_element_index, 1)
        self.assertEqual(first.nodes[1].native_locator.paragraph_index, 1)
        self.assertEqual(first.nodes[10].raw_text, "表后正文")
        self.assertEqual(first.nodes[10].native_locator.body_element_index, 3)
        self.assertEqual(first.nodes[10].native_locator.paragraph_index, 2)

        table_node = first.nodes[2]
        table = table_node.table_payload
        assert table is not None
        self.assertEqual(table_node.native_locator.body_element_index, 2)
        self.assertEqual(table_node.native_locator.table_index, 0)
        self.assertEqual((table.row_count, table.column_count), (3, 3))
        cells = {(cell.row, cell.column): cell for cell in table.cells}
        self.assertEqual(len(cells), 7)
        self.assertEqual(cells[(0, 0)].raw_text, "共享验收条件")
        self.assertEqual(cells[(0, 0)].column_span, 2)
        self.assertNotIn((0, 1), cells)
        self.assertEqual(cells[(1, 0)].raw_text, "共同材料")
        self.assertEqual(cells[(1, 0)].row_span, 2)
        self.assertNotIn((2, 0), cells)
        self.assertEqual(
            {
                (region.row, region.column, region.row_span, region.column_span)
                for region in table.merged_regions
            },
            {(0, 0, 1, 2), (1, 0, 2, 1)},
        )
        for (row, column), cell in cells.items():
            self.assertEqual(cell.native_locator.body_element_index, 2)
            self.assertEqual(cell.native_locator.table_index, 0)
            self.assertEqual(cell.native_locator.row_index, row)
            self.assertEqual(cell.native_locator.column_index, column)

        table_paragraphs = first.nodes[3:10]
        self.assertEqual(
            [node.raw_text for node in table_paragraphs],
            [
                "共享验收条件",
                "判定",
                "共同材料",
                "参数 A",
                "阈值 A",
                "参数 B",
                "阈值 B",
            ],
        )
        self.assertTrue(
            all(
                node.parent_source_node_id == table_node.source_node_id
                for node in table_paragraphs
            )
        )

        image_node = first.nodes[11]
        self.assertEqual(image_node.native_locator.body_element_index, 4)
        self.assertEqual(image_node.native_locator.paragraph_index, 3)
        self.assertEqual(len(first.assets), 1)
        self.assertEqual(image_node.asset_ref, first.assets[0].asset_id)
        self.assertEqual(first.assets[0].media_type, "image/png")
        self.assertTrue(first.assets[0].storage_key.startswith("source-assets/"))
        self.assertIn(first.assets[0].content_sha256, first.assets[0].storage_key)
        self.assertEqual(
            {item.code for item in first.diagnostics},
            {"docx_comments_ignored_by_policy"},
        )
        self.assertEqual(
            sorted(item.name for item in self.root.iterdir()),
            source_directory_before,
        )

    def test_table_cell_physical_paragraphs_are_parallel_ordered_facts(self) -> None:
        path = self._write_dht120_cell_shape("dht120-2.1.22.docx")

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(
            [node.kind for node in snapshot.nodes],
            [
                SourceNodeKind.TABLE,
                SourceNodeKind.HEADING,
                SourceNodeKind.LIST_ITEM,
                SourceNodeKind.LIST_ITEM,
                SourceNodeKind.IMAGE,
                SourceNodeKind.PARAGRAPH,
            ],
        )
        self.assertEqual(
            [node.source_order for node in snapshot.nodes],
            list(range(len(snapshot.nodes))),
        )

        table_node = snapshot.nodes[0]
        assert table_node.table_payload is not None
        self.assertEqual(
            table_node.table_payload.cells[0].raw_text,
            "2.1.22 高低温特性\n"
            "样本数量：6件/OTS\n"
            "前置条件：试验前静置2小时\n"
            "验收判定：性能参数应满足限值",
        )
        paragraph_nodes = (
            snapshot.nodes[1],
            snapshot.nodes[2],
            snapshot.nodes[3],
            snapshot.nodes[5],
        )
        self.assertEqual(
            [node.raw_text for node in paragraph_nodes],
            [
                "2.1.22 高低温特性",
                "样本数量：6件/OTS",
                "前置条件：试验前静置2小时",
                "验收判定：性能参数应满足限值",
            ],
        )
        self.assertEqual(
            [node.native_locator.paragraph_index for node in paragraph_nodes],
            [0, 1, 2, 3],
        )
        for node in paragraph_nodes:
            self.assertEqual(node.parent_source_node_id, table_node.source_node_id)
            self.assertEqual(node.native_locator.body_element_index, 0)
            self.assertEqual(node.native_locator.table_index, 0)
            self.assertEqual(node.native_locator.row_index, 0)
            self.assertEqual(node.native_locator.column_index, 0)

        self.assertEqual(snapshot.nodes[2].docx_list_metadata.num_id, 42)
        self.assertEqual(snapshot.nodes[2].docx_list_metadata.ilvl, 2)
        self.assertEqual(snapshot.nodes[3].docx_list_metadata.num_id, 42)
        self.assertIsNone(snapshot.nodes[3].docx_list_metadata.ilvl)
        self.assertIsNone(snapshot.nodes[1].docx_list_metadata)
        self.assertIsNone(snapshot.nodes[5].docx_list_metadata)

        image_node = snapshot.nodes[4]
        self.assertEqual(image_node.parent_source_node_id, snapshot.nodes[3].source_node_id)
        self.assertEqual(image_node.native_locator.paragraph_index, 2)
        self.assertEqual(image_node.native_locator.object_index, 0)

    def test_table_cell_paragraphs_inside_block_wrappers_are_not_lost(self) -> None:
        path = self.root / "wrapped-table-cell.docx"
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)

        controlled = cell.paragraphs[0]
        controlled.add_run("内容控件中的单元格正文")
        cell._tc.remove(controlled._p)
        content_control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        content.append(controlled._p)
        content_control.append(content)
        cell._tc.append(content_control)

        revised = cell.add_paragraph("修订包装中的验收判定")
        cell._tc.remove(revised._p)
        custom_xml = OxmlElement("w:customXml")
        insertion = OxmlElement("w:ins")
        insertion.append(revised._p)
        custom_xml.append(insertion)
        cell._tc.append(custom_xml)
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        table_node = snapshot.nodes[0]
        self.assertEqual(table_node.kind, SourceNodeKind.TABLE)
        assert table_node.table_payload is not None
        self.assertEqual(
            table_node.table_payload.cells[0].raw_text,
            "内容控件中的单元格正文\n修订包装中的验收判定",
        )
        self.assertEqual(
            [node.raw_text for node in snapshot.nodes[1:]],
            ["内容控件中的单元格正文", "修订包装中的验收判定"],
        )
        self.assertEqual(
            [node.native_locator.paragraph_index for node in snapshot.nodes[1:]],
            [0, 1],
        )
        self.assertTrue(
            all(
                node.parent_source_node_id == table_node.source_node_id
                for node in snapshot.nodes[1:]
            )
        )

    def test_table_rows_and_cells_inside_block_wrappers_are_not_lost(self) -> None:
        path = self.root / "wrapped-table-grid.docx"
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "直接行包装单元格"
        table.cell(0, 1).text = "直接行第二格"
        table.cell(1, 0).text = "包装行第一格"
        table.cell(1, 1).text = "包装行第二格"

        first_row = table.rows[0]._tr
        first_cell = first_row.findall(qn("w:tc"))[0]
        first_cell_index = first_row.index(first_cell)
        first_row.remove(first_cell)
        cell_control = OxmlElement("w:sdt")
        cell_content = OxmlElement("w:sdtContent")
        cell_content.append(first_cell)
        cell_control.append(cell_content)
        first_row.insert(first_cell_index, cell_control)

        second_row = table.rows[1]._tr
        table._tbl.remove(second_row)
        row_control = OxmlElement("w:sdt")
        row_content = OxmlElement("w:sdtContent")
        row_content.append(second_row)
        row_control.append(row_content)
        table._tbl.append(row_control)
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        table_node = snapshot.nodes[0]
        assert table_node.table_payload is not None
        payload = table_node.table_payload
        self.assertEqual((payload.row_count, payload.column_count), (2, 2))
        self.assertEqual(
            {(cell.row, cell.column): cell.raw_text for cell in payload.cells},
            {
                (0, 0): "直接行包装单元格",
                (0, 1): "直接行第二格",
                (1, 0): "包装行第一格",
                (1, 1): "包装行第二格",
            },
        )
        self.assertEqual(
            [node.raw_text for node in snapshot.nodes[1:]],
            [
                "直接行包装单元格",
                "直接行第二格",
                "包装行第一格",
                "包装行第二格",
            ],
        )

    def test_alternate_content_selects_one_visible_branch_without_duplicate_text(
        self,
    ) -> None:
        path = self.root / "alternate-content.docx"
        document = Document()
        body_paragraph = document.add_paragraph()
        body_paragraph._p.append(
            self._alternate_text("正文只应出现一次", "正文只应出现一次")
        )
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0]._p.append(
            self._alternate_text("单元格只应出现一次", "单元格只应出现一次")
        )
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(snapshot.nodes[0].raw_text, "正文只应出现一次")
        table_node = snapshot.nodes[1]
        assert table_node.table_payload is not None
        self.assertEqual(
            table_node.table_payload.cells[0].raw_text,
            "单元格只应出现一次",
        )
        self.assertEqual(snapshot.nodes[2].raw_text, "单元格只应出现一次")
        serialized_text = "\n".join(node.raw_text for node in snapshot.nodes)
        self.assertNotIn("正文只应出现一次正文只应出现一次", serialized_text)
        self.assertNotIn("单元格只应出现一次单元格只应出现一次", serialized_text)

    def test_alternate_content_uses_fallback_for_unsupported_choice(self) -> None:
        path = self.root / "alternate-content-fallback.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph._p.append(
            self._alternate_text(
                "不应读取的未来格式",
                "VISIBLE FALLBACK REQUIREMENT",
                required_prefix="future",
                required_namespace="urn:reqpilotx:test:future",
            )
        )
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(snapshot.nodes[0].raw_text, "VISIBLE FALLBACK REQUIREMENT")
        self.assertNotIn(
            "不应读取的未来格式",
            "\n".join(node.raw_text for node in snapshot.nodes),
        )

    def test_alternate_content_without_supported_choice_or_fallback_is_partial(
        self,
    ) -> None:
        path = self.root / "alternate-content-unrepresented.docx"
        document = Document()
        paragraph = document.add_paragraph()
        alternate = self._alternate_text(
            "不可安全解释的未来格式",
            None,
            required_prefix="future",
            required_namespace="urn:reqpilotx:test:future",
        )
        paragraph._p.append(alternate)
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.PARTIAL)
        self.assertEqual(snapshot.nodes, ())
        self.assertIn(
            "docx_alternate_content_not_represented",
            {item.code for item in snapshot.diagnostics},
        )

    def test_table_cell_paragraph_nodes_are_included_in_source_node_budget(self) -> None:
        path = self._write_dht120_cell_shape("dht120-node-budget.docx")

        with patch(
            "app.adapters.documents.docx_source_ir_adapter._MAX_DOCX_SOURCE_NODES",
            1,
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight_document(path),
                created_at=CREATED_AT,
            )

        self._assert_resource_budget_failure(snapshot)

    def test_style_inferred_list_item_does_not_fabricate_direct_metadata(self) -> None:
        path = self.root / "style-only-list.docx"
        document = Document()
        document.add_paragraph("仅由样式识别的列表项", style="List Bullet")
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(len(snapshot.nodes), 1)
        self.assertEqual(snapshot.nodes[0].kind, SourceNodeKind.LIST_ITEM)
        self.assertIsNone(snapshot.nodes[0].docx_list_metadata)

    def test_unparsed_header_relationship_is_explicitly_partial(self) -> None:
        path = self.root / "header.docx"
        document = Document()
        document.add_paragraph("正文")
        document.sections[0].header.paragraphs[0].text = "页眉上下文"
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.PARTIAL)
        self.assertIn(
            "docx_related_content_not_parsed",
            {item.code for item in snapshot.diagnostics},
        )

    def test_repeated_image_placements_keep_two_nodes_and_one_asset(self) -> None:
        path = self._write_repeated_image_docx("repeated-images.docx", placement_count=2)
        preflight = preflight_document(path)
        original_read = zipfile.ZipFile.read
        original_sha256 = hashlib.sha256
        media_reads: list[str] = []
        media_hashes: list[bytes] = []

        def counting_read(
            archive: zipfile.ZipFile,
            member: str | zipfile.ZipInfo,
            *args,
            **kwargs,
        ) -> bytes:
            member_name = member.filename if isinstance(member, zipfile.ZipInfo) else member
            if member_name.startswith("word/media/"):
                media_reads.append(member_name)
            return original_read(archive, member, *args, **kwargs)

        def counting_sha256(payload: bytes = b"", *args, **kwargs):
            if payload == _ONE_PIXEL_PNG:
                media_hashes.append(payload)
            return original_sha256(payload, *args, **kwargs)

        with (
            patch.object(zipfile.ZipFile, "read", new=counting_read),
            patch(
                "app.adapters.documents.docx_source_ir_adapter.hashlib.sha256",
                side_effect=counting_sha256,
            ),
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight,
                created_at=CREATED_AT,
            )

        image_nodes = [node for node in snapshot.nodes if node.kind is SourceNodeKind.IMAGE]
        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(len(image_nodes), 2)
        self.assertEqual(len(snapshot.assets), 1)
        self.assertEqual(
            [node.native_locator.object_index for node in image_nodes],
            [0, 1],
        )
        self.assertEqual(image_nodes[0].asset_ref, image_nodes[1].asset_ref)
        self.assertNotEqual(image_nodes[0].source_node_id, image_nodes[1].source_node_id)
        self.assertEqual(len(media_reads), 1)
        self.assertEqual(len(media_hashes), 1)

    def test_source_node_budget_fails_closed_with_dedicated_diagnostic(self) -> None:
        path = self.root / "too-many-nodes.docx"
        document = Document()
        document.add_paragraph("节点一")
        document.add_paragraph("节点二")
        document.save(path)

        with patch(
            "app.adapters.documents.docx_source_ir_adapter._MAX_DOCX_SOURCE_NODES",
            1,
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight_document(path),
                created_at=CREATED_AT,
            )

        self._assert_resource_budget_failure(snapshot)

    def test_physical_table_cell_budget_includes_nested_cells_and_fails_closed(self) -> None:
        path = self.root / "too-many-table-cells.docx"
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "外层单元格"
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套单元格"
        document.save(path)

        with patch(
            "app.adapters.documents.docx_source_ir_adapter._MAX_DOCX_PHYSICAL_TABLE_CELLS",
            1,
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight_document(path),
                created_at=CREATED_AT,
            )

        self._assert_resource_budget_failure(snapshot)

    def test_image_placement_budget_fails_closed(self) -> None:
        path = self._write_repeated_image_docx("too-many-images.docx", placement_count=2)

        with patch(
            "app.adapters.documents.docx_source_ir_adapter._MAX_DOCX_IMAGE_PLACEMENTS",
            1,
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight_document(path),
                created_at=CREATED_AT,
            )

        self._assert_resource_budget_failure(snapshot)

    def test_unique_media_byte_budget_fails_before_media_decompression(self) -> None:
        path = self._write_repeated_image_docx("media-too-large.docx", placement_count=1)
        preflight = preflight_document(path)
        original_read = zipfile.ZipFile.read
        media_reads: list[str] = []

        def counting_read(
            archive: zipfile.ZipFile,
            member: str | zipfile.ZipInfo,
            *args,
            **kwargs,
        ) -> bytes:
            member_name = member.filename if isinstance(member, zipfile.ZipInfo) else member
            if member_name.startswith("word/media/"):
                media_reads.append(member_name)
            return original_read(archive, member, *args, **kwargs)

        with (
            patch(
                "app.adapters.documents.docx_source_ir_adapter._MAX_DOCX_UNIQUE_MEDIA_BYTES",
                0,
            ),
            patch.object(zipfile.ZipFile, "read", new=counting_read),
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight,
                created_at=CREATED_AT,
            )

        self._assert_resource_budget_failure(snapshot)
        self.assertEqual(media_reads, [])

    def test_svg_and_fallback_in_one_drawing_are_one_placement(self) -> None:
        path = self._write_svg_with_png_fallback("svg-fallback.docx")

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        image_nodes = [node for node in snapshot.nodes if node.kind is SourceNodeKind.IMAGE]
        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(len(image_nodes), 1)
        self.assertEqual(image_nodes[0].native_locator.object_index, 0)
        self.assertEqual(len(snapshot.assets), 1)
        self.assertEqual(snapshot.assets[0].media_type, "image/svg+xml")
        self.assertEqual(image_nodes[0].asset_ref, snapshot.assets[0].asset_id)
        self.assertIn(
            "docx_image_alternate_renditions_collapsed",
            {item.code for item in snapshot.diagnostics},
        )

    def test_extract_parses_the_same_bytes_used_for_identity_validation(self) -> None:
        path = self._write_rich_docx("identity-bound.docx")
        preflight = preflight_document(path)
        original_extract_package = self.adapter._extract_package
        consumed_payloads: list[bytes] = []

        def mutate_path_then_parse(source_bytes: bytes, source_sha256: str):
            consumed_payloads.append(source_bytes)
            path.write_bytes(b"replaced after identity validation")
            return original_extract_package(source_bytes, source_sha256)

        with patch.object(
            self.adapter,
            "_extract_package",
            side_effect=mutate_path_then_parse,
        ):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight,
                created_at=CREATED_AT,
            )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(len(consumed_payloads), 1)
        self.assertEqual(
            hashlib.sha256(consumed_payloads[0]).hexdigest(),
            preflight.identity.source_sha256,
        )
        self.assertEqual(path.read_bytes(), b"replaced after identity validation")

    def test_content_control_is_flattened_without_losing_visible_paragraph(self) -> None:
        path = self.root / "content-control.docx"
        document = Document()
        paragraph = document.add_paragraph("内容控件中的真实正文")
        body = document._body._element
        content_control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        content.append(paragraph._p)
        content_control.append(content)
        body.insert(0, content_control)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), "0")
        body.insert(1, bookmark_end)
        document.save(path)

        snapshot = self.adapter.extract(
            path,
            preflight=preflight_document(path),
            created_at=CREATED_AT,
        )

        self.assertEqual(snapshot.status, SourceIRStatus.COMPLETE)
        self.assertEqual(
            [node.raw_text for node in snapshot.nodes if node.kind is SourceNodeKind.PARAGRAPH],
            ["内容控件中的真实正文"],
        )
        self.assertIn(
            "docx_nonvisual_markers_ignored",
            {item.code for item in snapshot.diagnostics},
        )
        self.assertNotIn(
            "docx_body_element_not_represented",
            {item.code for item in snapshot.diagnostics},
        )

    def test_invalid_preflight_source_change_and_naive_timestamp_are_rejected(self) -> None:
        source = self._write_rich_docx("source.docx")
        disguised = self.root / "disguised.pdf"
        shutil.copyfile(source, disguised)
        invalid_preflight = preflight_document(disguised)

        with self.assertRaises(SourceIRAdapterRejected) as invalid_context:
            self.adapter.extract(
                disguised,
                preflight=invalid_preflight,
                created_at=CREATED_AT,
            )
        self.assertEqual(
            invalid_context.exception.code,
            SourceIRAdapterRejectionCode.INVALID_PREFLIGHT,
        )

        valid_preflight = preflight_document(source)
        with source.open("ab") as handle:
            handle.write(b"changed-after-preflight")
        with self.assertRaises(SourceIRAdapterRejected) as changed_context:
            self.adapter.extract(source, preflight=valid_preflight, created_at=CREATED_AT)
        self.assertEqual(
            changed_context.exception.code,
            SourceIRAdapterRejectionCode.SOURCE_CHANGED,
        )

        timestamp_source = self._write_rich_docx("timestamp.docx")
        with self.assertRaises(SourceIRAdapterRejected) as timestamp_context:
            self.adapter.extract(
                timestamp_source,
                preflight=preflight_document(timestamp_source),
                created_at=datetime(2026, 7, 20, 12, 0),
            )
        self.assertEqual(
            timestamp_context.exception.code,
            SourceIRAdapterRejectionCode.INVALID_TIMESTAMP,
        )

    def test_native_parser_failure_after_valid_preflight_returns_failed_snapshot(self) -> None:
        path = self._write_rich_docx("failure.docx")
        with patch.object(self.adapter, "_extract_package", side_effect=ValueError("boom")):
            snapshot = self.adapter.extract(
                path,
                preflight=preflight_document(path),
                created_at=CREATED_AT,
            )

        self.assertEqual(snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(snapshot.nodes, ())
        self.assertEqual(snapshot.assets, ())
        self.assertEqual(
            {item.code for item in snapshot.diagnostics},
            {"docx_parse_failed"},
        )

    def test_adapter_satisfies_source_ir_port(self) -> None:
        self.assertIsInstance(self.adapter, SourceIRAdapterPort)
        self.assertEqual(self.adapter.document_format, DocumentFormat.DOCX)

    def _assert_resource_budget_failure(self, snapshot) -> None:
        self.assertEqual(snapshot.status, SourceIRStatus.FAILED)
        self.assertEqual(snapshot.nodes, ())
        self.assertEqual(snapshot.assets, ())
        self.assertEqual(
            {item.code for item in snapshot.diagnostics},
            {"docx_resource_budget_exceeded"},
        )

    def _write_repeated_image_docx(self, name: str, *, placement_count: int) -> Path:
        image_path = self.root / f"{name}.png"
        image_path.write_bytes(_ONE_PIXEL_PNG)
        path = self.root / name
        document = Document()
        paragraph = document.add_paragraph("重复图片引用")
        for _ in range(placement_count):
            paragraph.add_run().add_picture(str(image_path))
        document.save(path)
        image_path.unlink()
        return path

    def _write_dht120_cell_shape(self, name: str) -> Path:
        image_path = self.root / f"{name}.png"
        image_path.write_bytes(_ONE_PIXEL_PNG)
        path = self.root / name

        document = Document()
        table = document.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = "2.1.22 高低温特性"
        cell.paragraphs[0].style = "Heading 3"

        sample_count = cell.add_paragraph("样本数量：6件/OTS")
        self._set_direct_numbering(sample_count, num_id=42, ilvl=2)
        precondition = cell.add_paragraph("前置条件：试验前静置2小时")
        self._set_direct_numbering(precondition, num_id=42)
        precondition.add_run().add_picture(str(image_path))
        cell.add_paragraph("验收判定：性能参数应满足限值")

        document.save(path)
        image_path.unlink()
        return path

    @staticmethod
    def _set_direct_numbering(
        paragraph,
        *,
        num_id: int | None = None,
        ilvl: int | None = None,
    ) -> None:
        properties = paragraph._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        if ilvl is not None:
            level = OxmlElement("w:ilvl")
            level.set(qn("w:val"), str(ilvl))
            numbering.append(level)
        if num_id is not None:
            number_id = OxmlElement("w:numId")
            number_id.set(qn("w:val"), str(num_id))
            numbering.append(number_id)
        properties.append(numbering)

    @staticmethod
    def _alternate_text(
        choice_text: str,
        fallback_text: str | None,
        *,
        required_prefix: str = "wps",
        required_namespace: str = _WPS_NS,
    ) -> etree._Element:
        alternate = etree.Element(
            f"{{{_MC_NS}}}AlternateContent",
            nsmap={required_prefix: required_namespace},
        )
        choice = etree.SubElement(alternate, f"{{{_MC_NS}}}Choice")
        choice.set("Requires", required_prefix)
        branches = [(choice, choice_text)]
        if fallback_text is not None:
            fallback = etree.SubElement(alternate, f"{{{_MC_NS}}}Fallback")
            branches.append((fallback, fallback_text))
        for branch, value in branches:
            run = etree.SubElement(branch, qn("w:r"))
            text = etree.SubElement(run, qn("w:t"))
            text.text = value
        return alternate

    def _write_rich_docx(self, name: str) -> Path:
        image_path = self.root / f"{name}.png"
        image_path.write_bytes(_ONE_PIXEL_PNG)
        path = self.root / name

        document = Document()
        first = document.add_paragraph()
        first_run = first.add_run("父需求原文")
        document.add_comment(first_run, text="工程师批注", author="Engineer")

        list_item = document.add_paragraph("同级编号 1")
        properties = list_item._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number_id = OxmlElement("w:numId")
        number_id.set(qn("w:val"), "1")
        numbering.extend((level, number_id))
        properties.append(numbering)

        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "共享验收条件"
        table.cell(0, 2).text = "判定"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "共同材料"
        table.cell(1, 1).text = "参数 A"
        table.cell(1, 2).text = "阈值 A"
        table.cell(2, 1).text = "参数 B"
        table.cell(2, 2).text = "阈值 B"

        document.add_paragraph("表后正文")
        document.add_picture(str(image_path))
        document.save(path)
        image_path.unlink()
        return path

    def _write_svg_with_png_fallback(self, name: str) -> Path:
        image_path = self.root / f"{name}.png"
        image_path.write_bytes(_ONE_PIXEL_PNG)
        path = self.root / name
        document = Document()
        document.add_picture(str(image_path))
        document.save(path)
        image_path.unlink()

        with zipfile.ZipFile(path) as archive:
            members = {member: archive.read(member) for member in archive.namelist()}

        document_root = etree.fromstring(members["word/document.xml"])
        fallback_blip = document_root.find(f".//{{{_A_NS}}}blip")
        assert fallback_blip is not None
        extension_list = etree.SubElement(fallback_blip, f"{{{_A_NS}}}extLst")
        extension = etree.SubElement(extension_list, f"{{{_A_NS}}}ext")
        extension.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        svg_blip = etree.SubElement(
            extension,
            f"{{{_ASVG_NS}}}svgBlip",
            nsmap={"asvg": _ASVG_NS},
        )
        svg_blip.set(f"{{{_R_NS}}}embed", "rIdVectorRendition")
        members["word/document.xml"] = etree.tostring(
            document_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        relationships_root = etree.fromstring(members["word/_rels/document.xml.rels"])
        relationship = etree.SubElement(
            relationships_root,
            f"{{{_PKG_REL_NS}}}Relationship",
        )
        relationship.set("Id", "rIdVectorRendition")
        relationship.set(
            "Type",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        )
        relationship.set("Target", "media/vector.svg")
        members["word/_rels/document.xml.rels"] = etree.tostring(
            relationships_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        content_types_root = etree.fromstring(members["[Content_Types].xml"])
        svg_default = etree.SubElement(content_types_root, f"{{{_CONTENT_TYPE_NS}}}Default")
        svg_default.set("Extension", "svg")
        svg_default.set("ContentType", "image/svg+xml")
        members["[Content_Types].xml"] = etree.tostring(
            content_types_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )
        members["word/media/vector.svg"] = (
            b'<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1">'
            b'<rect width="1" height="1"/></svg>'
        )

        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for member, payload in members.items():
                archive.writestr(member, payload)
        return path


if __name__ == "__main__":
    unittest.main()
