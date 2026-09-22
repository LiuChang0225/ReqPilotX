from __future__ import annotations

import codecs
import hashlib
import logging
import warnings
import zipfile
from dataclasses import dataclass
from os import PathLike
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image as PILImage
from pypdf import PasswordType, PdfReader
from pypdf.errors import DependencyError

from reqpilotx.domain.extraction.preflight import (
    AccessStatus,
    AuthenticityStatus,
    DiagnosticCode,
    DiagnosticSeverity,
    DocumentFormat,
    DocumentPreflightReport,
    IntegrityStatus,
    PreflightDiagnostic,
    RawDocumentIdentity,
    SupportStatus,
)
from reqpilotx.domain.extraction.text import declared_plain_text_encoding


_PDF_HEADER = b"%PDF-"
_OLE_COMPOUND_HEADER = bytes.fromhex("D0CF11E0A1B11AE1")
_TIFF_HEADERS = (b"II*\x00", b"MM\x00*")
_ZIP_HEADERS = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
_HEAD_BYTES = 65_536
_HASH_CHUNK_BYTES = 1024 * 1024
_MAX_ZIP_MEMBERS = 20_000
_MAX_ZIP_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
_MAX_ZIP_MEMBER_BYTES = 128 * 1024 * 1024
_MAX_ZIP_COMPRESSION_RATIO = 500.0
_MAX_PDF_PAGES = 5_000
_MAX_TIF_FRAMES = 5_000
_MAX_TIF_TOTAL_PIXELS = 500_000_000

_EXPECTED_FORMAT_BY_EXTENSION = {
    ".pdf": DocumentFormat.PDF,
    ".docx": DocumentFormat.DOCX,
    ".xlsx": DocumentFormat.XLSX,
    ".xlsb": DocumentFormat.XLSB,
    ".doc": DocumentFormat.LEGACY_DOC,
    ".tif": DocumentFormat.TIF,
    ".tiff": DocumentFormat.TIF,
    ".txt": DocumentFormat.TXT,
}
_SUPPORTED_FORMATS = frozenset(
    {DocumentFormat.PDF, DocumentFormat.DOCX, DocumentFormat.XLSX, DocumentFormat.TXT}
)
_MIME_TYPES = {
    DocumentFormat.PDF: "application/pdf",
    DocumentFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    DocumentFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    DocumentFormat.XLSB: "application/vnd.ms-excel.sheet.binary.macroenabled.12",
    DocumentFormat.LEGACY_DOC: "application/msword",
    DocumentFormat.TIF: "image/tiff",
    DocumentFormat.TXT: "text/plain",
    DocumentFormat.UNKNOWN: "application/octet-stream",
}


@dataclass(frozen=True, slots=True)
class _SignatureProbe:
    detected_format: DocumentFormat
    corrupt_container: bool = False
    resource_limit_reason: str | None = None


class LocalDocumentPreflight:
    """Filesystem adapter that performs read-only, fail-closed document checks."""

    def inspect(self, source_path: str | PathLike[str]) -> DocumentPreflightReport:
        path = Path(source_path)
        if not path.is_file():
            raise FileNotFoundError(f"Document does not exist: {path}")

        source_size, source_sha256, head = _read_file_facts(path)
        declared_extension = path.suffix.lower()
        probe = _detect_format(path, head)
        identity = RawDocumentIdentity(
            source_name=path.name,
            source_sha256=source_sha256,
            source_size=source_size,
            declared_extension=declared_extension,
            detected_format=probe.detected_format,
            mime_type=_MIME_TYPES[probe.detected_format],
        )
        authenticity_status = _authenticity_status(
            declared_extension, probe.detected_format
        )
        support_status = (
            SupportStatus.SUPPORTED
            if probe.detected_format in _SUPPORTED_FORMATS
            else SupportStatus.UNSUPPORTED
        )
        diagnostics: list[PreflightDiagnostic] = []

        if source_size == 0:
            diagnostics.append(
                _diagnostic(DiagnosticCode.EMPTY_FILE, "The source document is empty.")
            )
            return DocumentPreflightReport(
                identity=identity,
                authenticity_status=authenticity_status,
                support_status=support_status,
                integrity_status=IntegrityStatus.EMPTY,
                access_status=AccessStatus.NOT_INSPECTED,
                diagnostics=tuple(diagnostics),
            )

        if authenticity_status is AuthenticityStatus.MISMATCHED:
            diagnostics.append(
                _diagnostic(
                    DiagnosticCode.EXTENSION_SIGNATURE_MISMATCH,
                    (
                        f"Declared extension {declared_extension or '(none)'} does not match "
                        f"detected format {probe.detected_format.value}."
                    ),
                )
            )

        if probe.corrupt_container:
            diagnostics.append(
                _diagnostic(
                    DiagnosticCode.CORRUPT_DOCUMENT,
                    "The ZIP-based Office container is damaged and cannot be inspected.",
                )
            )
            return DocumentPreflightReport(
                identity=identity,
                authenticity_status=authenticity_status,
                support_status=support_status,
                integrity_status=IntegrityStatus.CORRUPT,
                access_status=AccessStatus.NOT_INSPECTED,
                diagnostics=tuple(diagnostics),
            )

        if probe.resource_limit_reason is not None:
            diagnostics.append(
                _diagnostic(
                    DiagnosticCode.RESOURCE_LIMIT_EXCEEDED,
                    (
                        "The document container exceeds a safe inspection budget "
                        f"({probe.resource_limit_reason})."
                    ),
                )
            )
            return DocumentPreflightReport(
                identity=identity,
                authenticity_status=authenticity_status,
                support_status=support_status,
                integrity_status=IntegrityStatus.UNVERIFIED,
                access_status=AccessStatus.NOT_INSPECTED,
                diagnostics=tuple(diagnostics),
            )
        if probe.detected_format is DocumentFormat.UNKNOWN:
            diagnostics.append(
                _diagnostic(
                    DiagnosticCode.UNKNOWN_BINARY_FORMAT,
                    "The source bytes do not match a supported document or strict text signature.",
                )
            )
            return DocumentPreflightReport(
                identity=identity,
                authenticity_status=authenticity_status,
                support_status=SupportStatus.UNSUPPORTED,
                integrity_status=IntegrityStatus.UNVERIFIED,
                access_status=AccessStatus.NOT_INSPECTED,
                diagnostics=tuple(diagnostics),
            )

        if support_status is SupportStatus.UNSUPPORTED:
            diagnostics.append(
                _diagnostic(
                    DiagnosticCode.UNSUPPORTED_FORMAT,
                    _unsupported_format_message(probe.detected_format),
                )
            )
            if probe.detected_format is DocumentFormat.TIF:
                (
                    integrity_status,
                    access_status,
                    page_count,
                    inspection_diagnostics,
                ) = _inspect_tif(path)
                diagnostics.extend(inspection_diagnostics)
                return DocumentPreflightReport(
                    identity=identity,
                    authenticity_status=authenticity_status,
                    support_status=support_status,
                    integrity_status=integrity_status,
                    access_status=access_status,
                    diagnostics=tuple(diagnostics),
                    page_count=page_count,
                )
            return DocumentPreflightReport(
                identity=identity,
                authenticity_status=authenticity_status,
                support_status=support_status,
                integrity_status=IntegrityStatus.VALID,
                access_status=AccessStatus.NOT_INSPECTED,
                diagnostics=tuple(diagnostics),
            )

        integrity_status, access_status, page_count, inspection_diagnostics = (
            _inspect_supported_format(path, probe.detected_format)
        )
        diagnostics.extend(inspection_diagnostics)
        return DocumentPreflightReport(
            identity=identity,
            authenticity_status=authenticity_status,
            support_status=support_status,
            integrity_status=integrity_status,
            access_status=access_status,
            diagnostics=tuple(diagnostics),
            page_count=page_count,
        )


def preflight_document(source_path: str | PathLike[str]) -> DocumentPreflightReport:
    return LocalDocumentPreflight().inspect(source_path)


def _read_file_facts(path: Path) -> tuple[int, str, bytes]:
    digest = hashlib.sha256()
    source_size = 0
    head = b""
    with path.open("rb") as handle:
        while chunk := handle.read(_HASH_CHUNK_BYTES):
            if len(head) < _HEAD_BYTES:
                head += chunk[: _HEAD_BYTES - len(head)]
            digest.update(chunk)
            source_size += len(chunk)
    return source_size, digest.hexdigest(), head


def _detect_format(path: Path, head: bytes) -> _SignatureProbe:
    pdf_offset = head.find(_PDF_HEADER, 0, 1024)
    if pdf_offset >= 0:
        return _SignatureProbe(DocumentFormat.PDF)
    if head.startswith(_OLE_COMPOUND_HEADER):
        return _SignatureProbe(DocumentFormat.LEGACY_DOC)
    if head.startswith(_TIFF_HEADERS):
        return _SignatureProbe(DocumentFormat.TIF)
    if head.startswith(_ZIP_HEADERS):
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                names = frozenset(item.filename for item in members)
        except (OSError, EOFError, zipfile.BadZipFile, zipfile.LargeZipFile):
            return _SignatureProbe(DocumentFormat.UNKNOWN, corrupt_container=True)
        detected_format = DocumentFormat.UNKNOWN
        if "xl/workbook.bin" in names:
            detected_format = DocumentFormat.XLSB
        elif "xl/workbook.xml" in names:
            detected_format = DocumentFormat.XLSX
        elif "word/document.xml" in names:
            detected_format = DocumentFormat.DOCX
        return _SignatureProbe(
            detected_format,
            resource_limit_reason=_zip_resource_limit_reason(members),
        )
    if _looks_like_text(head):
        return _SignatureProbe(DocumentFormat.TXT)
    return _SignatureProbe(DocumentFormat.UNKNOWN)


def _looks_like_text(value: bytes) -> bool:
    if not value:
        return False
    encoding = declared_plain_text_encoding(value)
    try:
        # The bounded signature window may end midway through a Unicode code
        # point.  Keep that trailing partial sequence for the full-file check.
        decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
        text = decoder.decode(value, final=False)
    except (LookupError, UnicodeDecodeError):
        return False
    return _text_is_printable(text)


def _text_is_printable(text: str) -> bool:
    if not text or "\x00" in text:
        return False
    printable = sum(
        character.isprintable() or character in "\r\n\t" for character in text
    )
    return printable / len(text) >= 0.92


def _zip_resource_limit_reason(members: list[zipfile.ZipInfo]) -> str | None:
    if len(members) > _MAX_ZIP_MEMBERS:
        return "zip_member_count"
    total_uncompressed = 0
    for member in members:
        total_uncompressed += max(int(member.file_size), 0)
        if member.file_size > _MAX_ZIP_MEMBER_BYTES:
            return "zip_member_uncompressed_size"
        if member.file_size > 0:
            if member.compress_size <= 0:
                return "zip_compression_ratio"
            if member.file_size / member.compress_size > _MAX_ZIP_COMPRESSION_RATIO:
                return "zip_compression_ratio"
        if total_uncompressed > _MAX_ZIP_UNCOMPRESSED_BYTES:
            return "zip_total_uncompressed_size"
    return None


def _authenticity_status(
    extension: str, detected_format: DocumentFormat
) -> AuthenticityStatus:
    expected_format = _EXPECTED_FORMAT_BY_EXTENSION.get(extension)
    if expected_format is not None:
        return (
            AuthenticityStatus.MATCHED
            if expected_format is detected_format
            else AuthenticityStatus.MISMATCHED
        )
    if detected_format is DocumentFormat.UNKNOWN:
        return AuthenticityStatus.UNVERIFIABLE
    return AuthenticityStatus.MISMATCHED


def _inspect_supported_format(
    path: Path, detected_format: DocumentFormat
) -> tuple[IntegrityStatus, AccessStatus, int | None, list[PreflightDiagnostic]]:
    if detected_format is DocumentFormat.PDF:
        return _inspect_pdf(path)
    if detected_format is DocumentFormat.DOCX:
        try:
            Document(path)
        except Exception as exc:
            return _corrupt_result("DOCX", exc)
        return IntegrityStatus.VALID, AccessStatus.READABLE, None, []
    if detected_format is DocumentFormat.XLSX:
        workbook = None
        try:
            workbook = load_workbook(
                path, read_only=True, data_only=False, keep_links=False
            )
            if not workbook.sheetnames:
                raise ValueError("The workbook contains no worksheets.")
        except Exception as exc:
            return _corrupt_result("XLSX", exc)
        finally:
            if workbook is not None:
                workbook.close()
        return IntegrityStatus.VALID, AccessStatus.READABLE, None, []
    if detected_format is DocumentFormat.TXT:
        try:
            _validate_text_file(path)
        except (LookupError, UnicodeDecodeError, ValueError) as exc:
            return _corrupt_result("TXT", exc)
        return IntegrityStatus.VALID, AccessStatus.READABLE, None, []
    raise AssertionError(f"No supported-format inspector for {detected_format.value}")


def _inspect_pdf(
    path: Path,
) -> tuple[IntegrityStatus, AccessStatus, int | None, list[PreflightDiagnostic]]:
    pypdf_logger = logging.getLogger("pypdf")
    previous_level = pypdf_logger.level
    try:
        # Some readable customer PDFs contain duplicate resource dictionary
        # keys.  pypdf can still resolve them, but emits thousands of warnings;
        # preflight reports the document state, not dependency console noise.
        pypdf_logger.setLevel(logging.ERROR)
        reader = PdfReader(str(path), strict=False)
        access_status = AccessStatus.READABLE
        diagnostics: list[PreflightDiagnostic] = []
        if reader.is_encrypted:
            try:
                password_type = reader.decrypt("")
            except DependencyError as exc:
                return (
                    IntegrityStatus.VALID,
                    AccessStatus.UNSUPPORTED_ENCRYPTION,
                    None,
                    [
                        _diagnostic(
                            DiagnosticCode.UNSUPPORTED_PDF_ENCRYPTION,
                            f"The PDF encryption algorithm cannot be opened by this runtime: {exc}",
                        )
                    ],
                )
            if password_type is PasswordType.NOT_DECRYPTED:
                return (
                    IntegrityStatus.VALID,
                    AccessStatus.PASSWORD_REQUIRED,
                    None,
                    [
                        _diagnostic(
                            DiagnosticCode.PASSWORD_REQUIRED,
                            "The PDF requires a non-empty password before it can be parsed.",
                        )
                    ],
                )
            access_status = AccessStatus.READABLE_WITH_EMPTY_PASSWORD
            diagnostics.append(
                _diagnostic(
                    DiagnosticCode.PDF_SECURITY_RESTRICTIONS,
                    (
                        "The PDF has a security dictionary but is readable with an empty password; "
                        "processing remains allowed."
                    ),
                    severity=DiagnosticSeverity.INFO,
                )
            )
        page_count = len(reader.pages)
        if page_count > _MAX_PDF_PAGES:
            return (
                IntegrityStatus.UNVERIFIED,
                AccessStatus.NOT_INSPECTED,
                None,
                [
                    _diagnostic(
                        DiagnosticCode.RESOURCE_LIMIT_EXCEEDED,
                        "The PDF page count exceeds the safe inspection budget.",
                    )
                ],
            )
        # Resolving every page catches broken page-tree objects without extracting content.
        for page in reader.pages:
            _ = page.mediabox
        return IntegrityStatus.VALID, access_status, page_count, diagnostics
    except Exception as exc:
        return _corrupt_result("PDF", exc)
    finally:
        pypdf_logger.setLevel(previous_level)


def _inspect_tif(
    path: Path,
) -> tuple[IntegrityStatus, AccessStatus, int | None, list[PreflightDiagnostic]]:
    try:
        with warnings.catch_warnings():
            # Pillow's generic warning threshold is lower than the explicit
            # attachment inspection budget below. We suppress only that warning
            # and still reject over-budget dimensions before accepting the file.
            warnings.simplefilter("ignore", PILImage.DecompressionBombWarning)
            with PILImage.open(path) as image:
                if image.format != "TIFF":
                    raise ValueError("The image container is not TIFF.")
                frame_count = int(getattr(image, "n_frames", 1))
                if frame_count < 1 or frame_count > _MAX_TIF_FRAMES:
                    raise ValueError(
                        "The TIFF frame count exceeds the inspection budget."
                    )
                total_pixels = 0
                for frame_index in range(frame_count):
                    image.seek(frame_index)
                    width, height = image.size
                    if width <= 0 or height <= 0:
                        raise ValueError("The TIFF contains an empty frame.")
                    total_pixels += int(width) * int(height)
                    if total_pixels > _MAX_TIF_TOTAL_PIXELS:
                        raise ValueError(
                            "The TIFF pixel count exceeds the inspection budget."
                        )
                    image.load()
    except Exception as exc:
        return _corrupt_result("TIF", exc)
    return IntegrityStatus.VALID, AccessStatus.READABLE, frame_count, []


def _validate_text_file(path: Path) -> None:
    with path.open("rb") as handle:
        prefix = handle.read(4)
        handle.seek(0)
        encoding = declared_plain_text_encoding(prefix)
        decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
        printable = 0
        total = 0
        while chunk := handle.read(_HASH_CHUNK_BYTES):
            text = decoder.decode(chunk, final=False)
            if "\x00" in text:
                raise ValueError(
                    "NUL bytes are not valid in the TXT compatibility path."
                )
            printable += sum(
                character.isprintable() or character in "\r\n\t" for character in text
            )
            total += len(text)
        tail = decoder.decode(b"", final=True)
        if "\x00" in tail:
            raise ValueError("NUL bytes are not valid in the TXT compatibility path.")
        printable += sum(
            character.isprintable() or character in "\r\n\t" for character in tail
        )
        total += len(tail)
    if total == 0 or printable / total < 0.92:
        raise ValueError("The file contains binary or non-printable content.")


def _corrupt_result(
    format_name: str, exc: Exception
) -> tuple[IntegrityStatus, AccessStatus, None, list[PreflightDiagnostic]]:
    return (
        IntegrityStatus.CORRUPT,
        AccessStatus.NOT_INSPECTED,
        None,
        [
            _diagnostic(
                DiagnosticCode.CORRUPT_DOCUMENT,
                f"The {format_name} document is damaged or incomplete: {type(exc).__name__}: {exc}",
            )
        ],
    )


def _unsupported_format_message(detected_format: DocumentFormat) -> str:
    if detected_format is DocumentFormat.XLSB:
        return "XLSB is not supported in the first extraction release; convert it to XLSX without changing content."
    if detected_format is DocumentFormat.LEGACY_DOC:
        return "Legacy OLE DOC is not supported in the first extraction release; convert it to DOCX without changing content."
    if detected_format is DocumentFormat.TIF:
        return (
            "TIF can be retained only through the opaque attachment lane; "
            "image requirement extraction is not enabled."
        )
    return f"Format {detected_format.value} is not supported."


def _diagnostic(
    code: DiagnosticCode,
    message: str,
    *,
    severity: DiagnosticSeverity = DiagnosticSeverity.ERROR,
) -> PreflightDiagnostic:
    return PreflightDiagnostic(code=code, severity=severity, message=message)
