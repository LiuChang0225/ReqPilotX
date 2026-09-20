from __future__ import annotations

# Test runtime isolation must be activated before importing application modules.
# ruff: noqa: E402

if __package__:
    from .runtime_isolation import activate_test_runtime
else:
    from runtime_isolation import activate_test_runtime

activate_test_runtime()

if __package__:
    from .private_eval import require_private_eval
else:
    from private_eval import require_private_eval

import copy
import hashlib
import json
import multiprocessing
import shutil
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook
from docx import Document as WordDocument
from PIL import Image as PILImage
import pymupdf
from pypdf import PdfWriter

from app.domain.extraction.structure import (
    stable_structure_node_id,
    stable_structure_relation_id,
)
from scripts import extraction_case_runner as runner_module
from scripts.extraction_case_runner import (
    CorpusPathError,
    ManifestValidationError,
    load_manifest,
    load_oracle_registry,
    require_report_destination_outside_corpus,
    resolve_document_path,
    run_manifest,
    validate_manifest,
    write_reports,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
MANIFEST_PATH = PROJECT_DIR / "eval" / "extraction_cases" / "manifest.json"
ORACLE_PATH = PROJECT_DIR / "eval" / "extraction_cases" / "oracles.json"


class ExtractionCaseRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory(
            ignore_cleanup_errors=True
        )
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)
        self.oracles = load_oracle_registry(require_private_eval("extraction_cases/oracles.json"))
        self.manifest = load_manifest(require_private_eval("extraction_cases/manifest.json"))

    def _oracles_for_manifest(self, manifest: dict) -> dict:
        registry = copy.deepcopy(self.oracles)
        for oracle in registry["oracles"].values():
            if "document_ids" in oracle:
                oracle["source_sha256_by_document"] = {
                    document_id: manifest["documents"][document_id]["source_sha256"]
                    for document_id in oracle["document_ids"]
                }
                continue
            oracle["source_sha256"] = manifest["documents"][oracle["document_id"]][
                "source_sha256"
            ]
            review_provenance = oracle.get("review_provenance")
            if (
                isinstance(review_provenance, dict)
                and "reviewed_source_sha256" in review_provenance
            ):
                review_provenance["reviewed_source_sha256"] = oracle[
                    "source_sha256"
                ]
        body_profiles = manifest["cases"]["FMT-005"]["assertions"][3]["evaluation"][
            "reviewed_body_profiles_by_document"
        ]
        for document_id, profile in body_profiles.items():
            profile["source_sha256"] = manifest["documents"][document_id][
                "source_sha256"
            ]
        return registry

    def _run_manifest(self, manifest: dict, **kwargs: object) -> dict:
        oracle_registry = kwargs.pop("oracle_registry", None)
        return run_manifest(
            manifest,
            oracle_registry=(
                oracle_registry
                if isinstance(oracle_registry, dict)
                else self._oracles_for_manifest(manifest)
            ),
            **kwargs,
        )

    def _configure_source_scope_oracle(
        self,
        manifest: dict,
        *,
        oracle_id: str,
        path: Path,
        selector: dict[str, object],
    ) -> dict:
        registry = self._oracles_for_manifest(manifest)
        oracle = registry["oracles"][oracle_id]
        expected_scope = copy.deepcopy(oracle["fact"]["scopes"][0])
        expected_scope["selector"] = selector
        snapshot = runner_module.build_source_ir(path)
        actual_scope, _ = runner_module._measure_source_scope_snapshot(
            snapshot,
            expected_scope,
        )
        oracle["fact"] = {
            "kind": "source_scope_group",
            "expected_document_count": 1,
            "scopes": [actual_scope],
            "group_fingerprint": runner_module._fingerprint_source_scope_group(
                [actual_scope]
            ),
        }
        return registry

    def _configure_str_oracle(
        self,
        manifest: dict,
        values: list[str],
    ) -> dict:
        registry = self._oracles_for_manifest(manifest)
        manifest["cases"]["STR-001"]["assertions"][2]["evaluation"][
            "expected_cell_range"
        ] = "H1"
        manifest["cases"]["STR-001"]["assertions"][3]["evaluation"][
            "expected_cell_range"
        ] = "G1"
        cell_records = [
            {
                "cell_range": f"{chr(ord('A') + index)}1",
                "row_span": 1,
                "column_span": 1,
                "raw_sha256": hashlib.sha256(value.encode("utf-8")).hexdigest(),
            }
            for index, value in enumerate(values)
        ]
        registry["oracles"]["ORACLE-STR-001-A1"]["fact"] = {
            "kind": "xlsx_cell_range",
            "locator": {
                "format": "xlsx",
                "workbook": "xl/workbook.xml",
                "sheet": "Scope",
                "cell_range": "A1:H1",
            },
            "expected_cell_count": len(cell_records),
            "cells": cell_records,
            "range_fingerprint": runner_module._fingerprint_xlsx_cell_records(
                cell_records
            ),
        }
        return registry

    def _configure_tab1_oracle(
        self,
        manifest: dict,
        path: Path,
    ) -> dict:
        registry = self._oracles_for_manifest(manifest)
        snapshot = runner_module.build_source_ir(path)
        table_node = next(
            node for node in snapshot.nodes if node.table_payload is not None
        )
        table = table_node.table_payload
        self.assertIsNotNone(table)
        native = table_node.native_locator
        grid_payload = runner_module._canonical_json_payload(
            [
                {
                    "row": cell.row,
                    "column": cell.column,
                    "row_span": cell.row_span,
                    "column_span": cell.column_span,
                    "raw_text": cell.raw_text,
                }
                for cell in sorted(
                    table.cells,
                    key=lambda item: (item.row, item.column),
                )
            ]
        )
        record = {
            "locator": {
                "format": "docx",
                "body_element_index": native.body_element_index,
                "table_index": native.table_index,
            },
            "row_count": table.row_count,
            "column_count": table.column_count,
            "cell_count": len(table.cells),
            "data_row_count": max(0, table.row_count - 1),
            "grid_sha256": runner_module._sha256_text(grid_payload),
        }
        registry["oracles"]["ORACLE-TAB-001-A1"]["fact"] = {
            "kind": "docx_table_group",
            "expected_table_count": 1,
            "expected_data_row_count": record["data_row_count"],
            "expected_cell_count": record["cell_count"],
            "tables": [record],
            "group_fingerprint": runner_module._fingerprint_docx_table_records(
                [record]
            ),
        }
        mixed_oracle = registry["oracles"].get("ORACLE-SEM-003-A1")
        if isinstance(mixed_oracle, dict):
            mixed_fact = mixed_oracle["fact"]
            mixed_fact["expected_parameter_table_count"] = 1
            mixed_fact["parameter_table_group_fingerprint"] = (
                registry["oracles"]["ORACLE-TAB-001-A1"]["fact"][
                    "group_fingerprint"
                ]
            )
            mixed_fact["mixed_fingerprint"] = (
                runner_module._fingerprint_docx_mixed_content_group(
                    narratives=mixed_fact["narratives"],
                    parameter_table_oracle_id=str(
                        mixed_fact["parameter_table_oracle_id"]
                    ),
                    expected_parameter_table_count=1,
                    parameter_table_group_fingerprint=str(
                        mixed_fact["parameter_table_group_fingerprint"]
                    ),
                )
            )
        return registry

    def _configure_str5_oracle(
        self,
        manifest: dict,
        path: Path,
    ) -> dict:
        registry = self._oracles_for_manifest(manifest)
        snapshot = runner_module.build_source_ir(path)
        table_node = next(
            node for node in snapshot.nodes if node.table_payload is not None
        )
        table = table_node.table_payload
        self.assertIsNotNone(table)
        native = table_node.native_locator
        row_index = 0
        cells = [
            {
                "column": int(cell.column),
                "row_span": int(cell.row_span),
                "column_span": int(cell.column_span),
                "raw_sha256": hashlib.sha256(
                    cell.raw_text.encode("utf-8")
                ).hexdigest(),
            }
            for cell in sorted(
                (cell for cell in table.cells if cell.row == row_index),
                key=lambda cell: (
                    cell.column,
                    cell.row_span,
                    cell.column_span,
                ),
            )
        ]
        fact = {
            "kind": "docx_table_row_scope",
            "locator": {
                "format": "docx",
                "body_element_index": native.body_element_index,
                "table_index": native.table_index,
                "row_index": row_index,
            },
            "anchor_cell_count": len(cells),
            "logical_slot_count": sum(
                int(cell["column_span"]) for cell in cells
            ),
            "cells": cells,
            "row_fingerprint": runner_module._fingerprint_docx_table_row_cells(
                cells
            ),
        }
        fact["scope_fingerprint"] = (
            runner_module._fingerprint_docx_table_row_scope(fact)
        )
        registry["oracles"]["ORACLE-STR-005-A1"]["fact"] = fact
        return registry

    def _configure_tab2_oracle(
        self,
        manifest: dict,
        path: Path,
    ) -> dict:
        registry = self._oracles_for_manifest(manifest)
        snapshot = runner_module.build_source_ir(path)
        table_node = next(
            node for node in snapshot.nodes if node.table_payload is not None
        )
        table = table_node.table_payload
        self.assertIsNotNone(table)
        native = table_node.native_locator
        self.assertIsNotNone(native.cell_range)
        table_bounds = runner_module._parse_a1_range(native.cell_range)
        reviewed_merge_bounds = runner_module._parse_a1_range("D2:E4")
        cells_by_position = {(cell.row, cell.column): cell for cell in table.cells}
        cells_by_range = {
            cell.native_locator.cell_range: cell
            for cell in table.cells
            if cell.native_locator.cell_range is not None
        }
        merge_records = []
        for region in table.merged_regions:
            anchor = cells_by_position[(region.row, region.column)]
            absolute_row = table_bounds[0] + region.row
            absolute_column = table_bounds[1] + region.column
            absolute_bounds = (
                absolute_row,
                absolute_column,
                absolute_row + region.row_span - 1,
                absolute_column + region.column_span - 1,
            )
            if not runner_module._bounds_contain(
                reviewed_merge_bounds,
                absolute_bounds,
            ):
                continue
            merge_records.append(
                {
                    "anchor": (
                        f"{runner_module._column_name(absolute_column)}{absolute_row}"
                    ),
                    "row_span": region.row_span,
                    "column_span": region.column_span,
                    "anchor_raw_sha256": runner_module._sha256_text(anchor.raw_text),
                }
            )
        merge_records.sort(key=lambda item: item["anchor"])
        identity_records = []
        for cell in table.cells:
            cell_range = cell.native_locator.cell_range
            if cell_range in {"A2", "A3", "A4"}:
                identity_records.append(
                    {
                        "cell_range": cell_range,
                        "raw_sha256": runner_module._sha256_text(cell.raw_text),
                    }
                )
        identity_records.sort(key=lambda item: item["cell_range"])
        registry["oracles"]["ORACLE-TAB-002-A1"]["fact"] = {
            "kind": "xlsx_merge_group",
            "table_locator": {
                "format": "xlsx",
                "workbook": native.workbook,
                "sheet": native.sheet,
                "cell_range": native.cell_range,
            },
            "scope_range": "D2:E4",
            "covered_row_count": 3,
            "expected_merge_count": len(merge_records),
            "merges": merge_records,
            "merge_fingerprint": runner_module._fingerprint_xlsx_merge_records(
                merge_records
            ),
            "row_identity": {
                "cell_range": "A2:A4",
                "expected_cell_count": len(identity_records),
                "cells": identity_records,
                "fingerprint": runner_module._fingerprint_xlsx_identity_records(
                    identity_records
                ),
            },
            "group_fingerprint": runner_module._fingerprint_xlsx_merge_group(
                merge_records,
                identity_records,
            ),
        }
        source_fact = registry["oracles"]["ORACLE-TAB-002-A1"]["fact"]
        role_bindings = []
        for role, header_range, anchor_range, page_number in (
            ("STANDARD", "D1", "D2", 2),
            ("ACCEPTANCE_CRITERIA", "E1", "E2", 3),
        ):
            header = cells_by_range[header_range]
            anchor = cells_by_range[anchor_range]
            binding = {
                "role": role,
                "header_cell_range": header_range,
                "header_raw_sha256": runner_module._sha256_text(
                    header.raw_text
                ),
                "anchor_cell_range": anchor_range,
                "anchor_raw_sha256": runner_module._sha256_text(
                    anchor.raw_text
                ),
                "merge_range": f"{anchor_range}:{anchor_range[0]}4",
                "expected_applies_to_count": 3,
                "rendered_page_number": page_number,
            }
            binding["binding_fingerprint"] = (
                runner_module._fingerprint_xlsx_material_role_binding(
                    binding
                )
            )
            role_bindings.append(binding)
        material_fact = {
            "kind": "xlsx_material_requirement_scope",
            "source_oracle_id": "ORACLE-TAB-002-A1",
            "source_merge_group_fingerprint": source_fact[
                "group_fingerprint"
            ],
            "table_locator": dict(source_fact["table_locator"]),
            "full_row_scope": (
                f"A2:{runner_module._column_name(table_bounds[3])}4"
            ),
            "identity_rendered_page_number": 1,
            "expected_requirement_count": 3,
            "expected_role_applies_to_count": 6,
            "row_expectations": [
                {
                    "identity_cell_range": record["cell_range"],
                    "identity_raw_sha256": record["raw_sha256"],
                }
                for record in identity_records
            ],
            "role_bindings": role_bindings,
        }
        material_fact["scope_fingerprint"] = (
            runner_module._fingerprint_xlsx_material_requirement_scope(
                material_fact
            )
        )
        material_oracle = registry["oracles"]["ORACLE-TAB-002-A4"]
        material_oracle["fact"] = material_fact
        synthetic_source_sha256 = manifest["documents"][
            "DOC-IRPS-TRL-B"
        ]["source_sha256"]
        material_oracle["review_provenance"] = {
            "label": "AI_REVIEWED_SILVER",
            "reviewer_role": "AI_REVIEWER",
            "review_method": (
                "XLSX_LIBREOFFICE_RENDER_AND_OPENXML_CROSS_CHECK"
            ),
            "reviewed_page_numbers": [1, 2, 3],
            "renderer_id": "libreoffice",
            "renderer_version": "synthetic-test",
            "render_dpi": 180,
            "rendered_pdf_sha256": runner_module._sha256_text(
                "synthetic-material-render"
            ),
            "reviewed_source_sha256": synthetic_source_sha256,
            "reviewed_sheet_part": "xl/worksheets/sheet1.xml",
            "reviewed_sheet_part_sha256": runner_module._sha256_text(
                "synthetic-material-sheet"
            ),
            "page_image_sha256_by_page": {
                str(page_number): runner_module._sha256_text(
                    f"synthetic-material-page-{page_number}"
                )
                for page_number in (1, 2, 3)
            },
            "is_human_gold": False,
            "is_serving_activation_gate": False,
        }
        return registry

    def test_repository_manifest_has_exact_confirmed_cross_references(self) -> None:
        validate_manifest(self.manifest)

        assertion_ids = {
            assertion["assertion_id"]
            for case in self.manifest["cases"].values()
            for assertion in case["assertions"]
        }
        self.assertEqual(len(self.manifest["cases"]), 29)
        self.assertEqual(len(self.manifest["documents"]), 25)
        self.assertEqual(len(assertion_ids), 117)

        extra_document = copy.deepcopy(self.manifest)
        extra_document["documents"]["DOC-EXTRA-UNREFERENCED"] = copy.deepcopy(
            extra_document["documents"]["DOC-QC-T-1067"]
        )
        with self.assertRaises(ManifestValidationError):
            validate_manifest(extra_document)

    def test_tab2_material_oracle_integrity_fails_closed(self) -> None:
        def recompute_fact(fact: dict) -> None:
            for binding in fact["role_bindings"]:
                binding["binding_fingerprint"] = (
                    runner_module._fingerprint_xlsx_material_role_binding(
                        binding
                    )
                )
            fact["scope_fingerprint"] = (
                runner_module._fingerprint_xlsx_material_requirement_scope(
                    fact
                )
            )

        reversed_roles = copy.deepcopy(self.oracles)
        fact = reversed_roles["oracles"]["ORACLE-TAB-002-A4"]["fact"]
        fact["role_bindings"].reverse()
        recompute_fact(fact)
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                self.manifest,
                oracle_registry=reversed_roles,
            )

        truncated_merge = copy.deepcopy(self.oracles)
        fact = truncated_merge["oracles"]["ORACLE-TAB-002-A4"]["fact"]
        fact["role_bindings"][0]["merge_range"] = "D2:D3"
        fact["role_bindings"][0]["expected_applies_to_count"] = 2
        fact["expected_role_applies_to_count"] = 5
        recompute_fact(fact)
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                self.manifest,
                oracle_registry=truncated_merge,
            )

        duplicate_role_ref = copy.deepcopy(self.oracles)
        fact = duplicate_role_ref["oracles"]["ORACLE-TAB-002-A4"]["fact"]
        acceptance = fact["role_bindings"][1]
        standard = fact["role_bindings"][0]
        for key in (
            "header_cell_range",
            "header_raw_sha256",
            "anchor_cell_range",
            "anchor_raw_sha256",
            "merge_range",
            "rendered_page_number",
        ):
            standard[key] = acceptance[key]
        recompute_fact(fact)
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                self.manifest,
                oracle_registry=duplicate_role_ref,
            )

        wrong_component_ref = copy.deepcopy(self.oracles)
        fact = wrong_component_ref["oracles"]["ORACLE-TAB-002-A4"]["fact"]
        fact["source_oracle_id"] = "ORACLE-STR-001-A1"
        recompute_fact(fact)
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                self.manifest,
                oracle_registry=wrong_component_ref,
            )

        count_contract = copy.deepcopy(self.manifest)
        count_contract["cases"]["TAB-002"]["assertions"][2]["expected"] = 2
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                count_contract,
                oracle_registry=self.oracles,
            )

        human_gold_provenance = copy.deepcopy(self.oracles)
        human_gold_provenance["oracles"]["ORACLE-TAB-002-A4"][
            "review_provenance"
        ]["is_human_gold"] = True
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                self.manifest,
                oracle_registry=human_gold_provenance,
            )

    def test_assertion_evaluation_metadata_is_scoped_and_contract_checked(self) -> None:
        outside_scope = copy.deepcopy(self.manifest)
        assertion = outside_scope["cases"]["FMT-003"]["assertions"][0]
        assertion["evaluation"]["applicable_document_ids"] = ["DOC-QC-T-1067"]
        with self.assertRaises(ManifestValidationError):
            validate_manifest(outside_scope)

        mismatched_contract = copy.deepcopy(self.manifest)
        assertion = mismatched_contract["cases"]["FMT-003"]["assertions"][0]
        assertion["operator"] = "preserves"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(mismatched_contract)

        dedup_scope_escape = copy.deepcopy(self.manifest)
        dedup_assertion = dedup_scope_escape["cases"]["ROLE-005"]["assertions"][1]
        dedup_assertion["evaluation"]["applicable_document_ids"] = [
            "DOC-PEF-SOR-DUPLICATE",
            "DOC-IRPS-TRL-B",
        ]
        with self.assertRaises(ManifestValidationError):
            validate_manifest(dedup_scope_escape)

        layer_mismatch = copy.deepcopy(self.manifest)
        structure_assertion = layer_mismatch["cases"]["STR-001"]["assertions"][1]
        structure_assertion["layer"] = "source_ir"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(layer_mismatch)

        structure_contract_mismatch = copy.deepcopy(self.manifest)
        structure_assertion = structure_contract_mismatch["cases"]["STR-001"][
            "assertions"
        ][1]
        structure_assertion["target"] = "unreviewed.structure.target"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(structure_contract_mismatch)

        unknown_source_scope = copy.deepcopy(self.manifest)
        structure_assertion = unknown_source_scope["cases"]["STR-001"]["assertions"][1]
        structure_assertion["evaluation"]["source_oracle_id"] = "ORACLE-STR-999-A1"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(unknown_source_scope)

        incompatible_source_scope = copy.deepcopy(self.manifest)
        structure_assertion = incompatible_source_scope["cases"]["STR-001"][
            "assertions"
        ][1]
        structure_assertion["evaluation"]["source_oracle_id"] = "ORACLE-TAB-001-A1"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(incompatible_source_scope)

        section_target_mismatch = copy.deepcopy(self.manifest)
        section_target_mismatch["cases"]["SEM-002"]["assertions"][0]["target"] = (
            "section_5_1_7_2"
        )
        with self.assertRaises(ManifestValidationError):
            validate_manifest(section_target_mismatch)

        duplicate_section_locator = copy.deepcopy(self.manifest)
        duplicate_section_locator["cases"]["SEM-002"]["locators"].append(
            copy.deepcopy(duplicate_section_locator["cases"]["SEM-002"]["locators"][0])
        )
        with self.assertRaises(ManifestValidationError):
            validate_manifest(duplicate_section_locator)

    def test_role002_workflow_roles_do_not_depend_on_xlsb_parsing(
        self,
    ) -> None:
        case = self.manifest["cases"]["ROLE-002"]
        document_results = {
            document_id: {
                "document_id": document_id,
                "document_role": self.manifest["documents"][document_id][
                    "document_role"
                ],
                "ingress_route": self.manifest["documents"][document_id][
                    "ingress_route"
                ],
                "source_ir_status": (
                    "BLOCKED_PREFLIGHT"
                    if self.manifest["documents"][document_id]["format"] == "xlsb"
                    else "PARTIAL"
                ),
            }
            for document_id in case["document_ids"]
        }
        outcomes = runner_module._evaluate_case_assertions(
            case=case,
            manifest_documents=self.manifest["documents"],
            document_results=document_results,
            source_ir_mode="native",
            blocking_document_ids=[],
            oracle_registry=self.oracles,
            source_ir_snapshots={},
            structure_snapshots={},
            semantic_snapshots={},
            requirement_view_snapshots={},
            attachment_snapshots={},
        )
        by_id = {item["assertion_id"]: item for item in outcomes}
        self.assertEqual(by_id["ROLE-002-A1"]["status"], "PASS")
        self.assertEqual(by_id["ROLE-002-A2"]["status"], "PASS")
        self.assertEqual(by_id["ROLE-002-A3"]["status"], "GAP")
        self.assertEqual(
            by_id["ROLE-002-A3"]["reason_code"],
            "spreadsheet_ai_silver_review_pending",
        )

        mutated_results = copy.deepcopy(document_results)
        mutated_results["DOC-IRPS-RFQ-9HAT"].update(
            {
                "document_role": "customer_requirement_source",
                "ingress_route": "requirement_source",
            }
        )
        mutated = runner_module._evaluate_case_assertions(
            case=case,
            manifest_documents=self.manifest["documents"],
            document_results=mutated_results,
            source_ir_mode="native",
            blocking_document_ids=[],
            oracle_registry=self.oracles,
            source_ir_snapshots={},
            structure_snapshots={},
            semantic_snapshots={},
            requirement_view_snapshots={},
            attachment_snapshots={},
        )
        mutated_by_id = {item["assertion_id"]: item for item in mutated}
        self.assertEqual(mutated_by_id["ROLE-002-A1"]["status"], "FAIL")
        self.assertEqual(mutated_by_id["ROLE-002-A2"]["status"], "FAIL")

    def test_sem002_evaluates_full_document_semantics_and_projection(self) -> None:
        def write_objective_pdf(path: Path, *, with_child_requirement: bool) -> str:
            private_text = "private objective evidence must not enter the report"
            lines = [
                "5 System requirements",
                "5.1 Test requirements",
                "5.1.7 Functional validation",
                "5.1.7.1 Test objective",
                private_text,
            ]
            if with_child_requirement:
                lines.extend(
                    [
                        "5.1.7.1.1 Execution requirement",
                        "The system shall record the result.",
                    ]
                )
            document = pymupdf.open()
            try:
                page = document.new_page()
                for index, line in enumerate(lines, start=1):
                    page.insert_text((72, 60 + index * 60), line, fontsize=12)
                document.save(path)
            finally:
                document.close()
            return private_text

        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-QJLY-J7111029"]
        source_path = self.root / "objective.pdf"
        private_text = write_objective_pdf(
            source_path,
            with_child_requirement=False,
        )
        document["document_ref"] = f"external://training-corpus/{source_path.name}"
        document["source_sha256"] = hashlib.sha256(source_path.read_bytes()).hexdigest()

        with (
            patch(
                "scripts.extraction_case_runner.classify_content",
                wraps=runner_module.classify_content,
            ) as classify_mock,
            patch(
                "scripts.extraction_case_runner.project_requirements",
                wraps=runner_module.project_requirements,
            ) as projection_mock,
        ):
            passed = self._run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["SEM-002"],
            )

        assertions = {
            item["assertion_id"]: item
            for item in passed["cases"]["SEM-002"]["assertions"]
        }
        self.assertEqual(classify_mock.call_count, 1)
        self.assertEqual(projection_mock.call_count, 1)
        self.assertEqual(passed["cases"]["SEM-002"]["status"], "PASS")
        self.assertEqual(assertions["SEM-002-A1"]["status"], "PASS")
        self.assertEqual(assertions["SEM-002-A2"]["status"], "PASS")
        self.assertGreater(
            assertions["SEM-002-A1"]["evidence"]["semantic_scope_count"],
            1,
        )
        self.assertEqual(
            assertions["SEM-002-A2"]["evidence"]["primary_direct_requirement_count"],
            1,
        )

        child_path = self.root / "objective-with-child.pdf"
        write_objective_pdf(child_path, with_child_requirement=True)
        document["document_ref"] = f"external://training-corpus/{child_path.name}"
        document["source_sha256"] = hashlib.sha256(child_path.read_bytes()).hexdigest()
        failed = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["SEM-002"],
        )
        failed_assertions = {
            item["assertion_id"]: item
            for item in failed["cases"]["SEM-002"]["assertions"]
        }
        self.assertEqual(failed_assertions["SEM-002-A1"]["status"], "PASS")
        self.assertEqual(failed_assertions["SEM-002-A2"]["status"], "FAIL")
        self.assertEqual(
            failed_assertions["SEM-002-A2"]["reason_code"],
            "section_primary_direct_projection_mismatch",
        )
        self.assertEqual(
            failed_assertions["SEM-002-A2"]["evidence"][
                "primary_direct_requirement_count"
            ],
            0,
        )
        serialized = json.dumps([passed, failed], ensure_ascii=False)
        self.assertNotIn(private_text, serialized)
        self.assertNotIn(str(source_path), serialized)
        self.assertNotIn(str(child_path), serialized)

    def test_attachment_route_requires_explicit_trusted_revision_and_role(self) -> None:
        missing_revision = copy.deepcopy(self.manifest)
        missing_revision["documents"]["DOC-DRAWING-45104848"].pop("upload_revision_id")
        with self.assertRaises(ManifestValidationError):
            validate_manifest(missing_revision)

        wrong_role = copy.deepcopy(self.manifest)
        wrong_role["documents"]["DOC-DRAWING-45104848"]["document_role"] = (
            "customer_requirement_source"
        )
        with self.assertRaises(ManifestValidationError):
            validate_manifest(wrong_role)

        leaked_revision = copy.deepcopy(self.manifest)
        leaked_revision["documents"]["DOC-QC-T-1067"]["upload_revision_id"] = (
            "UPLOAD-REV-NOT-AN-ATTACHMENT"
        )
        with self.assertRaises(ManifestValidationError):
            validate_manifest(leaked_revision)

    def test_oracle_registry_rejects_unknown_binding_tamper_and_out_of_scope_fact(
        self,
    ) -> None:
        with self.assertRaises(ManifestValidationError):
            validate_manifest(self.manifest, oracle_registry={})

        unknown_reference = copy.deepcopy(self.manifest)
        unknown_reference["cases"]["STR-001"]["assertions"][0]["evaluation"][
            "oracle_id"
        ] = "ORACLE-STR-999-A1"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(unknown_reference, oracle_registry=self.oracles)

        source_mismatch = copy.deepcopy(self.oracles)
        source_mismatch["oracles"]["ORACLE-STR-001-A1"]["source_sha256"] = "0" * 64
        with self.assertRaises(ManifestValidationError):
            validate_manifest(self.manifest, oracle_registry=source_mismatch)

        fingerprint_tamper = copy.deepcopy(self.oracles)
        fingerprint_tamper["oracles"]["ORACLE-STR-001-A1"]["fact"][
            "range_fingerprint"
        ] = "0" * 64
        with self.assertRaises(ManifestValidationError):
            validate_manifest(self.manifest, oracle_registry=fingerprint_tamper)

        mixed_component_tamper = copy.deepcopy(self.oracles)
        mixed_fact = mixed_component_tamper["oracles"]["ORACLE-SEM-003-A1"]["fact"]
        mixed_fact["parameter_table_group_fingerprint"] = "0" * 64
        mixed_fact["mixed_fingerprint"] = (
            runner_module._fingerprint_docx_mixed_content_group(
                narratives=mixed_fact["narratives"],
                parameter_table_oracle_id=str(
                    mixed_fact["parameter_table_oracle_id"]
                ),
                expected_parameter_table_count=int(
                    mixed_fact["expected_parameter_table_count"]
                ),
                parameter_table_group_fingerprint=str(
                    mixed_fact["parameter_table_group_fingerprint"]
                ),
            )
        )
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                self.manifest,
                oracle_registry=mixed_component_tamper,
            )

        out_of_scope = copy.deepcopy(self.oracles)
        out_of_scope["oracles"]["ORACLE-STR-001-A1"]["fact"]["cells"][0][
            "cell_range"
        ] = "Z99"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(self.manifest, oracle_registry=out_of_scope)

        extra_property = copy.deepcopy(self.oracles)
        extra_property["oracles"]["ORACLE-STR-001-A1"]["customer_text"] = "forbidden"
        with self.assertRaises(ManifestValidationError):
            validate_manifest(self.manifest, oracle_registry=extra_property)

    def test_duplicate_relationships_require_existing_acyclic_nonself_targets(
        self,
    ) -> None:
        duplicate_id = "DOC-PEF-SOR-DUPLICATE"
        canonical_id = "DOC-PEF-SOR"

        invalid_manifests: list[dict[str, object]] = []
        unknown = copy.deepcopy(self.manifest)
        unknown["documents"][duplicate_id]["duplicate_of"] = "DOC-UNKNOWN"
        invalid_manifests.append(unknown)

        self_reference = copy.deepcopy(self.manifest)
        self_reference["documents"][duplicate_id]["duplicate_of"] = duplicate_id
        invalid_manifests.append(self_reference)

        cycle = copy.deepcopy(self.manifest)
        cycle["documents"][canonical_id]["duplicate_of"] = duplicate_id
        invalid_manifests.append(cycle)

        for invalid_manifest in invalid_manifests:
            with self.subTest(), self.assertRaises(ManifestValidationError):
                validate_manifest(invalid_manifest)

    def test_external_reference_cannot_escape_corpus_root(self) -> None:
        with self.assertRaises(CorpusPathError):
            resolve_document_path(
                self.root,
                "external://training-corpus/../../outside.pdf",
            )

    def test_strict_schema_rejects_malicious_document_id_before_preview_path_use(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"].pop("DOC-QC-T-1067")
        manifest["documents"]["../../escape"] = document
        with self.assertRaises(ManifestValidationError):
            validate_manifest(manifest)

    def test_report_destination_cannot_mutate_controlled_corpus(self) -> None:
        with self.assertRaises(CorpusPathError):
            require_report_destination_outside_corpus(
                self.root / "reports",
                self.root,
            )
        outside = self.root.parent / f"{self.root.name}-reports"
        self.assertEqual(
            require_report_destination_outside_corpus(outside, self.root),
            outside.resolve(),
        )
        with self.assertRaises(CorpusPathError):
            write_reports({}, self.root / "reports", corpus_root=self.root)

    def test_native_build_exception_is_visible_to_strict_source_ir_gate(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "safe.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "fixture"
        workbook.save(path)
        workbook.close()
        document["document_ref"] = "external://training-corpus/safe.xlsx"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

        with patch(
            "scripts.extraction_case_runner._run_source_ir_observation",
            return_value={
                "source_ir_status": "BUILD_EXCEPTION",
                "source_ir_counts": {},
                "source_ir_error_type": "RuntimeError",
            },
        ):
            report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
            )

        self.assertTrue(report["summary"]["identity_gate_pass"])
        self.assertFalse(report["summary"]["source_ir_observation_ok"])

    def test_native_snapshot_must_match_the_manifest_verified_source_identity(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "safe.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "fixture"
        workbook.save(path)
        workbook.close()
        document["document_ref"] = "external://training-corpus/safe.xlsx"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        mismatched_snapshot = SimpleNamespace(source_sha256="f" * 64)

        with patch(
            "scripts.extraction_case_runner.build_source_ir",
            return_value=mismatched_snapshot,
        ):
            report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
            )

        document_result = report["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        self.assertTrue(report["summary"]["identity_gate_pass"])
        self.assertFalse(report["summary"]["source_ir_observation_ok"])
        self.assertEqual(
            document_result["source_ir_status"],
            "SOURCE_IDENTITY_CHANGED",
        )
        self.assertEqual(document_result["source_ir_counts"], {})

    def test_native_observation_reuses_the_verified_preflight_report(self) -> None:
        path = self.root / "preflight-reuse.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "fixture"
        workbook.save(path)
        workbook.close()
        report = runner_module.preflight_document(path)
        snapshot = SimpleNamespace(
            source_sha256=report.identity.source_sha256,
            status=SimpleNamespace(value="COMPLETE"),
            nodes=(),
            assets=(),
            diagnostics=(),
            adapter_id="test-adapter",
            adapter_version="1.0.0",
        )

        with patch(
            "scripts.extraction_case_runner.build_source_ir",
            return_value=snapshot,
        ) as builder:
            outcome = runner_module._run_source_ir_observation(
                path,
                expected_source_sha256=report.identity.source_sha256,
                preflight_report=report,
            )

        builder.assert_called_once_with(path, preflight_report=report)
        self.assertEqual(outcome["source_ir_status"], "COMPLETE")

    def test_required_missing_and_hash_mismatch_are_hard_blockers(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        document["document_ref"] = "external://training-corpus/missing.xlsx"
        missing = self._run_manifest(
            manifest,
            corpus_root=self.root,
            selected_case_ids=["STR-001"],
        )
        self.assertFalse(missing["summary"]["identity_gate_pass"])
        self.assertEqual(
            missing["hard_blockers"][0]["code"], "required_document_missing"
        )

        path = self.root / "changed.xlsx"
        path.write_bytes(b"changed")
        document["document_ref"] = "external://training-corpus/changed.xlsx"
        mismatched = self._run_manifest(
            manifest,
            corpus_root=self.root,
            selected_case_ids=["STR-001"],
        )
        self.assertEqual(mismatched["hard_blockers"][0]["code"], "source_hash_mismatch")

    def test_source_change_between_manifest_hash_and_preflight_is_a_hard_blocker(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "changed-during-preflight.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "fixture"
        workbook.save(path)
        workbook.close()
        stale_hash = "0" * 64
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = stale_hash

        with patch(
            "scripts.extraction_case_runner.sha256_file",
            return_value=stale_hash,
        ):
            report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
            )

        self.assertFalse(report["summary"]["identity_gate_pass"])
        self.assertEqual(
            report["hard_blockers"],
            [
                {
                    "document_id": "DOC-LONG-AXLE-PDS-XLSX",
                    "code": "source_changed_during_preflight",
                }
            ],
        )
        self.assertEqual(
            report["documents"]["DOC-LONG-AXLE-PDS-XLSX"]["identity_status"],
            "SOURCE_CHANGED",
        )

    def test_source_changed_during_legacy_observation_is_not_reported_as_verified(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-QC-T-1067"]
        path = self.root / "source.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with path.open("wb") as handle:
            writer.write(handle)
        document["document_ref"] = "external://training-corpus/source.pdf"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

        def mutate_source(*_: object, **__: object) -> dict[str, object]:
            with path.open("ab") as handle:
                handle.write(b"\n")
            return {"legacy_status": "OBSERVED_CORE", "legacy_counts": {}}

        with patch(
            "scripts.extraction_case_runner._run_legacy_observation",
            side_effect=mutate_source,
        ):
            report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                legacy_mode="core",
                selected_case_ids=["STR-003"],
            )

        self.assertFalse(report["summary"]["identity_gate_pass"])
        self.assertEqual(
            report["hard_blockers"],
            [
                {
                    "document_id": "DOC-QC-T-1067",
                    "code": "source_changed_during_legacy_observation",
                }
            ],
        )

    def test_legacy_observation_timeout_terminates_isolated_worker(self) -> None:
        if "fork" not in multiprocessing.get_all_start_methods():
            self.skipTest("timeout worker test requires fork-compatible test isolation")
        source = self.root / "source.pdf"
        source.write_bytes(b"approved source bytes")
        expected_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        fork_context = multiprocessing.get_context("fork")

        def slow_observation(*_: object, **__: object) -> dict[str, object]:
            time.sleep(1)
            return {"legacy_status": "OBSERVED_CORE", "legacy_counts": {}}

        with (
            patch.object(
                runner_module.multiprocessing,
                "get_context",
                return_value=fork_context,
            ),
            patch.object(
                runner_module,
                "_run_staged_legacy_observation",
                side_effect=slow_observation,
            ),
        ):
            observation = runner_module._run_legacy_observation(
                source,
                document_id="DOC-QC-T-1067",
                legacy_mode="core",
                expected_source_sha256=expected_hash,
                timeout_seconds=0.02,
            )

        self.assertEqual(observation["legacy_status"], "BLOCKED_TIMEOUT")
        self.assertEqual(observation["legacy_counts"], {})
        self.assertEqual(observation["legacy_timeout_seconds"], 0.02)

    def test_programmatic_legacy_run_uses_the_bounded_default(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-QC-T-1067"]
        path = self.root / "bounded-default.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with path.open("wb") as handle:
            writer.write(handle)
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

        with patch(
            "scripts.extraction_case_runner._run_legacy_observation",
            return_value={"legacy_status": "OBSERVED_CORE", "legacy_counts": {}},
        ) as observation:
            report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                legacy_mode="core",
                selected_case_ids=["STR-003"],
            )

        self.assertEqual(
            observation.call_args.kwargs["timeout_seconds"],
            runner_module.DEFAULT_LEGACY_TIMEOUT_SECONDS,
        )
        self.assertEqual(
            report["legacy_timeout_seconds"],
            runner_module.DEFAULT_LEGACY_TIMEOUT_SECONDS,
        )

    def test_legacy_timeout_rejects_nonfinite_values(self) -> None:
        for timeout in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(timeout=timeout):
                with self.assertRaisesRegex(ValueError, "finite and positive"):
                    self._run_manifest(
                        copy.deepcopy(self.manifest),
                        corpus_root=self.root,
                        legacy_mode="core",
                        selected_case_ids=["STR-003"],
                        legacy_timeout_seconds=timeout,
                    )

    def test_safe_report_does_not_contain_customer_text_or_absolute_path(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "safe.xlsx"
        secret_text = "customer secret requirement 12345"
        workbook = Workbook()
        workbook.active.append([secret_text])
        workbook.save(path)
        workbook.close()
        document["document_ref"] = "external://training-corpus/safe.xlsx"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

        report = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["STR-001"],
        )
        output_dir = self.root.parent / f"{self.root.name}-reports"
        self.addCleanup(shutil.rmtree, output_dir, True)
        json_path, markdown_path = write_reports(
            report,
            output_dir,
            corpus_root=self.root,
        )
        serialized = json_path.read_text(encoding="utf-8") + markdown_path.read_text(
            encoding="utf-8"
        )

        self.assertTrue(report["summary"]["identity_gate_pass"])
        self.assertNotIn(secret_text, serialized)
        self.assertNotIn(str(self.root), serialized)

    def test_hash_only_oracle_passes_from_same_native_snapshot_without_report_leak(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "oracle-source.xlsx"
        secret_text = "customer private high-temperature requirement"
        values = [secret_text, "2", "3", "4", "5", "6", "7", "8"]
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Scope"
        worksheet.append(values)
        workbook.save(path)
        workbook.close()
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_str_oracle(manifest, values)
        real_builder = runner_module.build_source_ir
        real_structure_builder = runner_module.build_structure_tree

        with (
            patch(
                "scripts.extraction_case_runner.build_source_ir",
                wraps=real_builder,
            ) as build_mock,
            patch(
                "scripts.extraction_case_runner.build_structure_tree",
                wraps=real_structure_builder,
            ) as structure_mock,
        ):
            report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
                oracle_registry=registry,
            )

        assertions = {
            item["assertion_id"]: item
            for item in report["cases"]["STR-001"]["assertions"]
        }
        self.assertEqual(build_mock.call_count, 1)
        self.assertEqual(structure_mock.call_count, 1)
        self.assertEqual(
            report["oracle_registry_sha256"],
            hashlib.sha256(
                runner_module._canonical_json_payload(registry).encode("utf-8")
            ).hexdigest(),
        )
        self.assertEqual(assertions["STR-001-A1"]["status"], "PASS")
        self.assertEqual(
            assertions["STR-001-A1"]["reason_code"],
            "oracle_fidelity_observed",
        )
        self.assertEqual(assertions["STR-001-A1"]["evidence"]["mismatch_count"], 0)
        self.assertEqual(assertions["STR-001-A2"]["status"], "PASS")
        self.assertEqual(
            assertions["STR-001-A2"]["reason_code"],
            "single_table_row_boundary_observed",
        )
        self.assertEqual(
            assertions["STR-001-A2"]["evidence"]["observed_row_boundary_count"],
            1,
        )

        def duplicate_scoped_row(source_snapshot: object, **_: object) -> object:
            tree = real_structure_builder(source_snapshot)
            row = next(
                node
                for node in tree.nodes
                if node.kind.value == "table_row" and node.table_row_index == 0
            )
            identity_key = f"model:{hashlib.sha256(b'duplicate-row').hexdigest()}"
            duplicate = row.model_copy(
                update={
                    "structure_node_id": stable_structure_node_id(
                        source_ir_id=tree.source_ir_id,
                        resolver_id=tree.resolver_id,
                        resolver_version=tree.resolver_version,
                        kind=row.kind,
                        identity_key=identity_key,
                    ),
                    "identity_key": identity_key,
                    "source_order": len(tree.nodes),
                }
            )
            return tree.model_copy(update={"nodes": (*tree.nodes, duplicate)})

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=duplicate_scoped_row,
        ):
            duplicate_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
                oracle_registry=registry,
            )
        duplicate_assertion = {
            item["assertion_id"]: item
            for item in duplicate_report["cases"]["STR-001"]["assertions"]
        }["STR-001-A2"]
        self.assertEqual(duplicate_assertion["status"], "FAIL")
        self.assertEqual(
            duplicate_assertion["evidence"]["observed_row_boundary_count"],
            2,
        )
        self.assertGreater(
            duplicate_assertion["evidence"]["invalid_boundary_count"],
            0,
        )

        def alternate_resolver(source_snapshot: object, **_: object) -> object:
            return real_structure_builder(
                source_snapshot,
                resolver_id="alternate-structure-resolver",
                resolver_version="9.9.9",
            )

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=alternate_resolver,
        ):
            alternate_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
                oracle_registry=registry,
            )
        alternate_assertion = {
            item["assertion_id"]: item
            for item in alternate_report["cases"]["STR-001"]["assertions"]
        }["STR-001-A2"]
        self.assertEqual(alternate_assertion["status"], "FAIL")
        self.assertEqual(
            alternate_assertion["reason_code"],
            "structure_resolver_identity_mismatch",
        )
        serialized = json.dumps(report, ensure_ascii=False)
        self.assertNotIn(secret_text, serialized)
        self.assertNotIn(str(path), serialized)

    def test_docx_single_row_scope_evaluators_pass_and_fail_closed(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-ESOW-DHT120-STATOR"]
        path = self.root / "single-row-scope.docx"
        secret_text = "private supplier plan requirement"
        mutated_text = "mutated private supplier plan requirement"
        word = WordDocument()
        table = word.add_table(rows=1, cols=3)
        table.cell(0, 0).text = "2.1.22"
        table.cell(0, 1).text = f"The supplier shall maintain {secret_text}."
        table.cell(0, 1).add_paragraph(
            "The plan shall describe each required communication method."
        )
        table.cell(0, 2).text = "SE"
        word.save(path)
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_str5_oracle(manifest, path)

        report = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["STR-005"],
            oracle_registry=registry,
        )
        assertions = {
            item["assertion_id"]: item
            for item in report["cases"]["STR-005"]["assertions"]
        }
        self.assertEqual(report["cases"]["STR-005"]["status"], "PASS")
        self.assertEqual(
            {item["status"] for item in assertions.values()},
            {"PASS"},
        )
        self.assertEqual(
            assertions["STR-005-A1"]["reason_code"],
            "docx_single_table_row_boundary_observed",
        )
        self.assertEqual(
            assertions["STR-005-A2"]["evidence"][
                "top_level_requirement_count"
            ],
            1,
        )
        self.assertEqual(
            assertions["STR-005-A3"]["evidence"][
                "sentence_or_parallel_verb_split_count"
            ],
            0,
        )

        missing_provenance = copy.deepcopy(registry)
        del missing_provenance["oracles"]["ORACLE-STR-005-A1"][
            "review_provenance"
        ]
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                manifest,
                oracle_registry=missing_provenance,
            )

        forged_human_gold = copy.deepcopy(registry)
        forged_human_gold["oracles"]["ORACLE-STR-005-A1"][
            "review_provenance"
        ]["is_human_gold"] = True
        with self.assertRaises(ManifestValidationError):
            validate_manifest(
                manifest,
                oracle_registry=forged_human_gold,
            )

        source_snapshot = runner_module.build_source_ir(path)
        nodes = list(source_snapshot.nodes)
        table_node_index = next(
            index
            for index, node in enumerate(nodes)
            if node.table_payload is not None
        )
        table_node = nodes[table_node_index]
        source_table = table_node.table_payload
        self.assertIsNotNone(source_table)
        changed_cells = list(source_table.cells)
        changed_cells[1] = changed_cells[1].model_copy(
            update={"raw_text": mutated_text}
        )
        nodes[table_node_index] = table_node.model_copy(
            update={
                "table_payload": source_table.model_copy(
                    update={"cells": tuple(changed_cells)}
                )
            }
        )
        changed_snapshot = source_snapshot.model_copy(
            update={"nodes": tuple(nodes)}
        )
        real_structure_builder = runner_module.build_structure_tree
        with (
            patch(
                "scripts.extraction_case_runner.build_source_ir",
                return_value=changed_snapshot,
            ),
            patch(
                "scripts.extraction_case_runner.build_structure_tree",
                wraps=real_structure_builder,
            ) as structure_mock,
        ):
            changed_report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-005"],
                oracle_registry=registry,
            )
        self.assertEqual(structure_mock.call_count, 0)
        changed_assertions = changed_report["cases"]["STR-005"]["assertions"]
        self.assertEqual(
            {item["status"] for item in changed_assertions},
            {"FAIL"},
        )
        self.assertEqual(
            {
                item["reason_code"]
                for item in changed_assertions
            },
            {"structure_source_scope_mismatch"},
        )

        def duplicate_scoped_row(source: object, **_: object) -> object:
            tree = real_structure_builder(source)
            row = next(
                node
                for node in tree.nodes
                if node.kind.value == "table_row"
                and node.table_row_index == 0
            )
            identity_key = f"model:{hashlib.sha256(b'docx-split-row').hexdigest()}"
            duplicate = row.model_copy(
                update={
                    "structure_node_id": stable_structure_node_id(
                        source_ir_id=tree.source_ir_id,
                        resolver_id=tree.resolver_id,
                        resolver_version=tree.resolver_version,
                        kind=row.kind,
                        identity_key=identity_key,
                    ),
                    "identity_key": identity_key,
                    "source_order": len(tree.nodes),
                }
            )
            return tree.model_copy(update={"nodes": (*tree.nodes, duplicate)})

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=duplicate_scoped_row,
        ):
            duplicate_report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-005"],
                oracle_registry=registry,
            )
        duplicate_assertions = {
            item["assertion_id"]: item
            for item in duplicate_report["cases"]["STR-005"]["assertions"]
        }
        self.assertEqual(
            {item["status"] for item in duplicate_assertions.values()},
            {"FAIL"},
        )
        self.assertEqual(
            duplicate_assertions["STR-005-A1"]["evidence"][
                "observed_row_boundary_count"
            ],
            2,
        )

        serialized = json.dumps(
            [report, changed_report, duplicate_report],
            ensure_ascii=False,
        )
        self.assertNotIn(secret_text, serialized)
        self.assertNotIn(mutated_text, serialized)
        self.assertNotIn(str(path), serialized)

    def test_canonical_grid_fingerprint_has_no_newline_delimiter_collision(
        self,
    ) -> None:
        one_cell = [
            {
                "row": 0,
                "column": 0,
                "row_span": 1,
                "column_span": 1,
                "raw_text": "x\n1,0,1,1:y",
            }
        ]
        two_cells = [
            {
                "row": 0,
                "column": 0,
                "row_span": 1,
                "column_span": 1,
                "raw_text": "x",
            },
            {
                "row": 1,
                "column": 0,
                "row_span": 1,
                "column_span": 1,
                "raw_text": "y",
            },
        ]

        self.assertNotEqual(
            runner_module._sha256_text(runner_module._canonical_json_payload(one_cell)),
            runner_module._sha256_text(
                runner_module._canonical_json_payload(two_cells)
            ),
        )

    def test_hash_only_oracle_fails_on_native_fact_mismatch_without_report_leak(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "oracle-source.xlsx"
        secret_text = "customer private acceptance requirement"
        mutated_text = "different private requirement"
        values = [secret_text, "2", "3", "4", "5", "6", "7", "8"]
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Scope"
        worksheet.append(values)
        workbook.save(path)
        workbook.close()
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_str_oracle(manifest, values)

        snapshot = runner_module.build_source_ir(path)
        nodes = list(snapshot.nodes)
        table_node_index = next(
            index for index, node in enumerate(nodes) if node.table_payload is not None
        )
        table_node = nodes[table_node_index]
        table = table_node.table_payload
        self.assertIsNotNone(table)
        cells = list(table.cells)
        cells[0] = cells[0].model_copy(update={"raw_text": mutated_text})
        changed_table = table.model_copy(update={"cells": tuple(cells)})
        nodes[table_node_index] = table_node.model_copy(
            update={"table_payload": changed_table}
        )
        changed_snapshot = snapshot.model_copy(update={"nodes": tuple(nodes)})
        real_structure_builder = runner_module.build_structure_tree

        with (
            patch(
                "scripts.extraction_case_runner.build_source_ir",
                return_value=changed_snapshot,
            ),
            patch(
                "scripts.extraction_case_runner.build_structure_tree",
                wraps=real_structure_builder,
            ) as structure_mock,
        ):
            report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
                oracle_registry=registry,
            )
        self.assertEqual(structure_mock.call_count, 0)

        assertions = {
            item["assertion_id"]: item
            for item in report["cases"]["STR-001"]["assertions"]
        }
        outcome = assertions["STR-001-A1"]
        self.assertEqual(outcome["status"], "FAIL")
        self.assertEqual(outcome["reason_code"], "oracle_fidelity_mismatch")
        self.assertGreater(outcome["evidence"]["mismatch_count"], 0)
        self.assertEqual(assertions["STR-001-A2"]["status"], "FAIL")
        self.assertEqual(
            assertions["STR-001-A2"]["reason_code"],
            "structure_source_scope_mismatch",
        )
        serialized = json.dumps(report, ensure_ascii=False)
        self.assertNotIn(secret_text, serialized)
        self.assertNotIn(mutated_text, serialized)
        self.assertNotIn(str(path), serialized)

    def test_docx_data_row_structure_evaluator_passes_and_fails_closed(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-IRPS-PRODUCT-INFO"]
        path = self.root / "parameter-table.docx"
        secret_text = "private parameter requirement"
        word = WordDocument()
        table = word.add_table(rows=3, cols=6)
        rows = (
            (
                "Parameter",
                "Unit",
                "Lower limit",
                "Typ.",
                "Upper limit",
                "Condition",
            ),
            (secret_text, "mm", "1", "2", "3", "ambient"),
            ("Second parameter", "mm", "4", "5", "6", "ambient"),
        )
        for row_index, values in enumerate(rows):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = value
        word.save(path)
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_tab1_oracle(manifest, path)
        real_structure_builder = runner_module.build_structure_tree

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            wraps=real_structure_builder,
        ) as structure_mock:
            passed = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-001"],
                oracle_registry=registry,
            )

        assertions = {
            item["assertion_id"]: item
            for item in passed["cases"]["TAB-001"]["assertions"]
        }
        self.assertEqual(structure_mock.call_count, 1)
        self.assertEqual(assertions["TAB-001-A1"]["status"], "PASS")
        self.assertEqual(assertions["TAB-001-A2"]["status"], "PASS")
        self.assertEqual(assertions["TAB-001-A3"]["status"], "PASS")
        self.assertEqual(assertions["TAB-001-A4"]["status"], "PASS")
        self.assertEqual(
            assertions["TAB-001-A2"]["evidence"]["observed_data_row_boundary_count"],
            2,
        )
        self.assertEqual(
            assertions["TAB-001-A2"]["evidence"]["missing_cell_grounding_count"],
            0,
        )
        self.assertEqual(
            assertions["TAB-001-A3"]["evidence"][
                "row_with_exact_semantic_requirement_count"
            ],
            2,
        )
        self.assertEqual(
            assertions["TAB-001-A4"]["evidence"][
                "row_with_exact_primary_direct_projection_count"
            ],
            2,
        )

        def remove_one_data_row(source_snapshot: object, **_: object) -> object:
            tree = real_structure_builder(source_snapshot)
            removed = next(
                node
                for node in tree.nodes
                if node.kind.value == "table_row" and node.table_row_index == 1
            )
            return tree.model_copy(
                update={
                    "nodes": tuple(
                        node
                        for node in tree.nodes
                        if node.structure_node_id != removed.structure_node_id
                    ),
                    "relations": tuple(
                        relation
                        for relation in tree.relations
                        if removed.structure_node_id
                        not in {
                            relation.source_structure_node_id,
                            relation.target_structure_node_id,
                        }
                    ),
                }
            )

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=remove_one_data_row,
        ):
            failed = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-001"],
                oracle_registry=registry,
            )
        failed_assertion = {
            item["assertion_id"]: item
            for item in failed["cases"]["TAB-001"]["assertions"]
        }["TAB-001-A2"]
        self.assertEqual(failed_assertion["status"], "FAIL")
        self.assertEqual(
            failed_assertion["reason_code"],
            "table_data_row_boundary_mismatch",
        )

        def duplicate_data_row(source_snapshot: object, **_: object) -> object:
            tree = real_structure_builder(source_snapshot)
            row = next(
                node
                for node in tree.nodes
                if node.kind.value == "table_row" and node.table_row_index == 1
            )
            identity_key = f"model:{hashlib.sha256(b'duplicate-data-row').hexdigest()}"
            duplicate = row.model_copy(
                update={
                    "structure_node_id": stable_structure_node_id(
                        source_ir_id=tree.source_ir_id,
                        resolver_id=tree.resolver_id,
                        resolver_version=tree.resolver_version,
                        kind=row.kind,
                        identity_key=identity_key,
                    ),
                    "identity_key": identity_key,
                    "source_order": len(tree.nodes),
                }
            )
            return tree.model_copy(update={"nodes": (*tree.nodes, duplicate)})

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=duplicate_data_row,
        ):
            duplicate_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-001"],
                oracle_registry=registry,
            )
        duplicate_assertion = {
            item["assertion_id"]: item
            for item in duplicate_report["cases"]["TAB-001"]["assertions"]
        }["TAB-001-A2"]
        self.assertEqual(duplicate_assertion["status"], "FAIL")
        self.assertEqual(
            duplicate_assertion["evidence"]["observed_data_row_boundary_count"],
            3,
        )
        self.assertGreater(
            duplicate_assertion["evidence"]["invalid_boundary_count"],
            0,
        )

        real_classifier = runner_module.classify_content

        def reject_one_parameter_row(
            source_snapshot: object,
            tree: object,
            **kwargs: object,
        ) -> object:
            semantics = real_classifier(source_snapshot, tree, **kwargs)
            target = next(
                annotation
                for annotation in semantics.annotations
                if annotation.requirement_subtype
                is runner_module.RequirementSubtype.PERFORMANCE_PARAMETER_REQUIREMENT
            )
            rejected = target.model_copy(
                update={
                    "status": runner_module.SemanticAnnotationStatus.REJECTED,
                }
            )
            return semantics.model_copy(
                update={
                    "annotations": tuple(
                        rejected
                        if annotation.semantic_annotation_id
                        == target.semantic_annotation_id
                        else annotation
                        for annotation in semantics.annotations
                    )
                }
            )

        with patch(
            "scripts.extraction_case_runner.classify_content",
            side_effect=reject_one_parameter_row,
        ):
            rejected_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-001"],
                oracle_registry=registry,
            )
        rejected_assertions = {
            item["assertion_id"]: item
            for item in rejected_report["cases"]["TAB-001"]["assertions"]
        }
        self.assertEqual(rejected_assertions["TAB-001-A3"]["status"], "FAIL")
        self.assertEqual(
            rejected_assertions["TAB-001-A3"]["reason_code"],
            "performance_parameter_row_semantic_mismatch",
        )
        self.assertEqual(rejected_assertions["TAB-001-A4"]["status"], "FAIL")
        self.assertEqual(
            rejected_assertions["TAB-001-A4"]["reason_code"],
            "requirement_per_parameter_row_mapping_mismatch",
        )
        self.assertEqual(
            rejected_assertions["TAB-001-A4"]["evidence"][
                "missing_projection_row_count"
            ],
            1,
        )

        real_projector = runner_module.project_requirements

        def drop_one_row_source_ref(
            *args: object,
            **kwargs: object,
        ) -> object:
            view = real_projector(*args, **kwargs)
            target = next(
                node
                for node in view.nodes
                if node.requirement_subtype
                is runner_module.RequirementSubtype.PERFORMANCE_PARAMETER_REQUIREMENT
            )
            dropped_ref = next(
                ref
                for ref in target.source_refs
                if ref.kind is runner_module.StructureSourceRefKind.TABLE_CELL
            )
            partial = target.model_copy(
                update={
                    "source_refs": tuple(
                        ref
                        for ref in target.source_refs
                        if ref != dropped_ref
                    )
                }
            )
            return view.model_copy(
                update={
                    "nodes": tuple(
                        partial
                        if node.requirement_view_node_id
                        == target.requirement_view_node_id
                        else node
                        for node in view.nodes
                    )
                }
            )

        with (
            patch(
                "scripts.extraction_case_runner.project_requirements",
                side_effect=drop_one_row_source_ref,
            ),
            patch(
                "scripts.extraction_case_runner."
                "validate_requirement_view_against_inputs",
            ),
        ):
            partial_ref_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-001"],
                oracle_registry=registry,
            )
        partial_ref_assertions = {
            item["assertion_id"]: item
            for item in partial_ref_report["cases"]["TAB-001"]["assertions"]
        }
        self.assertEqual(partial_ref_assertions["TAB-001-A3"]["status"], "PASS")
        self.assertEqual(partial_ref_assertions["TAB-001-A4"]["status"], "FAIL")
        self.assertEqual(
            partial_ref_assertions["TAB-001-A4"]["reason_code"],
            "requirement_per_parameter_row_mapping_mismatch",
        )
        self.assertGreater(
            partial_ref_assertions["TAB-001-A4"]["evidence"][
                "missing_row_source_grounding_count"
            ],
            0,
        )
        self.assertNotIn(secret_text, json.dumps(passed, ensure_ascii=False))
        self.assertNotIn(secret_text, json.dumps(failed, ensure_ascii=False))
        self.assertNotIn(
            secret_text,
            json.dumps(rejected_report, ensure_ascii=False),
        )
        self.assertNotIn(
            secret_text,
            json.dumps(partial_ref_report, ensure_ascii=False),
        )

    def test_sem003_mixed_content_evaluators_pass_and_fail_closed(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-IRPS-PRODUCT-INFO"]
        path = self.root / "sem003-mixed-content.docx"
        secret_marker = "SEM003_PRIVATE_SYNTHETIC_MARKER_9F4A"
        narrative_texts = (
            (
                "The adjustment unit "
                f"{secret_marker} consists of two mechanical stages."
            ),
            (
                "The stroke calculation "
                f"{secret_marker} is performed from the sensor values."
            ),
            (
                "The internal sensor board "
                f"{secret_marker} provides a converted output signal."
            ),
        )
        private_parameter = f"{secret_marker} performance parameter"
        word = WordDocument()
        for text in narrative_texts:
            word.add_paragraph(text)
        table = word.add_table(rows=3, cols=6)
        rows = (
            (
                "Parameter",
                "Unit",
                "Lower limit",
                "Typ.",
                "Upper limit",
                "Condition",
            ),
            (private_parameter, "mm", "1", "2", "3", "ambient"),
            ("Second synthetic parameter", "mm", "4", "5", "6", "ambient"),
        )
        for row_index, values in enumerate(rows):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = value
        word.save(path)
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

        registry = self._configure_tab1_oracle(manifest, path)
        source_snapshot = runner_module.build_source_ir(path)
        narrative_nodes = sorted(
            (
                node
                for node in source_snapshot.nodes
                if node.kind.value == "paragraph"
                and node.native_locator.format.value == "docx"
                and node.native_locator.body_element_index is not None
                and node.native_locator.table_index is None
            ),
            key=lambda node: int(node.native_locator.body_element_index),
        )
        self.assertEqual(
            tuple(str(node.raw_text) for node in narrative_nodes),
            narrative_texts,
        )
        narrative_records = [
            {
                "locator": {
                    "format": "docx",
                    "body_element_index": int(
                        node.native_locator.body_element_index
                    ),
                },
                "kind": node.kind.value,
                "raw_sha256": runner_module._sha256_text(str(node.raw_text)),
            }
            for node in narrative_nodes
        ]
        parameter_fact = registry["oracles"]["ORACLE-TAB-001-A1"]["fact"]
        mixed_fact = registry["oracles"]["ORACLE-SEM-003-A1"]["fact"]
        mixed_fact.update(
            {
                "expected_narrative_count": len(narrative_records),
                "narratives": narrative_records,
                "expected_parameter_table_count": parameter_fact[
                    "expected_table_count"
                ],
                "parameter_table_group_fingerprint": parameter_fact[
                    "group_fingerprint"
                ],
            }
        )
        mixed_fact["mixed_fingerprint"] = (
            runner_module._fingerprint_docx_mixed_content_group(
                narratives=mixed_fact["narratives"],
                parameter_table_oracle_id=str(
                    mixed_fact["parameter_table_oracle_id"]
                ),
                expected_parameter_table_count=int(
                    mixed_fact["expected_parameter_table_count"]
                ),
                parameter_table_group_fingerprint=str(
                    mixed_fact["parameter_table_group_fingerprint"]
                ),
            )
        )

        passed = run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["SEM-003"],
            oracle_registry=registry,
        )
        passed_assertions = {
            item["assertion_id"]: item
            for item in passed["cases"]["SEM-003"]["assertions"]
        }
        self.assertEqual(
            {
                assertion_id: outcome["status"]
                for assertion_id, outcome in passed_assertions.items()
            },
            {
                "SEM-003-A1": "PASS",
                "SEM-003-A2": "PASS",
                "SEM-003-A3": "PASS",
                "SEM-003-A4": "PASS",
            },
        )
        self.assertEqual(
            passed_assertions["SEM-003-A2"]["evidence"][
                "valid_technical_narrative_count"
            ],
            3,
        )
        self.assertEqual(
            passed_assertions["SEM-003-A4"]["evidence"][
                "active_parameter_requirement_owner_count"
            ],
            2,
        )

        real_classifier = runner_module.classify_content

        def leave_one_narrative_unresolved(
            source_ir: object,
            tree: object,
            **kwargs: object,
        ) -> object:
            semantics = real_classifier(source_ir, tree, **kwargs)
            technical_annotations = [
                annotation
                for annotation in semantics.annotations
                if annotation.reason_codes == ("technical_explanation_signal",)
            ]
            self.assertEqual(len(technical_annotations), 3)
            target = technical_annotations[0]
            unresolved = target.model_copy(
                update={
                    "semantic_class": runner_module.SemanticClass.UNRESOLVED,
                    "status": runner_module.SemanticAnnotationStatus.UNRESOLVED,
                    "reason_codes": ("forced_unresolved_for_test",),
                }
            )
            return semantics.model_copy(
                update={
                    "annotations": tuple(
                        unresolved
                        if annotation.semantic_annotation_id
                        == target.semantic_annotation_id
                        else annotation
                        for annotation in semantics.annotations
                    )
                }
            )

        with patch(
            "scripts.extraction_case_runner.classify_content",
            side_effect=leave_one_narrative_unresolved,
        ):
            semantic_failure = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["SEM-003"],
                oracle_registry=registry,
            )
        semantic_failure_assertions = {
            item["assertion_id"]: item
            for item in semantic_failure["cases"]["SEM-003"]["assertions"]
        }
        self.assertEqual(
            semantic_failure_assertions["SEM-003-A2"]["status"],
            "FAIL",
        )
        self.assertEqual(
            semantic_failure_assertions["SEM-003-A2"]["reason_code"],
            "technical_narrative_disposition_mismatch",
        )
        self.assertEqual(
            semantic_failure_assertions["SEM-003-A2"]["evidence"][
                "invalid_technical_narrative_count"
            ],
            1,
        )
        for assertion_id in ("SEM-003-A1", "SEM-003-A3", "SEM-003-A4"):
            self.assertEqual(
                semantic_failure_assertions[assertion_id]["status"],
                "PASS",
            )

        wrong_role_manifest = copy.deepcopy(manifest)
        wrong_role_document = wrong_role_manifest["documents"][
            "DOC-IRPS-PRODUCT-INFO"
        ]
        wrong_role_document["document_role"] = "customer_requirement_source"
        wrong_role_document["ingress_route"] = "requirement_source"
        role_failure = run_manifest(
            wrong_role_manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["SEM-003"],
            oracle_registry=registry,
        )
        role_failure_assertions = {
            item["assertion_id"]: item
            for item in role_failure["cases"]["SEM-003"]["assertions"]
        }
        for assertion_id in ("SEM-003-A1", "SEM-003-A2", "SEM-003-A3"):
            self.assertEqual(
                role_failure_assertions[assertion_id]["status"],
                "PASS",
            )
        self.assertEqual(
            role_failure_assertions["SEM-003-A4"]["status"],
            "FAIL",
        )
        self.assertEqual(
            role_failure_assertions["SEM-003-A4"]["reason_code"],
            "source_role_semantic_exclusion_detected",
        )
        self.assertEqual(
            role_failure_assertions["SEM-003-A4"]["evidence"]["document_role"],
            "customer_requirement_source",
        )
        self.assertEqual(
            role_failure_assertions["SEM-003-A4"]["evidence"]["ingress_route"],
            "requirement_source",
        )

        serialized = json.dumps(
            [passed, semantic_failure, role_failure],
            ensure_ascii=False,
        )
        self.assertNotIn(secret_marker, serialized)
        self.assertNotIn(private_parameter, serialized)
        for narrative_text in narrative_texts:
            self.assertNotIn(narrative_text, serialized)
        self.assertNotIn(str(path), serialized)

    def test_merged_applicability_evaluator_checks_exact_relation_endpoints(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-IRPS-TRL-B"]
        path = self.root / "merged-scope.xlsx"
        secret_text = "private shared acceptance criterion"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Scope"
        worksheet.append(
            [
                "Material",
                "Description",
                "Category",
                "Standard",
                "Acceptance criteria",
                "Material requirement",
            ]
        )
        worksheet.append(
            ["M1", "D1", "C1", "Shared", secret_text, "Constraint"]
        )
        worksheet.append(["M2", "D2", "C2", None, None, None])
        worksheet.append(["M3", "D3", "C3", None, None, None])
        worksheet.merge_cells("D2:D4")
        worksheet.merge_cells("E2:E4")
        worksheet.merge_cells("F2:F4")
        workbook.save(path)
        workbook.close()
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_tab2_oracle(manifest, path)
        real_structure_builder = runner_module.build_structure_tree

        passed = run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["TAB-002"],
            oracle_registry=registry,
        )
        assertions = {
            item["assertion_id"]: item
            for item in passed["cases"]["TAB-002"]["assertions"]
        }
        self.assertEqual(assertions["TAB-002-A1"]["status"], "PASS")
        self.assertEqual(assertions["TAB-002-A2"]["status"], "PASS")
        self.assertEqual(assertions["TAB-002-A3"]["status"], "PASS")
        self.assertEqual(assertions["TAB-002-A4"]["status"], "PASS")
        self.assertFalse(
            assertions["TAB-002-A4"]["evidence"]["is_human_gold"]
        )
        self.assertFalse(
            assertions["TAB-002-A4"]["evidence"][
                "is_serving_activation_gate"
            ]
        )
        self.assertEqual(
            assertions["TAB-002-A2"]["evidence"]["expected_applies_to_count"],
            6,
        )
        self.assertEqual(
            assertions["TAB-002-A2"]["evidence"]["observed_applies_to_count"],
            6,
        )

        real_projector = runner_module.project_requirements

        def drop_one_material_projection(
            source_snapshot: object,
            tree: object,
            semantics: object,
            **kwargs: object,
        ) -> object:
            view = real_projector(
                source_snapshot,
                tree,
                semantics,
                **kwargs,
            )
            return view.model_copy(update={"nodes": view.nodes[:-1]})

        with patch(
            "scripts.extraction_case_runner.project_requirements",
            side_effect=drop_one_material_projection,
        ):
            count_failure_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-002"],
                oracle_registry=registry,
            )
        count_failure_assertion = {
            item["assertion_id"]: item
            for item in count_failure_report["cases"]["TAB-002"]["assertions"]
        }["TAB-002-A3"]
        self.assertEqual(count_failure_assertion["status"], "FAIL")

        def remove_one_applies_to(source_snapshot: object, **_: object) -> object:
            tree = real_structure_builder(source_snapshot)
            removed = False
            relations = []
            for relation in tree.relations:
                if not removed and relation.kind.value == "applies_to":
                    removed = True
                    continue
                relations.append(relation)
            return tree.model_copy(update={"relations": tuple(relations)})

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=remove_one_applies_to,
        ):
            failed = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-002"],
                oracle_registry=registry,
            )
        failed_assertion = {
            item["assertion_id"]: item
            for item in failed["cases"]["TAB-002"]["assertions"]
        }["TAB-002-A2"]
        self.assertEqual(failed_assertion["status"], "FAIL")
        self.assertEqual(
            failed_assertion["reason_code"],
            "merged_cell_applicability_mismatch",
        )

        def retarget_one_applies_to(
            source_snapshot: object,
            **_: object,
        ) -> object:
            tree = real_structure_builder(source_snapshot)
            relations = list(tree.relations)
            relation_index = next(
                index
                for index, relation in enumerate(relations)
                if relation.kind.value == "applies_to"
            )
            relation = relations[relation_index]
            wrong_target = next(
                node.structure_node_id
                for node in tree.nodes
                if node.kind.value == "table"
            )
            relations[relation_index] = relation.model_copy(
                update={
                    "relation_id": stable_structure_relation_id(
                        source_ir_id=tree.source_ir_id,
                        resolver_id=tree.resolver_id,
                        resolver_version=tree.resolver_version,
                        kind=relation.kind,
                        source_structure_node_id=(relation.source_structure_node_id),
                        target_structure_node_id=wrong_target,
                        source_refs=relation.source_refs,
                    ),
                    "target_structure_node_id": wrong_target,
                }
            )
            return tree.model_copy(update={"relations": tuple(relations)})

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=retarget_one_applies_to,
        ):
            wrong_endpoint_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-002"],
                oracle_registry=registry,
            )
        wrong_endpoint_assertion = {
            item["assertion_id"]: item
            for item in wrong_endpoint_report["cases"]["TAB-002"]["assertions"]
        }["TAB-002-A2"]
        self.assertEqual(wrong_endpoint_assertion["status"], "FAIL")
        self.assertEqual(
            wrong_endpoint_assertion["evidence"]["observed_applies_to_count"],
            wrong_endpoint_assertion["evidence"]["expected_applies_to_count"],
        )
        self.assertGreater(
            wrong_endpoint_assertion["evidence"]["invalid_relation_count"],
            0,
        )

        def replace_one_applies_to_source_ref(
            source_snapshot: object,
            **_: object,
        ) -> object:
            tree = real_structure_builder(source_snapshot)
            relations = list(tree.relations)
            relation_index = next(
                index
                for index, relation in enumerate(relations)
                if relation.kind.value == "applies_to"
                and relation.source_refs
            )
            relation = relations[relation_index]
            wrong_cell_id = next(
                cell_id
                for cell_id in tree.source_cell_ids
                if cell_id != relation.source_refs[0].source_id
            )
            wrong_refs = (
                runner_module.StructureSourceRef(
                    kind=runner_module.StructureSourceRefKind.TABLE_CELL,
                    source_id=wrong_cell_id,
                ),
            )
            relations[relation_index] = relation.model_copy(
                update={
                    "relation_id": stable_structure_relation_id(
                        source_ir_id=tree.source_ir_id,
                        resolver_id=tree.resolver_id,
                        resolver_version=tree.resolver_version,
                        kind=relation.kind,
                        source_structure_node_id=(
                            relation.source_structure_node_id
                        ),
                        target_structure_node_id=(
                            relation.target_structure_node_id
                        ),
                        source_refs=wrong_refs,
                    ),
                    "source_refs": wrong_refs,
                }
            )
            return tree.model_copy(update={"relations": tuple(relations)})

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=replace_one_applies_to_source_ref,
        ):
            wrong_ref_report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["TAB-002"],
                oracle_registry=registry,
            )
        wrong_ref_assertion = {
            item["assertion_id"]: item
            for item in wrong_ref_report["cases"]["TAB-002"]["assertions"]
        }["TAB-002-A4"]
        self.assertEqual(wrong_ref_assertion["status"], "FAIL")
        self.assertNotIn(secret_text, json.dumps(passed, ensure_ascii=False))
        self.assertNotIn(secret_text, json.dumps(failed, ensure_ascii=False))
        self.assertNotIn(
            secret_text,
            json.dumps(count_failure_report, ensure_ascii=False),
        )
        self.assertNotIn(
            secret_text,
            json.dumps(wrong_ref_report, ensure_ascii=False),
        )

    def test_structure_resolution_failure_reports_only_safe_error_type(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-LONG-AXLE-PDS-XLSX"]
        path = self.root / "structure-error.xlsx"
        secret_text = "private structure failure requirement"
        values = [secret_text, "2", "3", "4", "5", "6", "7", "8"]
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Scope"
        worksheet.append(values)
        workbook.save(path)
        workbook.close()
        document["document_ref"] = f"external://training-corpus/{path.name}"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_str_oracle(manifest, values)
        private_error = f"do not leak {secret_text} at {path}"

        with patch(
            "scripts.extraction_case_runner.build_structure_tree",
            side_effect=RuntimeError(private_error),
        ):
            report = run_manifest(
                manifest,
                corpus_root=self.root,
                source_ir_mode="native",
                selected_case_ids=["STR-001"],
                oracle_registry=registry,
            )
        assertion = {
            item["assertion_id"]: item
            for item in report["cases"]["STR-001"]["assertions"]
        }["STR-001-A2"]
        self.assertEqual(assertion["status"], "FAIL")
        self.assertEqual(assertion["reason_code"], "structure_resolution_failed")
        self.assertEqual(
            assertion["evidence"]["structure_error_type"],
            "RuntimeError",
        )
        serialized = json.dumps(report, ensure_ascii=False)
        self.assertNotIn(secret_text, serialized)
        self.assertNotIn(private_error, serialized)
        self.assertNotIn(str(path), serialized)

        not_run = run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="off",
            selected_case_ids=["STR-001"],
            oracle_registry=registry,
        )
        not_run_assertion = {
            item["assertion_id"]: item
            for item in not_run["cases"]["STR-001"]["assertions"]
        }["STR-001-A2"]
        self.assertEqual(not_run_assertion["status"], "GAP")
        self.assertEqual(not_run_assertion["reason_code"], "structure_not_run")

    def test_cli_returns_nonzero_when_an_oracle_assertion_fails(self) -> None:
        report = {
            "summary": {
                "identity_gate_pass": True,
                "source_ir_observation_ok": True,
                "assertion_status_counts": {"FAIL": 1, "GAP": 2},
            }
        }
        output_root = self.root.parent / "runner-cli-output"

        with (
            patch(
                "scripts.extraction_case_runner.load_oracle_registry",
                return_value=self.oracles,
            ),
            patch(
                "scripts.extraction_case_runner.load_manifest",
                return_value=self.manifest,
            ),
            patch(
                "scripts.extraction_case_runner.run_manifest",
                return_value=report,
            ),
            patch(
                "scripts.extraction_case_runner.write_reports",
                return_value=(output_root / "run.json", output_root / "summary.md"),
            ),
            patch("builtins.print"),
        ):
            exit_code = runner_module.main(
                [
                    "--corpus-root",
                    str(self.root),
                    "--out-dir",
                    str(output_root),
                ]
            )

        self.assertEqual(exit_code, 4)

    def test_cli_requires_pdf_body_review_paths_as_an_explicit_pair(self) -> None:
        with patch("builtins.print"):
            exit_code = runner_module.main(
                [
                    "--pdf-body-review-artifact-root",
                    str(self.root),
                ]
            )

        self.assertEqual(exit_code, 2)

    def test_case_visual_windows_narrow_only_with_complete_case_scope(self) -> None:
        narrowed = runner_module._selected_pdf_visual_page_scopes(
            manifest=self.manifest,
            selected_case_ids=["STR-006", "STR-007", "SEM-004"],
        )
        self.assertEqual(
            narrowed["DOC-PEF-SOR"],
            frozenset({6, 9, 10, 12}),
        )

        fallback = runner_module._selected_pdf_visual_page_scopes(
            manifest=self.manifest,
            selected_case_ids=["STR-006", "FMT-003"],
        )
        self.assertNotIn("DOC-PEF-SOR", fallback)

    def test_manifest_rejects_unordered_case_visual_page_scope(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        manifest["cases"]["SEM-004"]["pdf_visual_page_numbers_by_document"][
            "DOC-PEF-SOR"
        ] = [10, 9]

        with self.assertRaises(ManifestValidationError):
            validate_manifest(manifest, oracle_registry=self.oracles)

    def test_annotation_policy_pass_requires_detected_and_excluded_annotation_layer(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-PEF-SOR"]
        path = self.root / "annotated.pdf"
        secret_annotation = "private annotation must not enter runner report"
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "visible fixture text")
        page.add_text_annot((90, 90), secret_annotation)
        pdf.save(path)
        pdf.close()
        document["document_ref"] = "external://training-corpus/annotated.pdf"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_source_scope_oracle(
            manifest,
            oracle_id="ORACLE-FMT-003-A2",
            path=path,
            selector={
                "format": "pdf",
                "channel_mode": "NATIVE",
                "page_numbers": [1],
            },
        )

        report = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["FMT-003"],
            oracle_registry=registry,
        )

        case = report["cases"]["FMT-003"]
        assertion_results = {item["assertion_id"]: item for item in case["assertions"]}
        self.assertEqual(assertion_results["FMT-003-A1"]["status"], "PASS")
        self.assertEqual(
            assertion_results["FMT-003-A1"]["reason_code"],
            "annotation_layer_detected_and_excluded",
        )
        self.assertEqual(assertion_results["FMT-003-A2"]["status"], "PASS")
        self.assertEqual(assertion_results["FMT-003-A3"]["status"], "PASS")
        self.assertEqual(assertion_results["FMT-003-A4"]["status"], "PASS")
        self.assertEqual(case["status"], "PASS")
        self.assertIsNone(case["first_gap_layer"])
        self.assertEqual(report["summary"]["assertion_status_counts"], {"PASS": 4})
        self.assertNotIn(secret_annotation, json.dumps(report, ensure_ascii=False))

    def test_annotation_note_scope_fails_closed_when_visible_source_changes(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-PEF-SOR"]
        path = self.root / "annotated-scope.pdf"

        def write_fixture(visible_text: str) -> None:
            pdf = pymupdf.open()
            page = pdf.new_page()
            page.insert_text((72, 72), visible_text)
            page.add_text_annot((90, 90), "excluded annotation")
            pdf.save(path)
            pdf.close()

        write_fixture("reviewed visible note")
        document["document_ref"] = "external://training-corpus/annotated-scope.pdf"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        registry = self._configure_source_scope_oracle(
            manifest,
            oracle_id="ORACLE-FMT-003-A2",
            path=path,
            selector={
                "format": "pdf",
                "channel_mode": "NATIVE",
                "page_numbers": [1],
            },
        )

        write_fixture("changed visible note")
        changed_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        document["source_sha256"] = changed_hash
        baseline_fact = copy.deepcopy(registry["oracles"]["ORACLE-FMT-003-A2"]["fact"])
        registry = self._oracles_for_manifest(manifest)
        registry["oracles"]["ORACLE-FMT-003-A2"]["fact"] = baseline_fact

        report = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["FMT-003"],
            oracle_registry=registry,
        )

        assertion_results = {
            item["assertion_id"]: item
            for item in report["cases"]["FMT-003"]["assertions"]
        }
        self.assertEqual(assertion_results["FMT-003-A1"]["status"], "PASS")
        self.assertEqual(assertion_results["FMT-003-A2"]["status"], "FAIL")
        self.assertEqual(
            assertion_results["FMT-003-A2"]["reason_code"],
            "oracle_fidelity_mismatch",
        )
        self.assertEqual(assertion_results["FMT-003-A3"]["status"], "FAIL")
        self.assertEqual(assertion_results["FMT-003-A4"]["status"], "FAIL")

    def test_hash_deduplication_has_positive_and_negative_assertion_evidence(
        self,
    ) -> None:
        manifest = copy.deepcopy(self.manifest)
        canonical_path = self.root / "canonical.pdf"
        duplicate_path = self.root / "duplicate.pdf"
        unique_path = self.root / "unique.xlsx"
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "canonical fixture")
        pdf.save(canonical_path)
        pdf.close()
        duplicate_path.write_bytes(canonical_path.read_bytes())
        workbook = Workbook()
        workbook.active["A1"] = "unique fixture"
        workbook.save(unique_path)
        workbook.close()

        replacements = {
            "DOC-PEF-SOR": canonical_path,
            "DOC-PEF-SOR-DUPLICATE": duplicate_path,
            "DOC-IRPS-TRL-B": unique_path,
        }
        for document_id, path in replacements.items():
            document = manifest["documents"][document_id]
            document["document_ref"] = f"external://training-corpus/{path.name}"
            document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

        passed = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["ROLE-005"],
        )
        passed_case = passed["cases"]["ROLE-005"]
        passed_assertions = {
            item["assertion_id"]: item for item in passed_case["assertions"]
        }
        self.assertEqual(passed_assertions["ROLE-005-A1"]["status"], "PASS")
        self.assertEqual(passed_assertions["ROLE-005-A2"]["status"], "PASS")
        self.assertEqual(passed_assertions["ROLE-005-A3"]["status"], "PASS")
        self.assertEqual(passed_case["status"], "PASS")
        self.assertIsNone(passed_case["first_gap_layer"])
        self.assertEqual(
            passed_assertions["ROLE-005-A2"]["evidence"],
            {
                "declared_duplicate_pair_count": 1,
                "reused_document_count": 1,
                "unique_source_hash_count": 2,
                "source_ir_unavailable_count": 0,
            },
        )

        duplicate_path.write_bytes(canonical_path.read_bytes() + b"\n")
        manifest["documents"]["DOC-PEF-SOR-DUPLICATE"]["source_sha256"] = (
            hashlib.sha256(duplicate_path.read_bytes()).hexdigest()
        )
        failed = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["ROLE-005"],
        )
        failed_case = failed["cases"]["ROLE-005"]
        failed_assertions = {
            item["assertion_id"]: item for item in failed_case["assertions"]
        }
        self.assertEqual(failed_assertions["ROLE-005-A2"]["status"], "FAIL")
        self.assertEqual(failed_case["status"], "FAIL")
        self.assertEqual(failed_case["first_gap_layer"], "source_ir")

    def test_fmt004_uses_real_attachment_pipeline_for_pdf_and_tif(self) -> None:
        manifest = copy.deepcopy(self.manifest)
        paths = {
            "DOC-DRAWING-45104848": self.root / "drawing-a.pdf",
            "DOC-DRAWING-PEF16C02": self.root / "drawing-b.pdf",
            "DOC-DRAWING-TQ250-EM1": self.root / "drawing-c.tif",
        }
        for path in tuple(paths.values())[:2]:
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with path.open("wb") as handle:
                writer.write(handle)
        PILImage.new("1", (4, 4), color=1).save(
            paths["DOC-DRAWING-TQ250-EM1"],
            format="TIFF",
        )
        original_hashes = {
            document_id: hashlib.sha256(path.read_bytes()).hexdigest()
            for document_id, path in paths.items()
        }
        for document_id, path in paths.items():
            document = manifest["documents"][document_id]
            document["document_ref"] = f"external://training-corpus/{path.name}"
            document["source_sha256"] = original_hashes[document_id]

        report = self._run_manifest(
            manifest,
            corpus_root=self.root,
            source_ir_mode="native",
            selected_case_ids=["FMT-004"],
        )

        case = report["cases"]["FMT-004"]
        self.assertEqual(case["status"], "PASS")
        self.assertEqual(case["assertion_status_counts"], {"PASS": 4})
        self.assertIsNone(case["first_gap_layer"])
        self.assertEqual(report["summary"]["hard_blocker_count"], 0)
        self.assertEqual(
            report["summary"]["preflight_status_counts"],
            {"KNOWN_UNSUPPORTED": 1, "READY": 2},
        )
        self.assertEqual(
            report["summary"]["source_ir_status_counts"],
            {"ATTACHMENT_ONLY": 3},
        )
        for document_id, path in paths.items():
            result = report["documents"][document_id]
            self.assertEqual(result["attachment_status"], "COMPLETE")
            self.assertEqual(
                result["source_ir_counts"],
                {
                    "attachment_artifacts": 1,
                    "attachment_references": 1,
                    "decision_requirements": 0,
                    "visual_text_snapshots": 0,
                },
            )
            self.assertEqual(
                hashlib.sha256(path.read_bytes()).hexdigest(),
                original_hashes[document_id],
            )
        serialized = json.dumps(report, ensure_ascii=False)
        self.assertNotIn(str(self.root), serialized)
        self.assertNotIn("drawing-a.pdf", serialized)
        self.assertNotIn("drawing-c.tif", serialized)

    def test_preview_mode_uses_isolated_snapshot_root_and_never_writes_by_source(
        self,
    ) -> None:
        if "fork" not in multiprocessing.get_all_start_methods():
            self.skipTest("preview worker mock requires fork-compatible test isolation")
        manifest = copy.deepcopy(self.manifest)
        document = manifest["documents"]["DOC-QC-T-1067"]
        path = self.root / "source.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with path.open("wb") as handle:
            writer.write(handle)
        document["document_ref"] = "external://training-corpus/source.pdf"
        document["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        observed_snapshot_root_path = (
            self.root.parent / f"{self.root.name}-snapshot-root.txt"
        )
        self.addCleanup(observed_snapshot_root_path.unlink, missing_ok=True)

        def fake_preview(
            file_path: str, *, document_id: str, snapshot_root: Path
        ) -> dict:
            # macOS exposes the temporary root through both /var and
            # /private/var; compare canonical paths so the isolation check is
            # about the directory, not the spelling of that alias.
            staged_path = Path(file_path).resolve()
            self.assertNotEqual(staged_path, path.resolve())
            self.assertEqual(staged_path.suffix, path.suffix)
            self.assertEqual(staged_path.read_bytes(), path.read_bytes())
            self.assertNotEqual(snapshot_root.parent.resolve(), path.parent.resolve())
            observed_snapshot_root_path.write_text(str(snapshot_root), encoding="utf-8")
            snapshot_root.mkdir(parents=True, exist_ok=True)
            (snapshot_root / "temporary.txt").write_text("temporary", encoding="utf-8")
            return {
                "summary": {
                    "requirements_count": 7,
                    "raw_requirement_rows_count": 9,
                    "non_empty_pages": 1,
                },
                "section_outline": [],
                "complete_requirement_entries": [{"entry_kind": "REQUIREMENT"}] * 7,
                "complete_requirement_entries_raw": [{"entry_kind": "REQUIREMENT"}] * 9,
            }

        fork_context = multiprocessing.get_context("fork")
        with (
            patch(
                "scripts.extraction_case_runner.build_extraction_preview",
                side_effect=fake_preview,
            ),
            patch.object(
                runner_module.multiprocessing,
                "get_context",
                return_value=fork_context,
            ),
        ):
            report = self._run_manifest(
                manifest,
                corpus_root=self.root,
                legacy_mode="preview",
                source_ir_mode="native",
                selected_case_ids=["STR-003"],
            )

        self.assertEqual(
            report["documents"]["DOC-QC-T-1067"]["legacy_status"], "OBSERVED_PREVIEW"
        )
        self.assertEqual(
            report["documents"]["DOC-QC-T-1067"]["source_ir_status"], "PARTIAL"
        )
        self.assertEqual(
            report["documents"]["DOC-QC-T-1067"]["source_ir_counts"]["nodes"],
            0,
        )
        self.assertEqual(
            report["documents"]["DOC-QC-T-1067"]["legacy_counts"]["requirements"],
            7,
        )
        observed_snapshot_root = Path(
            observed_snapshot_root_path.read_text(encoding="utf-8")
        )
        self.assertFalse(observed_snapshot_root.exists())
        self.assertEqual(
            sorted(item.name for item in self.root.iterdir()), ["source.pdf"]
        )


if __name__ == "__main__":
    unittest.main()
